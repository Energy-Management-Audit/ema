"""Render a checked draft over its section's own region of the base."""

# pyright: reportPrivateUsage=false

from __future__ import annotations

import re
from copy import deepcopy
from decimal import Decimal
from functools import partial
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ema.audit.base_anchor import MARKER
from ema.audit.base_units import heading_spans_document
from ema.audit.catalogue import CATALOGUE
from ema.audit.draft_checks import (
    CITE_GAP,
    TOKEN,
    DraftCheck,
    DraftReview,
    check_draft,
    sentence_parts,
    token_only,
)
from ema.audit.draft_schema import DraftText, SectionDraft
from ema.audit.section_body import replace_section_body
from ema.core.errors import EmaError
from ema.core.office.blocks import Block, Missing, Paragraph
from ema.core.office.numbers_ro import format_number
from ema.core.review.models import Field

# Internal codes a reader never sees: the enum value's wording, and units the prose names itself.
ENUM_TEXT = {"below_1000_tep": "sub 1.000 tep", "at_least_1000_tep": "de cel puţin 1.000 tep"}
UNSPOKEN_UNITS = frozenset({"persons"})
UNSUPPORTED = "unsupported"


def rendered_value(field: Field) -> str:
    """A fact as the deliverable prints it: Romanian number format, with its unit."""
    if field.value_type in {"number", "year"}:
        return format_number(
            Decimal(str(field.value)),
            0 if field.value_type == "year" else field.decimals,
            None if field.unit in UNSPOKEN_UNITS else field.unit,
            False if field.value_type == "year" else field.grouping,
        )
    return ENUM_TEXT.get(str(field.value), str(field.value))


def support_text(text: str) -> str:
    return " ".join(text.split()).rstrip(".!?").rstrip()


def _resolved(text: str, facts: dict[str, Field]) -> str:
    """The text as printed: a value for each fact token, nothing for a citation."""

    def replace(match: re.Match[str]) -> str:
        value = rendered_value(facts[match.group(1)])
        suffix = match.group(2)
        return value if suffix and value.endswith((".", "!", "?")) else value + suffix

    return re.sub(TOKEN.pattern + r"(\.?)", replace, CITE_GAP.sub("", text))


def kept(
    item: DraftText, location: str, issues: tuple[DraftReview, ...], facts: dict[str, Field]
) -> list[str]:
    """The item's sentences the render keeps: all but the unsupported ones, which the draft
    review lists (D4); none when a flag covers the whole item."""
    relevant = [
        issue for issue in issues if issue.location == location and issue.code == UNSUPPORTED
    ]
    if any(issue.sentence is None for issue in relevant):
        return []
    flagged = [support_text(issue.sentence or "") for issue in relevant]
    return [
        sentence
        for sentence in sentence_parts(item.text, facts)
        # A verified fact standing alone is the source's own text; a flag never erases it.
        if token_only(sentence)
        or not any(
            flag and (flag in support_text(sentence) or support_text(sentence) in flag)
            for flag in flagged
        )
    ]


def _paragraph(
    item: DraftText, location: str, issues: tuple[DraftReview, ...], facts: dict[str, Field]
) -> Block | None:
    if item.kind == "missing":
        return Missing("body", MARKER)
    sentences = [_resolved(sentence, facts) for sentence in kept(item, location, issues, facts)]
    return Paragraph(item.kind, [" ".join(sentences)]) if sentences else None


def draft_blocks(
    draft: SectionDraft,
    facts: dict[str, Field],
    issues: tuple[DraftReview, ...],
    unit: int | None = None,
) -> list[Block]:
    """The section as blocks, unsupported sentences dropped; with a unit, only that unit's
    text. A missing item, or a section left empty, is the marker."""
    blocks: list[Block] = [
        block
        for index, item in enumerate(draft.paragraphs)
        if unit is None or item.unit == unit
        if (block := _paragraph(item, f"paragraph:{index}", issues, facts)) is not None
    ]
    if draft.status == "missing" or not blocks:
        blocks.append(Missing("body", MARKER))
    return blocks


def _clone_paragraph(prototype: Any, text: str) -> Any:
    cloned = deepcopy(prototype)
    for child in list(cloned):
        if child.tag != qn("w:pPr"):
            cloned.remove(child)
    run = next((item for item in prototype if item.tag == qn("w:r")), None)
    copied_run = deepcopy(run) if run is not None else OxmlElement("w:r")
    for child in list(copied_run):
        if child.tag != qn("w:rPr"):
            copied_run.remove(child)
    value = OxmlElement("w:t")
    value.text = text
    copied_run.append(value)
    cloned.append(copied_run)
    return cloned


def _insert_absent_section(document: Any, section_id: str, spans: list[Any]) -> bool:
    catalogue = {section.id: section for section in CATALOGUE}
    target = catalogue[section_id]
    if target.chapter not in {2, 3} or target.parent is None:
        return False
    if any(item.section_id == section_id for item, _, _ in spans):
        return False
    order = {section.id: index for index, section in enumerate(CATALOGUE)}
    siblings = [
        (item, start, end)
        for item, start, end in spans
        if item.section_id in catalogue and catalogue[item.section_id].parent == target.parent
    ]
    following = next(
        (
            (item, start, end)
            for item, start, end in siblings
            if order[item.section_id] > order[section_id]
        ),
        None,
    )
    preceding = next(
        (
            (item, start, end)
            for item, start, end in reversed(siblings)
            if order[item.section_id] < order[section_id]
        ),
        None,
    )
    sibling = following or preceding
    if sibling is None:
        raise EmaError("draft_prototype", "Secţiunea lipseşte din bază.", section_id)
    body = list(document.element.body)
    _, start, end = sibling
    body_prototype = next(
        (element for element in body[start + 1 : end] if element.tag == qn("w:p")),
        None,
    )
    if body_prototype is None:
        raise EmaError("draft_prototype", "Secţiunea nu are model de paragraf.", section_id)
    anchor_index = following[1] if following else end
    anchor = body[anchor_index]
    anchor.addprevious(_clone_paragraph(body[start], target.title))
    anchor.addprevious(_clone_paragraph(body_prototype, MARKER))
    return True


def render_section(
    base: Path,
    output: Path,
    draft: SectionDraft,
    facts: dict[str, Field],
    flags: tuple[DraftReview, ...],
    *,
    job: str,
) -> DraftCheck:
    """Write the checked draft over the section's own region of the base."""
    if output.resolve() == base.resolve():
        raise ValueError("output must not overwrite audit base")
    check = check_draft(draft, facts, job)
    if check.fatal:
        detail = f"{draft.section}: " + "; ".join(
            f"{issue.code} {issue.detail}" for issue in check.fatal
        )
        if any(issue.code == "passage_verbatim" for issue in check.fatal):
            # A draft stored under an earlier prompt printed its passages whole.
            detail += ". Refaceţi redactarea capitolului (versiunea nouă a promptului)."
        raise EmaError("draft_invalid", "Redactarea nu a trecut verificările.", detail)
    issues = (*check.review, *flags)
    output.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory() as directory:
        source = base
        document = Document(str(base))
        spans = heading_spans_document(document)
        if _insert_absent_section(document, draft.section, spans):
            source = Path(directory) / "with-section.docx"
            document.save(str(source))
        process = draft.section == "ch3.process"
        # #155 D2: a draft over the single pool of unitless passages belongs to unit 1 alone.
        pooled = (
            process
            and bool(draft.paragraphs)
            and all(item.unit in {None, 1} for item in draft.paragraphs)
        )
        by_unit = process and not pooled and any(item.unit is not None for item in draft.paragraphs)
        replace_section_body(
            source,
            output,
            draft.section,
            draft_blocks(draft, facts, issues),
            keep_base=True,
            unit_blocks=partial(draft_blocks, draft, facts, issues) if by_unit else None,
            pooled=pooled,
        )
    return check


def review_payload(
    draft: SectionDraft, check: DraftCheck, flags: tuple[DraftReview, ...]
) -> dict[str, object]:
    return {
        "section": draft.section,
        "coverage": check.coverage,
        "cited_sentences": check.cited_sentences,
        "total_sentences": check.total_sentences,
        "review": [item.__dict__ for item in (*check.review, *flags)],
    }
