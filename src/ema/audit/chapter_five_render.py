"""Replace chapter five in the audit base with confirmed meter and thermal figures."""

# pyright: reportPrivateUsage=false

from __future__ import annotations

import json
from collections.abc import Mapping
from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from lxml import etree

from ema.audit.base_anchor import MARKER
from ema.audit.base_package import package_issues
from ema.audit.base_prototypes import import_formatting
from ema.audit.base_toc import refresh_toc
from ema.audit.base_units import heading_spans_document
from ema.audit.chapter_five import ChapterFivePlan, PlannedPhoto, PlannedReading
from ema.audit.chapter_five_fixed import fixed_elements
from ema.core.config import Settings
from ema.core.office.blocks import (
    Block,
    BulletList,
    Caption,
    ElementLocator,
    Figure,
    Missing,
    Num,
    Paragraph,
    Prototypes,
    Ref,
    RenderReport,
    Retained,
    Segment,
)
from ema.core.office.package import read_parts, xml
from ema.core.office.region import replace_region
from ema.core.resources import resource_path

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _text(element: etree._Element) -> str:
    return "".join(node.text or "" for node in element.iter(W + "t"))


def _positions(base: Path) -> tuple[dict[str, int], list[etree._Element]]:
    spans = heading_spans_document(Document(str(base)))
    positions = {item.section_id: start for item, start, _ in spans}
    if "ch5" not in positions or "ch6" not in positions:
        raise ValueError("audit base lacks a bounded chapter-five region")
    root = xml(read_parts(base), "word/document.xml")
    body = root.find(W + "body")
    if body is None:
        raise ValueError("audit base lacks document body")
    return positions, list(body)


def _model_elements(
    base: Path,
    model: Path,
    prototype: Path,
    prepared: Path,
    identities: tuple[str, ...],
    required: frozenset[str],
) -> tuple[dict[str, etree._Element], dict[str, list[str]], list[str]]:
    target = Document(str(base))
    source = Document(str(model))
    paragraphs = source.paragraphs
    picture = next((item for item in paragraphs if list(item._p.iter(A + "blip"))), None)
    if picture is None:
        raise ValueError("measurement sheet model lacks a picture prototype")
    caption = next((item for item in paragraphs if _text(item._p).strip().startswith("Fig.")), None)
    bullet = next(
        (
            item
            for item in paragraphs
            if item.style is not None
            and (item.style.name or "").lower().startswith("list")
            and len(list(item._p.iter(W + "t"))) == 1
        ),
        None,
    )
    if caption is None or bullet is None:
        raise ValueError("measurement sheet model lacks caption or bullet prototype")
    copies = [deepcopy(item._p) for item in (picture, caption, bullet)]
    import_formatting(target, source, copies)
    fixed_nodes, fixed_keys, issues = fixed_elements(target, prototype, identities, required)
    target.save(str(prepared))
    return (
        {**dict(zip(("picture", "caption", "bullet"), copies, strict=True)), **fixed_nodes},
        fixed_keys,
        issues,
    )


def _prototypes(
    positions: dict[str, int], body: list[etree._Element], source: dict[str, etree._Element]
) -> Prototypes:
    chapter = positions["ch5"]
    following = positions["ch6"]
    body_paragraph = next(
        (
            node
            for node in body[chapter + 1 : following]
            if node.tag == W + "p" and len(list(node.iter(W + "t"))) == 1
        ),
        None,
    )
    if body_paragraph is None:
        raise ValueError("audit base lacks paragraph prototype")
    elements = {"body": body_paragraph, **source}
    for index in range(chapter + 1, following):
        elements[f"retain:{index}"] = body[index]
    return Prototypes(elements, 5, missing_text=MARKER)


def _retain(start: int, end: int, body: list[etree._Element]) -> list[Block]:
    return [
        Retained(f"retain:{index}")
        for index in range(start, end)
        if body[index].tag in {W + "p", W + "tbl"}
    ]


def _fixed(keys: Mapping[str, list[str]], name: str) -> list[Block]:
    return (
        [Retained(key, fresh=True) for key in keys[name]]
        if name in keys
        else [Missing("body", MARKER)]
    )


def _narrative(plan: ChapterFivePlan, key: str) -> Block:
    value = plan.narratives.get(key)
    return Paragraph("body", [value]) if value else Missing("body", MARKER)


def _electric_rules_pass(plan: ChapterFivePlan) -> bool:
    return bool(plan.panels) and all(
        photo.readings
        and all(reading.value is not None for reading in photo.readings)
        and photo.norm is not None
        and photo.narrative_key is None
        for panel in plan.panels
        for photo in panel.photos
    )


def _bullet(reading: PlannedReading) -> list[Segment]:
    parts = reading.key.split(".")
    quantity, phase = parts[-2:]
    label = {
        ("voltage_ll", "l12"): "U 12 (Faza 1 - Faza 2): ",
        ("voltage_ll", "l23"): "U 23 (Faza 2 - Faza 3): ",
        ("voltage_ll", "l31"): "U 31 (Faza 3 - Faza 1): ",
        ("current", "l1"): "I1 (Curentul pe Faza 1): ",
        ("current", "l2"): "I2 (Curentul pe Faza 2): ",
        ("current", "l3"): "I3 (Curentul pe Faza 3): ",
        ("thd_u", "l1"): "V1 Total HD (Faza 1): ",
        ("thd_i", "l1"): "I1 Total HD (Faza 1): ",
    }.get((quantity, phase), f"{quantity} {phase}: ")
    if reading.value is None:
        return [label, Num(None, 0, reading.unit, fact=reading.key), ";"]
    number = Decimal(reading.value)
    decimals = max(0, -int(number.as_tuple().exponent))
    return [label, Num(number, decimals, reading.unit, fact=reading.key), ";"]


def _photo_blocks(
    photo: PlannedPhoto,
    panel_id: str,
    device: str | None,
    images: Mapping[str, Path],
    phrases: dict[str, str],
    plan: ChapterFivePlan,
) -> list[Block]:
    image = images.get(photo.slot) or images.get(photo.sha)
    if image is None:
        raise ValueError(f"chapter-five image missing: {photo.slot}")
    figure_id = f"meter:{panel_id}:{photo.sha[:8]}"
    blocks: list[Block] = [
        Figure(
            "picture",
            Caption(
                "caption", "fig", figure_id, ["Fig. ", Ref("fig", figure_id), " " + photo.caption]
            ),
            image=image,
        ),
    ]
    if device:
        blocks.append(Paragraph("body", [phrases["screen_intro"].format(device=device)]))
    else:
        blocks.append(Missing("body", MARKER))
    if photo.readings:
        blocks.append(BulletList("bullet", [_bullet(reading) for reading in photo.readings]))
    else:
        blocks.append(Missing("body", MARKER))
    if photo.norm:
        blocks.extend(
            Paragraph("body", [phrases[f"norm_{rule}"]])
            for rule in photo.norm.split(", ")
            if f"norm_{rule}" in phrases
        )
    if photo.narrative_key:
        blocks.append(_narrative(plan, photo.narrative_key))
    return blocks


def _blocks(  # noqa: C901, PLR0912
    plan: ChapterFivePlan,
    positions: dict[str, int],
    body: list[etree._Element],
    images: Mapping[str, Path],
    phrases: dict[str, str],
    fixed: Mapping[str, list[str]],
) -> list[Block]:
    blocks: list[Block] = []
    if plan.panels and "ch5.electric" in positions:
        electric = positions["ch5.electric"]
        fisa = positions["ch5.electric_fisa"]
        blocks.extend(_retain(electric, electric + 1, body))
        blocks.extend(_fixed(fixed, "method"))
        blocks.extend(_retain(fisa, fisa + 1, body))
        for panel in plan.panels:
            date = plan.visit_date or MARKER
            blocks.append(
                Paragraph("body", [phrases["fisa_intro"].format(panel=panel.label, date=date)])
            )
            for photo in panel.photos:
                blocks.extend(_photo_blocks(photo, panel.id, panel.device, images, phrases, plan))
        results = positions.get("ch5.electric_rezultate")
        conclusions = positions.get("ch5.electric_concluzii")
        if results is not None:
            blocks.extend(_retain(results, results + 1, body))
            if _electric_rules_pass(plan):
                blocks.append(Paragraph("body", [phrases["results_intro"]]))
                rules = {
                    rule
                    for panel in plan.panels
                    for photo in panel.photos
                    for rule in (photo.norm or "").split(", ")
                }
                blocks.append(
                    BulletList(
                        "bullet",
                        [[phrases[f"results_{rule}"]] for rule in sorted(rules) if rule],
                    )
                )
            else:
                blocks.append(_narrative(plan, "narrative.ch5.electric_rezultate"))
        if conclusions is not None:
            blocks.extend(_retain(conclusions, conclusions + 1, body))
            if _electric_rules_pass(plan):
                blocks.append(Paragraph("body", [phrases["conclusions_first"]]))
            else:
                blocks.append(_narrative(plan, "narrative.ch5.electric_concluzii"))
            blocks.extend(_fixed(fixed, "harmonics"))
    if plan.thermal and "ch5.termic" in positions:
        thermal = positions["ch5.termic"]
        fisa = positions["ch5.termic_fisa"]
        blocks.extend(_retain(thermal, thermal + 1, body))
        blocks.extend(_fixed(fixed, "thermal"))
        blocks.extend(_retain(fisa, fisa + 1, body))
        blocks.append(
            Paragraph(
                "body",
                [
                    phrases["thermal_intro"].format(
                        date=plan.visit_date or MARKER, client=plan.client or MARKER
                    )
                ],
            )
        )
        blocks.append(
            Paragraph(
                "body",
                [
                    phrases["thermal_intro_second"].format(
                        date=plan.visit_date or MARKER, client=plan.client or MARKER
                    )
                ],
            )
        )
        for index, photo in enumerate(plan.thermal):
            image = images.get(photo.slot) or images.get(photo.sha)
            if image is None:
                raise ValueError(f"chapter-five image missing: {photo.slot}")
            identifier = f"thermal:{photo.sha[:8]}"
            letter = chr(ord("a") + index)
            caption = photo.component or MARKER
            blocks.append(
                Figure(
                    "picture",
                    Caption(
                        "caption",
                        "fig",
                        identifier,
                        ["Fig. ", Ref("fig", identifier), f" {letter}) Termografierea {caption}"],
                    ),
                    image=image,
                )
            )
        results = positions.get("ch5.termic_rezultate")
        if results is not None:
            blocks.extend(_retain(results, results + 1, body))
            blocks.append(_narrative(plan, "narrative.ch5.termic_rezultate"))
    return blocks


def render_chapter_five(
    base: Path,
    output: Path,
    plan: ChapterFivePlan,
    images: Mapping[str, Path],
    base_identity: tuple[str, ...],
) -> RenderReport:
    if not base_identity:
        raise ValueError("base identity denylist is required")
    settings = Settings()
    model = settings.audit_measurement_sheet_model
    prototype = settings.audit_measurement_prototype
    if model is None or prototype is None:
        raise ValueError("audit measurement sheet model or prototype is not configured")
    with TemporaryDirectory() as temporary:
        prepared = Path(temporary) / "base.docx"
        required: frozenset[str] = frozenset(
            ({"method", "harmonics"} if plan.panels else set[str]())
            | ({"thermal"} if plan.thermal else set[str]())
        )
        model_elements, fixed, fixed_issues = _model_elements(
            base, model, prototype, prepared, base_identity, required
        )
        positions, body = _positions(prepared)
        with resource_path("audit", "measurement_phrases.json").open(encoding="utf-8") as handle:
            phrases = {item["id"]: item["template"] for item in json.load(handle)}
        report = replace_region(
            prepared,
            output,
            ElementLocator(positions["ch5"] + 1),
            ElementLocator(positions["ch6"] + 1),
            _blocks(plan, positions, body, images, phrases, fixed),
            _prototypes(positions, body, model_elements),
        )
    document = Document(str(output))
    refresh_toc(document)
    document.save(str(output))
    issues = package_issues(output, base_identity)
    if issues:
        output.unlink(missing_ok=True)
        raise ValueError("chapter-five package invalid: " + "; ".join(issues[:8]))
    return replace(report, issues=[*report.issues, *fixed_issues])
