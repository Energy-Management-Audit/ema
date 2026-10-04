"""The 3.1.x process units of a dossier: its flow schemes, else the Fişa's `Flux` blocks (D3).

A process passage belongs to a unit by the source of its quote, never by a stage name: a scheme
file's text is that scheme's unit, a Fişa paragraph is the unit of the `Flux` block it sits in.
Every other source (the permit, a description) is the ch3.flux overview.
"""

from __future__ import annotations

import re
import sqlite3
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from docx import Document

from ema.audit.catalogue_types import AuditFact, fact_key
from ema.core.review.fields import fields
from ema.core.review.models import Evidence
from ema.core.workspace import Workspace
from ema.energy_data.source import normal

ProcessesSource = Literal["schemes", "fisa", "default"]

_SCHEME = re.compile(r"^5\.(\d+)\.")
_DOSSIER_ROWS = (
    "SELECT v.slot, v.file_sha, f.relative_path, f.added_at FROM slots s JOIN slot_versions v "
    "ON v.job_id=s.job_id AND v.slot=s.name AND v.version=s.active_version "
    "JOIN jobs j ON j.id=s.job_id "
    "LEFT JOIN files f ON f.sha=v.file_sha AND f.client_slug=j.client_slug "
    "WHERE s.job_id=? AND s.name LIKE 'dossier/%' ORDER BY s.name"
)


@dataclass(frozen=True)
class ProcessUnits:
    source: ProcessesSource = "default"
    # Per unit in scheme-id order: its scheme files as (dossier name, file version), by name.
    schemes: tuple[tuple[tuple[str, str], ...], ...] = ()
    fisa_sha: str | None = None
    # The Fişa's paragraphs and, per unit, the index of the `Flux` paragraph that opens it.
    paragraphs: tuple[str, ...] = ()
    starts: tuple[int, ...] = ()

    @property
    def count(self) -> int:
        return len(self.schemes) or len(self.starts) or 1

    def block(self, number: int) -> tuple[str, ...]:
        """The paragraphs of the Fişa's unit `number`: its `Flux` paragraph up to the next one."""
        ends = (*self.starts[1:], len(self.paragraphs))
        return self.paragraphs[self.starts[number - 1] : ends[number - 1]]


def flux_starts(paragraphs: Sequence[str]) -> tuple[int, ...]:
    return tuple(
        index for index, text in enumerate(paragraphs) if text.strip().casefold().startswith("flux")
    )


def process_units(
    slots: Iterable[tuple[str, str]], fisa: tuple[str, Sequence[str]] | None = None
) -> ProcessUnits:
    """Flow schemes (checklist item 5) first, then the Fişa's `Flux` paragraphs, else one unit.

    `slots` are the dossier's (name, file version) pairs, `fisa` the Fişa's version and paragraphs.
    """
    schemes: dict[int, list[tuple[str, str]]] = {}
    for name, sha in slots:
        if (match := _SCHEME.match(Path(name).name)) is not None:
            schemes.setdefault(int(match.group(1)), []).append((name, sha))
    if schemes:
        return ProcessUnits(
            "schemes", tuple(tuple(sorted(schemes[key])) for key in sorted(schemes))
        )
    if fisa is not None and (starts := flux_starts(fisa[1])):
        return ProcessUnits("fisa", fisa_sha=fisa[0], paragraphs=tuple(fisa[1]), starts=starts)
    return ProcessUnits()


def unit_of(evidence: Evidence, units: ProcessUnits) -> int | None:
    """The unit, from 1, whose source holds a passage's quote; None for the overview."""
    if units.source == "schemes":
        return next(
            (
                number
                for number, files in enumerate(units.schemes, 1)
                if evidence.file_sha in {sha for _, sha in files}
            ),
            None,
        )
    if units.source != "fisa" or evidence.file_sha != units.fisa_sha or not evidence.quote:
        return None
    text = "\n".join(units.paragraphs)
    at = text.find(evidence.quote)
    if at < 0:
        return None
    paragraph = text.count("\n", 0, at)
    return sum(start <= paragraph for start in units.starts) or None


def dossier_units(ws: Workspace, rows: Sequence[Mapping[str, object]]) -> ProcessUnits:
    """The units of the active dossier slot versions; the Fişa is the last one added."""
    fisas = sorted(
        (
            row
            for row in rows
            if Path(str(row["slot"])).suffix.lower() == ".docx"
            and normal(Path(str(row["slot"])).name).startswith("fisa")
            and row["relative_path"] is not None
        ),
        key=lambda row: float(str(row["added_at"])),
    )
    fisa = None
    if fisas:
        document = Document(str(ws.path(str(fisas[-1]["relative_path"]))))
        fisa = (str(fisas[-1]["file_sha"]), [paragraph.text for paragraph in document.paragraphs])
    slots = ((str(row["slot"]).removeprefix("dossier/"), str(row["file_sha"])) for row in rows)
    return process_units(slots, fisa)


def job_process_units(
    ws: Workspace, job: str, db: sqlite3.Connection | None = None
) -> ProcessUnits:
    if db is None:
        with ws.connect() as connection:
            return job_process_units(ws, job, connection)
    return dossier_units(ws, [dict(row) for row in db.execute(_DOSSIER_ROWS, (job,))])


def passage_units(
    ws: Workspace, job: str, units: ProcessUnits | None = None
) -> dict[str, int | None]:
    """Each found process passage's unit, by the source of its quote: what a draft groups by."""
    units = job_process_units(ws, job) if units is None else units
    passages = [
        item
        for item in fields(ws, job)
        if fact_key(item.key) == AuditFact.PROCESS_SECTIONS and item.presence == "found"
    ]
    with ws.connect() as db:
        evidence = {
            str(row["id"]): Evidence.model_validate_json(row["data"])
            for row in db.execute("SELECT id, data FROM evidence WHERE job_id=?", (job,))
        }
    return {
        item.key: next(
            (
                unit
                for ref in item.evidence
                if ref in evidence and (unit := unit_of(evidence[ref], units)) is not None
            ),
            None,
        )
        for item in passages
    }
