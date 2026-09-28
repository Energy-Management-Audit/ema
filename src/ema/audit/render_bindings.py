"""Write sourced values into the base anchors whose text the anchorer bound to a source."""

# pyright: reportPrivateUsage=false

from __future__ import annotations

import json
from collections.abc import Iterable
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.image.exceptions import UnrecognizedImageError
from docx.image.image import Image
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.text.paragraph import Paragraph

from ema.audit.base_anchor import MARKER, Binding, _paragraph_text, _replace_text
from ema.audit.chapter_four import MONTHS
from ema.core.errors import EmaError
from ema.core.jobs import StageContext
from ema.core.review.models import Field

# The one slot a job's own cover photo is uploaded to, and what its absence reads as.
COVER_SLOT = "cover/photo"
COVER_PHOTO_MISSING = "Fotografia sediului lipseşte"
_PHOTO_TYPES = {"image/jpeg", "image/png"}


def _text(field: Field | None) -> str | None:
    """A text value, unless rejected or blank: whitespace is not an address."""
    if field is None or field.value is None or field.review == "rejected":
        return None
    value = str(field.value)
    return value if value.strip() else None


def binding_values(
    client_name: str, job_fields: Iterable[Field], today: date, cover_photo: Path | None = None
) -> dict[Binding, str | None]:
    """Each binding's value: the job's client, the address, the date, the uploaded photo."""
    address = next((field for field in job_fields if field.key == "audit.address"), None)
    return {
        "client_name": client_name if client_name.strip() else None,
        "address": _text(address),
        "report_month": f"{MONTHS[today.month - 1]} {today.year}",
        "report_year": str(today.year),
        "cover_photo": str(cover_photo) if cover_photo is not None else None,
    }


def checked_photo(path: Path) -> Path:
    """A JPEG or PNG; any other file is refused (an item failure, the marker stays)."""
    try:
        kind = Image.from_file(str(path)).content_type
    except UnrecognizedImageError:
        kind = ""
    if kind not in _PHOTO_TYPES:
        raise EmaError("cover_photo_type", "Fotografia sediului nu este JPEG sau PNG.", kind)
    return path


def cover_photo(ctx: StageContext, client: str) -> Path | None:
    """The job's uploaded cover photo, read so a later upload makes the render stale."""
    slot = next((item for item in ctx.read_slots("cover") if item.slot == COVER_SLOT), None)
    return checked_photo(ctx.ws.file_path(client, slot.file_sha)) if slot is not None else None


def _roots(document: Any) -> list[Any]:
    roots = [document.element]
    for section in document.sections:
        roots.extend((section.header._element, section.footer._element))
    return roots


def _place_photo(document: Any, paragraph: Any, photo: Path, box: list[int]) -> None:
    """Her photo box, now holding the job's photo scaled to fit and never stretched."""
    image = Image.from_file(str(photo))
    scale = min(box[0] / image.width, box[1] / image.height)
    _replace_text(paragraph, MARKER, "")
    run = Paragraph(paragraph, document._body).add_run()
    run.add_picture(str(photo), Emu(round(image.width * scale)), Emu(round(image.height * scale)))


def fill_bindings(docx: Path, anchors: Path, values: dict[Binding, str | None]) -> list[str]:
    """Write each bound anchor's value in place of its marker; return the bindings left open.

    An anchor whose paragraph a later writer replaced (a ch. 4 caption) has no bookmark left
    and is skipped: nothing of the base remains there to fill.
    """
    entries = json.loads(anchors.read_text(encoding="utf-8"))["anchors"]
    bound: dict[str, dict[str, Any]] = {
        f"_ema_{item['slot']}": item for item in entries if item.get("binding")
    }
    document = Document(str(docx))
    left: list[str] = []
    for root in _roots(document):
        for node in root.iter(qn("w:bookmarkStart")):
            entry = bound.get(node.get(qn("w:name"), ""))
            paragraph = node.getparent()
            if entry is None or paragraph is None or MARKER not in _paragraph_text(paragraph):
                continue
            binding: Binding = entry["binding"]
            value = values.get(binding)
            if value is None:
                left.append(binding)
            elif binding == "cover_photo":
                if not entry.get("box"):
                    left.append(binding)
                    continue
                _place_photo(document, paragraph, Path(value), entry["box"])
            else:
                _replace_text(paragraph, MARKER, value)
    document.save(str(docx))
    return list(dict.fromkeys(left))


def cover_labels(anchors: Path) -> dict[str, str]:
    """Bookmark -> the label its marker is reported by, when the section title says too little."""
    entries = json.loads(anchors.read_text(encoding="utf-8"))["anchors"]
    return {
        f"_ema_{item['slot']}": COVER_PHOTO_MISSING
        for item in entries
        if item.get("binding") == "cover_photo"
    }
