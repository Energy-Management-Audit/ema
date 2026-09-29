"""The whole audit rendered into the configured base: a draft at any time, a final when ready."""

from __future__ import annotations

import json
import shutil
import sqlite3
from collections.abc import Iterable
from datetime import date
from functools import partial
from pathlib import Path
from typing import Literal

from docx import Document
from pydantic import BaseModel

from ema.audit.base import build_base
from ema.audit.base_package import package_issues, scrub_package, unreviewed_bullets
from ema.audit.base_toc import refresh_toc
from ema.audit.base_units import heading_spans_document
from ema.audit.catalogue import CATALOGUE
from ema.audit.draft_schema import SECTION_FACTS
from ema.audit.render_bindings import binding_values, cover_labels, cover_photo, fill_bindings
from ema.audit.render_plan import JobUnitPlan, ProcessesSource, unit_plan
from ema.audit.render_sections import mark_render_drafted, refuse_changed_base
from ema.audit.render_steps import (
    AuditBase,
    ai_wording_hits,
    body_counts,
    configured_base,
    drop_na_sections,
    droppable,
    sentence,
    snapshot_base,
    toc_pages,
)
from ema.audit.render_writers import (
    Chain,
    RenderFailure,
    drafted_sections,
    ready_runs,
    text_of,
    write_draft,
    write_five,
    write_four,
    write_intro,
    write_six,
)
from ema.audit.sections import refresh_staleness
from ema.audit.staleness import current_inputs
from ema.core.config import Settings, load_settings
from ema.core.errors import EmaError
from ema.core.jobs import StageContext, StageOutcome, get_job, run_stage
from ema.core.jobs.reads import revision
from ema.core.office.word_api import word_automation, word_available
from ema.core.review.models import Field
from ema.core.review.section_transition import SectionState, Status
from ema.core.workspace import Workspace

Kind = Literal["draft", "final"]
STAGES: dict[Kind, str] = {"draft": "audit_render", "final": "audit_final"}
NAMES: dict[Kind, str] = {"draft": "Audit-ciorna", "final": "Audit-final"}
INPUTS = "render-inputs.json"
_CHAPTERS = {section.id: section for section in CATALOGUE if section.parent is None}


class RenderChapter(BaseModel):
    number: int
    title: str
    section_id: str
    page: int | None = None


class RenderMarker(BaseModel):
    section_id: str
    label: str


class RenderUnitPlan(BaseModel):
    client_name: str
    processes: int
    processes_source: ProcessesSource
    carriers: list[str]
    measured_panels: int
    thermal_measurements: bool
    equipment_tables: int
    measures: int


class RenderSummary(BaseModel):
    kind: Kind
    chapters: list[RenderChapter]
    tables: int
    charts: int
    markers: list[RenderMarker]
    fields_total: int
    fields_confirmed: int
    fields_manual: int
    unit_plan: RenderUnitPlan
    pdf: bool
    toc_pages_set: bool
    dropped: list[str]
    failures: list[RenderFailure]
    charts_skipped: list[str] = []


def _plan_view(plan: JobUnitPlan) -> RenderUnitPlan:
    return RenderUnitPlan(
        client_name=plan.client_name,
        processes=plan.processes,
        processes_source=plan.processes_source,
        carriers=sorted(plan.carriers),
        measured_panels=plan.measured_panels,
        thermal_measurements=plan.thermal_measurements,
        equipment_tables=plan.equipment_tables,
        measures=plan.measures,
    )


# The field families the render iterates (the reviewed ch. 4 dataset, the unit plan's carriers,
# the texts and the AI gate) and the single keys it looks up (the unit plan, ch. 2-3 facts).
ITERATED = ("carrier.", "carrier_tep.", "production.", "turnover.", "energy_costs.", "narrative.")
LOOKED_UP = ("audit.company_name", "audit.address", "audit_measure.count")


def _read_fields(
    ctx: StageContext, keys: Iterable[str], db: sqlite3.Connection | None = None
) -> list[Field]:
    """The fields the render uses, with their membership bound: a key looked up is recorded even
    when absent, and each iterated family by the (key, revision) set it holds, so an insertion,
    a deletion or a decision in any of them makes the render stale."""
    looked_up = set(keys)
    if db is None:
        with ctx.ws.connect() as connection:
            connection.execute("BEGIN")
            return _read_fields(ctx, looked_up, connection)
    rows = db.execute(
        "SELECT key,data FROM fields WHERE job_id=? ORDER BY key", (ctx.job,)
    ).fetchall()
    reads = [("fields.key", f"{ctx.job}:{key}") for key in sorted(looked_up)]
    reads += [("fields.prefix", f"{ctx.job}:{prefix}") for prefix in ITERATED]
    recorded = [(table, row_id, revision(db, table, row_id) or 0) for table, row_id in reads]
    for table, row_id, value in recorded:
        ctx.record_read(table, row_id, value)
    return [
        Field.model_validate_json(row["data"])
        for row in rows
        if str(row["key"]) in looked_up or str(row["key"]).startswith(ITERATED)
    ]


def _record_sections(ctx: StageContext, db: sqlite3.Connection | None = None) -> None:
    """Every section's decision is an input: a later done -> n/a makes this render stale."""
    if db is None:
        with ctx.ws.connect() as connection:
            return _record_sections(ctx, connection)
    rows = db.execute(
        "SELECT section_id,revision FROM section_states WHERE job_id=?", (ctx.job,)
    ).fetchall()
    revisions = {str(row["section_id"]): int(row["revision"]) for row in rows}
    for section in CATALOGUE:
        ctx.record_read("section_states", f"{ctx.job}:{section.id}", revisions.get(section.id, 0))


def _statuses(ws: Workspace, job: str) -> dict[str, Status]:
    with ws.connect() as db:
        rows = db.execute(
            "SELECT section_id,data FROM section_states WHERE job_id=?", (job,)
        ).fetchall()
    return {str(row["section_id"]): SectionState.parse(row["data"]).status for row in rows}


def _chapters(docx: Path) -> list[str]:
    spans = heading_spans_document(Document(str(docx)))
    return [item.section_id for item, _, _ in spans if item.section_id in _CHAPTERS]


def _printed(chapters: list[str]) -> dict[str, int]:
    # As refresh_toc numbers them: without ch. 5 the later chapters move up by one.
    measured = "ch5" in chapters
    return {
        section_id: chapter if measured or chapter < 6 else chapter - 1
        for section_id in chapters
        if (chapter := _CHAPTERS[section_id].chapter)
    }


def render_audit(
    ctx: StageContext, kind: Kind, base: AuditBase, settings: Settings
) -> StageOutcome:
    """Render from one snapshot of the configured base, taken when the run starts."""
    folder = ctx.artifact_dir()
    base = snapshot_base(base, folder)
    try:
        return _render(ctx, kind, base, settings, folder)
    finally:
        base.document.unlink(missing_ok=True)


def _render(  # noqa: C901, PLR0912, PLR0915
    ctx: StageContext, kind: Kind, base: AuditBase, settings: Settings, folder: Path
) -> StageOutcome:
    ws, job = ctx.ws, ctx.job
    base_sha = base.inputs["audit_base_document"]
    if kind == "final":
        refuse_changed_base(ctx, base_sha)
    client = str(get_job(ws, job)["client_slug"])
    dossier = ctx.read_slots("dossier")
    visit = ctx.read_slots("visit")
    drafted = drafted_sections(ws, job, ctx=ctx)
    facts = (key for section in drafted for key in SECTION_FACTS.get(section, ()))
    with ws.connect() as db:
        db.execute("BEGIN IMMEDIATE")
        refresh_staleness(ws, job, db, base_sha=base_sha)
        job_fields = _read_fields(ctx, (*LOOKED_UP, *facts), db)
        _record_sections(ctx, db)
        plan = unit_plan(ws, job, db=db, ctx=ctx)
        composition_facts, _ = current_inputs(db, job)
        ctx.record_read("audit_materials", job, revision(db, "audit_materials", job) or 0)

    ctx.record_input(template=base_sha)
    (folder / INPUTS).write_text(json.dumps(base.inputs, indent=2), encoding="utf-8")
    start = build_base(
        plan,
        base_document=base.document,
        measurement_prototype=base.prototype,
        output=folder / "base.docx",
        base_identity=base.identity,
    )
    by_key = {field.key: field for field in job_fields}
    chain = Chain(folder, start)
    charts_skipped: list[str] = []
    try:
        photo = cover_photo(ctx, client)
    except EmaError as exc:  # a wrong file keeps the marker; the rest renders (R21)
        chain.failures.append(RenderFailure(section_id="front", code=exc.code))
        photo = None
    anchors = start.with_suffix(".anchors.json")
    fill_bindings(start, anchors, binding_values(plan.client_name, job_fields, date.today(), photo))
    chapters = _chapters(start)
    failed_intros: set[str] = set()
    for index, chapter_id in enumerate(chapters, 1):
        title = sentence(_CHAPTERS[chapter_id].title)
        ctx.progress(index - 1, len(chapters), f"Capitolul {index} din {len(chapters)} — {title}")
        if ctx.cancelled():
            return StageOutcome()
        if chapter_id in {"ch3", "ch6"} and not chain.apply(
            chapter_id,
            partial(
                write_intro,
                section_id=chapter_id,
                text=text_of(by_key, f"narrative.{chapter_id}"),
            ),
        ):
            failed_intros.add(chapter_id)
        if chapter_id in {"ch2", "ch3"}:
            for section in CATALOGUE:
                if (
                    section.parent is not None
                    and section.id in drafted
                    and (section.chapter == _CHAPTERS[chapter_id].chapter)
                ):
                    chain.apply(
                        section.id,
                        partial(
                            write_draft,
                            ws=ws,
                            job=job,
                            section=section.id,
                            by_key=by_key,
                            ctx=ctx,
                        ),
                    )
        elif chapter_id == "ch4":
            chain.apply(
                "ch4",
                partial(
                    write_four,
                    ws=ws,
                    client=client,
                    dossier=dossier,
                    job_fields=job_fields,
                    identity=base.identity,
                    chart_source=base.document,
                    client_name=plan.client_name,
                    charts_skipped=charts_skipped,
                ),
            )
        elif chapter_id == "ch5" and (runs := ready_runs(ws, job, "measurements", ctx=ctx)):
            images = {slot.slot: ws.file_path(client, slot.file_sha) for slot in visit}
            chain.apply("ch5", partial(write_five, runs[0], images, base.identity))
        elif chapter_id == "ch6" and (runs := ready_runs(ws, job, "measures", ctx=ctx)):
            chain.apply("ch6", partial(write_six, runs[0], base.identity))
    ctx.progress(len(chapters), len(chapters), f"Capitolul {len(chapters)} din {len(chapters)}")
    name = NAMES[kind]
    docx = folder / f"{name}.docx"
    shutil.copyfile(chain.current, docx)
    statuses = _statuses(ws, job)
    not_applicable = [key for key, value in statuses.items() if value == Status.NA]
    dropped = drop_na_sections(docx, droppable(not_applicable))
    document = Document(str(docx))
    refresh_toc(document)
    document.save(str(docx))
    # Dropping sections leaves their pictures and charts related but unused: none may ship.
    scrub_package(docx)
    issues = package_issues(docx, base.identity)
    if kind == "final":
        issues += unreviewed_bullets(docx)
    if issues:
        raise EmaError(
            "audit_package", "Pachetul Word al auditului este invalid.", "; ".join(issues[:3])
        )
    if kind == "final" and (hits := ai_wording_hits(docx, job_fields)):
        raise EmaError(
            "audit_ai_wording", "Raportul final conţine formulări despre AI.", "; ".join(hits)
        )
    counts = body_counts(docx, cover_labels(anchors))
    if kind == "final" and counts.markers:
        sections = dict.fromkeys(section_id for section_id, _ in counts.markers)
        raise EmaError(
            "audit_markers", "Raportul final are câmpuri necompletate.", "; ".join(sections)
        )
    pdf = folder / f"{name}.pdf"
    pages: dict[int, int] = {}
    if word_available(settings):
        office = word_automation(settings)
        office.update_toc_pages(docx)
        office.render_pdf(docx, pdf)
        office.open_check(docx)
        pages = toc_pages(docx)
    elif kind == "final":
        raise EmaError("word_unavailable", "Microsoft Word nu este disponibil.", "")
    final_chapters = _chapters(docx)
    printed = _printed(final_chapters)
    found = [field for field in job_fields if field.presence == "found"]
    summary = RenderSummary(
        kind=kind,
        chapters=[
            RenderChapter(
                number=printed[section_id],
                title=sentence(_CHAPTERS[section_id].title),
                section_id=section_id,
                page=pages.get(printed[section_id]),
            )
            for section_id in final_chapters
        ],
        tables=counts.tables,
        charts=counts.charts,
        markers=[RenderMarker(section_id=key, label=label) for key, label in counts.markers],
        fields_total=len(found),
        fields_confirmed=sum(
            field.review in {"accepted", "corrected"} or field.state == "manual" for field in found
        ),
        fields_manual=sum(field.state == "manual" for field in found),
        unit_plan=_plan_view(plan),
        pdf=pdf.is_file(),
        toc_pages_set=bool(pages),
        dropped=dropped,
        failures=chain.failures,
        charts_skipped=charts_skipped,
    )
    (folder / "render.json").write_text(summary.model_dump_json(indent=2), encoding="utf-8")
    if summary.pdf:
        ctx.save_output(pdf, pdf.name)
    ctx.save_output(docx, docx.name, kind=kind)
    if kind == "draft":
        failed = failed_intros | {
            item.section_id for item in chain.failures if item.section_id == "ch4"
        }
        mark_render_drafted(
            ctx,
            docx,
            base_sha,
            failed,
            facts=composition_facts,
            client_revision=plan.client_revision,
        )
    return StageOutcome(item_failures=[item.section_id for item in chain.failures])


def start_audit_render(
    ws: Workspace, job: str, kind: Kind, *, on_revision: int | None = None
) -> str:
    """Start a draft or final render; refusals that need no document come before the run."""
    if get_job(ws, job)["type"] != "audit":
        raise EmaError("wrong_job_type", "Lucrarea nu este un audit.", job)
    settings = load_settings(ws)
    base = configured_base(settings)
    if kind == "final" and not word_available(settings):
        raise EmaError("word_unavailable", "Microsoft Word nu este disponibil.", "")
    return run_stage(
        ws,
        job,
        STAGES[kind],
        lambda ctx: render_audit(ctx, kind, base, settings),
        on_revision=on_revision,
    )
