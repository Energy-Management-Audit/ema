"""Dossier files as plain text for the intake and Fill agents."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from ema.audit.fill_tools import FillDocument
from ema.core.office.convert import stored_file
from ema.core.office.sheets import open_book
from ema.core.office.sniff import FileKind, sniff
from ema.core.pdf import text as pdf_text
from ema.core.workspace import Workspace
from ema.core.workspace.conversion import active_version


def read_document(name: str, path: Path, file_sha: str | None = None) -> FillDocument:
    kind = sniff(path).kind
    if kind in {FileKind.TEXT, FileKind.HTML}:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        return FillDocument(name, text, file_sha=file_sha)
    if kind == FileKind.DOCX:
        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return FillDocument(name, "\n".join([*paragraphs, *cells]), file_sha=file_sha)
    if kind == FileKind.PDF:
        pages = pdf_text(path)
        return FillDocument(
            name,
            "\n".join(page.text for page in pages),
            page_images=tuple(f"{name}#page={page.page}" for page in pages),
            page_texts=tuple(page.text for page in pages),
            file_sha=file_sha,
        )
    if kind in {FileKind.XLS, FileKind.XLSX}:
        book = open_book(path)
        try:
            values: list[str] = []
            for sheet_name in book.sheet_names:
                sheet = book.sheet(sheet_name)
                for row in range(1, min(sheet.max_row, 200) + 1):
                    for col in range(1, min(sheet.max_col, 25) + 1):
                        value = sheet.value(row, col).value
                        if value is not None:
                            values.append(str(value))
            return FillDocument(name, "\n".join(values), file_sha=file_sha)
        finally:
            book.close()
    return FillDocument(name, "", page_images=(f"{name}#page=1",), file_sha=file_sha)


def dossier_documents(
    ws: Workspace, job: str, slots: dict[str, str] | None = None
) -> dict[str, FillDocument]:
    """Read each named slot's active version; by default every dossier slot by file name."""
    named = (
        slots
        if slots is not None
        else {slot.removeprefix("dossier/"): slot for slot in ws.list_slots(job, "dossier")}
    )
    documents: dict[str, FillDocument] = {}
    for name, slot in named.items():
        # A converted .doc is the slot's active version, so it is read as .docx here.
        version = active_version(ws, job, slot)
        if version is None:
            continue
        _, path = stored_file(ws, job, version.file_sha)
        documents[name] = read_document(name, path, version.file_sha)
    return documents
