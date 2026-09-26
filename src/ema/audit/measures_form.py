"""The auditor's labelled measures form and its located values."""

# pyright: reportPrivateUsage=false

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from ema.core.office.sheets import CellRef, Sheet, open_book
from ema.energy_data.carriers import Carrier, _normalize
from ema.energy_data.source import Located, ReaderIssue, filled, normal, number

HEADERS = (
    "Nr.",
    "Măsura propusă",
    "Efect",
    "Purtător de energie",
    "Economie de energie",
    "UM",
    "Investiție (mii lei)",
    "Economie de cost (mii lei/an)",
    "Notă privind costurile",
)
CARRIER_LABELS = {
    "Energie electrică": Carrier.electricity_grid,
    "Energie electrică fotovoltaică": Carrier.electricity_pv,
    "Gaze naturale": Carrier.natural_gas,
    "Motorină": Carrier.diesel,
    "Benzină": Carrier.petrol,
    "GPL": Carrier.lpg,
    "Păcură": Carrier.fuel_oil,
    "CLU": Carrier.clu,
    "Cărbune": Carrier.coal,
    "Cocs": Carrier.coke,
    "Lemn": Carrier.wood,
    "Biomasă": Carrier.biomass,
    "Coji de floarea soarelui": Carrier.sunflower_husks,
    "Biogaz": Carrier.biogas,
    "CTL": Carrier.ctl,
    "Energie termică": Carrier.purchased_heat,
}
NORMAL_CARRIERS = {_normalize(label): value for label, value in CARRIER_LABELS.items()}
UNITS = ("MWh", "t", "tep")


@dataclass(frozen=True)
class MeasureRow:
    title: Located
    effect: Located | None
    carrier: Located
    carrier_value: Carrier
    saving_amount: Located
    unit: Located
    investment: Located | None
    cost_saving: Located | None
    cost_note: Located | None


@dataclass(frozen=True)
class MeasuresForm:
    rows: tuple[MeasureRow, ...]
    issues: tuple[ReaderIssue, ...]
    header: Located | None


def write_measures_form(dest: Path) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Măsuri propuse"
    sheet.append(HEADERS)
    for col in range(1, len(HEADERS) + 1):
        cell = sheet.cell(1, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="23415C")
    for col, width in enumerate((7, 42, 34, 32, 22, 10, 24, 31, 42), 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = "C2"
    lists = book.create_sheet("Liste")
    for row, label in enumerate(CARRIER_LABELS, 1):
        lists.cell(row, 1, label)
    for row, unit in enumerate(UNITS, 1):
        lists.cell(row, 2, unit)
    lists.sheet_state = "hidden"
    book.defined_names.add(DefinedName("EmaCarriers", attr_text="'Liste'!$A$1:$A$16"))
    book.defined_names.add(DefinedName("EmaUnits", attr_text="'Liste'!$B$1:$B$3"))
    for col, formula in (("D", "=EmaCarriers"), ("F", "=EmaUnits")):
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.error = "Alegeți o valoare din listă."
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(f"{col}2:{col}501")  # pyright: ignore[reportUnknownMemberType]
    dest.parent.mkdir(parents=True, exist_ok=True)
    book.save(dest)
    return dest


def _header(sheet: Sheet) -> tuple[int, dict[str, int]] | None:
    for row in range(1, sheet.max_row + 1):
        for col in range(1, sheet.max_col + 1):
            value = sheet.value(row, col).value
            if not isinstance(value, str) or normal(value) != normal("Măsura propusă"):
                continue
            columns = {
                normal(str(cell.value)): index
                for index in range(1, sheet.max_col + 1)
                if (cell := sheet.value(row, index)).value is not None
            }
            if all(normal(label) in columns for label in HEADERS):
                return row, columns
    return None


def _row(
    sheet: Sheet, index: int, columns: dict[str, int], issues: list[ReaderIssue]
) -> MeasureRow | None:
    cells = [sheet.value(index, columns[normal(label)]) for label in HEADERS]
    values = [filled(cell) for cell in cells]
    if all(value is None for value in values):
        return None
    title = values[1]
    if title is not None and normal(str(title.value)).startswith("total"):
        return None
    carrier = values[3]
    unit = values[5]
    parsed: list[ReaderIssue] = []
    amount = number(cells[4], parsed)
    investment = number(cells[6], parsed)
    cost_saving = number(cells[7], parsed)
    mapped = NORMAL_CARRIERS.get(_normalize(str(carrier.value))) if carrier else None
    if (
        title is None
        or carrier is None
        or mapped is None
        or amount is None
        or unit is None
        or str(unit.value) not in UNITS
        or parsed
    ):
        issues.append(
            ReaderIssue("measure_row_invalid", f"row {index}", CellRef(sheet.name, index, 1))
        )
        return None
    return MeasureRow(
        title, values[2], carrier, mapped, amount, unit, investment, cost_saving, values[8]
    )


def read_measures_form(path: Path) -> MeasuresForm:
    book = open_book(path)
    try:
        for name in book.sheet_names:
            sheet = book.sheet(name)
            found = _header(sheet)
            if found is None:
                continue
            header_row, columns = found
            header = filled(sheet.value(header_row, columns[normal("Măsura propusă")]))
            rows: list[MeasureRow] = []
            issues: list[ReaderIssue] = []
            for index in range(header_row + 1, sheet.max_row + 1):
                if all(
                    filled(sheet.value(index, columns[normal(label)])) is None for label in HEADERS
                ):
                    break
                row = _row(sheet, index, columns, issues)
                if row is not None:
                    rows.append(row)
            return MeasuresForm(tuple(rows), tuple(issues), header)
        return MeasuresForm((), (), None)
    finally:
        book.close()
