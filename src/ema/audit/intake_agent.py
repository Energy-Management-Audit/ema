"""Replay-only agent pass for unnumbered intake documents."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from ema.audit.checklist import ChecklistItem
from ema.audit.intake_tools import IntakeDocument, IntakeTools
from ema.core.jobs import StageContext
from ema.core.llm import (
    AgentContext,
    Limits,
    ReplayProvider,
    agent_state,
    curated_models,
    run_agent,
)
from ema.core.office.convert import stored_file
from ema.core.office.sheets import open_book
from ema.core.office.sniff import FileKind, sniff
from ema.core.pdf import text as pdf_text
from ema.core.workspace.conversion import active_version


def _document(name: str, path: Path) -> IntakeDocument:
    kind = sniff(path).kind
    if kind in {FileKind.TEXT, FileKind.HTML}:
        return IntakeDocument(name, path.read_text(encoding="utf-8-sig", errors="replace"))
    if kind == FileKind.DOCX:
        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return IntakeDocument(name, "\n".join([*paragraphs, *cells]))
    if kind == FileKind.PDF:
        pages = pdf_text(path)
        return IntakeDocument(
            name,
            "\n".join(page.text for page in pages),
            page_images=tuple(f"{name}#page={page.page}" for page in pages),
        )
    if kind in {FileKind.XLS, FileKind.XLSX}:
        book = open_book(path)
        try:
            values: list[str] = []
            for sheet_name in book.sheet_names:
                sheet = book.sheet(sheet_name)
                for row in range(1, min(sheet.max_row, 200) + 1):
                    for col in range(1, min(sheet.max_col, 25) + 1):
                        value = sheet.value(row, col).value
                        if value is not None:
                            values.append(str(value))
            return IntakeDocument(name, "\n".join(values))
        finally:
            book.close()
    return IntakeDocument(name, page_images=(f"{name}#page=1",))


def classify_unplaced(
    ctx: StageContext,
    slots: dict[str, str],
    checklist: tuple[ChecklistItem, ...],
    replay: ReplayProvider,
    limits: Limits,
) -> tuple[IntakeTools, str]:
    documents: dict[str, IntakeDocument] = {}
    for name, slot in slots.items():
        version = active_version(ctx.ws, ctx.job, slot)
        if version is None:
            continue
        _, path = stored_file(ctx.ws, ctx.job, version.file_sha)
        documents[name] = _document(name, path)
    tools = IntakeTools(documents, {item.number: item.text for item in checklist})
    model = next(
        model.id
        for model in curated_models()
        if model.provider == "gemini" and model.tier == "standard"
    )
    context = AgentContext(ctx.ws, ctx.job, "intake", replay, model, "audit-intake-v1")
    previous = agent_state(ctx.ws, ctx.job, "intake")
    if previous is not None:
        tools.restore(previous.messages)
    state = run_agent(context, "Classify each file using verbatim evidence.", tools.tools(), limits)
    return tools, state.status
