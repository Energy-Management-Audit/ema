"""Document steps of the full audit render: base settings, n/a sections, counts, TOC pages."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
from collections.abc import Iterable, Iterator, Mapping
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn

from ema.audit.ai_wording import ai_wording
from ema.audit.base_anchor import MARKER
from ema.audit.base_toc import refresh_toc
from ema.audit.base_units import heading_spans_document
from ema.audit.catalogue import CATALOGUE
from ema.core.config import Settings
from ema.core.errors import EmaError
from ema.core.office.package import read_parts, xml
from ema.core.review.models import Field

_CHART = "{http://schemas.openxmlformats.org/drawingml/2006/chart}chart"
_PAGE_PARTS = re.compile(r"word/(?:header|footer)\d*\.xml")
# The cover, the TOC and anything else before the first heading, as the base anchors name it.
FRONT = "front"
_TITLES = {section.id: section.title for section in CATALOGUE} | {
    FRONT: "Prima pagină şi cuprinsul"
}
_BASE_SETTINGS = (
    "audit_base_document",
    "audit_measurement_prototype",
    "audit_measurement_sheet_model",
    "audit_base_identity",
)


@dataclass(frozen=True)
class AuditBase:
    document: Path
    prototype: Path
    identity: tuple[str, ...]
    # sha256 of every configured base file: a render made from other files is not current.
    inputs: dict[str, str]


def _missing(setting: str) -> EmaError:
    return EmaError("audit_base_missing", "Baza auditului nu este configurată.", setting)


def configured_base(settings: Settings) -> AuditBase:
    """The base files the render needs; the first missing or unreadable one is named."""
    paths: dict[str, Path] = {}
    for name in _BASE_SETTINGS:
        path: Path | None = getattr(settings, name)
        if path is None or not path.is_file():
            raise _missing(name)
        paths[name] = path
    try:
        terms = json.loads(paths["audit_base_identity"].read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise _missing("audit_base_identity") from exc
    if not isinstance(terms, list) or not terms:
        raise _missing("audit_base_identity")
    identity = tuple(str(term) for term in terms if str(term).strip())  # type: ignore[union-attr]
    if not identity:
        raise _missing("audit_base_identity")
    inputs = {name: hashlib.sha256(path.read_bytes()).hexdigest() for name, path in paths.items()}
    return AuditBase(
        paths["audit_base_document"], paths["audit_measurement_prototype"], identity, inputs
    )


def snapshot_base(base: AuditBase, folder: Path) -> AuditBase:
    """The base as this run reads it: one copy, hashed and built from, whatever the configured
    path holds later. The caller removes the copy when the run ends."""
    copy = folder / "audit-base-snapshot.docx"
    shutil.copyfile(base.document, copy)
    sha = hashlib.sha256(copy.read_bytes()).hexdigest()
    return replace(base, document=copy, inputs={**base.inputs, "audit_base_document": sha})


def sentence(title: str) -> str:
    """A title written in capitals reads as a sentence; any other title is kept."""
    title = title.translate(str.maketrans("șțȘȚ", "şţŞŢ"))
    if any(char.islower() for char in title):
        return title
    return title[:1].upper() + title[1:].lower()


def section_ids(docx: Path) -> list[str]:
    return [item.section_id for item, _, _ in heading_spans_document(Document(str(docx)))]


def droppable(not_applicable: Iterable[str]) -> list[str]:
    """N/a sections the document loses: never a fixed one, only when every descendant is n/a."""
    na = set(not_applicable)
    parents = {section.id: section.parent for section in CATALOGUE}

    def ancestors(section_id: str) -> Iterator[str]:
        parent = parents.get(section_id)
        while parent is not None:
            yield parent
            parent = parents.get(parent)

    kept = {
        ancestor
        for section in CATALOGUE
        if section.id not in na
        for ancestor in ancestors(section.id)
    }
    return [
        section.id
        for section in CATALOGUE
        if section.id in na and section.kind != "fixed" and section.id not in kept
    ]


def drop_na_sections(docx: Path, sections: Iterable[str]) -> list[str]:
    """Remove each n/a section, heading and body with its sub-sections, then refresh the TOC."""
    wanted = set(sections)
    document: Any = Document(str(docx))
    spans = [item for item in heading_spans_document(document) if item[0].section_id in wanted]
    kept: list[tuple[str, int, int]] = []
    for item, start, end in sorted(spans, key=lambda span: (span[1], -span[2])):
        if not any(start >= outer_start and end <= outer_end for _, outer_start, outer_end in kept):
            kept.append((item.section_id, start, end))
    if not kept:
        return []
    body = document.element.body
    elements = list(body)
    for _, start, end in reversed(kept):
        for element in elements[start:end]:
            body.remove(element)
    refresh_toc(document)
    document.save(str(docx))
    return [section_id for section_id, _, _ in kept]


@dataclass(frozen=True)
class Counts:
    markers: list[tuple[str, str]]
    tables: int
    charts: int


def _text(element: Any) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def body_counts(docx: Path, labels: Mapping[str, str] | None = None) -> Counts:
    """Markers left in the body with their innermost section, tables and native charts.

    ``labels`` names a marker by the bookmark of its paragraph (the cover photo) instead of
    by its section's title.
    """
    document: Any = Document(str(docx))
    spans = heading_spans_document(document)
    elements = list(document.element.body)
    markers: list[tuple[str, str]] = []
    for index, element in enumerate(elements):
        count = _text(element).count(MARKER)
        if not count:
            continue
        owner = max(
            (span for span in spans if span[1] <= index < span[2]),
            key=lambda span: span[1],
            default=None,
        )
        section_id = owner[0].section_id if owner else FRONT
        named = labels or {}
        for paragraph in element.iter(qn("w:p")):
            names = [str(node.get(qn("w:name"))) for node in paragraph.iter(qn("w:bookmarkStart"))]
            label = next((named[name] for name in names if name in named), None)
            if label is not None and (own := _text(paragraph).count(MARKER)):
                markers.extend([(section_id, label)] * own)
                count -= own
        label = sentence(_TITLES.get(section_id, section_id))
        markers.extend([(section_id, label)] * count)
    # A header or footer repeats on every page; its marker belongs to the cover's review.
    parts = read_parts(docx)
    for name in sorted(part for part in parts if _PAGE_PARTS.fullmatch(part)):
        count = _text(xml(parts, name)).count(MARKER)
        markers.extend([(FRONT, _TITLES[FRONT])] * count)
    tables = sum(element.tag == qn("w:tbl") for element in elements)
    charts = sum(1 for element in elements for _ in element.iter(_CHART))
    return Counts(markers, tables, charts)


def toc_pages(docx: Path) -> dict[int, int]:
    """Printed chapter number -> page, from the TOC1 entries Word has numbered."""
    pages: dict[int, int] = {}
    document: Any = Document(str(docx))
    for element in document.element.body.iter(qn("w:p")):
        style = element.find(f"./{qn('w:pPr')}/{qn('w:pStyle')}")
        if style is None or style.get(qn("w:val")) != "TOC1":
            continue
        match = re.match(r"^\s*(\d+)\..*?(\d+)\s*$", _text(element), flags=re.S)
        if match and int(match.group(2)) > 0:
            pages[int(match.group(1))] = int(match.group(2))
    return pages


_TEXT_PARTS = re.compile(r"word/(?:document|header\d*|footer\d*|footnotes|endnotes)\.xml")


def ai_wording_hits(docx: Path, job_fields: Iterable[Field]) -> list[str]:
    """Where the document mentions AI: the text field that brought it in, else its section.

    Every paragraph of the body, its tables, headers, footers and notes is read, so any text
    path into a final is covered, not only the ones known today.
    """
    texts = [
        field.key
        for field in job_fields
        if isinstance(field.value, str) and field.review != "rejected" and ai_wording(field.value)
    ]
    spans = heading_spans_document(Document(str(docx)))
    hits: list[str] = []
    parts = read_parts(docx)
    for name in sorted(part for part in parts if _TEXT_PARTS.fullmatch(part)):
        root = xml(parts, name)
        body = root.find(qn("w:body"))
        elements = list(body) if body is not None else [root]
        for index, element in enumerate(elements):
            for paragraph in element.iter(qn("w:p")):
                if not ai_wording(_text(paragraph)):
                    continue
                owner = max(
                    (span for span in spans if span[1] <= index < span[2]),
                    key=lambda span: span[1],
                    default=None,
                )
                where = (
                    (owner[0].section_id if owner else FRONT)
                    if body is not None
                    else name.removeprefix("word/").removesuffix(".xml")
                )
                hits.append(where)
    return list(dict.fromkeys([*texts, *hits]))
