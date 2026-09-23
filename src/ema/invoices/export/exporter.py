from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule  # pyright: ignore[reportUnknownVariableType]
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ema.invoices.configuration.field_catalog import (
    ACTIVE_ENERGY,
    ACTIVE_ENERGY_LOSSES,
    ACTIVE_ENERGY_LOSSES_PRICE,
    ACTIVE_ENERGY_PRICE,
    CONSUMPTION_PERIOD,
    GREEN_CERTIFICATES,
    GREEN_CERTIFICATES_PRICE,
    INVOICE_NUMBER,
    LOCATION_IDENTIFIER,
    METER_IDENTIFIER,
    REACTIVE_CAPACITIVE,
    REACTIVE_CAPACITIVE_PRICE,
    REACTIVE_INDUCTIVE,
    REACTIVE_INDUCTIVE_PRICE,
)
from ema.invoices.export.atomic_workbook_writer import (
    save_workbook_atomically,
)
from ema.invoices.export.verifier import (
    OpenpyxlWorkbookVerifier,
)
from ema.invoices.export.workbook_contract import (
    DETAIL_DRAFT_ID_COLUMN,
    DETAIL_HEADERS,
    DETAIL_NORMALIZED_QUANTITY_COLUMN,
    DETAIL_NUMBER_FORMAT,
    DETAIL_SHEET,
    SOURCE_BASIS_CATEGORIES,
    SUMMARY_COLUMNS,
    SUMMARY_PRICE_COLUMNS,
    SUMMARY_PRICE_FORMAT,
    SUMMARY_QUANTITY_COLUMNS,
    SUMMARY_QUANTITY_FORMAT,
    SUMMARY_SHEET,
    SUMMARY_SOURCE_UNIT_COLUMNS,
    SUMMARY_VALUE_COLUMNS,
    SUMMARY_VALUE_FORMAT,
    summary_basis,
    summary_price_formula,
    summary_quantity_formula,
    summary_value_formula,
)
from ema.invoices.models import (
    EnergyCategory,
    FieldStatus,
    InvoiceDraft,
)

_CATEGORY_BY_FIELD = {
    ACTIVE_ENERGY: EnergyCategory.ACTIVE_ENERGY,
    ACTIVE_ENERGY_PRICE: EnergyCategory.ACTIVE_ENERGY,
    ACTIVE_ENERGY_LOSSES: EnergyCategory.ACTIVE_ENERGY_LOSSES,
    ACTIVE_ENERGY_LOSSES_PRICE: EnergyCategory.ACTIVE_ENERGY_LOSSES,
    REACTIVE_CAPACITIVE: EnergyCategory.REACTIVE_CAPACITIVE,
    REACTIVE_CAPACITIVE_PRICE: EnergyCategory.REACTIVE_CAPACITIVE,
    REACTIVE_INDUCTIVE: EnergyCategory.REACTIVE_INDUCTIVE,
    REACTIVE_INDUCTIVE_PRICE: EnergyCategory.REACTIVE_INDUCTIVE,
    GREEN_CERTIFICATES: EnergyCategory.GREEN_CERTIFICATES,
    GREEN_CERTIFICATES_PRICE: EnergyCategory.GREEN_CERTIFICATES,
}

_PRICE_FIELDS = {
    ACTIVE_ENERGY_PRICE,
    ACTIVE_ENERGY_LOSSES_PRICE,
    REACTIVE_CAPACITIVE_PRICE,
    REACTIVE_INDUCTIVE_PRICE,
    GREEN_CERTIFICATES_PRICE,
}


class OpenpyxlWorkbookExporter:
    def __init__(self, verifier: OpenpyxlWorkbookVerifier | None = None) -> None:
        self._verifier = verifier or OpenpyxlWorkbookVerifier()

    def export(self, invoices: list[InvoiceDraft], destination: Path) -> Path:
        if not invoices:
            raise ValueError("Cannot export an empty invoice batch")
        rejected = [invoice.source_filename for invoice in invoices if not invoice.is_exportable]
        if rejected:
            raise ValueError(f"Cannot export invoices requiring review: {', '.join(rejected)}")

        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook: Any = Workbook()
        summary: Any = workbook.active
        if summary is None:
            raise RuntimeError("Openpyxl did not create an active worksheet.")
        summary.title = SUMMARY_SHEET
        details: Any = workbook.create_sheet(DETAIL_SHEET)

        try:
            self._write_details(details, invoices)
            detail_last_row = max(details.max_row, 2)
            self._write_summary(summary, invoices, detail_last_row)
            self._style_summary(summary)
            self._style_details(details)
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        except Exception:
            workbook.close()
            raise
        return save_workbook_atomically(
            workbook,
            destination,
            lambda path: self._verifier.verify(path, invoices),
        )

    def _write_summary(
        self,
        worksheet: Any,
        invoices: list[InvoiceDraft],
        detail_last_row: int,
    ) -> None:
        headers = [header for _, header in SUMMARY_COLUMNS]
        worksheet.append([*headers, "Draft ID"])
        draft_id_column = len(SUMMARY_COLUMNS) + 1
        draft_id_letter = get_column_letter(draft_id_column)
        for row_number, invoice in enumerate(invoices, start=2):
            draft_id_cell = f"${draft_id_letter}{row_number}"
            values: list[object] = []
            for column_number, (field_id, _) in enumerate(SUMMARY_COLUMNS, start=1):
                if column_number in SUMMARY_SOURCE_UNIT_COLUMNS:
                    category = SUMMARY_SOURCE_UNIT_COLUMNS[column_number]
                    values.append(summary_basis(invoice, category).unit)
                elif column_number == 12:
                    values.append("kWh")
                elif column_number in SUMMARY_VALUE_COLUMNS:
                    values.append(
                        summary_value_formula(
                            SUMMARY_VALUE_COLUMNS[column_number],
                            detail_last_row,
                            draft_id_cell,
                        )
                    )
                elif field_id in _CATEGORY_BY_FIELD:
                    values.append(
                        self._summary_value(
                            invoice,
                            field_id,
                            draft_id_cell,
                            detail_last_row,
                        )
                    )
                else:
                    field = invoice.fields.get(field_id) if field_id else None
                    values.append(field.value if field else None)
            worksheet.append([*values, invoice.document_id])
        worksheet.column_dimensions[draft_id_letter].hidden = True

    def _summary_value(
        self,
        invoice: InvoiceDraft,
        field_id: str,
        draft_id_cell: str,
        detail_last_row: int,
    ) -> object:
        field = invoice.fields[field_id]
        if field.status is FieldStatus.MANUALLY_CORRECTED:
            return field.value
        category = _CATEGORY_BY_FIELD[field_id].value
        basis = summary_basis(invoice, category)
        quantity_detail_column = (
            basis.quantity_detail_column
            if category in SOURCE_BASIS_CATEGORIES
            else DETAIL_NORMALIZED_QUANTITY_COLUMN
        )
        if field_id in _PRICE_FIELDS:
            if category in SOURCE_BASIS_CATEGORIES and basis.unique_unit_price is not None:
                return basis.unique_unit_price
            return summary_price_formula(
                category,
                detail_last_row,
                draft_id_cell,
                quantity_detail_column,
            )
        quantity_formula = summary_quantity_formula(
            category,
            detail_last_row,
            draft_id_cell,
            quantity_detail_column,
        )
        return f"={quantity_formula}"

    def _write_details(self, worksheet: Any, invoices: list[InvoiceDraft]) -> None:
        worksheet.append(DETAIL_HEADERS)
        for invoice in invoices:
            invoice_number = invoice.fields[INVOICE_NUMBER].value
            location = invoice.fields[LOCATION_IDENTIFIER].value
            meter = invoice.fields.get(METER_IDENTIFIER)
            consumption_period = invoice.fields.get(CONSUMPTION_PERIOD)
            for detail in invoice.price_details:
                worksheet.append(
                    (
                        invoice.source_filename,
                        invoice.supplier,
                        invoice_number,
                        location,
                        meter.value if meter else None,
                        consumption_period.value if consumption_period else None,
                        detail.category.value,
                        detail.description,
                        detail.source_quantity,
                        detail.source_unit,
                        detail.source_unit_price,
                        detail.normalized_quantity,
                        detail.normalized_unit,
                        detail.normalized_unit_price,
                        detail.net_value,
                        detail.evidence.page_number,
                        detail.evidence.snippet,
                        invoice.document_id,
                    )
                )
        worksheet.column_dimensions[DETAIL_DRAFT_ID_COLUMN].hidden = True

    def _style_summary(self, worksheet: Any) -> None:
        self._style_header(worksheet)
        self._configure_printing(worksheet)
        worksheet.freeze_panes = "A2"
        visible_last_column = get_column_letter(len(SUMMARY_COLUMNS))
        worksheet.auto_filter.ref = f"A1:{visible_last_column}{max(worksheet.max_row, 1)}"
        widths = (
            18,
            14,
            25,
            34,
            20,
            25,
            22,
            15,
            24,
            25,
            19,
            9,
            25,
            24,
            25,
            24,
            25,
            24,
            19,
            25,
            28,
        )
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = width
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row, 1).number_format = "@"
            worksheet.cell(row, 2).number_format = "dd.mm.yyyy"
            for column in (4, 5):
                worksheet.cell(row, column).number_format = "@"
            for column in SUMMARY_QUANTITY_COLUMNS:
                worksheet.cell(row, column).number_format = SUMMARY_QUANTITY_FORMAT
            for column in SUMMARY_PRICE_COLUMNS:
                worksheet.cell(row, column).number_format = SUMMARY_PRICE_FORMAT
            for column in SUMMARY_VALUE_COLUMNS:
                worksheet.cell(row, column).number_format = SUMMARY_VALUE_FORMAT
        self._add_negative_rule(worksheet, f"G2:{visible_last_column}{worksheet.max_row}")

    def _style_details(self, worksheet: Any) -> None:
        self._style_header(worksheet)
        self._configure_printing(worksheet)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:Q{max(worksheet.max_row, 1)}"
        widths = (
            30,
            22,
            18,
            26,
            20,
            25,
            24,
            48,
            18,
            14,
            22,
            20,
            18,
            22,
            18,
            10,
            60,
        )
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = width
        for row in range(2, worksheet.max_row + 1):
            for column in (3, 4, 5):
                worksheet.cell(row, column).number_format = "@"
            for column in (9, 11, 12, 14, 15):
                worksheet.cell(row, column).number_format = DETAIL_NUMBER_FORMAT
        if worksheet.max_row >= 2:
            self._add_negative_rule(worksheet, f"I2:O{worksheet.max_row}")

    @staticmethod
    def _style_header(worksheet: Any) -> None:
        fill = PatternFill("solid", fgColor="1F4E78")
        font = Font(color="FFFFFF", bold=True)
        for cell in worksheet[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[1].height = 46

    @staticmethod
    def _add_negative_rule(worksheet: Any, cell_range: str) -> None:
        worksheet.conditional_formatting.add(
            cell_range,
            CellIsRule(operator="lessThan", formula=["0"], font=Font(color="FF0000")),
        )

    @staticmethod
    def _configure_printing(worksheet: Any) -> None:
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A3
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.print_title_rows = "1:1"
        worksheet.sheet_view.showGridLines = False
