from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from ema.invoices.export.errors import (
    WorkbookContractError,
    WorkbookVerificationCode,
)
from ema.invoices.export.workbook_contract import (
    DETAIL_DRAFT_ID_COLUMN,
    DETAIL_HEADERS,
    DETAIL_NORMALIZED_QUANTITY_COLUMN,
    DETAIL_NUMBER_FORMAT,
    DETAIL_SHEET,
    SOURCE_BASIS_CATEGORIES,
    SUMMARY_COLUMNS,
    SUMMARY_PRICE_COLUMNS,
    SUMMARY_PRICE_FORMAT,
    SUMMARY_QUANTITY_COLUMNS,
    SUMMARY_QUANTITY_FORMAT,
    SUMMARY_SHEET,
    SUMMARY_SOURCE_UNIT_COLUMNS,
    SUMMARY_VALUE_COLUMNS,
    SUMMARY_VALUE_FORMAT,
    summary_basis,
    summary_price_formula,
    summary_quantity_formula,
    summary_value_formula,
)
from ema.invoices.models import FieldStatus, InvoiceDraft


class OpenpyxlWorkbookVerifier:
    """Verify the structural workbook contract without evaluating Excel formulas."""

    def verify(self, path: Path, invoices: list[InvoiceDraft]) -> None:
        try:
            workbook: Any = load_workbook(path, data_only=False)
        except Exception as error:
            raise WorkbookContractError(
                WorkbookVerificationCode.UNREADABLE,
                "The generated workbook cannot be reopened.",
            ) from error

        try:
            self._verify_sheets(workbook.sheetnames)
            summary = workbook[SUMMARY_SHEET]
            details = workbook[DETAIL_SHEET]
            self._verify_headers(summary, details)
            self._verify_row_counts(summary, details, invoices)
            self._verify_links(summary, details, invoices)
            self._verify_formulas(summary, invoices)
            self._verify_number_formats(summary, details)
            self._verify_negative_formatting(summary, details)
        finally:
            workbook.close()

    @staticmethod
    def _verify_sheets(sheetnames: list[str]) -> None:
        if sheetnames != [SUMMARY_SHEET, DETAIL_SHEET]:
            raise WorkbookContractError(
                WorkbookVerificationCode.SHEETS,
                "The workbook sheets do not match the approved contract.",
            )

    @staticmethod
    def _verify_headers(summary: Any, details: Any) -> None:
        summary_headers = tuple(cell.value for cell in summary[1])
        if summary_headers != (*[header for _, header in SUMMARY_COLUMNS], "Draft ID"):
            raise WorkbookContractError(
                WorkbookVerificationCode.SUMMARY_HEADERS,
                "The summary headers do not match the approved contract.",
            )
        detail_headers = tuple(cell.value for cell in details[1])
        if detail_headers != DETAIL_HEADERS:
            raise WorkbookContractError(
                WorkbookVerificationCode.DETAIL_HEADERS,
                "The detail headers do not match the approved contract.",
            )

    @staticmethod
    def _verify_row_counts(summary: Any, details: Any, invoices: list[InvoiceDraft]) -> None:
        if summary.max_row != len(invoices) + 1:
            raise WorkbookContractError(
                WorkbookVerificationCode.SUMMARY_ROW_COUNT,
                "The summary row count does not match the approved export subset.",
            )
        expected_detail_rows = sum(len(invoice.price_details) for invoice in invoices)
        if details.max_row != expected_detail_rows + 1:
            raise WorkbookContractError(
                WorkbookVerificationCode.DETAIL_ROW_COUNT,
                "The detail row count does not match the approved export subset.",
            )

    @staticmethod
    def _verify_links(summary: Any, details: Any, invoices: list[InvoiceDraft]) -> None:
        summary_id_column = len(SUMMARY_COLUMNS) + 1
        summary_id_letter = get_column_letter(summary_id_column)
        if (
            not summary.column_dimensions[summary_id_letter].hidden
            or not details.column_dimensions[DETAIL_DRAFT_ID_COLUMN].hidden
        ):
            raise WorkbookContractError(
                WorkbookVerificationCode.LINK_IDENTIFIERS,
                "The workbook linkage columns are not hidden.",
            )
        expected_summary_ids = [invoice.document_id for invoice in invoices]
        actual_summary_ids = [
            summary.cell(row, summary_id_column).value for row in range(2, summary.max_row + 1)
        ]
        expected_detail_ids = [
            invoice.document_id for invoice in invoices for _detail in invoice.price_details
        ]
        detail_id_column = column_index_from_string(DETAIL_DRAFT_ID_COLUMN)
        actual_detail_ids = [
            details.cell(row, detail_id_column).value for row in range(2, details.max_row + 1)
        ]
        if actual_summary_ids != expected_summary_ids or actual_detail_ids != expected_detail_ids:
            raise WorkbookContractError(
                WorkbookVerificationCode.LINK_IDENTIFIERS,
                "The workbook linkage identifiers do not match the exported drafts.",
            )

    @staticmethod
    def _verify_formulas(summary: Any, invoices: list[InvoiceDraft]) -> None:
        detail_last_row = max(sum(len(invoice.price_details) for invoice in invoices) + 1, 2)
        draft_id_letter = get_column_letter(len(SUMMARY_COLUMNS) + 1)
        for row, invoice in enumerate(invoices, start=2):
            OpenpyxlWorkbookVerifier._verify_quantity_formulas(
                summary, row, invoice, detail_last_row, draft_id_letter
            )
            OpenpyxlWorkbookVerifier._verify_price_formulas(
                summary, row, invoice, detail_last_row, draft_id_letter
            )
            OpenpyxlWorkbookVerifier._verify_unit_formulas(summary, row, invoice)
            OpenpyxlWorkbookVerifier._verify_value_formulas(
                summary, row, detail_last_row, draft_id_letter
            )

    @staticmethod
    def _verify_quantity_formulas(
        summary: Any, row: int, invoice: InvoiceDraft, detail_last_row: int, draft_id_letter: str
    ) -> None:
        for column, category in SUMMARY_QUANTITY_COLUMNS.items():
            field_id = SUMMARY_COLUMNS[column - 1][0]
            assert field_id is not None
            if invoice.fields[field_id].status is FieldStatus.MANUALLY_CORRECTED:
                continue
            basis = summary_basis(invoice, category)
            quantity_detail_column = (
                basis.quantity_detail_column
                if category in SOURCE_BASIS_CATEGORIES
                else DETAIL_NORMALIZED_QUANTITY_COLUMN
            )
            quantity_formula = summary_quantity_formula(
                category,
                detail_last_row,
                f"${draft_id_letter}{row}",
                quantity_detail_column,
            )
            expected = f"={quantity_formula}"
            if summary.cell(row, column).value != expected:
                raise WorkbookContractError(
                    WorkbookVerificationCode.FORMULAS,
                    "A summary quantity formula does not match the approved contract.",
                )

    @staticmethod
    def _verify_price_formulas(
        summary: Any, row: int, invoice: InvoiceDraft, detail_last_row: int, draft_id_letter: str
    ) -> None:
        for column, category in SUMMARY_PRICE_COLUMNS.items():
            field_id = SUMMARY_COLUMNS[column - 1][0]
            assert field_id is not None
            if invoice.fields[field_id].status is FieldStatus.MANUALLY_CORRECTED:
                continue
            basis = summary_basis(invoice, category)
            quantity_detail_column = (
                basis.quantity_detail_column
                if category in SOURCE_BASIS_CATEGORIES
                else DETAIL_NORMALIZED_QUANTITY_COLUMN
            )
            expected: object
            if category in SOURCE_BASIS_CATEGORIES and basis.unique_unit_price is not None:
                expected = basis.unique_unit_price
            else:
                expected = summary_price_formula(
                    category,
                    detail_last_row,
                    f"${draft_id_letter}{row}",
                    quantity_detail_column,
                )
            actual = summary.cell(row, column).value
            if not _cell_values_match(actual, expected):
                raise WorkbookContractError(
                    WorkbookVerificationCode.FORMULAS,
                    "A summary price formula does not match the approved contract "
                    f"at {summary.cell(row, column).coordinate}: "
                    f"expected {expected!r}, found {actual!r}.",
                )

    @staticmethod
    def _verify_unit_formulas(summary: Any, row: int, invoice: InvoiceDraft) -> None:
        for column, category in SUMMARY_SOURCE_UNIT_COLUMNS.items():
            if not _cell_values_match(
                summary.cell(row, column).value,
                summary_basis(invoice, category).unit,
            ):
                raise WorkbookContractError(
                    WorkbookVerificationCode.FORMULAS,
                    "A summary unit does not match the invoice source unit.",
                )

    @staticmethod
    def _verify_value_formulas(
        summary: Any, row: int, detail_last_row: int, draft_id_letter: str
    ) -> None:
        for column, category in SUMMARY_VALUE_COLUMNS.items():
            expected = summary_value_formula(
                category,
                detail_last_row,
                f"${draft_id_letter}{row}",
            )
            if summary.cell(row, column).value != expected:
                raise WorkbookContractError(
                    WorkbookVerificationCode.FORMULAS,
                    "A summary net-value formula does not match the approved contract.",
                )

    @staticmethod
    def _verify_number_formats(summary: Any, details: Any) -> None:
        for row in range(2, summary.max_row + 1):
            if summary.cell(row, 1).number_format != "@" or any(
                summary.cell(row, column).number_format != "@" for column in (4, 5)
            ):
                raise WorkbookContractError(
                    WorkbookVerificationCode.NUMBER_FORMATS,
                    "Invoice and location identifiers must use the Excel text format.",
                )
            if summary.cell(row, 2).number_format != "dd.mm.yyyy":
                raise WorkbookContractError(
                    WorkbookVerificationCode.NUMBER_FORMATS,
                    "The invoice date format does not match the approved contract.",
                )
            if (
                any(
                    summary.cell(row, column).number_format != SUMMARY_QUANTITY_FORMAT
                    for column in SUMMARY_QUANTITY_COLUMNS
                )
                or any(
                    summary.cell(row, column).number_format != SUMMARY_PRICE_FORMAT
                    for column in SUMMARY_PRICE_COLUMNS
                )
                or any(
                    summary.cell(row, column).number_format != SUMMARY_VALUE_FORMAT
                    for column in SUMMARY_VALUE_COLUMNS
                )
            ):
                raise WorkbookContractError(
                    WorkbookVerificationCode.NUMBER_FORMATS,
                    "A summary number format does not match the approved contract.",
                )
        for row in range(2, details.max_row + 1):
            if any(details.cell(row, column).number_format != "@" for column in (3, 4, 5)):
                raise WorkbookContractError(
                    WorkbookVerificationCode.NUMBER_FORMATS,
                    "Detail invoice and location identifiers must use the Excel text format.",
                )
            if any(
                details.cell(row, column).number_format != DETAIL_NUMBER_FORMAT
                for column in (9, 11, 12, 14, 15)
            ):
                raise WorkbookContractError(
                    WorkbookVerificationCode.NUMBER_FORMATS,
                    "A detail number format does not match the approved contract.",
                )

    @staticmethod
    def _verify_negative_formatting(summary: Any, details: Any) -> None:
        summary_ranges = {str(rule.sqref) for rule in summary.conditional_formatting}
        detail_ranges = {str(rule.sqref) for rule in details.conditional_formatting}
        expected_summary = {f"G2:{get_column_letter(len(SUMMARY_COLUMNS))}{summary.max_row}"}
        expected_details: set[str] = {f"I2:O{details.max_row}"} if details.max_row >= 2 else set()
        if summary_ranges != expected_summary or detail_ranges != expected_details:
            raise WorkbookContractError(
                WorkbookVerificationCode.NEGATIVE_FORMATTING,
                "Negative-value formatting does not match the approved contract.",
            )


def _cell_values_match(actual: object, expected: object) -> bool:
    if expected == "" and actual is None:
        return True
    if isinstance(expected, Decimal) and isinstance(actual, int | float):
        actual_decimal = Decimal(str(actual))
        tolerance = max(Decimal("1e-12"), abs(expected) * Decimal("1e-12"))
        return abs(actual_decimal - expected) <= tolerance
    return actual == expected
