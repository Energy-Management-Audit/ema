"""Write the energy-manager workbook with source and exception registers."""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ema.core.office.workbook import save_workbook
from ema.reporting import ReportResult

HEADERS = (
    "Nr. crt.",
    "Beneficiarul contractului de management energetic",
    "Obiectul principal de activitate",
    "Consum anual de energie înregistrat în anul {year} (tep/an)",
    "Măsuri de eficiență energetică realizate",
    "Economii de energie realizate (tep/an)",
    "Costul investiției (mii lei)",
)
CONTROL_HEADERS = (
    "Nr.",
    "Fișier sursă",
    "SHA-256",
    "Format",
    "Beneficiar",
    "CUI",
    "CAEN / activitate",
    "Sursa consumului",
    "Date lunare",
    "Date anuale",
    "Diferență",
    "Valoare utilizată",
    "Structură măsuri",
)


def _complete_total(values: Iterable[float | None]) -> float | None:
    items = list(values)
    return None if any(item is None for item in items) else sum(item or 0 for item in items)


def write_report(result: ReportResult, path: Path, *, firm_name: str | None = None) -> Path:
    book = Workbook()
    active = book.active
    assert active is not None
    book.remove(active)
    for year in result.years:
        sheet = book.create_sheet(str(year))
        sheet.append(
            [f"Raport activitate management energetic — măsuri puse în funcțiune în {year}"]
        )
        sheet.append(
            [value.format(year=result.consumption_year or max(result.years)) for value in HEADERS]
        )
        number = 0
        for company in result.companies:
            measures = company.measures.get(year, ())
            if not measures:
                continue
            number += 1
            for index, measure in enumerate(measures):
                sheet.append(
                    [
                        number if index == 0 else None,
                        company.name if index == 0 else None,
                        company.activity if index == 0 else None,
                        company.consumption_tep if index == 0 else None,
                        measure.description,
                        measure.saving_tep,
                        measure.cost_thousand_lei,
                    ]
                )
            sheet.append([])
        sheet.append(
            [
                "TOTAL",
                None,
                None,
                None
                if any(
                    company.consumption_tep is None
                    for company in result.companies
                    if company.measures.get(year)
                )
                else sum(
                    company.consumption_tep or 0
                    for company in result.companies
                    if company.measures.get(year)
                ),
                "Total măsuri",
                _complete_total(
                    measure.saving_tep
                    for company in result.companies
                    for measure in company.measures.get(year, ())
                ),
                _complete_total(
                    measure.cost_thousand_lei
                    for company in result.companies
                    for measure in company.measures.get(year, ())
                ),
            ]
        )
        _format_year(sheet)
    _control(book, result)
    _exceptions(book, result)
    path.parent.mkdir(parents=True, exist_ok=True)
    save_workbook(book, path, firm_name=firm_name)
    return path


def _format_year(sheet: Worksheet) -> None:
    widths = (9, 38, 35, 22, 64, 18, 18)
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    thin = Side(style="thin", color="B7C6D6")
    for row in sheet:
        for cell in row[:7]:
            row_number = cell.row or 0
            first = row_number == 1
            header = row_number == 2
            company = isinstance(sheet.cell(row_number, 1).value, int)
            total = sheet.cell(row_number, 1).value == "TOTAL"
            cell.alignment = Alignment(
                vertical="center" if first or header else "top", wrap_text=not first and not total
            )
            if first and cell.column == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17365D")
            elif header:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            elif company or total:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D6E4F5" if total else "EAF2F8")
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            elif row_number > 2 and any(item.value is not None for item in row):
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if row_number > 2 and cell.column in (4, 6):
                cell.number_format = "#,##0.000000"
            elif row_number > 2 and cell.column == 7:
                cell.number_format = "#,##0.00"
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 68
    sheet.row_dimensions[sheet.max_row].height = 24


def _control(book: Workbook, result: ReportResult) -> None:
    sheet = book.create_sheet("Control")
    sheet.append([f"Registru de control — {len(result.companies)} surse Anexa 2–3"])
    counts = {
        year: sum(len(company.measures.get(year, ())) for company in result.companies)
        for year in result.years
    }
    sheet.append(
        [
            "Surse: "
            + str(len(result.companies))
            + " | Măsuri: "
            + ", ".join(f"{year}={count}" for year, count in counts.items())
        ]
    )
    sheet.append([])
    sheet.append(
        [*CONTROL_HEADERS, *(f"Măsuri {year}" for year in result.years), "Status", "Observații"]
    )
    for index, company in enumerate(result.companies, 1):
        sheet.append(
            [
                index,
                company.source.name,
                company.sha256,
                company.source.suffix.lower(),
                company.name,
                company.cui,
                company.activity,
                company.consumption_source,
                company.monthly_tep,
                company.annual_tep,
                company.monthly_tep - company.annual_tep
                if company.monthly_tep is not None and company.annual_tep is not None
                else None,
                company.consumption_tep,
                company.measure_sheet,
                *(len(company.measures.get(year, ())) for year in result.years),
                company.status,
                company.observations,
            ]
        )
    widths = (6, 38, 20, 9, 34, 16, 38, 32, 17, 13, 15, 17, 38, *(13 for _ in result.years), 13, 70)
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.row_dimensions[1].height = 30
    for cell in sheet[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in sheet.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column in (9, 10, 11, 12):
                cell.number_format = "#,##0.000000"


def _exceptions(book: Workbook, result: ReportResult) -> None:
    sheet = book.create_sheet("Exceptions")
    sheet.append(["Excepții și decizii aplicate"])
    sheet.append([])
    sheet.append(["Severitate", "Fișier sursă", "Beneficiar", "Situație", "Decizie aplicată"])
    for item in result.exceptions:
        sheet.append(
            [item.severity, item.source.name, item.beneficiary, item.situation, item.decision]
        )
    for col, width in enumerate((14, 42, 34, 74, 48), 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.row_dimensions[1].height = 30
    for cell in sheet[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in sheet.iter_rows(min_row=4):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
