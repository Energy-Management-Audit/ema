"""Deterministic energy-manager reporting from Anexa 2–3."""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, replace
from pathlib import Path
from zipfile import BadZipFile

from openpyxl.utils import get_column_letter

from ema.core.office.errors import OfficeError
from ema.energy_data.anexa import parse_anexa
from ema.energy_data.anexa_cells import Measure
from ema.energy_data.source import Located


@dataclass(frozen=True)
class ReportMeasure:
    description: str
    saving_tep: float | None
    cost_thousand_lei: float | None


@dataclass(frozen=True)
class ReportCompany:
    source: Path
    sha256: str
    name: str
    cui: str | None
    activity: str | None
    consumption_tep: float | None
    consumption_source: str | None
    monthly_tep: float | None
    annual_tep: float | None
    measure_sheet: str | None
    measures: dict[int, tuple[ReportMeasure, ...]]
    status: str
    observations: str | None


@dataclass(frozen=True)
class ReportException:
    severity: str
    source: Path
    beneficiary: str | None
    situation: str
    decision: str
    ref: str | None = None


@dataclass(frozen=True)
class ReportResult:
    years: tuple[int, ...]
    companies: tuple[ReportCompany, ...]
    exceptions: tuple[ReportException, ...]
    consumption_year: int | None


def collect_annexes(sources: list[Path]) -> list[Path]:
    """Expand source folders into deterministically ordered annex workbooks."""
    return sorted(
        (
            item
            for source in sources
            for item in (source.iterdir() if source.is_dir() else (source,))
            if item.suffix.lower() in {".xls", ".xlsx"}
        ),
        key=lambda path: path.name.casefold(),
    )


def _numeric(item: Located | None) -> float | None:
    return float(item.value) if item is not None and isinstance(item.value, int | float) else None


def _text(item: Located | None) -> str | None:
    return str(item.value).strip() if item is not None else None


def _measure(item: Measure) -> ReportMeasure:
    return ReportMeasure(
        " ".join(str(item.description.value).split()),
        _numeric(item.values.get("saving_tep")),
        _numeric(item.values.get("investment_thousand_lei")),
    )


def _measure_structure(item: Measure) -> str:
    cost = item.values.get("investment_thousand_lei")
    saving = item.values.get("saving_tep")
    columns = {
        "descriere": get_column_letter(item.description.ref.col),
        "an": get_column_letter(item.commissioning_year.ref.col)
        if item.commissioning_year
        else "—",
        "cost": get_column_letter(cost.ref.col) if cost else "—",
        "economie": get_column_letter(saving.ref.col) if saving else "—",
    }
    return (
        item.description.ref.sheet
        + ": "
        + ", ".join(f"{label} {col}" for label, col in columns.items())
    )


def generate(  # noqa: C901, PLR0912, PLR0915
    annex_paths: list[Path], years: tuple[int, ...], source_names: dict[Path, Path] | None = None
) -> ReportResult:
    """Parse each source independently; a failed annex is recorded, not fatal."""
    if not years or years != tuple(sorted(set(years))):
        raise ValueError("years must be unique and ascending")
    companies: list[ReportCompany] = []
    exceptions: list[ReportException] = []
    consumption_year: int | None = None
    for input_path in annex_paths:
        path = (source_names or {}).get(input_path, input_path)
        digest = ""
        try:
            digest = hashlib.sha256(input_path.read_bytes()).hexdigest()
            annex = parse_anexa(input_path)
        except (OSError, ValueError, RuntimeError, OfficeError, BadZipFile) as exc:
            companies.append(
                ReportCompany(
                    path,
                    digest,
                    path.stem,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    {},
                    "EROARE",
                    str(exc),
                )
            )
            exceptions.append(
                ReportException("EROARE", path, None, "Fișierul nu poate fi citit.", str(exc))
            )
            continue
        name = _text(annex.identity.get("declared_name") or annex.identity.get("name"))
        if annex.name_origin is not None:
            exceptions.append(
                ReportException(
                    "INFORMARE",
                    path,
                    name,
                    "Denumirea din celula principală este o adresă poștală.",
                    f"S-a folosit {annex.name_origin}.",
                    annex.name_origin,
                )
            )
        if name is None:
            fallback = re.search(r"(?:anul[_ ]?\d{4}|anual)[_ ]+(.+)$", path.stem, re.I)
            name = fallback.group(1).strip() if fallback else path.stem
            if any(issue.code == "name_is_address" for issue in annex.issues):
                exceptions.append(
                    ReportException(
                        "INFORMARE",
                        path,
                        name,
                        "Denumirea din celula principală este o adresă poștală.",
                        "Denumirea a fost preluată din numele fișierului.",
                    )
                )
            else:
                exceptions.append(
                    ReportException(
                        "REVIZUIRE",
                        path,
                        name,
                        "Denumirea beneficiarului lipsește din anexă.",
                        "Denumirea provizorie provine din numele fișierului.",
                    )
                )
        activity = _text(annex.identity.get("caen_description"))
        caen = _text(annex.identity.get("caen_code"))
        if caen and activity:
            if caen.endswith(".0") and caen[:-2].isdigit():
                caen = caen[:-2]
            activity = f"CAEN {caen}: {activity}"
        monthly = _numeric(annex.monthly_total_tep)
        annual = _numeric(annex.annual.get("total_tep"))
        source = annex.monthly_total_tep or annex.annual.get("total_tep")
        source_year = annex.year.value if annex.year is not None else None
        if isinstance(source_year, int):
            if consumption_year is None:
                consumption_year = source_year
            elif source_year != consumption_year:
                exceptions.append(
                    ReportException(
                        "REVIZUIRE",
                        path,
                        name,
                        "Anul raportării diferă de celelalte anexe.",
                        "Sursa a fost exclusă din foile anuale.",
                        annex.year.ref.a1 if annex.year else None,
                    )
                )
        if monthly is None and annual is None:
            exceptions.append(
                ReportException(
                    "REVIZUIRE",
                    path,
                    name,
                    "Consumul anual lipsește din anexă.",
                    "Se completează din sursă.",
                    source.ref.a1 if source else None,
                )
            )
        elif monthly is None:
            exceptions.append(
                ReportException(
                    "INFORMARE",
                    path,
                    name,
                    "Date lunare: totalul în tep/an este indisponibil.",
                    f"S-a folosit {source.ref.a1}." if source else "",
                    source.ref.a1 if source else None,
                )
            )
        elif annual is not None and abs(monthly - annual) > 1e-9:
            exceptions.append(
                ReportException(
                    "REVIZUIRE",
                    path,
                    name,
                    "Totalurile lunar și anual diferă.",
                    "S-a păstrat totalul lunar.",
                    annex.monthly_total_tep.ref.a1 if annex.monthly_total_tep else None,
                )
            )
        grouped: dict[int, list[ReportMeasure]] = {year: [] for year in years}
        for measure in annex.existing_measures:
            year = _numeric(measure.commissioning_year)
            if year is None:
                exceptions.append(
                    ReportException(
                        "REVIZUIRE",
                        path,
                        name,
                        "Anul punerii în funcțiune lipsește.",
                        f"Verificați {measure.description.ref.a1}.",
                        measure.description.ref.a1,
                    )
                )
            elif int(year) in grouped:
                grouped[int(year)].append(_measure(measure))
                for key, message in (
                    ("investment_thousand_lei", "Costul investiției lipsește."),
                    ("saving_tep", "Economia de energie lipsește."),
                ):
                    if _numeric(measure.values.get(key)) is None:
                        exceptions.append(
                            ReportException(
                                "REVIZUIRE",
                                path,
                                name,
                                message,
                                f"Verificați {measure.description.ref.a1}.",
                                measure.description.ref.a1,
                            )
                        )
        if not annex.existing_measures:
            exceptions.append(
                ReportException(
                    "REVIZUIRE",
                    path,
                    name,
                    "Nu există măsuri de eficiență energetică.",
                    "Verificați anexa.",
                )
            )
        if monthly is None and annual is None:
            for year, items in grouped.items():
                if items:
                    exceptions.append(
                        ReportException(
                            "REVIZUIRE",
                            path,
                            name,
                            f"Totalul consumului pentru {year} este necunoscut.",
                            "Totalul din foaia anuală a fost lăsat gol.",
                        )
                    )
        measure_sheet = (
            _measure_structure(annex.existing_measures[0]) if annex.existing_measures else None
        )
        companies.append(
            ReportCompany(
                path,
                digest,
                name,
                _text(annex.identity.get("cui")),
                activity,
                monthly if monthly is not None else annual,
                (source.ref.a1 + " (sursă alternativă; totalul din Date lunare este indisponibil)")
                if source is not None and monthly is None
                else source.ref.a1
                if source
                else None,
                monthly,
                annual,
                measure_sheet,
                {year: tuple(items) for year, items in grouped.items()}
                if source_year is None or source_year == consumption_year
                else {},
                "REVIZUIRE"
                if any(e.source == path and e.severity == "REVIZUIRE" for e in exceptions)
                else "OK",
                None,
            )
        )
    companies.sort(key=lambda company: company.name.casefold())
    cui_groups: dict[str, list[ReportCompany]] = {}
    for company in companies:
        if company.cui:
            cui_groups.setdefault(
                "".join(char for char in company.cui if char.isdigit()), []
            ).append(company)
    duplicate_paths = {
        company.source for group in cui_groups.values() if len(group) > 1 for company in group
    }
    for company in companies:
        if company.source in duplicate_paths:
            exceptions.append(
                ReportException(
                    "REVIZUIRE",
                    company.source,
                    company.name,
                    "CUI apare în mai multe fișiere sursă.",
                    "Punctele de lucru au fost păstrate separat.",
                )
            )
        if (
            company.consumption_tep is not None
            and company.consumption_tep < 1000
            and "peste 1000" in company.source.name
        ):
            exceptions.append(
                ReportException(
                    "REVIZUIRE",
                    company.source,
                    company.name,
                    "Consumul anual este sub 1000 tep.",
                    "S-a păstrat valoarea din anexă.",
                )
            )
    review_paths = {item.source for item in exceptions if item.severity == "REVIZUIRE"}
    review_paths.update(
        company.source
        for company in companies
        if company.monthly_tep is None and company.annual_tep is not None
    )
    companies = [
        replace(company, status="REVIZUIRE") if company.source in review_paths else company
        for company in companies
    ]
    return ReportResult(years, tuple(companies), tuple(exceptions), consumption_year)


from ema.reporting.writer import write_report  # noqa: E402

__all__ = ["ReportResult", "collect_annexes", "generate", "write_report"]
