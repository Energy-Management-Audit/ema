"""Chart formula ranges, caches and embedded workbook construction."""

# pyright: reportPrivateUsage=false

from __future__ import annotations

import io
import re

from lxml import etree
from openpyxl import Workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from ema.core.office.errors import OfficeError
from ema.core.office.package import C

REF = re.compile(r"^(?:'((?:[^']|'')+)'|([^!]+))!\$([A-Z]+)\$(\d+)(?::\$([A-Z]+)\$(\d+))?$")


def cache_values(ref: etree._Element | None) -> list[str | None]:
    if ref is None:
        return []
    cache = ref.find(f"{{{C}}}strCache")
    if cache is None:
        cache = ref.find(f"{{{C}}}numCache")
    if cache is None:
        return []
    count = cache.find(f"{{{C}}}ptCount")
    size = int(count.get("val", "0")) if count is not None else 0
    values: list[str | None] = [None] * size
    for point in cache.findall(f"{{{C}}}pt"):
        index = int(point.get("idx", "0"))
        if index < size:
            values[index] = point.findtext(f"{{{C}}}v")
    return values


def formula_cells(formula: str) -> tuple[str, list[tuple[int, int]]]:
    match = REF.fullmatch(formula)
    if match is None:
        raise OfficeError("chart_formula", f"Unsupported chart formula: {formula}")
    quoted, plain, first_col, first_row, last_col, last_row = match.groups()
    sheet = (quoted or plain or "").replace("''", "'")
    first_c = column_index_from_string(first_col)
    last_c = column_index_from_string(last_col or first_col)
    first_r, last_r = int(first_row), int(last_row or first_row)
    if last_c < first_c or last_r < first_r:
        raise OfficeError("chart_formula", f"Reversed chart formula: {formula}")
    return sheet, [
        (row, col) for row in range(first_r, last_r + 1) for col in range(first_c, last_c + 1)
    ]


def formula_at(formula: str, row: int, col: int, size: int, axis: str) -> str:
    sheet, _ = formula_cells(formula)
    last_row = row + size - 1 if axis == "row" else row
    last_col = col + size - 1 if axis == "column" else col
    label = f"'{sheet.replace(chr(39), chr(39) * 2)}'"
    first = f"${get_column_letter(col)}${row}"
    last = f"${get_column_letter(last_col)}${last_row}"
    return f"{label}!{first}:{last}" if size > 1 else f"{label}!{first}"


def extend_formula(formula: str, size: int, axis: str = "row") -> str:
    _, cells = formula_cells(formula)
    if axis not in ("row", "column"):
        raise OfficeError("chart_formula", f"Unsupported formula axis: {axis}")
    rows = {row for row, _ in cells}
    cols = {col for _, col in cells}
    if len(rows) != 1 and len(cols) != 1:
        raise OfficeError("chart_formula", f"Cannot extend two-dimensional range: {formula}")
    first_row, first_col = cells[0]
    if len(cells) > 1:
        axis = "row" if len(cols) == 1 else "column"
    return formula_at(formula, first_row, first_col, max(size, len(cells)), axis)


def assign_series_formulas(
    fields: list[tuple[str, etree._Element, list[str] | list[float | None]]],
    occupied: dict[tuple[str, int, int], str | float | None],
    next_column: int,
) -> int:
    axis = "row"
    for name, ref, _ in fields:
        formula = ref.findtext(f"{{{C}}}f")
        if name != "tx" and formula:
            _, cells = formula_cells(formula)
            if len(cells) > 1:
                axis = "row" if cells[0][1] == cells[-1][1] else "column"
                break
    proposed: list[
        tuple[str, etree._Element, str, str, list[tuple[int, int]], list[str] | list[float | None]]
    ] = []
    for name, ref, values in fields:
        formula = ref.findtext(f"{{{C}}}f")
        if formula:
            extended = extend_formula(formula, len(values), axis)
            sheet, cells = formula_cells(extended)
            proposed.append((name, ref, extended, sheet, cells, values))
    # Series may intentionally share category cells, while their value cells stay independent.
    conflict = any(
        (sheet, row, col) in occupied and (name != "cat" or occupied[sheet, row, col] != value)
        for name, _, _, sheet, cells, values in proposed
        for (row, col), value in zip(cells, values, strict=True)
    )
    first_col = min((col for _, _, _, _, cells, _ in proposed for _, col in cells), default=0)
    offset = next_column - first_col if conflict else 0
    for _, ref, extended, sheet, cells, values in proposed:
        assigned_formula, assigned_sheet, assigned_cells = extended, sheet, cells
        if offset:
            row, col = cells[0]
            assigned_formula = formula_at(extended, row, col + offset, len(cells), axis)
            assigned_sheet, assigned_cells = formula_cells(assigned_formula)
        formula_node = ref.find(f"{{{C}}}f")
        if formula_node is not None:
            formula_node.text = assigned_formula
        for (row, col), value in zip(assigned_cells, values, strict=True):
            occupied[assigned_sheet, row, col] = value
        next_column = max(next_column, max(col for _, col in assigned_cells) + 1)
    return next_column


def build_workbook(root: etree._Element) -> bytes:
    book = Workbook()
    if book.active is not None:
        book.remove(book.active)
    for ref in root.iter(f"{{{C}}}strRef", f"{{{C}}}numRef"):
        formula = ref.findtext(f"{{{C}}}f")
        if not formula:
            continue
        sheet, cells = formula_cells(formula)
        sheet_obj = book[sheet] if sheet in book.sheetnames else book.create_sheet(sheet)
        values = cache_values(ref)
        if len(cells) != len(values):
            raise OfficeError("chart_formula", f"Cache/range length differs: {formula}")
        numeric = ref.tag == f"{{{C}}}numRef"
        cache = ref.find(f"{{{C}}}numCache")
        number_format = cache.findtext(f"{{{C}}}formatCode") if cache is not None else None
        if number_format is None:
            parent = ref.getparent()
            series = parent.getparent() if parent is not None else None
            format_node = series.find(f"{{{C}}}numFmt") if series is not None else None
            number_format = format_node.get("formatCode") if format_node is not None else None
        for (row, col), value in zip(cells, values, strict=True):
            cell = sheet_obj.cell(
                row, col, float(value) if numeric and value is not None else value
            )
            if numeric and number_format:
                cell.number_format = number_format
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
