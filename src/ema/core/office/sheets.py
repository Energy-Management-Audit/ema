"""Label-based, provenance-carrying reads from legacy Excel workbooks."""

from __future__ import annotations

import unicodedata
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Literal, Protocol, cast

import xlrd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from ema.core.office.errors import OfficeError
from ema.core.office.sheet_formulas import XlsFormulas, formula_inputs_blank
from ema.core.office.sniff import FileKind, sniff

type Scalar = str | float | int | datetime | None


@dataclass(frozen=True)
class CellRef:
    sheet: str
    row: int  # one-based, as displayed by Excel
    col: int

    @property
    def a1(self) -> str:
        return f"{self.sheet}!{get_column_letter(self.col)}{self.row}"


@dataclass(frozen=True)
class CellRange:
    first: CellRef
    last: CellRef

    def __post_init__(self) -> None:
        if (
            self.first.sheet != self.last.sheet
            or self.first.row > self.last.row
            or self.first.col > self.last.col
        ):
            raise ValueError("Invalid cell range")

    @property
    def a1(self) -> str:
        return f"{self.first.a1}:{get_column_letter(self.last.col)}{self.last.row}"


@dataclass(frozen=True)
class CellValue:
    value: Scalar
    ref: CellRef
    hyperlink: str | None = None
    number_format: str | None = None


def _normal(text: str) -> str:
    decomposed = unicodedata.normalize("NFKD", text)
    bare = "".join(char for char in decomposed if not unicodedata.combining(char))
    return " ".join(bare.casefold().split())


class Sheet(Protocol):
    name: str
    max_row: int
    max_col: int

    def value(self, row: int, col: int) -> CellValue: ...

    def formula_inputs_blank(self, row: int, col: int) -> bool: ...

    def find_label(self, labels: Sequence[str], *, within: CellRange | None = None) -> CellRef: ...

    def read_right(self, anchor: CellRef, offset: int = 1) -> CellValue: ...

    def read_below(self, anchor: CellRef, offset: int = 1) -> CellValue: ...

    def read_row(self, anchor: CellRef, count: int, *, start: int = 0) -> list[CellValue]: ...

    def read_block(
        self, anchor: CellRef, rows: int, cols: int, *, down: int = 0, right: int = 0
    ) -> list[list[CellValue]]: ...


class Book(Protocol):
    @property
    def sheet_names(self) -> tuple[str, ...]: ...

    def sheet(self, name: str) -> Sheet: ...

    def close(self) -> None: ...


class _Sheet:
    def __init__(self, name: str, max_row: int, max_col: int) -> None:
        self.name, self.max_row, self.max_col = name, max_row, max_col

    def value(self, row: int, col: int) -> CellValue:
        if row < 1 or col < 1:
            raise ValueError("Excel coordinates are one-based")
        return CellValue(
            self._raw(row, col),
            CellRef(self.name, row, col),
            self._hyperlink(row, col),
            self._number_format(row, col),
        )

    def formula_inputs_blank(self, row: int, col: int) -> bool:
        return False

    def _number_format(self, row: int, col: int) -> str | None:
        return None

    def _hyperlink(self, row: int, col: int) -> str | None:
        return None

    def _raw(self, row: int, col: int) -> Scalar:
        raise NotImplementedError

    def find_label(self, labels: Sequence[str], *, within: CellRange | None = None) -> CellRef:
        if not labels:
            raise ValueError("At least one label is required")
        if within is not None and within.first.sheet != self.name:
            raise ValueError("Search range belongs to a different sheet")
        first_row = within.first.row if within else 1
        last_row = min(within.last.row, self.max_row) if within else self.max_row
        first_col = within.first.col if within else 1
        last_col = min(within.last.col, self.max_col) if within else self.max_col
        matches: dict[str, list[CellRef]] = {_normal(label): [] for label in labels}
        for row in range(first_row, last_row + 1):
            for col in range(first_col, last_col + 1):
                value = self._raw(row, col)
                if isinstance(value, str) and (key := _normal(value)) in matches:
                    matches[key].append(CellRef(self.name, row, col))
        for label in labels:
            found = matches[_normal(label)]
            if len(found) > 1:
                refs = ", ".join(ref.a1 for ref in found)
                raise OfficeError("label_ambiguous", f"{self.name}: {label}: {refs}")
            if found:
                return found[0]
        raise OfficeError("label_missing", f"{self.name}: {', '.join(labels)}: no matches")

    def _check(self, anchor: CellRef) -> None:
        if anchor.sheet != self.name:
            raise ValueError("Anchor belongs to a different sheet")

    def read_right(self, anchor: CellRef, offset: int = 1) -> CellValue:
        self._check(anchor)
        return self.value(anchor.row, self._edge(anchor)[1] + offset)

    def read_below(self, anchor: CellRef, offset: int = 1) -> CellValue:
        self._check(anchor)
        return self.value(self._edge(anchor)[0] + offset, anchor.col)

    def _edge(self, anchor: CellRef) -> tuple[int, int]:
        return anchor.row, anchor.col

    def read_row(self, anchor: CellRef, count: int, *, start: int = 0) -> list[CellValue]:
        self._check(anchor)
        if count < 0:
            raise ValueError("count must be nonnegative")
        return [self.value(anchor.row, anchor.col + start + col) for col in range(count)]

    def read_block(
        self, anchor: CellRef, rows: int, cols: int, *, down: int = 0, right: int = 0
    ) -> list[list[CellValue]]:
        self._check(anchor)
        if rows < 0 or cols < 0:
            raise ValueError("block dimensions must be nonnegative")
        return [
            [self.value(anchor.row + down + row, anchor.col + right + col) for col in range(cols)]
            for row in range(rows)
        ]


class _XlsxSheet(_Sheet):
    def __init__(self, sheet: Worksheet, formulas: Worksheet) -> None:
        super().__init__(sheet.title, sheet.max_row, sheet.max_column)
        self._sheet = sheet
        self._formulas = formulas

    def _edge(self, anchor: CellRef) -> tuple[int, int]:
        for merged in self._sheet.merged_cells.ranges:
            if (
                merged.min_row <= anchor.row <= merged.max_row
                and merged.min_col <= anchor.col <= merged.max_col
            ):
                return merged.max_row, merged.max_col
        return anchor.row, anchor.col

    def _raw(self, row: int, col: int) -> Scalar:
        cached = self._sheet.cell(row, col)
        value: object = cached.value
        formula = self._formulas.cell(row, col).value
        if cached.data_type == "e":
            raise OfficeError("cell_error", f"{self.name}!{cached.coordinate}: {value}")
        if value is None and isinstance(formula, str) and formula.startswith("="):
            raise OfficeError("formula_uncached", f"{self.name}!{cached.coordinate}: {formula}")
        return cast(Scalar, value if isinstance(value, str | int | float | datetime) else None)

    def _hyperlink(self, row: int, col: int) -> str | None:
        link = self._formulas.cell(row, col).hyperlink
        return link.target if link is not None else None

    def _number_format(self, row: int, col: int) -> str | None:
        return self._formulas.cell(row, col).number_format

    def formula_inputs_blank(self, row: int, col: int) -> bool:
        formula = self._formulas.cell(row, col).value
        workbook = cast(Workbook, self._formulas.parent)
        return formula_inputs_blank(
            formula if isinstance(formula, str) and formula.startswith("=") else None,
            self.name,
            lambda name, r, c: workbook[name].cell(r, c).value,
        )


class _XlsSheet(_Sheet):
    def __init__(
        self, sheet: xlrd.sheet.Sheet, book: xlrd.book.Book, formulas: XlsFormulas
    ) -> None:
        super().__init__(sheet.name, sheet.nrows, sheet.ncols)
        self._sheet = sheet
        self._book = book
        self._datemode: Literal[0, 1] = book.datemode
        self._formulas = formulas

    def formula_inputs_blank(self, row: int, col: int) -> bool:
        return formula_inputs_blank(
            self._formulas.formula(self.name, row, col), self.name, self._dependency
        )

    def _dependency(self, name: str, row: int, col: int) -> Scalar:
        if self._formulas.contains(name, row, col):
            return "=formula"
        sheet = self._book.sheet_by_name(name)
        if row > sheet.nrows or col > sheet.ncols:
            return None
        return cast(Scalar, sheet.cell_value(row - 1, col - 1))

    def _number_format(self, row: int, col: int) -> str | None:
        if row > self.max_row or col > self.max_col:
            return None
        format_key = self._book.xf_list[self._sheet.cell_xf_index(row - 1, col - 1)].format_key
        item = self._book.format_map.get(format_key)
        return str(item.format_str) if item is not None else None

    def _raw(self, row: int, col: int) -> Scalar:
        if row > self.max_row or col > self.max_col:
            return None
        cell = self._sheet.cell(row - 1, col - 1)
        if cell.ctype == xlrd.XL_CELL_ERROR:
            raise OfficeError(
                "cell_error", f"{self.name}!{get_column_letter(col)}{row}: {cell.value}"
            )
        if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            return None
        if cell.ctype == xlrd.XL_CELL_DATE:
            return xlrd.xldate.xldate_as_datetime(float(cell.value), self._datemode)
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return int(cell.value)
        return cast(Scalar, cell.value)

    def _hyperlink(self, row: int, col: int) -> str | None:
        link = self._sheet.hyperlink_map.get((row - 1, col - 1))
        if link is None:
            return None
        return cast(str | None, link.url_or_path or link.target or None)


class _XlsxBook:
    def __init__(self, book: Workbook, formula_book: Workbook) -> None:
        self._book = book
        self._formula_book = formula_book

    @property
    def sheet_names(self) -> tuple[str, ...]:
        return tuple(self._book.sheetnames)

    def sheet(self, name: str) -> Sheet:
        if name not in self._book.sheetnames:
            raise OfficeError(
                "sheet_missing", f"{name}; available: {', '.join(self._book.sheetnames)}"
            )
        return _XlsxSheet(self._book[name], self._formula_book[name])

    def close(self) -> None:
        self._book.close()
        self._formula_book.close()


class _XlsBook:
    def __init__(self, book: xlrd.book.Book, formulas: XlsFormulas) -> None:
        self._book = book
        self._formulas = formulas

    @property
    def sheet_names(self) -> tuple[str, ...]:
        return tuple(self._book.sheet_names())

    def sheet(self, name: str) -> Sheet:
        if name not in self._book.sheet_names():
            raise OfficeError(
                "sheet_missing", f"{name}; available: {', '.join(self._book.sheet_names())}"
            )
        return _XlsSheet(self._book.sheet_by_name(name), self._book, self._formulas)

    def close(self) -> None:
        self._book.release_resources()


def open_book(path: Path) -> Book:
    detected = sniff(path)
    if detected.kind == FileKind.XLSX:
        return _XlsxBook(
            load_workbook(path, read_only=False, data_only=True),
            load_workbook(path, read_only=False, data_only=False),
        )
    if detected.kind == FileKind.XLS:
        book = xlrd.open_workbook(str(path), formatting_info=True)
        return _XlsBook(book, XlsFormulas(path, book))
    raise OfficeError(
        "unsupported_format", f"{path.name}: {detected.kind.value}; {detected.detail}"
    )
