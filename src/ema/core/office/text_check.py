"""Check that Word's DOC conversion retained text and tables."""

from __future__ import annotations

import re
import unicodedata
import zipfile
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from ema.core.office.word import DocText

_WORD = re.compile(r"[^\W_]+(?:['’-][^\W_]+)*", re.UNICODE)
_REFLOW_HYPHEN = re.compile(r"(?<=\w)-\s*\n\s*(?=\w)")


@dataclass(frozen=True)
class TextCheck:
    original_words: int
    converted_words: int
    shape_words: int
    original_tables: int
    converted_tables: int
    within_tolerance: bool
    warning: str | None


def check_text(original: DocText, converted: Path) -> TextCheck:
    """Compare ordered main-story words and tables; flag unchecked shape text."""
    Document(str(converted))  # Verify python-docx can open the result.
    with zipfile.ZipFile(converted) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    main_paragraphs, shape_text = _story_text(root)
    before = _words(original.text)
    after = _words(" ".join(main_paragraphs))
    words_before = len(before)
    words_after = len(after)
    shape_words = len(_WORD.findall(" ".join(shape_text)))
    tables_after = sum(1 for _ in root.iter(qn("w:tbl")))
    spans = [
        (before[a:b], after[c:d])
        for tag, a, b, c, d in SequenceMatcher(None, before, after, autojunk=False).get_opcodes()
        if tag != "equal"
    ]
    word_ok = not spans
    table_ok = original.tables == tables_after
    warning = None
    if not word_ok or not table_ok:
        excerpts = "; ".join(
            f"{' '.join(old[:5]) or '[missing]'} → {' '.join(new[:5]) or '[missing]'}"
            for old, new in spans
        )
        warning = (
            f"Text/tables differ: words {words_before}/{words_after}, "
            f"tables {original.tables}/{tables_after}. "
            f"Unmatched spans: {excerpts or 'none'}. Compare with original in Word."
        )
    elif shape_words and words_before == words_after == 0:
        warning = "Fidelity check not applicable (text in shapes): compare in Word."
    elif shape_words:
        warning = f"Shape text ({shape_words} words) could not be checked: compare in Word."
    return TextCheck(
        words_before,
        words_after,
        shape_words,
        original.tables,
        tables_after,
        word_ok and table_ok and shape_words == 0,
        warning,
    )


def _story_text(root: Any) -> tuple[list[str], list[str]]:
    main: list[str] = []
    for paragraph in root.iter(qn("w:p")):
        if _inside(paragraph, qn("w:txbxContent")):
            continue
        pieces: list[str] = []
        for node in paragraph.iter():
            if node.tag == qn("w:txbxContent") or _inside(node, qn("w:txbxContent"), paragraph):
                continue
            if node.tag == qn("w:t"):
                pieces.append(node.text or "")
            elif node.tag in {qn("w:tab"), qn("w:br"), qn("w:cr")}:
                pieces.append(" ")
        main.append("".join(pieces))
    shapes = [
        node.text or "" for box in root.iter(qn("w:txbxContent")) for node in box.iter(qn("w:t"))
    ]
    return main, shapes


def _inside(node: Any, tag: str, stop: Any | None = None) -> bool:
    parent = node.getparent()
    while parent is not None and parent is not stop:
        if parent.tag == tag:
            return True
        parent = parent.getparent()
    return False


def _words(text: str) -> list[str]:
    # These are known layout/typography normalizations, not missing or changed words.
    normalized = unicodedata.normalize("NFC", text).replace("\u00ad", "").replace("\u00a0", " ")
    normalized = normalized.translate(str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"'}))
    return [word.casefold() for word in _WORD.findall(_REFLOW_HYPHEN.sub("", normalized))]
