"""Write a source-backed Prelucrare workbook with live spreadsheet formulas."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ema.energy_data.carriers import Carrier
from ema.energy_data.factors import FACTORS_2026, FactorTable
from ema.energy_data.model import CarrierSeries, EnergyDataset
from ema.energy_data.prelucrare import MONTHS

NUMBER = "#,##0.00;[Red](#,##0.00);–"
NAMES = {
    Carrier.electricity_grid: ("Consum Electric", "[MWh]"),
    Carrier.natural_gas: ("Consum Gaz", "[MWh]"),
    Carrier.electricity_pv: ("Consum electrica fotovoltaic", "[MWh]"),
    Carrier.diesel: ("Consum Carburanti", "Motorina [t]"),
    Carrier.petrol: ("Consum Carburanti", "Benzina [t]"),
    Carrier.lpg: ("Consum Carburanti", "GPL [t]"),
    Carrier.ctl: ("Consum Carburanti", "CTL [t]"),
    Carrier.sunflower_husks: ("Consum Coji floarea soarelui", "[Gcal]"),
    Carrier.purchased_heat: ("Consum Energie termica terti", "[Gcal]"),
    Carrier.water_potable: ("consum apa potabila", "[m3]"),
    Carrier.water_industrial: ("consum apa industriala", "[m3]"),
}
SPECIFIC_LABELS = {
    Carrier.electricity_grid: "energie electrica",
    Carrier.electricity_pv: "energie electrica fotovoltaica",
    Carrier.natural_gas: "gaze naturale",
    Carrier.diesel: "motorina",
    Carrier.petrol: "benzina",
    Carrier.lpg: "GPL",
    Carrier.ctl: "CTL",
    Carrier.sunflower_husks: "coji de floarea soarelui",
    Carrier.purchased_heat: "energie termica terti",
}


def _sheet(book: Workbook, name: str) -> Worksheet:
    return book[name] if name in book else book.create_sheet(name)


def _header(sheet: Worksheet, year: int, row: int) -> None:
    sheet.cell(row, 1, year)
    for month, label in enumerate(MONTHS, 4):
        cell = sheet.cell(row, month, label.capitalize())
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="33516C")
    sheet.cell(row, 16, "TOTAL")
    sheet.column_dimensions["C"].width = 30
    for col in range(4, 17):
        sheet.column_dimensions[get_column_letter(col)].width = 14


def _physical(
    book: Workbook, ds: EnergyDataset, years: tuple[int, ...]
) -> dict[tuple[Carrier, int, int | None], str]:
    references: dict[tuple[Carrier, int, int | None], str] = {}
    groups: dict[str, list[tuple[Carrier, str]]] = {}
    for carrier in ds.carriers:
        if carrier in NAMES:
            sheet_name, label = NAMES[carrier]
            groups.setdefault(sheet_name, []).append((carrier, label))
    for sheet_name, entries in groups.items():
        sheet = _sheet(book, sheet_name)
        for index, year in enumerate(years):
            header = 2 + index * (len(entries) + 3)
            _header(sheet, year, header)
            for offset, (carrier, label) in enumerate(entries, 1):
                row = header + offset
                sheet.cell(row, 3, label)
                series = ds.carriers.get(carrier, {}).get(year, CarrierSeries())
                for month in range(1, 13):
                    reading = series.months.get(month)
                    if reading is not None and reading.value is not None:
                        cell = sheet.cell(row, month + 3, reading.value)
                        cell.number_format = NUMBER
                        references[carrier, year, month] = f"'{sheet_name}'!{cell.coordinate}"
                if series.annual is not None and series.annual.value is not None:
                    sheet.cell(row, 16, series.annual.value)
                elif len(series.months) == 12 and all(
                    r.value is not None for r in series.months.values()
                ):
                    sheet.cell(row, 16, f"=SUM(D{row}:O{row})")
                sheet.cell(row, 16).number_format = NUMBER
                references[carrier, year, None] = f"'{sheet_name}'!P{row}"
    return references


def _factor_sheet(
    book: Workbook, ds: EnergyDataset, years: tuple[int, ...], factors: FactorTable
) -> dict[tuple[Carrier, int], str]:
    sheet = _sheet(book, "Principali factori de conversie")
    sheet.append(("Purtător", "Unitate", "An", "tep/unitate", "Sursă"))
    result: dict[tuple[Carrier, int], str] = {}
    for carrier, series in ds.carriers.items():
        for year in years:
            item = series.get(year)
            if item is None:
                continue
            reading = item.annual or next(iter(item.months.values()), None)
            if reading is None:
                continue
            factor = factors.tep_factor(carrier, reading.unit, year)
            if factor is None:
                continue
            sheet.append((carrier.value, reading.unit, year, factor.per_unit, factor.source))
            result[carrier, year] = f"'Principali factori de conversie'!D{sheet.max_row}"
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["E"].width = 48
    return result


def _mwh_factor_sheet(
    book: Workbook, factors: FactorTable, factor_refs: dict[tuple[Carrier, int], str]
) -> None:
    sheet = _sheet(book, "Factori de conversie in MWh")
    sheet.append(("Purtător", "Unitate", "An", "Factor [tep/unitate]", "Sursă", "MWh/unitate"))
    for factor in factors.tep:
        sheet.append(
            (factor.carrier.value, factor.unit, factor.year, factor.per_unit, factor.source)
        )
        row = sheet.max_row
        grid = factor_refs.get((Carrier.electricity_grid, factor.year or 0))
        if grid:
            sheet.cell(row, 6, f"=D{row}/{grid}").number_format = NUMBER
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["E"].width = 48


def _tep(
    book: Workbook,
    ds: EnergyDataset,
    years: tuple[int, ...],
    sources: dict[tuple[Carrier, int, int | None], str],
    factors: dict[tuple[Carrier, int], str],
) -> dict[tuple[Carrier | None, int], str]:
    sheet = _sheet(book, "TEP")
    annual: dict[tuple[Carrier | None, int], str] = {}
    for index, year in enumerate(years):
        start = 2 + index * (len(ds.carriers) + 4)
        sheet.cell(start, 3, year)
        _header(sheet, year, start + 1)
        rows: list[int] = []
        for offset, carrier in enumerate(ds.carriers, 2):
            row = start + offset
            sheet.cell(row, 3, f"{carrier.value} [tep]")
            factor = factors.get((carrier, year))
            for month in range(1, 13):
                source = sources.get((carrier, year, month))
                if source and factor:
                    sheet.cell(row, month + 3, f"={source}*{factor}").number_format = NUMBER
            source = sources.get((carrier, year, None))
            if source and factor:
                sheet.cell(row, 16, f"={source}*{factor}").number_format = NUMBER
                annual[carrier, year] = f"'TEP'!P{row}"
            rows.append(row)
        total_row = start + len(ds.carriers) + 2
        sheet.cell(total_row, 3, "TOTAL [tep]")
        for col in range(4, 17):
            letter = get_column_letter(col)
            refs = ",".join(
                f"{letter}{row}" for row in rows if sheet.cell(row, col).value is not None
            )
            if refs:
                sheet.cell(total_row, col, f"=SUM({refs})").number_format = NUMBER
        annual[None, year] = f"'TEP'!P{total_row}"
    return annual


def _production(book: Workbook, ds: EnergyDataset, years: tuple[int, ...]) -> dict[int, str]:
    sheet = _sheet(book, "Productii")
    result: dict[int, str] = {}
    for index, year in enumerate(years):
        header = 2 + index * 4
        _header(sheet, year, header)
        for offset, (product, series_by_year) in enumerate(ds.production.items(), 1):
            series = series_by_year.get(year)
            if series is None:
                continue
            row = header + offset
            unit = ds.production_unit.get(product, "")
            label = (
                "kWh gaz vehiculat"
                if unit == "mii MWh gaz vehiculat"
                else "tone/luna"
                if unit == "mii tone"
                else product
            )
            sheet.cell(row, 3, label)
            multiplier = (
                1_000_000 if unit == "mii MWh gaz vehiculat" else 1000 if unit == "mii tone" else 1
            )
            for month, reading in series.months.items():
                sheet.cell(
                    row,
                    month + 3,
                    reading.value * multiplier if reading.value is not None else None,
                ).number_format = NUMBER
            sheet.cell(
                row,
                16,
                series.annual.value * multiplier
                if series.annual and series.annual.value is not None
                else f"=SUM(D{row}:O{row})",
            ).number_format = NUMBER
            if offset == 1:
                result[year] = f"('Productii'!P{row}/{multiplier})"
    return result


def _economics(
    book: Workbook,
    ds: EnergyDataset,
    years: tuple[int, ...],
    tep: dict[tuple[Carrier | None, int], str],
) -> None:
    sheet = _sheet(book, "Chelt-Cifra afaceri")
    sheet.cell(2, 3, "Anul")
    sheet.cell(3, 3, "Cifra de afaceri [lei]")
    sheet.cell(4, 3, "Cheltuieli cu energia [lei]")
    sheet.cell(5, 3, "Intensitate energetica [tep/1000 lei]")
    for col, year in enumerate(years, 4):
        sheet.cell(2, col, year)
        turnover = ds.turnover_lei.get(year)
        costs = ds.energy_costs_lei.get(year)
        if turnover is not None:
            sheet.cell(3, col, turnover.value).number_format = NUMBER
        if costs is not None:
            sheet.cell(4, col, costs.value).number_format = NUMBER
        if turnover is not None and tep.get((None, year)):
            sheet.cell(
                5, col, f"={tep[None, year]}/({get_column_letter(col)}3/1000)"
            ).number_format = NUMBER
    sheet.column_dimensions["C"].width = 42


def _specific(
    book: Workbook,
    years: tuple[int, ...],
    tep: dict[tuple[Carrier | None, int], str],
    production: dict[int, str],
) -> None:
    sheet = _sheet(book, "Consumuri specifice")
    sheet.cell(2, 2, "Anul")
    sheet.cell(2, 3, "Consum specific anual de energie [tep]")
    carriers = sorted(
        {carrier for carrier, _ in tep if carrier is not None and carrier in SPECIFIC_LABELS},
        key=lambda carrier: carrier.value,
    )
    for col, carrier in enumerate(carriers, 4):
        sheet.cell(2, col, f"Consum specific anual de {SPECIFIC_LABELS[carrier]} [tep]")
        sheet.column_dimensions[get_column_letter(col)].width = 35
    for row, year in enumerate(years, 3):
        sheet.cell(row, 2, year)
        if year in production and (None, year) in tep:
            sheet.cell(row, 3, f"={tep[None, year]}/{production[year]}").number_format = NUMBER
            for col, carrier in enumerate(carriers, 4):
                if (carrier, year) in tep:
                    sheet.cell(
                        row, col, f"={tep[carrier, year]}/{production[year]}"
                    ).number_format = NUMBER


def _impact(
    book: Workbook,
    ds: EnergyDataset,
    years: tuple[int, ...],
    sources: dict[tuple[Carrier, int, int | None], str],
    factors: FactorTable,
) -> None:
    sheet = _sheet(book, "impact de mediu")
    for index, year in enumerate(years):
        start = 2 + index * (len(ds.carriers) + 3)
        sheet.cell(start, 2, year)
        rows: list[int] = []
        for offset, carrier in enumerate(ds.carriers, 1):
            row = start + offset
            sheet.cell(row, 2, carrier.value)
            series = ds.carriers[carrier].get(year)
            if series is None:
                continue
            reading = series.annual or next(iter(series.months.values()), None)
            source = sources.get((carrier, year, None))
            factor = factors.co2_factor(carrier, reading.unit, year) if reading else None
            if factor and source and reading is not None:
                factor_col = 6 if reading.unit == "MWh" else 8
                sheet.cell(row, factor_col, factor.per_unit)
                factor_ref = f"{get_column_letter(factor_col)}{row}"
                sheet.cell(row, 9, f"={source}*{factor_ref}").number_format = NUMBER
                rows.append(row)
        total = start + len(ds.carriers) + 1
        sheet.cell(total, 2, "Indicator global prin suprapunerea efectelor")
        if rows:
            sheet.cell(
                total, 9, "=SUM(" + ",".join(f"I{row}" for row in rows) + ")"
            ).number_format = NUMBER


def write_prelucrare(
    dataset: EnergyDataset, years: tuple[int, ...], path: Path, factors: FactorTable = FACTORS_2026
) -> None:
    if (
        not years
        or tuple(sorted(set(years))) != years
        or any(year not in dataset.years for year in years)
    ):
        raise ValueError("years must be present in the dataset, unique and ascending")
    book = Workbook()
    active = book.active
    assert active is not None
    book.remove(active)
    sources = _physical(book, dataset, years)
    factor_refs = _factor_sheet(book, dataset, years, factors)
    tep = _tep(book, dataset, years, sources, factor_refs)
    production = _production(book, dataset, years)
    _economics(book, dataset, years, tep)
    _specific(book, years, tep, production)
    _impact(book, dataset, years, sources, factors)
    _mwh_factor_sheet(book, factors, factor_refs)
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    book.save(path)
