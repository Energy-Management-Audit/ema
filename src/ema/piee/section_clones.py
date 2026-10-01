"""Compose sourced carrier sections from the approved PIEE prototypes."""

# pyright: reportPrivateUsage=false

from __future__ import annotations

import copy
import json
from pathlib import Path
from tempfile import TemporaryDirectory

from docx.oxml.ns import qn
from lxml import etree

from ema.core.office.anchors import AnchorLedger, find
from ema.core.office.cell_text import set_cell_text
from ema.core.office.package import C, encoded, read_parts, write_parts, xml
from ema.core.office.run_range import TextSpan, replace_spans, visible_text
from ema.energy_data.calc import tep
from ema.energy_data.carriers import WATER_CARRIERS, Carrier
from ema.piee.dataset import PieeData
from ema.piee.number import prototype_number
from ema.piee.section_carrier_group import _carrier_group
from ema.piee.section_clone_tools import (
    ORDER,
    VOCABULARY,
    Figure,
    _add_chart,
    _clone_figures,
    _figures,
    _insert_before,
    _paragraph,
    _present,
    _remove_before,
    _role,
    _take_charts,
    _track,
)
from ema.piee.section_water_body import render_potable_water

COGEN_MIX_SENTENCE = (
    "În această pondere nu este prezentată distinct cantitatea de energie electrică produsă "
    "prin cogenerare întrucât, pentru obținerea ei se utilizează gazul natural care este prezentat."
)
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _chart_unit(data: PieeData, item: Figure) -> str:
    production = next(iter(data.dataset.production_unit.values()), "")
    if item.kind == "specific":
        return (
            f"m3/{production}"
            if item.carrier in WATER_CARRIERS
            else (f"tep/{'mii tone' if production == 'tone' else production}")
        )
    if item.kind == "total":
        return "tep"
    if item.carrier in WATER_CARRIERS:
        return "m3"
    assert item.carrier is not None
    return next(
        (
            reading.annual.unit
            for reading in data.dataset.carriers[item.carrier].values()
            if reading.annual
        ),
        "MWh",
    )


def _set_chart_unit(parts: dict[str, bytes], part: str, unit: str) -> None:
    chart = xml(parts, part)
    for title in chart.findall(f".//{{{C}}}valAx/{{{C}}}title"):
        texts = list(title.iter(A + "t"))
        for index, node in enumerate(texts):
            node.text = unit if index == 0 else ""
    parts[part] = encoded(chart)


def _append_caption_unit(paragraph: etree._Element, unit: str) -> None:
    content = visible_text(paragraph)
    if unit in content:
        return
    for old in ("tep/mii tone", "tep/tone", "m3/tone"):
        if old in content:
            start = content.index(old)
            replace_spans(paragraph, (TextSpan(start, start + len(old), unit),))
            return
    texts = list(paragraph.iter(qn("w:t")))
    if texts:
        texts[-1].text = (texts[-1].text or "") + f", {unit}"


def _coke_centralizer(
    root: etree._Element, data: PieeData, ledger: AnchorLedger, next_id: list[int]
) -> None:
    marker = find([root], "table_277_r1_c3")
    table = next(marker.iterancestors(qn("w:tbl")), None)
    if table is None:
        raise ValueError("equivalent-energy table missing")
    grid = table.find(qn("w:tblGrid"))
    if grid is not None and len(grid) >= 4:
        grid.insert(len(grid) - 1, copy.deepcopy(grid[-2]))
    rows = table.findall(qn("w:tr"))
    if len(rows) != 4:
        raise ValueError("equivalent-energy table has unexpected rows")
    for row_index, row in enumerate(rows):
        cells = row.findall(qn("w:tc"))
        cloned = copy.deepcopy(cells[-2])
        for mark in tuple(cloned.iter()):
            if mark.tag in {qn("w:bookmarkStart"), qn("w:bookmarkEnd")}:
                parent = mark.getparent()
                if parent is not None:
                    parent.remove(mark)
        year = data.year - 3 + row_index
        amount = tep(data.dataset, data.factors, Carrier.coke, year).value if row_index else None
        content = (
            "Cocs"
            if row_index == 0
            else prototype_number(amount, 2)
            if amount is not None
            else "n.d."
        )
        set_cell_text(cloned, content, missing=content == "n.d.")
        _track(cloned, f"coke_centralizer_{row_index}", ledger=ledger, next_id=next_id)
        row.insert(len(cells) - 1, cloned)


def render_section_clones(  # noqa: C901, PLR0912, PLR0915
    source: Path,
    manifest_path: Path,
    data: PieeData,
    output: Path,
    ledger: AnchorLedger,
) -> None:
    """Add missing carrier and water sections, then remove absent base groups."""
    manifest: dict[str, dict[str, str | list[str]]] = json.loads(
        manifest_path.read_text(encoding="utf-8")
    )
    figures = _figures(data)
    with TemporaryDirectory() as directory:
        current, generated = _clone_figures(source, data, figures, Path(directory))
        parts = read_parts(current)
        root = xml(parts, "word/document.xml")
        charts = _take_charts(parts, root, generated)
        next_id = [
            max(
                (int(node.get(qn("w:id"), "0")) for node in root.iter(qn("w:bookmarkStart"))),
                default=0,
            )
            + 1
        ]
        boundaries = manifest["boundaries"]
        if not _present(data, Carrier.electricity_pv):
            _remove_before(
                root, str(boundaries["pv_heading"]), str(boundaries["gas_heading"]), ledger
            )
        if not data.layout.water_monthly and _present(data, Carrier.water_potable):
            _remove_before(
                root,
                _role(manifest, "water", "monthly_intro"),
                _role(manifest, "water", "annual_intro"),
                ledger=ledger,
            )
        for carrier in ORDER:
            if not _present(data, carrier):
                continue
            boundary = (
                boundaries["fuel_heading"]
                if carrier == Carrier.electricity_cogen
                else boundaries["water_heading"]
                if carrier == Carrier.coke
                else boundaries["equivalent_heading"]
            )
            _insert_before(
                root,
                str(boundary),
                _carrier_group(
                    root, manifest, data, carrier, charts, ledger=ledger, next_id=next_id
                ),
            )
        if _present(data, Carrier.electricity_cogen):
            sentence = _paragraph(
                root,
                manifest,
                "gas_annual",
                "annual_observation",
                COGEN_MIX_SENTENCE,
                "electricity_cogen_mix_observation",
                ledger=ledger,
                next_id=next_id,
            )
            _insert_before(root, str(boundaries["specific_group_heading"]), [sentence])
        if data.layout.total_energy_figure:
            item = Figure("word/charts/chart16.xml", None, "total", None)
            nodes = [
                _paragraph(
                    root,
                    manifest,
                    "gas_annual",
                    "annual_intro",
                    "Evoluția consumului total de energie echivalentă este prezentată "
                    "în figura următoare.",
                    "total_energy_annual_intro",
                    ledger=ledger,
                    next_id=next_id,
                )
            ]
            _add_chart(nodes, charts[item], None, "annual", ledger=ledger, next_id=next_id)
            nodes.append(
                _paragraph(
                    root,
                    manifest,
                    "gas_annual",
                    "annual_observation",
                    "Consumul total de energie echivalentă este prezentat mai sus.",
                    "total_energy_annual_observation",
                    ledger=ledger,
                    next_id=next_id,
                )
            )
            _insert_before(root, str(boundaries["specific_group_heading"]), nodes)
        available = data.layout.specific_carriers
        for carrier, first, last in (
            (Carrier.electricity_grid, "specific_grid_heading", "specific_gas_heading"),
            (Carrier.natural_gas, "specific_gas_heading", "specific_fuel_heading"),
            (Carrier.diesel, "specific_fuel_heading", "specific_total_heading"),
        ):
            if not _present(data, carrier) or (available is not None and carrier not in available):
                _remove_before(root, str(boundaries[first]), str(boundaries[last]), ledger)
        for carrier in (Carrier.coke, Carrier.water_industrial, Carrier.water_storm):
            item = Figure(
                "word/charts/chart26.xml" if carrier == Carrier.coke else "word/charts/chart29.xml",
                carrier,
                "specific",
                None,
            )
            if item not in charts:
                continue
            group = "specific_gas"
            noun = VOCABULARY[carrier][0]
            heading = f"Consumul specific de {noun}"
            nodes = [
                _paragraph(
                    root,
                    manifest,
                    group,
                    "heading",
                    heading,
                    f"{carrier.value}_specific_heading",
                    ledger=ledger,
                    next_id=next_id,
                ),
                _paragraph(
                    root,
                    manifest,
                    group,
                    "intro",
                    f"Consumul specific de {noun} este prezentat în figura următoare.",
                    f"{carrier.value}_specific_intro",
                    ledger=ledger,
                    next_id=next_id,
                ),
            ]
            _add_chart(nodes, charts[item], carrier, "specific", ledger=ledger, next_id=next_id)
            nodes.append(
                _paragraph(
                    root,
                    manifest,
                    group,
                    "annual_observation",
                    f"Evoluția consumului specific de {noun} este prezentată mai sus.",
                    f"{carrier.value}_specific_observation",
                    ledger=ledger,
                    next_id=next_id,
                )
            )
            boundary = (
                boundaries["specific_total_heading"]
                if carrier == Carrier.coke
                else boundaries["intensity_heading"]
            )
            _insert_before(root, str(boundary), nodes)
        render_potable_water(root, manifest, data, ledger)
        if _present(data, Carrier.coke):
            _coke_centralizer(root, data, ledger=ledger, next_id=next_id)
        for item, part in generated.items():
            unit = _chart_unit(data, item)
            _set_chart_unit(parts, part, unit)
            if item.kind == "specific":
                _append_caption_unit(charts[item][1], unit)
        original_specific = (
            (25, Carrier.electricity_grid),
            (26, Carrier.natural_gas),
            (27, Carrier.diesel),
            (28, None),
            (29, Carrier.water_potable),
        )
        present = {node.get(qn("w:name")) for node in root.iter(qn("w:bookmarkStart"))}
        for number, carrier in original_specific:
            if f"_ema_chart_{number}" not in present:
                continue
            unit = _chart_unit(
                data, Figure(f"word/charts/chart{number}.xml", carrier, "specific", None)
            )
            _set_chart_unit(parts, f"word/charts/chart{number}.xml", unit)
            chart_paragraph = find([root], f"chart_{number}")
            caption = chart_paragraph.getnext()
            if caption is not None and "tone" in unit:
                _append_caption_unit(caption, unit)
        parts["word/document.xml"] = encoded(root)
        write_parts(parts, output)
