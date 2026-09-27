"""Reporting selects indexed beneficiaries and leaves other client uploads alone."""

from __future__ import annotations

import threading
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from tests.unit.energy_data.test_annex_index import annex

from ema.clients.files import store_upload
from ema.clients.registry import create_client
from ema.core.errors import EmaError
from ema.core.workspace import Workspace
from ema.energy_data.annex_index import import_annexes, indexed
from ema.reporting.runs import get_run, list_runs, preview, start_run


def _import(ws: Workspace, path: Path) -> None:
    with path.open("rb") as stream:
        result = import_annexes(ws, [(path.name, stream)])
    assert len(result.imported) == 1


def _finish(ws: Workspace, run_id: str) -> dict[str, object]:
    for worker in threading.enumerate():
        if worker.name.startswith("ema-reporting-"):
            worker.join(timeout=30)
            assert not worker.is_alive()
    return get_run(ws, run_id)


def test_newest_per_beneficiary_keeps_two_sites_and_preview(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    client = create_client(ws, "Companie Exemplu", "12345678")
    missing = create_client(ws, "Fără anexă", "87654321")
    first = tmp_path / "Anexa – à ^.xlsx"
    annex(first, name="Punct de lucru A", total=10)
    _import(ws, first)
    newer = tmp_path / "Anexa A nouà.xlsx"
    annex(newer, name=" PUNCT  DE   LUCRU A ", total=12)
    _import(ws, newer)
    second = tmp_path / "Anexa B.xlsx"
    annex(second, name="Punct de lucru B", total=20)
    _import(ws, second)
    unrelated = tmp_path / "Necesar info.xlsx"
    annex(unrelated, name="Alte date", total=100)
    with unrelated.open("rb") as stream:
        store_upload(ws, str(client["id"]), stream, unrelated.name)
    started = start_run(ws, [2025], [str(client["id"]), str(missing["id"])])
    run = _finish(ws, str(started["id"]))
    assert run["state"] == "ready"
    assert run["job_id"]
    assert run["created_at"]
    data = preview(ws, str(run["id"]))
    assert data["years"] == [2025]
    assert data["read"] == 2
    assert data["companies_per_year"] == {"2025": 2}
    assert len(data["rows"]["2025"]) == 2
    assert {row["beneficiary"] for row in data["rows"]["2025"]} == {
        "PUNCT  DE   LUCRU A",
        "Punct de lucru B",
    }
    assert any(item["code"] == "annex_missing" for item in run["exceptions"])
    assert list_runs(ws)[0]["id"] == run["id"]
    with ws.connect() as db:
        row = db.execute(
            "SELECT relative_path FROM outputs WHERE id=?", (run["output_id"],)
        ).fetchone()
    book = load_workbook(ws.path(str(row["relative_path"])), read_only=True, data_only=True)
    control_files = [row[1] for row in list(book["Control"].values)[4:]]
    assert unicodedata.normalize("NFC", newer.name) in control_files
    assert any(row.data["file_name"] == newer.name for row in indexed(ws)[str(client["id"])])
    assert "Anexa B.xlsx" in control_files
    assert "Necesar info.xlsx" not in control_files


def test_preview_not_ready_and_missing_cost_ref(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    client = create_client(ws, "Companie Exemplu", "12345678")
    path = tmp_path / "annex.xlsx"
    annex(path)
    book = load_workbook(path)
    book["Solutii EE existente"]["E5"] = None
    book.save(path)
    _import(ws, path)
    started = start_run(ws, [2025], [str(client["id"])])
    try:
        preview(ws, str(started["id"]))
    except EmaError as error:
        assert error.code == "run_not_ready"
    run = _finish(ws, str(started["id"]))
    assert run["state"] == "ready"
    assert any(item["ref"] == "Solutii EE existente!B5" for item in run["exceptions"])
