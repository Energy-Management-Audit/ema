"""Incomplete report totals, source traceability and deleted-history recovery."""

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from tests.unit.energy_data.test_annex_index import annex
from tests.unit.reporting.test_report import _annex
from tests.unit.reporting.test_runs_index import _finish, _import
from typer.testing import CliRunner

from ema.cli import _app
from ema.clients.registry import create_client
from ema.core.workspace import Workspace
from ema.core.workspace.cleanup import delete_job_rows
from ema.reporting import generate, write_report
from ema.reporting.runs import generate_from_sources, list_runs, start_run


@pytest.mark.parametrize("missing", ["cost", "saving"])
def test_missing_measure_values_leave_totals_empty(tmp_path: Path, missing: str) -> None:
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    _annex(first)
    _annex(second)
    book = load_workbook(second)
    book["Solutii EE existente"]["E5" if missing == "cost" else "G5"] = None
    book.save(second)
    result = generate([first, second], (2025,))
    output = write_report(result, tmp_path / "report.xlsx")
    book = load_workbook(output, data_only=True)
    total = next(row for row in book["2025"].values if row[0] == "TOTAL")
    assert total[6 if missing == "cost" else 5] is None
    assert total[5 if missing == "cost" else 6] is not None
    message = (
        "Costul investiției lipsește." if missing == "cost" else "Economia de energie lipsește."
    )
    assert any(item.situation == message and item.source == second for item in result.exceptions)
    assert any(row[3] == message for row in book["Exceptions"].values)


def test_control_follows_requested_years_and_never_guesses_columns(tmp_path: Path) -> None:
    path = tmp_path / "annex.xlsx"
    _annex(path, monthly=None)
    source = load_workbook(path)
    source["Solutii EE existente"]["E5"] = None
    source["Solutii EE existente"]["G5"] = None
    source.save(path)
    result = generate([path], (2024, 2025, 2026, 2027))
    output = write_report(result, tmp_path / "report.xlsx")
    book = load_workbook(output, data_only=True)
    control = book["Control"]
    assert list(control.values)[3][13:17] == (
        "Măsuri 2024",
        "Măsuri 2025",
        "Măsuri 2026",
        "Măsuri 2027",
    )
    assert list(control.values)[4][13:17] == (0, 1, 0, 0)
    assert "surse Anexa 2–3" in control.cell(1, 1).value
    assert result.companies[0].measure_sheet.endswith("cost —, economie —")
    assert "fallback" not in result.companies[0].consumption_source
    assert "M3" not in result.companies[0].consumption_source
    assert book["2025"].column_dimensions["G"].width == 18
    assert control.column_dimensions["S"].width == 70


def _report(ws: Workspace, tmp_path: Path) -> dict[str, object]:
    client = create_client(ws, "Synthetic", "12345678")
    path = tmp_path / "annex.xlsx"
    annex(path)
    _import(ws, path)
    return _finish(ws, str(start_run(ws, [2025], [str(client["id"])])["id"]))


def test_delete_completed_reporting_job_cleans_metadata_and_events(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    run = _report(ws, tmp_path)
    assert run["state"] == "ready"
    job = str(run["job_id"])
    ws.delete_job(job)
    assert list_runs(ws) == []
    with ws.connect() as db:
        for table in ("reporting_runs", "job_events"):
            assert (
                db.execute(f"SELECT COUNT(*) FROM {table} WHERE job_id=?", (job,)).fetchone()[0]
                == 0
            )


def test_legacy_orphan_is_skipped_and_cleanup_rolls_back(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    run = _report(ws, tmp_path)
    job = str(run["job_id"])
    with pytest.raises(RuntimeError, match="interrupted"), ws.connect() as db:
        db.execute("BEGIN IMMEDIATE")
        delete_job_rows(db, job)
        raise RuntimeError("interrupted")
    assert list_runs(ws)[0]["id"] == run["id"]
    with ws.connect() as db:
        db.execute(
            "INSERT INTO reporting_runs(id,job_id,years_json,client_ids_json,state) "
            "VALUES ('orphan','deleted','[2025]','[]','ready')"
        )
    assert [item["id"] for item in list_runs(ws)] == [run["id"]]


def test_delete_recovers_when_job_folder_already_removed(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    run = _report(ws, tmp_path)
    with ws.connect() as db:
        row = db.execute("SELECT relative_path FROM jobs WHERE id=?", (run["job_id"],)).fetchone()
        db.execute("UPDATE jobs SET deleted=1 WHERE id=?", (run["job_id"],))
    shutil.rmtree(ws.path(str(row["relative_path"])))
    ws.finish_deletes()
    assert list_runs(ws) == []


def test_cli_and_ui_share_job_output_and_source_deduplication(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    folder = tmp_path / "sources"
    folder.mkdir()
    annex(folder / "a-old.xlsx", name="Punct A", total=10)
    annex(folder / "z-new.xlsx", name=" PUNCT  A ", total=20)
    root = tmp_path / "ws"
    monkeypatch.setenv("EMA_WORKSPACE", str(root))
    output = tmp_path / "cli.xlsx"
    response = CliRunner().invoke(
        _app, ["reporting", "generate", str(folder), "--years", "2024-2026", "--out", str(output)]
    )
    assert response.exit_code == 0, response.output
    assert output.exists()
    ws = Workspace(root)
    runs = list_runs(ws)
    assert len(runs) == 1 and runs[0]["state"] == "ready" and runs[0]["output_id"]
    cli = load_workbook(output, data_only=True)
    control = list(cli["Control"].values)
    assert len(control[4:]) == 1
    assert control[4][1] == "z-new.xlsx" and control[4][11] == 20
    ui = _finish(ws, str(start_run(ws, [2024, 2025, 2026], runs[0]["client_ids"])["id"]))
    with ws.connect() as db:
        row = db.execute(
            "SELECT relative_path FROM outputs WHERE id=?", (ui["output_id"],)
        ).fetchone()
    workbook = load_workbook(ws.path(str(row["relative_path"])), data_only=True)
    for sheet in cli.sheetnames:
        assert list(cli[sheet].values) == list(workbook[sheet].values)


def test_folder_report_keeps_valid_sources_when_one_workbook_fails(tmp_path: Path) -> None:
    good, bad = tmp_path / "good.xlsx", tmp_path / "bad.xlsx"
    annex(good)
    bad.write_bytes(b"invalid workbook")
    ws = Workspace(tmp_path / "ws")
    output = generate_from_sources(ws, [good, bad], [2025], tmp_path / "report.xlsx")
    assert list_runs(ws)[0]["state"] == "ready"
    book = load_workbook(output, data_only=True)
    assert len(list(book["Control"].values)[4:]) == 1
    assert any(row[1] == bad.name for row in book["Exceptions"].values)
