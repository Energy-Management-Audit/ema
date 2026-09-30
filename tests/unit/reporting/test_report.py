"""Reporting behaviour on synthetic annexes."""

from pathlib import Path

from openpyxl import Workbook, load_workbook

from ema.energy_data.anexa import parse_anexa
from ema.reporting import collect_annexes, generate, write_report


def test_collect_annexes_expands_filters_and_orders(tmp_path: Path) -> None:
    folder = tmp_path / "sources"
    folder.mkdir()
    for name in ("z.XLS", "A.xlsx", "ignore.txt"):
        (folder / name).touch()
    direct = tmp_path / "b.xls"
    direct.touch()
    assert [path.name for path in collect_annexes([folder, direct])] == ["A.xlsx", "b.xls", "z.XLS"]


def _annex(
    path: Path,
    *,
    name_is_address: bool = False,
    alternate_name: bool = False,
    measures: bool = True,
    cost: float | None = 12.0,
    monthly: float | None = 25.0,
    annual: float = 24.0,
    name: str = "Companie Exemplu",
) -> None:
    book = Workbook()
    general = book.active
    assert general is not None
    general.title = "Date generale"
    general["A1"] = "Denumirea operatorului economic"
    general["B1"] = "Strada Exemplu 1" if name_is_address else name
    general["A2"] = "Adresa poștală"
    general["B2"] = "Strada Exemplu 1"
    general["A3"] = "Cod CAEN"
    general["B3"] = 1234
    general["A4"] = "Sector de activitate"
    general["B4"] = "Industrie"
    if alternate_name:
        other = book.create_sheet("Identitate")
        other["A1"] = "Denumire"
        other["B1"] = "Nume Alternativ"
    monthly_sheet = book.create_sheet("Date lunare")
    monthly_sheet["A2"] = "CONSUM DE ENERGIE TOTAL ANUAL"
    monthly_sheet["J2"] = "[ tep / an ]"
    monthly_sheet["M3"] = monthly
    annual_sheet = book.create_sheet("Date anuale")
    annual_sheet["A2"] = "CONSUM DE ENERGIE TOTAL ANUAL"
    annual_sheet["F3"] = "tep/an"
    annual_sheet["G3"] = annual
    existing = book.create_sheet("Solutii EE existente")
    existing["B3"] = "Descrierea măsurii aplicate"
    existing["C3"] = "Data punerii în funcţiune"
    existing["E3"] = "Costul investiției"
    existing["G4"] = "tep/an"
    if measures:
        existing["B5"] = "Măsură de test"
        existing["C5"] = 2025
        existing["E5"] = cost
        existing["G5"] = 3.5
    book.save(path)


def test_two_branch_declaration_keeps_both_name_views(tmp_path: Path) -> None:
    path = tmp_path / "anexa.xlsx"
    declaration = "Operator Exemplu SRL, SUCURSALELE SITE-1 si SITE-2"
    _annex(path, name=declaration)

    identity = parse_anexa(path).identity
    assert identity["declared_name"].value == declaration
    assert identity["declared_name"].ref == identity["name"].ref
    assert identity["name"].value == "Operator Exemplu SRL"
    assert identity["site_1_name"].value == "Site-1"
    assert identity["site_2_name"].value == "Site-2"
    assert generate([path], (2025,)).companies[0].name == declaration


def test_missing_cost_and_disagreement(tmp_path: Path) -> None:
    path = tmp_path / "anexa.xlsx"
    _annex(path, cost=None)
    result = generate([path], (2025,))
    assert result.companies[0].consumption_tep == 25
    assert result.companies[0].measures[2025][0].cost_thousand_lei is None
    assert {item.situation for item in result.exceptions} == {
        "Costul investiției lipsește.",
        "Totalurile lunar și anual diferă.",
    }
    target = tmp_path / "report.xlsx"
    write_report(result, target)
    assert target.is_file()


def test_no_measures_and_bad_file_are_isolated(tmp_path: Path) -> None:
    good = tmp_path / "good.xlsx"
    bad = tmp_path / "bad.xlsx"
    _annex(good, measures=False)
    bad.write_text("bad workbook")
    result = generate([good, bad], (2025,))
    assert len(result.companies) == 2
    assert (
        len(next(company for company in result.companies if company.source == good).measures[2025])
        == 0
    )
    assert next(company for company in result.companies if company.source == bad).status == "EROARE"


def test_name_is_address_uses_alternate_label(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    _annex(path, name_is_address=True, alternate_name=True)
    result = generate([path], (2025,))
    assert result.companies[0].name == "Nume Alternativ"
    assert result.exceptions[0].decision.endswith("Identitate!B1.")


def test_name_is_address_uses_filename_with_exception(tmp_path: Path) -> None:
    path = tmp_path / "anul 2025_Companie din fișier.xlsx"
    _annex(path, name_is_address=True)
    result = generate([path], (2025,))
    assert result.companies[0].name == "Companie din fișier"
    assert result.exceptions[0].decision == "Denumirea a fost preluată din numele fișierului."


def test_monthly_total_follows_label_when_columns_shift(tmp_path: Path) -> None:
    path = tmp_path / "shifted.xlsx"
    _annex(path, monthly=25.0, annual=24.0)
    book = load_workbook(path)
    sheet = book["Date lunare"]
    sheet["J2"] = None
    sheet["M3"] = None
    sheet["K2"] = "[ tep / an ]"
    sheet["N3"] = 25.0
    book.save(path)

    result = generate([path], (2025,))
    assert result.companies[0].monthly_tep == 25.0
    assert result.companies[0].consumption_tep == 25.0


def test_ambiguous_monthly_totals_are_not_selected(tmp_path: Path) -> None:
    path = tmp_path / "ambiguous.xlsx"
    _annex(path)
    book = load_workbook(path)
    sheet = book["Date lunare"]
    sheet["K2"] = "[ tep / an ]"
    sheet["N3"] = 26.0
    book.save(path)

    result = generate([path], (2025,))
    assert result.companies[0].monthly_tep is None
    assert result.companies[0].consumption_tep == 24.0


def test_missing_consumption_does_not_become_zero_in_total(tmp_path: Path) -> None:
    path = tmp_path / "missing-consumption.xlsx"
    _annex(path, monthly=None, annual=None)
    result = generate([path], (2025,))
    assert result.companies[0].consumption_tep is None

    output = write_report(result, tmp_path / "report.xlsx")
    sheet = load_workbook(output, data_only=True)["2025"]
    total = next(row for row in sheet.iter_rows(values_only=True) if row[0] == "TOTAL")
    assert total[3] is None
    assert any(
        item.source == path and "Totalul consumului pentru 2025" in item.situation
        for item in result.exceptions
    )


def test_annex_with_different_reporting_year_is_isolated(tmp_path: Path) -> None:
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    _annex(first)
    _annex(second)
    for path, year in ((first, 2025), (second, 2024)):
        book = load_workbook(path)
        sheet = book["Date anuale"]
        sheet["A5"] = "Datele anului anterior"
        sheet["B5"] = year
        book.save(path)

    result = generate([first, second], (2025,))
    assert len(result.companies) == 2
    excluded = next(company for company in result.companies if company.source == second)
    assert excluded.status == "REVIZUIRE"
    assert not excluded.measures
    assert any(
        item.source == second and "Anul raportării diferă" in item.situation
        for item in result.exceptions
    )
