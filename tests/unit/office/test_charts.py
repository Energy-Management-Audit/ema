"""Synthetic package checks without client documents or Word."""

from io import BytesIO
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
from test_chart_rewrite_package import package

from ema.core.office.blocks import ElementLocator, NativeChart, Paragraph, Prototypes, Retained
from ema.core.office.chart_blocks import build_column_chart_detached, clone_chart_detached
from ema.core.office.charts import (
    Series,
    build_column_chart,
    clone_chart,
    embed_all_data,
    embed_data,
    read_series,
)
from ema.core.office.package import C, P, R, check_standalone, inspect, read_parts, write_parts
from ema.core.office.region import replace_region
from ema.core.office.workbook import extend_formula, formula_cells


def test_embed_link_and_cache(tmp_path):
    source = package(tmp_path)
    result = tmp_path / "embedded.docx"
    embed_data(source, "word/charts/chart1.xml", result)
    report = inspect(result)
    assert check_standalone(result) == []
    assert read_series(result, report.charts[0].part)[0].values == [1.234567, 2.5]
    with ZipFile(result) as archive:
        book = load_workbook(BytesIO(archive.read(report.charts[0].embedded)), data_only=True)
    assert book["Sheet1"]["B1"].value == "Șir"
    assert book["Sheet1"]["B2"].value == 1.234567
    assert b'Extension="xlsx"' in read_parts(result)["[Content_Types].xml"]


def test_embed_all_links(tmp_path):
    source = package(tmp_path)
    out = tmp_path / "all.docx"
    embed_all_data(source, out)
    assert check_standalone(out) == []


def test_replace_region_removes_old_chart_and_embedded_resources(tmp_path):
    source = package(tmp_path)
    embedded = tmp_path / "embedded.docx"
    embed_data(source, "word/charts/chart1.xml", embedded)
    parts = read_parts(embedded)
    root = etree.fromstring(parts["word/document.xml"])
    body = root.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body")
    assert body is not None
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def paragraph(value):
        node = etree.Element(w + "p")
        etree.SubElement(etree.SubElement(node, w + "r"), w + "t").text = value
        return node

    before = paragraph("Before")
    start = paragraph("Start")
    end = paragraph("End")
    body.insert(0, before)
    body.insert(1, start)
    body.insert(3, end)
    parts["word/document.xml"] = etree.tostring(root)
    write_parts(parts, embedded)
    out = tmp_path / "replaced.docx"
    replace_region(
        embedded,
        out,
        ElementLocator(2),
        ElementLocator(4),
        [Paragraph("body", ["Replacement"])],
        Prototypes({"body": paragraph("Prototype")}, 4),
    )
    result = read_parts(out)
    doc = etree.fromstring(result["word/document.xml"])
    texts = [item.text for item in doc.iter(w + "t")]
    assert texts == ["Before", "Start", "Replacement", "End"]
    assert inspect(out).charts == []
    assert not any(part.startswith("word/embeddings/") for part in result)
    assert check_standalone(out) == []

    chart_out = tmp_path / "replaced-chart.docx"
    replace_region(
        embedded,
        chart_out,
        ElementLocator(2),
        ElementLocator(4),
        [
            NativeChart(
                "chart", "word/charts/chart1.xml", [Series("Gaz", ["Ian", "Feb"], [3.5, 4.5])]
            )
        ],
        Prototypes({"chart": body[2]}, 4),
    )
    charts = inspect(chart_out).charts
    assert len(charts) == 1
    assert charts[0].part != "word/charts/chart1.xml"
    assert read_series(chart_out, charts[0].part)[0].values == [3.5, 4.5]
    assert check_standalone(chart_out) == []

    retained_out = tmp_path / "retained-chart.docx"
    replace_region(
        embedded,
        retained_out,
        ElementLocator(2),
        ElementLocator(4),
        [Retained("chart")],
        Prototypes({"chart": body[2]}, 4),
    )
    assert len(inspect(retained_out).charts) == 1
    assert check_standalone(retained_out) == []


def test_clone_and_build_have_unique_parts_relationships_and_ids(tmp_path):
    source = package(tmp_path)
    first = tmp_path / "clone.docx"
    second = tmp_path / "built.docx"
    series = [Series("PV ș ț ă î â", ["Ian", "Feb"], [3.125, 4.5])]
    cloned = clone_chart(source, "word/charts/chart1.xml", series, "Producție ș ț ă î â", first)
    built = build_column_chart(first, "word/charts/chart1.xml", series, "MWh", second)
    report = inspect(second)
    assert {chart.part for chart in report.charts} == {"word/charts/chart1.xml", cloned, built}
    assert len({chart.rel_id for chart in report.charts}) == 3
    parts = read_parts(second)
    doc = etree.fromstring(parts["word/document.xml"])
    ids = [node.get("id") for node in doc.xpath("//*[local-name()='docPr']")]
    assert len(ids) == len(set(ids))
    assert read_series(second, cloned)[0].name == series[0].name
    assert read_series(second, built)[0].values == series[0].values
    assert len({chart.embedded for chart in report.charts if chart.embedded}) == 2
    assert "Producție ș ț ă î â".encode() in parts[cloned]
    assert "word/charts/style2.xml" in parts and "word/charts/style3.xml" in parts
    types = etree.fromstring(parts["[Content_Types].xml"])
    assert {
        item.get("PartName") for item in types if "chartstyle" in item.get("ContentType", "")
    } == {"/word/charts/style1.xml", "/word/charts/style2.xml", "/word/charts/style3.xml"}


def test_detached_charts_reuse_s0_builders_without_inserting(tmp_path):
    source = package(tmp_path)
    root = etree.fromstring(read_parts(source)["word/document.xml"])
    prototype = root.xpath("//*[local-name()='p' and .//*[local-name()='chart']]")[0]
    series = [Series("Gas", ["Jan", "Feb"], [1.0, 2.0])]
    cloned_file = tmp_path / "detached-clone.docx"
    built_file = tmp_path / "detached-built.docx"
    cloned, clone_paragraph = clone_chart_detached(
        source, "word/charts/chart1.xml", series, None, cloned_file, prototype
    )
    built, built_paragraph = build_column_chart_detached(
        cloned_file, "word/charts/chart1.xml", series, "tep", built_file, prototype
    )
    report = inspect(built_file)
    assert {cloned, built}.issubset({chart.part for chart in report.charts})
    assert all(
        chart.embedded and chart.external is None
        for chart in report.charts
        if chart.part in {cloned, built}
    )
    assert len(root.xpath("//*[local-name()='p' and .//*[local-name()='chart']]")) == 1
    assert (
        len(
            etree.fromstring(read_parts(built_file)["word/document.xml"]).xpath(
                "//*[local-name()='p' and .//*[local-name()='chart']]"
            )
        )
        == 1
    )
    assert clone_paragraph.xpath(".//*[local-name()='chart']")
    assert built_paragraph.xpath(".//*[local-name()='chart']")


def test_inspect_planted_orphan_and_missing_workbook(tmp_path):
    source = package(tmp_path)
    parts = read_parts(source)
    parts["word/embeddings/orphan.xlsx"] = b"orphan"
    write_parts(parts, source)
    report = inspect(source)
    assert len(report.external_relationships) == 1
    assert report.charts_without_workbook == ["word/charts/chart1.xml"]
    assert report.orphan_parts == ["word/embeddings/orphan.xlsx"]


def test_many_clones_have_valid_unique_paragraph_ids_and_no_bookmarks(tmp_path):
    current = package(tmp_path)
    series = [Series("PV", ["Ian", "Feb"], [1.0, 2.0])]
    for index in range(64):
        following = tmp_path / f"clone-{index}.docx"
        clone_chart(current, "word/charts/chart1.xml", series, None, following)
        current = following
    doc = etree.fromstring(read_parts(current)["word/document.xml"])
    namespace = "{http://schemas.microsoft.com/office/word/2010/wordml}"
    ids = [
        int(value, 16)
        for paragraph in doc.xpath("//*[local-name()='p']")
        for key in ("paraId", "textId")
        if (value := paragraph.get(namespace + key)) is not None
    ]
    assert len(ids) == 130  # original pair plus 64 cloned pairs
    assert len(ids) == len(set(ids))
    assert all(1 <= value <= 0x7FFFFFFF for value in ids)
    assert len(doc.xpath("//*[local-name()='bookmarkStart']")) == 1
    assert len(doc.xpath("//*[local-name()='bookmarkEnd']")) == 1


def test_formula_extension_is_one_dimensional_and_quotes_sheet():
    for name in ("Consum-Gaz", "Consum Gaz", "Șir Ștefan", "O'Brien"):
        formula = "'" + name.replace("'", "''") + "'!$B$2"
        expanded = extend_formula(formula, 3, "row")
        assert expanded.startswith("'" + name.replace("'", "''") + "'!")
        assert formula_cells(expanded) == (name, [(2, 2), (3, 2), (4, 2)])
    assert formula_cells(extend_formula("Sheet1!$B$2", 3, "column"))[1] == [(2, 2), (2, 3), (2, 4)]


def test_two_series_clone_has_distinct_workbook_cells_and_formulas(tmp_path):
    source = package(tmp_path)
    parts = read_parts(source)
    root = etree.fromstring(parts["word/charts/chart1.xml"])
    bar = root.find(f"{{{C}}}chart/{{{C}}}plotArea/{{{C}}}barChart")
    assert bar is not None
    first = bar.find(f"{{{C}}}ser")
    assert first is not None
    second = etree.fromstring(etree.tostring(first))
    second.find(f"{{{C}}}idx").set("val", "1")
    second.find(f"{{{C}}}order").set("val", "1")
    bar.insert(list(bar).index(first) + 1, second)
    parts["word/charts/chart1.xml"] = etree.tostring(root)
    write_parts(parts, source)
    result = tmp_path / "two-series.docx"
    items = [Series("Gaz", ["Ian", "Feb"], [0.25, 0.5]), Series("PV", ["Mar", "Apr"], [0.75, 1.0])]
    cloned = clone_chart(source, "word/charts/chart1.xml", items, None, result)
    series = read_series(result, cloned)
    assert series[0].refs.values != series[1].refs.values
    assert series[0].refs.categories != series[1].refs.categories
    chart = next(item for item in inspect(result).charts if item.part == cloned)
    with ZipFile(result) as archive:
        book = load_workbook(BytesIO(archive.read(chart.embedded)), data_only=True)
    for item in series:
        for formula, expected in (
            (item.refs.name, [item.name]),
            (item.refs.categories, item.categories),
            (item.refs.values, item.values),
        ):
            sheet, cells = formula_cells(formula)
            assert [book[sheet].cell(row, col).value for row, col in cells] == expected
    assert book["Sheet1"]["B2"].number_format == "0.00"

    same = tmp_path / "same-values.docx"
    same_part = clone_chart(source, "word/charts/chart1.xml", [items[0], items[0]], None, same)
    same_series = read_series(same, same_part)
    assert same_series[0].refs.values != same_series[1].refs.values


def test_percentage_cache_sets_workbook_number_format(tmp_path):
    source = package(tmp_path)
    parts = read_parts(source)
    parts["word/charts/chart1.xml"] = parts["word/charts/chart1.xml"].replace(
        b"<c:formatCode>0.00</c:formatCode>", b"<c:formatCode>0.0%</c:formatCode>"
    )
    write_parts(parts, source)
    result = tmp_path / "percent.docx"
    embed_data(source, "word/charts/chart1.xml", result)
    chart = inspect(result).charts[0]
    with ZipFile(result) as archive:
        book = load_workbook(BytesIO(archive.read(chart.embedded)))
    assert book["Sheet1"]["B2"].number_format == "0.0%"

    parts = read_parts(source)
    root = etree.fromstring(parts["word/charts/chart1.xml"])
    cache = root.find(f".//{{{C}}}numCache")
    assert cache is not None
    format_code = cache.find(f"{{{C}}}formatCode")
    assert format_code is not None
    cache.remove(format_code)
    ser = root.find(f".//{{{C}}}ser")
    assert ser is not None
    etree.SubElement(ser, f"{{{C}}}numFmt", formatCode="0%")
    parts["word/charts/chart1.xml"] = etree.tostring(root)
    write_parts(parts, source)
    embed_data(source, "word/charts/chart1.xml", result)
    with ZipFile(result) as archive:
        book = load_workbook(BytesIO(archive.read(inspect(result).charts[0].embedded)))
    assert book["Sheet1"]["B2"].number_format == "0%"


def test_clone_copies_only_target_drawing_run(tmp_path):
    source = package(tmp_path)
    parts = read_parts(source)
    doc = etree.fromstring(parts["word/document.xml"])
    paragraph = doc.xpath("//*[local-name()='p' and .//*[local-name()='chart']]")[0]
    text_run = etree.Element("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r")
    etree.SubElement(
        text_run, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
    ).text = "Keep me"
    paragraph.insert(0, text_run)
    other_run = etree.fromstring(
        etree.tostring(paragraph.xpath("./*[local-name()='r' and .//*[local-name()='chart']]")[0])
    )
    other_chart = other_run.xpath(".//*[local-name()='chart']")[0]
    other_chart.set(f"{{{R}}}id", "rId998")
    paragraph.append(other_run)
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    etree.SubElement(
        rels, f"{{{P}}}Relationship", Id="rId998", Type=f"{R}/chart", Target="charts/chart2.xml"
    )
    parts["word/charts/chart2.xml"] = parts["word/charts/chart1.xml"]
    parts["word/_rels/document.xml.rels"] = etree.tostring(rels)
    parts["word/document.xml"] = etree.tostring(doc)
    write_parts(parts, source)
    result = tmp_path / "one-drawing.docx"
    cloned = clone_chart(
        source, "word/charts/chart1.xml", [Series("PV", ["Ian", "Feb"], [1, 2])], None, result
    )
    final = etree.fromstring(read_parts(result)["word/document.xml"])
    paragraphs = final.xpath("//*[local-name()='p' and .//*[local-name()='chart']]")
    assert len(paragraphs) == 2
    assert [
        node.get(f"{{{R}}}id") for node in paragraphs[0].xpath(".//*[local-name()='chart']")
    ] == ["rId999", "rId998"]
    assert paragraphs[0].xpath(".//*[local-name()='t']/text()") == ["Keep me"]
    assert len(paragraphs[1].xpath(".//*[local-name()='chart']")) == 1
    assert paragraphs[1].xpath(".//*[local-name()='t']") == []
    assert next(item for item in inspect(result).charts if item.part == cloned).rel_id != "rId998"
