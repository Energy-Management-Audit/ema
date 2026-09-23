"""Single-slot legacy intake without launching Word."""

from types import SimpleNamespace

import pytest
from docx import Document
from typer.testing import CliRunner

from ema.cli import _app
from ema.core.config import Settings
from ema.core.intake import ItemOutcome, intake_file, intake_legacy, intake_stage
from ema.core.jobs import StageContext, create_job, status
from ema.core.office import convert as conversion_module
from ema.core.office.convert import ConversionFailed, convert_doc
from ema.core.office.errors import OfficeError
from ema.core.office.sniff import FileKind, Sniffed
from ema.core.office.word import DocText
from ema.core.workspace import Workspace
from ema.core.workspace.conversion import active_version


def _case(tmp_path, monkeypatch):
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    source = tmp_path / "original.doc"
    source.write_bytes(b"synthetic DOC")
    sha = ws.add_file("synthetic", source)
    original = ws.set_slot(job, "dossier/0001", sha)
    monkeypatch.setattr(
        conversion_module,
        "sniff",
        lambda _path: Sniffed(FileKind.DOC, False, "doc"),
    )
    monkeypatch.setattr(
        "ema.core.intake.sniff",
        lambda _path: Sniffed(FileKind.DOC, False, "doc"),
    )
    return ws, job, original


def test_unavailable_word_flags_one_file(tmp_path, monkeypatch):
    ws, job, original = _case(tmp_path, monkeypatch)
    monkeypatch.setattr(
        "ema.core.intake.load_settings",
        lambda _ws: Settings(word_path=tmp_path / "missing.app"),
    )
    ctx = StageContext(ws, job, "synthetic", "intake")
    result = intake_file(ctx, original.slot)
    assert result.status == "needs_conversion"
    assert result.error_code == "needs_conversion"
    assert "save it as .docx" in (result.detail or "")
    assert ws.list_versions(job, original.slot) == [original]
    outcome = intake_stage(ctx, original.slot)
    assert outcome.item_failures == []
    assert outcome.warnings
    with ws.connect() as db:
        log_path = ws.job_path(db, job) / "log.jsonl"
    assert "needs_conversion" in log_path.read_text(encoding="utf-8")


def test_conversion_adds_version_and_rerun_is_noop(tmp_path, monkeypatch):
    ws, job, original = _case(tmp_path, monkeypatch)
    app = tmp_path / "Word.app"
    app.mkdir()
    monkeypatch.setattr("ema.core.intake.load_settings", lambda _ws: Settings(word_path=app))
    monkeypatch.setattr(conversion_module, "sys", SimpleNamespace(platform="darwin"))

    class FakeWord:
        def __init__(self, app, timeout_s):
            assert app == tmp_path / "Word.app"
            assert timeout_s == 120

        def convert_doc(self, _source, target):
            doc = Document()
            doc.add_paragraph("Alpha beta gamma")
            doc.save(target)

        def doc_text(self, _source):
            return DocText("Alpha beta gamma", 0)

    monkeypatch.setattr(conversion_module, "WordMac", FakeWord)
    ctx = StageContext(ws, job, "synthetic", "intake")
    result = intake_file(ctx, original.slot)
    versions = ws.list_versions(job, original.slot)
    assert result.status == "converted"
    assert result.original_words == result.converted_words == 3
    assert len(versions) == 2
    assert versions[0] == original
    assert versions[1].converted_from == original.file_sha
    assert versions[1].origin == "converted"
    assert intake_file(ctx, original.slot).status == "already_converted"
    assert ws.list_versions(job, original.slot) == versions


def test_new_upload_during_conversion_supersedes_result(tmp_path, monkeypatch):
    ws, job, original = _case(tmp_path, monkeypatch)
    app = tmp_path / "Word.app"
    app.mkdir()
    monkeypatch.setattr("ema.core.intake.load_settings", lambda _ws: Settings(word_path=app))
    monkeypatch.setattr(conversion_module, "sys", SimpleNamespace(platform="darwin"))
    replacement = tmp_path / "new.doc"
    replacement.write_bytes(b"new DOC")
    new_sha = ws.add_file("synthetic", replacement)

    class FakeWord:
        def __init__(self, app, timeout_s):
            pass

        def convert_doc(self, _source, target):
            doc = Document()
            doc.add_paragraph("Alpha beta gamma")
            doc.save(target)
            ws.set_slot(job, original.slot, new_sha)

        def doc_text(self, _source):
            return DocText("Alpha beta gamma", 0)

    monkeypatch.setattr(conversion_module, "WordMac", FakeWord)
    result = intake_file(StageContext(ws, job, "synthetic", "intake"), original.slot)
    assert result.status == "superseded"
    versions = ws.list_versions(job, original.slot)
    assert len(versions) == 2
    assert active_version(ws, job, original.slot) == versions[1]
    with ws.connect() as db:
        assert db.execute("SELECT COUNT(*) FROM files").fetchone()[0] == 2
    assert len(list((ws.root / "clients/synthetic/files").iterdir())) == 2


def test_superseded_result_is_stage_warning_with_rerun_step(tmp_path, monkeypatch):
    ws, job, original = _case(tmp_path, monkeypatch)
    monkeypatch.setattr(
        "ema.core.intake.intake_file",
        lambda _ctx, _slot: ItemOutcome(
            original.slot,
            original.version,
            original.file_sha,
            FileKind.DOC,
            "superseded",
            warning=(
                "Fișierul a fost înlocuit în timpul conversiei; rulați din nou preluarea "
                "pentru noua versiune (R21)."
            ),
        ),
    )
    outcome = intake_stage(StageContext(ws, job, "synthetic", "intake"), original.slot)
    assert outcome.item_failures == []
    assert len(outcome.warnings) == 1
    assert "R21" in outcome.warnings[0]
    assert "rulați din nou" in outcome.warnings[0]


def test_conversion_timeout_preserves_original(tmp_path, monkeypatch):
    ws, job, original = _case(tmp_path, monkeypatch)
    app = tmp_path / "Word.app"
    app.mkdir()

    class TimeoutWord:
        def __init__(self, app, timeout_s):
            pass

        def convert_doc(self, _source, _target):
            raise OfficeError("word_timeout", "hung twice")

    monkeypatch.setattr(conversion_module, "WordMac", TimeoutWord)
    monkeypatch.setattr(conversion_module, "sys", SimpleNamespace(platform="darwin"))
    with pytest.raises(ConversionFailed) as error:
        convert_doc(ws, job, original.slot, original, Settings(word_path=app))
    assert error.value.code == "convert_timeout"
    assert "hung twice" in error.value.detail
    assert ws.list_versions(job, original.slot) == [original]


def test_cancelled_before_conversion(tmp_path, monkeypatch):
    ws, job, original = _case(tmp_path, monkeypatch)
    ctx = StageContext(ws, job, "synthetic", "intake")
    monkeypatch.setattr(ctx, "cancelled", lambda: True)
    assert intake_file(ctx, original.slot).status == "cancelled"
    assert ws.list_versions(job, original.slot) == [original]


def test_cli_records_html_saved_as_xls(tmp_path, monkeypatch):
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    source = tmp_path / "misnamed.xls"
    source.write_text("<!doctype html><html><body><table></table></body></html>")
    sha = ws.add_file("synthetic", source)
    ws.set_slot(job, "meters/0001", sha)
    monkeypatch.setenv("EMA_WORKSPACE", str(ws.root))

    response = CliRunner().invoke(_app, ["intake", job, "meters"])

    assert response.exit_code == 0, response.output
    assert '"kind": "html"' in response.output
    assert '"status": "html_as_xls"' in response.output
    assert '"slot": "meters/0001"' in response.output
    assert status(ws, job).runs[-1]["state"] == "ready"


def test_collection_reports_each_item_and_keeps_failed_file_local(tmp_path, monkeypatch):
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    for number in (1, 2, 3):
        source = tmp_path / f"item-{number}.txt"
        source.write_text(f"item {number}")
        ws.set_slot(job, f"dossier/{number:04}", ws.add_file("synthetic", source))
    original_intake = intake_file

    def one_failure(ctx, slot):
        if slot.endswith("0002"):
            raise OSError("unreadable source")
        return original_intake(ctx, slot)

    monkeypatch.setattr("ema.core.intake.intake_file", one_failure)
    results = []
    ctx = StageContext(ws, job, "synthetic", "intake")
    outcome = intake_legacy(ctx, "dossier", results.append)
    assert [item.status for item in results] == ["detected", "failed", "detected"]
    assert len(outcome.item_failures) == 1
    assert "unreadable source" in outcome.item_failures[0]
    assert set(ctx.inputs) == {f"slot:dossier/{number:04}" for number in (1, 2, 3)}


def test_collection_stops_before_next_file_on_cancellation(tmp_path, monkeypatch):
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    source = tmp_path / "item.txt"
    source.write_text("item")
    sha = ws.add_file("synthetic", source)
    for number in (1, 2):
        ws.set_slot(job, f"dossier/{number:04}", sha)
    ctx = StageContext(ws, job, "synthetic", "intake")
    checks = iter((False, False, True))
    monkeypatch.setattr(ctx, "cancelled", lambda: next(checks))
    results = []
    intake_legacy(ctx, "dossier", results.append)
    assert [item.slot for item in results] == ["dossier/0001"]


def test_html_with_unrelated_extension_is_detected(tmp_path):
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    source = tmp_path / "misnamed.pdf"
    source.write_text("<!doctype html><html><body>hello</body></html>")
    sha = ws.add_file("synthetic", source)
    ws.set_slot(job, "meters/0001", sha)
    result = intake_file(StageContext(ws, job, "synthetic", "intake"), "meters/0001")
    assert result.status == "detected"
    assert result.kind == FileKind.HTML
    assert "HTML saved as .pdf" in (result.warning or "")
