"""Mapped text spans preserve surrounding Word runs and isolate red missing values."""

from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn

from ema.core.office.run_range import TextSpan, replace_spans, visible_text


def test_replaces_mapped_span_across_mixed_runs() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Client: ")
    paragraph.add_run("SYNTHETIC").bold = True
    paragraph.add_run("; year 2025")
    replace_spans(paragraph._p, (TextSpan(8, 17, "NEW COMPANY"),))
    assert visible_text(paragraph._p) == "Client: NEW COMPANY; year 2025"
    assert paragraph.runs[1].bold


def test_missing_value_is_red_without_colouring_prefix_or_suffix() -> None:
    document = Document()
    paragraph = document.add_paragraph("Number: 12345; filed")
    replace_spans(paragraph._p, (TextSpan(8, 13, "n.d.", missing=True),))
    assert visible_text(paragraph._p) == "Number: n.d.; filed"
    runs = list(paragraph._p.iter(qn("w:r")))
    assert len(runs) == 3
    assert runs[0].find(f"{qn('w:rPr')}/{qn('w:color')}") is None
    assert runs[1].find(f"{qn('w:rPr')}/{qn('w:color')}").get(qn("w:val")) == "FF0000"
    assert runs[2].find(f"{qn('w:rPr')}/{qn('w:color')}") is None


def test_multiple_spans_in_one_original_run() -> None:
    document = Document()
    paragraph = document.add_paragraph("A 111 and 222 Z")
    replace_spans(paragraph._p, (TextSpan(2, 5, "first"), TextSpan(10, 13, "n.d.", True)))
    assert visible_text(paragraph._p) == "A first and n.d. Z"
