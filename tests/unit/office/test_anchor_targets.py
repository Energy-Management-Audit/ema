"""Synthetic cell, row and relationship anchors without client documents."""

from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from ema.core.office.anchor_targets import (
    chart_target,
    find_cell,
    find_row,
    hyperlink_target,
    stamp_cell,
    stamp_row,
)
from ema.core.office.anchors import stamp
from ema.core.office.package import P, R, encoded


def test_cell_and_row_targets() -> None:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    root = document.element
    stamp_cell(table.cell(0, 1)._tc, "identity", 1)
    stamp_row(table.rows[1]._tr, "measure_row", 2)
    assert find_cell([root], "identity") is table.cell(0, 1)._tc
    assert find_row([root], "measure_row") is table.rows[1]._tr


def test_chart_and_hyperlink_resolve_from_bookmark_not_position() -> None:
    document = Document()
    chart = document.add_paragraph()
    stamp(chart._p, "annual_chart", 1)
    run = OxmlElement("w:r")
    drawing = OxmlElement("w:drawing")
    chart_node = etree.SubElement(
        drawing, "{http://schemas.openxmlformats.org/drawingml/2006/chart}chart"
    )
    chart_node.set(f"{{{R}}}id", "rIdChart")
    run.append(drawing)
    chart._p.append(run)
    link = document.add_paragraph()
    stamp(link._p, "website", 2)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), "rIdLink")
    link._p.append(hyperlink)
    rels = etree.Element(f"{{{P}}}Relationships")
    etree.SubElement(
        rels, f"{{{P}}}Relationship", Id="rIdChart", Type=f"{R}/chart", Target="charts/chart7.xml"
    )
    etree.SubElement(
        rels,
        f"{{{P}}}Relationship",
        Id="rIdLink",
        Type=f"{R}/hyperlink",
        Target="https://example.test",
        TargetMode="External",
    )
    parts = {
        "word/document.xml": encoded(document.element),
        "word/_rels/document.xml.rels": encoded(rels),
    }
    target = chart_target(parts, "annual_chart")
    assert (target.owner, target.relationship_id, target.part) == (
        "word/document.xml",
        "rIdChart",
        "word/charts/chart7.xml",
    )
    assert hyperlink_target(parts, "website").relationship_id == "rIdLink"
