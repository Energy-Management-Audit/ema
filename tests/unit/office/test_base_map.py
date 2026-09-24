"""A base version requires complete classification before it can be stamped."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from ema.core.office.anchors import find
from ema.core.office.base_map import classify, inventory, load, save, stamp_base, validate
from ema.core.office.package import read_parts, xml


def test_complete_map_stamps_only_variable_slots(tmp_path: Path) -> None:
    base = tmp_path / "base.docx"
    document = Document()
    document.add_paragraph("Fixed contract text")
    document.add_paragraph("Synthetic client")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Synthetic measure"
    document.save(base)
    items = inventory(base)
    assert {item.kind for item in items} >= {"paragraph", "cell", "row", "relationship"}
    variable_paragraph = next(
        item for item in items if item.kind == "paragraph" and item.selector.endswith("}p[2]")
    )
    variable_row = next(item for item in items if item.kind == "row")
    mapping = classify(base, {variable_paragraph.id: "client", variable_row.id: "measure"})
    assert len(mapping.elements) == len(items)
    manifest = tmp_path / "map.json"
    save(mapping, manifest)
    assert load(manifest) == mapping
    stamped = tmp_path / "stamped.docx"
    preview = tmp_path / "preview.docx"
    stamp_base(base, mapping, stamped)
    stamp_base(base, mapping, preview, preview=True)
    roots = [xml(read_parts(stamped), "word/document.xml")]
    assert find(roots, "client") is not None
    assert find(roots, "measure") is not None
    assert "[client]" not in xml(read_parts(stamped), "word/document.xml").xpath("string()")
    assert "[client]" in xml(read_parts(preview), "word/document.xml").xpath("string()")
    changed = Document(base)
    changed.add_paragraph("New unmapped item")
    changed.save(base)
    with pytest.raises(ValueError, match="hash"):
        validate(base, mapping)


def test_package_target_requires_bookmark_companion(tmp_path: Path) -> None:
    base = tmp_path / "base.docx"
    document = Document()
    document.add_paragraph("Synthetic")
    document.save(base)
    relation = next(item for item in inventory(base) if item.kind == "relationship")
    mapping = classify(base, {relation.id: "orphan"})
    with pytest.raises(ValueError, match="bookmark companion"):
        validate(base, mapping)
