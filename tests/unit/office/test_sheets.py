"""The same label semantics on both workbook formats."""

from pathlib import Path

import pytest
import xlwt
from openpyxl import Workbook

from ema.core.office.errors import OfficeError
from ema.core.office.sheets import CellRange, CellRef, open_book


@pytest.fixture(params=["xls", "xlsx"])
def workbook(tmp_path: Path, request: pytest.FixtureRequest) -> Path:
    suffix = request.param
    path = tmp_path / f"shifted.{suffix}"
    if suffix == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "Date"
        sheet["B3"] = "  C.U.I.  "
        sheet["C3"] = "12345"
        sheet["B7"] = "Știință și Țevi"
        sheet["C7"] = 12.5
        sheet["B11"] = "CUI"
        sheet["D11"] = "cui"
        book.save(path)
    else:
        book = xlwt.Workbook()
        sheet = book.add_sheet("Date")
        for row, col, value in (
            (2, 1, "  C.U.I.  "),
            (2, 2, "12345"),
            (6, 1, "Știință și Țevi"),
            (6, 2, 12.5),
            (10, 1, "CUI"),
            (10, 3, "cui"),
        ):
            sheet.write(row, col, value)
        book.save(str(path))
    return path


def test_alias_diacritics_relative_reads_and_provenance(workbook: Path) -> None:
    book = open_book(workbook)
    sheet = book.sheet("Date")
    label = sheet.find_label(["unavailable", " ştiinţă  şi ţevi "])
    assert label.a1 == "Date!B7"
    assert sheet.read_right(label).value == 12.5
    assert sheet.read_right(label).ref.a1 == "Date!C7"
    assert sheet.read_below(label).value is None
    assert [item.ref.a1 for item in sheet.read_row(label, 2)] == ["Date!B7", "Date!C7"]
    assert sheet.read_block(label, 1, 2)[0][1].value == 12.5
    book.close()


def test_missing_duplicate_and_region(workbook: Path) -> None:
    book = open_book(workbook)
    sheet = book.sheet("Date")
    with pytest.raises(OfficeError) as error:
        sheet.find_label(["nothing"])
    assert error.value.code == "label_missing"
    with pytest.raises(OfficeError) as error:
        sheet.find_label(["CUI"])
    assert error.value.code == "label_ambiguous"
    assert "Date!B11" in error.value.detail and "Date!D11" in error.value.detail
    assert (
        sheet.find_label(
            ["CUI"], within=CellRange(CellRef("Date", 9, 1), CellRef("Date", 12, 2))
        ).a1
        == "Date!B11"
    )
    book.close()


def test_content_sniff_rejects_html(tmp_path: Path) -> None:
    path = tmp_path / "html.xls"
    path.write_text("<html>not an Excel workbook</html>")
    with pytest.raises(OfficeError) as error:
        open_book(path)
    assert error.value.code == "unsupported_format"


def test_xlsx_formula_without_cache_is_not_blank(tmp_path: Path) -> None:
    path = tmp_path / "formula.xlsx"
    book = Workbook()
    book.active.title = "Calc"
    book.active["A1"] = "=1+1"
    book.active["B1"] = ""
    book.save(path)
    loaded = open_book(path)
    sheet = loaded.sheet("Calc")
    with pytest.raises(OfficeError) as error:
        sheet.value(1, 1)
    assert error.value.code == "formula_uncached"
    assert "Calc!A1" in error.value.detail
    assert sheet.value(1, 2).value is None
    loaded.close()


def test_read_right_starts_after_merged_label(tmp_path: Path) -> None:
    path = tmp_path / "merged.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Calc"
    sheet.merge_cells("D2:E3")
    sheet["D2"] = "Label"
    sheet["F2"] = 42
    sheet["D4"] = 43
    book.save(path)
    loaded = open_book(path)
    loaded_sheet = loaded.sheet("Calc")
    result = loaded_sheet.read_right(CellRef("Calc", 2, 4))
    assert result.value == 42
    assert result.ref.a1 == "Calc!F2"
    below = loaded_sheet.read_below(CellRef("Calc", 2, 4))
    assert below.value == 43
    assert below.ref.a1 == "Calc!D4"
    loaded.close()


def test_xlsx_error_cell_is_not_blank(tmp_path: Path) -> None:
    path = tmp_path / "error.xlsx"
    book = Workbook()
    book.active.title = "Calc"
    book.active["A1"] = "#DIV/0!"
    book.active["A1"].data_type = "e"
    book.save(path)
    loaded = open_book(path)
    with pytest.raises(OfficeError) as error:
        loaded.sheet("Calc").value(1, 1)
    assert error.value.code == "cell_error"
    assert "Calc!A1" in error.value.detail
    loaded.close()


def test_xlsx_hyperlink_target_is_preserved(tmp_path: Path) -> None:
    path = tmp_path / "website.xlsx"
    book = Workbook()
    book.active.title = "Date"
    book.active["B2"] = "Company website"
    book.active["B2"].hyperlink = "https://example.test/contact"
    book.save(path)

    loaded = open_book(path)
    try:
        cell = loaded.sheet("Date").value(2, 2)
        assert cell.value == "Company website"
        assert cell.hyperlink == "https://example.test/contact"
    finally:
        loaded.close()
