"""The imported chapter-four chart supplies style only and leaves no source link."""

from docx import Document
from lxml import etree
from test_chart_rewrite_package import package

from ema.audit.chapter_four_charts import chart_blocks
from ema.core.office.blocks import ElementLocator, Missing, NativeChart, Prototypes
from ema.core.office.chart_blocks import chart_caption_prototype, import_chart_style
from ema.core.office.chart_series import Series
from ema.core.office.package import (
    C,
    check_standalone,
    inspect,
    read_parts,
    rels_path,
    write_parts,
    xml,
)
from ema.core.office.region import replace_region
from ema.energy_data.carriers import Carrier
from ema.energy_data.factors import FACTORS_2026
from ema.energy_data.model import CarrierSeries, EnergyDataset, Reading

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def test_import_style_and_retire_it_after_region_replacement(tmp_path) -> None:
    source = package(tmp_path)
    source_parts = read_parts(source)
    source_root = xml(source_parts, "word/document.xml")
    source_body = source_root.find(W + "body")
    assert source_body is not None
    caption = etree.Element(W + "p")
    etree.SubElement(etree.SubElement(caption, W + "r"), W + "t").text = "Fig. nr. 4.1 Test"
    source_body.insert(1, caption)
    source_parts["word/document.xml"] = etree.tostring(source_root)
    chart = xml(source_parts, "word/charts/chart1.xml")
    chart_body = chart.find(f"{{{C}}}chart")
    assert chart_body is not None
    chart_body.insert(0, etree.Element(f"{{{C}}}title"))
    source_parts["word/charts/chart1.xml"] = etree.tostring(chart)
    write_parts(source_parts, source)

    target = tmp_path / "target.docx"
    document = Document()
    document.add_paragraph("Start")
    document.add_paragraph("Old")
    document.add_paragraph("End")
    document.save(target)
    parts = read_parts(target)
    style, drawing = import_chart_style(source, parts)
    caption_text = chart_caption_prototype(source).xpath(".//*[local-name()='t']")[0].text
    assert caption_text == "Fig. nr. 4.1 Test"
    assert xml(parts, style).find(f"{{{C}}}externalData") is None
    assert xml(parts, style).find(f"{{{C}}}chart/{{{C}}}title") is None
    assert all(rel.get("TargetMode") != "External" for rel in xml(parts, rels_path(style)))
    assert any(rel.get("Type", "").endswith("/chartStyle") for rel in xml(parts, rels_path(style)))
    root = xml(parts, "word/document.xml")
    body = root.find(W + "body")
    assert body is not None
    body.insert(1, drawing)
    parts["word/document.xml"] = etree.tostring(root)
    write_parts(parts, target)
    out = tmp_path / "result.docx"
    replace_region(
        target,
        out,
        ElementLocator(1),
        ElementLocator(4),
        [NativeChart("chart", style, [Series("Consum", ["2024", "2025"], [3.5, None])])],
        Prototypes({"chart": drawing}, 4),
    )
    result = read_parts(out)
    assert style not in result
    assert rels_path(style) not in result
    assert len(inspect(out).charts) == 1
    assert check_standalone(out) == []


def test_missing_chart_caption_renders_from_mixed_authored_runs(tmp_path) -> None:
    source = package(tmp_path)
    parts = read_parts(source)
    root = xml(parts, "word/document.xml")
    body = root.find(W + "body")
    assert body is not None
    caption = etree.Element(W + "p")
    for index, text in enumerate(("Fig. nr. 4.1 ", "Authored subject", " ", "Client")):
        run = etree.SubElement(caption, W + "r")
        props = etree.SubElement(run, W + "rPr")
        etree.SubElement(props, W + "sz").set(W + "val", "24")
        if index != 2:
            etree.SubElement(props, W + "i")
        if index == 1:
            etree.SubElement(props, W + "iCs")
        etree.SubElement(run, W + "t").text = text
    body.insert(1, caption)
    parts["word/document.xml"] = etree.tostring(root)
    write_parts(parts, source)
    target = tmp_path / "target.docx"
    document = Document()
    for text in ("Start", "Old", "End"):
        document.add_paragraph(text)
    document.save(target)
    dataset = EnergyDataset(
        (2025,),
        {Carrier.electricity_grid: {2025: CarrierSeries(annual=Reading(0, "MWh"))}},
        {},
        {},
    )
    missing, _ = chart_blocks("ch4.electricitate", dataset, FACTORS_2026, "Client")
    assert all(isinstance(block, Missing) and block.proto == "body" for block in missing)
    body = etree.Element(W + "p")
    run = etree.SubElement(body, W + "r")
    etree.SubElement(etree.SubElement(run, W + "rPr"), W + "sz").set(W + "val", "24")
    etree.SubElement(run, W + "t").text = "Body prototype"
    out = tmp_path / "missing.docx"
    replace_region(
        target,
        out,
        ElementLocator(1),
        ElementLocator(3),
        missing,
        Prototypes({"body": body, "chart_caption": chart_caption_prototype(source)}, 4),
    )
    paragraphs = Document(out).paragraphs[1:-1]
    assert [paragraph.text for paragraph in paragraphs] == [block.text for block in missing]
    for paragraph in paragraphs:
        assert all(run.font.size.pt == 12 for run in paragraph.runs)
        assert all(str(run.font.color.rgb) == "FF0000" for run in paragraph.runs)
