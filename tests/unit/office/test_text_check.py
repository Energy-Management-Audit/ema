"""Conversion fidelity checks use content, not docx package size."""

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from ema.core.office.text_check import check_text
from ema.core.office.word import DocText


def test_matching_words_and_table_count(tmp_path):
    target = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("Alpha beta gamma")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Delta"
    doc.save(target)
    result = check_text(DocText("Alpha beta gamma Delta", 1), target)
    assert result.within_tolerance
    assert result.original_words == result.converted_words == 4
    assert result.original_tables == result.converted_tables == 1


def test_missing_table_is_warning(tmp_path):
    target = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("Alpha beta gamma")
    doc.save(target)
    result = check_text(DocText("Alpha beta gamma", 1), target)
    assert not result.within_tolerance
    assert "tables 1/0" in (result.warning or "")


def test_text_box_is_reported_as_unchecked(tmp_path):
    target = tmp_path / "shapes.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph._p.append(
        parse_xml(
            f'<w:pict {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml">'
            "<v:shape><v:textbox><w:txbxContent>"
            "<w:p><w:r><w:t>Flow chart words</w:t></w:r></w:p>"
            "</w:txbxContent></v:textbox></v:shape></w:pict>"
        )
    )
    doc.save(target)

    result = check_text(DocText("", 0), target)

    assert result.original_words == result.converted_words == 0
    assert result.shape_words == 3
    assert not result.within_tolerance
    assert result.warning == "Fidelity check not applicable (text in shapes): compare in Word."


def test_dropped_paragraph_reports_its_span(tmp_path):
    target = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("Alpha beta gamma delta epsilon.")
    doc.add_paragraph("Closing sentence remains here.")
    doc.save(target)
    original = DocText(
        "Alpha beta gamma delta epsilon. Missing short section of five words. "
        "Closing sentence remains here.",
        0,
    )
    result = check_text(original, target)
    assert not result.within_tolerance
    assert "missing short section of five" in (result.warning or "")


def test_equal_count_substitution_is_reported(tmp_path):
    target = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("Alpha beta gamma delta epsilon zeta eta ending")
    doc.save(target)
    result = check_text(DocText("Alpha beta one two three four five ending", 0), target)
    assert result.original_words == result.converted_words
    assert not result.within_tolerance
    assert "one two three four five" in (result.warning or "")


def test_two_word_deletion_in_short_document_is_reported(tmp_path):
    target = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("one two three four five six seven eight nine ten")
    doc.save(target)
    result = check_text(
        DocText("one two three four five six seven eight nine ten eleven twelve", 0), target
    )
    assert not result.within_tolerance
    assert "eleven twelve" in (result.warning or "")


def test_word_split_across_formatting_runs_counts_once(tmp_path):
    target = tmp_path / "out.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    for text in ("inter", "na", "tional"):
        paragraph.add_run(text)
    doc.save(target)
    result = check_text(DocText("international", 0), target)
    assert result.within_tolerance
    assert result.original_words == result.converted_words == 1


def test_hyphenation_and_whitespace_reflow_pass(tmp_path):
    target = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("International   energy\tusage")
    doc.save(target)
    result = check_text(DocText("Inter-\nnational energy\n usage", 0), target)
    assert result.within_tolerance
