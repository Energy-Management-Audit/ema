"""A bookmarked static pie becomes a native chart without orphaned image data."""

import base64
from io import BytesIO
from pathlib import Path

from docx import Document

from ema.core.office.anchors import stamp
from ema.core.office.package import check_standalone, read_parts, write_parts
from ema.core.office.pie_replace import remove_pie_picture, replace_pie_picture

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9aH6kAAAAASUVORK5CYII="
)


def _picture(path: Path) -> Path:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run().add_picture(BytesIO(_PNG))
    stamp(paragraph._p, "pie_test", 1)
    document.save(path)
    return path


def test_replace_picture_keeps_native_chart_and_removes_unused_image(tmp_path: Path) -> None:
    parts = read_parts(_picture(tmp_path / "source.docx"))
    assert any(name.startswith("word/media/") for name in parts)
    chart = replace_pie_picture(parts, "pie_test", "pv", ("Grid", "Solar"), (0.8, 0.2))
    assert chart.startswith("word/charts/chart")
    assert b"pie3DChart" in parts[chart]
    assert not any(name.startswith("word/media/") for name in parts)
    assert any(name.startswith("word/embeddings/") for name in parts)
    output = tmp_path / "native.docx"
    write_parts(parts, output)
    assert check_standalone(output) == []


def test_remove_missing_picture_cleans_relationship_and_paragraph(tmp_path: Path) -> None:
    parts = read_parts(_picture(tmp_path / "source.docx"))
    remove_pie_picture(parts, "pie_test")
    assert b"_ema_pie_test" not in parts["word/document.xml"]
    assert not any(name.startswith("word/media/") for name in parts)
