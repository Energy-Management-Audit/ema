"""Synthetic document blocks exercise style preservation and numbering."""

from pathlib import Path

import pytest
from docx import Document
from docx.shared import RGBColor
from lxml import etree

from ema.core.office.blocks import (
    BulletList,
    Caption,
    ElementLocator,
    Missing,
    Num,
    Paragraph,
    Prototypes,
    Ref,
    Table,
    render,
)
from ema.core.office.errors import OfficeError
from ema.core.office.numbers_ro import format_number
from ema.core.office.package import read_parts, write_parts

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _body(path: Path) -> etree._Element:
    root = etree.fromstring(read_parts(path)["word/document.xml"])
    body = root.find(W + "body")
    assert body is not None
    return body


def _text(node: etree._Element) -> str:
    return "".join(text.text or "" for text in node.iter(W + "t"))


def _source(tmp_path: Path) -> tuple[Path, Prototypes, ElementLocator]:
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Tabel nr. 4.6 Exemplu")
    doc.add_paragraph("Fig. nr. 4.6 Exemplu")
    doc.add_heading("4.3.2 Test", level=2)
    body = doc.add_paragraph()
    body.add_run("Prima ").bold = True
    body.add_run("parte").italic = True
    caption = doc.add_paragraph("Tabel nr. 4.7 Test")
    figure_caption = doc.add_paragraph("Fig. nr. 4.7 Test")
    item = doc.add_paragraph("Un punct", style="List Bullet")
    missing = doc.add_paragraph("date indisponibile")
    missing.runs[0].font.color.rgb = RGBColor(255, 0, 0)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "An"
    table.cell(0, 1).text = "tep"
    table.cell(1, 0).text = "2023"
    table.cell(1, 1).text = "5,55"
    doc.save(path)
    prototypes = Prototypes(
        {
            "body": body._p,
            "tabcap": caption._p,
            "figcap": figure_caption._p,
            "item": item._p,
            "missing": missing._p,
            "table": table._tbl,
        },
        chapter=4,
    )
    xml_body = _body(path)
    children = list(xml_body)
    for node in children[3:-1]:
        xml_body.remove(node)
    parts = read_parts(path)
    parts["word/document.xml"] = etree.tostring(xml_body.getroottree(), encoding="UTF-8")
    write_parts(parts, path)
    return path, prototypes, ElementLocator(3)


@pytest.mark.parametrize("row_count", [0, 1, 5])
def test_render_table_rows_styles_and_refs(tmp_path: Path, row_count: int) -> None:
    source, prototypes, locator = _source(tmp_path)
    out = tmp_path / "out.docx"
    rows = [[[str(2023 + i)], [Num(1234.5 + i, 2, "tep", fact=f"f{i}")]] for i in range(row_count)]
    report = render(
        source,
        out,
        locator,
        [
            Paragraph("body", ["Conform tabelului numărul ", Ref("tab", "total")]),
            Caption("tabcap", "tab", "total", ["Tabel nr. ", Ref("tab", "total"), " Test"]),
            Table("table", rows),
            Caption("figcap", "fig", "gas", ["Fig. nr. ", Ref("fig", "gas"), " Test"]),
            BulletList("item", [["Punct A"], ["Punct B"]]),
        ],
        prototypes,
    )
    body = _body(out)
    assert report.numbers[0].number == "4.7"
    assert report.numbers[1].number == "4.7"
    assert "Conform tabelului numărul 4.7" in _text(body)
    table = next(node for node in body if node.tag == W + "tbl")
    assert len(table.findall(W + "tr")) == row_count + 1
    assert _text(table.findall(W + "tr")[0]) == "Antep"
    assert len(report.values) == row_count
    assert all(value.fact for value in report.values)
    paragraph = list(body)[3]
    assert len(list(paragraph.iter(W + "rPr"))) == 2
    assert [_text(run) for run in paragraph.findall(W + "r")] == [
        "Conform tabelului numărul ",
        "4.7",
    ]


def test_mixed_run_replacement_requires_slot_alignment(tmp_path: Path) -> None:
    source, prototypes, locator = _source(tmp_path)
    out = tmp_path / "out.docx"
    with pytest.raises(OfficeError) as error:
        render(source, out, locator, [Paragraph("body", ["Rezultat kg"])], prototypes)
    assert error.value.code == "mixed_run_replacement"
    assert "Prima " in error.value.detail
    assert "parte" in error.value.detail


def test_identical_run_properties_form_one_slot(tmp_path: Path) -> None:
    source, prototypes, locator = _source(tmp_path)
    body = prototypes.elements["body"]
    runs = body.findall(W + "r")
    first_props = runs[0].find(W + "rPr")
    second_props = runs[1].find(W + "rPr")
    assert first_props is not None and second_props is not None
    second_props.clear()
    for child in first_props:
        second_props.append(etree.fromstring(etree.tostring(child)))
    out = tmp_path / "out.docx"
    render(source, out, locator, [Paragraph("body", ["Text mai lung"])], prototypes)
    paragraph = list(_body(out))[3]
    assert _text(paragraph) == "Text mai lung"


def test_missing_reference_and_red_marker(tmp_path: Path) -> None:
    source, prototypes, locator = _source(tmp_path)
    out = tmp_path / "out.docx"
    with pytest.raises(OfficeError) as error:
        render(source, out, locator, [Paragraph("body", [Ref("fig", "unknown")])], prototypes)
    assert error.value.code == "block_reference"
    render(
        source,
        out,
        locator,
        [Missing("missing", "n.d."), Paragraph("tabcap", ["before ", Num(None, 2), " after"])],
        prototypes,
    )
    body = _body(out)
    assert "n.d." in _text(body)
    assert "date indisponibile" in _text(body)
    assert b"FF0000" in read_parts(out)["word/document.xml"]
    inline = list(body)[4]
    red_text = "".join(
        _text(run)
        for run in inline.iter(W + "r")
        if (colour := run.find(W + "rPr/" + W + "color")) is not None
        and colour.get(W + "val") == "FF0000"
    )
    assert red_text == "date indisponibile"


def test_table_formats_numeric_value_and_marks_missing_cell(tmp_path: Path) -> None:
    source, prototypes, locator = _source(tmp_path)
    out = tmp_path / "out.docx"
    report = render(
        source,
        out,
        locator,
        [Table("table", [[["2023"], [Num(1234.5, 2)]], [["2024"], [Num(None, 2)]]])],
        prototypes,
    )
    table = next(node for node in _body(out) if node.tag == W + "tbl")
    rows = table.findall(W + "tr")
    assert _text(rows[1].findall(W + "tc")[1]) == "1.234,50"
    missing_cell = rows[2].findall(W + "tc")[1]
    assert _text(missing_cell) == "date indisponibile"
    assert any(colour.get(W + "val") == "FF0000" for colour in missing_cell.iter(W + "color"))
    assert [value.text for value in report.values] == ["1.234,50", "date indisponibile"]


@pytest.mark.parametrize(
    ("value", "decimals", "expected"),
    [(1234.5, 2, "1.234,50"), (-1234.5, 2, "-1.234,50"), (1234.5, 0, "1.235")],
)
def test_romanian_format(value: float, decimals: int, expected: str) -> None:
    assert format_number(value, decimals) == expected


def test_romanian_format_without_grouping() -> None:
    assert format_number(1234.5, 2, grouping=False) == "1234,50"
    assert format_number(1234.5, 2) == "1.234,50"


def test_romanian_format_rounds_negative_zero_and_large_values() -> None:
    assert format_number(-0.001, 2) == "0,00"
    assert format_number(1e30, 2) == "1.000.000.000.000.000.000.000.000.000.000,00"
