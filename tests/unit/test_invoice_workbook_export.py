"""The invoice workbook contract is checked on a synthetic export."""

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ema.invoices.configuration.field_catalog import (
    ACTIVE_ENERGY,
    ACTIVE_ENERGY_PRICE,
    INVOICE_NUMBER,
    LOCATION_IDENTIFIER,
)
from ema.invoices.export.errors import WorkbookContractError, WorkbookVerificationCode
from ema.invoices.export.exporter import OpenpyxlWorkbookExporter
from ema.invoices.export.verifier import OpenpyxlWorkbookVerifier
from ema.invoices.export.workbook_contract import SUMMARY_COLUMNS
from ema.invoices.models import (
    EnergyCategory,
    FieldStatus,
    FieldValue,
    InvoiceDraft,
    PriceDetail,
    SourceEvidence,
)


def _invoice() -> InvoiceDraft:
    evidence = SourceEvidence(1, "energie activa 10 kWh")
    fields = {
        field_id: FieldValue(Decimal("0"), FieldStatus.APPROVED)
        for field_id, _header in SUMMARY_COLUMNS
        if field_id is not None
    }
    fields.update(
        {
            INVOICE_NUMBER: FieldValue("A-1", FieldStatus.APPROVED),
            LOCATION_IDENTIFIER: FieldValue("POD-1", FieldStatus.APPROVED),
            ACTIVE_ENERGY: FieldValue(Decimal("10"), FieldStatus.APPROVED),
            ACTIVE_ENERGY_PRICE: FieldValue(Decimal("2"), FieldStatus.APPROVED),
        }
    )
    return InvoiceDraft(
        "draft-1",
        "synthetic.pdf",
        "Supplier",
        fields,
        [
            PriceDetail(
                EnergyCategory.ACTIVE_ENERGY,
                "active",
                Decimal("10"),
                "kWh",
                Decimal("2"),
                Decimal("10"),
                "kWh",
                Decimal("2"),
                Decimal("20"),
                evidence,
            )
        ],
    )


def test_export_links_summary_to_source_detail_and_verifies_it(tmp_path: Path) -> None:
    invoice = _invoice()
    path = tmp_path / "invoice.xlsx"
    assert OpenpyxlWorkbookExporter().export([invoice], path) == path
    book = load_workbook(path)
    summary = book["Centralizator"]
    details = book["Detalii preturi"]
    assert summary["A2"].value == "A-1"
    assert summary["G2"].value.startswith("=")
    assert details["C2"].value == "A-1"
    assert details["Q2"].value == "energie activa 10 kWh"
    assert summary["V2"].value == details["R2"].value == "draft-1"
    assert summary.column_dimensions["V"].hidden
    assert details.column_dimensions["R"].hidden
    book.close()
    OpenpyxlWorkbookVerifier().verify(path, [invoice])


def test_verifier_rejects_changed_source_link(tmp_path: Path) -> None:
    invoice = _invoice()
    path = OpenpyxlWorkbookExporter().export([invoice], tmp_path / "invoice.xlsx")
    book = load_workbook(path)
    book["Detalii preturi"]["R2"] = "another-draft"
    book.save(path)
    book.close()
    with pytest.raises(WorkbookContractError) as error:
        OpenpyxlWorkbookVerifier().verify(path, [invoice])
    assert error.value.code is WorkbookVerificationCode.LINK_IDENTIFIERS


def test_export_refuses_unreviewed_invoice_and_empty_batch(tmp_path: Path) -> None:
    exporter = OpenpyxlWorkbookExporter()
    with pytest.raises(ValueError, match="empty invoice batch"):
        exporter.export([], tmp_path / "empty.xlsx")
    invoice = _invoice()
    invoice.fields[INVOICE_NUMBER] = FieldValue(None, FieldStatus.MISSING)
    with pytest.raises(ValueError, match="requiring review"):
        exporter.export([invoice], tmp_path / "unreviewed.xlsx")
    assert not (tmp_path / "unreviewed.xlsx").exists()


@pytest.mark.parametrize(
    "change,code",
    [
        ("sheet", WorkbookVerificationCode.SHEETS),
        ("summary_header", WorkbookVerificationCode.SUMMARY_HEADERS),
        ("detail_header", WorkbookVerificationCode.DETAIL_HEADERS),
        ("summary_count", WorkbookVerificationCode.SUMMARY_ROW_COUNT),
        ("detail_count", WorkbookVerificationCode.DETAIL_ROW_COUNT),
        ("hidden_link", WorkbookVerificationCode.LINK_IDENTIFIERS),
        ("quantity", WorkbookVerificationCode.FORMULAS),
        ("price", WorkbookVerificationCode.FORMULAS),
        ("unit", WorkbookVerificationCode.FORMULAS),
        ("value", WorkbookVerificationCode.FORMULAS),
        ("summary_format", WorkbookVerificationCode.NUMBER_FORMATS),
        ("detail_format", WorkbookVerificationCode.NUMBER_FORMATS),
        ("negative_format", WorkbookVerificationCode.NEGATIVE_FORMATTING),
    ],
)
def test_verifier_rejects_contract_changes(  # noqa: C901, PLR0912
    tmp_path: Path, change: str, code: WorkbookVerificationCode
) -> None:
    invoice = _invoice()
    path = OpenpyxlWorkbookExporter().export([invoice], tmp_path / "invoice.xlsx")
    book = load_workbook(path)
    summary = book["Centralizator"]
    details = book["Detalii preturi"]
    if change == "sheet":
        summary.title = "Wrong"
    elif change == "summary_header":
        summary["A1"] = "Wrong"
    elif change == "detail_header":
        details["A1"] = "Wrong"
    elif change == "summary_count":
        summary["A3"] = "extra"
    elif change == "detail_count":
        details["A3"] = "extra"
    elif change == "hidden_link":
        details.column_dimensions["R"].hidden = False
    elif change == "quantity":
        summary["G2"] = 10
    elif change == "price":
        summary["I2"] = 3
    elif change == "unit":
        summary["H2"] = "MWh"
    elif change == "value":
        summary["J2"] = 20
    elif change == "summary_format":
        summary["A2"].number_format = "General"
    elif change == "detail_format":
        details["I2"].number_format = "General"
    else:
        summary.conditional_formatting._cf_rules.clear()
    book.save(path)
    book.close()
    with pytest.raises(WorkbookContractError) as error:
        OpenpyxlWorkbookVerifier().verify(path, [invoice])
    assert error.value.code is code
