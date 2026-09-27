"""Synthetic cross-platform package and content comparison."""

import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pypdfium2
import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from openpyxl import Workbook
from scripts import piee_parity
from scripts.piee_parity import compare


def _docx(path: Path, number: str = "42", page: str = "3") -> None:
    document = Document()
    document.add_paragraph(f"Value {number}")
    document.styles.add_style("TOC1", WD_STYLE_TYPE.PARAGRAPH)
    paragraph = document.add_paragraph(style="TOC1")
    paragraph.add_run("Chapter ")
    paragraph.add_run(page)
    document.save(path)


def _xlsx(path: Path, number: int = 42) -> None:
    book = Workbook()
    book.active["A1"] = number
    book.save(path)


def _pdf(path: Path, pages: int = 1) -> None:
    document = pypdfium2.PdfDocument.new()
    for _ in range(pages):
        document.new_page(595, 842)
    document.save(path)


def _timestamp_only(path: Path) -> None:
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["docProps/core.xml"] = b"<different timestamp='today'/>"
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def test_equal_and_timestamp_only(tmp_path: Path, capsys) -> None:
    left, right = tmp_path / "mac", tmp_path / "windows"
    left.mkdir()
    right.mkdir()
    for folder in (left, right):
        _docx(folder / "PIEE-draft.docx")
        _xlsx(folder / "Prelucrare-date.xlsx")
        _docx(folder / "PIEE-final.docx")
        _docx(folder / "Audit-ciorna.docx")
        _docx(folder / "Audit-final.docx")
        _pdf(folder / "PIEE-final.pdf")
        _pdf(folder / "Audit-final.pdf")
    _timestamp_only(right / "PIEE-draft.docx")
    assert compare(left, right)
    assert capsys.readouterr().out.count("| equal") == 7


def test_changed_number_toc_pdf_and_missing(tmp_path: Path, capsys) -> None:
    left, right = tmp_path / "mac", tmp_path / "windows"
    left.mkdir()
    right.mkdir()
    for folder in (left, right):
        _xlsx(folder / "Prelucrare-date.xlsx")
        _docx(folder / "PIEE-final.docx")
        _pdf(folder / "PIEE-final.pdf")
    _xlsx(right / "Prelucrare-date.xlsx", 43)
    _docx(right / "PIEE-final.docx", page="4")
    _pdf(right / "PIEE-final.pdf", pages=2)
    _docx(left / "PIEE-draft.docx")
    assert not compare(left, right)
    output = capsys.readouterr().out
    assert "Prelucrare-date.xlsx | DIFFERENT:" in output
    assert "PIEE-final.docx | DIFFERENT: TOC" in output
    assert "PIEE-final.pdf | DIFFERENT: page count" in output
    assert "PIEE-draft.docx | DIFFERENT: missing" in output


def test_empty_directories_are_not_evidence(tmp_path: Path, capsys) -> None:
    left, right = tmp_path / "mac", tmp_path / "windows"
    left.mkdir()
    right.mkdir()
    assert not compare(left, right)
    assert "no files" in capsys.readouterr().out


def test_command_failure_includes_stderr() -> None:
    with pytest.raises(RuntimeError) as error:
        piee_parity._run(
            [
                sys.executable,
                "-c",
                "import sys; print('decision failed', file=sys.stderr); sys.exit(3)",
            ],
            {},
        )
    assert "(3)" in str(error.value)
    assert "decision failed" in str(error.value)


def test_conflict_loop_has_a_cap(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    inputs = tmp_path / "received"
    inputs.mkdir()
    for name in ("Anexa.xlsx", "Necesar.xls", "Prelucrare.xls"):
        (inputs / name).touch()
    monkeypatch.setenv("EMA_PIEE_BASE_DOCUMENT", str(tmp_path / "base.docx"))
    monkeypatch.setenv("EMA_PIEE_BASE_DIRECTORY", str(tmp_path / "base"))
    calls: list[list[str]] = []

    def fake_run(command: list[str], _env: dict[str, str], *, interactive: bool = False) -> str:
        assert not interactive
        calls.append(command)
        if "generate" in command:
            return '{"job": "synthetic-job"}'
        if "fields" in command:
            return '[{"id":"stuck-field","alternatives":[{"id":"candidate"}],"revision":1}]'
        return "{}"

    monkeypatch.setattr(piee_parity, "_run", fake_run)
    with pytest.raises(RuntimeError, match="conflicts remain after 10 decisions: stuck-field"):
        piee_parity.run("ema", inputs, tmp_path / "out")
    assert sum("decide" in command for command in calls) == 10
