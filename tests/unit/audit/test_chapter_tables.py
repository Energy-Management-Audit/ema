"""The ch. 2-3 tables take their rows from the reviewed fields: one row per data row."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from tests.unit.audit.test_draft_checks import _fact

from ema.audit.catalogue import CATALOGUE
from ema.audit.chapter_tables import (
    BOILERS_TABLE,
    CAPTIONS,
    EMPLOYEES_TABLE,
    MISSING,
    TURNOVER_TABLE,
    VEHICLES_TABLE,
    TableSpec,
    fill_table,
    write_tables,
)
from ema.audit.chapter_tables import (
    _table as find_table,
)
from ema.audit.chapter_tables_data import FAMILIES, boiler_rows, vehicle_rows, yearly
from ema.audit.draft_render import render_section
from ema.audit.draft_schema import SECTION_FACTS, DraftText, SectionDraft
from ema.audit.render import ITERATED
from ema.audit.section_body import REMOVABLE_PHOTOS, _removable
from ema.core.errors import EmaError
from ema.core.office.blocks import Table
from ema.core.review.models import Field

TITLES = {section.id: section.title for section in CATALOGUE}


def _field(key: str, value: object, *, review: str = "pending", decimals: int = 0) -> Field:
    number = isinstance(value, int | Decimal)
    return Field.model_validate(
        {
            "id": key,
            "job_id": "job",
            "key": key,
            "label": key,
            "value_type": "number" if number else "text",
            "decimals": decimals,
            "value": value,
            "state": "supplied",
            "presence": "found",
            "review": review,
        }
    )


def _row(family: str, number: int, **cells: object) -> list[Field]:
    return [_field(f"{family}{number}.{role}", value) for role, value in cells.items()]


def test_yearly_follows_the_years_held_and_a_rejected_value_is_missing() -> None:
    table = yearly(
        [
            _field("audit.employees.2025", 379),
            _field("audit.employees.2023", 377),
            _field("audit.employees.2024", 376, review="rejected"),
            _field("audit.employees", 379),
        ],
        "audit.employees.",
    )
    assert table.years == [2023, 2024, 2025]
    assert table.rows == [["2023", "377"], ["2024", None], ["2025", "379"]]
    assert table.values == [377.0, None, 379.0]


def test_money_prints_as_in_her_documents() -> None:
    field = _field("turnover.2025", Decimal("405087623"), decimals=2)
    assert yearly([field], "turnover.").rows == [["2025", "405.087.623,00"]]


def test_equipment_rows_keep_sheet_order_resource_and_missing_cells() -> None:
    fields = [
        *_row("audit.equipment_row.", 2, name="Compresor", process="Aer", count=1, power=25),
        *_row(
            "audit.equipment_row.",
            1,
            name="Centrala A",
            process="Incalzire",
            count=2,
            power=500,
            resource="gaz",
        ),
        *_row("audit.vehicle.", 1, name="Autoturism", maker="Marca Z", type="308", count=42),
        *_row("audit.forklift.", 1, name="MARCA : Marca X", fuel="electric"),
        *_row("audit.forklift.", 1, type="TIP : ERE 220 / SERIA : 90120354"),
        *_row("audit.forklift.", 2, type="TIP : TFG 316"),
    ]
    assert boiler_rows(fields) == [
        ["Centrala A", "Incalzire", "2", "500", "gaz"],
        ["Compresor", "Aer", "1", "25", None],
    ]
    assert vehicle_rows(fields) == [
        ["Autoturism Marca Z 308", "42", None],
        ["Autostivuitor Marca X ERE 220 (electric)", "1", None],
        ["Autostivuitor TFG 316", "1", None],
    ]


def _table() -> object:
    document = Document()
    table = document.add_table(rows=4, cols=3)
    for index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.paragraphs[0].add_run("[de completat]" if index != 2 else "x")
    return table._tbl


def test_fill_clones_the_prototype_row_per_data_row_and_marks_missing_red() -> None:
    table = _table()
    spec = TableSpec("ch3.x", ("Nr", "Denumire", "Buc."), numbered=True)
    fill_table(table, spec, [["A", "1"], [None, "2"], ["C", None], ["D", "4"], ["E", "5"]])  # type: ignore[arg-type]
    rows = table.findall(qn("w:tr"))  # type: ignore[attr-defined]
    text = [
        ["".join(t.text or "" for t in c.iter(qn("w:t"))) for c in r.findall(qn("w:tc"))]
        for r in rows
    ]
    # header, then the data rows: the Nr column is the base's own and stays as it was
    assert text[0] == ["Nr", "Denumire", "Buc."]
    assert [row[1:] for row in text[1:]] == [
        ["A", "1"],
        [MISSING, "2"],
        ["C", MISSING],
        ["D", "4"],
        ["E", "5"],
    ]
    red = [
        run.find(qn("w:rPr")).find(qn("w:color")).get(qn("w:val"))  # type: ignore[union-attr]
        for run in rows[2].iter(qn("w:r"))
        if "".join(t.text or "" for t in run.iter(qn("w:t"))) == MISSING
    ]
    assert red == ["FF0000"]


def test_fewer_data_rows_remove_the_unused_prototype_rows() -> None:
    table = _table()
    spec = TableSpec("ch2.x", ("An", "Valoare", "Alta"))
    fill_table(table, spec, [["2023", "1", "2"]])  # type: ignore[arg-type]
    assert len(table.findall(qn("w:tr"))) == 2  # type: ignore[attr-defined]


def test_the_render_binds_every_family_it_reads() -> None:
    """A new row, or a decision on one, must make the render stale."""
    assert set(FAMILIES) <= set(ITERATED)


def test_table_with_rows_and_no_place_is_a_render_failure(tmp_path: Path) -> None:
    titles = {section.id: section.title for section in CATALOGUE}
    document = Document()
    document.add_paragraph(titles["ch2"], style="Heading 1")
    document.add_paragraph(titles["ch2.date_generale"], style="Heading 2")
    document.add_paragraph("proza")
    document.add_paragraph(titles["ch3"], style="Heading 1")
    source = tmp_path / "base.docx"
    document.save(source)
    with pytest.raises(EmaError) as failure:
        write_tables(
            source, tmp_path / "out.docx", chapter="ch2", fields=[_field("audit.employees.2025", 7)]
        )
    assert (failure.value.code, failure.value.detail) == ("table_slot", "ch2.date_generale")


def test_drafted_sections_keep_base_tables_captions_and_chart_slots(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph(TITLES["ch2"], style="Heading 1")
    for section, spec in (("ch2.date_generale", EMPLOYEES_TABLE), ("ch2.istorie", TURNOVER_TABLE)):
        document.add_paragraph(TITLES[section], style="Heading 2")
        document.add_paragraph("proza veche")
        document.add_paragraph("Tabelul [de completat]", style="Caption")
        table = document.add_table(rows=2, cols=len(spec.header))
        for row in table.rows:
            for cell in row.cells:
                cell.text = "antet"
        for cell, label in zip(table.rows[0].cells, spec.header, strict=True):
            cell.text = label
        document.add_paragraph("[de completat]")
        document.add_paragraph("Fig. [de completat]", style="Caption")
    document.add_paragraph(TITLES["ch3"], style="Heading 1")
    document.add_paragraph(TITLES["ch3.flux"], style="Heading 2")
    document.add_paragraph(TITLES["ch3.equipment"], style="Heading 3")
    document.add_paragraph("proza veche")
    document.add_paragraph("Tabelul [de completat]", style="Caption")
    table = document.add_table(rows=3, cols=6)
    for row in table.rows:
        for cell in row.cells:
            cell.text = "antet"
    for cell, label in zip(table.rows[0].cells, BOILERS_TABLE.header, strict=True):
        cell.text = label
    document.add_paragraph(TITLES["ch4"], style="Heading 1")
    current = tmp_path / "base.docx"
    document.save(current)
    for section in ("ch2.date_generale", "ch3.equipment"):
        output = tmp_path / f"{section}.docx"
        key = sorted(SECTION_FACTS[section])[0]
        draft = SectionDraft(
            section=section,
            status="drafted",
            paragraphs=[DraftText(text=f"proza {{{{f:{key}}}}}", fact_ids=[key])],
        )
        render_section(current, output, draft, {key: _fact(key, "nouă")}, (), job="synthetic")
        current = output
    filled = tmp_path / "filled.docx"
    fields = [
        _field("audit.employees.2025", 7),
        _field("turnover.2025", 900),
        _field("audit.equipment_row.1.name", "Centrală"),
        _field("audit.equipment_row.1.count", 2),
    ]
    write_tables(current, filled, chapter="ch2", fields=fields)
    write_tables(filled, current, chapter="ch3", fields=fields)
    result = Document(current)
    assert len(result.tables) == 3
    assert result.tables[0].rows[1].cells[1].text == "7"
    assert result.tables[1].rows[1].cells[1].text == "900"
    assert result.tables[2].rows[2].cells[1].text == "Centrală"
    texts = [paragraph.text for paragraph in result.paragraphs]
    assert texts.count("proza nouă") == 2
    assert texts.count("Tabelul [de completat]") == 0
    assert texts.count("Fig. [de completat]") == 0
    assert "Tabelul Numărul mediu de angajați" in texts
    assert "Fig. Evoluția numărului mediu de angajați" in texts
    assert "Tabelul Cifra de afaceri (lei)" in texts
    assert "Fig. Evoluția cifrei de afaceri (lei)" in texts
    assert "Tabelul Echipamente și utilaje" in texts
    assert CAPTIONS[BOILERS_TABLE][0].title == "Echipamente și utilaje"
    assert CAPTIONS[VEHICLES_TABLE][0].title == "Parcul auto"
    assert CAPTIONS[BOILERS_TABLE][0].unit is None
    assert CAPTIONS[VEHICLES_TABLE][0].unit is None
    assert texts.count("[de completat]") == 2


def test_unfilled_table_keeps_its_caption_marker(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph(TITLES["ch2"], style="Heading 1")
    document.add_paragraph(TITLES["ch2.date_generale"], style="Heading 2")
    document.add_paragraph("Tabelul [de completat]", style="Caption")
    table = document.add_table(rows=2, cols=2)
    for row in table.rows:
        for cell in row.cells:
            cell.text = "antet"
    for cell, label in zip(table.rows[0].cells, EMPLOYEES_TABLE.header, strict=True):
        cell.text = label
    document.add_paragraph("[de completat]")
    document.add_paragraph("Fig. [de completat]", style="Caption")
    document.add_paragraph(TITLES["ch3"], style="Heading 1")
    source, target = tmp_path / "base.docx", tmp_path / "out.docx"
    document.save(source)
    write_tables(source, target, chapter="ch2", fields=[])
    texts = [paragraph.text for paragraph in Document(target).paragraphs]
    assert "Tabelul [de completat]" in texts
    assert "Fig. [de completat]" in texts


def test_filled_caption_uses_catalogue_title(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph(TITLES["ch2"], style="Heading 1")
    document.add_paragraph(TITLES["ch2.date_generale"], style="Heading 2")
    table_caption = document.add_paragraph("Tabelul 2.1 Numărul personalului", style="Caption")
    table = document.add_table(rows=2, cols=2)
    for row in table.rows:
        for cell in row.cells:
            cell.text = "antet"
    for cell, label in zip(table.rows[0].cells, EMPLOYEES_TABLE.header, strict=True):
        cell.text = label
    document.add_paragraph("[de completat]")
    chart_caption = document.add_paragraph("Fig. nr. 2.1 Evoluția personalului", style="Caption")
    document.add_paragraph(TITLES["ch3"], style="Heading 1")
    source, target = (tmp_path / name for name in ("source.docx", "out.docx"))
    table_caption.text = "Tabelul 2.1 [de completat]"
    chart_caption.text = "Fig. nr. 2.1 [de completat]"
    document.save(source)
    write_tables(
        source,
        target,
        chapter="ch2",
        fields=[_field("audit.employees.2025", 7)],
    )
    texts = [paragraph.text for paragraph in Document(target).paragraphs]
    assert "Tabelul 2.1 Numărul mediu de angajați" in texts
    assert "Fig. nr. 2.1 Evoluția numărului mediu de angajați" in texts


@pytest.mark.parametrize("spec", (EMPLOYEES_TABLE, TURNOVER_TABLE, BOILERS_TABLE, VEHICLES_TABLE))
def test_each_table_spec_selects_its_header_instead_of_the_first_table(spec: TableSpec) -> None:
    document = Document()
    decoy = document.add_table(rows=2, cols=len(spec.header))
    for cell in decoy.rows[0].cells:
        cell.text = "alt antet"
    target = document.add_table(rows=2, cols=len(spec.header))
    for cell, label in zip(target.rows[0].cells, spec.header, strict=True):
        cell.text = label
    body = list(document.element.body)
    assert find_table(body, {spec.section: (0, len(body))}, spec) == 1


@pytest.mark.parametrize(
    ("section", "anchor", "header"),
    (
        ("ch3.electricitate", "Tabel 4. [de completat]", "Tipul lămpii"),
        ("ch3.gaz", "Tabel 5. [de completat]", "Parametrii"),
    ),
)
def test_unfilled_base_table_is_removed_by_its_caption_anchor(
    section: str, anchor: str, header: str
) -> None:
    document = Document()
    document.add_paragraph(anchor)
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = header
    region = list(document.element.body)[:2]
    assert _removable(region, section, []) == {0, 1}
    assert _removable(region, section, [Table("table", [])]) == set()


@pytest.mark.parametrize(
    ("section", "number"),
    [(section, number) for section, numbers in REMOVABLE_PHOTOS.items() for number in numbers],
)
def test_each_unfilled_base_photo_caption_is_removed(section: str, number: int) -> None:
    document = Document()
    document.add_paragraph("[de completat]")
    document.add_paragraph(f"Fig. {number}. [de completat]")
    assert _removable(list(document.element.body)[:2], section, []) == {0, 1}


def test_caption_catalogue_has_verbatim_titles_and_units() -> None:
    assert {
        spec: [(caption.key, caption.title, caption.unit) for caption in captions]
        for spec, captions in CAPTIONS.items()
    } == {
        EMPLOYEES_TABLE: [
            ("employees.table", "Numărul mediu de angajați", None),
            ("employees.chart", "Evoluția numărului mediu de angajați", None),
        ],
        TURNOVER_TABLE: [
            ("turnover.table", "Cifra de afaceri", "lei"),
            ("turnover.chart", "Evoluția cifrei de afaceri", "lei"),
        ],
        BOILERS_TABLE: [("boilers.table", "Echipamente și utilaje", None)],
        VEHICLES_TABLE: [("vehicles.table", "Parcul auto", None)],
    }
