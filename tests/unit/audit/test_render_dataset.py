"""Ch. 4's dataset is the parsed workbook with the auditor's decisions laid over it."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from ema.audit.render_dataset import reviewed_dataset
from ema.core.review.models import Field
from ema.energy_data.carriers import Carrier
from ema.energy_data.model import EnergyDataset, FiledValue, Reading
from ema.energy_data.necesar import parse_necesar_info, to_dataset

MONTHS = (
    "Ianuarie",
    "Februarie",
    "Martie",
    "Aprilie",
    "Mai",
    "Iunie",
    "Iulie",
    "August",
    "Septembrie",
    "Octombrie",
    "Noiembrie",
    "Decembrie",
)
GAS = Carrier.natural_gas


def _necesar(path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Cons energetice"
    row = 1
    for label in ("Consum gaze naturale", "Consum energie electrica din SEN"):
        sheet.cell(row, 1, label)
        row += 1
        for year in (2024, 2025):
            for column, value in enumerate((year, *MONTHS, "Total"), 1):
                sheet.cell(row, column, value)
            sheet.cell(row + 1, 1, "[MWh]")
            sheet.cell(row + 2, 1, "[tep]")
            for month in range(12):
                sheet.cell(row + 1, 2 + month, 10)
                sheet.cell(row + 2, 2 + month, 0.86)
            sheet.cell(row + 1, 14, 120)
            sheet.cell(row + 2, 14, 10.32)
            row += 3
    book.save(path)
    return path


def _field(key: str, value: object, *, review: str, unit: str | None = "MWh") -> Field:
    return Field.model_validate(
        {
            "id": key,
            "job_id": "job",
            "key": key,
            "label": key,
            "value_type": "number",
            "unit": unit,
            "value": value,
            "state": "manual" if review == "corrected" else "extracted",
            "presence": "found",
            "review": review,
            "evidence": ["manual"] if review == "corrected" else [],
        }
    )


def _parsed(tmp_path: Path) -> EnergyDataset:
    return to_dataset(parse_necesar_info(_necesar(tmp_path / "necesar.xlsx")))


def test_nothing_decided_leaves_the_parse_as_it_is(tmp_path: Path) -> None:
    dataset = _parsed(tmp_path)
    assert dataset.filed_indicators["tep_total"][2025].value > 0
    pending = [
        _field("carrier.natural_gas.2025.03", Decimal(99), review="pending"),
        _field("carrier.natural_gas.2025.04", Decimal(10), review="accepted"),
    ]
    assert reviewed_dataset(dataset, pending) == dataset


def test_corrected_month_recomputes_only_its_year(tmp_path: Path) -> None:
    dataset = _parsed(tmp_path)
    corrected = [_field("carrier.natural_gas.2025.03", Decimal("12.5"), review="corrected")]
    result = reviewed_dataset(dataset, corrected)
    assert result.carriers[GAS][2025].months[3] == Reading(12.5, "MWh")
    assert 2025 not in result.filed_indicators["tep_total"]
    assert 2025 not in result.filed_indicators["tep.natural_gas"]
    assert result.filed_indicators["tep_total"][2024] == dataset.filed_indicators["tep_total"][2024]
    kept = result.filed_indicators["tep.electricity_grid"][2025]
    assert kept == dataset.filed_indicators["tep.electricity_grid"][2025]
    assert not kept.source.startswith("review:")


def test_rejected_month_becomes_absent(tmp_path: Path) -> None:
    dataset = _parsed(tmp_path)
    result = reviewed_dataset(
        dataset,
        [
            _field("carrier.natural_gas.2024.02", Decimal(10), review="rejected"),
            _field("carrier.natural_gas.2024.03", Decimal(10), review="pending"),
        ],
    )
    assert result.carriers[GAS][2024].months[2] == Reading(None, "MWh")
    assert 2024 not in result.filed_indicators["tep_total"]


def test_corrected_carrier_tep_is_filed_with_its_review_source(tmp_path: Path) -> None:
    dataset = _parsed(tmp_path)
    fields = [
        _field("carrier_tep.natural_gas.2025", Decimal("11.5"), review="corrected", unit="tep"),
        _field("carrier.natural_gas.2025.01", Decimal(11), review="corrected"),
    ]
    result = reviewed_dataset(dataset, fields)
    assert result.filed_indicators["tep.natural_gas"][2025] == FiledValue(
        11.5, "tep", 2, source="review:carrier_tep.natural_gas.2025"
    )
    assert 2025 not in result.filed_indicators["tep_total"]
