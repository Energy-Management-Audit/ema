"""Regression case: separate dossier slots can have the same basename."""

import json
from pathlib import Path

from openpyxl import Workbook

from ema.audit.intake import audit_intake
from ema.audit.intake_tools import IntakeDocument, IntakeTools
from ema.core.jobs import create_job, run_stage, status, subscribe
from ema.core.llm import ReplayProvider
from ema.core.workspace import Workspace


def test_agent_keeps_same_named_files_in_distinct_slots(tmp_path: Path) -> None:
    tools = IntakeTools(
        {
            "first/evidence.txt": IntakeDocument("first/evidence.txt", "Permit granted"),
            "second/evidence.txt": IntakeDocument("second/evidence.txt", "Meter export"),
        },
        {2: "Permit", 13: "Meter"},
    )
    assert tools.list_files({}) == ["first/evidence.txt", "second/evidence.txt"]
    checklist = tmp_path / "0.Necesar info.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "diverse"
    for number in range(1, 14):
        sheet.cell(number, 1, number)
        sheet.cell(number, 2, f"Document {number}")
    book.save(checklist)
    first = tmp_path / "first" / "evidence.txt"
    second = tmp_path / "second" / "evidence.txt"
    for path, content in ((first, "Permit granted"), (second, "Meter export")):
        path.parent.mkdir()
        path.write_text(content, encoding="utf-8")
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    for slot, path in (
        ("dossier/0.Necesar info.xlsx", checklist),
        ("dossier/first/evidence.txt", first),
        ("dossier/second/evidence.txt", second),
    ):
        ws.set_slot(job, slot, ws.add_file("synthetic", path))
    recording = tmp_path / "replay.json"
    recording.write_text(
        json.dumps(
            {
                "source": "hand-authored",
                "format": "openai-chat-completions",
                "responses": [
                    {
                        "choices": [{"message": {"content": "done"}}],
                        "usage": {"prompt_tokens": 1, "completion_tokens": 1},
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    run = run_stage(
        ws, job, "audit_intake", lambda ctx: audit_intake(ctx, replay=ReplayProvider(recording))
    )
    for _ in subscribe(ws, job):
        pass
    outcome = next(row for row in status(ws, job).runs if row["id"] == run)
    assert outcome["state"] == "ready", outcome["error"]
    with ws.connect() as db:
        report_path = ws.artifact_dir(db, job, "audit_intake", run) / "completeness.json"
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert len(report["unclassified"]) == 2
    assert len(set(report["unclassified"])) == 2
