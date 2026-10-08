"""#163 fix round 1: a previous audit is refused unless it is a .docx with ch. 2-3 sections, on
the CLI and HTTP paths alike; replacing it leaves ch. 2-3 drafts stale; a draft that drops or
reorders the reference's parts fails like the other rules."""

from pathlib import Path

import pytest
from docx import Document
from tests.audit_replay import audit_job_with_facts, draft_recording, support_recording
from tests.unit.audit.section_marks_seams import session
from tests.unit.audit.test_draft_reference import _audit, covering
from tests.unit.audit.test_draft_structured import DraftProvider, SupportProvider

from ema.audit.draft_agent import draft_section_run
from ema.audit.draft_chapter import part_issues
from ema.audit.draft_plan import SectionPlan
from ema.audit.draft_prompt import rule_text
from ema.audit.draft_schema import DraftText, SectionDraft
from ema.audit.draft_stage import draft_section
from ema.audit.draft_style import PREVIOUS_AUDIT_SLOT, ExamplePart
from ema.audit.sections import get_status
from ema.audit.stages import add_document
from ema.core.errors import EmaError
from ema.core.llm.models import default_model
from ema.core.workspace import Workspace
from ema.core.workspace.conversion import active_version

SECTION = "ch2.date_generale"


def _no_chapters(path: Path) -> Path:
    """A .docx audit whose only chapter is the first: no ch. 2-3 section to follow."""
    doc = Document()
    doc.add_paragraph("INTRODUCERE", style="Heading 1")
    doc.add_paragraph("Auditul energetic se realizează conform legislației.")
    doc.save(path)
    return path


def _unrelated(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Lista de cumpărături pentru atelier.")
    doc.save(path)
    return path


def _text(path: Path) -> Path:
    path.write_text("Audit energetic", encoding="utf-8")
    return path


@pytest.mark.parametrize(
    ("make", "code"),
    [
        (_text, "previous_audit_type"),
        (_no_chapters, "previous_audit_unusable"),
        (_unrelated, "previous_audit_unusable"),
    ],
)
def test_the_cli_path_refuses_an_unusable_previous_audit(
    tmp_path: Path, make: object, code: str
) -> None:
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    source = make(tmp_path / "previous.docx")  # type: ignore[operator]
    with pytest.raises(EmaError) as refused:
        add_document(ws, job, source, PREVIOUS_AUDIT_SLOT)
    assert refused.value.code == code
    assert active_version(ws, job, PREVIOUS_AUDIT_SLOT) is None


def test_the_http_path_refuses_an_unusable_previous_audit(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    client, headers = session(ws)
    url = f"/jobs/{job}/slots/{PREVIOUS_AUDIT_SLOT}"
    for make, status, code in (
        (_text, 415, "previous_audit_type"),
        (_no_chapters, 422, "previous_audit_unusable"),
        (_unrelated, 422, "previous_audit_unusable"),
    ):
        sha = ws.add_file("synthetic", make(tmp_path / f"{make.__name__}.docx"))
        refused = client.put(url, json={"file_sha": sha}, headers=headers)
        assert (refused.status_code, refused.json()["type"]) == (status, f"urn:ema:error:{code}")
        assert active_version(ws, job, PREVIOUS_AUDIT_SLOT) is None
    valid = _audit(tmp_path / "valid.docx", ("body", "Societatea are sediul în localitate."))
    accepted = client.put(url, json={"file_sha": ws.add_file("synthetic", valid)}, headers=headers)
    assert accepted.status_code == 200, accepted.json()
    assert active_version(ws, job, PREVIOUS_AUDIT_SLOT) is not None


def _drafted(tmp_path: Path) -> tuple[Workspace, str]:
    """A job whose ch2.date_generale was drafted, by replay, from its previous audit."""
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    previous = _audit(tmp_path / "first.docx", ("body", "Societatea are sediul în localitate."))
    add_document(ws, job, previous, PREVIOUS_AUDIT_SLOT)
    draft = covering(1)
    draft_section(
        ws,
        job,
        SECTION,
        draft_recording=draft_recording(ws, job, draft, tmp_path / "draft.json"),
        support_recording=support_recording(ws, job, draft, tmp_path / "support.json", None),
    )
    state = get_status(ws, job, SECTION)
    assert not state.stale
    assert any(item.startswith("reference:previous_audit@1:") for item in state.fingerprint)
    return ws, job


def _replacement(tmp_path: Path) -> Path:
    return _audit(
        tmp_path / "second.docx",
        ("body", "Societatea are sediul în localitate."),
        ("body", "Regimul de lucru este continuu."),
    )


def test_replacing_the_previous_audit_on_the_cli_leaves_the_draft_stale(tmp_path: Path) -> None:
    ws, job = _drafted(tmp_path)
    add_document(ws, job, _replacement(tmp_path), PREVIOUS_AUDIT_SLOT)
    state = get_status(ws, job, SECTION)
    assert state.stale
    assert (state.changed_input or "").startswith("reference:previous_audit@1:")


def test_replacing_the_previous_audit_over_http_leaves_the_draft_stale(tmp_path: Path) -> None:
    ws, job = _drafted(tmp_path)
    client, headers = session(ws)
    sha = ws.add_file("synthetic", _replacement(tmp_path))
    replaced = client.put(
        f"/jobs/{job}/slots/{PREVIOUS_AUDIT_SLOT}", json={"file_sha": sha}, headers=headers
    )
    assert replaced.status_code == 200, replaced.json()
    assert get_status(ws, job, SECTION).stale


PARTS = (
    ExamplePart("body", "{{…}} are sediul social."),
    ExamplePart("bullet", "denumirea;"),
    ExamplePart("bullet", "codul;"),
    ExamplePart("table", "{{…}} | {{…}}"),
    ExamplePart("body", "{{…}} de lucru este continuu."),
)


def _draft(*parts: int | None, section: str = SECTION, unit: int | None = None) -> SectionDraft:
    return SectionDraft(
        section=section,
        status="drafted",
        paragraphs=[DraftText(text="Text.", part=part, unit=unit) for part in parts],
    )


def test_the_parts_in_order_with_a_longer_list_pass() -> None:
    plan = SectionPlan(SECTION, {}, 100, PARTS)
    assert part_issues(_draft(1, 2, 2, 2, 3), plan) == []
    # A missing item holds its part's place like a written one.
    assert part_issues(_draft(1, 2, 3, 3), plan) == []
    # Without a reference there is nothing to follow.
    assert part_issues(_draft(None), SectionPlan(SECTION, {}, None, ())) == []


def test_an_omitted_or_reordered_part_fails_with_its_rule() -> None:
    plan = SectionPlan(SECTION, {}, 100, PARTS)
    omitted = part_issues(_draft(1, 3), plan)
    assert [(issue.code, issue.detail) for issue in omitted] == [("part_omitted", "2")]
    reordered = part_issues(_draft(1, 3, 2), plan)
    assert [issue.code for issue in reordered] == ["part_order"]
    unknown = part_issues(_draft(1, 2, 3, 4), plan)
    assert [(issue.code, issue.detail) for issue in unknown] == [("part_invalid", "4")]
    assert rule_text("part_omitted").startswith("Fiecare element are câmpul part")


def test_each_process_unit_follows_the_reference_from_its_first_part() -> None:
    plan = SectionPlan("ch3.process", {}, 100, PARTS[:1] + PARTS[4:])
    paragraphs = [
        DraftText(text="Text.", part=part, unit=unit) for unit in (1, 2) for part in (1, 2)
    ]
    draft = SectionDraft(section="ch3.process", status="drafted", paragraphs=paragraphs)
    assert part_issues(draft, plan) == []
    dropped = draft.model_copy(update={"paragraphs": paragraphs[:3]})
    assert [(issue.code, issue.location) for issue in part_issues(dropped, plan)] == [
        ("part_omitted", "unit:2")
    ]


def test_a_draft_that_drops_a_part_goes_back_with_rule_nineteen(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    base = _audit(
        tmp_path / "base.docx",
        ("body", "Societatea are sediul în localitate."),
        ("body", "Regimul de lucru este continuu."),
    )
    identity = tmp_path / "identity.json"
    identity.write_text("[]", encoding="utf-8")
    monkeypatch.setenv("EMA_AUDIT_BASE_DOCUMENT", str(base))
    monkeypatch.setenv("EMA_AUDIT_BASE_IDENTITY", str(identity))
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    provider = DraftProvider([covering(1), covering(2)])
    _, accepted, _, _ = draft_section_run(
        ws,
        job,
        SECTION,
        provider,
        SupportProvider(),
        model_id=default_model("openai").id,
        synthetic=True,
    )
    assert accepted == covering(2)
    (error,) = provider.requests[1]["errors"]
    assert (error["rule"], error["detail"]) == ("part_omitted", "2")
    assert error["rule_text"].startswith("Fiecare element are câmpul part")
