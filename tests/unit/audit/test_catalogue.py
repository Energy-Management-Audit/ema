"""Synthetic outline, mapping and repeatable-unit behavior."""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ema.audit import headings as heading_module
from ema.audit.catalogue import CATALOGUE, AuditFact
from ema.audit.catalogue_types import CarrierPattern
from ema.audit.headings import headings, map_headings
from ema.audit.inventory import inventory


def _outline(node, level: int) -> None:  # type: ignore[no-untyped-def]
    ppr = node.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    ppr.append(outline)


def test_outline_from_paragraph_and_style_chain(tmp_path: Path) -> None:
    doc = Document()
    parent = doc.styles.add_style("OwnHeading", WD_STYLE_TYPE.PARAGRAPH)
    _outline(parent.element, 0)
    child = doc.styles.add_style("Child", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    doc.add_paragraph("DESCRIEREA ŞI SCOPUL AUDITULUI", style=child)
    p = doc.add_paragraph("OBIECTIVE URMĂRITE")
    _outline(p._p, 1)
    doc.add_paragraph(" ", style=child)
    ignored = doc.add_paragraph("Figura 1 Caption")
    _outline(ignored._p, 7)
    path = tmp_path / "synthetic.docx"
    doc.save(path)
    result = headings(path)
    assert [(h.level, h.path, h.index) for h in result] == [
        (0, (), 0),
        (1, ("DESCRIEREA ŞI SCOPUL AUDITULUI",), 1),
    ]
    assert [h.text for h in map_headings(path, "synthetic").excluded_captions] == [
        "Figura 1 Caption"
    ]


def test_high_outline_levels_nest_by_nearest_lower_level(tmp_path: Path) -> None:
    doc = Document()
    root = doc.add_paragraph("DESCRIEREA ŞI SCOPUL AUDITULUI")
    _outline(root._p, 0)
    parent = doc.add_paragraph("CONŢINUTUL AUDITULUI")
    _outline(parent._p, 1)
    for text in ("Detail A", "Detail B"):
        child = doc.add_paragraph(text)
        _outline(child._p, 8)
    path = tmp_path / "deep.docx"
    doc.save(path)
    assert [item.path for item in headings(path)][2:] == [
        (root.text, parent.text),
        (root.text, parent.text),
    ]


def test_alias_diacritics_duplicate_and_unmapped(tmp_path: Path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    doc = Document()
    root = doc.add_paragraph("DESCRIEREA SI SCOPUL AUDITULUI")
    _outline(root._p, 0)
    alias = doc.add_paragraph("CONŢINUTUL AUDITULUI")
    _outline(alias._p, 1)
    unknown = doc.add_paragraph("Unknown heading")
    _outline(unknown._p, 1)
    path = tmp_path / "synthetic.docx"
    doc.save(path)
    mapped = map_headings(path, "synthetic")
    assert [item.section_id for item in mapped.mapped] == ["ch1", "ch1.continut"]
    assert [item.text for item in mapped.unmapped] == ["Unknown heading"]
    duplicate = replace(CATALOGUE[2], aliases=("CONȚINUTUL AUDITULUI",))
    monkeypatch.setattr(heading_module, "CATALOGUE", (*CATALOGUE, duplicate))
    mapped = map_headings(path, "synthetic")
    assert alias.text in [item.text for item in mapped.unmapped]


def _process_doc(path: Path, count: int) -> None:
    doc = Document()
    root = doc.add_paragraph("DESCRIEREA SITUAŢIEI EXISTENTE")
    _outline(root._p, 0)
    for index in range(count):
        paragraph = doc.add_paragraph(f"DESCRIEREA SECȚIEI UNIT {index + 1}")
        _outline(paragraph._p, 1)
        doc.add_paragraph("Synthetic description")
    doc.save(path)


def test_inventory_scales_process_units(tmp_path: Path) -> None:
    for count in (2, 5):
        path = tmp_path / f"synthetic-{count}.docx"
        _process_doc(path, count)
        result = inventory(path)
        assert len(result.processes) == count
        assert all(unit.element_range[1] > unit.element_range[0] for unit in result.processes)


def test_every_audit_fact_has_a_catalogue_user() -> None:
    used = {ref for section in CATALOGUE for ref in section.facts if isinstance(ref, AuditFact)}
    assert used == set(AuditFact)
    assert any(isinstance(ref, CarrierPattern) for section in CATALOGUE for ref in section.facts)
