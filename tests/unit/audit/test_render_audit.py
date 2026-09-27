"""The full audit render on a synthetic base: order, failures, markers, Word and outputs."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from tests.unit.audit.render_seams import (
    FakeWord,
    fill_writer,
    outputs,
    run_render,
    summary_of,
    synthetic_render,
)

from ema.audit import render
from ema.audit.catalogue import CATALOGUE
from ema.audit.render import start_audit_render
from ema.audit.render_steps import body_counts
from ema.audit.sections import Status, set_status
from ema.core.errors import EmaError
from ema.core.workspace import Workspace

TITLES = {section.id: section.title for section in CATALOGUE}


@pytest.fixture
def job(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[Workspace, str, list[str]]:
    return synthetic_render(tmp_path, monkeypatch)


def _progress(ws: Workspace, job: str) -> list[tuple[int, int, str]]:
    with ws.connect() as db:
        rows = db.execute(
            "SELECT payload FROM job_events WHERE job_id=? AND type='stage_progress' ORDER BY seq",
            (job,),
        ).fetchall()
    payloads = [json.loads(row["payload"]) for row in rows]
    return [(item["done"], item["total"], item["message"]) for item in payloads]


def failed_payload(ws: Workspace, job: str) -> dict[str, object]:
    with ws.connect() as db:
        row = db.execute(
            "SELECT payload FROM job_events WHERE job_id=? AND type='stage_failed' "
            "ORDER BY seq DESC LIMIT 1",
            (job,),
        ).fetchone()
    return json.loads(row["payload"])


def test_draft_without_word_writes_every_chapter_in_order(
    job: tuple[Workspace, str, list[str]],
) -> None:
    ws, job_id, calls = job
    record = run_render(ws, job_id)
    assert record["state"] == "ready", record["error"]
    assert calls == ["draft", "draft", "ch4", "ch5", "ch6"]
    progress = _progress(ws, job_id)
    assert progress[0] == (0, 7, "Capitolul 1 din 7 — Descrierea şi scopul auditului")
    assert progress[3] == (
        3,
        7,
        "Capitolul 4 din 7 — Analiza modului în care se realizează consumurile energetice pe "
        "platforma societății",
    )
    assert progress[-1][:2] == (7, 7)
    summary = summary_of(ws, job_id, str(record["id"]))
    assert [(item.number, item.section_id, item.page) for item in summary.chapters] == [
        (number, f"ch{number}", None) for number in range(1, 8)
    ]
    assert summary.chapters[1].title == "Descrierea şi istoricul societăţii"
    assert [marker.section_id for marker in summary.markers] == [
        "ch2.date_generale",
        "ch3.flux",
        "ch4.bilant_real",
    ]
    assert summary.markers[0].label == "Date generale"
    assert (summary.tables, summary.charts) == (1, 0)
    assert (summary.pdf, summary.toc_pages_set, summary.failures) == (False, False, [])
    assert summary.unit_plan.processes_source == "default"
    assert outputs(ws, job_id) == [("Audit-ciorna.docx", "draft")]


def test_failed_section_keeps_its_markers_and_the_rest_continues(
    job: tuple[Workspace, str, list[str]], monkeypatch: pytest.MonkeyPatch
) -> None:
    ws, job_id, calls = job

    def broken(*_: object, **__: object) -> None:
        raise EmaError("draft_slots", "Secţiunea nu are ancore.", "ch4")

    monkeypatch.setattr(render, "write_four", broken)
    record = run_render(ws, job_id)
    assert record["state"] == "ready"
    assert json.loads(str(record["outcome"]))["item_failures"] == ["ch4"]
    summary = summary_of(ws, job_id, str(record["id"]))
    assert [(item.section_id, item.code) for item in summary.failures] == [("ch4", "draft_slots")]
    assert "ch4.concluzii" in [marker.section_id for marker in summary.markers]
    assert calls == ["draft", "draft", "ch5", "ch6"]


def test_package_issue_fails_the_run_and_publishes_nothing(
    job: tuple[Workspace, str, list[str]], monkeypatch: pytest.MonkeyPatch
) -> None:
    ws, job_id, _ = job
    monkeypatch.setattr(render, "package_issues", lambda path, identity: ["leftover identity"])
    record = run_render(ws, job_id)
    assert record["state"] == "failed"
    assert record["error"] == "leftover identity"
    assert failed_payload(ws, job_id) == {
        "code": "audit_package",
        "message": "Pachetul Word al auditului este invalid.",
    }
    assert outputs(ws, job_id) == []


def test_na_sections_are_dropped_before_counting(
    job: tuple[Workspace, str, list[str]],
) -> None:
    ws, job_id, _ = job
    set_status(ws, job_id, "ch4.bilant_real", Status.NA, "user")
    summary = summary_of(ws, job_id, str(run_render(ws, job_id)["id"]))
    assert summary.dropped == ["ch4.bilant_real"]
    assert "ch4.bilant_real" not in [marker.section_id for marker in summary.markers]


def test_word_sets_pages_and_pdf_comes_before_docx(
    job: tuple[Workspace, str, list[str]], monkeypatch: pytest.MonkeyPatch
) -> None:
    ws, job_id, _ = job
    word = FakeWord()
    monkeypatch.setattr(render, "word_available", lambda settings: True)
    monkeypatch.setattr(render, "word_automation", lambda settings: word)
    summary = summary_of(ws, job_id, str(run_render(ws, job_id)["id"]))
    assert word.calls == ["toc", "pdf", "open"]
    assert (summary.pdf, summary.toc_pages_set) == (True, True)
    assert all(item.page is not None for item in summary.chapters)
    assert outputs(ws, job_id) == [("Audit-ciorna.pdf", "draft"), ("Audit-ciorna.docx", "draft")]


def test_final_needs_word_and_no_marker(
    job: tuple[Workspace, str, list[str]], monkeypatch: pytest.MonkeyPatch
) -> None:
    ws, job_id, _ = job
    with pytest.raises(EmaError) as refused:
        start_audit_render(ws, job_id, "final")
    assert (refused.value.code, refused.value.user_message_ro) == (
        "word_unavailable",
        "Microsoft Word nu este disponibil.",
    )
    word = FakeWord()
    monkeypatch.setattr(render, "word_available", lambda settings: True)
    monkeypatch.setattr(render, "word_automation", lambda settings: word)
    record = run_render(ws, job_id, "final")
    assert record["state"] == "failed"
    assert record["error"] == "ch2.date_generale; ch3.flux; ch4.bilant_real"
    assert failed_payload(ws, job_id)["code"] == "audit_markers"
    assert word.calls == []
    set_status(ws, job_id, "ch4.bilant_real", Status.NA, "user")
    monkeypatch.setattr(render, "write_draft", fill_writer("ch2.date_generale", "ch3.flux"))
    record = run_render(ws, job_id, "final")
    assert record["state"] == "ready", record["error"]
    assert outputs(ws, job_id) == [("Audit-final.pdf", "draft"), ("Audit-final.docx", "final")]


def test_missing_base_setting_is_named(
    job: tuple[Workspace, str, list[str]], monkeypatch: pytest.MonkeyPatch
) -> None:
    ws, job_id, _ = job
    monkeypatch.delenv("EMA_AUDIT_BASE_IDENTITY")
    with pytest.raises(EmaError) as refused:
        start_audit_render(ws, job_id, "draft")
    assert (refused.value.code, refused.value.detail) == (
        "audit_base_missing",
        "audit_base_identity",
    )


def test_markers_before_the_first_heading_belong_to_front(tmp_path: Path) -> None:
    docx = tmp_path / "cover.docx"
    document = Document()
    document.add_paragraph("[de completat]")
    document.add_paragraph(TITLES["ch1"], style="Heading 1")
    document.add_paragraph("[de completat] şi [de completat]")
    document.save(str(docx))
    counts = body_counts(docx)
    assert counts.markers == [
        ("front", "Prima pagină şi cuprinsul"),
        ("ch1", "Descrierea şi scopul auditului"),
        ("ch1", "Descrierea şi scopul auditului"),
    ]
