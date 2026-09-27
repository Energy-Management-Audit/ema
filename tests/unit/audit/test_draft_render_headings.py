"""A drafted paragraph never lands in a heading, whatever the base calls its heading style."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.unit.audit.test_draft_checks import _draft, _fact

from ema.audit.draft_render import render_section
from ema.core.office.anchors import stamp


def _outline(style: object) -> None:
    properties = style.element.get_or_add_pPr()  # type: ignore[attr-defined]
    level = OxmlElement("w:outlineLvl")
    level.set(qn("w:val"), "1")
    properties.append(level)


def test_outline_styles_and_their_children_are_skipped(tmp_path: Path) -> None:
    base, output, anchors = (tmp_path / name for name in ("b.docx", "o.docx", "b.anchors.json"))
    document = Document()
    heading = document.styles.add_style("Stil heading", WD_STYLE_TYPE.PARAGRAPH)
    _outline(heading)
    child = document.styles.add_style("Stil copil", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = heading
    document.styles.add_style("Stil corp", WD_STYLE_TYPE.PARAGRAPH)
    records = []
    for index, style in enumerate(("Stil heading", "Stil copil", "Stil corp"), 1):
        paragraph = document.add_paragraph("[de completat]", style=style)
        stamp(paragraph._p, f"slot{index}", index)
        records.append(
            {
                "slot": f"slot{index}",
                "section": "ch2.date_generale",
                "classification": "variable",
                "part": "word/document.xml",
            }
        )
    document.save(str(base))
    anchors.write_text(json.dumps({"version": 1, "anchors": records}), encoding="utf-8")
    facts = {"audit.company_name": _fact("audit.company_name", "Atelier Exemplu")}
    draft = _draft("Societatea {{f:audit.company_name}} produce bunuri.", ["audit.company_name"])
    render_section(base, anchors, output, draft, facts, (), job="synthetic")
    assert [paragraph.text for paragraph in Document(str(output)).paragraphs] == [
        "[de completat]",
        "[de completat]",
        "Societatea Atelier Exemplu produce bunuri.",
    ]


def test_based_on_cycle_does_not_hang(tmp_path: Path) -> None:
    base, output, anchors = (tmp_path / name for name in ("b.docx", "o.docx", "b.anchors.json"))
    document = Document()
    first = document.styles.add_style("Stil A", WD_STYLE_TYPE.PARAGRAPH)
    second = document.styles.add_style("Stil B", WD_STYLE_TYPE.PARAGRAPH)
    second.base_style = first
    first.base_style = second
    paragraph = document.add_paragraph("[de completat]", style="Stil A")
    stamp(paragraph._p, "slot", 1)
    document.save(str(base))
    anchors.write_text(
        json.dumps(
            {
                "version": 1,
                "anchors": [
                    {
                        "slot": "slot",
                        "section": "ch2.date_generale",
                        "classification": "variable",
                        "part": "word/document.xml",
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    facts = {"audit.company_name": _fact("audit.company_name", "Atelier Exemplu")}
    draft = _draft("Societatea {{f:audit.company_name}} produce bunuri.", ["audit.company_name"])
    render_section(base, anchors, output, draft, facts, (), job="synthetic")
    assert Document(str(output)).paragraphs[0].text.startswith("Societatea")
