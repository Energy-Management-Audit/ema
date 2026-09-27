"""Documents view never assigns an intake outcome to a different file version."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from ema.audit.documents import documents
from ema.audit.intake import audit_intake
from ema.core.jobs import StageContext, StageOutcome, create_job, run_stage, subscribe
from ema.core.workspace import Workspace


def _checklist(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "diverse"
    for number in range(1, 14):
        sheet.cell(number, 1, number)
        sheet.cell(number, 2, f"Document {number}")
    book.save(path)


def test_status_is_bound_to_active_sha_and_old_report_is_unread(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = create_job(ws, "audit", "synthetic", 2026)
    checklist = tmp_path / "0. Necesar info.xlsx"
    _checklist(checklist)
    source = tmp_path / "1. Source.txt"
    source.write_text("old", encoding="utf-8")
    for path in (checklist, source):
        ws.set_slot(job, f"dossier/{path.name}", ws.add_file("synthetic", path))
    run = run_stage(ws, job, "intake", audit_intake)
    for _ in subscribe(ws, job):
        pass
    view = documents(ws, job)
    assert view.checklist is not None and view.checklist[0].received
    assert view.files[1].sha == ws.list_versions(job, "dossier/1. Source.txt")[0].file_sha
    displayed_revision = view.files[1].slot_revision
    assert displayed_revision > 0
    assert view.files[1].status == "unread"
    with ws.connect() as db:
        report_path = ws.artifact_dir(db, job, "intake", run) / "completeness.json"
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["files"][1]["file_sha"] == view.files[1].sha
    run_stage(ws, job, "read", lambda _ctx: StageOutcome())
    for _ in subscribe(ws, job):
        pass
    assert documents(ws, job).files[0].status == "read"
    source.write_text("new", encoding="utf-8")
    ws.set_slot(job, "dossier/1. Source.txt", ws.add_file("synthetic", source))
    replacement = documents(ws, job).files[1]
    assert replacement.status == "unread"
    assert replacement.slot_revision > displayed_revision
    report["files"][1].pop("file_sha")
    report_path.write_text(json.dumps(report), encoding="utf-8")
    assert documents(ws, job).files[1].status == "unread"


def test_status_reason_and_running_cover_form_rows(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = create_job(ws, "audit", "synthetic", 2026)
    checklist = tmp_path / "0. Necesar info.xlsx"
    _checklist(checklist)
    source = tmp_path / "1. Source.txt"
    source.write_text("synthetic", encoding="utf-8")
    for slot, path in (
        (f"dossier/{checklist.name}", checklist),
        (f"dossier/{source.name}", source),
        ("anexa", source),
        ("measures", source),
    ):
        ws.set_slot(job, slot, ws.add_file("synthetic", path))
    run = run_stage(ws, job, "intake", audit_intake)
    for _ in subscribe(ws, job):
        pass
    with ws.connect() as db:
        path = ws.artifact_dir(db, job, "intake", run) / "completeness.json"
    report = json.loads(path.read_text(encoding="utf-8"))
    record = next(item for item in report["files"] if item["slot"] == "dossier/1. Source.txt")
    record.update(status="failed", error_code="file_type")
    path.write_text(json.dumps(report), encoding="utf-8")
    view = documents(ws, job)
    assert view.files[1].status == "failed"
    assert view.files[1].reason == "Tipul fişierului nu este acceptat."
    assert view.anexa is not None and view.anexa.status == "unread"
    assert view.measures is not None and view.measures.status == "unread"
    assert view.missing == list(range(2, 14))
    with ws.connect() as db:
        db.execute(
            "INSERT INTO runs (id,job_id,stage,owner,state,started_at) VALUES (?,?,?,?,?,?)",
            ("running-read", job, "read", "synthetic", "running", "9999-01-01T00:00:00Z"),
        )
    view = documents(ws, job)
    assert view.files[1].status == "reading"
    assert view.anexa is not None and view.anexa.status == "reading"


def test_form_status_tracks_the_slot_read_by_each_stage(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = create_job(ws, "audit", "synthetic", 2026)
    source = tmp_path / "form.xlsx"
    source.write_text("first", encoding="utf-8")
    sha = ws.add_file("synthetic", source)
    ws.set_slot(job, "anexa", sha)
    ws.set_slot(job, "measures", sha)

    def read_anexa(ctx: StageContext) -> StageOutcome:
        ctx.read_slot("anexa")
        return StageOutcome()

    def read_measures(ctx: StageContext) -> StageOutcome:
        ctx.read_slot("measures")
        return StageOutcome()

    run_stage(ws, job, "read", read_anexa)
    for _ in subscribe(ws, job):
        pass
    view = documents(ws, job)
    assert view.anexa is not None and view.anexa.status == "read"
    assert view.measures is not None and view.measures.status == "unread"

    run_stage(ws, job, "measures", read_measures)
    for _ in subscribe(ws, job):
        pass
    view = documents(ws, job)
    assert view.measures is not None and view.measures.status == "read"

    source.write_text("replacement", encoding="utf-8")
    ws.set_slot(job, "anexa", ws.add_file("synthetic", source))
    view = documents(ws, job)
    assert view.anexa is not None and view.anexa.status == "unread"
    assert view.measures is not None and view.measures.status == "read"

    with ws.connect() as db:
        db.execute(
            "INSERT INTO runs (id,job_id,stage,owner,state,started_at) VALUES (?,?,?,?,?,?)",
            ("failed-measures", job, "measures", "synthetic", "failed", "9999-01-01T00:00:00Z"),
        )
    view = documents(ws, job)
    assert view.measures is not None and view.measures.status == "unread"
    with ws.connect() as db:
        db.execute(
            "INSERT INTO runs (id,job_id,stage,owner,state,started_at) VALUES (?,?,?,?,?,?)",
            ("running-measures", job, "measures", "synthetic", "running", "9999-01-02T00:00:00Z"),
        )
    view = documents(ws, job)
    assert view.measures is not None and view.measures.status == "reading"
