"""the auditor's ch. 4 texts are fields; an n/a section leaves the document and never blocks."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from tests.audit_structure import add_audit_toc, number_audit_headings

from ema.audit.catalogue import CATALOGUE
from ema.audit.chapter_four_blocks import chapter_four_blocks
from ema.audit.content_checks import content_issues
from ema.audit.read import read_dossier
from ema.audit.render_steps import drop_na_sections, droppable
from ema.audit.sections import Status, set_status
from ema.core.jobs import create_job
from ema.core.office.blocks import Missing, Paragraph
from ema.core.review import decide, fields
from ema.core.workspace import Workspace
from ema.energy_data.factors import FACTORS_2026
from ema.energy_data.necesar import parse_necesar_info, to_dataset

FIXTURE = Path(__file__).parents[2] / "fixtures/audit/synthetic_necesar.xlsx"
TITLES = {section.id: section.title for section in CATALOGUE}
NARRATIVES = ("narrative.ch4.bilant_real", "narrative.ch4.concluzii", "narrative.ch4.eficienta")


def _narratives(ws: Workspace, job: str) -> set[str]:
    return {field.key for field in fields(ws, job) if field.key in NARRATIVES}


def test_read_creates_the_three_texts_absent_and_keeps_a_written_one(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = create_job(ws, "audit", "synthetic", 2026)
    read_dossier(ws, job, FIXTURE)
    created = _narratives(ws, job)
    assert sorted(created) == list(NARRATIVES)
    concluzii = next(field for field in fields(ws, job) if field.key == "narrative.ch4.concluzii")
    assert concluzii.value is None and concluzii.presence == "not_found"
    assert (concluzii.chapter, concluzii.value_type) == ("ch4", "text")
    assert concluzii.label == "Concluziile privind analiza consumului echivalent de energie"
    decide(ws, job, concluzii.id, "correct", concluzii.revision, "user", value="Text scris.")
    read_dossier(ws, job, FIXTURE)
    kept = next(field for field in fields(ws, job) if field.key == "narrative.ch4.concluzii")
    assert (kept.value, kept.state, kept.review) == ("Text scris.", "manual", "corrected")


def test_written_text_replaces_the_marker() -> None:
    dataset = to_dataset(parse_necesar_info(FIXTURE))
    blocks = chapter_four_blocks(
        dataset, FACTORS_2026, texts={"ch4.concluzii": "Primul paragraf.\n\n  Al doilea.  "}
    )
    heading = blocks.index(Paragraph("heading:ch4.concluzii", [TITLES["ch4.concluzii"]]))
    assert blocks[heading + 1 : heading + 3] == [
        Paragraph("body", ["Primul paragraf."]),
        Paragraph("body", ["Al doilea."]),
    ]
    real = blocks.index(Paragraph("heading:ch4.bilant_real", [TITLES["ch4.bilant_real"]]))
    assert blocks[real + 1] == Missing("body", "[de completat]")
    assert chapter_four_blocks(dataset, FACTORS_2026) == chapter_four_blocks(
        dataset, FACTORS_2026, texts={}
    )


def test_na_section_text_never_blocks_the_final(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = create_job(ws, "audit", "synthetic", 2026)
    read_dossier(ws, job, FIXTURE)
    set_status(ws, job, "ch4.bilant_real", Status.NA, "user")
    with ws.connect() as db:
        blocking = {issue.message for issue in content_issues(db, job)}
    labels = {message.removeprefix("Textul lipseşte: ") for message in blocking}
    assert labels == {
        "Concluziile privind analiza consumului echivalent de energie",
        "Analiza eficienței utilizării energiei",
        "Introducerea capitolului 3",
        "Introducerea capitolului 6",
    }


def _base(path: Path) -> Path:
    document = Document()
    add_audit_toc(document)
    document.add_paragraph(TITLES["ch4"], style="Heading 1")
    document.add_paragraph(TITLES["ch4.concluzii"], style="Heading 2")
    document.add_paragraph("Concluzii scrise")
    document.add_paragraph(TITLES["ch4.bilant_real"], style="Heading 2")
    document.add_paragraph("Bilanț de eliminat")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Tabel de eliminat"
    document.add_paragraph(TITLES["ch7"], style="Heading 1")
    document.add_paragraph("Finanţare")
    number_audit_headings(document)
    document.save(str(path))
    return path


def test_drop_removes_heading_body_and_toc_entry(tmp_path: Path) -> None:
    docx = _base(tmp_path / "base.docx")
    assert drop_na_sections(docx, ["ch4.bilant_real", "ch6"]) == ["ch4.bilant_real"]
    document = Document(str(docx))
    texts = [paragraph.text for paragraph in document.paragraphs]
    body = [text for text in texts if not text.startswith(("4.", "5.", "6.", "7."))]
    assert "Bilanț de eliminat" not in texts
    assert not document.tables
    assert "Concluzii scrise" in body and "Finanţare" in body
    assert not any(TITLES["ch4.bilant_real"].casefold() in text.casefold() for text in texts)
    toc = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style is not None and paragraph.style.name in {"TOC 1", "TOC 2"}
    ]
    assert any(TITLES["ch4.concluzii"].casefold() in text.casefold() for text in toc)
    assert drop_na_sections(docx, []) == []


def test_droppable_keeps_fixed_sections_and_parents_of_kept_children() -> None:
    chapter_three = [section.id for section in CATALOGUE if section.id.startswith("ch3")]
    chapter_four = [section.id for section in CATALOGUE if section.id.startswith("ch4")]
    assert droppable(["ch1", "ch1.scop", "ch2", "ch7"]) == []
    without_flux = [section for section in chapter_three if section != "ch3.flux"]
    result = droppable(without_flux)
    assert "ch3" not in result and "ch3.flux" not in result
    assert "ch3.process" in result
    assert "ch4" in droppable(chapter_four)
    assert "ch4" not in droppable([item for item in chapter_four if item != "ch4.gaz"])


def test_drafted_child_inside_an_na_parent_survives(tmp_path: Path) -> None:
    document = Document()
    add_audit_toc(document)
    document.add_paragraph(TITLES["ch3"], style="Heading 1")
    document.add_paragraph("Introducere capitol")
    document.add_paragraph(TITLES["ch3.flux"], style="Heading 2")
    document.add_paragraph("Flux redactat")
    document.add_paragraph(TITLES["ch3.process"], style="Heading 2")
    document.add_paragraph("Proces de eliminat")
    document.add_paragraph(TITLES["ch7"], style="Heading 1")
    docx = tmp_path / "base.docx"
    number_audit_headings(document)
    document.save(str(docx))
    chapter_three = [section.id for section in CATALOGUE if section.id.startswith("ch3")]
    na = [section for section in chapter_three if section != "ch3.flux"]
    assert drop_na_sections(docx, droppable(na)) == ["ch3.process"]
    texts = [paragraph.text for paragraph in Document(str(docx)).paragraphs]
    assert "Introducere capitol" in texts and "Flux redactat" in texts
    assert "Proces de eliminat" not in texts
