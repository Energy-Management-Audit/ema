"""Synthetic checklist and intake agent behavior."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from tests.workspace_jobs import create_job

from ema.audit.checklist import read_checklist
from ema.audit.intake import audit_intake
from ema.audit.intake_tools import IntakeDocument, IntakeTools
from ema.core.jobs import run_stage, status, subscribe
from ema.core.llm import AgentContext, Limits, ReplayProvider, agent_state, run_agent
from ema.core.workspace import Workspace

REPLAY = Path(__file__).resolve().parents[2] / "fixtures/llm/intake_openai.json"
GEMINI_REPLAY = REPLAY.with_name("intake_gemini.json")
STAGE_REPLAY = REPLAY.with_name("intake_stage_openai.json")


def _checklist(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "diverse"
    for number in range(1, 14):
        sheet.cell(number + 3, 4, number)
        sheet.cell(number + 3, 5, f"Requested document {number}")
    book.save(path)


def _documents() -> dict[str, IntakeDocument]:
    return {
        "permit.txt": IntakeDocument("permit.txt", "Permit granted for the workshop"),
        "flow.txt": IntakeDocument("flow.txt", "Flow scheme for paint line"),
        "meter.txt": IntakeDocument("meter.txt", "Meter export for July"),
        "photo.txt": IntakeDocument("photo.txt", "Photo of boiler room"),
    }


def test_checklist_uses_number_labels_and_intake_reports_every_file(tmp_path: Path) -> None:
    checklist = tmp_path / "0.Necesar info.xlsx"
    _checklist(checklist)
    items = read_checklist(checklist)
    assert [item.number for item in items] == list(range(1, 14))
    assert items[0].row == 4
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    for source in (checklist, tmp_path / "2.1.permit.txt", tmp_path / "loose.txt"):
        if not source.exists():
            source.write_text("synthetic", encoding="utf-8")
        ws.set_slot(job, f"dossier/{source.name}", ws.add_file("synthetic", source))
    run = run_stage(ws, job, "audit_intake", audit_intake)
    for _ in subscribe(ws, job):
        pass
    record = next(row for row in status(ws, job).runs if row["id"] == run)
    assert record["state"] == "ready", record["error"]
    with ws.connect() as db:
        report_path = ws.artifact_dir(db, job, "audit_intake", run) / "completeness.json"
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert len(report["files"]) == 3
    assert report["received"]["2"] == ["2.1.permit.txt"]
    assert report["missing"] == [1, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13]
    assert report["unclassified"] == ["loose.txt"]
    repeat = run_stage(ws, job, "audit_intake", audit_intake)
    for _ in subscribe(ws, job):
        pass
    rerun = next(row for row in status(ws, job).runs if row["id"] == repeat)
    assert rerun["state"] == "ready" and rerun["publication"] == "current"


@pytest.mark.parametrize("fixture", [REPLAY, GEMINI_REPLAY])
def test_agent_replay_rejects_fabricated_quote_and_resumes(tmp_path: Path, fixture: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    replay = ReplayProvider(fixture)
    tools = IntakeTools(_documents(), {number: f"Item {number}" for number in range(1, 14)})
    context = AgentContext(ws, job, "intake", replay, "gemini-3.6-flash", "intake-v1")
    state = run_agent(context, "Classify each file", tools.tools(), Limits(2))
    assert state.status == "step_limit" and state.steps == 2
    persisted = agent_state(ws, job, "intake")
    assert persisted is not None and persisted.steps == 2 and persisted.status == "step_limit"
    state = run_agent(context, "Classify each file", tools.tools(), Limits(20))
    assert state.status == "done" and replay.calls == 12
    assert tools.classifications == {
        "permit.txt": 2,
        "flow.txt": 5,
        "meter.txt": 13,
        "photo.txt": 4,
    }
    assert tools.missing == {6}
    assert "evidence_quote" in [
        message["content"].get("error")
        for message in state.messages
        if message["role"] == "tool" and isinstance(message["content"], dict)
    ]
    with ws.connect() as db:
        count = db.execute("SELECT count(*) FROM llm_calls WHERE job_id=?", (job,)).fetchone()[0]
    assert count == 12
    ws.delete_job(job)
    with ws.connect() as db:
        assert db.execute("SELECT count(*) FROM agent_sessions").fetchone()[0] == 0
        assert db.execute("SELECT count(*) FROM llm_calls").fetchone()[0] == 0


def test_intake_stage_classifies_unnumbered_synthetic_files(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    checklist = tmp_path / "0.Necesar info.xlsx"
    _checklist(checklist)
    sources = [checklist]
    for name, document in _documents().items():
        path = tmp_path / name
        path.write_text(document.text, encoding="utf-8")
        sources.append(path)
    for source in sources:
        ws.set_slot(job, f"dossier/{source.name}", ws.add_file("synthetic", source))
    replay = ReplayProvider(STAGE_REPLAY)
    run = run_stage(ws, job, "audit_intake", lambda ctx: audit_intake(ctx, replay=replay))
    for _ in subscribe(ws, job):
        pass
    outcome = next(row for row in status(ws, job).runs if row["id"] == run)
    assert outcome["state"] == "ready", outcome["error"]
    with ws.connect() as db:
        path = ws.artifact_dir(db, job, "audit_intake", run) / "completeness.json"
    report = json.loads(path.read_text(encoding="utf-8"))
    assert report["unclassified"] == []
    assert {int(number): len(names) for number, names in report["received"].items() if names} == {
        2: 1,
        4: 1,
        5: 1,
        13: 1,
    }
    assert replay.calls == 12


def test_intake_keeps_unclassified_file_when_replay_fails(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    checklist = tmp_path / "0.Necesar info.xlsx"
    _checklist(checklist)
    loose = tmp_path / "loose.txt"
    loose.write_text("Unnumbered synthetic note", encoding="utf-8")
    for source in (checklist, loose):
        ws.set_slot(job, f"dossier/{source.name}", ws.add_file("synthetic", source))
    empty = tmp_path / "empty.json"
    empty.write_text(
        '{"source":"hand-authored","format":"openai-chat-completions","responses":[]}',
        encoding="utf-8",
    )
    replay = ReplayProvider(empty)
    run = run_stage(ws, job, "audit_intake", lambda ctx: audit_intake(ctx, replay=replay))
    for _ in subscribe(ws, job):
        pass
    outcome = next(row for row in status(ws, job).runs if row["id"] == run)
    assert outcome["state"] == "ready"
    assert "replay_invalid" in outcome["outcome"]
    with ws.connect() as db:
        path = ws.artifact_dir(db, job, "audit_intake", run) / "completeness.json"
    report = json.loads(path.read_text(encoding="utf-8"))
    assert report["unclassified"] == ["loose.txt"]


def test_spend_cap_stops_before_a_provider_call_and_is_resumable(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    replay = ReplayProvider(REPLAY)
    tools = IntakeTools(_documents(), {number: f"Item {number}" for number in range(1, 14)})
    context = AgentContext(ws, job, "intake", replay, "gemini-3.6-flash", "intake-v1")
    state = run_agent(context, "Classify each file", tools.tools(), Limits(20, 0.001))
    assert state.status == "spend_cap" and replay.calls == 0
    assert state.spend_cap_usd == 0.001
    state = run_agent(context, "Classify each file", tools.tools(), Limits(20, 1.0))
    assert state.status == "done" and replay.calls == 12
