"""Audit package checks reject broken links and leaked base content."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from ema.audit.base_package import package_issues, scrub_package
from ema.core.office.package import P, R, read_parts, write_parts


def test_package_issues_report_relationship_bookmark_and_identity_failures(tmp_path: Path) -> None:
    path = tmp_path / "audit.docx"
    document = Document()
    document.add_paragraph("Acme used 1234 MWh")
    document.save(path)
    parts = read_parts(path)
    relationships = etree.fromstring(parts["word/_rels/document.xml.rels"])
    existing = relationships[0]
    etree.SubElement(
        relationships,
        f"{{{P}}}Relationship",
        Id=existing.get("Id"),
        Type=f"{R}/hyperlink",
        Target="https://example.test",
        TargetMode="External",
    )
    etree.SubElement(
        relationships,
        f"{{{P}}}Relationship",
        Id="rId999",
        Type=f"{R}/image",
        Target="media/missing.png",
    )
    parts["word/_rels/document.xml.rels"] = etree.tostring(relationships)
    root = etree.fromstring(parts["word/document.xml"])
    paragraph = next(root.iter(qn("w:p")))
    for name, number in (("first", "8"), ("second", "8")):
        etree.SubElement(
            paragraph, qn("w:bookmarkStart"), attrib={qn("w:id"): number, qn("w:name"): name}
        )
    etree.SubElement(paragraph, qn("w:bookmarkEnd"), attrib={qn("w:id"): "9"})
    parts["word/document.xml"] = etree.tostring(root)
    parts["word/media/orphan.png"] = b"synthetic"
    write_parts(parts, path)

    issues = package_issues(path, ("Acme",), numeric_leftovers=("Acme used 1234 MWh",))
    assert any(issue.startswith("duplicate relationship ID") for issue in issues)
    assert any(issue.startswith("external relationship") for issue in issues)
    assert any(issue.startswith("dangling relationship") for issue in issues)
    assert any(issue.startswith("duplicate bookmark ID") for issue in issues)
    assert any(issue.startswith("unpaired bookmarks") for issue in issues)
    assert any(issue.startswith("orphan asset") for issue in issues)
    assert any(issue.startswith("numeric base paragraph remains") for issue in issues)
    assert any(issue.startswith("base identity remains") for issue in issues)


def test_scrub_removes_unreferenced_assets_and_document_properties(tmp_path: Path) -> None:
    path = tmp_path / "audit.docx"
    Document().save(path)
    parts = read_parts(path)
    parts["word/media/orphan.png"] = b"synthetic"
    write_parts(parts, path)
    scrub_package(path)
    cleaned = read_parts(path)
    assert "word/media/orphan.png" not in cleaned
    assert not any(name.startswith("docProps/") for name in cleaned)
