"""Repeating an audit unit remaps OOXML identifiers and supports omission."""

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ema.audit.base_units import UnitPlan, _resize
from ema.audit.inventory import BaseUnit


def _document() -> Document:
    document = Document()
    document.add_paragraph("Before")
    unit = document.add_paragraph("Process description")
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "7")
    start.set(qn("w:name"), "process")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "7")
    unit._p.insert(0, start)
    unit._p.append(end)
    document.add_paragraph("After")
    return document


def test_resize_repeats_process_with_unique_bookmarks() -> None:
    document = _document()
    _resize(document, (BaseUnit("process", (), (1, 2)),), 3)
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Before",
        "Process description",
        "Process description",
        "Process description",
        "After",
    ]
    starts = [node.get(qn("w:id")) for node in document.element.iter(qn("w:bookmarkStart"))]
    ends = [node.get(qn("w:id")) for node in document.element.iter(qn("w:bookmarkEnd"))]
    assert len(set(starts)) == 3
    assert sorted(starts) == sorted(ends)


def test_resize_can_remove_unit_and_reject_missing_prototype() -> None:
    document = _document()
    _resize(document, (BaseUnit("process", (), (1, 2)),), 0)
    assert [paragraph.text for paragraph in document.paragraphs] == ["Before", "After"]
    with pytest.raises(ValueError, match="no prototype"):
        _resize(document, (), 1)


@pytest.mark.parametrize(
    "kwargs", [{"client_name": ""}, {"processes": -1}, {"carriers": frozenset({"nuclear"})}]
)
def test_unit_plan_rejects_invalid_input(kwargs: dict[str, object]) -> None:
    defaults = {
        "client_name": "Test",
        "processes": 1,
        "carriers": frozenset({"gas"}),
        "measured_panels": 0,
        "thermal_measurements": False,
        "equipment_tables": 0,
        "measures": 0,
    }
    with pytest.raises(ValueError):
        UnitPlan(**{**defaults, **kwargs})
