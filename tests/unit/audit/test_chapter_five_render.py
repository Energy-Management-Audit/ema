"""Synthetic chapter-five rendering replaces figure media and preserves fixed text."""

from __future__ import annotations

import hashlib
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from tests.audit_structure import add_audit_toc, number_audit_headings

import ema.audit.chapter_five_render as render_module
from ema.audit.catalogue import CATALOGUE
from ema.audit.chapter_five import (
    ChapterFivePlan,
    PlannedPanel,
    PlannedPhoto,
    PlannedReading,
    PlannedThermal,
)
from ema.audit.chapter_five_fixed import fixed_elements
from ema.audit.chapter_five_render import render_chapter_five
from ema.core.office.package import read_parts, xml


def _outline(paragraph, level: int) -> None:  # type: ignore[no-untyped-def]
    node = OxmlElement("w:outlineLvl")
    node.set(qn("w:val"), str(level))
    paragraph._p.get_or_add_pPr().append(node)


def _base(path: Path) -> None:
    document = Document()
    add_audit_toc(document)
    titles = {section.id: section.title for section in CATALOGUE}
    for section_id in (
        "ch5",
        "ch5.electric",
        "ch5.electric_fisa",
        "ch5.electric_rezultate",
        "ch5.electric_concluzii",
        "ch5.termic",
        "ch5.termic_fisa",
        "ch5.termic_rezultate",
        "ch6",
    ):
        heading = document.add_paragraph(titles[section_id])
        _outline(heading, 0 if section_id in {"ch5", "ch6"} else 1)
        if section_id == "ch5.electric":
            document.add_paragraph("Fixed method text")
            document.add_paragraph("Pentru realizarea măsurătorilor electrice – old instrument")
        elif section_id == "ch5.electric_concluzii":
            document.add_paragraph("Client-specific conclusion")
            document.add_paragraph("Fixed harmonics text")
        elif section_id == "ch5.termic":
            document.add_paragraph("Fixed thermal method")
    number_audit_headings(document)
    document.save(path)


def _model(path: Path, image: Path) -> None:
    document = Document()
    document.add_picture(str(image))
    document.add_paragraph("Fig. 5.1 Prototype caption")
    document.add_paragraph("U 12: 230 V", style="List Paragraph")
    document.save(path)


def test_figure_uses_source_image_and_removes_instrument(tmp_path: Path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    base = tmp_path / "base.docx"
    model = tmp_path / "model.docx"
    image = tmp_path / "display.png"
    output = tmp_path / "out.docx"
    Image.new("RGB", (100, 50), "white").save(image)
    _base(base)
    _model(model, image)
    monkeypatch.setenv("EMA_AUDIT_MEASUREMENT_SHEET_MODEL", str(model))
    monkeypatch.setenv("EMA_AUDIT_MEASUREMENT_PROTOTYPE", str(base))
    original_replace = render_module.replace_region

    def with_region_issue(*args, **kwargs):  # type: ignore[no-untyped-def]
        return replace(original_replace(*args, **kwargs), issues=["synthetic region issue"])

    monkeypatch.setattr(render_module, "replace_region", with_region_issue)
    plan = ChapterFivePlan(
        panels=[
            PlannedPanel(
                id="panel-1",
                label="Panel 1",
                device="Synthetic meter",
                photos=[
                    PlannedPhoto(
                        slot="visit/meter/Panel 1/display.png",
                        sha="a" * 64,
                        name="display.png",
                        caption="Valorile Tensiunii de fază",
                        display="voltage_ln",
                        readings=[
                            PlannedReading(
                                key="meter.panel-1.aaaaaaaa.voltage_ln.l1",
                                label="U1",
                                value="230.00",
                                unit="V",
                            )
                        ],
                        norm="voltage",
                        narrative_key=None,
                    )
                ],
            )
        ],
        thermal=[],
        visit_date="2026-09-27",
        client="Atelier Exemplu",
        missing_narratives=[],
    )
    report = render_chapter_five(
        base, output, plan, {plan.panels[0].photos[0].slot: image}, ("Forbidden Client",)
    )
    text = "\n".join(paragraph.text for paragraph in Document(str(output)).paragraphs)
    assert "Fixed method text" in text
    assert "Fixed harmonics text" in text
    assert "old instrument" not in text
    assert "Fig. 5.1" in text
    parts = read_parts(output)
    assert any(
        hashlib.sha256(content).digest() == hashlib.sha256(image.read_bytes()).digest()
        for name, content in parts.items()
        if name.startswith("word/media/")
    )
    assert report.numbers[0].number == "5.1"
    assert report.issues == ["synthetic region issue"]
    document = Document(str(output))
    assert any(
        paragraph.style.name == "List Paragraph" and "230,00 V" in paragraph.text
        for paragraph in document.paragraphs
    )
    picture = next(
        xml(parts, "word/document.xml").iter(
            "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
        )
    )
    assert abs(int(picture.get("cx")) / int(picture.get("cy")) - 2) < 0.01


def test_missing_instrument_label_rejects_partial_fixed_method(tmp_path: Path) -> None:
    prototype = tmp_path / "prototype.docx"
    _base(prototype)
    source = Document(prototype)
    paragraph = next(item for item in source.paragraphs if "old instrument" in item.text)
    paragraph.text = "Unlabelled instrument text"
    source.save(prototype)
    nodes, keys, issues = fixed_elements(
        Document(), prototype, ("Forbidden Client",), frozenset({"method"})
    )
    assert not nodes and not keys
    assert issues == ["fixed method instrument label missing in measurement prototype"]


def _photo(panel: str, sha: str) -> PlannedPhoto:
    return PlannedPhoto(
        slot=f"visit/meter/{panel}/display.png",
        sha=sha,
        name="display.png",
        caption="Valorile Tensiunii de fază",
        display="voltage_ln",
        readings=[PlannedReading(key="meter.x.voltage_ln.l1", label="U1", value="230", unit="V")],
        norm="voltage",
        narrative_key=None,
    )


def test_equipment_notes_render_in_place_and_only_when_present(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    base, model, image = tmp_path / "base.docx", tmp_path / "model.docx", tmp_path / "d.png"
    Image.new("RGB", (100, 50), "white").save(image)
    _base(base)
    _model(model, image)
    monkeypatch.setenv("EMA_AUDIT_MEASUREMENT_SHEET_MODEL", str(model))
    monkeypatch.setenv("EMA_AUDIT_MEASUREMENT_PROTOTYPE", str(base))
    plan = ChapterFivePlan(
        panels=[
            PlannedPanel(
                id="panel-1", label="Panel 1", device="Meter", photos=[_photo("Panel 1", "a" * 64)]
            ),
            PlannedPanel(
                id="panel-2", label="Panel 2", device="Meter", photos=[_photo("Panel 2", "b" * 64)]
            ),
        ],
        thermal=[
            PlannedThermal(
                slot="visit/thermal/c.png", sha="c" * 64, name="c.png", component="Comp C"
            ),
            PlannedThermal(
                slot="visit/thermal/d.png", sha="d" * 64, name="d.png", component="Comp D"
            ),
        ],
        visit_date="2026-09-27",
        client="Atelier Exemplu",
        missing_narratives=[],
        narratives={
            "narrative.ch5.equipment.panel-1": "Tabloul alimentează hala.",
            "narrative.ch5.equipment.thermal:cccccccc": "Componenta produce căldura.",
            # A rejected note reaches the plan as None.
            "narrative.ch5.equipment.thermal:dddddddd": None,
        },
    )
    images = {key: image for key in ("a" * 64, "b" * 64, "c" * 64, "d" * 64)}

    def texts(rendered: ChapterFivePlan, name: str) -> list[str]:
        output = tmp_path / name
        render_chapter_five(base, output, rendered, images, ("Forbidden Client",))
        # Table-of-contents lines carry a tab; this base numbers every body paragraph there.
        paragraphs = Document(str(output)).paragraphs
        return [paragraph.text for paragraph in paragraphs if "\t" not in paragraph.text]

    noted = texts(plan, "noted.docx")
    plain = texts(plan.model_copy(update={"narratives": {}}), "plain.docx")
    first = noted.index("Fișa de măsurători electroenergetice pentru Panel 1, la data 2026-09-27.")
    assert noted[first + 1] == "Importanța echipamentului: Tabloul alimentează hala."
    second = noted.index("Fișa de măsurători electroenergetice pentru Panel 2, la data 2026-09-27.")
    assert noted[second + 1] == plain[plain.index(noted[second]) + 1]
    caption_c = next(i for i, text in enumerate(noted) if text.endswith("a) Termografierea Comp C"))
    caption_d = next(i for i, text in enumerate(noted) if text.endswith("b) Termografierea Comp D"))
    note_c = noted.index("Importanța echipamentului: Componenta produce căldura.")
    assert noted[note_c + 1 : caption_c] == [""]
    assert noted[caption_c + 1 : caption_d] == [""]
    assert [text for text in noted if not text.startswith("Importanța echipamentului: ")] == plain
    assert sum(text.startswith("Importanța echipamentului: ") for text in noted) == 2
