"""Review regressions for heading case, catalogue metadata and readable evidence."""

from dataclasses import replace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ema.audit.base_cleanup import clean_base
from ema.audit.catalogue import CATALOGUE
from ema.audit.chapter_readiness import empty_chapters
from ema.audit.heading_titles import MARKER, heading_blocks, heading_snapshot, title_changes
from ema.audit.headings import Heading, MappedHeading
from ema.core.errors import EmaError
from ema.core.office.blocks import Paragraph, Prototypes
from ema.core.review.section_transition import SectionState, Status


@pytest.mark.parametrize(
    "prototype", ["UPPER CASE", f"UPPER CASE {MARKER}", "Sentence case", "lower case"]
)
def test_rendered_headings_keep_the_named_level_three_prototype_case(prototype):
    document = Document()
    elements = {
        "heading:ch4.echiv_pv": document.add_paragraph(prototype)._p,
        "heading:ch4.bilant_real": document.add_paragraph("lower base title")._p,
    }
    blocks = [
        Paragraph("heading:ch4.echiv_pv", ["New heading"]),
        Paragraph("heading:ch4.bilant_real", [f"bilanțul energetic real {MARKER}"]),
    ]
    rendered = heading_blocks(blocks, Prototypes(elements, 4))
    expected = {
        "UPPER CASE": "NEW HEADING",
        f"UPPER CASE {MARKER}": "NEW HEADING",
        "Sentence case": "New heading",
        "lower case": "new heading",
    }
    assert rendered[0].segments == [expected[prototype]]
    assert rendered[1].segments == [f"BILANȚUL ENERGETIC REAL {MARKER}"]
    assert blocks[0].segments == ["New heading"]


def test_intro_content_is_catalogue_driven(monkeypatch):
    chapter = next(section for section in CATALOGUE if section.id == "ch2")
    monkeypatch.setattr(
        "ema.audit.chapter_readiness.CATALOGUE",
        tuple(
            replace(section, has_intro_content=True) if section == chapter else section
            for section in CATALOGUE
        ),
    )
    states = [
        SectionState(section.id, Status.DONE if section == chapter else Status.NA)
        for section in CATALOGUE
    ]
    assert not any(
        issue.message.startswith(chapter.title + ":") for issue in empty_chapters(states)
    )
    assert next(section for section in CATALOGUE if section.id == "ch1").has_intro_content


def test_fixed_chapter_marked_na_still_counts_as_empty():
    states = [SectionState(section.id, Status.NA) for section in CATALOGUE]
    for chapter in (
        section for section in CATALOGUE if section.parent is None and section.kind == "fixed"
    ):
        assert any(
            issue.code == "chapter_empty"
            and issue.message == f"{chapter.title}: capitolul nu are conținut"
            for issue in empty_chapters(states)
        )


def test_cover_without_extent_raises_an_ema_error():
    document = Document()
    inline = OxmlElement("wp:inline")
    text = OxmlElement("w:t")
    text.text = "AUDIT ENERGETIC"
    inline.append(text)
    document.add_paragraph().add_run()._r.append(inline)
    with pytest.raises(EmaError) as error:
        clean_base(document)
    assert error.value.code == "cover_extent_missing"
    assert "wp:extent" in error.value.detail


def test_digest_lists_actual_heading_text_and_corrected_measurement_catalogue_titles():
    document = Document()
    paragraph = document.add_paragraph("4.7. bilanțul energetic real")
    item = MappedHeading(Heading(1, paragraph.text, (), 0), "ch4.bilant_real")
    before = heading_snapshot(document, [(item, 0, 1)])
    paragraph.runs[0].text = "4.7. BILANȚUL ENERGETIC REAL"
    changes = title_changes(before)
    assert (
        changes[0] == "G2/G3 ch4.bilant_real heading text: "
        "4.7. bilanțul energetic real → 4.7. BILANȚUL ENERGETIC REAL"
    )
    corrected = {
        section.id: section.title
        for section in CATALOGUE
        if section.chapter == 5 and section.id.endswith(("_fisa", "_rezultate"))
    }
    assert set(corrected) == {
        "ch5.electric_fisa",
        "ch5.electric_rezultate",
        "ch5.termic_fisa",
        "ch5.termic_rezultate",
    }
    for section_id, title in corrected.items():
        assert (
            f"G1 {section_id} catalogue title: {title.replace('măsurător', 'măsurator')} → {title}"
            in changes
        )
    assert "MĂSURĂTORI" in corrected["ch5.electric_fisa"].upper()
    assert paragraph._p.find(qn("w:r")) is not None
