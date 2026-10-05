"""The 3.1.x process units of a draft render: a pooled draft and an empty unit (#155 D2)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from tests.unit.audit.test_draft_checks import _fact

from ema.audit.catalogue import CATALOGUE
from ema.audit.draft_render import render_section
from ema.audit.draft_schema import DraftText, SectionDraft

TITLES = {section.id: section.title for section in CATALOGUE}
UNITS = ("Alpha", "Beta", "Gamma")
KEYS = ("audit.process_sections", "audit.process_sections.2")
FACTS = {key: _fact(key, f"etapa {index}") for index, key in enumerate(KEYS, 1)}


def _base(path: Path) -> Path:
    """Three copies of her unit text, each with base sentences, a table and a figure caption."""
    document = Document()
    document.add_paragraph(TITLES["ch3"], style="Heading 1")
    document.add_paragraph(TITLES["ch3.flux"], style="Heading 2")
    document.add_paragraph("[de completat]", style="Body Text")
    for name in UNITS:
        document.add_paragraph(f"DESCRIEREA SECȚIEI {name}", style="Heading 3")
        document.add_paragraph("Secția [de completat] are o linie.", style="Body Text")
        document.add_paragraph("Piesele trec prin [de completat].", style="Body Text")
        document.add_paragraph("Tabelul 3. [de completat]", style="Caption")
        document.add_table(rows=1, cols=2).rows[0].cells[0].text = "[de completat]"
        document.add_paragraph("Fig. 4. [de completat]", style="Caption")
    document.add_paragraph(TITLES["ch4"], style="Heading 1")
    document.save(str(path))
    return path


def _render(tmp_path: Path, draft: SectionDraft) -> Document:
    output = tmp_path / "out.docx"
    render_section(_base(tmp_path / "base.docx"), output, draft, FACTS, (), job="synthetic")
    return Document(str(output))


def _paragraph(key: str, unit: int | None) -> DraftText:
    return DraftText(text=f"proces {{{{c:{key}}}}}.", fact_ids=[key], unit=unit)


def _after(texts: list[str], heading: str) -> list[str]:
    """The paragraphs from the heading up to the next heading of the base."""
    start = texts.index(heading)
    following = [
        index
        for index, text in enumerate(texts)
        if index > start and (text.startswith("DESCRIEREA") or text == TITLES["ch4"])
    ]
    return texts[start + 1 : following[0]]


def test_a_pooled_draft_is_unit_one_alone(tmp_path: Path) -> None:
    paragraphs = [_paragraph(KEYS[0], 1), _paragraph(KEYS[1], None)]
    document = _render(
        tmp_path, SectionDraft(section="ch3.process", status="drafted", paragraphs=paragraphs)
    )
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert _after(texts, "DESCRIEREA SECȚIEI Alpha")[:2] == ["proces.", "proces."]
    # Units 2..N go whole: no heading, no base text, no table.
    assert not any("Beta" in text or "Gamma" in text for text in texts)
    assert len(document.tables) == 1
    assert texts[-1] == TITLES["ch4"]


def test_an_empty_unit_is_its_heading_and_one_marker(tmp_path: Path) -> None:
    paragraphs = [_paragraph(KEYS[0], 1), _paragraph(KEYS[1], 2)]
    document = _render(
        tmp_path, SectionDraft(section="ch3.process", status="drafted", paragraphs=paragraphs)
    )
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert _after(texts, "DESCRIEREA SECȚIEI Gamma") == ["[de completat]"]
    # Alpha and Beta keep their base table; Gamma keeps none of its base content.
    assert len(document.tables) == 2


def test_a_missing_process_draft_marks_each_unit_once(tmp_path: Path) -> None:
    draft = SectionDraft(
        section="ch3.process", status="missing", missing_fact_ids=["audit.equipment"]
    )
    document = _render(tmp_path, draft)
    texts = [paragraph.text for paragraph in document.paragraphs]
    for name in UNITS:
        assert _after(texts, f"DESCRIEREA SECȚIEI {name}") == ["[de completat]"]
    assert not document.tables
