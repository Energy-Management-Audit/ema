"""A carrier with no quantity is not used unless its spend says so; the item names the place."""

from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from tests.workspace_jobs import create_job

from ema.audit.carrier_gap import blocked_message
from ema.audit.chapter_four_blocks import chapter_four_blocks
from ema.audit.chapter_four_sentences import sentence_plan
from ema.audit.content_checks import content_issues
from ema.audit.read import read_dossier
from ema.audit.render_dataset import reviewed_dataset
from ema.core.office.blocks import Missing
from ema.core.review import decide, fields
from ema.core.review.models import Field
from ema.core.workspace import Workspace
from ema.energy_data.calc import shares, tep_total
from ema.energy_data.carriers import Carrier
from ema.energy_data.factors import AUDIT_FACTORS_2026, FACTORS_2026
from ema.energy_data.model import CarrierSeries, EnergyDataset, Reading

MONTHS = (
    "Ianuarie",
    "Februarie",
    "Martie",
    "Aprilie",
    "Mai",
    "Iunie",
    "Iulie",
    "August",
    "Septembrie",
    "Octombrie",
    "Noiembrie",
    "Decembrie",
)
FILE = "necesar-sintetic.xlsx"
LPG = Carrier.lpg


def _block(sheet: object, row: int, label: str, unit: str, years: dict[int, float | None]) -> int:
    sheet.cell(row, 1, label)  # type: ignore[attr-defined]
    row += 1
    for year, month_value in years.items():
        for column, value in enumerate((year, *MONTHS, "Total"), 1):
            sheet.cell(row, column, value)  # type: ignore[attr-defined]
        sheet.cell(row + 1, 1, f"[{unit}]")  # type: ignore[attr-defined]
        sheet.cell(row + 2, 1, "[tep]")  # type: ignore[attr-defined]
        if month_value is not None:
            for month in range(12):
                sheet.cell(row + 1, 2 + month, month_value)  # type: ignore[attr-defined]
            sheet.cell(row + 1, 14, month_value * 12)  # type: ignore[attr-defined]
        else:
            sheet.cell(row + 1, 14, 0)  # type: ignore[attr-defined]
        row += 3
    return row + 1


def _necesar(path: Path, costs: dict[str, dict[int, int]]) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Cons energetice"
    row = _block(
        sheet, 1, "Consum energie electrica din SEN", "MWh", {2023: 10, 2024: 10, 2025: 10}
    )
    row = _block(sheet, row, "Consum motorina", "tone", {2023: None, 2024: 2, 2025: None})
    _block(sheet, row, "Consum GPL", "tone", {2023: None, 2024: None, 2025: None})
    economic = book.create_sheet("Cifre economice")
    economic.cell(3, 1, "Anul")
    for column, year in enumerate((2023, 2024, 2025), 2):
        economic.cell(3, column, year)
    for line, (label, by_year) in enumerate(costs.items(), 4):
        economic.cell(line, 1, f"Cheltuieli {label} [lei]")
        for year, cost in by_year.items():
            economic.cell(line, 2 + year - 2023, cost)
    book.save(path)
    return path


def _job(tmp_path: Path, costs: dict[str, dict[int, int]]) -> tuple[Workspace, str]:
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2025)
    path = _necesar(tmp_path / "source.xlsx", costs)
    read_dossier(ws, job, path)
    ws.set_slot(job, "dossier/necesar", ws.add_file("synthetic", path), original_name=FILE)
    return ws, job


def _blocked(ws: Workspace, job: str) -> list[str]:
    with ws.connect() as db:
        return [i.message for i in content_issues(db, job) if i.code == "data_total_blocked"]


def _number(key: str, value: int, *, review: str = "pending", unit: str = "t") -> Field:
    return Field.model_validate(
        {
            "id": key,
            "job_id": "job",
            "key": key,
            "label": key,
            "value_type": "number",
            "unit": unit,
            "value": Decimal(value),
            "state": "extracted",
            "presence": "found",
            "review": review,
        }
    )


def test_absent_quantity_and_no_cost_is_not_used_and_totals_compute(tmp_path: Path) -> None:
    ws, job = _job(tmp_path, {})
    assert _blocked(ws, job) == []
    dataset = EnergyDataset(
        (2023,),
        {
            Carrier.electricity_grid: {2023: CarrierSeries(annual=Reading(10, "MWh"))},
            LPG: {2023: CarrierSeries(annual=Reading(0, "t"))},
            Carrier.diesel: {2023: CarrierSeries(annual=Reading(None, "t"))},
        },
    )
    result = reviewed_dataset(dataset, fields(ws, job))
    assert set(result.carriers) == {Carrier.electricity_grid}
    assert tep_total(result, AUDIT_FACTORS_2026, 2023).value is not None


def test_zero_quantity_with_positive_cost_blocks_the_total(tmp_path: Path) -> None:
    ws, job = _job(tmp_path, {"GPL": {2023: 12345}})
    messages = _blocked(ws, job)
    assert len(messages) == 1
    assert messages[0].startswith("Totalul de energie din 2023 lipseşte: cantitatea de GPL")


def test_rejected_quantity_is_marked_not_used(tmp_path: Path) -> None:
    ws, job = _job(tmp_path, {"GPL": {2023: 12345}})
    total = next(f for f in fields(ws, job) if f.key == "carrier.lpg.2023")
    decide(ws, job, total.id, "reject", total.revision, "user")
    assert _blocked(ws, job) == []


def test_all_rejected_month_fields_mark_not_used_but_one_is_not_enough() -> None:
    dataset = EnergyDataset((2024,), {LPG: {2024: CarrierSeries(annual=Reading(None, "t"))}})
    cost = _number("audit.economics.lpg_costs_lei.2024", 500, unit="lei")
    rejected = [_number(f"carrier.lpg.2024.{m:02d}", 1, review="rejected") for m in (1, 2)]
    assert LPG not in reviewed_dataset(dataset, [*rejected, cost]).carriers
    pending = _number("carrier.lpg.2024.02", 1)
    assert LPG in reviewed_dataset(dataset, [rejected[0], pending, cost]).carriers


def test_carrier_used_in_one_year_is_excluded_in_the_others(tmp_path: Path) -> None:
    ws, job = _job(tmp_path, {})
    dataset = EnergyDataset(
        (2023, 2024, 2025),
        {
            Carrier.diesel: {
                2023: CarrierSeries(annual=Reading(0, "t")),
                2024: CarrierSeries(annual=Reading(2, "t")),
                2025: CarrierSeries(annual=Reading(None, "t")),
            }
        },
    )
    assert set(reviewed_dataset(dataset, fields(ws, job)).carriers[Carrier.diesel]) == {2024}
    assert _blocked(ws, job) == []


def test_message_names_file_sheet_row_and_cost_cell(tmp_path: Path) -> None:
    ws, job = _job(tmp_path, {"GPL": {2023: 12345}})
    assert _blocked(ws, job) == [
        "Totalul de energie din 2023 lipseşte: cantitatea de GPL nu este completată în "
        f"«{FILE}», foaia «Cons energetice», rândul 25 (lunile goale, total 0); societatea are "
        f"cheltuieli cu GPL de 12.345,00 lei în 2023 («{FILE}», foaia «Cifre economice», "
        "celula B4). Completaţi cantitatea din facturi sau marcaţi combustibilul ca neutilizat."
    ]


def test_message_drops_a_missing_place_and_falls_back_without_cost(tmp_path: Path) -> None:
    ws, job = _job(tmp_path, {"GPL": {2023: 12345}})
    with ws.connect() as db:
        db.execute("UPDATE slot_versions SET original_name=NULL WHERE job_id=?", (job,))
    assert _blocked(ws, job) == [
        "Totalul de energie din 2023 lipseşte: cantitatea de GPL nu este completată; societatea "
        "are cheltuieli cu GPL de 12.345,00 lei în 2023. Completaţi cantitatea din facturi sau "
        "marcaţi combustibilul ca neutilizat."
    ]


def test_plain_message_when_there_is_no_cost_or_the_quantity_is_partly_filled(
    tmp_path: Path,
) -> None:
    ws, job = _job(tmp_path, {"GPL": {2023: 12345}})
    plain = "Totalul de energie din 2023 lipseşte: completaţi cantitatea de GPL."
    by_key = {f.key: f for f in fields(ws, job)}
    filled = CarrierSeries({1: Reading(3, "t")})
    with ws.connect() as db:
        assert blocked_message(db, job, by_key, filled, LPG, 2023) == plain
        assert blocked_message(db, job, by_key, CarrierSeries(), LPG, 2025) == plain.replace(
            "2023", "2025"
        )


def _ch4_dataset(*, lpg_cost: bool) -> tuple[EnergyDataset, list[Field]]:
    dataset = EnergyDataset(
        (2025,),
        {
            Carrier.electricity_grid: {2025: CarrierSeries(annual=Reading(100, "MWh"))},
            Carrier.natural_gas: {2025: CarrierSeries(annual=Reading(200, "MWh"))},
            LPG: {2025: CarrierSeries(annual=Reading(0, "t"))},
        },
        {"product": {2025: CarrierSeries(annual=Reading(1_000_000, "kg"))}},
        {"product": "kg"},
        {2025: Reading(10_000_000, "lei")},
    )
    cost = [
        Field.model_validate(
            {
                "id": "cost",
                "job_id": "job",
                "key": "audit.economics.lpg_costs_lei.2025",
                "label": "Cheltuieli cu GPL 2025",
                "value_type": "number",
                "unit": "lei",
                "value": Decimal(900),
                "state": "supplied",
                "presence": "found",
            }
        )
    ]
    return dataset, cost if lpg_cost else []


def _missing(dataset: EnergyDataset) -> list[Missing]:
    return [b for b in chapter_four_blocks(dataset, FACTORS_2026) if isinstance(b, Missing)]


def test_ch4_with_a_not_used_carrier_computes_totals_shares_and_conclusions() -> None:
    raw, no_cost = _ch4_dataset(lpg_cost=False)
    unused = reviewed_dataset(raw, no_cost)
    assert LPG not in unused.carriers
    total = tep_total(unused, FACTORS_2026, 2025)
    assert total.value is not None
    assert all(part.value is not None for part in shares(unused, FACTORS_2026, 2025).values())
    without = replace(raw, carriers={c: y for c, y in raw.carriers.items() if c != LPG})
    assert _missing(unused) == _missing(without)
    assert sentence_plan(unused, FACTORS_2026).sections.get("ch4.concluzii")


def test_ch4_with_a_used_carrier_without_quantity_shows_the_gap() -> None:
    raw, cost = _ch4_dataset(lpg_cost=True)
    used = reviewed_dataset(raw, cost)
    assert LPG in used.carriers
    assert tep_total(used, FACTORS_2026, 2025).value is None
    assert len(_missing(used)) > len(_missing(reviewed_dataset(raw, [])))
