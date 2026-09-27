"""Audit-only PDF intake classification."""

from __future__ import annotations

import json
from pathlib import Path

import pypdfium2
from openpyxl import Workbook
from tests.unit.test_invoice_pdf import _pdf

from ema.audit.intake import audit_intake
from ema.audit.pdf_state import pdf_state
from ema.core.jobs import create_job, run_stage, subscribe
from ema.core.workspace import Workspace


def test_blank_pdf_is_scanned(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    document = pypdfium2.PdfDocument.new()
    document.new_page(595, 842)
    document.save(path)
    assert pdf_state(path) == "scanned"


def test_embedded_text_pdf_is_text(tmp_path: Path) -> None:
    assert pdf_state(_pdf(tmp_path / "text.pdf")) == "text"


def test_password_error_is_protected(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "protected.pdf"

    def encrypted(_path: Path) -> None:
        raise pypdfium2.PdfiumError("Incorrect password error", err_code=4)

    monkeypatch.setattr("ema.audit.pdf_state.pypdfium2.PdfDocument", encrypted)
    assert pdf_state(path) == "protected"


def _intake_with_pdf(tmp_path: Path, source: Path) -> tuple[Workspace, str, str]:
    ws = Workspace(tmp_path / "ws")
    job = create_job(ws, "audit", "synthetic", 2026)
    checklist = tmp_path / "0. Necesar info.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "diverse"
    for number in range(1, 14):
        sheet.cell(number, 1, number)
        sheet.cell(number, 2, f"Document {number}")
    book.save(checklist)
    for path in (checklist, source):
        ws.set_slot(job, f"dossier/{path.name}", ws.add_file("synthetic", path))
    run = run_stage(ws, job, "intake", audit_intake)
    for _ in subscribe(ws, job):
        pass
    return ws, job, run


def _report(ws: Workspace, job: str, run: str) -> dict[str, object]:
    with ws.connect() as db:
        path = ws.artifact_dir(db, job, "intake", run) / "completeness.json"
    return json.loads(path.read_text(encoding="utf-8"))


def test_audit_intake_marks_image_pdf_scanned_without_failure(tmp_path: Path) -> None:
    source = tmp_path / "1. Blank.pdf"
    document = pypdfium2.PdfDocument.new()
    document.new_page(595, 842)
    document.save(source)
    ws, job, run = _intake_with_pdf(tmp_path, source)
    report = _report(ws, job, run)
    record = next(item for item in report["files"] if item["slot"] == "dossier/1. Blank.pdf")
    assert record["status"] == "scanned"
    assert record["error_code"] is None


def test_audit_intake_records_protected_error(tmp_path: Path, monkeypatch) -> None:
    source = _pdf(tmp_path / "1. Protected.pdf")
    monkeypatch.setattr("ema.audit.intake.pdf_state", lambda _path: "protected")
    ws, job, run = _intake_with_pdf(tmp_path, source)
    report = _report(ws, job, run)
    record = next(item for item in report["files"] if item["slot"] == "dossier/1. Protected.pdf")
    assert record["status"] == "protected"
    assert record["error_code"] == "pdf_protected"
