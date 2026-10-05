"""The equipment rows are found by their column labels and recorded with their cells."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from tests.workspace_jobs import create_job

from ema.audit.chapter_tables_data import boiler_rows
from ema.audit.draft_render import rendered_value
from ema.audit.read import read_dossier
from ema.core.review import fields
from ema.core.review.models import Cell, Evidence, Field
from ema.core.workspace import Workspace
from ema.energy_data.necesar import parse_necesar_info
from ema.energy_data.necesar_equipment import read_equipment


def _transformer_facts(sheet) -> None:
    for row, label, value in (
        (15, "An fabricaţie", 1999),
        (16, "An punere în funcţiune", 2000),
        (17, "Tensiune [kV]", 0.4),
    ):
        sheet.cell(row, 3, label)
        for column in (4, 5):
            cell = sheet.cell(row, column, value)
            if row == 17:
                cell.number_format = "0.0"


def _necesar(path: Path) -> Path:
    """The sheets as a later client may send them: the boiler list under "Echipamente 2" and
    every header a few rows lower than usual."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Echipamente 2"
    for column, label in enumerate(
        ("Nr. Crt", "Denumire", "Proces de fabricație deservit", "Nr. buc", "An PIF"), 2
    ):
        sheet.cell(7, column, label)
    sheet.cell(7, 8, "Putere instalată")
    sheet.cell(7, 9, "Combustibil utilizat")
    sheet.cell(8, 3, "Tip echipament")
    sheet.cell(8, 8, "(kW)")
    for row, (name, process, count, year, power, fuel) in enumerate(
        (
            ("Centrala A", "Incalzire hala", 1, 2007, 500, "gaz natural"),
            ("Centrala B", "Apa calda", 2, 1997, 55, None),
            ("Compresor", "Aer comprimat", 1, 2010, 25, "energie electrică"),
        ),
        9,
    ):
        for column, value in zip(
            (2, 3, 4, 5, 6, 8, 9), (row - 8, name, process, count, year, power, fuel), strict=True
        ):
            sheet.cell(row, column, value)
    sheet.cell(12, 3, "Total")
    sheet.cell(12, 5, 4)
    forklifts = book.create_sheet("Echipamente 1")
    for column, label in enumerate(
        ("Nr.", "Denumire", "Comb.", "Greutate/inaltime", "An fabricaţie", "Ore de funcţionare"), 2
    ):
        forklifts.cell(5, column, label)
    forklifts.cell(6, 2, "Crt.")
    for column, value in enumerate((1, "MARCA : Marca X", "electric", "Q=2000KG", 2004, 14367), 2):
        forklifts.cell(7, column, value)
    forklifts.cell(8, 3, "TIP : ERE 220 / SERIA : 90120354")
    for column, value in enumerate((2, "MARCA : Marca Y", "propan", "Q=1600KG", 2008, 9673), 2):
        forklifts.cell(9, column, value)
    forklifts.cell(10, 3, "TIP : TFG 316 / SERIA : FN372205")
    forklifts.cell(13, 3, "Tip")
    forklifts.cell(13, 4, "Transformator uscat")
    forklifts.cell(13, 5, "Transformator ulei")
    forklifts.cell(14, 3, "Putere aparentă nominală [kVA]")
    forklifts.cell(14, 4, 1000)
    forklifts.cell(14, 5, 1000)
    _transformer_facts(forklifts)
    vehicles = book.create_sheet("Autovehicule")
    for column, label in enumerate(
        ("Denumire autovehicul", "Producător", "Tip", "Nr buc", "An fabricaţie", "Tip combustibil"),
        2,
    ):
        vehicles.cell(6, column, label)
    for column, value in enumerate(("Autoturism", "Marca Z", 308, 42, "2015-2016", "Diesel"), 2):
        vehicles.cell(8, column, value)
    book.save(path)
    return path


def test_rows_are_found_by_label_and_kept_apart(tmp_path: Path) -> None:
    equipment = read_equipment(parse_necesar_info(_necesar(tmp_path / "necesar.xlsx")))
    assert [row["name"].value for row in equipment.boilers] == [
        "Centrala A",
        "Centrala B",
        "Compresor",
    ]
    assert [row["count"].value for row in equipment.boilers] == [1, 2, 1]
    assert equipment.boilers[0]["resource"].value == "gaz natural"
    assert [row["name"].value for row in equipment.rows] == [
        "Centrala A",
        "Centrala B",
        "Compresor",
    ]
    assert equipment.boilers[0]["power"].unit == "kW"
    assert [row["name"].value for row in equipment.forklifts] == [
        "MARCA : Marca X",
        "MARCA : Marca Y",
    ]
    assert equipment.forklifts[1]["type"].value == "TIP : TFG 316 / SERIA : FN372205"
    # A model number the sheet fills as a number is a name, not a quantity.
    assert equipment.vehicles[0]["type"].value == "308"
    assert len(equipment.transformers) == 2
    assert equipment.transformers[1]["Putere aparentă nominală [kVA]"].value == 1000


def test_power_not_in_kilowatts_is_not_read(tmp_path: Path) -> None:
    path = _necesar(tmp_path / "necesar.xlsx")
    book = load_workbook(path)
    book["Echipamente 2"].cell(8, 8, "(CP)")
    book.save(path)
    equipment = read_equipment(parse_necesar_info(path))
    assert equipment.boilers and all("power" not in row for row in equipment.boilers)
    assert equipment.rows and all("power" not in row for row in equipment.rows)


def test_transformer_years_render_as_ungrouped_integers(tmp_path: Path) -> None:
    workspace = Workspace(tmp_path / "workspace")
    job = create_job(workspace, "audit", "synthetic", 2025)
    read_dossier(workspace, job, _necesar(tmp_path / "necesar.xlsx"))
    saved = {field.key: field for field in fields(workspace, job)}
    for number, key, expected in (
        (1, "an_fabricatie", "1999"),
        (2, "an_fabricatie", "1999"),
        (1, "an_punere_in_functiune", "2000"),
        (2, "an_punere_in_functiune", "2000"),
    ):
        field = saved[f"audit.transformer.{number}.{key}"]
        assert field.value_type == "year"
        assert rendered_value(field) == expected


def test_transformer_voltage_uses_source_precision_and_unit(tmp_path: Path) -> None:
    workspace = Workspace(tmp_path / "workspace")
    job = create_job(workspace, "audit", "synthetic", 2025)
    read_dossier(workspace, job, _necesar(tmp_path / "necesar.xlsx"))
    field = next(field for field in fields(workspace, job) if field.key.endswith("tensiune_kv"))
    assert field.value_type == "number"
    assert field.decimals == 1
    assert field.unit == "kV"
    assert rendered_value(field) == "0,4 kV"


def test_shifted_headers_fill_the_general_table_in_sheet_order(tmp_path: Path) -> None:
    path = _necesar(tmp_path / "necesar.xlsx")
    book = load_workbook(path)
    book["Echipamente 2"].insert_rows(7)
    book["Echipamente 1"].insert_rows(5)
    book.save(path)
    workspace = Workspace(tmp_path / "workspace")
    job = create_job(workspace, "audit", "synthetic", 2025)
    read_dossier(workspace, job, path)
    assert boiler_rows(fields(workspace, job)) == [
        ["Centrala A", "Incalzire hala", "1", "500", "gaz natural"],
        ["Centrala B", "Apa calda", "2", "55", None],
        ["Compresor", "Aer comprimat", "1", "25", "energie electrică"],
    ]


def test_read_records_each_row_with_its_cell(tmp_path: Path) -> None:
    workspace = Workspace(tmp_path / "workspace")
    job = create_job(workspace, "audit", "synthetic", 2025)
    read_dossier(workspace, job, _necesar(tmp_path / "necesar.xlsx"))
    saved: dict[str, Field] = {field.key: field for field in fields(workspace, job)}
    assert saved["audit.equipment_row.2.name"].value == "Centrala B"
    assert saved["audit.equipment_row.2.count"].value == 2
    assert saved["audit.equipment_row.1.power"].value == 500
    assert saved["audit.equipment_row.1.power"].unit == "kW"
    assert not any(key.startswith("audit.boiler.") for key in saved)
    assert saved["audit.equipment_row.1.resource"].value == "gaz natural"
    assert saved["audit.equipment_row.3.name"].value == "Compresor"
    assert saved["audit.forklift.2.fuel"].value == "propan"
    assert saved["audit.vehicle.1.type"].value == "308"
    assert saved["audit.vehicle.1.year"].value == "2015-2016"
    assert saved["audit.vehicle.1.count"].value == 42
    assert saved["audit.transformer.1.putere_aparenta_nominala_kva"].value == Decimal("1000")
    with workspace.connect() as db:
        row = db.execute(
            "SELECT data FROM evidence WHERE id=?",
            (saved["audit.equipment_row.2.name"].evidence[0],),
        ).fetchone()
    locator = Evidence.model_validate_json(row["data"]).locator
    assert isinstance(locator, Cell)
    assert (locator.sheet, locator.ref) == ("Echipamente 2", "C10")
    with workspace.connect() as db:
        resource = db.execute(
            "SELECT data FROM evidence WHERE id=?",
            (saved["audit.equipment_row.1.resource"].evidence[0],),
        ).fetchone()
    fuel_locator = Evidence.model_validate_json(resource["data"]).locator
    assert isinstance(fuel_locator, Cell)
    assert (fuel_locator.sheet, fuel_locator.ref) == ("Echipamente 2", "I9")
