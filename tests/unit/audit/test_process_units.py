"""The 3.1.x process units (#111 A, D3): which passage belongs to which unit, and unit names."""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document
from tests.unit.audit.test_fill_extract import Scripted, log_events, sections, values
from tests.unit.audit.test_fill_stage import OPENAI_MODEL, use_provider
from tests.workspace_jobs import create_job

from ema.audit.catalogue_types import process_unit_name
from ema.audit.dossier import dossier_documents
from ema.audit.fill_extract import extract_facts
from ema.audit.fill_files import file_ids
from ema.audit.fill_stage import fill_sections
from ema.audit.fill_tools import FillDocument, FillTools
from ema.audit.process_units import (
    ProcessUnits,
    job_process_units,
    passage_units,
    process_units,
    unit_of,
)
from ema.core.errors import EmaError
from ema.core.review.fields import propose
from ema.core.review.models import Evidence, FieldSpec, TextLoc
from ema.core.workspace import Workspace


def evidence(sha: str, quote: str) -> Evidence:
    return Evidence(
        id=quote,
        provenance="document",
        file_sha=sha,
        locator=TextLoc(span=quote),
        method="questionnaire",
        retrieved_at=datetime(2026, 10, 4, tzinfo=UTC),
        quote=quote,
        highlight="exact",
    )


SCHEMES = [
    ("5.10. Flux ambalare.pdf", "c"),
    ("5.2. Flux solvent.pdf", "b"),
    ("5.1. Flux vopsire.pdf", "a"),
    ("5.1. Flux vopsire anexa.pdf", "a2"),
    ("autorizatie.pdf", "permit"),
]


def test_a_three_scheme_plan_maps_each_passage_by_its_source_file() -> None:
    units = process_units(SCHEMES)

    assert (units.source, units.count) == ("schemes", 3)
    # Scheme ids sort as numbers: 5.10 is the third unit, after 5.2.
    assert [unit_of(evidence(sha, "x"), units) for sha in ("a", "a2", "b", "c")] == [1, 1, 2, 3]
    # The permit's text is the ch3.flux overview: it belongs to no unit.
    assert unit_of(evidence("permit", "x"), units) is None


FISA = (
    "Fişa de date a instalaţiei",
    "Societatea are două fluxuri tehnologice.",
    "Flux 1: vopsire în câmp electrostatic",
    "Piesele sunt degresate şi vopsite.",
    "Piesele sunt uscate în cuptor.",
    "Flux 2: recuperarea solvenţilor",
    "Solventul uzat este distilat.",
    "Piesele sunt uscate în cuptor.",
    "Descrierea utilităţilor",
)


def fisa_unit(quote: str, paragraphs: tuple[str, ...] = FISA, sha: str = "fisa") -> int | None:
    units = process_units([("Fisa de date.docx", "fisa")], ("fisa", paragraphs))
    return unit_of(evidence(sha, quote), units)


def test_a_fisa_passage_belongs_to_the_one_flux_block_that_wholly_holds_its_quote() -> None:
    units = process_units([("Fisa de date.docx", "fisa")], ("fisa", FISA))

    assert (units.source, units.count) == ("fisa", 2)
    assert units.block(1) == FISA[2:5] and units.block(2) == FISA[5:]
    assert fisa_unit("Piesele sunt degresate şi vopsite.") == 1
    assert fisa_unit("Flux 2: recuperarea solvenţilor\nSolventul uzat este distilat.") == 2
    assert fisa_unit("Societatea are două fluxuri tehnologice.") is None
    assert fisa_unit("Un text din tabelul fişei.") is None
    assert fisa_unit("Piesele sunt degresate şi vopsite.", sha="permit") is None
    assert unit_of(evidence("a", "x"), ProcessUnits()) is None


def test_a_quote_across_a_block_end_or_in_two_blocks_belongs_to_no_unit() -> None:
    # Found first in block 1, it used to be block 1's; it is in block 2 as well.
    assert fisa_unit("Piesele sunt uscate în cuptor.") is None
    assert fisa_unit("Piesele sunt uscate în cuptor.\nFlux 2: recuperarea solvenţilor") is None
    assert fisa_unit("tehnologice.\nFlux 1: vopsire") is None


def test_a_word_line_break_inside_a_paragraph_keeps_each_quote_in_its_block() -> None:
    # A Word line break is a newline inside one paragraph: counting newlines put the last
    # paragraph of block 1 into block 2.
    paragraphs = (
        "Fişa de date a instalaţiei\nrevizia 2",
        "Flux 1: vopsire",
        "Piesele sunt vopsite.",
        "Flux 2: solvent",
        "Solventul este\ndistilat.",
        "Descrierea utilităţilor",
    )
    assert fisa_unit("Piesele sunt vopsite.", paragraphs) == 1
    assert fisa_unit("Solventul este\ndistilat.", paragraphs) == 2
    assert fisa_unit("Descrierea utilităţilor", paragraphs) == 2
    assert fisa_unit("revizia 2", paragraphs) is None


def write_fisa(path: Path, paragraphs: tuple[str, ...]) -> Path:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))
    return path


def test_an_ambiguous_fisa_passage_is_the_overview_and_logged(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    job_id = create_job(ws, "audit", "made-up", 2026)
    sha = ws.add_file("made-up", write_fisa(tmp_path / "Fisa de date.docx", FISA))
    ws.set_slot(job_id, "dossier/Fisa de date.docx", sha)
    spec = FieldSpec(key="audit.process_sections", label="Flux", value_type="text")
    for key, quote in (
        ("audit.process_sections", "Piesele sunt degresate şi vopsite."),
        ("audit.process_sections.2", "Piesele sunt uscate în cuptor."),
    ):
        keyed = spec.model_copy(update={"key": key})
        propose(ws, job_id, keyed, quote, [evidence(sha, quote)], state="extracted")

    assert passage_units(ws, job_id) == {
        "audit.process_sections": 1,
        "audit.process_sections.2": None,
    }
    (event,) = log_events(ws, job_id, "passage_unit_ambiguous")
    assert event["keys"] == ["audit.process_sections.2"]


def test_schemes_are_found_before_the_fisa_is_opened(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    job_id = create_job(ws, "audit", "made-up", 2026)
    broken = tmp_path / "Fisa de date.docx"
    broken.write_bytes(b"PK\x03\x04 not a document")
    ws.set_slot(job_id, "dossier/Fisa de date.docx", ws.add_file("made-up", broken))
    with pytest.raises(Exception):  # noqa: B017
        job_process_units(ws, job_id)
    scheme = tmp_path / "5.1. Flux vopsire.txt"
    scheme.write_text("Linia de vopsire", encoding="utf-8")
    ws.set_slot(job_id, "dossier/5.1. Flux vopsire.txt", ws.add_file("made-up", scheme))

    units = job_process_units(ws, job_id)

    assert (units.source, units.count) == ("schemes", 1)


SCHEME_TEXTS = {
    "5.1. Flux vopsire.txt": "Document aprobat la 01.01.2026\nLinia de vopsire\n"
    "Etapa 1. Piesele sunt degresate.",
    "5.2. Flux solvent.txt": "Recuperarea solvenţilor\nEtapa 1. Solventul este distilat.",
    "5.3. Flux ambalare.txt": "Etapa 1. Piesele sunt ambalate.",
    "autorizatie.txt": "Instalaţia are trei fluxuri tehnologice.",
}


def scheme_dossier(tmp_path: Path) -> tuple[Workspace, str]:
    ws = Workspace(tmp_path / "workspace")
    job_id = create_job(ws, "audit", "made-up", 2026)
    for name, text in SCHEME_TEXTS.items():
        source = tmp_path / name
        source.write_text(text, encoding="utf-8")
        ws.set_slot(job_id, f"dossier/{name}", ws.add_file("made-up", source))
    return ws, job_id


def passage(name: str, quote: str) -> dict[str, object]:
    return {"key": "audit.process_sections", "value": "", "file": name, "page": 1, "quote": quote}


def name(number: int, file: str, value: str, quote: str | None = None) -> dict[str, object]:
    key = process_unit_name(number)
    return {"key": key, "value": value, "file": file, "page": 1, "quote": quote or value}


def test_extraction_names_each_unit_only_from_its_own_files(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("EMA_AI_CLIENT_LIVE", "1")
    ws, job_id = scheme_dossier(tmp_path)
    ids = {name: file_id for file_id, name in file_ids(dossier_documents(ws, job_id)).items()}
    paint, solvent = ids["5.1. Flux vopsire.txt"], ids["5.2. Flux solvent.txt"]
    returned = [
        passage("autorizatie.txt", "Instalaţia are trei fluxuri tehnologice."),
        passage("5.1. Flux vopsire.txt", "Etapa 1. Piesele sunt degresate."),
        passage("5.2. Flux solvent.txt", "Etapa 1. Solventul este distilat."),
        # The title is the scheme's second line, not its first ("Document aprobat ...").
        name(1, paint, "Linia de vopsire"),
        # Verbatim, but from unit 2's scheme: never unit 3's name.
        name(3, solvent, "Recuperarea solvenţilor"),
    ]
    retried = {"facts": [name(2, solvent, "Recuperarea solvenţilor")], "missing": []}
    provider = Scripted([{"facts": returned, "missing": []}, retried])
    use_provider(monkeypatch, provider, OPENAI_MODEL)

    fill_sections(ws, job_id, ["ch3.process"])

    content = provider.sent[0][1]["content"]
    for number, files in ((1, paint), (2, solvent), (3, ids["5.3. Flux ambalare.txt"])):
        line = f"{process_unit_name(number)} — Fluxul tehnologic {number} – denumire"
        assert f"{line} — ch3.process — titlu ({files})\n" in content
    found = values(ws, job_id)
    assert found[process_unit_name(1)] == ("Linia de vopsire", "found")
    assert found[process_unit_name(2)] == ("Recuperarea solvenţilor", "found")
    # No verified proposal: the heading keeps its marker; nothing is taken from a first line.
    assert found[process_unit_name(3)][1] == "not_found"
    by_key = {key: found[key][0] for key in found if key.startswith("audit.process_sections")}
    units = {by_key[key]: unit for key, unit in passage_units(ws, job_id).items()}
    assert units == {
        "Instalaţia are trei fluxuri tehnologice.": None,
        "Etapa 1. Piesele sunt degresate.": 1,
        "Etapa 1. Solventul este distilat.": 2,
    }


def test_a_fisa_unit_name_is_quoted_from_its_own_flux_block_and_page(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    job_id = create_job(ws, "audit", "made-up", 2026)
    documents = {"Fisa.docx": FillDocument("Fisa.docx", "\n".join(FISA), file_sha="fisa")}
    units = process_units([("Fisa.docx", "fisa")], ("fisa", FISA))
    tools = FillTools(ws, job_id, "ch3.process", documents, units)

    def rejected(args: dict[str, object], code: str, section: str = "ch3.process") -> None:
        with pytest.raises(EmaError) as error:
            FillTools(ws, job_id, section, documents, units).record_fact(args)
        assert error.value.code == code

    title = "Flux 2: recuperarea solvenţilor"
    record = {"key": process_unit_name(2), "value": title, "name": "F1", "quote": title}
    rejected(record, "page_missing")
    rejected({**record, "key": process_unit_name(1), "page": 1}, "fact_unit")
    rejected({**record, "page": 1}, "fact_section", section="ch3.flux")
    rejected({**record, "quote": "Piesele sunt uscate în cuptor.", "page": 1}, "value_unverified")
    # A name is a non-empty text: an empty or numeric value is rejected, never raised unchecked.
    for value in ("", " ", 1):
        rejected({**record, "value": value, "page": 1}, "fact_type")
    assert process_unit_name(2) not in values(ws, job_id)

    tools.record_fact({**record, "page": 1})

    assert values(ws, job_id)[process_unit_name(2)] == (title, "found")


def test_a_rerun_with_fewer_units_drops_the_stale_names(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("EMA_AI_CLIENT_LIVE", "1")
    ws, job_id = scheme_dossier(tmp_path)
    documents = dossier_documents(ws, job_id)
    ids = {name: file_id for file_id, name in file_ids(documents).items()}
    answer = [name(1, ids["5.1. Flux vopsire.txt"], "Linia de vopsire")]
    answer.append(name(2, ids["5.2. Flux solvent.txt"], "Recuperarea solvenţilor"))
    use_provider(monkeypatch, Scripted([{"facts": answer, "missing": []}]), OPENAI_MODEL)
    fill_sections(ws, job_id, ["ch3.process"])
    assert values(ws, job_id)[process_unit_name(2)][1] == "found"
    first = documents["5.1. Flux vopsire.txt"]

    extract_facts(
        ws,
        job_id,
        sections("ch3.process"),
        documents=documents,
        provider=Scripted([{"facts": answer[:1], "missing": []}]),
        model_id=OPENAI_MODEL,
        artifacts=ws.root / "artifacts",
        client_live=True,
        units=process_units([("5.1. Flux.txt", first.sha)]),
    )

    found = values(ws, job_id)
    assert found[process_unit_name(1)] == ("Linia de vopsire", "found")
    assert found[process_unit_name(2)][1] == found[process_unit_name(3)][1] == "not_found"
