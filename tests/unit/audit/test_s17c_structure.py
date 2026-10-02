"""S17c corrections preserve source text and reject missing heading prototypes."""

from dataclasses import replace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips
from lxml import etree

from ema.audit.base_cleanup import CORRECT_LAW, PASTE_SLIP, clean_base
from ema.audit.base_numbering import effective_indent, insert_chapter_numbering, printed_numbers
from ema.audit.catalogue import CATALOGUE
from ema.audit.chapter_four import _prototypes
from ema.audit.chapter_readiness import empty_chapters
from ema.audit.headings import Heading, MappedHeading
from ema.core.errors import EmaError
from ema.core.review.section_transition import SectionState, Status


def _node(tag, **attributes):
    node = OxmlElement(tag)
    for key, value in attributes.items():
        node.set(key, value)
    return node


def test_cover_wordart_anchor_fits_margin_text_area(monkeypatch):
    document = Document()
    inline = _node("wp:inline")
    inline.append(_node("wp:extent", cx="8413750", cy="571500"))
    inline.append(_node("wp:docPr", id="1", name="WordArt"))
    graphic = _node("a:graphic")
    transform = _node("a:xfrm", rot="5400000")
    transform.append(_node("a:ext", cx="8413750", cy="571500"))
    graphic.append(transform)
    text = _node("w:t")
    text.text = "AUDIT ENERGETIC"
    graphic.append(text)
    inline.append(graphic)
    drawing = _node("w:drawing")
    drawing.append(inline)
    document.add_paragraph().add_run()._r.append(drawing)
    monkeypatch.setattr("ema.audit.base_cleanup.heading_spans_document", lambda _: [])
    changes = clean_base(document)
    anchor = document.element.xpath(".//wp:anchor")[0]
    extent = anchor.find(qn("wp:extent"))
    section = document.sections[0]
    assert int(extent.get("cx")) <= int(
        section.page_height - section.top_margin - section.bottom_margin
    )
    assert int(extent.get("cy")) == 571500
    assert anchor.find(qn("wp:wrapNone")) is not None
    assert anchor.find(qn("wp:positionV")).get("relativeFrom") == "margin"
    displacement = (int(extent.get("cx")) - int(extent.get("cy"))) // 2
    assert anchor.find(qn("wp:positionV"))[0].text == str(displacement)
    assert anchor.find(qn("wp:positionH"))[0].text == str(-displacement)
    assert transform[0].get("cx") == extent.get("cx")
    assert transform.get("rot") == "5400000"
    assert len(changes) == 1 and changes[0].startswith("C1 WordArt XML:")
    assert f"cx 8413750 → {extent.get('cx')}" in changes[0]
    assert clean_base(document) == []


def test_paste_slip_is_corrected_across_runs_without_losing_format(monkeypatch):
    document = Document()
    paragraph = document.add_paragraph()
    text = "Public legislation: " + PASTE_SLIP + " Next sentence."
    cut = text.index("ulterioare") + 4
    paragraph.add_run(text[:cut]).bold = True
    paragraph.add_run(text[cut:]).italic = True
    properties = [
        etree.tostring(run._r.rPr, method="c14n", exclusive=True) for run in paragraph.runs
    ]
    monkeypatch.setattr("ema.audit.base_cleanup.heading_spans_document", lambda _: [])
    changes = clean_base(document)
    assert paragraph.text == "Public legislation: " + CORRECT_LAW + " Next sentence."
    assert [
        etree.tostring(run._r.rPr, method="c14n", exclusive=True) for run in paragraph.runs
    ] == properties
    assert changes[0].startswith("G1 fixed text:") and " → " in changes[0]
    assert clean_base(document) == []


def _chapter_prototypes():
    document = Document()
    missing = {"ch4.electricitate_pv", "ch4.echiv_pv", "ch4.specific_pv", "ch4.bilant_real"}
    positions = {}
    for section in CATALOGUE:
        if section.id.startswith("ch4") and section.id not in missing:
            positions[section.id] = len(document.element.body) - 1
            paragraph = document.add_paragraph(section.title)
            style = _node("w:pStyle")
            style.set(qn("w:val"), section.id)
            paragraph._p.get_or_add_pPr().append(style)
    document.add_paragraph("Body prototype")
    document.add_paragraph("Tabelul 4.1")
    table = document.add_table(rows=1, cols=7)
    for cell in table.rows[0].cells:
        cell.text = "Prototype"
    return positions, list(document.element.body)


def test_added_headings_borrow_only_the_named_sibling():
    positions, body = _chapter_prototypes()
    prototypes = _prototypes(positions, body, 0, len(body), (2025,))
    for section, sibling in {
        "ch4.electricitate_pv": "ch4.electricitate",
        "ch4.echiv_pv": "ch4.echiv_electric",
        "ch4.specific_pv": "ch4.specific_electric",
        "ch4.bilant_real": "ch4.mediu",
    }.items():
        assert prototypes.elements["heading:" + section] is body[positions[sibling]]


@pytest.mark.parametrize(
    "missing", ["ch4.echiv_electric", "ch4.electricitate", "ch4.specific_electric", "ch4.mediu"]
)
def test_added_heading_without_its_sibling_fails_loudly(missing):
    positions, body = _chapter_prototypes()
    del positions[missing]
    with pytest.raises(EmaError) as error:
        _prototypes(positions, body, 0, len(body), (2025,))
    assert error.value.code == "heading_prototype_missing"
    assert error.value.detail == missing


@pytest.mark.parametrize("section_id", ["ch2", "ch4"])
def test_empty_chapter_title_does_not_count_as_content(section_id):
    states = [SectionState(section.id, Status.NA) for section in CATALOGUE]
    states = [
        replace(state, status=Status.DONE) if state.section_id == section_id else state
        for state in states
    ]
    issues = empty_chapters(states)
    chapter = next(section for section in CATALOGUE if section.id == section_id)
    assert any(
        issue.code == "chapter_empty"
        and issue.message
        == f"{chapter.title}: capitolul nu are conținut".translate(str.maketrans("șțȘȚ", "şţŞŢ"))
        for issue in issues
    )
    states = [
        replace(state, status=Status.DONE)
        if state.section_id == ("ch2.date_generale" if section_id == "ch2" else "ch4.productie")
        else state
        for state in states
    ]
    assert not any(
        issue.message.startswith(chapter.title + ":") for issue in empty_chapters(states)
    )
    assert not any(issue.message.startswith("BILANȚURILE ENERGETICE:") for issue in issues)


def test_catalogue_subsection_titles_are_sentence_case_and_keep_acronyms():
    for section in CATALOGUE:
        if section.parent and "{" not in section.title:
            assert section.title[0].isupper()
            assert not section.title.isupper()
    assert "SEN" in next(
        section.title for section in CATALOGUE if section.id == "ch4.echiv_electric"
    )


def test_chapter_one_fixed_introduction_counts_as_content():
    states = [
        SectionState(section.id, Status.DONE if section.id == "ch1" else Status.NA)
        for section in CATALOGUE
    ]
    title = next(section.title for section in CATALOGUE if section.id == "ch1")
    assert not any(issue.message.startswith(title + ":") for issue in empty_chapters(states))


def test_imported_measurement_chapter_continues_the_body_counter():
    document = Document()
    numbering = document.part.numbering_part.element
    for num_id, pattern in (("1200", "%1."), ("1201", "%1."), ("1202", "5.%1.")):
        abstract = _node("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), num_id)
        level = _node("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for name, value in (("start", "1"), ("lvlText", pattern)):
            child = _node("w:" + name)
            child.set(qn("w:val"), value)
            level.append(child)
        abstract.append(level)
        numbering.find(qn("w:num")).addprevious(abstract)
        num = _node("w:num")
        num.set(qn("w:numId"), num_id)
        child = _node("w:abstractNumId")
        child.set(qn("w:val"), num_id)
        num.append(child)
        numbering.append(num)
    paragraphs = []
    for num_id in ("1200", "1200", "1200", "1200", "1201", "1200", "1202", "1202", "1200"):
        paragraph = document.add_paragraph("Chapter or section")._p
        paragraph.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = int(num_id)
        paragraphs.append(paragraph)
    insert_chapter_numbering(document, paragraphs[4], paragraphs[5:8])
    numbers = printed_numbers(document)
    assert [numbers[p] for p in paragraphs] == [
        "1.",
        "2.",
        "3.",
        "4.",
        "5.",
        "6.",
        "6.1.",
        "6.2.",
        "7.",
    ]
    assert (
        numbering.xpath('./w:abstractNum[@w:abstractNumId="1202"]/w:lvl/w:lvlText')[0].get(
            qn("w:val")
        )
        == "5.%1."
    )


def test_numbering_restarts_children_per_parent():
    document = Document()
    abstract = _node("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "1200")
    for index, pattern in enumerate(("%1.", "%1.%2.", "%1.%2.%3.")):
        level = _node("w:lvl")
        level.set(qn("w:ilvl"), str(index))
        for name, value in (("start", "1"), ("lvlText", pattern)):
            child = _node("w:" + name)
            child.set(qn("w:val"), value)
            level.append(child)
        abstract.append(level)
    number = _node("w:num")
    number.set(qn("w:numId"), "1200")
    child = _node("w:abstractNumId")
    child.set(qn("w:val"), "1200")
    number.append(child)
    document.part.numbering_part.element.extend([abstract, number])
    paragraphs = []
    for index in (0, 1, 2, 2, 1, 2):
        paragraph = document.add_paragraph("Heading")._p
        numpr = _node("w:numPr")
        for name, value in (("numId", "1200"), ("ilvl", str(index))):
            child = _node("w:" + name)
            child.set(qn("w:val"), value)
            numpr.append(child)
        paragraph.get_or_add_pPr().append(numpr)
        paragraphs.append(paragraph)
    numbers = printed_numbers(document)
    assert [numbers[paragraph] for paragraph in paragraphs] == [
        "1.",
        "1.1.",
        "1.1.1.",
        "1.1.2.",
        "1.2.",
        "1.2.1.",
    ]


def test_heading_alignment_and_indent_are_normalized_once(monkeypatch):
    document = Document()
    paragraphs = [
        document.add_paragraph(text, style="Heading 1")._p
        for text in ("Chapter", "Section", "Detail")
    ]
    numbering = document.part.numbering_part.element
    abstract = _node("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "1200")
    for index, pattern in enumerate(("%1.", "%1.%2.", "%1.%2.%3.")):
        level = _node("w:lvl")
        level.set(qn("w:ilvl"), str(index))
        text = _node("w:lvlText")
        text.set(qn("w:val"), pattern)
        properties = _node("w:pPr")
        indent = _node("w:ind")
        indent.set(qn("w:left"), "720")
        indent.set(qn("w:hanging"), "360")
        properties.append(indent)
        level.extend([text, properties])
        abstract.append(level)
    num = _node("w:num")
    num.set(qn("w:numId"), "1200")
    child = _node("w:abstractNumId")
    child.set(qn("w:val"), "1200")
    num.append(child)
    numbering.extend([abstract, num])
    for index, paragraph in enumerate(paragraphs):
        paragraph.get_or_add_pPr().get_or_add_jc().set(qn("w:val"), "both")
        paragraph.pPr.get_or_add_ind().set(qn("w:left"), "1134")
        numpr = paragraph.pPr.get_or_add_numPr()
        numpr.get_or_add_numId().val = 1200
        numpr.get_or_add_ilvl().val = index
    spans = [
        (MappedHeading(Heading(index, "", (), index), section), index, index + 1)
        for index, section in enumerate(("ch1", "ch1.obiective", "ch1.continut"))
    ]
    monkeypatch.setattr("ema.audit.base_cleanup.heading_spans_document", lambda _: spans)
    changes = clean_base(document)
    assert paragraphs[0].pPr.ind is None
    assert [paragraph.pPr.ind.get(qn("w:left")) for paragraph in paragraphs[1:]] == ["426", "720"]
    assert [
        left - hanging
        for left, hanging in (effective_indent(document, paragraph) for paragraph in paragraphs)
    ] == [360, 66, 360]
    assert [paragraph.pPr.jc.get(qn("w:val")) for paragraph in paragraphs[1:]] == ["left", "left"]
    assert (
        "left=1134, hanging=inherited, jc=both → left=inherited, hanging=inherited, jc=both"
        in changes[0]
    )
    assert "left=1134, hanging=inherited, jc=both → left=720, hanging=360, jc=left" in changes[2]
    assert len(changes) == 3
    assert clean_base(document) == []

    # Word may elide the direct ind when the paragraph style already holds that value.
    paragraphs[1].pPr.remove(paragraphs[1].pPr.ind)
    paragraphs[1].pPr.pStyle.val = "Heading2"
    document.styles["Heading 2"].paragraph_format.left_indent = Twips(426)
    left, hanging = effective_indent(document, paragraphs[1])
    assert left - hanging == 66
