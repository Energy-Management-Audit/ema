"""A synthetic audit render: a small catalogue base, recording chapter writers, a fake Word."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from tests.audit_structure import add_audit_toc, number_audit_headings

from ema.audit import render
from ema.audit.catalogue import CATALOGUE
from ema.audit.render import RenderSummary, start_audit_render
from ema.audit.render_plan import JobUnitPlan
from ema.core.jobs import create_job, status, subscribe
from ema.core.review import decide, mark_absent
from ema.core.review.models import FieldSpec
from ema.core.workspace import Workspace

TITLES = {section.id: section.title for section in CATALOGUE}
MARKER = "[de completat]"


def _base(path: Path) -> None:
    document = Document()
    add_audit_toc(document)
    for chapter, children in (
        ("ch1", ("ch1.obiective",)),
        ("ch2", ("ch2.date_generale",)),
        ("ch3", ("ch3.flux",)),
        ("ch4", ("ch4.concluzii", "ch4.bilant_real")),
        ("ch5", ()),
        ("ch6", ()),
        ("ch7", ()),
    ):
        document.add_paragraph(TITLES[chapter], style="Heading 1")
        document.add_paragraph(f"Text fix {chapter}")
        for child in children:
            document.add_paragraph(TITLES[child], style="Heading 2")
            document.add_paragraph("Text fix" if child == "ch1.obiective" else MARKER)
    document.add_table(rows=1, cols=1)
    number_audit_headings(document)
    document.save(str(path))


def _fill(source: Path, target: Path, *sections: str) -> None:
    document = Document(str(source))
    paragraphs = document.paragraphs
    for section in sections:
        index = next(i for i, p in enumerate(paragraphs) if p.text == TITLES[section])
        paragraphs[index + 1].text = "Scris"
    document.save(str(target))


class FakeWord:
    calls: list[str]

    def __init__(self) -> None:
        self.calls = []

    def update_toc_pages(self, docx: Path) -> None:
        self.calls.append("toc")
        document = Document(str(docx))
        for number, node in enumerate(
            (node for node in document.element.body.iter(qn("w:t")) if node.text == "0"),
            3,
        ):
            node.text = str(number)
        document.save(str(docx))

    def render_pdf(self, docx: Path, pdf: Path) -> None:
        self.calls.append("pdf")
        pdf.write_bytes(b"%PDF-1.7 synthetic")

    def open_check(self, docx: Path) -> None:
        self.calls.append("open")

    def convert_doc(self, doc: Path, out_docx: Path) -> None:
        raise AssertionError("not used")

    def doc_text(self, doc: Path) -> object:
        raise AssertionError("not used")


def synthetic_render(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> tuple[Workspace, str, list[str]]:
    """An audit job whose base and chapter writers are synthetic; Word is absent."""
    ws = Workspace(tmp_path / "ws")
    job = create_job(ws, "audit", "synthetic", 2026)
    for name in ("base.docx", "prototype.docx", "sheet.docx"):
        (tmp_path / name).write_bytes(b"synthetic")
    (tmp_path / "identity.json").write_text('["Forbidden Base SRL"]', encoding="utf-8")
    monkeypatch.setenv("EMA_AUDIT_BASE_DOCUMENT", str(tmp_path / "base.docx"))
    monkeypatch.setenv("EMA_AUDIT_MEASUREMENT_PROTOTYPE", str(tmp_path / "prototype.docx"))
    monkeypatch.setenv("EMA_AUDIT_MEASUREMENT_SHEET_MODEL", str(tmp_path / "sheet.docx"))
    monkeypatch.setenv("EMA_AUDIT_BASE_IDENTITY", str(tmp_path / "identity.json"))
    monkeypatch.setenv("EMA_WORD_PATH", str(tmp_path / "No Word.app"))
    calls: list[str] = []

    def build(plan: object, *, output: Path, **_: object) -> Path:
        _base(output)
        output.with_suffix(".anchors.json").write_text('{"version": 1, "anchors": []}', "utf-8")
        return output

    def writer(name: str, fill: str | None = None):
        def write(*args: object, **_: object) -> None:
            source, target = args[-2], args[-1]
            assert isinstance(source, Path) and isinstance(target, Path)
            calls.append(name)
            shutil.copyfile(source, target)
            if fill:
                _fill(target, target, fill)

        return write

    plan = JobUnitPlan("Client", 1, frozenset({"gas"}), 1, False, 0, 1, "default")
    monkeypatch.setattr(render, "build_base", build)
    monkeypatch.setattr(render, "unit_plan", lambda ws, job: plan)
    monkeypatch.setattr(
        render, "drafted_sections", lambda ws, job: {"ch2.date_generale", "ch3.flux"}
    )
    monkeypatch.setattr(render, "ready_runs", lambda ws, job, stage: [tmp_path])
    monkeypatch.setattr(render, "write_draft", writer("draft"))
    monkeypatch.setattr(render, "write_four", writer("ch4", "ch4.concluzii"))
    monkeypatch.setattr(render, "write_five", writer("ch5"))
    monkeypatch.setattr(render, "write_six", writer("ch6"))
    return ws, job, calls


def run_render(ws: Workspace, job: str, kind: render.Kind = "draft") -> dict[str, object]:
    run = start_audit_render(ws, job, kind)
    for _ in subscribe(ws, job):
        pass
    return next(item for item in status(ws, job).runs if item["id"] == run)


def summary_of(ws: Workspace, job: str, run: str, stage: str = "audit_render") -> RenderSummary:
    with ws.connect() as db:
        folder = ws.artifact_dir(db, job, stage, run)
    return RenderSummary.model_validate_json((folder / "render.json").read_text("utf-8"))


def outputs(ws: Workspace, job: str) -> list[tuple[str, str]]:
    with ws.connect() as db:
        rows = db.execute(
            "SELECT relative_path,kind FROM outputs WHERE job_id=? ORDER BY seq", (job,)
        ).fetchall()
    # Stored names carry an id prefix; the download name is what follows it.
    return [
        (Path(str(row["relative_path"])).name.split("-", 1)[1], str(row["kind"])) for row in rows
    ]


def fill_writer(*sections: str):
    """A chapter writer that writes text in place of each section's marker."""

    def write(*args: object, **_: object) -> None:
        source, target = args[-2], args[-1]
        assert isinstance(source, Path) and isinstance(target, Path)
        _fill(source, target, *sections)

    return write


def narrative_writer(section: str):
    """The ch. 4 writer as the real one behaves for texts: the field's value replaces the marker."""

    def write(source: Path, target: Path, **kwargs: object) -> None:
        job_fields = kwargs["job_fields"]
        assert isinstance(job_fields, list)
        value = next(
            str(field.value)  # type: ignore[attr-defined]
            for field in job_fields
            if field.key == f"narrative.{section}"  # type: ignore[attr-defined]
        )
        document = Document(str(source))
        paragraphs = document.paragraphs
        index = next(i for i, p in enumerate(paragraphs) if p.text == TITLES[section])
        paragraphs[index + 1].text = value
        document.save(str(target))

    return write


def write_narrative(ws: Workspace, job: str, section: str, text: str) -> None:
    """the auditor writes a text: the field exists absent after read, then she corrects it."""
    spec = FieldSpec(
        key=f"narrative.{section}",
        label=TITLES[section],
        value_type="text",
        chapter=section.split(".", maxsplit=1)[0],
    )
    field = mark_absent(ws, job, spec, "not_found")
    decide(ws, job, field.id, "correct", field.revision, "user", value=text)


def write_intros(ws: Workspace, job: str) -> None:
    """The introductions of chapters 3 and 6, which a final needs written."""
    for section in ("ch3", "ch6"):
        write_narrative(ws, job, section, f"Introducere {section}.")
