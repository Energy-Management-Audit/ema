"""Fix round 1: the cover photo slot, a blank address, and the base snapshot of a run."""

from __future__ import annotations

import hashlib
import io
import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from tests.unit.audit.render_seams import TITLES, outputs, run_render, summary_of
from tests.unit.audit.section_marks_seams import by_id, draft, marked_job, session
from tests.unit.audit.test_base_switch import _ready_for_final

from ema.api.job_routes import validate_slot
from ema.audit import render
from ema.audit.base_package import package_issues
from ema.audit.render_bindings import COVER_PHOTO_MISSING, COVER_SLOT
from ema.audit.sections import Status, set_status
from ema.audit.workflow import AuditWorkflow
from ema.core.errors import EmaError
from ema.core.jobs import subscribe
from ema.core.office.anchors import stamp
from ema.core.office.charts import embed_data
from ema.core.office.package import (
    CT_CHART,
    REL_PACKAGE,
    C,
    P,
    R,
    W,
    encoded,
    read_parts,
    target_part,
    write_parts,
    xml,
)
from ema.core.review import fields, propose
from ema.core.workspace import Workspace

BOX = (3_520_440, 2_640_330)
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _png(size: tuple[int, int], color: str = "green") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, "PNG")
    return buffer.getvalue()


@pytest.fixture
def job(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[Workspace, str, list[Path]]:
    """A job whose base cover has a bound address and a bound photo box (as her base has)."""
    ws, job_id = marked_job(tmp_path, monkeypatch)
    built = render.build_base
    sources: list[Path] = []

    def with_cover(plan: object, *, output: Path, base_document: Path, **kwargs: object) -> Path:
        sources.append(base_document)
        path = built(plan, output=output, base_document=base_document, **kwargs)
        document = Document(str(path))
        first = document.paragraphs[0]
        for index in (0, 1):
            paragraph = first.insert_paragraph_before("[de completat]")
            stamp(paragraph._p, f"body_{index}_0", 900 + index)
        document.save(str(path))
        entries = [
            {"slot": "body_0_0", "section": "front", "part": "word/document.xml"},
            {"slot": "body_1_0", "section": "front", "part": "word/document.xml"},
        ]
        entries[0].update(classification="variable", binding="address", box=None)
        entries[1].update(classification="variable", binding="cover_photo", box=list(BOX))
        anchors = {"version": 1, "anchors": entries}
        path.with_suffix(".anchors.json").write_text(json.dumps(anchors), "utf-8")
        return path

    monkeypatch.setattr(render, "build_base", with_cover)
    propose(ws, job_id, "audit.address", "Str. Exemplu nr. 1", [], state="supplied")
    return ws, job_id, sources


def _upload(ws: Workspace, job: str, data: bytes, name: str = "photo.png") -> None:
    path = ws.root.parent / name
    path.write_bytes(data)
    ws.set_slot(job, COVER_SLOT, ws.add_file("synthetic", path))


def _newest(ws: Workspace, job: str, stage: str) -> Path:
    with ws.connect() as db:
        run = db.execute(
            "SELECT id FROM runs WHERE job_id=? AND stage=? ORDER BY started_at DESC LIMIT 1",
            (job, stage),
        ).fetchone()[0]
        return ws.artifact_dir(db, job, stage, str(run))


def test_an_uploaded_photo_fills_her_box_scaled_to_fit(
    job: tuple[Workspace, str, list[Path]],
) -> None:
    ws, job_id, _ = job
    photo = _png((400, 100))
    _upload(ws, job_id, photo)
    record = draft(ws, job_id)
    document = Document(str(_newest(ws, job_id, "audit_render") / "Audit-ciorna.docx"))
    cover = document.paragraphs[1]
    assert "[de completat]" not in cover.text
    extent = next(cover._p.iter(WP + "extent"))
    width, height = int(extent.get("cx")), int(extent.get("cy"))
    assert width == BOX[0] and height == BOX[0] // 4  # the wide photo fits the box's width
    blip = next(cover._p.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"))
    image = document.part.related_parts[blip.get(qn("r:embed"))]
    assert hashlib.sha256(image.blob).hexdigest() == hashlib.sha256(photo).hexdigest()
    markers = summary_of(ws, job_id, str(record["id"])).markers
    assert [item for item in markers if item.section_id == "front"] == []


def test_no_photo_is_a_named_marker_and_refuses_the_final(
    job: tuple[Workspace, str, list[Path]],
) -> None:
    ws, job_id, _ = job
    record = draft(ws, job_id)
    markers = summary_of(ws, job_id, str(record["id"])).markers
    cover = [(item.section_id, item.label) for item in markers if item.section_id == "front"]
    assert cover == [("front", COVER_PHOTO_MISSING)]
    _ready_for_final(ws, job_id)
    final = run_render(ws, job_id, "final")
    assert (final["state"], final["error"]) == ("failed", "front")
    assert [name for name, kind in outputs(ws, job_id) if kind == "final"] == []


def test_a_file_that_is_not_a_photo_is_an_item_failure(
    job: tuple[Workspace, str, list[Path]],
) -> None:
    ws, job_id, _ = job
    _upload(ws, job_id, b"not an image", "photo.txt")
    record = draft(ws, job_id)
    summary = summary_of(ws, job_id, str(record["id"]))
    assert [(item.section_id, item.code) for item in summary.failures] == [
        ("front", "cover_photo_type")
    ]
    cover = [item.label for item in summary.markers if item.section_id == "front"]
    assert cover == [COVER_PHOTO_MISSING]


def test_a_replaced_photo_makes_the_final_stale(job: tuple[Workspace, str, list[Path]]) -> None:
    ws, job_id, _ = job
    _upload(ws, job_id, _png((100, 100)))
    _ready_for_final(ws, job_id)
    assert run_render(ws, job_id, "final")["state"] == "ready"
    assert AuditWorkflow().readiness(ws, job_id).final_ok
    _upload(ws, job_id, _png((100, 100), "blue"), "other.png")
    blocking = AuditWorkflow().readiness(ws, job_id).blocking
    assert "final_stale" in [issue.code for issue in blocking]


def test_a_blank_address_through_the_api_refuses_the_final(
    job: tuple[Workspace, str, list[Path]],
) -> None:
    ws, job_id, _ = job
    _upload(ws, job_id, _png((100, 100)))
    client, headers = session(ws)
    address = next(field for field in fields(ws, job_id) if field.key == "audit.address")
    decided = client.post(
        f"/jobs/{job_id}/fields/{address.id}/decide",
        json={"action": "correct", "on_revision": address.revision, "value": " "},
        headers=headers,
    )
    assert decided.status_code == 200, decided.json()
    _ready_for_final(ws, job_id)
    revision = client.get(f"/jobs/{job_id}").json()["revision"]
    started = client.post(
        f"/jobs/{job_id}/stages/audit_final", json={"on_revision": revision}, headers=headers
    )
    assert started.status_code == 202, started.json()
    list(subscribe(ws, job_id))
    final = client.get(f"/jobs/{job_id}/audit/report").json()["final"]
    assert final["state"] == "failed"
    with ws.connect() as db:
        event = db.execute(
            "SELECT payload FROM job_events WHERE run_id=? AND type='stage_failed'",
            (final["run_id"],),
        ).fetchone()
        error = db.execute("SELECT error FROM runs WHERE id=?", (final["run_id"],)).fetchone()[0]
    assert json.loads(event["payload"])["code"] == "audit_markers"
    assert error == "front"


def test_the_run_builds_from_its_own_snapshot_of_the_base(
    job: tuple[Workspace, str, list[Path]], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    ws, job_id, sources = job
    _upload(ws, job_id, _png((100, 100)))
    configured = tmp_path / "base.docx"
    original = configured.read_bytes()
    built = render.build_base

    def replaced_on_open(plan: object, *, base_document: Path, **kwargs: object) -> Path:
        configured.write_bytes(original + b" replaced after the hash")
        assert base_document != configured and base_document.read_bytes() == original
        return built(plan, base_document=base_document, **kwargs)

    monkeypatch.setattr(render, "build_base", replaced_on_open)
    draft(ws, job_id)
    folder = _newest(ws, job_id, "audit_render")
    made_from = json.loads((folder / render.INPUTS).read_text("utf-8"))
    assert made_from["audit_base_document"] == hashlib.sha256(original).hexdigest()
    assert sources[-1].parent == folder and not sources[-1].exists()  # removed after the run
    readiness = AuditWorkflow().readiness(ws, job_id)
    assert "stale" in [issue.code for issue in readiness.blocking]
    assert by_id(ws, job_id)["ch1"].stale


def test_the_cover_slot_is_one_file_of_audit_jobs_only(
    job: tuple[Workspace, str, list[Path]],
) -> None:
    validate_slot("audit", "cover/photo")
    for job_type, slot in (
        ("audit", "cover/second.jpg"),
        ("audit", "cover"),
        ("piee", "cover/photo"),
        ("invoices", "cover/photo"),
    ):
        with pytest.raises(EmaError) as refused:
            validate_slot(job_type, slot)
        assert refused.value.code == "invalid_slot"
    ws, job_id, _ = job
    client, headers = session(ws)
    path = ws.root.parent / "second.png"
    path.write_bytes(_png((10, 10)))
    sha = ws.add_file("synthetic", path)
    second = client.put(
        f"/jobs/{job_id}/slots/cover/second.png", json={"file_sha": sha}, headers=headers
    )
    assert (second.status_code, second.json()["type"]) == (400, "urn:ema:error:invalid_slot")
    text = ws.root.parent / "photo.txt"
    text.write_bytes(b"not an image")
    refused = client.put(
        f"/jobs/{job_id}/slots/cover/photo",
        json={"file_sha": ws.add_file("synthetic", text)},
        headers=headers,
    )
    assert (refused.status_code, refused.json()["type"]) == (415, "urn:ema:error:cover_photo_type")
    assert ws.list_slots(job_id, "cover") == []  # refused at the boundary: nothing was set
    first = client.put(f"/jobs/{job_id}/slots/cover/photo", json={"file_sha": sha}, headers=headers)
    assert first.status_code == 200, first.json()


def test_a_dropped_section_leaves_no_picture_in_the_package(
    job: tuple[Workspace, str, list[Path]], monkeypatch: pytest.MonkeyPatch
) -> None:
    ws, job_id, _ = job
    _upload(ws, job_id, _png((100, 100)))
    set_status(ws, job_id, "ch4.bilant_real", Status.NA, "user", "nu se aplică")
    built = render.build_base

    def with_picture(plan: object, **kwargs: object) -> Path:
        path = built(plan, **kwargs)
        document = Document(str(path))
        heading = next(p for p in document.paragraphs if p.text == TITLES["ch4.bilant_real"])
        picture = heading.insert_paragraph_before()
        picture.add_run().add_picture(io.BytesIO(_png((50, 50), "red")))
        heading._p.addnext(picture._p)  # the picture belongs to the section
        document.save(str(path))
        return path

    monkeypatch.setattr(render, "build_base", with_picture)
    draft(ws, job_id)
    docx = _newest(ws, job_id, "audit_render") / "Audit-ciorna.docx"
    with zipfile.ZipFile(docx) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert len(media) == 1  # only the cover photo; the dropped section's picture is gone


def _cover_image(docx: Path) -> str:
    document = Document(str(docx))
    blip = next(document.paragraphs[1]._p.iter(A + "blip"))
    return hashlib.sha256(document.part.related_parts[blip.get(qn("r:embed"))].blob).hexdigest()


def test_a_photo_replaced_after_confirmation_blocks_the_final_until_redrafted(
    job: tuple[Workspace, str, list[Path]],
) -> None:
    """Round 2: a final never prints a cover photo that no draft showed."""
    ws, job_id, _ = job
    _upload(ws, job_id, _png((100, 100)))
    _ready_for_final(ws, job_id)  # drafted with the first photo, then confirmed
    replaced = _png((100, 100), "blue")
    _upload(ws, job_id, replaced, "other.png")
    blocking = AuditWorkflow().readiness(ws, job_id).blocking
    assert [(issue.code, issue.message) for issue in blocking] == [
        ("cover_photo_changed", "Fotografia sediului s-a schimbat după ciornă. Refaceţi ciorna.")
    ]
    with pytest.raises(EmaError) as refused:
        AuditWorkflow().start_final(ws, job_id)
    assert refused.value.code == "not_ready"
    draft(ws, job_id)
    assert AuditWorkflow().readiness(ws, job_id).final_ok
    AuditWorkflow().start_final(ws, job_id)
    list(subscribe(ws, job_id))
    final = _newest(ws, job_id, "audit_final") / "Audit-final.docx"
    assert _cover_image(final) == hashlib.sha256(replaced).hexdigest()


def _chart(path: Path, number: int, after: str | None) -> None:
    """A bar chart with its own embedded workbook, after the paragraph reading ``after`` (or
    first in the body), as ch. 4 writes them."""
    parts = read_parts(path)
    part = f"word/charts/chart{number}.xml"
    parts[part] = (
        f'<c:chartSpace xmlns:c="{C}"><c:chart><c:plotArea><c:barChart><c:barDir val="col"/>'
        '<c:ser><c:idx val="0"/><c:order val="0"/><c:tx><c:strRef><c:f>Sheet1!$B$1</c:f>'
        '<c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>Consum</c:v></c:pt></c:strCache>'
        "</c:strRef></c:tx><c:cat><c:strRef><c:f>Sheet1!$A$2</c:f><c:strCache>"
        '<c:ptCount val="1"/><c:pt idx="0"><c:v>Ian</c:v></c:pt></c:strCache></c:strRef></c:cat>'
        '<c:val><c:numRef><c:f>Sheet1!$B$2</c:f><c:numCache><c:ptCount val="1"/><c:pt idx="0">'
        "<c:v>1</c:v></c:pt></c:numCache></c:numRef></c:val></c:ser></c:barChart></c:plotArea>"
        "</c:chart></c:chartSpace>"
    ).encode()
    rels = xml(parts, "word/_rels/document.xml.rels")
    etree.SubElement(
        rels,
        f"{{{P}}}Relationship",
        Id=f"rIdChart{number}",
        Type=f"{R}/chart",
        Target=f"charts/chart{number}.xml",
    )
    parts["word/_rels/document.xml.rels"] = encoded(rels)
    types = xml(parts, "[Content_Types].xml")
    etree.SubElement(
        types, f"{{{types.nsmap[None]}}}Override", PartName=f"/{part}", ContentType=CT_CHART
    )
    parts["[Content_Types].xml"] = encoded(types)
    document = xml(parts, "word/document.xml")
    drawing = etree.fromstring(
        f'<w:p xmlns:w="{W}" xmlns:r="{R}" xmlns:c="{C}" xmlns:a="{A[1:-1]}" '
        f'xmlns:wp="{WP[1:-1]}"><w:r><w:drawing><wp:inline><wp:extent cx="1" cy="1"/>'
        f'<wp:docPr id="{900 + number}" name="Chart {number}"/><a:graphic><a:graphicData '
        f'uri="{C}"><c:chart r:id="rIdChart{number}"/></a:graphicData></a:graphic></wp:inline>'
        "</w:drawing></w:r></w:p>"
    )
    paragraphs = list(document.iter(f"{{{W}}}p"))
    if after is None:
        paragraphs[0].addprevious(drawing)
    else:
        text = f"{{{W}}}t"
        heading = next(
            p for p in paragraphs if "".join(t.text or "" for t in p.iter(text)) == after
        )
        heading.addnext(drawing)
    parts["word/document.xml"] = encoded(document)
    write_parts(parts, path)
    embed_data(path, part, path)


def test_a_dropped_chart_goes_with_its_workbook_and_a_kept_one_keeps_its_own(
    job: tuple[Workspace, str, list[Path]], monkeypatch: pytest.MonkeyPatch
) -> None:
    """Round 2: the scrub after the n/a drop removes the dropped section's chart and workbook,
    and keeps a chart that stays together with its workbook relationship."""
    ws, job_id, _ = job
    _upload(ws, job_id, _png((100, 100)))
    set_status(ws, job_id, "ch4.bilant_real", Status.NA, "user", "nu se aplică")
    built = render.build_base

    def with_charts(plan: object, **kwargs: object) -> Path:
        path = built(plan, **kwargs)
        _chart(path, 1, None)
        _chart(path, 2, TITLES["ch4.bilant_real"])
        return path

    monkeypatch.setattr(render, "build_base", with_charts)
    draft(ws, job_id)
    docx = _newest(ws, job_id, "audit_render") / "Audit-ciorna.docx"
    parts = read_parts(docx)
    charts = sorted(name for name in parts if name.startswith("word/charts/"))
    assert charts == ["word/charts/_rels/chart1.xml.rels", "word/charts/chart1.xml"]
    workbooks = [name for name in parts if name.startswith("word/embeddings/")]
    assert len(workbooks) == 1
    relation = next(iter(xml(parts, "word/charts/_rels/chart1.xml.rels")))
    assert relation.get("Type") == REL_PACKAGE
    assert target_part("word/charts/chart1.xml", relation.get("Target", "")) == workbooks[0]
    assert package_issues(docx, ("Client Secret",)) == []
