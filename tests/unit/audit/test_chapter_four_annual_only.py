"""Annual-only carriers keep their yearly evidence without empty monthly output."""

import re
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from ema.audit.chapter_four_annual import annual_carrier_table
from ema.audit.chapter_four_blocks import chapter_four_blocks
from ema.audit.chapter_four_chart_placement import place_chart_groups
from ema.audit.chapter_four_charts import chapter_chart_groups, chart_blocks
from ema.consumption_analysis.analysis import Metric, value
from ema.core.office.blocks import (
    Caption,
    ElementLocator,
    Missing,
    NativeChart,
    Num,
    Paragraph,
    Prototypes,
    Ref,
    Table,
    render,
)
from ema.core.office.missing_text import TABLE_MISSING_TEXT
from ema.core.office.package import read_parts, write_parts
from ema.energy_data.carriers import CARRIER_NAMES_RO, Carrier
from ema.energy_data.factors import FACTORS_2026
from ema.energy_data.model import CarrierSeries, EnergyDataset, Reading


def _dataset(*, with_fuel_months: bool = False) -> EnergyDataset:
    years = (2023, 2024, 2025)
    carriers = {
        Carrier.petrol: {
            year: CarrierSeries({1: Reading(2, "t")} if with_fuel_months else {}, Reading(2, "t"))
            for year in years
        },
        Carrier.diesel: {
            year: CarrierSeries({1: Reading(3, "t")} if with_fuel_months else {}, Reading(3, "t"))
            for year in years
        },
        Carrier.natural_gas: {
            year: CarrierSeries({1: Reading(100, "MWh")}, Reading(100, "MWh")) for year in years
        },
    }
    return EnergyDataset(
        years,
        carriers,
        {"product": {year: CarrierSeries(annual=Reading(1000, "t")) for year in years}},
        {"product": "t"},
    )


def _section(blocks: list[object], name: str) -> list[object]:
    start = next(
        i
        for i, block in enumerate(blocks)
        if isinstance(block, Paragraph) and block.proto == f"heading:{name}"
    )
    end = next(
        (
            i
            for i in range(start + 1, len(blocks))
            if isinstance(blocks[i], Paragraph) and blocks[i].proto.startswith("heading:")
        ),
        len(blocks),
    )
    return blocks[start + 1 : end]


def test_annual_only_fuels_have_annual_tables_charts_and_contiguous_captions() -> None:
    dataset = _dataset()
    groups, skipped = chapter_chart_groups(dataset, FACTORS_2026, "Client")
    blocks = place_chart_groups(chapter_four_blocks(dataset, FACTORS_2026), groups)
    assert not any(key.startswith("ch4.carburant:") for key in skipped)
    for section, unit in (("ch4.carburant", "t"), ("ch4.echiv_carburant", "tep")):
        body = _section(blocks, section)
        tables = [block for block in body if isinstance(block, Table)]
        captions = [block for block in body if isinstance(block, Caption)]
        charts = [block for block in body if isinstance(block, NativeChart)]
        assert len(tables) == len(captions) == 2
        assert len(charts) == 1
        assert all(table.header == [["Anul", "2023", "2024", "2025"]] for table in tables)
        assert all(
            len(table.rows) == 1 and table.rows[0][0] == [f"Valoare ({unit})"] for table in tables
        )
        assert all(
            isinstance(cell[0], Num) and cell[0].value is not None
            for table in tables
            for cell in table.rows[0][1:]
        )
        for carrier in (Carrier.petrol, Carrier.diesel):
            sentence = (
                f"Pentru {CARRIER_NAMES_RO[carrier]} au fost transmise numai consumurile anuale."
            )
            assert Paragraph("body", [sentence]) in body
            caption = next(c for c in captions if carrier.value in c.id)
            assert caption.segments[:2] == ["Tabelul ", Ref("tab", caption.id)]
            assert body.index(Paragraph("body", [sentence])) < body.index(caption)
        assert {series.name for series in charts[0].series} == {
            "Consumul de benzină",
            "Consumul de motorină",
        }
        assert not any(isinstance(block, Table) and block.proto != "emissions" for block in body)
    specific = _section(blocks, "ch4.specific_carburant")
    specific_tables = [block for block in specific if isinstance(block, Table)]
    specific_captions = [block for block in specific if isinstance(block, Caption)]
    assert len(specific_tables) == len(specific_captions) == 2
    assert all(table.header == [["Anul", "2023", "2024", "2025"]] for table in specific_tables)
    assert all(
        table.rows == [[["Valoare (tep/mii tone)"], *table.rows[0][1:]]]
        for table in specific_tables
    )
    assert len([block for block in specific if isinstance(block, NativeChart)]) == 1
    for carrier in (Carrier.petrol, Carrier.diesel):
        sentence = Paragraph(
            "body",
            [f"Pentru {CARRIER_NAMES_RO[carrier]} au fost transmise numai consumurile anuale."],
        )
        caption = next(c for c in specific_captions if carrier.value in c.id)
        assert specific.index(sentence) < specific.index(caption)
        assert caption.segments[:2] == ["Tabelul ", Ref("tab", caption.id)]
    table_ids = [block.id for block in blocks if isinstance(block, Caption)]
    assert len(table_ids) == len(set(table_ids))
    figure_numbers = [
        int(match.group(1))
        for block in blocks
        if (isinstance(block, Paragraph) and block.proto == "chart_caption")
        or isinstance(block, Missing)
        for match in [
            re.match(
                r"Fig\. nr\. 4\.(\d+)",
                block.segments[0] if isinstance(block, Paragraph) else block.text,
            )
        ]
        if match
    ]
    assert sorted(set(figure_numbers)) == list(range(1, max(figure_numbers) + 1))


def test_mixed_fuel_chart_keeps_monthly_carrier_and_all_annual_carriers() -> None:
    dataset = _dataset()
    dataset.carriers[Carrier.lpg] = {
        year: CarrierSeries({1: Reading(1, "t")}, Reading(1, "t")) for year in dataset.years
    }
    blocks, skipped = chart_blocks("ch4.carburant", dataset, FACTORS_2026, "Client")
    charts = [block for block in blocks if isinstance(block, NativeChart)]
    section = _section(chapter_four_blocks(dataset, FACTORS_2026), "ch4.carburant")
    assert len([block for block in section if isinstance(block, Table)]) == 8
    assert not skipped
    assert len(charts) == 4
    assert all(
        [series.name for series in chart.series] == ["Consumul de GPL"] for chart in charts[:-1]
    )
    assert {series.name for series in charts[-1].series} == {
        "Consumul de benzină",
        "Consumul de motorină",
        "Consumul de GPL",
    }


def test_annual_only_fuels_leave_annual_tep_total_unchanged() -> None:
    annual_dataset = _dataset()
    monthly_dataset = _dataset(with_fuel_months=True)
    for year in annual_dataset.years:
        expected = value(monthly_dataset, FACTORS_2026, Metric("tep_total"), year, filed=False)[0]
        actual = value(annual_dataset, FACTORS_2026, Metric("tep_total"), year, filed=False)[0]
        assert actual == pytest.approx(expected)
        blocks = _section(chapter_four_blocks(annual_dataset, FACTORS_2026), "ch4.echiv_total")
        printed = [
            segment
            for block in blocks
            if isinstance(block, Paragraph)
            for segment in block.segments
            if isinstance(segment, Num) and segment.unit == "tep"
        ]
        assert [number.value for number in printed] == pytest.approx(
            [
                value(annual_dataset, FACTORS_2026, Metric("tep_total"), y, filed=False)[0]
                for y in annual_dataset.years
            ]
        )


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _emissions_base(path: Path, columns: int) -> tuple[Prototypes, ElementLocator]:
    """A base whose "emissions" prototype is horizontal: Sursa plus one column per year."""
    doc = Document()
    doc.add_heading("4.3 Test", level=2)
    caption = doc.add_paragraph("Tabelul 4.1. Exemplu")
    table = doc.add_table(rows=2, cols=columns)
    for row in table.rows:
        for cell in row.cells:
            cell.text = "x"
    doc.save(path)
    prototypes = Prototypes({"caption": caption._p, "emissions": table._tbl}, chapter=4)
    parts = read_parts(path)
    root = etree.fromstring(parts["word/document.xml"])
    body = root.find(W + "body")
    assert body is not None
    for node in list(body)[1:-1]:
        body.remove(node)
    parts["word/document.xml"] = etree.tostring(root, encoding="UTF-8")
    write_parts(parts, path)
    return prototypes, ElementLocator(1)


@pytest.mark.parametrize("years", [(2023, 2024, 2025), (2024, 2025)])
def test_the_annual_table_renders_on_the_horizontal_emissions_prototype(
    tmp_path: Path, years: tuple[int, ...]
) -> None:
    dataset = EnergyDataset(
        years,
        {
            Carrier.diesel: {
                year: CarrierSeries(annual=None if year == 2024 else Reading(3, "t"))
                for year in years
            }
        },
        {},
        {},
    )
    blocks = annual_carrier_table(
        dataset,
        FACTORS_2026,
        Metric("carrier", (Carrier.diesel,)),
        years,
        "t",
        label="Motorină",
        section="ch4.carburant",
    )
    source, out = tmp_path / "base.docx", tmp_path / "out.docx"
    prototypes, locator = _emissions_base(source, len(years) + 1)
    render(source, out, locator, blocks, prototypes)
    body = etree.fromstring(read_parts(out)["word/document.xml"]).find(W + "body")
    assert body is not None
    table = next(node for node in body if node.tag == W + "tbl")
    rows = [
        ["".join(text.text or "" for text in cell.iter(W + "t")) for cell in row.findall(W + "tc")]
        for row in table.findall(W + "tr")
    ]
    printed = ["3,00" if year != 2024 else TABLE_MISSING_TEXT for year in years]
    assert rows == [["Anul", *map(str, years)], ["Valoare (t)", *printed]]
