"""TOC field boundaries, printed list labels, and readable entries."""

import re

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ema.audit.base_numbering import printed_numbers
from ema.audit.base_toc import DEFAULT_TOC, refresh_toc
from ema.audit.headings import Heading, MappedHeading


def _node(tag: str, **attributes: str):
    node = OxmlElement(tag)
    for key, value in attributes.items():
        node.set(qn("w:" + key), value)
    return node


def _field(paragraph, kind: str, instruction: str = "") -> None:
    run = _node("w:r")
    run.append(_node("w:fldChar", fldCharType=kind))
    paragraph._p.append(run)
    if instruction:
        run = _node("w:r")
        code = _node("w:instrText")
        code.text = instruction
        run.append(code)
        paragraph._p.append(run)


def _toc(document, level: int, text: str):
    paragraph = document.add_paragraph(text)
    paragraph._p.get_or_add_pPr().append(_node("w:pStyle", val=f"TOC{level}"))
    return paragraph


def _fixture():
    document = Document()
    document.add_paragraph("Preface")
    page_break = _toc(document, 1, "")
    page_break.add_run().add_break(WD_BREAK.PAGE)
    heading = document.add_paragraph("CUPRINS")
    self_entry = _toc(document, 1, "Cuprins 99")
    _field(self_entry, "begin", DEFAULT_TOC)
    _field(self_entry, "separate")
    first = _toc(document, 1, "stale chapter 99")
    first._p.get_or_add_pPr().append(_node("w:spacing", line="360"))
    first._p.pPr.append(_node("w:tabs"))
    first._p.pPr.tabs.append(_node("w:tab", val="right", leader="dot", pos="8500"))
    second = _toc(document, 2, "stale section 99")
    second._p.get_or_add_pPr().append(_node("w:ind", left="300"))
    closing = document.add_paragraph("After the TOC")
    _field(closing, "end")
    outside = _toc(document, 2, "Outside field — preserve")
    return document, page_break, heading, outside


def _headings(document, monkeypatch, specifications):
    headings = []
    for section, level, title in specifications:
        paragraph = document.add_paragraph(title)
        headings.append((section, level, paragraph._p))

    def spans(doc):
        body = list(doc.element.body)
        return [
            (
                MappedHeading(
                    Heading(level, "", (), body.index(paragraph)),
                    section,
                    template="DESCRIEREA SECȚIEI {process}" if section == "ch3.process" else None,
                ),
                body.index(paragraph),
                body.index(paragraph) + 1,
            )
            for section, level, paragraph in headings
        ]

    monkeypatch.setattr("ema.audit.base_toc.heading_spans_document", spans)
    return [paragraph for _, _, paragraph in headings]


def _entries(document):
    return [
        paragraph._p for paragraph in document.paragraphs if paragraph._p.xpath(".//w:hyperlink")
    ]


def _text(element):
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def _add_list(document, paragraphs, pattern="4.3.%1.", start="1"):
    abstract = _node("w:abstractNum", abstractNumId="1200")
    level = _node("w:lvl", ilvl="0")
    level.extend(
        [
            _node("w:start", val=start),
            _node("w:numFmt", val="decimal"),
            _node("w:lvlText", val=pattern),
        ]
    )
    abstract.append(level)
    number = _node("w:num", numId="1200")
    number.append(_node("w:abstractNumId", val="1200"))
    document.part.numbering_part.element.extend([abstract, number])
    parent = document.styles.add_style("Numbered parent", WD_STYLE_TYPE.PARAGRAPH)
    parent.element.get_or_add_pPr().append(_node("w:numPr"))
    parent.element.pPr.numPr.append(_node("w:numId", val="1200"))
    child = document.styles.add_style("Numbered child", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    for paragraph in paragraphs:
        paragraph.get_or_add_pPr().append(_node("w:pStyle", val=child.style_id))


def test_refresh_preserves_page_break_heading_self_entry_and_outside_paragraph(monkeypatch):
    document, page_break, heading, outside = _fixture()
    paragraphs = _headings(
        document, monkeypatch, [("ch6", 0, "6. Indicatori"), ("ch6.indicatori", 1, "6.1. Consum")]
    )
    bookmark = _node("w:bookmarkStart", id="12", name="_ema_approved")
    paragraphs[0].append(bookmark)
    paragraphs[0].append(_node("w:bookmarkEnd", id="12"))
    for _ in range(2):
        refresh_toc(document)
        body = list(document.element.body)
        entries = _entries(document)
        assert body.index(page_break._p) < body.index(heading._p) < body.index(entries[0])
        assert body.index(entries[0]) == body.index(heading._p) + 1
        assert _text(entries[0]) == "Cuprins0"
        assert _text(entries[1]) == "5.INDICATORI0"
        assert outside._p in body
        assert paragraphs[0].xpath('./w:bookmarkStart[@w:name="_ema_approved"]')
        assert not any("stale" in _text(entry) for entry in entries)
        assert len(document.element.xpath('.//w:instrText[starts-with(., " TOC ")]')) == 1
        refs = [
            node.text
            for node in document.element.iter(qn("w:instrText"))
            if "PAGEREF" in (node.text or "")
        ]
        names = {node.get(qn("w:name")) for node in document.element.iter(qn("w:bookmarkStart"))}
        assert all(re.search(r"PAGEREF\s+(\S+)", ref)[1] in names for ref in refs)
        depth = 0
        for char in document.element.iter(qn("w:fldChar")):
            kind = char.get(qn("w:fldCharType"))
            depth += 1 if kind == "begin" else -1 if kind == "end" else 0
            assert depth >= 0
        assert depth == 0


def test_entries_use_list_counters_catalogue_case_and_explicit_formatting(monkeypatch):
    document, *_ = _fixture()
    paragraphs = _headings(
        document,
        monkeypatch,
        [
            ("ch4", 0, "4. Analiza energiei"),
            ("ch4.echiv_electric", 1, "HEADING IN CAPITALS"),
            ("ch4.echiv_pv", 1, "heading in lower case"),
        ],
    )
    _add_list(document, paragraphs[1:], pattern="4.%1.", start="3")
    for _ in range(2):
        refresh_toc(document)
        entries = _entries(document)
        assert [_text(entry) for entry in entries] == [
            "Cuprins0",
            "4.ANALIZA ENERGIEI0",
            "4.3.Analiza consumului total echivalent de energie electrică din SEN0",
            "4.4.Analiza consumului echivalent de energie electrică fotovoltaică0",
        ]
        for entry, level in zip(entries, [1, 1, 2, 2], strict=True):
            for run in entry.iter(qn("w:r")):
                properties = run.find(qn("w:rPr"))
                assert properties is not None
                assert properties.find(qn("w:rFonts")).get(qn("w:ascii")) == "Times New Roman"
                assert properties.find(qn("w:sz")).get(qn("w:val")) == "24"
                assert properties.find(qn("w:smallCaps")).get(qn("w:val")) == "0"
                assert properties.find(qn("w:caps")).get(qn("w:val")) == "0"
                assert properties.find(qn("w:b")).get(qn("w:val")) == ("1" if level == 1 else "0")
                assert properties.find(qn("w:i")).get(qn("w:val")) == "0"
            offset = 600 if level == 1 else 1000
            left = 0 if level == 1 else 300
            assert entry.pPr.ind.get(qn("w:hanging")) == str(offset)
            assert entry.pPr.ind.get(qn("w:left")) == str(left + offset)
            assert entry.pPr.tabs[0].get(qn("w:pos")) == str(left + offset)
        assert entries[1].pPr.spacing.get(qn("w:line")) == "360"
        assert entries[1].xpath("./w:hyperlink/w:r/w:tab")
        assert entries[1].pPr.tabs[-1].get(qn("w:leader")) == "dot"


def test_refresh_clears_italic_toc3_style_and_generated_runs(monkeypatch):
    document, *_ = _fixture()
    toc3 = document.styles.add_style("TOC3", WD_STYLE_TYPE.PARAGRAPH)
    toc3.font.italic = True
    _headings(document, monkeypatch, [("ch4.echiv_electric", 1, "4.3. Energia")])

    refresh_toc(document)

    style = next(
        style
        for style in document.styles.element.findall(qn("w:style"))
        if style.get(qn("w:styleId")) == "TOC3"
    )
    assert style.find(f"{qn('w:rPr')}/{qn('w:i')}").get(qn("w:val")) == "0"
    entries = _entries(document)
    assert any(entry.pPr.pStyle.get(qn("w:val")) == "TOC2" for entry in entries)
    for entry in entries:
        for run in entry.iter(qn("w:r")):
            properties = run.find(qn("w:rPr"))
            assert properties is not None
            italic = properties.find(qn("w:i"))
            assert italic is None or italic.get(qn("w:val")) != "1"


def test_numbering_start_override_and_toc_exclusion():
    document = Document()
    first = document.add_paragraph("First")._p
    ignored = _toc(document, 2, "Ignored")._p
    second = document.add_paragraph("Second")._p
    _add_list(document, [first, second], start="3")
    second.get_or_add_pPr().append(_node("w:numPr"))
    second.pPr.numPr.append(_node("w:ilvl", val="0"))
    ignored.pPr.append(_node("w:numPr"))
    ignored.pPr.numPr.append(_node("w:numId", val="1200"))
    number = document.part.numbering_part.element.xpath('./w:num[@w:numId="1200"]')[0]
    override = _node("w:lvlOverride", ilvl="0")
    override.append(_node("w:startOverride", val="7"))
    number.append(override)
    assert printed_numbers(document) == {first: "4.3.7.", second: "4.3.8."}


def test_refresh_preserves_page_break_inside_the_toc_field(monkeypatch):
    document, *_ = _fixture()
    retained = _toc(document, 1, "")
    retained.add_run().add_break(WD_BREAK.PAGE)
    self_entry = next(p for p in document.paragraphs if p.text == "Cuprins 99")
    self_entry._p.addnext(retained._p)
    _headings(document, monkeypatch, [("ch4", 0, "4. Analiza energiei")])
    refresh_toc(document)
    assert retained._p in list(document.element.body)
    assert retained._p.xpath('.//w:br[@w:type="page"]')


def test_refresh_requires_an_actual_toc_field(monkeypatch):
    document = Document()
    _toc(document, 1, "Outside")
    _toc(document, 2, "Outside")
    monkeypatch.setattr("ema.audit.base_toc.heading_spans_document", lambda doc: [])
    with pytest.raises(ValueError, match="no TOC prototypes"):
        refresh_toc(document)


def test_repeated_process_toc_entries_keep_each_process_name(monkeypatch):
    document, *_ = _fixture()
    _headings(
        document,
        monkeypatch,
        [
            ("ch3", 0, "3. Descrierea situației existente"),
            ("ch3.process", 1, "3.2. DESCRIEREA SECȚIEI Atelier Alpha"),
            ("ch3.process", 1, "3.3. DESCRIEREA SECȚIEI Atelier Beta"),
            ("ch6.measure", 1, "6.1. Lighting retrofit"),
        ],
    )
    for _ in range(2):
        refresh_toc(document)
        assert [_text(entry) for entry in _entries(document)][2:] == [
            "3.2.Descrierea secției ATELIER ALPHA0",
            "3.3.Descrierea secției ATELIER BETA0",
            "5.1.Lighting retrofit0",
        ]


def test_body_levels_one_and_two_are_uppercase_and_level_three_keeps_base_case(monkeypatch):
    document, *_ = _fixture()
    paragraphs = _headings(
        document,
        monkeypatch,
        [
            ("ch4", 0, "4. Analiza energiei"),
            ("ch4.bilant_real", 1, "4.7. bilanțul energetic real"),
            ("ch4.echiv_electric", 2, "4.3.1. ANALIZA ENERGIEI DIN SEN"),
            ("ch4.echiv_pv", 2, "4.3.2. Analiza energiei fotovoltaice"),
        ],
    )
    refresh_toc(document)
    assert [_text(paragraph) for paragraph in paragraphs] == [
        "4. ANALIZA ENERGIEI",
        "4.7. BILANȚUL ENERGETIC REAL",
        "4.3.1. ANALIZA ENERGIEI DIN SEN",
        "4.3.2. Analiza energiei fotovoltaice",
    ]
    assert "Analiza bilanțului energetic real" in _text(_entries(document)[2])


def test_toc_lists_levels_one_and_two_only_but_still_normalises_level_three(monkeypatch):
    document, *_ = _fixture()
    paragraphs = _headings(
        document,
        monkeypatch,
        [
            ("ch4", 0, "4. Analiza energiei"),
            ("ch4.bilant_real", 1, "4.7. bilanțul energetic real"),
            ("ch4.echiv_electric", 2, "4.3.1. ANALIZA ENERGIEI DIN SEN"),
        ],
    )
    refresh_toc(document)
    entries = _entries(document)
    assert [entry.pPr.pStyle.get(qn("w:val")) for entry in entries] == ["TOC1", "TOC1", "TOC2"]
    assert not any("SEN" in _text(entry) for entry in entries)
    assert not paragraphs[2].xpath("./w:bookmarkStart")
    assert _text(paragraphs[2]) == "4.3.1. ANALIZA ENERGIEI DIN SEN"
    instructions = [
        node.text
        for node in document.element.iter(qn("w:instrText"))
        if (node.text or "").strip().startswith("TOC")
    ]
    assert len(instructions) == 1
    assert '\\o "1-2"' in instructions[0]
    assert '\\o "1-3"' not in instructions[0]
