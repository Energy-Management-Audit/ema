"""The unit plan of an audit job: every count from its own source."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from ema.audit.render_plan import JobUnitPlan, process_count, unit_plan
from ema.core.errors import EmaError
from ema.core.jobs import create_job
from ema.core.review import decide, mark_absent, propose
from ema.core.review.models import FieldSpec
from ema.core.workspace import Workspace


def _fisa(path: Path, flows: int) -> Path:
    document = Document()
    document.add_paragraph("Fişa de date")
    for index in range(flows):
        document.add_paragraph(f"  Flux {index + 1}: prelucrare")
    document.add_paragraph("Descriere: flux tehnologic")
    document.save(str(path))
    return path


def _necesar(path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Cons energetice"
    sheet.cell(5, 1, "Consum energie electrica din SEN")
    sheet.cell(6, 1, 2025)
    for name in ("Echipamente 1", "Echipamente 2", "Autovehicule"):
        table = book.create_sheet(name)
        table.cell(1, 1, "Denumire")
        table.cell(2, 1, "Pompa")
    book.save(path)
    return path


def _add(ws: Workspace, job: str, slot: str, source: Path) -> None:
    ws.set_slot(job, slot, ws.add_file("synthetic", source))


def test_process_count_prefers_schemes_then_fisa_then_one(tmp_path: Path) -> None:
    schemes = ["dossier/5.1. Flux A.pdf", "dossier/5.1. Flux A v2.pdf", "dossier/5.3. Flux C.pdf"]
    assert process_count(schemes, None) == (2, "schemes")
    fisa = _fisa(tmp_path / "Fisa.docx", 2)
    assert process_count(["dossier/0. Necesar.xlsx"], fisa) == (2, "fisa")
    assert process_count(schemes, fisa) == (2, "schemes")
    assert process_count([], _fisa(tmp_path / "None.docx", 0)) == (1, "default")
    assert process_count([], None) == (1, "default")


def test_unit_plan_counts_every_unit_from_the_job(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = create_job(ws, "audit", "synthetic", 2026)
    image = tmp_path / "photo.jpg"
    image.write_bytes(b"\xff\xd8\xff synthetic")
    thermal = tmp_path / "thermal.jpg"
    thermal.write_bytes(b"\xff\xd8\xff thermal")
    _add(ws, job, "dossier/0. Necesar info.xlsx", _necesar(tmp_path / "necesar.xlsx"))
    _add(ws, job, "dossier/Fisa de date.docx", _fisa(tmp_path / "fisa.docx", 2))
    _add(ws, job, "visit/meter/TG 1/a.jpg", image)
    _add(ws, job, "visit/meter/TG 2/a.jpg", image)
    _add(ws, job, "visit/thermal/t.jpg", thermal)
    spec = FieldSpec(key="audit.company_name", label="Denumire", value_type="text")
    mark_absent(ws, job, spec, "not_found")
    for key, value in (
        ("carrier.natural_gas.2025", Decimal(5)),
        ("carrier.diesel.2025.01", Decimal(1)),
        ("carrier.water_potable.2025", Decimal(2)),
        ("audit_measure.count", Decimal(3)),
    ):
        propose(ws, job, key, value, [], state="extracted")
    mark_absent(ws, job, "carrier.electricity_grid.2025", "not_found")
    with ws.connect() as db:
        db.execute("UPDATE clients SET name='Client Sintetic SRL' WHERE id='synthetic'")

    plan = unit_plan(ws, job)

    assert plan == JobUnitPlan(
        client_name="Client Sintetic SRL",
        processes=2,
        carriers=frozenset({"gas", "fuel", "water"}),
        measured_panels=2,
        thermal_measurements=True,
        equipment_tables=2,
        measures=3,
        processes_source="fisa",
        client_revision=("synthetic", 1),
    )
    named = propose(ws, job, spec, "Nume din anexă SA", [], state="extracted")
    assert unit_plan(ws, job).client_name == "Nume din anexă SA"
    # A name she rejected is not the client's: the client record names it, as without one (D2).
    decide(ws, job, named.id, "reject", named.revision, "user")
    assert unit_plan(ws, job).client_name == "Client Sintetic SRL"


def test_unit_plan_without_photos_form_or_name(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = create_job(ws, "audit", "synthetic", 2026)
    with pytest.raises(EmaError) as caught:
        unit_plan(ws, job)
    assert caught.value.code == "audit_client_name"
    with ws.connect() as db:
        db.execute("UPDATE clients SET name='Client' WHERE id='synthetic'")
    plan = unit_plan(ws, job)
    assert (plan.processes, plan.processes_source) == (1, "default")
    assert (plan.measured_panels, plan.thermal_measurements, plan.measures) == (0, False, 0)
    assert (plan.equipment_tables, plan.carriers) == (0, frozenset())
