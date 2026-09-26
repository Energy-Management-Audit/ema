"""The blank form and label-based measure reader."""

from pathlib import Path

from openpyxl import load_workbook

from ema.audit.measures_form import HEADERS, read_measures_form, write_measures_form
from ema.energy_data.carriers import Carrier


def test_form_round_trip_and_moved_header(tmp_path: Path) -> None:
    path = write_measures_form(tmp_path / "form.xlsx")
    book = load_workbook(path)
    assert book.sheetnames == ["Măsuri propuse", "Liste"]
    assert tuple(cell.value for cell in book["Măsuri propuse"][1]) == HEADERS
    assert {item.formula1 for item in book["Măsuri propuse"].data_validations.dataValidation} == {
        "=EmaCarriers",
        "=EmaUnits",
    }
    assert book.defined_names["EmaCarriers"].attr_text == "'Liste'!$A$1:$A$16"
    sheet = book["Măsuri propuse"]
    sheet.insert_rows(1, 3)
    sheet.append([1, "LED replacement", "Lower use", "Energie electrică", 100, "MWh", 40, 10, None])
    book.create_sheet("Other", 0)
    book.save(path)
    result = read_measures_form(path)
    assert result.header is not None and result.header.ref.row == 4
    assert len(result.rows) == 1
    assert result.rows[0].title.ref.row == 5
    assert result.rows[0].carrier_value == Carrier.electricity_grid
    assert result.rows[0].saving_amount.ref.a1 == "Măsuri propuse!E5"
    assert not result.issues


def test_populated_nameless_row_is_issue_and_blank_row_stops(tmp_path: Path) -> None:
    path = write_measures_form(tmp_path / "form.xlsx")
    book = load_workbook(path)
    sheet = book["Măsuri propuse"]
    sheet.append([1, "Lighting", "Lower use", "Energie electrică", 1, "MWh", 2, 1])
    sheet.append([2, None, "Lower use", "Energie electrică", 1, "MWh", 2, 1])
    sheet.append([3, "Heating", "Lower use", "Gaze naturale", 2, "MWh", 3, 1])
    sheet.append([4, "Total", None, None, None, None, None, None])
    sheet.append([None] * 9)
    sheet.append([5, "After blank", "Lower use", "Gaze naturale", 2, "MWh", 3, 1])
    book.save(path)
    result = read_measures_form(path)
    assert [row.title.value for row in result.rows] == ["Lighting", "Heating"]
    assert [(item.code, item.detail) for item in result.issues] == [
        ("measure_row_invalid", "row 3")
    ]


def test_unknown_carrier_and_unit_are_item_issues(tmp_path: Path) -> None:
    path = write_measures_form(tmp_path / "form.xlsx")
    book = load_workbook(path)
    sheet = book["Măsuri propuse"]
    sheet.append([1, "Unknown carrier", None, "Solar", 1, "MWh", None, None])
    sheet.append([2, "Unknown unit", None, "Motorină", 1, "litre", None, None])
    sheet.append([3, "Valid", None, "Motorină", 1, "t", None, None])
    book.save(path)
    result = read_measures_form(path)
    assert [item.code for item in result.issues] == ["measure_row_invalid"] * 2
    assert [row.title.value for row in result.rows] == ["Valid"]
