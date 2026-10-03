"""The audit draft stage: replay only, persisted as a run, refused before any write otherwise."""

import json
import sys
from pathlib import Path

import pytest
from docx import Document
from tests.audit_replay import (
    CH2_DRAFT,
    audit_job_with_facts,
    draft_recording,
    render_draft_section,
    support_recording,
)
from tests.workspace_jobs import create_job
from typer.testing import CliRunner

from ema.audit.catalogue import CATALOGUE
from ema.audit.draft_agent import recorded_facts
from ema.audit.draft_render import review_payload
from ema.audit.draft_stage import draft_section
from ema.audit.sections import get_status
from ema.cli import _app, app
from ema.core.errors import EmaError
from ema.core.jobs import get_job
from ema.core.review import log
from ema.core.workspace import Workspace

SECTION = "ch2.date_generale"
TITLES = {section.id: section.title for section in CATALOGUE}


def _recordings(ws: Workspace, job: str, folder: Path) -> tuple[Path, Path]:
    folder.mkdir(parents=True, exist_ok=True)
    return (
        draft_recording(ws, job, CH2_DRAFT, folder / "draft.json"),
        support_recording(ws, job, CH2_DRAFT, folder / "support.json", None),
    )


def _counts(ws: Workspace, job: str) -> tuple[int, int, int]:
    with ws.connect() as db:
        runs = db.execute("SELECT COUNT(*) FROM runs WHERE job_id=?", (job,)).fetchone()[0]
        sessions = db.execute("SELECT COUNT(*) FROM agent_sessions").fetchone()[0]
    return runs, sessions, int(str(get_job(ws, job)["revision"]))


def test_replay_draft_is_a_ready_run_with_artifacts_and_reads(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    drafts, support = _recordings(ws, job, tmp_path / "rec")

    result = draft_section(ws, job, SECTION, draft_recording=drafts, support_recording=support)
    assert any(item.actor == "agent" and item.field_id == SECTION for item in log(ws, job))

    assert (result.draft_status, result.section_status, result.coverage) == (
        "drafted",
        "drafted",
        1.0,
    )
    with ws.connect() as db:
        run = db.execute("SELECT state FROM runs WHERE id=?", (result.run,)).fetchone()
        reads = db.execute(
            "SELECT row_id,revision FROM run_reads WHERE run_id=? AND table_name='fields'",
            (result.run,),
        ).fetchall()
        facts = db.execute("SELECT id,revision FROM fields WHERE job_id=?", (job,)).fetchall()
        folder = ws.job_path(db, job) / "work" / "draft" / result.run / "sections"
    assert run["state"] == "ready"
    assert {(row[0], row[1]) for row in reads} == {(row[0], row[1]) for row in facts}
    assert result.draft_path == folder / f"{SECTION}.json"
    assert result.review_path == folder / f"{SECTION}.draft-review.json"
    assert json.loads(result.draft_path.read_text(encoding="utf-8"))["section"] == SECTION
    assert list(json.loads(result.review_path.read_text(encoding="utf-8"))) == [
        "section",
        "coverage",
        "cited_sentences",
        "total_sentences",
        "review",
    ]


def test_live_switch_off_or_half_a_recording_refuses_before_any_write(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    constructed: list[str] = []
    monkeypatch.setenv("EMA_LLM_LIVE", "1")
    monkeypatch.setenv("EMA_GEMINI_API_KEY", "synthetic-key")
    monkeypatch.setenv("EMA_OPENAI_API_KEY", "synthetic-key")
    monkeypatch.setattr(
        "ema.core.llm.providers.genai.Client", lambda **_: constructed.append("gemini")
    )
    monkeypatch.setattr("ema.core.llm.providers.OpenAI", lambda **_: constructed.append("openai"))
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    drafts, _ = _recordings(ws, job, tmp_path / "rec")
    before = _counts(ws, job)

    for recordings, code in (
        ((None, None), "ai_client_disabled"),
        ((drafts, None), "replay_invalid"),
        ((None, drafts), "replay_invalid"),
    ):
        with pytest.raises(EmaError) as refused:
            draft_section(
                ws, job, SECTION, draft_recording=recordings[0], support_recording=recordings[1]
            )
        assert refused.value.code == code

    assert _counts(ws, job) == before
    assert constructed == []


def test_wrong_job_unknown_section_and_invalid_recording_start_no_run(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    piee = create_job(ws, "piee", "synthetic", 2026)
    invalid = tmp_path / "invalid.json"
    invalid.write_text(
        json.dumps({"source": "live", "format": "openai-chat-completions"}), encoding="utf-8"
    )
    not_json = tmp_path / "not-json.json"
    not_json.write_text("{ not json", encoding="utf-8")
    no_responses = tmp_path / "no-responses.json"
    no_responses.write_text(
        json.dumps({"source": "recorded", "format": "openai-chat-completions"}), encoding="utf-8"
    )
    listed = tmp_path / "listed.json"
    listed.write_text(json.dumps([{"source": "recorded"}]), encoding="utf-8")
    scalar_responses = tmp_path / "scalar-responses.json"
    scalar_responses.write_text(
        json.dumps({"source": "recorded", "format": "openai-chat-completions", "responses": 3}),
        encoding="utf-8",
    )
    cases = (
        (piee, SECTION, invalid, "wrong_job_type"),
        (job, "ch4.unknown", invalid, "section_missing"),
        (job, SECTION, invalid, "replay_invalid"),
        (job, SECTION, not_json, "replay_invalid"),
        (job, SECTION, no_responses, "replay_invalid"),
        (job, SECTION, listed, "replay_invalid"),
        (job, SECTION, scalar_responses, "replay_invalid"),
        (job, SECTION, tmp_path / "absent.json", "replay_invalid"),
        (job, SECTION, tmp_path, "replay_invalid"),
    )
    for target, section, recording, code in cases:
        with pytest.raises(EmaError) as refused:
            draft_section(
                ws, target, section, draft_recording=recording, support_recording=recording
            )
        assert refused.value.code == code
        assert _counts(ws, target)[0] == 0


def test_mismatched_recording_fails_the_run_and_keeps_the_section(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    drafts, support = _recordings(ws, job, tmp_path / "rec")
    recording = json.loads(drafts.read_text(encoding="utf-8"))
    recording["responses"][0]["request_hashes"] = {"messages": "0" * 64}
    drafts.write_text(json.dumps(recording), encoding="utf-8")
    before = get_status(ws, job, SECTION)

    with pytest.raises(EmaError) as failed:
        draft_section(ws, job, SECTION, draft_recording=drafts, support_recording=support)

    assert failed.value.code == "draft_failed"
    assert get_status(ws, job, SECTION) == before
    with ws.connect() as db:
        assert db.execute("SELECT state FROM runs WHERE job_id=?", (job,)).fetchone()[0] == "failed"


def test_render_writes_the_review_payload(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    base, output = tmp_path / "base.docx", tmp_path / "o.docx"
    document = Document()
    document.add_paragraph(TITLES["ch2"], style="Heading 1")
    document.add_paragraph(TITLES[SECTION], style="Heading 2")
    document.add_paragraph("[de completat]")
    document.add_paragraph(TITLES["ch2.istorie"], style="Heading 2")
    document.save(str(base))
    facts = recorded_facts(ws, job, SECTION)
    check = render_draft_section(ws, job, base, output, draft=CH2_DRAFT, facts=facts, flags=())

    written = json.loads(output.with_suffix(".draft-review.json").read_text(encoding="utf-8"))
    assert written == json.loads(json.dumps(review_payload(CH2_DRAFT, check, ())))
    assert get_status(ws, job, SECTION).status.value == "drafted"


def test_cli_prints_the_draft_in_contract_order(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    ws = Workspace(tmp_path / "ws")
    monkeypatch.setenv("EMA_WORKSPACE", str(ws.root))
    job = audit_job_with_facts(ws)
    drafts, support = _recordings(ws, job, tmp_path / "rec")
    args = ["audit", "draft", job, SECTION, "--draft-recording", str(drafts)]

    printed = CliRunner().invoke(_app, [*args, "--support-recording", str(support)])

    assert printed.exit_code == 0, printed.output
    assert list(json.loads(printed.output)) == [
        "job",
        "run",
        "section",
        "draft_status",
        "section_status",
        "coverage",
        "cited_sentences",
        "total_sentences",
        "review",
        "draft_path",
        "review_path",
    ]


def test_cli_without_recordings_and_live_switch_off_exits_with_the_disabled_message(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    ws = Workspace(tmp_path / "ws")
    monkeypatch.setenv("EMA_WORKSPACE", str(ws.root))
    job = audit_job_with_facts(ws)
    monkeypatch.setattr(sys, "argv", ["ema", "audit", "draft", job, "ch3.flux"])

    with pytest.raises(SystemExit) as exited:
        app()

    assert exited.value.code == 1
    assert "Documentele clientului nu pot fi trimise la AI." in capsys.readouterr().err


def test_cli_with_a_missing_recording_exits_with_the_invalid_message(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    ws = Workspace(tmp_path / "ws")
    monkeypatch.setenv("EMA_WORKSPACE", str(ws.root))
    job = audit_job_with_facts(ws)
    absent = str(tmp_path / "absent.json")
    arguments = ["--draft-recording", absent, "--support-recording", absent]
    monkeypatch.setattr(sys, "argv", ["ema", "audit", "draft", job, SECTION, *arguments])

    with pytest.raises(SystemExit) as exited:
        app()

    err = capsys.readouterr().err
    assert exited.value.code == 1
    assert "Înregistrarea AI este invalidă." in err and "Traceback" not in err
    assert _counts(ws, job)[0] == 0
