"""Draft trust boundary: references, missing facts, names and render isolation."""

import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from ema.audit.draft_checks import DraftReview, check_draft
from ema.audit.draft_render import render_section, unrendered_items
from ema.audit.draft_schema import DraftFigure, DraftTable, DraftText, SectionDraft
from ema.core.errors import EmaError
from ema.core.office.anchors import stamp
from ema.core.review.models import Field


def _fact(key: str, value: str | int, kind: str = "text") -> Field:
    return Field(
        id=key,
        job_id="synthetic",
        key=key,
        label=key,
        value_type=kind,
        value=value,
        state="supplied",
        presence="found",
        evidence=["synthetic-evidence"],
    )


def _draft(text: str, ids: list[str]) -> SectionDraft:
    return SectionDraft(
        section="ch2.date_generale",
        status="drafted",
        paragraphs=[DraftText(text=text, fact_ids=ids)],
    )


def test_rejects_literal_number_name_ai_and_unknown_fact() -> None:
    facts = {"audit.company_name": _fact("audit.company_name", "Atelier Exemplu")}
    for text, ids, code in (
        (
            "Societatea {{f:audit.company_name}} are 500 angajați.",
            ["audit.company_name"],
            "literal_number",
        ),
        ("Societatea Inventata are sediul aici.", [], "literal_name"),
        ("Societatea {{f:audit.company_name}} folosește AI.", ["audit.company_name"], "ai_mention"),
        ("Societatea {{f:audit.cui}} este aici.", ["audit.cui"], "fact_missing"),
    ):
        assert code in {
            issue.code for issue in check_draft(_draft(text, ids), facts, "synthetic").fatal
        }


def test_uncited_and_nonrenderable_items_are_reviewed() -> None:
    facts = {"audit.company_name": _fact("audit.company_name", "Atelier Exemplu")}
    draft = _draft(
        "Societatea {{f:audit.company_name}} produce bunuri. Este modernă.", ["audit.company_name"]
    )
    checked = check_draft(draft, facts, "synthetic")
    assert checked.coverage == 0.5
    assert [issue.code for issue in checked.review] == ["uncited_sentence"]
    rich = SectionDraft(
        section="ch2.date_generale",
        status="drafted",
        tables=[
            DraftTable(
                caption=DraftText(
                    text="Date despre {{f:audit.company_name}}.", fact_ids=["audit.company_name"]
                ),
                rows=[
                    [DraftText(text="{{f:audit.company_name}}", fact_ids=["audit.company_name"])]
                ],
            )
        ],
        figures=[
            DraftFigure(
                fact_id="audit.company_name",
                caption=DraftText(text="{{f:audit.company_name}}", fact_ids=["audit.company_name"]),
            )
        ],
    )
    assert not check_draft(rich, facts, "synthetic").fatal
    assert {item.code for item in check_draft(rich, facts, "synthetic").review} == {
        "unrendered_table",
        "unrendered_figure",
    }
    assert len(unrendered_items(rich)) == 2


def test_render_changes_only_selected_paragraph_and_blocks_unsupported(tmp_path: Path) -> None:
    base, output, anchors = (
        tmp_path / name for name in ("base.docx", "output.docx", "base.anchors.json")
    )
    document = Document()
    first = document.add_paragraph("[de completat]")
    second = document.add_paragraph("[de completat]")
    stamp(first._p, "one", 1)
    stamp(second._p, "two", 2)
    document.save(str(base))
    anchors.write_text(
        json.dumps(
            {
                "version": 1,
                "anchors": [
                    {
                        "slot": "one",
                        "section": "ch2.date_generale",
                        "classification": "variable",
                        "part": "word/document.xml",
                    },
                    {
                        "slot": "two",
                        "section": "ch2.istorie",
                        "classification": "variable",
                        "part": "word/document.xml",
                    },
                ],
            }
        )
    )
    facts = {"audit.company_name": _fact("audit.company_name", "Atelier Exemplu")}
    draft = _draft("Societatea {{f:audit.company_name}} produce bunuri.", ["audit.company_name"])
    render_section(base, anchors, output, draft, facts, (), job="synthetic")
    rendered = Document(str(output))
    assert [p.text for p in rendered.paragraphs] == [
        "Societatea Atelier Exemplu produce bunuri.",
        "[de completat]",
    ]
    assert next(rendered.paragraphs[0]._p.iter(qn("w:bookmarkStart")), None) is not None
    rerendered = tmp_path / "rerendered.docx"
    render_section(output, anchors, rerendered, draft, facts, (), job="synthetic")
    assert Document(str(rerendered)).paragraphs[1].text == "[de completat]"
    flagged = (DraftReview("unsupported", "paragraph:0", "unsupported claim"),)
    render_section(base, anchors, output, draft, facts, flagged, job="synthetic")
    assert Document(str(output)).paragraphs[0].text == "[de completat]"
    with pytest.raises(EmaError) as error:
        render_section(
            base,
            anchors,
            output,
            SectionDraft(
                section="ch2.date_generale", status="drafted", paragraphs=draft.paragraphs * 2
            ),
            facts,
            (),
            job="synthetic",
        )
    assert error.value.code == "draft_slots"
