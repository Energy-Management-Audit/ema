"""Synthetic Draft v4 single-section path: the chapter request, checker retry, fail-closed
support and runtime style redaction."""

import json
import re
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from tests.audit_replay import CH2_DRAFT, audit_job_with_facts

from ema.audit.draft_agent import draft_section_run
from ema.audit.draft_prompt import rule_text
from ema.audit.draft_schema import DraftText, SectionDraft
from ema.audit.draft_style import style_example
from ema.core.errors import EmaError
from ema.core.llm.models import default_model
from ema.core.llm.types import Exchange
from ema.core.workspace import Workspace

SECTION = "ch2.date_generale"
Answer = SectionDraft | list[SectionDraft] | str


class DraftProvider:
    """Answers each chapter call in turn: a draft, the drafts of a call, or a raw text."""

    name = "openai"

    def __init__(self, drafts: list[Answer]) -> None:
        self.drafts = drafts
        self.requests: list[dict[str, Any]] = []
        self.limits: list[int] = []

    def respond(self, *args: Any, **kwargs: Any) -> Exchange:
        self.requests.append(json.loads(args[1][1]["content"]))
        self.limits.append(args[4])
        answer = self.drafts.pop(0)
        if isinstance(answer, str):
            return Exchange(answer, (), 1, 1)
        drafts = answer if isinstance(answer, list) else [answer]
        return Exchange(json.dumps({"sections": [item.model_dump() for item in drafts]}), (), 1, 1)


class SupportProvider:
    """Supports every sentence it is asked about, except the refused (location, index) pairs;
    or answers a fixed text."""

    name = "openai"

    def __init__(
        self, response: str | None = None, refused: frozenset[tuple[str, int]] = frozenset()
    ) -> None:
        self.response = response
        self.refused = refused
        self.calls = 0
        self.requests: list[list[dict[str, Any]]] = []
        self.limits: list[int] = []

    def respond(self, *args: Any, **kwargs: Any) -> Exchange:
        self.calls += 1
        request = json.loads(args[1][1]["content"])
        self.requests.append(request)
        self.limits.append(args[4])
        if self.response is not None:
            return Exchange(self.response, (), 1, 1)
        verdicts = [
            {
                "location": item["location"],
                "sentence_index": item["sentence_index"],
                "supported": (item["location"], item["sentence_index"]) not in self.refused,
                "reason": "claim",
            }
            for item in request
        ]
        return Exchange(json.dumps({"verdicts": verdicts}), (), 1, 1)


def _run(
    tmp_path: Path, drafts: list[Answer], support: SupportProvider | None = None
) -> tuple[DraftProvider, SupportProvider, SectionDraft, tuple[Any, ...]]:
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    provider = DraftProvider(drafts)
    support = support or SupportProvider()
    _, accepted, _, flags = draft_section_run(
        ws,
        job,
        SECTION,
        provider,
        support,
        model_id=default_model("openai").id,
        synthetic=True,
    )
    return provider, support, accepted, flags


def test_accepted_draft_uses_one_structured_call(tmp_path: Path) -> None:
    provider, support, accepted, flags = _run(tmp_path, [CH2_DRAFT])
    assert accepted == CH2_DRAFT and flags == ()
    assert len(provider.requests) == support.calls == 1
    assert [item["section"] for item in provider.requests[0]["sections"]] == [SECTION]


def test_literal_name_retry_carries_rule(tmp_path: Path) -> None:
    invalid = SectionDraft(
        section=SECTION,
        status="drafted",
        paragraphs=[DraftText(text="Atelier Exemplu are activitate.")],
    )
    provider, support, accepted, _ = _run(tmp_path, [invalid, CH2_DRAFT])
    assert accepted == CH2_DRAFT
    assert len(provider.requests) == 2
    assert support.calls == 1
    rules = {(error["rule"], error["rule_text"]) for error in provider.requests[1]["errors"]}
    assert rules == {
        ("literal_name", rule_text("literal_name")),
        ("uncited_sentence", rule_text("uncited_sentence")),
    }
    assert rule_text("literal_name").startswith("Orice nume, număr")
    assert [item["section"] for item in provider.requests[1]["request"]["sections"]] == [SECTION]


def test_second_fatal_draft_stops_after_two_calls(tmp_path: Path) -> None:
    invalid = SectionDraft(
        section=SECTION, status="drafted", paragraphs=[DraftText(text="Atelier Exemplu.")]
    )
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    provider, support = DraftProvider([invalid, invalid]), SupportProvider()
    with pytest.raises(EmaError) as error:
        draft_section_run(
            ws, job, SECTION, provider, support, model_id=default_model("openai").id, synthetic=True
        )
    assert error.value.code == "draft_incomplete"
    assert len(provider.requests) == 2
    assert support.calls == 0


def test_a_sentence_without_a_verdict_is_unsupported(tmp_path: Path) -> None:
    sentence = "Societatea {{f:audit.company_name}} are {{f:audit.employees}} angajați."
    _, _, accepted, flags = _run(tmp_path, [CH2_DRAFT], SupportProvider('{"verdicts": []}'))
    assert accepted == CH2_DRAFT
    assert [(flag.code, flag.location, flag.detail, flag.sentence) for flag in flags] == [
        ("unsupported", "paragraph:0", "no verdict", sentence)
    ]


def test_a_refused_sentence_is_flagged_alone(tmp_path: Path) -> None:
    draft = SectionDraft(
        section=SECTION,
        status="drafted",
        paragraphs=[
            DraftText(
                text=(
                    "Societatea {{f:audit.company_name}} există. "
                    "Are {{f:audit.employees}} angajați."
                ),
                fact_ids=["audit.company_name", "audit.employees"],
            )
        ],
    )
    support = SupportProvider(refused=frozenset({(f"{SECTION}:paragraph:0", 1)}))
    _, _, _, flags = _run(tmp_path, [draft], support)
    assert [(flag.code, flag.sentence) for flag in flags] == [
        ("unsupported", "Are {{f:audit.employees}} angajați.")
    ]
    assert [item["sentence_index"] for item in support.requests[0]] == [0, 1]
    assert support.requests[0][1]["facts"] == {"audit.employees": "85"}


def test_unavailable_support_marks_only_sentences_that_cite(tmp_path: Path) -> None:
    cited = "Societatea are personal propriu {{c:audit.employees}}."
    mixed = "Firma {{f:audit.company_name}} are personal calificat {{c:audit.employees}}."
    valued = "Societatea {{f:audit.company_name}} are {{f:audit.employees}} angajați."
    draft = SectionDraft(
        section=SECTION,
        status="drafted",
        paragraphs=[
            DraftText(text=cited, fact_ids=["audit.employees"]),
            DraftText(text=mixed, fact_ids=["audit.company_name", "audit.employees"]),
            DraftText(text=valued, fact_ids=["audit.company_name", "audit.employees"]),
        ],
    )
    _, _, accepted, flags = _run(tmp_path, [draft], SupportProvider("not json"))
    assert accepted == draft
    assert [(flag.code, flag.location, flag.detail) for flag in flags] == [
        ("unsupported", "paragraph:0", "support_unavailable"),
        ("unsupported", "paragraph:1", "support_unavailable"),
        ("support_unavailable", "section", "ai_schema"),
    ]


def test_support_output_is_bounded_per_sentence(tmp_path: Path) -> None:
    _, support, _, _ = _run(tmp_path, [CH2_DRAFT])
    assert support.limits == [16_000 + 40]


def test_an_answer_off_the_schema_takes_the_one_retry(tmp_path: Path) -> None:
    provider, support, accepted, _ = _run(tmp_path, ['{"sections": [', CH2_DRAFT])
    assert accepted == CH2_DRAFT
    assert [error["rule"] for error in provider.requests[1]["errors"]] == ["omitted"]
    assert (len(provider.requests), support.calls) == (2, 1)
    ws = Workspace(tmp_path / "again")
    job = audit_job_with_facts(ws)
    provider, support = DraftProvider(["not json", "not json"]), SupportProvider()
    with pytest.raises(EmaError) as error:
        draft_section_run(
            ws, job, SECTION, provider, support, model_id=default_model("openai").id, synthetic=True
        )
    assert error.value.code == "ai_schema"
    assert (len(provider.requests), support.calls) == (2, 0)


def test_style_example_masks_identity_and_every_number(tmp_path: Path) -> None:
    base = tmp_path / "synthetic.docx"
    doc = Document()
    doc.add_paragraph("DESCRIEREA ȘI ISTORICUL SOCIETĂȚII", style="Heading 1")
    doc.add_paragraph("Date generale", style="Heading 2")
    doc.add_paragraph("Atelier Exemplu are 1.234 angajați în 2026.")
    doc.add_paragraph("Istoria companiei", style="Heading 2")
    doc.save(base)
    example = style_example(base, ("Atelier Exemplu",), SECTION)
    assert "Atelier Exemplu" not in example
    assert not re.search(r"\d", example)
    assert example.count("{{…}}") == 3


def test_style_example_masks_unlisted_names_variants_and_acronyms(tmp_path: Path) -> None:
    base = tmp_path / "synthetic.docx"
    doc = Document()
    doc.add_paragraph("DESCRIEREA ȘI ISTORICUL SOCIETĂȚII", style="Heading 1")
    doc.add_paragraph("Date generale", style="Heading 2")
    doc.add_paragraph(
        "Clientul Ştefan\nIonescu lucrează cu Ana Popescu în Brașov la ACME. "
        "În 2026-02-03 avea 12,5%."
    )
    doc.add_paragraph("Istoria companiei", style="Heading 2")
    doc.save(base)
    example = style_example(base, ("Ștefan Ionescu",), SECTION)
    for leaked in ("Ştefan", "Ionescu", "Ana", "Popescu", "Brașov", "ACME", "2026", "12,5"):
        assert leaked not in example
    assert example.count("{{…}}") >= 7


def test_provider_receives_only_redacted_base_example(tmp_path: Path, monkeypatch: Any) -> None:
    base = tmp_path / "synthetic.docx"
    doc = Document()
    doc.add_paragraph("DESCRIEREA ȘI ISTORICUL SOCIETĂȚII", style="Heading 1")
    doc.add_paragraph("Date generale", style="Heading 2")
    doc.add_paragraph("Atelier Exemplu avea 12 angajați în 2025.")
    doc.add_paragraph("Istoria companiei", style="Heading 2")
    doc.save(base)
    identity = tmp_path / "identity.json"
    identity.write_text(json.dumps(["Atelier Exemplu"]), encoding="utf-8")
    monkeypatch.setenv("EMA_AUDIT_BASE_DOCUMENT", str(base))
    monkeypatch.setenv("EMA_AUDIT_BASE_IDENTITY", str(identity))
    provider, _, _, _ = _run(tmp_path, [CH2_DRAFT])
    example = provider.requests[0]["sections"][0]["style_example"]
    assert example
    assert "Atelier Exemplu" not in example
    assert not re.search(r"\d", example)


def test_preflight_refuses_before_draft_call(tmp_path: Path, monkeypatch: Any) -> None:
    monkeypatch.setenv("EMA_AI_JOB_BUDGET_USD", "0.000001")
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    provider = DraftProvider([CH2_DRAFT])
    with pytest.raises(EmaError) as error:
        draft_section_run(
            ws,
            job,
            SECTION,
            provider,
            SupportProvider(),
            model_id=default_model("openai").id,
            synthetic=True,
        )
    assert error.value.code == "ai_budget"
    assert provider.requests == []
