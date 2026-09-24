"""Synthetic regression cases for the S10b audit base review."""

from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement

from ema.audit.base_anchor import _classification
from ema.audit.base_package import _numbering_issues, package_issues


def test_client_number_in_fixed_chapter_is_variable() -> None:
    document = Document()
    paragraph = document.add_paragraph("Annual production was 12345 MWh.")
    assert _classification("ch1", paragraph._p, ("ACME",), heading=False) == "variable"


def test_numeric_leftover_is_reported_after_package_save(tmp_path) -> None:
    path = tmp_path / "base.docx"
    document = Document()
    document.add_paragraph("Annual production was 12345 MWh.")
    document.save(path)
    issues = package_issues(
        path, ("ACME",), numeric_leftovers=("Annual production was 12345 MWh.",)
    )
    assert any("numeric base paragraph remains" in issue for issue in issues)


def test_numbering_definitions_follow_required_child_order() -> None:
    root = OxmlElement("w:numbering")
    for tag in ("w:numPicBullet", "w:abstractNum", "w:num", "w:numIdMacAtCleanup"):
        root.append(OxmlElement(tag))
    assert not _numbering_issues(root)
    root.insert(0, OxmlElement("w:abstractNum"))
    assert _numbering_issues(root) == ["invalid numbering element order"]
