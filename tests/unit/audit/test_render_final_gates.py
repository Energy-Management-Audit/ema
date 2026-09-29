"""A final never carries AI wording: the rendered document's whole text is read before it ships."""

from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path

import pytest
from docx import Document
from lxml import etree
from PIL import Image
from tests.unit.audit.render_seams import (
    TITLES,
    FakeWord,
    fill_writer,
    narrative_writer,
    outputs,
    run_render,
    summary_of,
    synthetic_render,
    write_intros,
    write_narrative,
)

from ema.audit import base_numeric, render
from ema.audit.render_steps import ai_wording_hits
from ema.audit.sections import Status, set_status
from ema.audit.workflow import AuditWorkflow
from ema.core.office.missing_text import MISSING_TEXT, TABLE_MISSING_NOTE, TABLE_MISSING_TEXT
from ema.core.office.package import P, R, W, encoded, read_parts, write_parts, xml
from ema.core.review.models import Field
from ema.core.workspace import Workspace


def test_every_text_part_is_read(tmp_path: Path) -> None:
    docx = tmp_path / "gate.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Redactat cu GPT-5"
    document.add_paragraph(TITLES["ch1"], style="Heading 1")
    document.add_paragraph("Angajaţi ai unor persoane juridice.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Tabel generat de IA"
    document.save(str(docx))
    assert ai_wording_hits(docx, []) == ["ch1", "header1"]
    clean = tmp_path / "clean.docx"
    Document().save(str(clean))
    assert ai_wording_hits(clean, []) == []


@pytest.fixture
def final(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[Workspace, str]:
    ws, job, _ = synthetic_render(tmp_path, monkeypatch)
    write_intros(ws, job)
    set_status(ws, job, "ch4.bilant_real", Status.NA, "user")
    monkeypatch.setattr(render, "write_draft", fill_writer("ch2.date_generale", "ch3.flux"))
    monkeypatch.setattr(render, "write_four", narrative_writer("ch4.concluzii"))
    monkeypatch.setattr(render, "word_available", lambda settings: True)
    monkeypatch.setattr(render, "word_automation", lambda settings: FakeWord())
    return ws, job


def _failure(ws: Workspace, job: str) -> dict[str, object]:
    with ws.connect() as db:
        row = db.execute(
            "SELECT payload FROM job_events WHERE job_id=? AND type='stage_failed' "
            "ORDER BY seq DESC LIMIT 1",
            (job,),
        ).fetchone()
    return json.loads(row["payload"])


def test_ai_wording_in_a_ch4_text_refuses_the_final_naming_the_field(
    final: tuple[Workspace, str],
) -> None:
    ws, job = final
    write_narrative(ws, job, "ch4.concluzii", "Acest raport a fost generat de AI")
    readiness = AuditWorkflow().readiness(ws, job)
    assert [issue.message for issue in readiness.blocking if issue.code == "ai_wording"] == [
        f"Textul menţionează AI („AI”): {TITLES['ch4.concluzii']}"
    ]
    record = run_render(ws, job, "final")
    assert record["state"] == "failed"
    assert _failure(ws, job) == {
        "code": "audit_ai_wording",
        "message": "Raportul conţine formulări despre AI.",
    }
    assert str(record["error"]).split("; ") == ["narrative.ch4.concluzii", "ch4.concluzii"]
    assert outputs(ws, job) == []


def test_clean_text_passes(final: tuple[Workspace, str]) -> None:
    ws, job = final
    write_narrative(ws, job, "ch4.concluzii", "Consumul a scăzut după modernizare.")
    record = run_render(ws, job, "final")
    assert record["state"] == "ready", record["error"]
    assert [name for name, _ in outputs(ws, job)] == ["Audit-final.pdf", "Audit-final.docx"]


@pytest.mark.parametrize("kind", ["draft", "final"])
@pytest.mark.parametrize("missing_text", [MISSING_TEXT, TABLE_MISSING_TEXT])
def test_sourced_missing_cells_keep_dev_counts_and_final_gate_outcome(
    final: tuple[Workspace, str], monkeypatch: pytest.MonkeyPatch, kind: str, missing_text: str
) -> None:
    ws, job = final
    write_narrative(ws, job, "ch4.concluzii", "Consumul a scăzut după modernizare.")
    original = render.write_four
    word = FakeWord()
    monkeypatch.setattr(render, "word_automation", lambda settings: word)

    def with_missing_cells(source: Path, target: Path, **kwargs: object) -> None:
        original(source, target, **kwargs)
        document = Document(target)
        heading = next(p for p in document.paragraphs if p.text == TITLES["ch4.concluzii"])
        table = document.add_table(rows=1, cols=3)
        for cell in table.rows[0].cells[:2]:
            cell.text = missing_text
        table.rows[0].cells[2].text = "Available"
        note = document.add_paragraph(
            TABLE_MISSING_NOTE if missing_text == TABLE_MISSING_TEXT else MISSING_TEXT
        )
        heading._p.addnext(table._tbl)
        table._tbl.addnext(note._p)
        document.save(target)

    monkeypatch.setattr(render, "write_four", with_missing_cells)
    record = run_render(ws, job, kind)
    assert record["state"] == "ready", record["error"]
    stage = "audit_render" if kind == "draft" else "audit_final"
    assert summary_of(ws, job, str(record["id"]), stage).markers == []
    assert word.calls == ["toc", "pdf", "open"]
    name = "ciorna" if kind == "draft" else "final"
    assert [output for output, _ in outputs(ws, job)] == [f"Audit-{name}.pdf", f"Audit-{name}.docx"]


def test_a_marker_in_a_header_refuses_the_final(
    final: tuple[Workspace, str], monkeypatch: pytest.MonkeyPatch
) -> None:
    ws, job = final
    write_narrative(ws, job, "ch4.concluzii", "Consumul a scăzut după modernizare.")
    built = render.build_base

    def with_header(plan: object, *, output: Path, **kwargs: object) -> Path:
        path = built(plan, output=output, **kwargs)
        document = Document(str(path))
        document.sections[0].header.paragraphs[0].text = "Audit energetic [de completat]"
        document.save(str(path))
        return path

    monkeypatch.setattr(render, "build_base", with_header)
    record = run_render(ws, job, "final")
    assert record["state"] == "failed"
    assert _failure(ws, job)["code"] == "audit_markers"
    assert record["error"] == "front"
    assert outputs(ws, job) == []


def _png(color: str) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (4, 4), color).save(buffer, "PNG")
    return buffer.getvalue()


def _picture_bullet(path: Path, picture: bytes) -> None:
    """Her numbering's picture bullet, as word/numbering.xml relates it to word/media."""
    parts = read_parts(path)
    parts["word/media/bullet.png"] = picture
    parts["word/_rels/numbering.xml.rels"] = (
        f'<Relationships xmlns="{P}"><Relationship Id="rIdBullet" Type="{R}/image" '
        'Target="media/bullet.png"/></Relationships>'
    ).encode()
    numbering = xml(parts, "word/numbering.xml")
    bullet = etree.fromstring(
        f'<w:numPicBullet xmlns:w="{W}" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:r="{R}" '
        'w:numPicBulletId="0"><w:pict><v:shape><v:imagedata r:id="rIdBullet"/></v:shape>'
        "</w:pict></w:numPicBullet>"
    )
    numbering.insert(0, bullet)
    parts["word/numbering.xml"] = encoded(numbering)
    types = xml(parts, "[Content_Types].xml")
    if not any(item.get("Extension") == "png" for item in types):
        etree.SubElement(types, f"{{{types.nsmap[None]}}}Default", Extension="png").set(
            "ContentType", "image/png"
        )
        parts["[Content_Types].xml"] = encoded(types)
    write_parts(parts, path)


@pytest.mark.parametrize(("color", "state"), [("red", "ready"), ("blue", "failed")])
def test_a_changed_picture_bullet_refuses_the_final(
    final: tuple[Workspace, str], monkeypatch: pytest.MonkeyPatch, color: str, state: str
) -> None:
    """Round 2: a picture bullet ships only by the digest of its own bytes."""
    ws, job = final
    write_narrative(ws, job, "ch4.concluzii", "Consumul a scăzut după modernizare.")
    reviewed = frozenset({hashlib.sha256(_png("red")).hexdigest()})
    monkeypatch.setitem(base_numeric._ALLOWED, "synthetic", reviewed)  # pyright: ignore[reportPrivateUsage]
    built = render.build_base

    def with_bullet(plan: object, *, output: Path, **kwargs: object) -> Path:
        path = built(plan, output=output, **kwargs)
        _picture_bullet(path, _png(color))
        return path

    monkeypatch.setattr(render, "build_base", with_bullet)
    record = run_render(ws, job, "final")
    assert record["state"] == state, record["error"]
    if state == "failed":
        assert _failure(ws, job)["code"] == "audit_package"
        assert record["error"] == "unreviewed picture bullet: word/media/bullet.png"
        assert outputs(ws, job) == []


@pytest.mark.parametrize("kind", ["draft", "final"])
def test_both_exports_refuse_ai_in_narrative_fields(final, kind) -> None:
    ws, job = final
    write_narrative(ws, job, "ch4.concluzii", "Text generat automat.")
    record = run_render(ws, job, kind)
    assert record["state"] == "failed"
    assert _failure(ws, job)["code"] == "audit_ai_wording"
    assert "narrative.ch4.concluzii" in str(record["error"])
    assert outputs(ws, job) == []


def test_scan_refuses_narrative_wording_even_if_the_rendered_body_omits_it(tmp_path: Path) -> None:
    path = tmp_path / "clean.docx"
    Document().save(path)
    field = Field(
        id="field",
        job_id="job",
        key="narrative.ch6.measure.1",
        label="Text",
        value_type="text",
        value="Generată automat.",
        state="supplied",
        presence="found",
    )
    assert ai_wording_hits(path, [field]) == ["narrative.ch6.measure.1"]
