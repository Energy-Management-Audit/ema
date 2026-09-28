"""Chapter 6 replaces only the measures region and keeps the base's table styling."""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from ema.audit.catalogue import CATALOGUE
from ema.audit.chapter_six import (
    MEASURE_HEADER,
    SYNTHESIS_HEADER,
    ChapterSixPlan,
    PlannedMeasure,
    render_chapter_six,
)

TITLES = {section.id: section.title for section in CATALOGUE}


def _base(path: Path, *, chapter_five: bool) -> Path:
    doc = Document()
    doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    doc.styles.add_style("TOC 2", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("TOC one", style="TOC 1")
    doc.add_paragraph("TOC two", style="TOC 2")
    if chapter_five:
        doc.add_paragraph(TITLES["ch5"], style="Heading 1")
    for key, style in (
        ("ch6", "Heading 1"),
        ("ch6.indicatori", "Heading 2"),
        ("ch6.generale", "Heading 2"),
        ("ch6.specifice", "Heading 2"),
    ):
        doc.add_paragraph(TITLES[key], style=style)
    doc.add_paragraph("Old introduction")
    doc.add_paragraph("Old bullet", style="List Bullet")
    doc.add_paragraph("Old measure", style="Heading 3")
    doc.add_paragraph("Old narrative")
    doc.add_paragraph("În tabelul numărul 5.1 se prezintă estimarea")
    doc.add_paragraph("Tabelul 5.1")
    doc.add_paragraph("Sinteza măsurii de eficiență energetică propusă pentru implementare")
    measure = doc.add_table(rows=3, cols=6)
    measure.cell(0, 2).merge(measure.cell(0, 3))
    for row in measure.rows:
        for cell in row.cells:
            cell.text = "Prototype"
    doc.add_paragraph("Costurile sunt estimate")
    doc.add_paragraph(TITLES["ch6.sinteza"], style="Heading 2")
    doc.add_paragraph("Pentru creșterea eficienței energetice")
    doc.add_paragraph("Tabelul 5.2")
    doc.add_paragraph("Sinteza măsurilor de eficiență energetică propuse pentru implementare")
    synthesis = doc.add_table(rows=3, cols=5)
    synthesis.cell(0, 2).merge(synthesis.cell(0, 3))
    for row in synthesis.rows:
        for cell in row.cells:
            cell.text = "Prototype"
    doc.add_paragraph("Retained closing paragraph")
    doc.add_paragraph(TITLES["ch7"], style="Heading 1")
    doc.add_paragraph("Financing preserved")
    doc.save(path)
    return path


def _plan(closing: str | None = "Concluzie verificată.") -> ChapterSixPlan:
    return ChapterSixPlan(
        company_name="Synthetic factory",
        closing=closing,
        measures=(
            PlannedMeasure(
                title="Lighting",
                effect="Less electricity",
                saving_tep=1234.5,
                co2_t=12.345,
                investment_thousand_lei=40,
                payback_years=0.5,
                cost_note=None,
                narrative=None,
            ),
            PlannedMeasure(
                title="Insulation",
                effect=None,
                saving_tep=None,
                co2_t=None,
                investment_thousand_lei=None,
                payback_years=None,
                cost_note="Quote pending",
                narrative="Verified description",
            ),
        ),
    )


def test_render_measures_with_and_without_chapter_five(tmp_path: Path) -> None:
    for chapter_five, expected in ((True, "6"), (False, "5")):
        base = _base(tmp_path / f"base-{expected}.docx", chapter_five=chapter_five)
        output = tmp_path / f"output-{expected}.docx"
        report = render_chapter_six(base, output, _plan(), ("Old measure",))
        doc = Document(output)
        paragraphs = [item.text for item in doc.paragraphs]
        text = "\n".join(paragraphs)
        assert "Synthetic factory" in text
        assert "lighting;" in text and "insulation." in text
        assert f"Tabelul {expected}.1" in text
        assert f"Tabelul {expected}.3" in text
        assert "Retained closing paragraph" not in text
        assert "Concluzie verificată." in text
        assert "Financing preserved" in text
        assert "Old measure" not in text
        assert text.count("[de completat]") == 1
        assert [number.number for number in report.numbers] == [
            f"{expected}.1",
            f"{expected}.2",
            f"{expected}.3",
        ]
        assert len(doc.tables) == 3
        assert len(doc.tables[0].rows) == 3
        assert len(doc.tables[2].rows) == 4
        assert doc.tables[0].cell(2, 2).text == "1.234,5"
        assert doc.tables[0].cell(2, 3).text == "12,35"
        assert doc.tables[0].cell(2, 5).text == "0,5"
        assert doc.tables[1].cell(2, 2).text == "[de completat]"


def _cells(table: object) -> list[list[str]]:
    rows = table._tbl.tr_lst  # type: ignore[attr-defined]
    return [
        ["".join(t.text or "" for t in tc.iter(tc.tag.replace("tc", "t"))) for tc in tr.tc_lst]
        for tr in rows
    ]


def test_headers_are_her_labels_and_the_closing_is_the_reviewed_text(tmp_path: Path) -> None:
    base = _base(tmp_path / "base.docx", chapter_five=True)
    output = tmp_path / "output.docx"
    render_chapter_six(base, output, _plan(closing=None), ("Old measure",))
    doc = Document(output)
    measure, synthesis = _cells(doc.tables[0]), _cells(doc.tables[2])
    assert measure[:2] == [list(row) for row in MEASURE_HEADER]
    assert synthesis[:2] == [list(row) for row in SYNTHESIS_HEADER]
    paragraphs = [item.text for item in doc.paragraphs]
    closing = paragraphs.index(TITLES["ch7"]) - 1
    assert paragraphs[closing] == "[de completat]"
    assert "Retained closing paragraph" not in paragraphs
