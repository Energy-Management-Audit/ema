"""Her cover, fixed prose, headings and header lines: reviewed digests and bound sources (D1-D3)."""

from __future__ import annotations

import hashlib
import io
import json
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from ema.audit import base_numeric
from ema.audit.base_anchor import MARKER, anchor_document, assert_markers, save_anchor_map
from ema.audit.render_bindings import binding_values, fill_bindings
from ema.audit.render_steps import body_counts
from ema.core.review.models import Field

IDENTITY = ("Acme Industrie SRL", "Oras")
REVIEWED = (
    "Elaborator:",
    "Autorizația nr. 12 din 01.02.2020",
    "PERIOADA 2021-2023",
    "Audit energetic pe întregul contur al {client}",
    "Audit energetic {year}",
)


def _digest(text: str | bytes) -> str:
    return hashlib.sha256(text if isinstance(text, bytes) else text.strip().encode()).hexdigest()


def _png(color: str) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (4, 4), color).save(buffer, "PNG")
    return buffer.getvalue()


@pytest.fixture(autouse=True)
def reviewed(monkeypatch: pytest.MonkeyPatch) -> None:
    group = frozenset({*(_digest(text) for text in REVIEWED), _digest(_png("red"))})
    monkeypatch.setitem(base_numeric._ALLOWED, "synthetic", group)  # pyright: ignore[reportPrivateUsage]


def _cover(picture: bytes = _png("red")) -> object:
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].text = "Acme Industrie  SRL"
    cell.add_paragraph("Sediul", style="Body Text")
    cell.add_paragraph("")
    cell.add_paragraph("Str. Exemplu 1, Oras Test")
    cell.add_paragraph("Elaborator:")
    cell.add_paragraph("Autorizația nr. 12 din 01.02.2020")
    cell.add_paragraph("Autorizația nr. 13 din 01.02.2020")
    cell.add_paragraph().add_run().add_picture(io.BytesIO(picture))
    cell.add_paragraph("August 2023")
    document.add_paragraph("Lucrare întocmită de noi.")
    document.add_paragraph("DESCRIEREA ŞI SCOPUL AUDITULUI", style="Heading 1")
    document.add_paragraph("Programul 2021 al Acme Industrie SRL.")
    return document


def _texts(document: object) -> list[str]:
    cell = document.tables[0].cell(0, 0)  # type: ignore[attr-defined]
    return [paragraph.text for paragraph in cell.paragraphs]


def test_front_is_fixed_by_style_or_reviewed_digest_never_with_identity_or_new_number() -> None:
    document = _cover()
    anchors = anchor_document(document, "Atelier Exemplu SRL", IDENTITY)
    texts = _texts(document)
    assert texts[1] == "Sediul"  # styled
    assert texts[4] == "Elaborator:"  # reviewed digest, no style
    assert texts[5] == "Autorizația nr. 12 din 01.02.2020"  # reviewed number
    assert texts[6] == MARKER  # a number nobody reviewed
    assert texts[3] == MARKER  # an identity term
    paragraphs = document.paragraphs  # type: ignore[attr-defined]
    assert paragraphs[0].text == MARKER  # unstyled, unreviewed
    front = {anchor.slot: anchor for anchor in anchors if anchor.section == "front"}
    assert front["body_0_4"].classification == "fixed"
    assert front["body_0_6"].classification == "variable"
    assert_markers(document, anchors)


def test_a_picture_is_kept_only_when_its_bytes_were_reviewed() -> None:
    kept, other = _cover(_png("red")), _cover(_png("blue"))
    anchor_document(kept, "Atelier Exemplu SRL", IDENTITY)
    anchor_document(other, "Atelier Exemplu SRL", IDENTITY)
    drawing = qn("w:drawing")
    assert next(kept.tables[0].cell(0, 0).paragraphs[7]._p.iter(drawing), None) is not None  # type: ignore[attr-defined]
    replaced = other.tables[0].cell(0, 0).paragraphs[7]  # type: ignore[attr-defined]
    assert next(replaced._p.iter(drawing), None) is None and replaced.text == MARKER


def test_bindings_are_judged_on_her_text_and_filled_from_sources(tmp_path: Path) -> None:
    document = _cover()
    anchors = anchor_document(document, "Atelier Exemplu SRL", IDENTITY)
    bound = {anchor.slot: anchor.binding for anchor in anchors if anchor.binding}
    assert bound == {"body_0_0": "client_name", "body_0_3": "address", "body_0_8": "report_month"}
    docx, mapping = tmp_path / "base.docx", tmp_path / "base.anchors.json"
    document.save(str(docx))  # type: ignore[attr-defined]
    save_anchor_map(mapping, "sha", anchors)
    assert {item["binding"] for item in json.loads(mapping.read_text())["anchors"]} == {
        None,
        "client_name",
        "address",
        "report_month",
    }
    address = Field(
        id="a",
        job_id="j",
        key="audit.address",
        label="Adresa",
        value_type="text",
        value="Str. Nouă 1",
        state="supplied",
        presence="found",
    )
    values = binding_values("Atelier Exemplu SRL", [address], date(2026, 9, 27))
    assert values == {
        "client_name": "Atelier Exemplu SRL",
        "address": "Str. Nouă 1",
        "report_month": "Septembrie 2026",
        "report_year": "2026",
        "cover_photo": None,
    }
    assert fill_bindings(docx, mapping, values) == []
    texts = _texts(Document(str(docx)))
    assert (texts[0], texts[3], texts[8]) == (
        "Atelier Exemplu SRL",
        "Str. Nouă 1",
        "Septembrie 2026",
    )


def test_a_missing_or_rejected_address_stays_a_marker(tmp_path: Path) -> None:
    document = _cover()
    anchors = anchor_document(document, "Atelier Exemplu SRL", IDENTITY)
    docx, mapping = tmp_path / "base.docx", tmp_path / "base.anchors.json"
    document.save(str(docx))  # type: ignore[attr-defined]
    save_anchor_map(mapping, "sha", anchors)
    rejected = Field(
        id="a",
        job_id="j",
        key="audit.address",
        label="Adresa",
        value_type="text",
        value="Str. Nouă 1",
        state="supplied",
        presence="found",
        review="rejected",
    )
    values = binding_values("Atelier Exemplu SRL", [rejected], date(2026, 1, 5))
    assert values["address"] is None and values["report_month"] == "Ianuarie 2026"
    assert binding_values("X", [], date(2026, 1, 5))["address"] is None
    blank = rejected.model_copy(update={"review": "corrected", "value": "  "})
    assert binding_values("X", [blank], date(2026, 1, 5))["address"] is None  # round 1 B2
    assert fill_bindings(docx, mapping, values) == ["address"]
    assert _texts(Document(str(docx)))[3] == MARKER


def _headings(period: str) -> object:
    document = Document()
    for text, style in (
        ("3. DESCRIEREA SITUAŢIEI EXISTENTE", "Heading 1"),
        ("3.1. DESCRIEREA UTILITĂȚILOR DIN CADRUL ACME INDUSTRIE SRL", "Heading 2"),
        ("3.2. DESCRIEREA FLUX VOPSIRE", "Heading 2"),
        ("3.2.1. DESCRIEREA SECȚIEI VOPSIRE", "Heading 3"),
        (f"6. MĂSURI DE CREȘTERE A EFICIENŢEI ENERGETICE {period}", "Heading 1"),
    ):
        document.add_paragraph(text, style=style)
        document.add_paragraph("Text.")
    return document


def test_heading_slots_take_the_client_her_reviewed_period_or_a_marker() -> None:
    document = _headings("PERIOADA 2021-2023")
    anchors = anchor_document(document, "Atelier Exemplu SRL", IDENTITY)
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]  # type: ignore[attr-defined]
    assert headings == [
        "3. DESCRIEREA SITUAŢIEI EXISTENTE",
        "3.1. DESCRIEREA UTILITĂȚILOR DIN CADRUL Atelier Exemplu SRL",
        "3.2. DESCRIEREA FLUXULUI TEHNOLOGIC",
        f"3.2.1. DESCRIEREA SECȚIEI {MARKER}",
        "6. MĂSURI DE CREȘTERE A EFICIENŢEI ENERGETICE PERIOADA 2021-2023",
    ]
    headings_slots = {f"body_{index}_0" for index in range(0, 10, 2)}
    variable = {
        anchor.slot
        for anchor in anchors
        if anchor.slot in headings_slots and anchor.classification == "variable"
    }
    assert variable == {"body_6_0"}  # only the heading where a marker remains is bookmarked
    unreviewed = _headings("PERIOADA 2019-2020")
    anchor_document(unreviewed, "Atelier Exemplu SRL", IDENTITY)
    assert (
        unreviewed.paragraphs[8].text == f"6. MĂSURI DE CREȘTERE A EFICIENŢEI ENERGETICE {MARKER}"
    )  # type: ignore[attr-defined]


def _field(paragraph: object, instruction: str, cached: str) -> None:
    for kind, text in (("begin", None), (None, instruction), ("separate", None), (None, cached)):
        run = OxmlElement("w:r")
        if kind is not None:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        elif text == instruction:
            node = OxmlElement("w:instrText")
            node.text = text
        else:
            node = OxmlElement("w:t")
            node.text = text
        run.append(node)
        paragraph._p.append(run)  # type: ignore[attr-defined]
    end = OxmlElement("w:r")
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), "end")
    end.append(node)
    paragraph._p.append(end)  # type: ignore[attr-defined]


def test_header_lines_keep_her_template_and_the_page_fields(tmp_path: Path) -> None:
    document = Document()
    header = document.sections[0].header
    header.paragraphs[0].text = "Audit energetic pe întregul contur al Acme Industrie SRL"
    header.add_paragraph("Audit energetic 2023")
    header.add_paragraph("Raport Acme Industrie SRL 2023")
    footer = document.sections[0].footer.paragraphs[0]
    footer.add_run("Pagina ")
    _field(footer, " PAGE ", "3")
    footer.add_run(" din ")
    _field(footer, " NUMPAGES ", "40")
    document.add_paragraph("DESCRIEREA ŞI SCOPUL AUDITULUI", style="Heading 1")
    anchors = anchor_document(document, "Atelier Exemplu SRL", IDENTITY)
    lines = [paragraph.text for paragraph in header.paragraphs]
    assert lines == [
        "Audit energetic pe întregul contur al Atelier Exemplu SRL",
        f"Audit energetic {MARKER}",
        MARKER,
    ]
    parts = {anchor.slot: (anchor.classification, anchor.binding) for anchor in anchors}
    assert parts["header_1_0"] == ("fixed", None)
    assert parts["header_1_1"] == ("variable", "report_year")
    assert parts["header_1_2"] == ("variable", None)
    assert parts["footer_1_0"] == ("fixed", None)
    instructions = [node.text for node in footer._p.iter(qn("w:instrText"))]
    assert instructions == [" PAGE ", " NUMPAGES "]
    docx, mapping = tmp_path / "base.docx", tmp_path / "base.anchors.json"
    document.save(str(docx))
    save_anchor_map(mapping, "sha", anchors)
    fill_bindings(docx, mapping, binding_values("Atelier Exemplu SRL", [], date(2026, 9, 27)))
    filled = Document(str(docx)).sections[0].header.paragraphs
    assert [paragraph.text for paragraph in filled][:2] == [
        "Audit energetic pe întregul contur al Atelier Exemplu SRL",
        "Audit energetic 2026",
    ]


def test_prose_next_to_a_page_field_is_a_marker(tmp_path: Path) -> None:
    """Round 2: a PAGE field keeps only her page numbering, never the prose around it."""
    document = Document()
    footer = document.sections[0].footer
    lines = [footer.paragraphs[0], *(footer.add_paragraph() for _ in range(4))]
    lines[0].add_run("Fabrica nouă — pagina ")
    _field(lines[0], " PAGE ", "3")
    lines[1].add_run("Pagina ")
    _field(lines[1], " PAGE ", "3")
    lines[1].add_run(" din ")
    _field(lines[1], " NUMPAGES ", "40")
    lines[1].add_run(" — linia nouă")
    _field(lines[2], " PAGE   \\* MERGEFORMAT ", "3")
    lines[3].add_run("Pagina ")
    _field(lines[3], " PAGE ", "3")
    lines[3].add_run(" din ")
    _field(lines[3], " NUMPAGES ", "40")
    _field(lines[4], " DATE ", "")  # blank until Word updates it: not a blank line
    document.add_paragraph("DESCRIEREA ŞI SCOPUL AUDITULUI", style="Heading 1")
    anchors = anchor_document(document, "Atelier Exemplu SRL", IDENTITY)
    parts = {anchor.slot: anchor.classification for anchor in anchors}
    assert [parts[f"footer_1_{index}"] for index in range(5)] == [
        "variable",
        "variable",
        "fixed",
        "fixed",
        "variable",
    ]
    assert [line.text for line in lines[:2]] == [MARKER, MARKER]
    assert _counted(document, tmp_path) == [("front", "Prima pagină şi cuprinsul")] * 3


def _counted(document: object, tmp_path: Path) -> list[tuple[str, str]]:
    path = tmp_path / "anchored.docx"
    document.save(str(path))  # type: ignore[attr-defined]
    return body_counts(path).markers


def test_a_picture_in_a_kept_paragraph_needs_its_own_digest(tmp_path: Path) -> None:
    """Round 1 B1: no style, text digest or fixed section keeps a picture nobody reviewed."""
    document = Document()
    styled = document.add_paragraph(style="Body Text")
    styled.add_run().add_picture(io.BytesIO(_png("blue")))
    mixed = document.add_paragraph("Elaborator:")
    mixed.add_run().add_picture(io.BytesIO(_png("blue")))
    reviewed = document.add_paragraph("Elaborator:")
    reviewed.add_run().add_picture(io.BytesIO(_png("red")))
    document.add_paragraph("DESCRIEREA ŞI SCOPUL AUDITULUI", style="Heading 1")
    fixed = document.add_paragraph("Text fix al capitolului.")
    fixed.add_run().add_picture(io.BytesIO(_png("blue")))
    anchor_document(document, "Atelier Exemplu SRL", IDENTITY)
    paragraphs = document.paragraphs
    drawing = qn("w:drawing")
    for index in (0, 1, 4):
        assert paragraphs[index].text == MARKER
        assert next(paragraphs[index]._p.iter(drawing), None) is None
    assert paragraphs[2].text == "Elaborator:"
    assert next(paragraphs[2]._p.iter(drawing), None) is not None
    # The final gate counts them: audit_markers refuses it.
    assert [section for section, _ in _counted(document, tmp_path)] == ["front", "front", "ch1"]


def test_an_unreviewed_header_line_or_picture_is_a_marker(tmp_path: Path) -> None:
    """Round 1 B4: only page-field lines, her templates and reviewed text stay in a header."""
    document = Document()
    header = document.sections[0].header
    header.paragraphs[0].text = "Fabrica de vopsele din zona industrială"
    header.add_paragraph("Elaborator:")
    header.add_paragraph().add_run().add_picture(io.BytesIO(_png("blue")))
    header.add_paragraph().add_run().add_picture(io.BytesIO(_png("red")))
    header.add_paragraph("")
    document.add_paragraph("DESCRIEREA ŞI SCOPUL AUDITULUI", style="Heading 1")
    anchors = anchor_document(document, "Atelier Exemplu SRL", IDENTITY)
    lines = header.paragraphs
    assert [line.text for line in lines[:2]] == [MARKER, "Elaborator:"]
    assert lines[2].text == MARKER and next(lines[2]._p.iter(qn("w:drawing")), None) is None
    assert next(lines[3]._p.iter(qn("w:drawing")), None) is not None
    parts = {anchor.slot: anchor.classification for anchor in anchors}
    assert [parts[f"header_1_{index}"] for index in range(5)] == [
        "variable",
        "fixed",
        "variable",
        "fixed",
        "fixed",
    ]
    assert _counted(document, tmp_path) == [("front", "Prima pagină şi cuprinsul")] * 2


def test_a_drawing_without_a_picture_is_judged_not_kept_as_blank() -> None:
    """A chart or shape in an empty paragraph is not a blank line: outside fixed text it goes."""
    document = Document()
    document.add_paragraph("DESCRIEREA SITUAŢIEI EXISTENTE", style="Heading 1")
    chart = document.add_paragraph()
    chart.add_run().add_picture(io.BytesIO(_png("red")))
    for blip in list(chart._p.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")):
        blip.getparent().remove(blip)  # a drawing that holds no picture, as a chart does
    anchor_document(document, "Atelier Exemplu SRL", IDENTITY)
    assert document.paragraphs[1].text == MARKER
    assert next(document.paragraphs[1]._p.iter(qn("w:drawing")), None) is None
