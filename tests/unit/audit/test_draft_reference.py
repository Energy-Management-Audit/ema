"""#163: ch. 2-3 drafts rewrite a reference audit's section in its order; the client's previous
audit is the reference when the job holds one, and a part without its fact stays a marker."""

import json
import re
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from tests.audit_replay import CH2_DRAFT, audit_job_with_facts
from tests.unit.audit.test_draft_structured import DraftProvider, SupportProvider

from ema.audit.base_anchor import MARKER
from ema.audit.draft_agent import draft_section_run, job_facts
from ema.audit.draft_checks import check_draft
from ema.audit.draft_prompt import PROMPT_VERSION, rule_text
from ema.audit.draft_render import draft_blocks
from ema.audit.draft_schema import DraftText, SectionDraft
from ema.audit.draft_style import PREVIOUS_AUDIT_SLOT, reference_examples, style_examples
from ema.audit.stages import add_document
from ema.clients.registry import get_client, update_client
from ema.core.errors import EmaError
from ema.core.llm.models import default_model
from ema.core.office.blocks import Missing, Paragraph
from ema.core.workspace import Workspace
from ema.core.workspace.slots import validate_slot

SECTION = "ch2.date_generale"


def _audit(path: Path, *parts: tuple[str, str]) -> Path:
    """A synthetic audit whose Date generale holds the parts in order: body, bullet or table."""
    doc = Document()
    doc.add_paragraph("DESCRIEREA ȘI ISTORICUL SOCIETĂȚII", style="Heading 1")
    doc.add_paragraph("Date generale", style="Heading 2")
    for kind, text in parts:
        if kind == "table":
            table = doc.add_table(rows=2, cols=2)
            for cell, value in zip(table.rows[0].cells, text.split(" | "), strict=True):
                cell.text = value
            doc.add_paragraph("Tabel 2.1 Date generale")
        else:
            doc.add_paragraph(text, style="List Bullet" if kind == "bullet" else None)
    doc.add_paragraph("Istoria companiei", style="Heading 2")
    doc.save(path)
    return path


def _base(tmp_path: Path, monkeypatch: Any, *parts: tuple[str, str]) -> None:
    base = _audit(tmp_path / "base.docx", *parts)
    identity = tmp_path / "identity.json"
    identity.write_text(json.dumps(["Atelier Exemplu"]), encoding="utf-8")
    monkeypatch.setenv("EMA_AUDIT_BASE_DOCUMENT", str(base))
    monkeypatch.setenv("EMA_AUDIT_BASE_IDENTITY", str(identity))


def covering(parts: int) -> SectionDraft:
    """The synthetic draft with one item for each of the reference's first `parts` parts."""
    general = DraftText(text="Activitatea respectă cadrul normativ în vigoare.")
    return CH2_DRAFT.model_copy(
        update={
            "paragraphs": [
                CH2_DRAFT.paragraphs[0].model_copy(update={"part": 1}),
                *(general.model_copy(update={"part": number}) for number in range(2, parts + 1)),
            ]
        }
    )


def _request(tmp_path: Path, parts: int) -> dict[str, Any]:
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    provider = DraftProvider([covering(parts)])
    draft_section_run(
        ws,
        job,
        SECTION,
        provider,
        SupportProvider(),
        model_id=default_model("openai").id,
        synthetic=True,
    )
    return provider.requests[0]["sections"][0]


def test_the_draft_request_keeps_the_reference_order(tmp_path: Path, monkeypatch: Any) -> None:
    _base(
        tmp_path,
        monkeypatch,
        ("body", "Societatea are sediul social în localitate."),
        ("body", "Datele de identificare sunt următoarele:"),
        ("bullet", "denumirea societății;"),
        ("bullet", "codul unic de înregistrare;"),
        ("table", "Indicator | Valoare"),
        ("body", "Regimul de lucru este continuu."),
    )
    reference = _request(tmp_path, 4)["reference"]
    # A list's items share its number; a table has none (#163 fix round 1).
    assert [part["part"] for part in reference] == [1, 2, 3, 3, None, 4]
    assert [part["kind"] for part in reference] == [
        "body",
        "body",
        "bullet",
        "bullet",
        "table",
        "body",
    ]
    assert reference[0]["text"].endswith("are sediul social în localitate.")
    assert reference[2]["text"] == "denumirea societății;"
    assert reference[4]["text"] == "{{…}} | {{…}}"
    assert reference[5]["text"].endswith("de lucru este continuu.")
    assert "style_example" not in _request(tmp_path / "again", 4)


def test_redaction_strips_numbers_and_names_from_every_part(tmp_path: Path) -> None:
    base = _audit(
        tmp_path / "base.docx",
        ("body", "Atelier Exemplu are 1.234 angajați din 2026."),
        ("bullet", "furnizorul Ana Popescu livrează 12,5%;"),
        ("table", "ACME 2025 | Brașov"),
    )
    example = style_examples(base, ("Atelier Exemplu",), (SECTION,))[SECTION]
    text = example.text
    for leaked in ("Atelier", "Exemplu", "1.234", "2026", "Ana", "Popescu", "12,5", "ACME"):
        assert leaked not in text
    assert "Brașov" not in text
    assert not re.search(r"\d", text)
    assert [part.kind for part in example.parts] == ["body", "bullet", "table"]
    # The draft writes no table: its header counts no words of the target.
    assert example.words == 7 + 5


def test_the_previous_audit_is_the_reference_over_the_base(
    tmp_path: Path, monkeypatch: Any
) -> None:
    _base(tmp_path, monkeypatch, ("body", "Textul bazei configurate."))
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    base = reference_examples(ws, job, (SECTION,))
    assert base.identity.startswith("base:") and base.identity != "base:none"
    assert [part.text for part in base.examples[SECTION].parts] == ["{{…}} bazei configurate."]
    client = get_client(ws, "synthetic")
    update_client(ws, "synthetic", {"name": "Fabrica Model", "cui": "RO98765"}, client["revision"])
    previous = _audit(
        tmp_path / "previous.docx",
        ("body", "Auditul anterior pentru fabrica model, cod RO98765."),
        ("bullet", "primul element;"),
    )
    add_document(ws, job, previous, PREVIOUS_AUDIT_SLOT)
    reference = reference_examples(ws, job, (SECTION,))
    assert reference.identity.startswith("previous_audit@1:")
    parts = reference.examples[SECTION].parts
    assert [(part.kind, part.text) for part in parts] == [
        ("body", "{{…}} anterior pentru {{…}}, cod {{…}}."),
        ("bullet", "primul element;"),
    ]


def test_the_previous_audit_slot_is_named_and_takes_only_a_docx(tmp_path: Path) -> None:
    assert PREVIOUS_AUDIT_SLOT == "previous_audit"
    validate_slot("audit", "previous_audit")
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    text = tmp_path / "audit.txt"
    text.write_text("Audit", encoding="utf-8")
    with pytest.raises(EmaError) as error:
        add_document(ws, job, text, PREVIOUS_AUDIT_SLOT)
    assert error.value.code == "previous_audit_type"


def test_the_request_lists_section_keys_never_recorded_as_missing(
    tmp_path: Path, monkeypatch: Any
) -> None:
    _base(tmp_path, monkeypatch, ("body", "Text."))
    missing = _request(tmp_path, 1)["missing"]
    assert {"audit.address", "audit.cui", "audit.work_regime"} <= set(missing)
    assert not {"audit.company_name", "audit.employees"} & set(missing)


def _with_missing(*paragraphs: DraftText) -> SectionDraft:
    return SectionDraft(section=SECTION, status="drafted", paragraphs=list(paragraphs))


FIRST = DraftText(
    text="Societatea {{f:audit.company_name}} are {{f:audit.employees}} angajați.",
    fact_ids=["audit.company_name", "audit.employees"],
)
LAST = DraftText(
    text="Datele sunt prezentate în tabelul următor {{c:audit.company_name}}.",
    fact_ids=["audit.company_name"],
)


def test_a_part_without_its_fact_renders_the_marker_in_its_place(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    facts = job_facts(ws, job)
    gap = DraftText(text="", kind="missing", missing_fact_ids=["audit.work_regime"])
    draft = _with_missing(FIRST, gap, LAST)
    check = check_draft(draft, facts, job)
    assert check.fatal == ()
    blocks = draft_blocks(draft, facts, check.review)
    assert isinstance(blocks[0], Paragraph)
    assert blocks[1] == Missing("body", MARKER)
    assert isinstance(blocks[2], Paragraph)
    assert len(blocks) == 3


@pytest.mark.parametrize(
    "item",
    [
        # A key with a value is written, not left missing.
        DraftText(text="", kind="missing", missing_fact_ids=["audit.employees"]),
        # A key outside the section.
        DraftText(text="", kind="missing", missing_fact_ids=["audit.history"]),
        # A missing item prints nothing but the marker.
        DraftText(text="Regim continuu.", kind="missing", missing_fact_ids=["audit.work_regime"]),
        DraftText(text="", kind="missing"),
        # Only a missing item names missing keys.
        DraftText(text="Text general.", missing_fact_ids=["audit.work_regime"]),
    ],
)
def test_an_invalid_missing_item_fails_with_its_rule(tmp_path: Path, item: DraftText) -> None:
    ws = Workspace(tmp_path / "ws")
    job = audit_job_with_facts(ws)
    check = check_draft(_with_missing(FIRST, item), job_facts(ws, job), job)
    assert "missing_item_invalid" in {issue.code for issue in check.fatal}
    assert rule_text("missing_item_invalid").startswith("O parte a referinței")


def test_the_prompt_is_v6() -> None:
    assert PROMPT_VERSION == "audit-draft-v6"
