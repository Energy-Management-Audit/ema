"""S19 data regressions: job ownership, Romanian labels, and source presentation."""

import ast
from decimal import Decimal
from pathlib import Path
from typing import get_args

import pytest
from docx import Document
from tests.unit.audit.test_measures import _form
from tests.unit.audit.test_read_synthetic import _necesar
from tests.workspace_jobs import create_job

from ema.audit.base_cleanup import clean_base
from ema.audit.catalogue import CATALOGUE, FACT_LABELS, AuditFact, fact_spec, field_label
from ema.audit.chapter_five import PlannedReading
from ema.audit.chapter_five_render import _bullet
from ema.audit.draft_render import _value
from ema.audit.fill_tools import FillDocument, FillTools
from ema.audit.measures import run_measures
from ema.audit.read import _record, read_dossier
from ema.audit.reading_labels import READING_LABELS
from ema.audit.readings import _proof
from ema.audit.readings_schema import Phase, Quantity
from ema.audit.visit import VisitPhoto
from ema.core.errors import EmaError
from ema.core.jobs import StageContext
from ema.core.office.sheets import CellRef
from ema.core.review import fields, propose
from ema.core.review.models import Field, Photo
from ema.core.review.store import save_evidence
from ema.core.workspace import Workspace
from ema.energy_data.source import Located


def test_two_jobs_share_questionnaire_and_rerun_without_evidence_collisions(tmp_path):
    ws = Workspace(tmp_path / "workspace")
    source = _necesar(tmp_path / "necesar.xlsx")
    jobs = [create_job(ws, "audit", "synthetic", 2026) for _ in range(2)]
    evidence_sets = []
    for job in jobs:
        first = read_dossier(ws, job, source)
        second = read_dossier(ws, job, source)
        assert first.fields == second.fields
        assert all(field.label != field.key for field in fields(ws, job))
        evidence_sets.append({eid for field in first.fields for eid in field.evidence})
    assert evidence_sets[0].isdisjoint(evidence_sets[1])


def test_two_different_measure_forms_with_same_calculated_result(tmp_path):
    ws = Workspace(tmp_path / "workspace")
    evidence_sets = []
    for index in range(2):
        job = create_job(ws, "audit", "synthetic", 2026)
        form = _form(
            tmp_path / f"measures-{index}.xlsx",
            [[1, f"Măsura {index}", "Economie", "Energie electrică", 100, "MWh", 40, 10, None]],
        )
        ws.set_slot(job, "measures", ws.add_file("synthetic", form))
        run_measures(ws, job)
        before = fields(ws, job)
        run_measures(ws, job)
        assert fields(ws, job) == before
        assert all(field.label != field.key for field in before)
        saving = next(field for field in before if field.key == "audit_measure.1.saving_tep")
        assert saving.value == Decimal("8.6")
        evidence_sets.append(set(saving.evidence))
    assert evidence_sets[0].isdisjoint(evidence_sets[1])


def test_photo_and_fill_evidence_are_job_owned_and_immutable(tmp_path):
    ws = Workspace(tmp_path / "workspace")
    jobs = [create_job(ws, "audit", "synthetic", 2026) for _ in range(2)]
    ids = []
    for job in jobs:
        ctx = StageContext(ws, job, "unused", "readings")
        photo = VisitPhoto(slot="visit/meter/panel/screen.png", sha="a" * 64, name="screen.png")
        proof = _proof(ctx, photo, "meter.panel.screen.voltage_ln.l1", "230", Photo())
        with ws.connect() as db:
            save_evidence(db, job, [proof])
        assert _proof(ctx, photo, "meter.panel.screen.voltage_ln.l1", "230", Photo()) == proof
        tools = FillTools(
            ws, job, "ch2.date_generale", {"fisa.txt": FillDocument("fisa.txt", "Firma Exemplu")}
        )
        tools.record_fact(
            {
                "key": "audit.company_name",
                "value": "Firma Exemplu",
                "name": "fisa.txt",
                "quote": "Firma Exemplu",
            }
        )
        before = fields(ws, job)
        tools.record_fact(
            {
                "key": "audit.company_name",
                "value": "Firma Exemplu",
                "name": "fisa.txt",
                "quote": "Firma Exemplu",
            }
        )
        assert fields(ws, job) == before
        ids.append({proof.id, *before[0].evidence})
    assert ids[0].isdisjoint(ids[1])
    with ws.connect() as db, pytest.raises(EmaError) as caught:
        save_evidence(db, jobs[0], [proof])
    assert caught.value.code == "evidence_immutable"


def test_every_schema_pair_has_romanian_label_unknown_pair_has_marker():
    assert set(READING_LABELS) == {
        (quantity, phase)
        for quantity in get_args(Quantity.__value__)
        for phase in get_args(Phase.__value__)
    }
    for quantity, phase in READING_LABELS:
        reading = PlannedReading(
            key=f"meter.panel.photo.{quantity}.{phase}", label="", value="231.2", unit="V"
        )
        assert _bullet(reading)[0] == READING_LABELS[quantity, phase] + ": "
    assert _bullet(
        PlannedReading(key="meter.panel.photo.unknown.l1", label="", value="231", unit="V")
    ) == ["[de completat]"]


@pytest.mark.parametrize(
    ("key", "value", "kind", "expected"),
    [
        ("audit.history", 1994, "year", "1994"),
        ("audit.employees", 1234, "number", "1.234"),
        ("audit.employees.2024", 1234, "number", "1.234"),
        ("audit.cui", 12345678, "number", "12345678"),
        ("audit.caen_code", 2511, "number", "2511"),
        ("audit.production", Decimal("14947.828020086"), "number", "14.947,83"),
    ],
)
def test_fact_presentation_is_persisted_and_identifiers_are_text(
    tmp_path, key, value, kind, expected
):
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    spec = fact_spec(key, kind)
    value = str(value) if spec.value_type == "text" else value
    field = propose(ws, job, spec, value, [], state="extracted")
    assert _value(Field.model_validate_json(field.model_dump_json())) == expected
    if key.startswith("audit.employees"):
        persisted = fields(ws, job)[0]
        assert persisted.decimals == 0 and persisted.grouping


def test_integer_xls_display_and_source_precision(tmp_path):
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    source = Located(250.0, CellRef("Necesar", 1, 1), displayed_decimals=0)
    field = _record(ws, job, "audit.employees.2024", source, "a" * 64, "questionnaire")
    assert type(field.value) is int and field.value == 250 and _value(field) == "250"
    assert type(Field.model_validate_json(field.model_dump_json()).value) is int
    assert field.label == "Angajaţi 2024"
    source = Located(12.3456, CellRef("Necesar", 1, 2), displayed_decimals=3)
    field = _record(ws, job, "production.produs.2024", source, "a" * 64, "questionnaire")
    assert _value(field) == "12,346"
    new_format = Located(12.3456, source.ref, displayed_decimals=1)
    reread = _record(ws, job, field.key, new_format, "a" * 64, "questionnaire")
    assert reread.value == field.value and reread.revision == field.revision + 1
    assert _value(reread) == "12,3"


def test_record_fact_stores_numeric_identifiers_as_text(tmp_path):
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    tools = FillTools(
        ws,
        job,
        "ch2.date_generale",
        {"fisa.txt": FillDocument("fisa.txt", "CUI 12345678, CAEN 2511")},
    )
    for key, value in (("audit.cui", 12345678), ("audit.caen_code", 2511)):
        tools.record_fact(
            {"key": key, "value": value, "name": "fisa.txt", "quote": "CUI 12345678, CAEN 2511"}
        )
    assert [(field.value_type, _value(field)) for field in fields(ws, job)] == [
        ("text", "2511"),
        ("text", "12345678"),
    ]


def test_every_audit_fact_and_measure_has_label():
    assert set(FACT_LABELS) == set(AuditFact)
    assert all(field_label(fact.value) != fact.value for fact in AuditFact)
    assert field_label("audit_measure.1.saving_tep") == "Măsura 1 – economie (tep)"
    assert field_label("thermal.photo.spot") == "Temperatura în punctul măsurat"
    assert field_label("thermal.photo.max") == "Temperatura maximă"
    assert field_label("thermal.photo.min") == "Temperatura minimă"


def test_every_audit_error_message_uses_handoff_diacritics():
    for source in (Path(__file__).parents[3] / "src/ema/audit").glob("*.py"):
        text = source.read_text(encoding="utf-8")
        for node in ast.walk(ast.parse(text)):
            if (
                isinstance(node, ast.Call)
                and isinstance(node.func, ast.Name)
                and node.func.id == "EmaError"
                and len(node.args) > 1
            ):
                message_source = ast.get_source_segment(text, node.args[1]) or ""
                assert not set("șțȘȚ") & set(message_source), source.name


def test_document_headings_normalised_with_digest_evidence(tmp_path):
    document = Document()
    document.add_heading("DESCRIEREA ŞI SCOPUL AUDITULUI", 1)
    changes = clean_base(document)
    assert document.paragraphs[0].text == "DESCRIEREA ȘI SCOPUL AUDITULUI"
    assert any("F17 ch1 heading:" in item for item in changes)
    assert clean_base(document) == []
    assert all(not set("şţŞŢ") & set(section.title) for section in CATALOGUE)
