"""The HTML-disguised spreadsheet accepted by intake can enter through the UI upload API."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from tests.unit.test_api_files import preview_pdf

from ema.audit.intake import audit_intake
from ema.clients.files import store_upload
from ema.clients.registry import create_client
from ema.core.errors import EmaError
from ema.core.jobs import create_job, run_stage, status, subscribe
from ema.core.workspace import Workspace


def test_html_xls_upload_keeps_name_and_intake_status(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    client = create_client(ws, "Synthetic")["id"]
    job = create_job(ws, "audit", client, 2026)
    checklist = tmp_path / "0. Necesar info.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "diverse"
    for number in range(1, 14):
        sheet.cell(number, 1, number)
        sheet.cell(number, 2, f"Document {number}")
    book.save(checklist)
    ws.set_slot(job, f"dossier/{checklist.name}", ws.add_file(client, checklist))
    stored = store_upload(
        ws, client, io.BytesIO(b"<html><body>synthetic</body></html>"), "1. Source.xls"
    )
    assert stored["name"] == "1. Source.xls"
    assert stored["kind"] == "html"
    ws.set_slot(job, "dossier/1. Source.xls", stored["sha"])
    run = run_stage(ws, job, "intake", audit_intake)
    for _ in subscribe(ws, job):
        pass
    record = next(row for row in status(ws, job).runs if row["id"] == run)
    assert record["state"] == "ready", record["error"]
    with ws.connect() as db:
        artifact = ws.artifact_dir(db, job, "intake", run) / "completeness.json"
    report = json.loads(artifact.read_text(encoding="utf-8"))
    assert (
        next(row for row in report["files"] if row["slot"] == "dossier/1. Source.xls")["status"]
        == "html_as_xls"
    )


@pytest.mark.parametrize(
    "name,payload",
    [
        ("source.pdf", b"<html><body>synthetic</body></html>"),
        ("source.docx", b"<html><body>synthetic</body></html>"),
        ("source.xls", preview_pdf()),
    ],
)
def test_other_mismatches_stay_rejected(tmp_path: Path, name: str, payload: bytes) -> None:
    ws = Workspace(tmp_path / "ws")
    client = create_client(ws, "Synthetic")["id"]
    with pytest.raises(EmaError) as caught:
        store_upload(ws, client, io.BytesIO(payload), name)
    assert caught.value.code == "file_type"
