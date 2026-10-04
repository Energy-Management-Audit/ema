"""A drafted section and a chapter introduction replace their own region of the base (D4)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.unit.audit.test_draft_checks import _fact

from ema.audit.catalogue import CATALOGUE
from ema.audit.draft_checks import DraftReview
from ema.audit.draft_render import render_section
from ema.audit.draft_schema import DraftText, SectionDraft
from ema.audit.render_writers import write_intro
from ema.audit.section_body import own_region
from ema.core.errors import EmaError

TITLES = {section.id: section.title for section in CATALOGUE}
FACTS = {"audit.company_name": _fact("audit.company_name", "Atelier Exemplu")}
NAME = DraftText(text="{{f:audit.company_name}}", fact_ids=["audit.company_name"])


def _numbered(paragraph: object) -> None:
    properties = paragraph._p.get_or_add_pPr()  # type: ignore[attr-defined]
    numbering = OxmlElement("w:numPr")
    level, number = OxmlElement("w:ilvl"), OxmlElement("w:numId")
    level.set(qn("w:val"), "0")
    number.set(qn("w:val"), "1")
    numbering.extend((level, number))
    properties.append(numbering)


def _base(path: Path, *, table: bool = True) -> Path:
    document = Document()
    document.add_paragraph(TITLES["ch2"], style="Heading 1")
    document.add_paragraph("Introducerea capitolului")
    document.add_paragraph("A doua frază a introducerii")
    document.add_paragraph(TITLES["ch2.date_generale"], style="Heading 2")
    document.add_paragraph("[de completat]", style="Body Text")
    _numbered(document.add_paragraph("[de completat]", style="List Paragraph"))
    document.add_paragraph("Tabelul 2. [de completat]", style="Caption")
    if table:
        grid = document.add_table(rows=2, cols=2)
        for row in grid.rows:
            for cell in row.cells:
                cell.text = "vechi"
    document.add_paragraph("[de completat] în plus")
    document.add_paragraph(TITLES["ch2.istorie"], style="Heading 2")
    document.add_paragraph("Istoria rămâne")
    document.save(str(path))
    return path


def _texts(path: Path) -> list[str]:
    return [paragraph.text for paragraph in Document(str(path)).paragraphs]


def _section(path: Path) -> list[str]:
    texts = _texts(path)
    start = texts.index(TITLES["ch2.date_generale"])
    return texts[start + 1 : texts.index(TITLES["ch2.istorie"])]


def _draft(**kwargs: object) -> SectionDraft:
    return SectionDraft(section="ch2.date_generale", status="drafted", **kwargs)  # type: ignore[arg-type]


def test_own_region_ends_at_the_next_heading_of_any_level(tmp_path: Path) -> None:
    document = Document(str(_base(tmp_path / "base.docx")))
    body = list(document.element.body)
    first, end = own_region(document, "ch2")
    assert [body[index].xpath("string(.)") for index in range(first, end)] == [
        "Introducerea capitolului",
        "A doua frază a introducerii",
    ]
    first, end = own_region(document, "ch2.date_generale")
    assert body[end].xpath("string(.)") == TITLES["ch2.istorie"]
    with pytest.raises(EmaError) as missing:
        own_region(document, "ch3")
    assert (missing.value.code, missing.value.detail) == ("draft_prototype", "ch3")


def test_paragraphs_replace_the_whole_region_whatever_its_length(tmp_path: Path) -> None:
    base, output = _base(tmp_path / "base.docx"), tmp_path / "out.docx"
    paragraphs = [
        DraftText(
            text="Societatea {{f:audit.company_name}} produce.", fact_ids=["audit.company_name"]
        ),
        DraftText(
            text="punct {{f:audit.company_name}}.", fact_ids=["audit.company_name"], kind="bullet"
        ),
        *[
            DraftText(
                text=f"rândul {{{{f:audit.company_name}}}} {index}.",
                fact_ids=["audit.company_name"],
            )
            for index in "abcdef"
        ],
    ]
    render_section(base, output, _draft(paragraphs=paragraphs), FACTS, (), job="synthetic")
    written = _section(output)
    assert written[:2] == ["Societatea Atelier Exemplu produce.", "punct Atelier Exemplu."]
    assert len(written) == 9 and "[de completat] în plus" not in written
    assert len(Document(str(output)).tables) == 1
    assert "Tabelul 2. [de completat]" in written
    document = Document(str(output))
    bullet = document.paragraphs[_texts(output).index("punct Atelier Exemplu.")]
    assert bullet._p.find(f".//{qn('w:numPr')}") is not None
    assert _texts(output)[-1] == "Istoria rămâne"
    assert _texts(output)[1] == "Introducerea capitolului"


def test_flagged_items_and_missing_status_become_markers(tmp_path: Path) -> None:
    base, output = _base(tmp_path / "base.docx"), tmp_path / "out.docx"
    draft = _draft(
        paragraphs=[
            DraftText(text="Societatea {{f:audit.company_name}}.", fact_ids=["audit.company_name"]),
            DraftText(text="Firma {{f:audit.company_name}}.", fact_ids=["audit.company_name"]),
        ],
    )
    flagged = (DraftReview("unsupported", "paragraph:0", "claim"),)
    render_section(base, output, draft, FACTS, flagged, job="synthetic")
    written = _section(output)
    assert written == ["[de completat]", "Firma Atelier Exemplu.", "Tabelul 2. [de completat]"]
    missing = SectionDraft(
        section="ch2.date_generale",
        status="missing",
        paragraphs=[
            DraftText(text="Societatea {{f:audit.company_name}}.", fact_ids=["audit.company_name"])
        ],
        missing_fact_ids=["audit.cui"],
    )
    render_section(base, output, missing, FACTS, (), job="synthetic")
    assert _section(output) == [
        "Societatea Atelier Exemplu.",
        "[de completat]",
        "Tabelul 2. [de completat]",
    ]


@pytest.mark.parametrize(
    ("flagged", "expected"),
    [
        ((1,), "prima Atelier Exemplu. [de completat] ultima Atelier Exemplu."),
        ((0, 1, 2), "[de completat]"),
        ((0,), "[de completat] mijloc Atelier Exemplu. ultima Atelier Exemplu."),
        ((2,), "prima Atelier Exemplu. mijloc Atelier Exemplu. [de completat]"),
        ((0, 1), "[de completat] ultima Atelier Exemplu."),
    ],
)
def test_only_flagged_sentences_become_markers(
    tmp_path: Path, flagged: tuple[int, ...], expected: str
) -> None:
    sentences = [f"{word} {{{{f:audit.company_name}}}}." for word in ("prima", "mijloc", "ultima")]
    draft = _draft(
        paragraphs=[DraftText(text=" ".join(sentences), fact_ids=["audit.company_name"])]
    )
    flags = tuple(DraftReview("unsupported", "paragraph:0", "claim", sentences[i]) for i in flagged)
    output = tmp_path / "out.docx"
    render_section(_base(tmp_path / "base.docx"), output, draft, FACTS, flags, job="synthetic")
    assert _section(output)[0] == expected


@pytest.mark.parametrize(
    ("flag", "expected"),
    [
        (
            "mijloc {{f:audit.company_name}}. ultima {{f:audit.company_name}}",
            "prima Atelier Exemplu. [de completat]",
        ),
        (
            "mijloc {{f:audit.company_name}}",
            "prima Atelier Exemplu. [de completat] ultima Atelier Exemplu.",
        ),
    ],
)
def test_fragment_flags_mark_overlapping_sentences(
    tmp_path: Path, flag: str, expected: str
) -> None:
    text = " ".join(
        f"{word} {{{{f:audit.company_name}}}}." for word in ("prima", "mijloc", "ultima")
    )
    draft = _draft(paragraphs=[DraftText(text=text, fact_ids=["audit.company_name"])])
    output = tmp_path / "out.docx"
    render_section(
        _base(tmp_path / "base.docx"),
        output,
        draft,
        FACTS,
        (DraftReview("unsupported", "paragraph:0", "claim", flag),),
        job="synthetic",
    )
    assert _section(output)[0] == expected


@pytest.mark.parametrize("ending", [".", "!", "?"])
def test_fact_sentence_punctuation_does_not_duplicate_period(tmp_path: Path, ending: str) -> None:
    facts = {"audit.company_name": _fact("audit.company_name", "solventi" + ending)}
    draft = _draft(
        paragraphs=[DraftText(text="{{f:audit.company_name}}.", fact_ids=["audit.company_name"])]
    )
    output = tmp_path / "out.docx"
    render_section(_base(tmp_path / "base.docx"), output, draft, facts, (), job="synthetic")
    assert _section(output)[0] == "solventi" + ending


def test_a_citation_renders_nothing(tmp_path: Path) -> None:
    base, output = _base(tmp_path / "base.docx"), tmp_path / "out.docx"
    text = "Societatea produce ambalaje {{c:audit.company_name}}. Firma {{f:audit.company_name}}."
    draft = _draft(paragraphs=[DraftText(text=text, fact_ids=["audit.company_name"])])
    render_section(base, output, draft, FACTS, (), job="synthetic")
    assert _section(output)[0] == "Societatea produce ambalaje. Firma Atelier Exemplu."


def test_a_stored_draft_with_tables_is_redone_but_empty_lists_load() -> None:
    stored = {
        "section": "ch2.date_generale",
        "status": "drafted",
        "paragraphs": [NAME.model_dump()],
    }
    assert SectionDraft.model_validate({**stored, "tables": [], "figures": []}).paragraphs
    with pytest.raises(ValueError, match="tables"):
        SectionDraft.model_validate({**stored, "tables": [{"caption": NAME.model_dump()}]})


def test_an_intro_is_her_text_or_one_marker(tmp_path: Path) -> None:
    base, output = _base(tmp_path / "base.docx"), tmp_path / "out.docx"
    write_intro(base, output, section_id="ch2", text="Primul paragraf.\n\nAl doilea.")
    texts = _texts(output)
    assert texts[1:4] == ["Primul paragraf.", "Al doilea.", TITLES["ch2.date_generale"]]
    write_intro(base, output, section_id="ch2", text=None)
    assert _texts(output)[1:3] == ["[de completat]", TITLES["ch2.date_generale"]]


def test_draft_creates_absent_activity_section(tmp_path: Path) -> None:
    base, output = _base(tmp_path / "base.docx"), tmp_path / "out.docx"
    draft = SectionDraft(
        section="ch2.activitate",
        status="drafted",
        paragraphs=[
            DraftText(
                text="activitatea {{f:audit.business_activity}}.",
                fact_ids=["audit.business_activity"],
            )
        ],
    )
    render_section(
        base,
        output,
        draft,
        {"audit.business_activity": _fact("audit.business_activity", "industrială")},
        (),
        job="synthetic",
    )
    texts = _texts(output)
    assert texts.index(TITLES["ch2.activitate"]) < texts.index(TITLES["ch2.istorie"])
    assert texts[texts.index(TITLES["ch2.activitate"]) + 1] == "activitatea industrială."


def test_two_absent_sections_clone_sibling_heading_and_body(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph(TITLES["ch2"], style="Heading 1")
    sibling = document.add_paragraph(TITLES["ch2.date_generale"], style="Heading 2")
    _numbered(sibling)
    sibling.runs[0].bold = True
    body = document.add_paragraph("original", style="Body Text")
    body.runs[0].italic = True
    history_heading = document.add_paragraph(TITLES["ch2.istorie"], style="Heading 2")
    _numbered(history_heading)
    history_heading.runs[0].bold = True
    document.add_paragraph("history")
    document.add_paragraph(TITLES["ch3"], style="Heading 1")
    document.add_paragraph(TITLES["ch3.flux"], style="Heading 2")
    document.add_paragraph("flow")
    document.add_paragraph(TITLES["ch4"], style="Heading 1")
    base = tmp_path / "base.docx"
    document.save(str(base))
    current = base
    for section_id, key in (
        ("ch2.manager", "audit.energy_manager"),
        ("ch2.activitate", "audit.business_activity"),
    ):
        output = tmp_path / f"{section_id}.docx"
        draft = SectionDraft(
            section=section_id,
            status="drafted",
            paragraphs=[DraftText(text=f"{{{{f:{key}}}}}", fact_ids=[key])],
        )
        render_section(current, output, draft, {key: _fact(key, "x")}, (), job="synthetic")
        current = output
    result = Document(str(current))
    headings = [paragraph for paragraph in result.paragraphs if paragraph.text in TITLES.values()]
    chapter_two = [
        paragraph.text
        for paragraph in headings
        if paragraph.text
        in {
            TITLES[key]
            for key in ("ch2", "ch2.date_generale", "ch2.manager", "ch2.activitate", "ch2.istorie")
        }
    ]
    assert chapter_two == [
        TITLES[key]
        for key in ("ch2", "ch2.date_generale", "ch2.manager", "ch2.activitate", "ch2.istorie")
    ]
    original = next(paragraph for paragraph in headings if paragraph.text == TITLES["ch2.istorie"])
    for section_id in ("ch2.manager", "ch2.activitate"):
        inserted = next(paragraph for paragraph in headings if paragraph.text == TITLES[section_id])
        assert inserted.style.style_id == original.style.style_id
        assert inserted._p.pPr.xml == original._p.pPr.xml
        assert inserted._p.pPr.numPr.xml == original._p.pPr.numPr.xml
        assert inserted._p.r_lst[0].rPr.xml == original._p.r_lst[0].rPr.xml
        marker = result.paragraphs[
            [paragraph.text for paragraph in result.paragraphs].index(TITLES[section_id]) + 1
        ]
        assert marker.style.name == "Normal"
    assert [paragraph.style.name for paragraph in headings] == [
        "Heading 1",
        "Heading 2",
        "Heading 2",
        "Heading 2",
        "Heading 2",
        "Heading 1",
        "Heading 2",
        "Heading 1",
    ]


def test_process_draft_fills_the_first_unit_and_marks_the_others(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph(TITLES["ch3"], style="Heading 1")
    document.add_paragraph(TITLES["ch3.flux"], style="Heading 2")
    document.add_paragraph("[de completat]", style="Body Text")
    for name in ("Alpha", "Beta"):
        document.add_paragraph(f"DESCRIEREA SECȚIEI {name}", style="Heading 3")
        document.add_paragraph("[de completat]", style="Body Text")
    document.add_paragraph(TITLES["ch4"], style="Heading 1")
    base, output = tmp_path / "base.docx", tmp_path / "out.docx"
    document.save(str(base))
    draft = SectionDraft(
        section="ch3.process",
        status="drafted",
        paragraphs=[
            DraftText(
                text="proces {{f:audit.process_sections}}.", fact_ids=["audit.process_sections"]
            )
        ],
    )
    render_section(
        base,
        output,
        draft,
        {"audit.process_sections": _fact("audit.process_sections", "test")},
        (),
        job="synthetic",
    )
    assert _texts(output).count("proces test.") == 1


def test_each_process_unit_renders_only_its_own_paragraphs(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph(TITLES["ch3"], style="Heading 1")
    document.add_paragraph(TITLES["ch3.flux"], style="Heading 2")
    document.add_paragraph("[de completat]", style="Body Text")
    for name in ("Alpha", "Beta", "Gamma"):
        document.add_paragraph(f"DESCRIEREA SECȚIEI {name}", style="Heading 3")
        document.add_paragraph("[de completat]", style="Body Text")
    document.add_paragraph(TITLES["ch4"], style="Heading 1")
    base, output = tmp_path / "base.docx", tmp_path / "out.docx"
    document.save(str(base))
    keys = ("audit.process_sections", "audit.process_sections.2")
    draft = SectionDraft(
        section="ch3.process",
        status="drafted",
        paragraphs=[
            DraftText(text=f"prima {{{{f:{keys[0]}}}}}.", fact_ids=[keys[0]], unit=1),
            DraftText(text=f"urmează {{{{f:{keys[1]}}}}}.", fact_ids=[keys[1]], unit=2),
        ],
    )
    facts = {key: _fact(key, f"etapa {index}") for index, key in enumerate(keys, 1)}
    render_section(base, output, draft, facts, (), job="synthetic")
    texts = _texts(output)
    units = [texts.index(f"DESCRIEREA SECȚIEI {name}") for name in ("Alpha", "Beta", "Gamma")]
    assert [texts[index + 1] for index in units] == [
        "prima etapa 1.",
        "urmează etapa 2.",
        "[de completat]",
    ]
