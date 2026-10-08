"""#163 D3: the ch. 2-3 scorer on synthetic documents: the reference's structure, her final's
length within ±25 % where both have the section, the rest listed only."""

from pathlib import Path

from docx import Document
from tests.golden.ch23_baseline import SectionScore, score


def _audit(path: Path, sections: list[tuple[str, int]]) -> Path:
    """Ch. 2 with the given subsections, each with its words of own text and a caption."""
    doc = Document()
    doc.add_paragraph("DESCRIEREA ȘI ISTORICUL SOCIETĂȚII", style="Heading 1")
    for title, words in sections:
        doc.add_paragraph(title, style="Heading 2")
        doc.add_paragraph(" ".join(["cuvânt"] * words))
        doc.add_paragraph("Tabel 2.1 Date care nu se numără")
    doc.save(path)
    return path


GENERAL, HISTORY, PLACE = "Date generale", "Istoria companiei", "Localizarea companiei"
MANAGER = "Date privind managerul energetic al activității economice"


def test_the_reference_structure_at_her_length_passes(tmp_path: Path) -> None:
    reference = _audit(tmp_path / "reference.docx", [(GENERAL, 80), (HISTORY, 10), (MANAGER, 5)])
    final = _audit(tmp_path / "final.docx", [(GENERAL, 100), (HISTORY, 40), (PLACE, 60)])
    draft = _audit(tmp_path / "draft.docx", [(GENERAL, 125), (HISTORY, 30), (MANAGER, 300)])
    baseline = score(draft, reference, final)
    assert baseline.draft_order == ("ch2", "ch2.date_generale", "ch2.istorie", "ch2.manager")
    assert baseline.structure_kept
    assert baseline.sections[1] == SectionScore("ch2.date_generale", 125, 100)
    # Not in her final: listed, never scored.
    assert baseline.sections[3] == SectionScore("ch2.manager", 300, None)
    assert baseline.absent == ("ch2.localizare",)
    assert baseline.passed
    report = baseline.report()
    assert "ch2.manager: words 300, not in her final" in report
    assert "ch2.localizare: in her final, not in the draft" in report


def test_a_shared_section_out_of_length_fails_and_is_reported(tmp_path: Path) -> None:
    final = _audit(tmp_path / "final.docx", [(GENERAL, 100), (HISTORY, 40)])
    draft = _audit(tmp_path / "draft.docx", [(GENERAL, 74), (HISTORY, 40)])
    baseline = score(draft, final, final)
    assert not baseline.passed
    assert "ch2.date_generale: words 74/100, ratio 0.74 OUT" in baseline.report()


def test_a_draft_off_the_reference_structure_fails(tmp_path: Path) -> None:
    reference = _audit(tmp_path / "reference.docx", [(GENERAL, 10), (HISTORY, 10)])
    for name, sections in (
        ("moved", [(HISTORY, 10), (GENERAL, 10)]),
        ("missing", [(GENERAL, 10)]),
        ("extra", [(GENERAL, 10), (HISTORY, 10), (PLACE, 10)]),
    ):
        draft = _audit(tmp_path / f"{name}.docx", sections)
        baseline = score(draft, reference, draft)
        assert not baseline.structure_kept, name
        assert not baseline.passed, name
        assert "structure CHANGED" in baseline.report()
    assert score(reference, reference, reference).passed
