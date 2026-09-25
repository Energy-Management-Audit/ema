"""Imported audit units retain their own style and numbering definitions."""

from copy import deepcopy

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ema.audit.base_prototypes import import_formatting


def _numbered_source() -> tuple[Document, object]:
    source = Document()
    style = source.styles.add_style("Synthetic Process", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = source.styles["Normal"]
    paragraph = source.add_paragraph("Process", style=style)
    numbering = source.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "900")
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "900")
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), "900")
    num.append(abstract_ref)
    numbering.append(abstract)
    numbering.append(num)
    num_props = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "900")
    num_props.append(num_id)
    paragraph._p.get_or_add_pPr().append(num_props)
    return source, paragraph._p


def test_import_formatting_remaps_style_and_numbering_ids() -> None:
    source, element = _numbered_source()
    target = Document()
    first, second = deepcopy(element), deepcopy(element)
    import_formatting(target, source, [first, second])
    styles = [
        node.get(qn("w:val")) for node in (first, second) for node in node.iter(qn("w:pStyle"))
    ]
    numbers = [
        node.get(qn("w:val")) for node in (first, second) for node in node.iter(qn("w:numId"))
    ]
    assert styles == ["EmaPrototype1", "EmaPrototype1"]
    assert numbers[0] == numbers[1] != "900"
    assert any(
        node.get(qn("w:styleId")) == "EmaPrototype1"
        for node in target.styles.element.iter(qn("w:style"))
    )
    assert any(
        node.get(qn("w:numId")) == numbers[0]
        for node in target.part.numbering_part.element.iter(qn("w:num"))
    )
