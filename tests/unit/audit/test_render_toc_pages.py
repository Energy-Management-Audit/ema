"""Read printed chapter pages from the TOC's page run."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ema.audit.render_steps import toc_pages


def test_title_ending_in_year_does_not_merge_with_page(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:pStyle"))
    paragraph._p.pPr[-1].set(qn("w:val"), "TOC1")
    run = paragraph.add_run("2. Date pentru anul 2030")
    tab = OxmlElement("w:tab")
    run._r.append(tab)
    text = OxmlElement("w:t")
    text.text = "89"
    run._r.append(text)
    output = tmp_path / "toc.docx"
    document.save(str(output))
    assert toc_pages(output) == {2: 89}


def test_pageref_result_without_tab(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("3. Date pentru anul 2031")
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:pStyle"))
    paragraph._p.pPr[-1].set(qn("w:val"), "TOC1")
    instruction = OxmlElement("w:instrText")
    instruction.text = " PAGEREF bookmark \\h "
    paragraph.add_run()._r.append(instruction)
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(separator)
    paragraph.add_run("91")
    output = tmp_path / "pageref.docx"
    document.save(str(output))
    assert toc_pages(output) == {3: 91}
