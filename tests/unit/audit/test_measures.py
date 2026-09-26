"""The measures stage is a sourced, repeatable audit use case."""

from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ema.audit.chapter_six import ChapterSixPlan
from ema.audit.measures import run_measures
from ema.audit.measures_form import write_measures_form
from ema.audit.sections import get_status
from ema.audit.workflow import AuditWorkflow
from ema.core.errors import EmaError
from ema.core.jobs import create_job
from ema.core.review import decide, fields, propose
from ema.core.review.models import Cell, Evidence, FieldSpec
from ema.core.workspace import Workspace


def _form(path: Path, rows: list[list[object]]) -> Path:
    write_measures_form(path)
    book = load_workbook(path)
    for row in rows:
        book["Măsuri propuse"].append(row)
    book.save(path)
    return path


def _job(tmp_path: Path, rows: list[list[object]]) -> tuple[Workspace, str]:
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    path = _form(tmp_path / "measures.xlsx", rows)
    ws.set_slot(job, "measures", ws.add_file("synthetic", path))
    return ws, job


def test_stage_facts_evidence_narrative_and_rerun(tmp_path: Path) -> None:
    ws, job = _job(
        tmp_path,
        [
            [1, "Lighting", "Less use", "Energie electrică", 100, "MWh", 40, 10, None],
            [2, "Insulation", "Less heat", "Gaze naturale", 20, "MWh", 30, None, "Estimate"],
        ],
    )
    result = run_measures(ws, job)
    assert (
        result.measures,
        result.payback_missing,
        result.factor_version,
        result.missing_narratives,
    ) == (2, 1, "2026", 2)
    assert result.plan_path.name == "ch6.json"
    plan = ChapterSixPlan.model_validate_json(result.plan_path.read_text("utf-8"))
    assert plan.measures[0].saving_tep == 8.6
    assert plan.measures[0].co2_t == 22.6
    assert plan.measures[0].payback_years == 4
    by_key = {field.key: field for field in fields(ws, job)}
    assert by_key["audit_measure.count"].value == 2
    assert by_key["audit_measure.1.carrier"].value == "electricity_grid"
    assert by_key["audit_measure.2.payback_years"].presence == "not_found"
    assert by_key["audit_measure.1.saving_tep"].state == "calculated"
    assert by_key["audit_measure.1.saving_tep"].derivation is not None
    assert by_key["audit_measure.1.saving_tep"].derivation.inputs == [
        "audit_measure.1.saving_amount"
    ]
    source = by_key["audit_measure.1.saving_amount"]
    with ws.connect() as db:
        stored = db.execute(
            "SELECT data FROM evidence WHERE id=?", (source.evidence[0],)
        ).fetchone()
        material = db.execute(
            "SELECT source FROM audit_materials WHERE job_id=? AND kind='measures_form'", (job,)
        ).fetchone()
    evidence = Evidence.model_validate_json(stored["data"])
    assert evidence.method == "form" and evidence.provenance == "document"
    assert evidence.locator is not None and evidence.locator.kind == "cell"
    assert material["source"].startswith("sha256:")
    assert get_status(ws, job, "ch6.measure").status.value == "drafted"
    revisions = {key: field.revision for key, field in by_key.items()}
    run_measures(ws, job)
    assert {field.key: field.revision for field in fields(ws, job)} == revisions

    narrative = by_key["narrative.ch6.measure.1"]
    decide(ws, job, narrative.id, "correct", narrative.revision, "user", value="Verified text")
    readiness = AuditWorkflow().readiness(ws, job)
    assert readiness.draft_ok
    assert [issue.code for issue in readiness.blocking].count("narrative_missing") == 1
    assert get_status(ws, job, "ch6.measure").stale


def test_shorter_form_scopes_old_narrative(tmp_path: Path) -> None:
    ws, job = _job(
        tmp_path,
        [
            [1, "Lighting", None, "Energie electrică", 1, "MWh", 2, 1],
            [2, "Heating", None, "Gaze naturale", 1, "MWh", 2, 1],
        ],
    )
    run_measures(ws, job)
    shorter = _form(
        tmp_path / "shorter.xlsx", [[1, "Lighting", None, "Energie electrică", 1, "MWh", 2, 1]]
    )
    ws.set_slot(job, "measures", ws.add_file("synthetic", shorter))
    result = run_measures(ws, job)
    assert result.measures == 1
    assert [issue.code for issue in AuditWorkflow().readiness(ws, job).blocking].count(
        "narrative_missing"
    ) == 1


def test_changed_form_cell_stales_chapter_six(tmp_path: Path) -> None:
    ws, job = _job(
        tmp_path,
        [[1, "Lighting", None, "Energie electrică", 1, "MWh", 2, 1]],
    )
    run_measures(ws, job)
    old = next(field for field in fields(ws, job) if field.key == "audit_measure.1.saving_amount")
    propose(
        ws,
        job,
        FieldSpec(key=old.key, label=old.label, value_type="number", unit="MWh"),
        2,
        [
            Evidence(
                id="changed-cell",
                provenance="document",
                file_sha="synthetic-revision",
                locator=Cell(sheet="Măsuri propuse", ref="E2"),
                method="form",
                retrieved_at=datetime(2026, 1, 1, tzinfo=UTC),
                quote="2",
                highlight="exact",
            )
        ],
        state="supplied",
    )
    AuditWorkflow().readiness(ws, job)
    assert get_status(ws, job, "ch6.measure").stale


def test_missing_slot_and_header_have_contract_errors(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    with pytest.raises(EmaError) as missing:
        run_measures(ws, job)
    assert missing.value.code == "measures_form_missing"
    path = tmp_path / "invalid.xlsx"
    write_measures_form(path)
    book = load_workbook(path)
    book["Măsuri propuse"]["B1"] = "Wrong"
    book.save(path)
    ws.set_slot(job, "measures", ws.add_file("synthetic", path))
    with pytest.raises(EmaError) as invalid:
        run_measures(ws, job)
    assert invalid.value.code == "measures_form_invalid"
