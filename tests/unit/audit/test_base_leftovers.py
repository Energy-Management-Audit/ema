"""Issue #68: the base's own client leaves nothing in fixed text, numbering or layout."""

from types import SimpleNamespace

from docx import Document
from docx.oxml.ns import qn
from tests.unit.audit.test_base_toc_synthetic import _fixture, _headings, _node

from ema.audit.base_cleanup import clean_base
from ema.audit.base_toc import refresh_toc
from ema.audit.chapter_four_blocks import chapter_four_blocks
from ema.audit.chapter_four_water import without_empty_water
from ema.core.office.blocks import Paragraph
from ema.energy_data.carriers import Carrier
from ema.energy_data.factors import FACTORS_2026
from ema.energy_data.model import CarrierSeries, EnergyDataset, Reading


def test_transport_funding_guide_is_cut_by_its_words(monkeypatch):
    document = Document()
    document.add_paragraph("Table of sources")
    document.add_paragraph("")
    document.add_paragraph(
        "Se urmărește ghidul pentru sectorul transportului naval și aerian și "
        "transportului feroviar elaborat de minister."
    )
    document.add_paragraph(
        "Prin Programul-cheie 9: Eficiență energetică în transporturi - sectorul naval, "
        "aerian și feroviar."
    )
    document.add_paragraph("Indicatorul I.1 - Număr de echipamente;")
    document.add_paragraph("")
    document.add_paragraph("O altă sursă de finanțare o reprezintă PNRR.")
    monkeypatch.setattr("ema.audit.base_cleanup.heading_spans_document", lambda _: [])
    changes = clean_base(document)
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert texts == ["Table of sources", "", "O altă sursă de finanțare o reprezintă PNRR."]
    assert len(changes) == 1 and changes[0].startswith("G4 fixed text:")
    assert clean_base(document) == []


def test_funding_chapter_names_the_measures_chapter_as_rendered(monkeypatch):
    document, *_ = _fixture()
    measures = document.add_paragraph("6. Măsuri")
    funding = document.add_paragraph("7. Surse de finanțare")
    reference = document.add_paragraph("planul de măsuri din ")
    reference.add_run("Capitolului 7")
    reference.add_run(".")
    body = list(document.element.body)
    spans = [
        (SimpleNamespace(section_id="ch6", template=None), body.index(measures._p), 0),
        (
            SimpleNamespace(section_id="ch7", template=None),
            body.index(funding._p),
            body.index(reference._p) + 1,
        ),
    ]
    monkeypatch.setattr("ema.audit.base_toc.heading_spans_document", lambda _: spans)
    refresh_toc(document)
    # no measurements chapter in this layout: the measures chapter prints as 5
    assert reference.text == "planul de măsuri din Capitolului 5."
    refresh_toc(document)
    assert reference.text == "planul de măsuri din Capitolului 5."


def test_water_carrier_without_readings_gets_no_block():
    empty = {2025: CarrierSeries({1: Reading(None, "m3")}, Reading(None, "m3"))}
    full = {2025: CarrierSeries(annual=Reading(5, "m3"))}
    base = {Carrier.electricity_grid: {2025: CarrierSeries(annual=Reading(10, "MWh"))}}
    production = ({"product": {2025: CarrierSeries(annual=Reading(100, "t"))}}, {"product": "t"})
    only_empty = EnergyDataset((2025,), {**base, Carrier.water_potable: empty}, *production)
    one_left = EnergyDataset(
        (2025,), {**base, Carrier.water_potable: empty, Carrier.water_industrial: full}, *production
    )

    assert Carrier.water_potable not in without_empty_water(only_empty).carriers
    headings = [
        block.proto
        for block in chapter_four_blocks(only_empty, FACTORS_2026)
        if isinstance(block, Paragraph) and block.proto.startswith("heading:")
    ]
    assert "heading:ch4.apa" not in headings
    assert "heading:ch4.specific_apa" not in headings
    assert "heading:ch4.gaz" in headings

    labels = [
        segment
        for block in chapter_four_blocks(one_left, FACTORS_2026)
        if isinstance(block, Paragraph)
        for segment in block.segments
        if isinstance(segment, str)
    ]
    assert "Consumul de apă industrială" in labels
    assert "Consumul de apă potabilă" not in labels


def test_blank_paragraph_after_the_toc_is_dropped(monkeypatch):
    document, _, _, _ = _fixture()
    document.paragraphs[-2].text = ""
    blank = document.add_paragraph("")
    paragraphs = _headings(document, monkeypatch, [("ch1", 1, "1. Descrierea")])
    refresh_toc(document)
    body = list(document.element.body)
    assert blank._p not in body
    assert body[body.index(paragraphs[0]) - 1].xpath(".//w:hyperlink")


def test_equipment_heading_takes_its_siblings_paragraph_style(monkeypatch):
    document = Document()
    headings = []
    for label, spaced in (
        ("3.1.1. Proces", True),
        ("3.1.2. Proces", False),
        ("3.1.3. Echipamente", False),
    ):
        paragraph = document.add_paragraph(label)
        if spaced:
            paragraph._p.get_or_add_pPr().append(_node("w:spacing", line="240", lineRule="auto"))
        headings.append(paragraph._p)
    body = list(document.element.body)
    items = [
        (SimpleNamespace(section_id=section), body.index(heading), body.index(heading) + 1)
        for section, heading in zip(
            ["ch3.process", "ch3.process", "ch3.equipment"], headings, strict=True
        )
    ]
    monkeypatch.setattr("ema.audit.base_cleanup.heading_spans_document", lambda _: items)
    clean_base(document)
    for heading in headings:
        spacing = heading.pPr.find(qn("w:spacing"))
        assert spacing is not None and spacing.get(qn("w:line")) == "240"
        order = [child.tag for child in heading.pPr]
        assert order.index(qn("w:spacing")) < order.index(qn("w:ind"))
