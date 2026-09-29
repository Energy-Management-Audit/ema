"""Audit Read stores source cells and calculates only supported classifications."""

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from tests.workspace_jobs import create_job

from ema.audit.read import read_dossier
from ema.core.review import fields
from ema.core.workspace import Workspace

MONTHS = (
    "Ianuarie",
    "Februarie",
    "Martie",
    "Aprilie",
    "Mai",
    "Iunie",
    "Iulie",
    "August",
    "Septembrie",
    "Octombrie",
    "Noiembrie",
    "Decembrie",
)


def _necesar(path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "Cons energetice"
    sheet.cell(5, 1, "Consum energie electrica din SEN")
    for column, value in enumerate((2025, *MONTHS, "Total"), 1):
        sheet.cell(6, column, value)
    sheet.cell(7, 1, "[MWh]")
    sheet.cell(7, 2, 12)
    sheet.cell(7, 14, 12)
    sheet.cell(8, 1, "[tep]")
    book.save(path)
    return path


def test_read_dossier_records_cell_evidence_and_calculated_threshold(tmp_path: Path) -> None:
    workspace = Workspace(tmp_path / "workspace")
    job = create_job(workspace, "audit", "synthetic", 2025)
    result = read_dossier(workspace, job, _necesar(tmp_path / "necesar.xlsx"))
    saved = {field.key: field for field in fields(workspace, job)}
    assert result.dataset.years == (2025,)
    assert saved["carrier.electricity_grid.2025.01"].value == Decimal("12")
    assert saved["carrier.electricity_grid.2025"].value == Decimal("12")
    assert saved["audit.tep_class"].value == "below_1000_tep"
    assert saved["carrier.electricity_grid.2025.01"].evidence
    assert saved["audit.tep_class"].derivation is not None
