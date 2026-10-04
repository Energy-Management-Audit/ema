"""Chapter-four formatting stays bounded and never moves or renumbers a heading."""

import pytest
from docx import Document
from docx.shared import Pt
from lxml import etree

from ema.audit.catalogue import CATALOGUE
from ema.audit.chapter_four_format import format_chapter_four
from ema.core.errors import EmaError

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TITLES = {section.id: section.title for section in CATALOGUE}


def test_keeps_chapter_chart_groups_and_centres_missing_cells_at_text_width():
    document = Document()
    document.add_heading(TITLES["ch3"], 1)
    untouched = document.add_table(rows=2, cols=7)
    before = etree.tostring(untouched._tbl)
    chapter = document.add_heading(TITLES["ch4"], 1)
    heading = document.add_heading(TITLES["ch4.electricitate"], 3)
    label = document.add_paragraph("Consumul de energie electrică din rețea")
    drawing = document.add_paragraph()
    etree.SubElement(etree.SubElement(drawing._p, W + "r"), W + "drawing")
    caption = document.add_paragraph("Fig. nr. 4.1 Evoluția lunară a consumului")
    table = document.add_table(rows=2, cols=7)
    for cell in table.rows[1].cells:
        cell.text = "date indisponibile"
    table.rows[1].cells[0].text = "2025"
    table.rows[1].cells[1].text = "1.234,50"
    document.add_heading(TITLES["ch5"], 1)
    body_before = list(document.element.body)
    format_chapter_four(document)
    assert list(document.element.body) == body_before
    assert etree.tostring(untouched._tbl) == before
    for paragraph in (chapter, heading, label, drawing):
        assert paragraph._p.find(W + "pPr/" + W + "keepNext").get(W + "val") == "1"
    assert caption._p.find(W + "pPr/" + W + "keepLines").get(W + "val") == "1"
    assert caption._p.find(W + "pPr/" + W + "keepNext").get(W + "val") == "1"
    section = document.sections[0]
    width = (section.page_width - section.left_margin - section.right_margin) // 635
    assert table._tbl.find(W + "tblPr/" + W + "tblW").get(W + "w") == str(width)
    assert (
        sum(int(c.get(W + "w")) for c in table._tbl.findall(W + "tblGrid/" + W + "gridCol"))
        == width
    )
    for cell in table.rows[1].cells:
        assert cell._tc.find(W + "tcPr/" + W + "vAlign").get(W + "val") == "center"
        if cell.text == "date indisponibile":
            assert cell._tc.find(W + "tcPr/" + W + "noWrap") is not None
        else:
            assert cell._tc.find(W + "tcPr/" + W + "noWrap") is None


def test_missing_chapter_bounds_raise_contextual_ema_error():
    document = Document()
    document.add_heading(TITLES["ch4"], 1)
    with pytest.raises(EmaError) as failure:
        format_chapter_four(document)
    assert failure.value.code == "chapter_four_layout"
    assert "next chapter=None" in failure.value.detail


def test_dash_cells_are_centred_without_wrapping_or_changing_font_size():
    document = Document()
    document.add_heading(TITLES["ch4"], 1)
    table = document.add_table(rows=2, cols=7)
    cell = table.rows[1].cells[1]
    run = cell.paragraphs[0].add_run("—")
    run.font.size = Pt(12)
    document.add_heading(TITLES["ch5"], 1)
    format_chapter_four(document)
    assert cell._tc.find(W + "tcPr/" + W + "noWrap").get(W + "val") == "1"
    assert cell._tc.find(W + "tcPr/" + W + "vAlign").get(W + "val") == "center"
    assert cell.paragraphs[0]._p.find(W + "pPr/" + W + "jc").get(W + "val") == "center"
    assert run.font.size.pt == 12


def test_merged_rows_are_skipped_and_custom_missing_text_stays_on_one_line():
    document = Document()
    document.add_heading(TITLES["ch4"], 1)
    table = document.add_table(rows=2, cols=7)
    table.rows[0].cells[0].merge(table.rows[0].cells[-1]).text = "Header"
    merged = etree.tostring(table._tbl.find(W + "tr"))
    table.rows[1].cells[1].text = "lipsă"
    document.add_heading(TITLES["ch5"], 1)
    format_chapter_four(document, missing_text="lipsă")
    assert etree.tostring(table._tbl.find(W + "tr")) == merged
    assert table.rows[1].cells[1]._tc.find(W + "tcPr/" + W + "noWrap") is not None


@pytest.mark.parametrize("width", ["invalid", "0"])
def test_invalid_table_grid_raises_contextual_ema_error(width):
    document = Document()
    document.add_heading(TITLES["ch4"], 1)
    table = document.add_table(rows=2, cols=7)
    table._tbl.find(W + "tblGrid/" + W + "gridCol").set(W + "w", width)
    document.add_heading(TITLES["ch5"], 1)
    with pytest.raises(EmaError) as failure:
        format_chapter_four(document)
    assert failure.value.code == "chapter_four_layout"
    assert "chapter-four table" in failure.value.detail


def test_long_production_values_get_width_from_shorter_columns_within_text_width():
    document = Document()
    document.add_heading(TITLES["ch4"], 1)
    document.add_heading(TITLES["ch4.productie"], 2)
    table = document.add_table(rows=3, cols=7)
    for row in table.rows[1:]:
        for cell, text in zip(
            row.cells,
            ("2024", "123.456.789,12", "12,34", "1.234.567,89", "2,00", "3,00", "4,00"),
            strict=True,
        ):
            cell.text = text
    document.add_heading(TITLES["ch4.consum"], 2)
    normal = document.add_table(rows=2, cols=7)
    document.add_heading(TITLES["ch5"], 1)
    text_before = [[cell.text for cell in row.cells] for row in table.rows]
    format_chapter_four(document)
    grid = [
        int(column.get(W + "w")) for column in table._tbl.findall(W + "tblGrid/" + W + "gridCol")
    ]
    section = document.sections[0]
    width = (section.page_width - section.left_margin - section.right_margin) // 635
    assert sum(grid) == width
    assert len(grid) == 7  # Her two six-month rows retain the year column and six value columns.
    assert grid[1] >= 10 * 120 + 3 * 60 + 216
    assert grid[1] > grid[3] > grid[2]
    assert [[cell.text for cell in row.cells] for row in table.rows] == text_before
    assert (
        max(
            int(column.get(W + "w"))
            for column in normal._tbl.findall(W + "tblGrid/" + W + "gridCol")
        )
        - min(
            int(column.get(W + "w"))
            for column in normal._tbl.findall(W + "tblGrid/" + W + "gridCol")
        )
        <= len(grid) - 1
    )
    for row in table.rows:
        assert [
            int(cell._tc.find(W + "tcPr/" + W + "tcW").get(W + "w")) for cell in row.cells
        ] == grid


def test_six_long_production_values_fit_without_changing_font_or_text():
    document = Document()
    section = document.sections[0]
    section.left_margin = section.right_margin = Pt(72)
    document.add_heading(TITLES["ch4"], 1)
    document.add_heading(TITLES["ch4.productie"], 2)
    tables = [document.add_table(rows=2, cols=7) for _ in range(2)]
    for table in tables:
        for column, cell in enumerate(table.rows[1].cells):
            cell.text = "2025" if column == 0 else "1.234.567,89"
            cell.paragraphs[0].runs[0].font.size = Pt(12)
    document.add_heading(TITLES["ch5"], 1)
    format_chapter_four(document)
    for table in tables:
        assert len(table.columns) == 7
        assert sum(column.width for column in table.columns) == (
            section.page_width - section.left_margin - section.right_margin
        )
        for cell in table.rows[1].cells[1:]:
            margins = cell._tc.find(W + "tcPr/" + W + "tcMar")
            padding = sum(int(margins.find(W + side).get(W + "w")) for side in ("left", "right"))
            assert cell.width // 635 - padding >= 1260  # 63 pt of text at her 12 pt font
            assert cell.text == "1.234.567,89"
            assert cell.paragraphs[0].runs[0].font.size == Pt(12)


def test_first_half_table_rows_link_to_second_half_caption() -> None:
    document = Document()
    document.add_heading(TITLES["ch4"], 1)
    first_caption = document.add_paragraph("Tabelul 4.1 – 2025")
    first = document.add_table(rows=2, cols=7)
    first.cell(0, 1).text = "Ianuarie"
    second_caption = document.add_paragraph("Tabelul 4.2 – 2025")
    second = document.add_table(rows=2, cols=7)
    second.cell(0, 1).text = "Iulie"
    document.add_heading(TITLES["ch5"], 1)
    format_chapter_four(document)
    for caption in (first_caption, second_caption):
        assert caption._p.find(W + "pPr/" + W + "keepNext").get(W + "val") == "1"
    for row in first.rows:
        for cell in row.cells:
            assert cell.paragraphs[0]._p.find(W + "pPr/" + W + "keepNext") is not None
    assert second.cell(1, 1).paragraphs[0]._p.find(W + "pPr/" + W + "keepNext") is None
