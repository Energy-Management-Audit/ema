"""The shared dossier reader used by the intake and Fill agents."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from tests.unit.audit.test_fill_stage import write_pdf
from tests.workspace_jobs import create_job

from ema.audit.dossier import dossier_documents
from ema.audit.fill_tools import FillDocument
from ema.core.workspace import Workspace
from ema.core.workspace.conversion import active_version, publish_conversion


def test_each_active_dossier_slot_is_read_with_its_file_sha(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "made-up", 2026)
    pdf = write_pdf(tmp_path / "permit.pdf", "Suprafata: 500 m2.")
    note = tmp_path / "fisa.txt"
    note.write_text("Firma Exemplu SRL", encoding="utf-8")
    for source in (pdf, note):
        ws.set_slot(job, f"dossier/{source.name}", ws.add_file("made-up", source))
    ws.set_slot(job, "anexa", ws.add_file("made-up", note))

    documents = dossier_documents(ws, job)

    assert set(documents) == {"permit.pdf", "fisa.txt"}
    permit = documents["permit.pdf"]
    assert permit.page_texts == (permit.text,)
    assert "500 m2" in permit.text
    assert permit.page_images == ("permit.pdf#page=1",)
    version = active_version(ws, job, "dossier/permit.pdf")
    assert version is not None
    assert permit.file_sha == version.file_sha
    assert documents["fisa.txt"] == FillDocument(
        "fisa.txt", "Firma Exemplu SRL", file_sha=documents["fisa.txt"].file_sha
    )


def test_a_converted_doc_is_read_from_its_active_docx(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "made-up", 2026)
    legacy = tmp_path / "fisa.doc"
    legacy.write_bytes(b"legacy word file")
    ws.set_slot(job, "dossier/fisa.doc", ws.add_file("made-up", legacy))
    converted = tmp_path / "converted.docx"
    document = Document()
    document.add_paragraph("Activitate: vopsire industrială.")
    document.save(str(converted))
    original = active_version(ws, job, "dossier/fisa.doc")
    assert original is not None
    published = publish_conversion(ws, job, "dossier/fisa.doc", original, converted)
    assert published is not None

    documents = dossier_documents(ws, job)

    assert documents["fisa.doc"].text == "Activitate: vopsire industrială."
    assert documents["fisa.doc"].file_sha == published.file_sha
