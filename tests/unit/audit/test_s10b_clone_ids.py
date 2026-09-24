"""Synthetic bookmark integrity check for cloned audit units."""

from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn

from ema.audit.base_units import _resize
from ema.audit.inventory import BaseUnit
from ema.core.office.anchors import stamp


def test_cloned_unit_has_unique_bookmark_ids() -> None:
    document = Document()
    paragraph = document.add_paragraph("Process prototype")
    stamp(paragraph._p, "prototype", 42)
    document.add_paragraph("Next section")
    unit = BaseUnit("process", ("Process",), (0, 1))
    _resize(document, (unit,), 2)
    ids = [node.get(qn("w:id")) for node in document.element.iter(qn("w:bookmarkStart"))]
    assert len(ids) == len(set(ids))
