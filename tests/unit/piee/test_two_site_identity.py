"""Two-site declarations keep distinct evidence and review gaps."""

# pyright: reportPrivateUsage=false

import json
from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from tests.unit.piee.synthetic_piee import piee_data
from tests.workspace_jobs import create_job

from ema.core.errors import EmaError
from ema.core.office.anchors import AnchorLedger, stamp
from ema.core.office.sheets import open_book
from ema.core.review import decide, fields
from ema.core.review.evidence import get_evidence
from ema.core.workspace import Workspace
from ema.energy_data.anexa_cells import AnexaData
from ema.energy_data.anexa_identity import read_identity
from ema.piee.identity import identity_values
from ema.piee.identity_map import render_identity
from ema.piee.intake import _record_identity
from ema.piee.review_overlay import apply_review
from ema.piee.review_workflow import PieeWorkflow


def _two_site_anexa(tmp_path: Path) -> AnexaData:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Date generale"
    sheet["B8"] = "Denumirea operatorului economic"
    sheet["D8"] = "Operator SRL, SUCURSALELE SITE-1 si SITE-2"
    sheet["B9"] = "Adresa poștală"
    sheet["D9"] = "Str. Exemplu 1, SITE-1"
    source = tmp_path / "anexa.xlsx"
    workbook.save(source)
    book = open_book(source)
    anexa = AnexaData()
    try:
        read_identity(book, anexa)
    finally:
        book.close()
    return anexa


def test_shifted_two_site_labels_keep_one_address_and_review_gaps(tmp_path: Path) -> None:
    anexa = _two_site_anexa(tmp_path)

    assert anexa.identity["name"].value == "Operator SRL"
    assert anexa.identity["site_1_name"].value == "Site-1"
    assert anexa.identity["site_2_name"].value == "Site-2"
    assert anexa.identity["site_1_name"].ref.a1 == "Date generale!D8"
    assert anexa.identity["site_1_address"].ref.a1 == "Date generale!D9"
    assert "site_2_address" not in anexa.identity
    values = identity_values(anexa, date(2026, 1, 2))
    assert values["client_name"] == "Operator SRL"
    assert values["site_2_address"] is None

    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "piee", "synthetic", 2026)
    _record_identity(ws, job, replace(piee_data(), anexa=anexa), "synthetic-sha")
    by_key = {field.key: field for field in fields(ws, job)}
    assert by_key["identity.site_1_address"].presence == "found"
    for key in (
        "identity.site_2_address",
        "identity.site_1_production_share",
        "identity.site_2_production_share",
    ):
        assert by_key[key].presence == "not_found"
        assert not by_key[key].required
    assert [warning.field_id for warning in PieeWorkflow().readiness(ws, job).warnings] == [
        by_key[key].id
        for key in (
            "identity.site_1_production_share",
            "identity.site_2_address",
            "identity.site_2_production_share",
        )
    ]


def test_manual_site_corrections_have_distinct_evidence_and_render(tmp_path: Path) -> None:
    anexa = _two_site_anexa(tmp_path)
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "piee", "synthetic", 2026)
    _record_identity(ws, job, replace(piee_data(), anexa=anexa), "synthetic-sha")
    share = next(
        field for field in fields(ws, job) if field.key == "identity.site_1_production_share"
    )
    with pytest.raises(EmaError) as invalid:
        decide(ws, job, share.id, "correct", share.revision, "user", value=101)
    assert invalid.value.code == "value_invalid"
    for key, value in (
        ("identity.site_2_address", "Str. Exemplu 2, SITE-2"),
        ("identity.site_1_production_share", 60),
        ("identity.site_2_production_share", 40),
    ):
        field = next(field for field in fields(ws, job) if field.key == key)
        decide(ws, job, field.id, "correct", field.revision, "user", value=value)
    by_key = {field.key: field for field in fields(ws, job)}
    evidence_ids = []
    for key in (
        "identity.site_2_address",
        "identity.site_1_production_share",
        "identity.site_2_production_share",
    ):
        field = by_key[key]
        assert field.presence == "found"
        evidence = get_evidence(ws, field.evidence[0])
        assert evidence.method == "manual"
        assert evidence.locator.kind == "manual"
        evidence_ids.append(evidence.id)
    assert len(set(evidence_ids)) == 3
    assert not PieeWorkflow().readiness(ws, job).warnings
    reviewed = apply_review(replace(piee_data(), anexa=anexa), list(by_key.values()), {})
    assert reviewed.dataset == piee_data().dataset
    values = identity_values(reviewed.anexa, date(2026, 1, 2), analysis_year=2025)
    assert values["site_2_address"] == "Str. Exemplu 2, SITE-2"
    assert values["site_1_production_share"] == "60,00"
    assert values["site_2_production_share"] == "40,00"

    document = Document()
    address = document.add_paragraph("Base: source address")
    stamp(address._p, "body_9", 1)
    preceding = document.add_paragraph("Production body")
    stamp(preceding._p, "number_629", 2)
    document.add_paragraph("Analiza consumului de energie")
    stamped = tmp_path / "stamped.docx"
    document.save(stamped)
    manifest = tmp_path / "identity-spans.json"
    manifest.write_text(
        json.dumps(
            [
                {
                    "owner": "word/document.xml",
                    "slot": "body_9",
                    "start": 6,
                    "end": 20,
                    "key": "address",
                }
            ]
        ),
        encoding="utf-8",
    )
    output = tmp_path / "reviewed.docx"
    render_identity(stamped, manifest, values, output, AnchorLedger(frozenset({"body_9"})))
    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "- Sucursala Site-2: Str. Exemplu 2, SITE-2." in text
    assert "sucursala Site-1 are un procent de 60,00% iar sucursala Site-2 40,00%." in text


@pytest.mark.parametrize(
    ("address", "expected"),
    [
        ("Str. Exemplu, SITE-1", "site_1_address"),
        ("Str. Exemplu, SITE-3", None),
        ("Str. Exemplu, SITE-1 și SITE-2", None),
    ],
)
def test_postal_address_is_attributed_only_to_one_named_site(
    tmp_path: Path, address: str, expected: str | None
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Date generale"
    sheet["A5"] = "Denumirea operatorului economic"
    sheet["C5"] = "Operator SRL, sucursalele SITE-1 și SITE-2"
    sheet["A6"] = "Adresa poștală"
    sheet["C6"] = address
    source = tmp_path / "anexa.xlsx"
    workbook.save(source)
    book = open_book(source)
    anexa = AnexaData()
    try:
        read_identity(book, anexa)
    finally:
        book.close()
    found = {key for key in anexa.identity if key in {"site_1_address", "site_2_address"}}
    assert found == ({expected} if expected else set())
