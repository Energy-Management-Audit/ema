"""Conditional headings replace authored TOC entries without losing page fields."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ema.core.office.package import read_parts, xml
from ema.core.office.run_range import visible_text
from ema.piee.toc import render_toc_from_headings


def test_toc_follows_inserted_and_removed_heading_three_entries(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Intro")
    toc = document.add_table(rows=1, cols=1)
    cell = toc.cell(0, 0)
    for index in range(26):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.text = f"2.2.{index} Old {index + 10}"
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), "_Old")
        paragraph._p.append(link)
        field = OxmlElement("w:r")
        instruction = OxmlElement("w:instrText")
        instruction.text = " PAGEREF _Old \\h "
        field.append(instruction)
        paragraph._p.append(field)
    for name in ("Grid", "Cogeneration", "Coke", "Echivalent de energie"):
        document.add_paragraph(name, style="Heading 3")
    document.add_paragraph("Consum specific", style="Heading 2")
    for name in ("Specific gas", "Specific coke", "Specific water"):
        document.add_paragraph(name, style="Heading 3")
    document.add_paragraph("Audits", style="Heading 1")
    source, output = tmp_path / "source.docx", tmp_path / "toc.docx"
    document.save(source)
    render_toc_from_headings(source, output)
    root = xml(read_parts(output), "word/document.xml")
    body = root.find(qn("w:body"))
    assert body is not None
    entries = list(list(body)[1].iter(qn("w:p")))
    text = [visible_text(entry) for entry in entries]
    assert text[9:13] == [
        "2.2.1.Grid 19",
        "2.2.2.Cogeneration 19",
        "2.2.3.Coke 19",
        "2.2.4.Echivalent de energie 19",
    ]
    assert text[13:16] == [
        "2.3.1.Specific gas 26",
        "2.3.2.Specific coke 26",
        "2.3.3.Specific water 26",
    ]
    assert all(
        any("PAGEREF _TocEma" in (node.text or "") for node in entry.iter(qn("w:instrText")))
        for entry in entries[9:16]
    )
