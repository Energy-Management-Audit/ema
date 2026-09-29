"""PIEE captions and references follow figure removal and addition."""

from pathlib import Path

from docx import Document

from ema.piee.figure_numbering import numbered_caption, renumber_figures


def test_renumbers_groups_subfigures_additions_and_references(tmp_path: Path) -> None:
    source, output = tmp_path / "in.docx", tmp_path / "out.docx"
    doc = Document()
    texts = [
        "În figura numărul 1 b) se prezintă producția.",
        "Fig. nr. 1 b) Producție în primul an disponibil",
        "Fig. nr. 1 d) Producție anuală",
        "Conform figurii numărul 1 d) producția crește.",
        "Fig. nr. 3 Consum de energie",
        "Fig. nr. 20 Un grafic adăugat",
        "Conform figurilor 4 ponderea este redusă.",
        "Fig. nr. 4 a) Pondere în primul an",
        "Fig. nr. 4 c) Pondere în al doilea an disponibil",
        "În figura numărul 13 se prezintă consumul specific.",
        "Fig. nr. 13 Consum specific",
        "Fig. nr. 14 Intensitate",
        "În figura numărul 13 se prezintă impactul.",
        "Fig. 13 Impact de mediu",
    ]
    for text in texts:
        paragraph = doc.add_paragraph()
        paragraph.add_run(text[:12]).bold = True
        paragraph.add_run(text[12:])
    doc.save(source)
    renumber_figures(source, output)
    result = [p.text for p in Document(output).paragraphs]
    assert result[0] == "În figura numărul 1 a) se prezintă producția."
    assert result[2] == "Fig. nr. 1 b) Producție anuală"
    assert result[3] == "Conform figurii numărul 1 b) producția crește."
    assert result[4] == "Fig. nr. 2 Consum de energie"
    assert result[5] == "Fig. nr. 3 Un grafic adăugat"
    assert result[6] == "Conform figurilor 4 ponderea este redusă."
    assert result[8] == "Fig. nr. 4 b) Pondere în al doilea an disponibil"
    assert result[9] == "În figura numărul 5 se prezintă consumul specific."
    assert result[12] == "În figura numărul 7 se prezintă impactul."
    assert result[13] == "Fig. 7 Impact de mediu"
    assert Document(output).paragraphs[0].runs[0].bold
    twice = tmp_path / "twice.docx"
    renumber_figures(output, twice)
    assert [p.text for p in Document(twice).paragraphs] == result


def test_removed_reference_is_red_missing_and_new_caption_is_numbered(tmp_path: Path) -> None:
    source, output = tmp_path / "in.docx", tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("Fig. nr. 2 Consum")
    doc.add_paragraph("Conform figurii numărul 1 a) datele sunt disponibile.")
    doc.save(source)
    assert numbered_caption(source, "Fig. Consum suplimentar") == "Fig. nr. 3 Consum suplimentar"
    renumber_figures(source, output)
    paragraph = Document(output).paragraphs[1]
    assert paragraph.text == "Conform figurii numărul n.d. datele sunt disponibile."
    assert any(run.font.color.rgb is not None for run in paragraph.runs)


def test_pv_reference_changes_when_an_earlier_group_is_removed(tmp_path: Path) -> None:
    source, output = tmp_path / "in.docx", tmp_path / "out.docx"
    doc = Document()
    for text in (
        "Fig. nr. 1 Producție",
        "Fig. nr. 4 a) Pondere",
        "Conform figurilor 4 ponderea este redusă.",
    ):
        doc.add_paragraph(text)
    doc.save(source)
    renumber_figures(source, output)
    assert Document(output).paragraphs[-1].text == "Conform figurilor 2 ponderea este redusă."
