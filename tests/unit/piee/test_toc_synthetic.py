"""Water TOC insertion keeps hyperlink and page field connected to its heading."""

import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ema.core.office.base_map import classify, stamp_base
from ema.core.office.package import read_parts, xml
from ema.piee.toc import build_toc_manifest, render_toc
from ema.piee.water import INDUSTRIAL_BOOKMARK


def test_missing_water_toc_inserts_linked_industrial_entry(tmp_path: Path) -> None:
    source = tmp_path / "toc.docx"
    document = Document()
    heading = document.add_paragraph("Industrial water")
    bookmark = OxmlElement("w:bookmarkStart")
    bookmark.set(qn("w:id"), "44")
    bookmark.set(qn("w:name"), INDUSTRIAL_BOOKMARK)
    heading._p.insert(0, bookmark)
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = "TOC 0"
    for index in range(1, 27):
        cell.add_paragraph(f"TOC {index}")
    water = cell.paragraphs[13]
    water.text = ""
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "_OldWater")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "2.2.6.Analiza consumului de apă potabilă 12"
    run.append(text)
    link.append(run)
    water._p.append(link)
    field = OxmlElement("w:r")
    instruction = OxmlElement("w:instrText")
    instruction.text = " PAGEREF _OldWater \\h "
    field.append(instruction)
    water._p.append(field)
    cell.paragraphs[14].text = "2.2.6.Echivalent 13"
    document.save(source)

    parts = read_parts(source)
    root = xml(parts, "word/document.xml")
    toc = list(root.find(qn("w:body")))[1]
    paragraphs = list(toc.iter(qn("w:p")))
    selectors = {
        f"paragraph:word/document.xml:{paragraph.getroottree().getelementpath(paragraph)}": slot
        for paragraph, slot in ((paragraphs[13], "toc_water"), (paragraphs[14], "toc_equivalent"))
    }
    mapping = classify(source, selectors)
    manifest = tmp_path / "toc.json"
    build_toc_manifest(parts, mapping, manifest)
    assert json.loads(manifest.read_text(encoding="utf-8"))["water_slot"] == "toc_water"
    stamped = tmp_path / "stamped.docx"
    stamp_base(source, mapping, stamped)
    output = tmp_path / "rendered.docx"
    render_toc(stamped, manifest, output, water_missing=True)
    updated = xml(read_parts(output), "word/document.xml")
    assert (
        sum(
            node.get(qn("w:anchor")) == INDUSTRIAL_BOOKMARK
            for node in updated.iter(qn("w:hyperlink"))
        )
        == 1
    )
    text = read_parts(output)["word/document.xml"].decode()
    assert "2.2.6.Analiza consumului de apă industrială" in text
    assert "2.2.7.Echivalent" in text
    assert "PAGEREF _TocEmaWaterIndustrial" in text
