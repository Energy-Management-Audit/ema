"""Identity rendering uses a private offset manifest and hidden bookmark."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from ema.core.office.anchors import AnchorLedger, stamp
from ema.piee.identity_map import render_identity


def test_replaces_identity_by_bookmark_and_preserves_emphasis(tmp_path: Path) -> None:
    stamped = tmp_path / "stamped.docx"
    manifest = tmp_path / "identity-spans.json"
    output = tmp_path / "draft.docx"
    document = Document()
    paragraph = document.add_paragraph("Client: ")
    paragraph.add_run("ACME").bold = True
    stamp(paragraph._p, "client", 1)
    document.save(stamped)
    manifest.write_text(
        json.dumps(
            [
                {
                    "owner": "word/document.xml",
                    "slot": "client",
                    "start": 8,
                    "end": 12,
                    "key": "client_name",
                }
            ]
        ),
        encoding="utf-8",
    )
    ledger = AnchorLedger(frozenset({"client"}))
    render_identity(stamped, manifest, {"client_name": "NEW"}, output, ledger)
    rendered = Document(output)
    assert rendered.paragraphs[0].text == "Client: NEW"
    assert rendered.paragraphs[0].runs[1].bold
    assert not ledger.untouched


def test_absent_identity_becomes_red_missing_marker(tmp_path: Path) -> None:
    stamped = tmp_path / "stamped.docx"
    manifest = tmp_path / "identity-spans.json"
    output = tmp_path / "draft.docx"
    document = Document()
    paragraph = document.add_paragraph("CUI: 123456")
    stamp(paragraph._p, "cui", 1)
    document.save(stamped)
    manifest.write_text(
        json.dumps(
            [{"owner": "word/document.xml", "slot": "cui", "start": 5, "end": 11, "key": "cui"}]
        ),
        encoding="utf-8",
    )
    render_identity(stamped, manifest, {"cui": None}, output, AnchorLedger(frozenset({"cui"})))
    rendered = Document(output)
    assert rendered.paragraphs[0].text == "CUI: n.d."
    assert rendered.paragraphs[0].runs[1].font.color.rgb is not None
