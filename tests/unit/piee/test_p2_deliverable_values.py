"""P2 preserves missing values, rounding and confirmation at the document boundary."""

import ast
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from tests.unit.energy_data.test_annex_index import annex
from tests.unit.piee.synthetic_piee import YEAR, piee_data
from tests.workspace_jobs import create_job

from ema.api import create_app
from ema.core.jobs import status
from ema.core.review import decide, fields, propose
from ema.core.review.models import Cell, Evidence, FieldSpec
from ema.core.workspace import Workspace
from ema.energy_data.anexa import parse_anexa
from ema.energy_data.carriers import Carrier
from ema.energy_data.model import CarrierSeries, Reading
from ema.piee.intake import _record_identity
from ema.piee.number import prototype_number
from ema.piee.payback_review import record_payback_check
from ema.piee.review_workflow import PieeWorkflow
from ema.piee.tables import _equivalent


@pytest.mark.parametrize("value,expected", [(0.125, "0,13"), (2.675, "2,68"), (-0.004, "0,00")])
def test_piee_uses_half_up_rounding(value: float, expected: str) -> None:
    assert prototype_number(value, 2) == expected
    assert prototype_number(1234.5, 2) == "1.234,50"
    assert prototype_number(2025, 0, grouping=False) == "2025"


def test_group_is_missing_when_one_present_carrier_is_missing() -> None:
    data = piee_data(prelucrare=False)
    data.dataset.carriers[Carrier.diesel] = {YEAR: CarrierSeries(annual=Reading(10, "t"))}
    data.dataset.carriers[Carrier.petrol] = {YEAR: CarrierSeries({1: Reading(None, "t")})}
    assert _equivalent(data, YEAR)[2] is None
    data.dataset.carriers[Carrier.petrol] = {YEAR - 1: CarrierSeries({1: Reading(None, "t")})}
    assert _equivalent(data, YEAR)[2] is not None
    del data.dataset.carriers[Carrier.diesel]
    assert _equivalent(data, YEAR)[2] is None


def test_electricity_group_requires_present_pv_and_grid() -> None:
    data = piee_data(prelucrare=False)
    assert _equivalent(data, YEAR)[0] is not None
    data.dataset.carriers[Carrier.electricity_pv] = {
        YEAR: CarrierSeries(annual=Reading(None, "MWh"))
    }
    assert _equivalent(data, YEAR)[0] is None


def test_pending_calculated_payback_blocks_until_confirmed(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = create_job(ws, "piee", "synthetic", YEAR + 1)
    data = piee_data()
    row = replace(
        data.anexa.planned_measures[0],
        values={
            key: item
            for key, item in data.anexa.planned_measures[0].values.items()
            if key != "payback_years"
        },
    )
    record_payback_check(ws, job, "planned", 1, row, "a" * 64)
    calculated = next(item for item in fields(ws, job) if item.key.endswith(".payback_years"))
    propose(
        ws,
        job,
        FieldSpec(key="measure.planned.1.description", label="Măsură", value_type="text"),
        row.description.value,
        [
            Evidence(
                id="description",
                provenance="document",
                file_sha="a" * 64,
                locator=Cell(sheet="Anexa", ref="B1"),
                method="anexa",
                retrieved_at=datetime.now(UTC),
                highlight="exact",
            )
        ],
        state="extracted",
    )
    issues = [
        item
        for item in PieeWorkflow().readiness(ws, job).blocking
        if item.code == "calculated_unconfirmed"
    ]
    assert len(issues) == 1
    assert issues[0].field_id == calculated.id
    assert issues[0].message == "Confirmaţi durata de recuperare calculată: Izolare"
    decide(ws, job, calculated.id, "accept", calculated.revision, "user")
    assert not any(
        item.code == "calculated_unconfirmed" for item in PieeWorkflow().readiness(ws, job).blocking
    )


def test_invalid_ownership_has_a_missing_review_field(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "ws")
    job = create_job(ws, "piee", "synthetic", YEAR + 1)
    data = piee_data()
    source = tmp_path / "annex.xlsx"
    annex(source)
    book = load_workbook(source)
    book["Date generale"].append(["Stat", "DA", "Privat", "100%"])
    book.save(source)
    data = replace(data, anexa=parse_anexa(source))
    _record_identity(ws, job, data, "a" * 64)
    field = next(item for item in fields(ws, job) if item.key == "identity.ownership_state")
    assert field.presence == "not_found" and field.value is None and field.required
    assert any(item.field_id == field.id for item in PieeWorkflow().readiness(ws, job).blocking)


@pytest.mark.parametrize(
    "year,code", [(2024, "piee_annex_year"), (2025, "piee_sources_incomplete")]
)
def test_source_validation_is_422_through_http(tmp_path: Path, year: int, code: str) -> None:
    ws = Workspace(tmp_path / "ws")
    job = create_job(ws, "piee", "synthetic", 2026)
    source = tmp_path / "annex.xlsx"
    annex(source, year=year)
    ws.set_slot(job, "anexa", ws.add_file("synthetic", source))
    with TestClient(
        create_app(ws, 8766, launch_code="code"), base_url="http://127.0.0.1:8766"
    ) as client:
        csrf = client.post("/session", json={"code": "code"}).json()["csrf"]
        revision = client.get(f"/jobs/{job}").json()["revision"]
        response = client.post(
            f"/jobs/{job}/piee/import", json={"on_revision": revision}, headers={"x-ema-csrf": csrf}
        )
    assert response.status_code == 422
    assert response.json()["type"] == f"urn:ema:error:{code}"
    assert not status(ws, job).runs


def test_piee_and_reporting_error_messages_use_handoff_diacritics() -> None:
    messages = []
    for folder in (Path("src/ema/piee"), Path("src/ema/reporting")):
        for path in folder.rglob("*.py"):
            tree = ast.parse(path.read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                if (
                    isinstance(node, ast.Call)
                    and isinstance(node.func, ast.Name)
                    and node.func.id == "EmaError"
                    and len(node.args) >= 2
                ):
                    messages.append(
                        "".join(
                            part.value
                            for part in ast.walk(node.args[1])
                            if isinstance(part, ast.Constant) and isinstance(part.value, str)
                        )
                    )
    assert messages
    assert not any(set(message) & set("șțȘȚ") for message in messages)
