"""Production presentation units are selected from a delivered programme."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from ema.core.office.package import encoded, read_parts, write_parts
from ema.core.office.pie_xml import pie_root
from ema.energy_data.model import CarrierSeries, EnergyDataset, Reading
from ema.piee.units import (
    convert,
    delivered_pie_representation,
    delivered_production_unit,
    delivered_separate_pv_figures,
    presentation_dataset,
)


@pytest.mark.parametrize(
    ("source", "target", "expected"),
    [
        ("mii tone", "tone", 1000.0),
        ("tone", "mii tone", 0.001),
        ("kWh", "MWh", 0.001),
        ("MWh", "mii MWh", 0.001),
        ("mii MWh gaz vehiculat", "kWh", 1_000_000.0),
    ],
)
def test_unit_conversion(source: str, target: str, expected: float) -> None:
    assert convert(1.0, source, target) == expected


def test_incompatible_units_are_rejected() -> None:
    with pytest.raises(ValueError, match="incompatible"):
        convert(1.0, "tone", "MWh")


def _programme(path: Path, *lines: str) -> None:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    document.save(path)


def test_previous_programme_sets_presentation_unit_and_records_derivation(tmp_path: Path) -> None:
    previous = tmp_path / "previous.docx"
    _programme(previous, "Figura 1. Producție tone/lună", "Total tone/an")
    dataset = EnergyDataset(
        (2025,),
        {},
        {"main": {2025: CarrierSeries({1: Reading(1.5, "mii tone")}, Reading(1.5, "mii tone"))}},
        {"main": "mii tone"},
    )

    presented, derivation = presentation_dataset(dataset, previous)

    assert delivered_production_unit(previous) == "tone"
    assert dataset.production["main"][2025].months[1].value == 1.5
    assert presented.production["main"][2025].months[1] == Reading(1500.0, "tone")
    assert presented.production["main"][2025].annual == Reading(1500.0, "tone")
    assert presented.production_unit == {"main": "tone"}
    assert derivation is not None
    assert (derivation.source_unit, derivation.presentation_unit, derivation.multiplier) == (
        "mii tone",
        "tone",
        1000.0,
    )


def test_ambiguous_delivered_unit_is_rejected(tmp_path: Path) -> None:
    previous = tmp_path / "previous.docx"
    _programme(previous, "Producție tone/an", "Producție mii tone/an")
    with pytest.raises(ValueError, match="unambiguous"):
        delivered_production_unit(previous)


@pytest.mark.parametrize("values,expected", [((0.2, 0.8), "normalized"), ((20, 80), "raw")])
def test_prior_native_pie_determines_value_scale_and_pv_visibility(
    tmp_path: Path, values: tuple[float, float], expected: str
) -> None:
    path = tmp_path / "previous.docx"
    Document().save(path)
    parts = read_parts(path)
    parts["word/charts/chart1.xml"] = encoded(pie_root("pv", ("Grid", "Solar"), values))
    write_parts(parts, path)
    assert delivered_pie_representation(path) == expected
    assert delivered_separate_pv_figures(path)


def test_prior_document_without_native_pie_has_no_pie_representation(tmp_path: Path) -> None:
    path = tmp_path / "previous.docx"
    Document().save(path)
    with pytest.raises(ValueError, match="no native pie"):
        delivered_pie_representation(path)
    assert not delivered_separate_pv_figures(path)
