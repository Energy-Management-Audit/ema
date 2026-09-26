"""Reviewed values as the generated document prints them: corrected and chosen values, n.d. for
rejected ones (the other build's rendered-result assertions, through our table renderer)."""

from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from pathlib import Path
from typing import Any

from docx import Document
from tests.unit.piee.synthetic_piee import YEAR, piee_data

from ema.core.office.anchor_targets import stamp_cell
from ema.core.office.anchors import AnchorLedger
from ema.core.review.models import Candidate, Cell, Field
from ema.piee.dataset import PieeData
from ema.piee.review_overlay import apply_review
from ema.piee.tables import MONTHLY_TABLES, render_tables

# Table 104 holds electricity from the grid, January to June; row 3 is the data year.
MARCH = "table_104_r3_c3"
APRIL = "table_104_r3_c4"
ELECTRICITY_TEP, GAS_TEP, TOTAL_TEP = "table_277_r3_c1", "table_277_r3_c2", "table_277_r3_c4"


def _tables() -> dict[int, int]:
    columns = {number: 7 for number, _, _ in MONTHLY_TABLES}
    return {**columns, 382: 2, 277: 5}


def _base(path: Path) -> frozenset[str]:
    document = Document()
    slots: set[str] = set()
    bookmark = 100
    for number, width in _tables().items():
        table = document.add_table(rows=3, cols=width)
        for row in range(1, 4):
            for col in range(width):
                slot = f"table_{number}_r{row}_c{col}"
                bookmark += 1
                stamp_cell(table.cell(row - 1, col)._tc, slot, bookmark)  # pyright: ignore[reportPrivateUsage]
                slots.add(slot)
        slots.add(f"table_{number}_rows")
    document.save(str(path))
    return frozenset(slots)


def _data() -> PieeData:
    data = piee_data()
    return replace(data, dataset=replace(data.dataset, years=(YEAR - 2, YEAR - 1, YEAR)))


def _rendered(data: PieeData, tmp_path: Path) -> dict[str, str]:
    tmp_path.mkdir(parents=True, exist_ok=True)
    base, output = tmp_path / "base.docx", tmp_path / "out.docx"
    expected = _base(base)
    render_tables(base, data, output, AnchorLedger(expected))
    document = Document(str(output))
    texts: dict[str, str] = {}
    names = iter(
        f"table_{number}_r{row}_c{col}"
        for number, width in _tables().items()
        for row in range(1, 4)
        for col in range(width)
    )
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts[next(names)] = cell.text
    return texts


def _field(key: str, value: Any, **changes: Any) -> Field:
    return Field.model_validate(
        {
            "id": key,
            "job_id": "job",
            "key": key,
            "label": key,
            "value_type": "number",
            "unit": "MWh" if key.startswith("carrier.") else "tep",
            "value": value,
            "state": "manual",
            "presence": "found",
            "review": "corrected",
            "evidence": ["manual"],
            **changes,
        }
    )


def test_the_document_prints_the_corrected_month_and_keeps_the_filed_annual(
    tmp_path: Path,
) -> None:
    before = _rendered(_data(), tmp_path / "before")
    month = _field(f"carrier.electricity_grid.{YEAR}.03", Decimal("140.5"))
    after = _rendered(apply_review(_data(), [month], {}), tmp_path / "after")
    assert (before[MARCH], after[MARCH]) == ("100,00", "140,50")
    assert after[APRIL] == before[APRIL] == "100,00"
    assert after[ELECTRICITY_TEP] == before[ELECTRICITY_TEP] == "103,20"
    assert after[TOTAL_TEP] == before[TOTAL_TEP] == "146,00"


def test_a_corrected_annual_recomputes_its_tep_and_the_total(tmp_path: Path) -> None:
    annual = _field(f"carrier.electricity_grid.{YEAR}", Decimal("1500"))
    after = _rendered(apply_review(_data(), [annual], {}), tmp_path)
    assert after[ELECTRICITY_TEP] == "129,00"
    assert after[TOTAL_TEP] not in ("146,00", "n.d.")


def test_the_document_prints_the_chosen_total(tmp_path: Path) -> None:
    chosen = _field(
        "annual.total_tep",
        Decimal("150.25"),
        review="accepted",
        state="extracted",
        alternatives=[
            Candidate(id="c-calc", value=Decimal("146"), evidence=["calc"]),
            Candidate(id="c-anexa", value=Decimal("150.25"), evidence=["cell"]),
        ],
        chosen="c-anexa",
        evidence=["cell"],
    )
    cells = {"cell": Cell(sheet="Date anuale", ref="Date anuale!F21")}
    after = _rendered(apply_review(_data(), [chosen], cells), tmp_path)
    assert after[TOTAL_TEP] == "150,25"


def test_rejected_values_print_as_nd(tmp_path: Path) -> None:
    rejected = [
        _field(f"carrier.electricity_grid.{YEAR}.03", Decimal("100"), review="rejected"),
        _field(f"carrier.natural_gas.{YEAR}", Decimal("500"), review="rejected"),
    ]
    before = _rendered(_data(), tmp_path / "before")
    after = _rendered(apply_review(_data(), rejected, {}), tmp_path / "after")
    assert before[MARCH] != "n.d." and after[MARCH] == "n.d."
    assert after[APRIL] == "100,00"
    assert before[GAS_TEP] != "n.d." and after[GAS_TEP] == "n.d."
