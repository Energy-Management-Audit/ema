"""Template measure rows are excluded without losing real measures."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from ema.energy_data.anexa import parse_anexa
from ema.energy_data.source import is_placeholder


def test_placeholder_description_matches_normalized_x_only() -> None:
    assert is_placeholder("x X x")
    assert is_placeholder("ＸＸＸ")
    assert not is_placeholder("x-ray")
    assert not is_placeholder("measure x")


@pytest.mark.parametrize(
    ("sheet_name", "attribute"),
    [
        ("Solutii EE existente", "existing_measures"),
        ("Solutii EE planificate", "planned_measures"),
    ],
)
def test_solution_placeholder_rows_are_skipped_before_number_conversion(
    tmp_path: Path, sheet_name: str, attribute: str
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = sheet_name
    for col, label in enumerate(
        (
            "Descrierea măsurii",
            "Termenul de aplicare",
            "Estimarea duratei de recuperare",
            "Costul aplicării măsurii",
            "Economia de energie",
            "Economia de energie",
            "Economia de cost",
        ),
        2,
    ):
        sheet.cell(3, col, label)
    if sheet_name == "Solutii EE existente":
        sheet["C3"] = "Data punerii în funcțiune"
    sheet["F4"] = "MWh/an"
    sheet["G4"] = "tep/an"
    sheet["B5"] = "xxx"
    sheet["D5"] = "#DIV/0!"
    sheet["E5"] = 0
    sheet["B6"] = "Înlocuirea pompei"
    sheet["C6"] = 2026
    sheet["D6"] = 2
    sheet["E6"] = 120
    sheet["F6"] = 30
    sheet["G6"] = 2.58
    sheet["H6"] = 60
    sheet["B7"] = "x x x"
    sheet["E7"] = 10
    path = tmp_path / "measures.xlsx"
    workbook.save(path)

    parsed = parse_anexa(path)

    rows = getattr(parsed, attribute)
    assert [row.description.value for row in rows] == ["Înlocuirea pompei", "x x x"]
    assert rows[0].commissioning_year is not None
    assert rows[0].commissioning_year.value == 2026
    assert rows[0].values["investment_thousand_lei"].value == 120
    assert rows[1].values["investment_thousand_lei"].value == 10
    assert [
        (issue.code, issue.detail, issue.ref.a1 if issue.ref else None)
        for issue in parsed.issues
        if issue.code == "measure_placeholder"
    ] == [("measure_placeholder", "x x x", f"{sheet_name}!B7")]
    assert not any(issue.ref is not None and issue.ref.row == 5 for issue in parsed.issues)


def test_audit_placeholder_rows_are_skipped_before_number_conversion(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Audit energetic"
    sheet["C3"] = "Măsura"
    sheet["E3"] = "Costuri (mii lei)"
    sheet["G3"] = "Economii (tep/an)"
    sheet["I3"] = "Economii de cost (mii lei/an)"
    for col in (5, 7, 9):
        sheet.cell(4, col, "estimate")
    sheet["C5"] = "xxx"
    sheet["E5"] = "#DIV/0!"
    sheet["G5"] = 0
    sheet["C6"] = "Izolarea rețelei"
    sheet["E6"] = 120
    sheet["G6"] = 14
    sheet["I6"] = 30
    sheet["C7"] = "X X X"
    sheet["E7"] = 10
    path = tmp_path / "audit.xlsx"
    workbook.save(path)

    parsed = parse_anexa(path)

    assert [row.description.value for row in parsed.audit_measures] == [
        "Izolarea rețelei",
        "X X X",
    ]
    assert parsed.audit_measures[0].values["saving_tep"].value == 14
    assert parsed.audit_measures[1].values["investment_thousand_lei"].value == 10
    assert [
        (issue.code, issue.detail, issue.ref.a1 if issue.ref else None)
        for issue in parsed.issues
        if issue.code == "measure_placeholder"
    ] == [("measure_placeholder", "X X X", "Audit energetic!C7")]
    assert not any(issue.ref is not None and issue.ref.row == 5 for issue in parsed.issues)
