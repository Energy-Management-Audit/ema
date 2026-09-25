"""Source-located Prelucrare factors are imported without guessed defaults."""

from pathlib import Path

from openpyxl import Workbook

from ema.core.office.sheets import open_book
from ema.energy_data.carriers import Carrier
from ema.energy_data.factors import FactorTable
from ema.energy_data.model import CarrierSeries, EnergyDataset, Reading
from ema.energy_data.prelucrare_factors import read_factors
from ema.energy_data.prelucrare_types import PrelucrareData


def test_imports_labeled_tep_and_co2_factors_with_source_cells(tmp_path: Path) -> None:
    path = tmp_path / "factors.xlsx"
    book = Workbook()
    electric = book.active
    electric.title = "Consum Electric"
    electric.append(("1 MWh", "=", 0.09, "tep"))
    primary = book.create_sheet("Principali factori de conversie")
    primary.append(("Purtător", "Factor", "Unitate"))
    primary.append(("electricitate", 0.09, "tep"))
    secondary = book.create_sheet("Factori de conversie in MWh")
    secondary.append(("electricitate", 11.1))
    impact = book.create_sheet("impact de mediu")
    impact.cell(2, 2, 2025)
    impact.cell(3, 2, "energie electrica")
    impact.cell(3, 6, 0.225)
    impact.cell(3, 9, 2.7)
    impact.cell(4, 2, "Indicator global prin suprapunerea efectelor")
    impact.cell(4, 9, 2.7)
    book.save(path)

    dataset = EnergyDataset(
        (2025,),
        {Carrier.electricity_grid: {2025: CarrierSeries(annual=Reading(12, "MWh"))}},
    )
    imported = PrelucrareData(dataset, FactorTable("empty", 2025, (), ()))
    factors = read_factors(open_book(path), (2025,), imported, path.name)
    assert factors.tep_factor(Carrier.electricity_grid, "MWh", 2025).per_unit == 0.09
    assert factors.co2_factor(Carrier.electricity_grid, "MWh", 2025).per_unit == 0.225
    assert imported.filed["co2.electricity_grid.2025"].value == 2.7
    assert imported.filed["co2.total.2025"].value == 2.7
    assert imported.located["factor.tep.electricity_grid"].ref.a1 == "Consum Electric!C1"
    assert factors.tep_factor(Carrier.natural_gas, "MWh", 2025) is None
