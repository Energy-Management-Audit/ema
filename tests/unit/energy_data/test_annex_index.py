"""Annex intake keeps each original source and matches only in-document IDs."""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl import Workbook

from ema.clients.registry import create_client
from ema.core.office.sniff import FileKind, Sniffed
from ema.core.workspace import Workspace
from ema.core.workspace.references import referenced_files
from ema.energy_data import annex_index
from ema.energy_data.annex_index import import_annexes, indexed


def annex(
    path: Path,
    *,
    name: str = "Exemplu Energie SA",
    cui: str | None = "RO 12345678",
    year: int | None = 2025,
    total: float = 12.5,
) -> None:
    book = Workbook()
    general = book.active
    assert general is not None
    general.title = "Date generale"
    general["A1"] = "Denumirea operatorului economic"
    general["B1"] = name
    general["A2"] = "Adresa poștală"
    general["B2"] = "Strada Exemplu 1"
    general["A3"] = "CUI"
    general["B3"] = cui
    general["A4"] = "Cod CAEN"
    general["B4"] = "1234"
    annual = book.create_sheet("Date anuale")
    annual["B4"] = "anului anterior"
    annual["D4"] = year
    annual["B7"] = "CONSUM DE ENERGIE TOTAL ANUAL"
    annual["C8"] = "tep/an"
    annual["D8"] = total
    measures = book.create_sheet("Solutii EE existente")
    measures["B3"] = "Descrierea măsurii aplicate"
    measures["C3"] = "Data punerii în funcţiune"
    measures["E3"] = "Costul investiției"
    measures["G4"] = "tep/an"
    measures["B5"] = "Măsură de test"
    measures["C5"] = 2025
    measures["E5"] = 10
    measures["G5"] = 1.5
    book.save(path)


def upload(ws: Workspace, path: Path, name: str | None = None):  # type: ignore[no-untyped-def]
    with path.open("rb") as stream:
        return import_annexes(ws, [(name or path.name, stream)])


def test_two_ids_keep_raw_evidence_and_reimport_matches(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    path = tmp_path / "annex.xlsx"
    annex(path, cui="CUI: 1234567 / 87654321")
    first = upload(ws, path, "folder/Anexa – à ^.xlsx")
    assert len(first.imported) == 1
    item = first.imported[0]
    assert item.created
    assert item.file_name == "Anexa – à ^.xlsx"
    assert len(indexed(ws)[item.client_id]) == 1
    with ws.connect() as db:
        client = db.execute("SELECT cui FROM clients WHERE id=?", (item.client_id,)).fetchone()
        assert client["cui"] == "1234567"
        assert any(sha == item.sha for _, sha, _ in referenced_files(db))
    data = indexed(ws)[item.client_id][0].data
    assert data["cui"] == "CUI: 1234567 / 87654321"
    assert data["file_name"] == "Anexa – à ^.xlsx"
    again = upload(ws, path)
    assert not again.imported[0].created
    assert again.imported[0].client_id == item.client_id
    assert len(indexed(ws)[item.client_id]) == 1


def test_cui_match_ambiguity_and_failed_file_isolation(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    first = create_client(ws, "Primul", "RO1234567")
    create_client(ws, "Al doilea", "87654321")
    path = tmp_path / "annex.xlsx"
    annex(path, cui="CUI: 1234567 / 98765432")
    result = upload(ws, path)
    assert result.imported[0].client_id == first["id"]
    annex(path, cui="CUI: 1234567 / 87654321")
    ambiguous = upload(ws, path)
    assert [item.code for item in ambiguous.ignored] == ["cui_ambiguous"]
    with path.open("rb") as stream:
        batch = import_annexes(ws, [("bad.txt", io.BytesIO(b"bad")), ("good.xlsx", stream)])
    assert [item.code for item in batch.ignored] == ["not_spreadsheet", "cui_ambiguous"]


def test_missing_year_cui_and_unreadable_are_independent(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "workspace")
    path = tmp_path / "annex.xlsx"
    annex(path, cui=None)
    missing_cui = upload(ws, path, "RO12345678.xlsx")
    assert [item.code for item in missing_cui.ignored] == ["cui_missing"]
    annex(path, year=None)
    assert [item.code for item in upload(ws, path).ignored] == ["year_missing"]
    assert [
        item.code for item in import_annexes(ws, [("bad.xlsx", io.BytesIO(b"bad"))]).ignored
    ] == ["unreadable"]


def test_large_file_is_item_error(tmp_path: Path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    ws = Workspace(tmp_path / "workspace")
    monkeypatch.setattr("ema.energy_data.annex_index.LIMIT", 3)
    result = import_annexes(ws, [("large.xlsx", io.BytesIO(b"1234"))])
    assert [item.code for item in result.ignored] == ["file_too_large"]


def test_stream_read_failure_does_not_stop_next_file(tmp_path: Path) -> None:
    class BadStream(io.BytesIO):
        def read(self, size: int = -1) -> bytes:
            raise OSError("synthetic read failure")

    ws = Workspace(tmp_path / "workspace")
    path = tmp_path / "annex.xlsx"
    annex(path)
    with path.open("rb") as good:
        result = import_annexes(ws, [("unreadable.xlsx", BadStream()), ("annex.xlsx", good)])
    assert [item.code for item in result.ignored] == ["unreadable"]
    assert len(result.imported) == 1


def test_misnamed_xls_and_unexpected_reader_error_do_not_abort_batch(
    tmp_path: Path, monkeypatch
) -> None:  # type: ignore[no-untyped-def]
    ws = Workspace(tmp_path / "workspace")
    first = tmp_path / "first.xlsx"
    last = tmp_path / "last.xlsx"
    annex(first, cui="1234567")
    annex(last, cui="2345678")
    with first.open("rb") as one, last.open("rb") as two:
        result = import_annexes(
            ws,
            [
                ("first.xlsx", one),
                ("misnamed.xls", io.BytesIO(first.read_bytes())),
                ("last.xlsx", two),
            ],
        )
    assert [item.file_name for item in result.imported] == ["first.xlsx", "last.xlsx"]
    assert [item.code for item in result.ignored] == ["unreadable"]
    assert "Extensia" in result.ignored[0].reason
    with ws.connect() as db:
        assert db.execute("SELECT COUNT(*) FROM clients").fetchone()[0] == 2

    reader = annex_index.parse_anexa
    calls = 0

    def intermittent(path: Path):  # type: ignore[no-untyped-def]
        nonlocal calls
        calls += 1
        if calls == 1:
            raise LookupError("synthetic reader failure")
        return reader(path)

    monkeypatch.setattr(annex_index, "parse_anexa", intermittent)
    with last.open("rb") as one, first.open("rb") as two:
        retried = import_annexes(ws, [("reader-fail.xlsx", one), ("good.xlsx", two)])
    assert len(retried.imported) == 1
    assert retried.ignored[0].code == "unreadable"
    assert "LookupError" in retried.ignored[0].reason


def test_ignored_annex_and_storage_failure_leave_no_client(tmp_path: Path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    ws = Workspace(tmp_path / "workspace")
    path = tmp_path / "annex.xlsx"
    annex(path)
    original_sniff = annex_index.sniff
    monkeypatch.setattr(
        annex_index,
        "sniff",
        lambda _path: Sniffed(FileKind.XLS, True, "synthetic mismatch"),
    )
    ignored = upload(ws, path)
    assert [item.code for item in ignored.ignored] == ["unreadable"]
    monkeypatch.setattr(annex_index, "sniff", original_sniff)

    def disk_failure(*_args: object) -> None:
        raise OSError("synthetic disk failure")

    monkeypatch.setattr(annex_index.os, "replace", disk_failure)
    failed = upload(ws, path)
    assert [item.code for item in failed.ignored] == ["unreadable"]
    with ws.connect() as db:
        for table in ("clients", "files", "client_uploads", "client_annexes"):
            assert db.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] == 0
