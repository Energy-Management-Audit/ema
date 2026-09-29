"""Blank formula inputs never fabricate an annual total; explicit zeros stay sourced."""

from pathlib import Path
from struct import pack, unpack_from
from zipfile import ZipFile

import olefile
import pytest
import xlwt
from lxml import etree
from openpyxl import Workbook, load_workbook
from tests.unit.energy_data.test_necesar_reader import MONTHS
from tests.unit.energy_data.test_necesar_reader import _book as necesar_book

from ema.energy_data.carriers import Carrier
from ema.energy_data.necesar import parse_necesar_info, to_dataset


def _cache_zero(path: Path) -> None:
    if path.suffix == ".xlsx":
        with ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        root = etree.fromstring(parts["xl/worksheets/sheet1.xml"])
        cell = root.find(f".//{ns}c[@r='N3']")
        cell.find(ns + "v").text = "0"
        parts["xl/worksheets/sheet1.xml"] = etree.tostring(root)
        with ZipFile(path, "w") as archive:
            for name, content in parts.items():
                archive.writestr(name, content)
    else:
        with olefile.OleFileIO(str(path), write_mode=True) as archive:
            data = bytearray(archive.openstream("Workbook").read())
            offset = 0
            while offset + 4 <= len(data):
                kind, length = unpack_from("<HH", data, offset)
                if kind == 6:
                    data[offset + 10 : offset + 18] = pack("<d", 0)
                offset += 4 + length
            archive.write_stream("Workbook", bytes(data))


def _book(path: Path, formula: bool, inputs: tuple[int, ...]) -> None:
    if path.suffix == ".xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "Cons energetice"

        def write(row, col, value):
            sheet.cell(row + 1, col + 1, value)

        total = "=SUM(B3:M3)" if formula else 0
    else:
        book = xlwt.Workbook()
        sheet = book.add_sheet("Cons energetice")
        write = sheet.write
        total = xlwt.Formula("SUM(B3:M3)") if formula else 0
    write(0, 0, "Consum GPL")
    for col, text in enumerate((2025, *MONTHS, "Total")):
        write(1, col, text)
    write(2, 0, "[tone]")
    for col, number in enumerate(inputs, 1):
        write(2, col, number)
    write(2, 13, total)
    book.save(str(path))
    if formula:
        _cache_zero(path)


@pytest.mark.parametrize("suffix", ["xls", "xlsx"])
@pytest.mark.parametrize(
    "formula,inputs,missing",
    [
        (True, (), True),
        (False, (), False),
        (True, (0,), False),
        (True, (1, -1), False),
    ],
)
def test_annual_totals_distinguish_blank_formulas_from_sourced_zero(
    tmp_path: Path, suffix: str, formula: bool, inputs: tuple[int, ...], missing: bool
) -> None:
    path = tmp_path / f"necesar.{suffix}"
    _book(path, formula, inputs)
    info = parse_necesar_info(path)
    total = info.carriers[Carrier.lpg].years[2025].total
    if missing:
        assert total is None
    else:
        assert total is not None and total.value == 0
        assert total.ref.a1 == "Cons energetice!N3"
    dataset = to_dataset(info)
    annual = dataset.carriers[Carrier.lpg][2025].annual
    assert annual is not None and annual.unit == "t"
    assert annual.value == (None if missing else 0)


def test_product_name_preserves_the_necesar_spelling(tmp_path: Path) -> None:
    path = necesar_book(tmp_path / "name.xlsx")
    book = load_workbook(path)
    book["Productia"]["B3"] = "kg lacuri și vopsele"
    book.save(path)
    dataset = to_dataset(parse_necesar_info(path))
    assert dataset.production_name == {"kg_lacuri_si_vopsele": "kg lacuri și vopsele"}


def test_empty_carrier_without_a_blank_sum_is_dropped(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    _book(path, False, ())
    book = load_workbook(path)
    book["Cons energetice"]["N3"] = None
    book.save(path)
    dataset = to_dataset(parse_necesar_info(path))
    assert Carrier.lpg not in dataset.carriers
