"""Generic tables must anchor to column labels, not the widest text row."""

from pathlib import Path

from openpyxl import Workbook

from ema.energy_data.necesar import parse_necesar_info


def test_table_title_with_more_text_does_not_replace_header(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Autovehicule"
    for col, value in enumerate(("Fleet", "for", "site", "A"), 1):
        sheet.cell(1, col, value)
    for col, value in enumerate(("Nr.", "Marca", "Consum"), 1):
        sheet.cell(3, col, value)
    sheet.cell(4, 1, 1)
    sheet.cell(4, 2, "Vehicle A")
    sheet.cell(4, 3, 7)
    path = tmp_path / "shifted-table.xlsx"
    book.save(path)

    info = parse_necesar_info(path)
    rows = info.tables["autovehicule"].rows
    assert len(rows) == 1
    assert rows[0].values["Marca"].value == "Vehicle A"
    assert rows[0].values["Consum"].value == 7
