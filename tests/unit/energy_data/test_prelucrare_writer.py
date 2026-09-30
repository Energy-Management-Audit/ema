"""A generated Prelucrare keeps source readings and live derivation formulas."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from ema.energy_data.carriers import Carrier
from ema.energy_data.factors import Factor, FactorTable
from ema.energy_data.model import CarrierSeries, EnergyDataset, Reading
from ema.energy_data.prelucrare import import_prelucrare
from ema.energy_data.prelucrare_writer import write_prelucrare


def _cell_for_labels(
    sheet,
    row_label: str,
    column_header: str,
    *,
    row_occurrence: int = 0,
    column_occurrence: int = 0,
    value_year: int | None = None,
    year: int | None = None,
):
    """Return a row's value under a matching header, optionally within a year block."""
    matches = [
        row_index
        for row_index in range(1, sheet.max_row + 1)
        if any(
            sheet.cell(row_index, col).value == row_label for col in range(1, sheet.max_column + 1)
        )
    ]
    if year is not None:
        matches = [
            row_index
            for row_index in matches
            if any(
                sheet.cell(r, col).value == year
                for r in range(1, row_index + 1)
                for col in range(1, sheet.max_column + 1)
            )
        ]
    row = matches[row_occurrence]
    columns = []
    for col_index in range(1, sheet.max_column + 1):
        header_rows = [
            row_index
            for row_index in range(1, sheet.max_row + 1)
            if sheet.cell(row_index, col_index).value == column_header
        ]
        if header_rows and header_rows[0] <= row:
            columns.append(col_index)
    if value_year is not None:
        columns = [
            col_index
            for col_index in range(1, sheet.max_column + 1)
            if any(
                sheet.cell(row_index, col_index).value == value_year
                for row_index in range(1, sheet.max_row + 1)
            )
        ]
    column = columns[column_occurrence]
    return sheet.cell(row, column)


def _dataset() -> EnergyDataset:
    months = {month: Reading(float(month), "MWh") for month in range(1, 13)}
    return EnergyDataset(
        (2024, 2025),
        {
            Carrier.electricity_grid: {
                2024: CarrierSeries(months),
                2025: CarrierSeries({1: Reading(8, "MWh")}, Reading(80, "MWh")),
            },
            Carrier.natural_gas: {2025: CarrierSeries(annual=Reading(20, "MWh"))},
            Carrier.water_potable: {2025: CarrierSeries(annual=Reading(4, "m3"))},
        },
        {"output": {2025: CarrierSeries({1: Reading(2, "mii tone")}, Reading(3, "mii tone"))}},
        {"output": "mii tone"},
        {2025: Reading(100_000, "lei")},
        {2025: Reading(12_000, "lei")},
    )


def test_writer_preserves_sources_and_links_derived_values(tmp_path: Path) -> None:
    path = tmp_path / "prelucrare.xlsx"
    write_prelucrare(_dataset(), (2024, 2025), path)
    book = load_workbook(path)

    physical = book["Consum Electric"]
    assert _cell_for_labels(physical, "[MWh]", "Ianuarie").value == 1
    assert _cell_for_labels(physical, "[MWh]", "Decembrie").value == 12
    assert _cell_for_labels(physical, "[MWh]", "TOTAL").value == "=SUM(D3:O3)"
    assert _cell_for_labels(physical, "[MWh]", "TOTAL", row_occurrence=1).value == 80
    production = book["Productii"]
    assert _cell_for_labels(production, "tone/luna", "Ianuarie").value == 2000
    assert _cell_for_labels(production, "tone/luna", "TOTAL").value == 3000
    assert (
        _cell_for_labels(
            book["Principali factori de conversie"], "electricity_grid", "Purtător"
        ).value
        == "electricity_grid"
    )

    tep = book["TEP"]
    tep_year = _cell_for_labels(tep, "electricity_grid [tep]", "Ianuarie")
    assert "'Consum Electric'!D3" in tep_year.value
    assert "'Principali factori de conversie'!D2" in tep_year.value
    tep_total = _cell_for_labels(tep, "TOTAL [tep]", "TOTAL")
    assert tep_total.value.startswith("=SUM(")
    economics = book["Chelt-Cifra afaceri"]
    turnover = _cell_for_labels(economics, "Cifra de afaceri [lei]", "Anul", value_year=2025)
    assert turnover.value == 100_000
    tep_total_address = _cell_for_labels(tep, "TOTAL [tep]", "TOTAL", year=2025).coordinate
    assert (
        f"'TEP'!{tep_total_address}"
        in _cell_for_labels(
            economics,
            "Intensitate energetica [tep/1000 lei]",
            "Anul",
            value_year=2025,
        ).value
    )
    specific = book["Consumuri specifice"]
    specific_cell = _cell_for_labels(
        specific, "Anul", "Consum specific anual de energie [tep]"
    ).offset(row=2)
    assert (
        f"'Productii'!{_cell_for_labels(production, 'tone/luna', 'TOTAL').coordinate}/1000"
        in specific_cell.value
    )
    impact = book["impact de mediu"]
    impact_cell = _cell_for_labels(
        impact, "electricity_grid", "electricity_grid", year=2025
    ).offset(column=7)
    assert impact_cell.value.startswith("='Consum Electric'!")
    assert book.calculation.fullCalcOnLoad is True


@pytest.mark.parametrize("years", [(), (2025, 2024), (2025, 2025), (2023,)])
def test_writer_rejects_invalid_year_selection(tmp_path: Path, years: tuple[int, ...]) -> None:
    with pytest.raises(ValueError, match="years must be present"):
        write_prelucrare(_dataset(), years, tmp_path / "invalid.xlsx")
    assert not (tmp_path / "invalid.xlsx").exists()


def test_explicit_water_with_missing_readings_survives_round_trip(tmp_path: Path) -> None:
    months = {month: Reading(None, "m3") for month in range(1, 13)}
    dataset = EnergyDataset((2025,), {Carrier.water_potable: {2025: CarrierSeries(months)}})
    path = tmp_path / "water.xlsx"
    write_prelucrare(dataset, (2025,), path)
    rebuilt = import_prelucrare(path)
    assert rebuilt.dataset.carriers[Carrier.water_potable][2025].months == months
    assert rebuilt.dataset.carriers[Carrier.water_potable][2025].annual is None


@pytest.mark.parametrize("carrier,unit", [(Carrier.diesel, "t"), (Carrier.natural_gas, "MWh")])
def test_explicit_energy_carrier_with_missing_readings_survives_round_trip(
    tmp_path: Path, carrier: Carrier, unit: str
) -> None:
    months = {month: Reading(None, unit) for month in range(1, 13)}
    dataset = EnergyDataset((2025,), {carrier: {2025: CarrierSeries(months)}})
    path = tmp_path / "missing.xlsx"
    write_prelucrare(dataset, (2025,), path)
    rebuilt = import_prelucrare(path)
    assert rebuilt.dataset.carriers[carrier][2025].months == months
    assert rebuilt.dataset.carriers[carrier][2025].annual is None


def test_writer_names_cogeneration_and_coke_and_omits_internal_generation_from_total(
    tmp_path: Path,
) -> None:
    dataset = EnergyDataset(
        (2025,),
        {
            Carrier.electricity_cogen: {2025: CarrierSeries(annual=Reading(20, "MWh"))},
            Carrier.coke: {2025: CarrierSeries(annual=Reading(3, "t"))},
        },
        {"main": {2025: CarrierSeries(annual=Reading(100, "tone"))}},
        {"main": "tone"},
    )
    factors = FactorTable(
        "synthetic",
        2025,
        (
            Factor(Carrier.electricity_cogen, "MWh", 0.086, "fixture"),
            Factor(Carrier.coke, "t", 0.762, "fixture"),
        ),
        (),
        2025,
    )
    path = tmp_path / "generated.xlsx"
    write_prelucrare(dataset, (2025,), path, factors)
    book = load_workbook(path)
    assert "energi electrica din cogenerar" in book.sheetnames
    assert "Consum Cocs" in book.sheetnames
    assert _cell_for_labels(book["energi electrica din cogenerar"], "[MWh]", "TOTAL").value == 20
    assert _cell_for_labels(book["Consum Cocs"], "[t]", "TOTAL").value == 3
    assert _cell_for_labels(book["Productii"], "tone", "TOTAL").value == 100
    labels = {book["TEP"].cell(row, 3).value for row in range(1, book["TEP"].max_row + 1)}
    assert "coke [tep]" in labels
    assert "electricity_cogen [tep]" not in labels
