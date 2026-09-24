"""Production year groups can contain multiple product data rows."""

from pathlib import Path

from openpyxl import Workbook

from ema.energy_data.necesar import parse_necesar_info, to_dataset

MONTHS = (
    "Ianuarie",
    "Februarie",
    "Martie",
    "Aprilie",
    "Mai",
    "Iunie",
    "Iulie",
    "August",
    "Septembrie",
    "Octombrie",
    "Noiembrie",
    "Decembrie",
)


def test_two_products_under_one_year_header_are_retained(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Productia"
    for col, value in enumerate((2025, *MONTHS, "Total"), 2):
        sheet.cell(2, col, value)
    sheet.cell(3, 1, "Product A")
    sheet.cell(3, 2, "kg")
    sheet.cell(3, 3, 2)
    sheet.cell(3, 15, 2)
    sheet.cell(4, 1, "Product B")
    sheet.cell(4, 2, "kg")
    sheet.cell(4, 3, 3)
    sheet.cell(4, 15, 3)
    path = tmp_path / "multiple-products.xlsx"
    book.save(path)

    info = parse_necesar_info(path)
    dataset = to_dataset(info)
    assert {item.name.value for item in info.production} == {"Product A", "Product B"}
    assert set(dataset.production) == {"product_a", "product_b"}
