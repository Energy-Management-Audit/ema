"""Prelucrare source sheets are located by labels and keep filed values separate."""

from pathlib import Path

from openpyxl import Workbook

from ema.energy_data.carriers import Carrier
from ema.energy_data.prelucrare import import_prelucrare

MONTHS = (
    "Ianuarie",
    "Februarie",
    "Martie",
    "Aprilie",
    "Mai",
    "Iunie",
    "Iulie",
    "August",
    "Septembrie",
    "Octombrie",
    "Noiembrie",
    "Decembrie",
)


def _source(path: Path) -> Path:
    book = Workbook()
    electric = book.active
    electric.title = "Consum Electric"
    electric.cell(2, 1, 2025)
    for column, month in enumerate(MONTHS, 4):
        electric.cell(2, column, month)
    electric.cell(3, 3, "[MWh]")
    electric.cell(3, 4, 10)
    electric.cell(3, 16, 10)
    electric.cell(5, 1, "1 MWh")
    electric.cell(5, 2, "=")
    electric.cell(5, 3, 0.086)
    electric.cell(5, 4, "tep")
    production = book.create_sheet("Productii")
    production.cell(2, 1, 2025)
    for column, month in enumerate(MONTHS, 4):
        production.cell(2, column, month)
    production.cell(3, 3, "tone/luna")
    production.cell(3, 4, 1000)
    production.cell(3, 16, 1000)
    economics = book.create_sheet("Chelt-Cifra afaceri")
    economics.cell(2, 3, "Anul")
    economics.cell(2, 4, 2025)
    economics.cell(3, 3, "Cifra de afaceri [lei]")
    economics.cell(3, 4, 100_000)
    economics.cell(4, 3, "CHELTUIELI ENERGETICE TOTALE [lei]")
    economics.cell(4, 4, 12_000)
    book.create_sheet("Principali factori de conversie")
    book.create_sheet("Factori de conversie in MWh")
    tep = book.create_sheet("TEP")
    tep.cell(2, 3, 2025)
    tep.cell(3, 3, "Ianuarie")
    tep.cell(4, 3, "TOTAL [tep]")
    tep.cell(4, 16, 0.86)
    impact = book.create_sheet("impact de mediu")
    impact.cell(2, 2, 2025)
    impact.cell(3, 2, "energie electrica")
    impact.cell(3, 6, 0.226)
    impact.cell(3, 9, 2.26)
    impact.cell(4, 2, "Indicator global prin suprapunerea efectelor")
    impact.cell(4, 9, 2.26)
    book.save(path)
    return path


def test_import_preserves_physical_readings_factors_and_economics(tmp_path: Path) -> None:
    imported = import_prelucrare(_source(tmp_path / "source.xlsx"))
    assert imported.dataset.years == (2025,)
    electric = imported.dataset.carriers[Carrier.electricity_grid][2025]
    assert electric.months[1].value == 10
    assert electric.annual.value == 10
    assert imported.located["carrier.electricity_grid.2025.01"].ref.a1 == "Consum Electric!D3"
    assert imported.dataset.production["main"][2025].months[1].value == 1
    assert imported.dataset.turnover_lei[2025].value == 100_000
    assert imported.dataset.energy_costs_lei[2025].value == 12_000
    assert imported.factors.tep_factor(Carrier.electricity_grid, "MWh", 2025).per_unit == 0.086
    assert imported.factors.co2_factor(Carrier.electricity_grid, "MWh", 2025).per_unit == 0.226
    assert imported.filed["co2.total.2025"].value == 2.26
