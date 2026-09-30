"""Shifted labels, units, factors, and ambiguous Prelucrare blocks."""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from ema.energy_data.calc import tep, tep_total
from ema.energy_data.carriers import Carrier
from ema.energy_data.prelucrare import import_prelucrare

MONTHS = (
    "Ianuarie",
    "Februarie",
    "Martie",
    "Aprilie",
    "Mai",
    "Iunie",
    "Iulie",
    "August",
    "Septembrie",
    "Octombrie",
    "Noiembrie",
    "Decembrie",
)


def _header(sheet, row: int, *, year: int | None = 2025, start_col: int = 4) -> None:
    if year is not None:
        sheet.cell(row, start_col - 1, year)
    for index, month in enumerate(MONTHS, start_col):
        sheet.cell(row, index, month)
    sheet.cell(row, start_col + 13, "TOTAL")


def _workbook(path: Path, *, coke_unit: str = "[MWh]") -> Path:
    book = Workbook()
    cogeneration = book.active
    cogeneration.title = "energi electrica din cogenerar "
    cogeneration.cell(3, 4, "1 MWh")
    cogeneration.cell(3, 5, "=")
    cogeneration.cell(3, 6, 0.086)
    cogeneration.cell(3, 7, "tep")
    _header(cogeneration, 9)
    cogeneration.cell(10, 3, "[MWh]")
    for col in range(4, 16):
        cogeneration.cell(10, col, 2)
    cogeneration.cell(10, 17, 24)

    coke = book.create_sheet("Consum Cocs")
    for row, label, factor in ((3, "1 tona", 0.762), (4, "1MWh", 0.086)):
        coke.cell(row, 4, label)
        coke.cell(row, 5, "=")
        coke.cell(row, 6, factor)
        coke.cell(row, 7, "tep")
    _header(coke, 12)
    coke.cell(13, 3, coke_unit)
    for col in range(4, 16):
        coke.cell(13, col, 3)
    coke.cell(13, 17, 36)

    fuel = book.create_sheet("Consum Carburanti")
    fuel.cell(2, 4, "1 t (motorina)")
    fuel.cell(2, 5, "=")
    fuel.cell(2, 6, 1.015)
    fuel.cell(2, 7, "tep")
    fuel.cell(5, 3, "tone")
    _header(fuel, 11)
    fuel.cell(12, 3, "Motorina")
    for col in range(4, 16):
        fuel.cell(12, col, 1)
    fuel.cell(12, 17, 12)

    production = book.create_sheet("Productii")
    _header(production, 14)
    production.cell(15, 3, "tone")
    for col in range(4, 16):
        production.cell(15, col, 10)
    production.cell(15, 17, 120)
    production.cell(27, 4, "total anual, tone")
    production.cell(28, 3, 2025)
    production.cell(28, 4, 120)

    water = book.create_sheet("Consum apa ")
    for title_row, label in (
        (4, "Consum apa potabilla"),
        (30, "industriala - m3"),
        (60, "Consum apa meteorica, etc"),
    ):
        water.cell(title_row, 2, label)
        header_row = title_row + 1
        for col, month in enumerate(MONTHS, 3):
            water.cell(header_row, col, month)
        water.cell(header_row, 15, "Total")
        water.cell(header_row + 1, 2, 2025)
        for col in range(3, 15):
            water.cell(header_row + 1, col, 5)
        water.cell(header_row + 1, 15, 60)
    book.save(path)
    return path


@pytest.mark.parametrize("coke_unit,factor", [("[MWh]", 0.086), ("tone", 0.762)])
def test_shifted_case_c_labels_and_applicable_coke_factor(
    tmp_path: Path, coke_unit: str, factor: float
) -> None:
    imported = import_prelucrare(_workbook(tmp_path / "source.xlsx", coke_unit=coke_unit))
    assert imported.dataset.production_unit == {"main": "tone"}
    assert imported.dataset.production["main"][2025].annual.value == 120
    assert imported.located["production.main.2025"].ref.a1 == "Productii!D28"
    assert imported.located["carrier.electricity_cogen.2025.01"].ref.a1.endswith("!D10")
    assert imported.located["carrier.coke.2025"].ref.a1 == "Consum Cocs!Q13"
    assert imported.dataset.carriers[Carrier.coke][2025].annual.unit == (
        "MWh" if coke_unit == "[MWh]" else "t"
    )
    assert tep(imported.dataset, imported.factors, Carrier.coke, 2025).value == pytest.approx(
        36 * factor
    )
    assert imported.dataset.carriers[Carrier.diesel][2025].annual.value == 12
    assert set(imported.dataset.carriers) >= {
        Carrier.water_potable,
        Carrier.water_industrial,
        Carrier.water_storm,
    }
    for carrier in (Carrier.water_potable, Carrier.water_industrial, Carrier.water_storm):
        assert imported.dataset.carriers[carrier][2025].annual.value == 60
    assert tep_total(imported.dataset, imported.factors, 2025).value == pytest.approx(
        36 * factor + 12 * 1.015
    )


def test_duplicate_water_and_diesel_labels_are_ambiguous(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "ambiguous.xlsx")
    book = load_workbook(path)
    water = book["Consum apa "]
    water.cell(90, 2, "industriala - m3")
    for col, month in enumerate(MONTHS, 3):
        water.cell(91, col, month)
    water.cell(92, 2, 2025)
    fuel = book["Consum Carburanti"]
    fuel.cell(13, 3, "Motorina")
    book.save(path)
    imported = import_prelucrare(path)
    assert Carrier.water_industrial not in imported.dataset.carriers
    assert Carrier.diesel not in imported.dataset.carriers
    assert {issue.detail for issue in imported.issues if issue.code == "label_ambiguous"} >= {
        "carrier.water_industrial",
        "diesel.2025",
    }


def test_ambiguous_factor_and_unknown_coke_unit_stay_missing(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "missing.xlsx", coke_unit="[unknown]")
    book = load_workbook(path)
    cogen = book["energi electrica din cogenerar "]
    cogen.cell(4, 4, "1 MWh")
    cogen.cell(4, 5, "=")
    cogen.cell(4, 6, 0.09)
    cogen.cell(4, 7, "tep")
    book.save(path)
    imported = import_prelucrare(path)
    assert Carrier.coke not in imported.dataset.carriers
    assert any(issue.code == "unit_missing" and issue.detail == "coke" for issue in imported.issues)
    assert imported.factors.tep_factor(Carrier.electricity_cogen, "MWh", 2025) is None
    assert any(issue.code == "factor_ambiguous" for issue in imported.issues)
    assert tep(imported.dataset, imported.factors, Carrier.electricity_cogen, 2025).value is None
