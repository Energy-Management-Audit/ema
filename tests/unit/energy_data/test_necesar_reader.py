"""Synthetic Necesar info layouts and missing-data behavior."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from ema.energy_data.carriers import Carrier
from ema.energy_data.necesar import parse_necesar_info, to_dataset

MONTHS = (
    "Ianuarie",
    "Februarie",
    "Martie",
    "Aprilie",
    "Mai",
    "Iunie",
    "Iulie",
    "August",
    "Septembrie",
    "Octombrie",
    "Noiembrie",
    "Decembrie",
)


def _year(sheet: object, row: int, year: int, values: dict[int, object], total: object) -> None:
    for col, value in enumerate((year, *MONTHS, "Total"), 1):
        sheet.cell(row, col, value)
    sheet.cell(row + 1, 1, "[MWh]")
    for month, value in values.items():
        sheet.cell(row + 1, month + 1, value)
    sheet.cell(row + 1, 14, total)
    sheet.cell(row + 2, 1, "[tep]")


def _book(path: Path, *, three_years: bool = False, unknown: bool = False) -> Path:
    book = Workbook()
    energy = book.active
    energy.title = "Cons energetice"
    energy.cell(5, 1, "Consum bezina")
    energy.cell(6, 2, "extra note")
    _year(energy, 7, 2023, {1: 1, 2: 2}, "3,00")
    if three_years:
        _year(energy, 11, 2024, {1: 4}, 4)
        _year(energy, 15, 2025, {1: 5}, 5)
    else:
        energy.cell(11, 1, "Consum energie electrica din SEN")
        _year(energy, 12, 2023, {month: 1 for month in range(1, 13)}, 13)
    if unknown:
        energy.cell(20, 1, "Consum unobtainium")
        _year(energy, 21, 2023, {1: 9}, 9)
    production = book.create_sheet("Productia")
    production.cell(2, 2, 2023)
    for col, month in enumerate(MONTHS, 3):
        production.cell(2, col, month)
    production.cell(2, 15, "Total")
    production.cell(3, 2, "kg product")
    production.cell(3, 3, 2)
    production.cell(3, 15, 2)
    economic = book.create_sheet("Cifre economice")
    economic.cell(3, 1, "Anul")
    economic.cell(3, 2, 2023)
    economic.cell(4, 1, "Cifra de afaceri neta [lei]")
    economic.cell(4, 2, 100)
    book.save(path)
    return path


def test_shifted_label_notes_sparse_comma_and_total_mismatch(tmp_path: Path) -> None:
    info = parse_necesar_info(_book(tmp_path / "input.xlsx"))
    petrol = info.carriers[Carrier.petrol].years[2023]
    assert petrol.months[0] is not None and petrol.months[0].ref.a1 == "Cons energetice!B8"
    assert petrol.months[2] is None
    assert petrol.total is not None and petrol.total.value == 3
    assert petrol.total.ref.a1 == "Cons energetice!N8"
    assert info.carriers[Carrier.electricity_grid].years[2023].total is not None
    assert [(i.code, i.ref.a1 if i.ref else None) for i in info.issues] == [
        ("total_mismatch", "Cons energetice!N13"),
        ("label_from_unit_cell", "Productia!B3"),
    ]
    dataset = to_dataset(info)
    assert set(dataset.carriers[Carrier.petrol][2023].months) == {1, 2}
    assert dataset.production["kg_product"][2023].annual is not None
    assert dataset.turnover_lei[2023].value == 100


def test_three_years_and_unknown_carrier_are_isolated(tmp_path: Path) -> None:
    info = parse_necesar_info(_book(tmp_path / "three.xlsx", three_years=True, unknown=True))
    assert set(info.carriers[Carrier.petrol].years) == {2023, 2024, 2025}
    assert any(i.code == "carrier_unknown" and i.ref is not None for i in info.issues)
    assert to_dataset(info).years == (2023, 2024, 2025)


def test_missing_sheet_is_issue_not_exception(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    book = Workbook()
    book.save(path)
    info = parse_necesar_info(path)
    assert [issue.code for issue in info.issues] == ["sheet_missing"] * 3
