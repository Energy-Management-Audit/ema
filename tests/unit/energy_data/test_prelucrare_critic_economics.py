"""Synthetic regressions for Prelucrare economic source fields."""

from pathlib import Path

from openpyxl import Workbook

from ema.energy_data.carriers import Carrier
from ema.energy_data.model import CarrierSeries, EnergyDataset, Reading
from ema.energy_data.prelucrare import import_prelucrare
from ema.energy_data.prelucrare_writer import write_prelucrare


def _book(
    path: Path,
    *,
    label_column: int = 3,
    turnover: int | None = 100,
    revenue: int | None = 200,
) -> Path:
    book = Workbook()
    electric = book.active
    electric.title = "Consum Electric"
    electric.cell(2, 1, 2025)
    for column, month in enumerate(
        (
            "Ianuarie",
            "Februarie",
            "Martie",
            "Aprilie",
            "Mai",
            "Iunie",
            "Iulie",
            "August",
            "Septembrie",
            "Octombrie",
            "Noiembrie",
            "Decembrie",
        ),
        4,
    ):
        electric.cell(2, column, month)
    electric.cell(3, 3, "[MWh]")
    electric.cell(3, 4, 1)
    economic = book.create_sheet("Chelt-Cifra afaceri")
    economic.cell(2, label_column, "Anul")
    economic.cell(2, label_column + 1, 2025)
    for row, label, value in (
        (3, "Cifra de afaceri [lei]", turnover),
        (4, "Valoarea veniturilor din exploatare [lei]", revenue),
        (5, "CHELTUIELI ENERGETICE TOTALE [lei]", 30),
        (6, "Cheltuieli de productie [lei]", 300),
    ):
        economic.cell(row, label_column, label)
        economic.cell(row, label_column + 1, value)
    book.save(path)
    return path


def test_turnover_precedes_operating_revenue(tmp_path: Path) -> None:
    imported = import_prelucrare(_book(tmp_path / "economic.xlsx"))
    assert imported.dataset.turnover_lei[2025].value == 100
    assert imported.located["turnover.2025"].label == "Cifra de afaceri [lei]"


def test_empty_turnover_falls_back_and_records_source_label(tmp_path: Path) -> None:
    imported = import_prelucrare(_book(tmp_path / "economic.xlsx", turnover=None))
    assert imported.dataset.turnover_lei[2025].value == 200
    assert imported.located["turnover.2025"].label == "Valoarea veniturilor din exploatare [lei]"


def test_production_costs_third_fallback_preserves_source_label(tmp_path: Path) -> None:
    imported = import_prelucrare(_book(tmp_path / "economic.xlsx", turnover=None, revenue=None))
    assert imported.dataset.turnover_lei[2025].value == 300
    assert imported.located["turnover.2025"].label == "Cheltuieli de productie [lei]"


def test_energy_costs_imported_from_labeled_row(tmp_path: Path) -> None:
    imported = import_prelucrare(_book(tmp_path / "economic.xlsx"))
    assert imported.dataset.energy_costs_lei[2025].value == 30


def test_economic_labels_can_move_column(tmp_path: Path) -> None:
    imported = import_prelucrare(_book(tmp_path / "economic.xlsx", label_column=4))
    assert imported.dataset.turnover_lei[2025].value == 100


def test_writer_round_trip_preserves_energy_costs(tmp_path: Path) -> None:
    dataset = EnergyDataset(
        (2025,),
        {Carrier.electricity_grid: {2025: CarrierSeries({1: Reading(1, "MWh")})}},
        turnover_lei={2025: Reading(100, "lei")},
        energy_costs_lei={2025: Reading(30, "lei")},
    )
    path = tmp_path / "economic.xlsx"
    write_prelucrare(dataset, (2025,), path)

    imported = import_prelucrare(path)
    assert imported.dataset.energy_costs_lei == dataset.energy_costs_lei
    assert imported.located["energy_costs.2025"].label == "Cheltuieli cu energia [lei]"
