"""Audit measures remain located when the annex table moves."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from ema.energy_data.anexa import parse_anexa


def test_audit_measure_headers_can_shift_without_changing_reader(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Audit energetic"
    header = 12
    sheet.cell(header, 3, "Măsura")
    sheet.cell(header, 5, "Costuri (mii lei)")
    sheet.cell(header, 7, "Economii (tep/an)")
    sheet.cell(header, 9, "Economii de cost (mii lei/an)")
    for col in (5, 7, 9):
        sheet.cell(header + 1, col, "estimate")
    sheet.cell(header + 2, 3, "Izolarea rețelei")
    sheet.cell(header + 2, 5, 120.0)
    sheet.cell(header + 2, 7, 14.0)
    sheet.cell(header + 2, 9, 30.0)
    path = tmp_path / "anexa.xlsx"
    workbook.save(path)

    result = parse_anexa(path)

    assert len(result.audit_measures) == 1
    measure = result.audit_measures[0]
    assert measure.description.ref.a1 == "Audit energetic!C14"
    assert measure.values["investment_thousand_lei"].value == 120.0
    assert measure.values["saving_tep"].ref.a1 == "Audit energetic!G14"
    assert measure.values["saving_thousand_lei"].value == 30.0
