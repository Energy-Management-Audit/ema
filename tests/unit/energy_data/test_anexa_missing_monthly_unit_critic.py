"""Monthly quantities without a source unit require an explicit issue."""

from pathlib import Path

from openpyxl import Workbook

from ema.energy_data.anexa import parse_anexa


def test_monthly_values_without_unit_are_reported(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Date lunare"
    sheet["A1"] = "ENERGIE ELECTRICĂ – Consumul total anual"
    sheet["A2"] = "Luna"
    for col, month in enumerate(
        ("Ian", "Feb", "Mar", "Apr", "Mai", "Iun", "Iul", "Aug", "Sep", "Oct", "Noi", "Dec"),
        start=2,
    ):
        sheet.cell(2, col, month)
        sheet.cell(3, col, col)
    path = tmp_path / "missing-unit.xlsx"
    workbook.save(path)

    parsed = parse_anexa(path)

    assert any(issue.code == "unit_missing" and issue.ref is not None for issue in parsed.issues)
    assert all(
        value.unit is not None for value in parsed.monthly.get("electricity_grid", {}).values()
    )
