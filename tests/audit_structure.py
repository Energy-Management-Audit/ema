"""Faithful complex TOC fields and printed numbering for synthetic audit bases."""

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ema.audit.catalogue import CATALOGUE
from ema.audit.sections import Status, get_status, mark_drafted, set_status


def _node(tag, **attributes):
    node = OxmlElement(tag)
    for key, value in attributes.items():
        node.set(qn("w:" + key), value)
    return node


def add_audit_toc(document) -> None:
    for level in (1, 2, 3):
        if f"TOC {level}" not in document.styles:
            document.styles.add_style(f"TOC {level}", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("CUPRINS")
    first = document.add_paragraph("Chapter prototype", style="TOC 1")
    for tag, attributes, text in (
        ("fldChar", {"fldCharType": "begin"}, None),
        ("instrText", {}, ' TOC \\o "1-3" \\h \\z \\u '),
        ("fldChar", {"fldCharType": "separate"}, None),
    ):
        node = _node("w:" + tag, **attributes)
        node.text = text
        first.add_run()._r.append(node)
    last = document.add_paragraph("Section prototype", style="TOC 2")
    last.add_run()._r.append(_node("w:fldChar", fldCharType="end"))


def number_audit_headings(document) -> None:
    """Add native constant-prefix lists, as in the authored base, without altering titles."""
    titles = {section.title: section for section in CATALOGUE}
    title_ch5 = next(section.title for section in CATALOGUE if section.id == "ch5")
    measured = any(paragraph.text == title_ch5 for paragraph in document.paragraphs)
    chapter = 0
    subsection = 0
    group_ids = {}
    numbering = document.part.numbering_part.element
    next_id = 1200
    for paragraph in document.paragraphs:
        properties = paragraph._p.pPr
        direct = properties.find(qn("w:outlineLvl")) if properties is not None else None
        style = paragraph.style.style_id if paragraph.style is not None else ""
        level = (
            int(direct.get(qn("w:val")))
            if direct is not None
            else (int(style[-1]) - 1 if style in {"Heading1", "Heading2", "Heading3"} else None)
        )
        if level is None:
            continue
        if level == 0:
            section = titles[paragraph.text]
            chapter = section.chapter - (1 if not measured and section.chapter > 5 else 0)
            subsection = 0
            pattern = f"{chapter}."
        elif level == 1:
            subsection += 1
            pattern = f"{chapter}.%1."
        else:
            pattern = f"{chapter}.{subsection}.%1."
        if pattern not in group_ids:
            number_id = str(next_id)
            next_id += 1
            abstract = _node("w:abstractNum", abstractNumId=number_id)
            definition = _node("w:lvl", ilvl="0")
            definition.extend(
                [
                    _node("w:start", val="1"),
                    _node("w:numFmt", val="decimal"),
                    _node("w:lvlText", val=pattern),
                ]
            )
            abstract.append(definition)
            num = _node("w:num", numId=number_id)
            num.append(_node("w:abstractNumId", val=number_id))
            numbering.find(qn("w:num")).addprevious(abstract)
            numbering.append(num)
            group_ids[pattern] = number_id
        numpr = _node("w:numPr")
        numpr.append(_node("w:numId", val=group_ids[pattern]))
        paragraph._p.get_or_add_pPr().append(numpr)


RETAINED_CONTENT = {"ch1.obiective", "ch2.date_generale", "ch7"}


def confirm_retained_content(ws, job) -> None:
    """Complete the minimum retained fixed-chapter content before the rest is marked N/A."""

    for section_id in RETAINED_CONTENT:
        state = get_status(ws, job, section_id)
        if state.status == Status.DONE:
            continue
        if state.status in {Status.MISSING, Status.READY}:
            set_status(ws, job, section_id, Status.READY, "ema")
            mark_drafted(ws, job, section_id, "ema", ())
        set_status(ws, job, section_id, Status.DONE, "user")
