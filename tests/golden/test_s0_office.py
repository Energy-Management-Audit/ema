"""S0 package and Word proofs against the local reference library."""

import io
import math
import os
import shutil
from pathlib import Path
from zipfile import ZipFile

import pypdfium2
import pytest
from lxml import etree
from openpyxl import load_workbook

from conftest import artifacts_path
from ema.core.config import Settings
from ema.core.office.charts import (
    Series,
    build_column_chart,
    clone_chart,
    embed_data,
    read_series,
)
from ema.core.office.package import check_standalone, inspect, read_parts, write_parts
from ema.core.office.word_api import word_automation
from ema.core.office.workbook import formula_cells

pytestmark = pytest.mark.golden
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
OUTPUT = artifacts_path("s0")


@pytest.fixture(scope="module")
def references():
    base = os.environ.get("EMA_REFERENCE")
    if not base:
        pytest.skip("EMA_REFERENCE unavailable; golden not verified")
    root = Path(base)
    piee = (
        root
        / "piee/cases/piee-case-a/generated"
        / ("Program de îmbunătățire a eficienței energetice CLIENT-P1 SA_2026.docx")
    )
    audit = root / "audit/finished-audits" / ("AUDIT ENERGETIC AUDIT-01 ORAS - 2026.docx")
    assert piee.is_file() and audit.is_file()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    return piee, audit


def _workbook_values(archive, embedded, formula):
    book = load_workbook(io.BytesIO(archive.read(embedded)), read_only=True, data_only=True)
    sheet, cells = formula_cells(formula)
    return [book[sheet].cell(row, col).value for row, col in cells]


def _toc(source):
    root = etree.fromstring(read_parts(source)["word/document.xml"])
    result = []
    for paragraph in root.iter(W + "p"):
        props = paragraph.find(W + "pPr")
        style = props.find(W + "pStyle") if props is not None else None
        if style is None or not (style.get(W + "val") or "").upper().startswith("TOC"):
            continue
        texts = list(paragraph.iter(W + "t"))
        if texts and texts[-1].text and texts[-1].text.isdecimal():
            result.append((paragraph, texts[-1]))
    return root, result


def _formatting(paragraph):
    props = paragraph.find(W + "pPr")
    runs = [node for node in paragraph.iter(W + "rPr")]
    return (
        etree.tostring(props, method="c14n") if props is not None else b"",
        [etree.tostring(node, method="c14n") for node in runs],
    )


def _visual_xml(docx, part, normalize_axis_title=False, generated=False):
    root = etree.fromstring(read_parts(docx)[part])
    for node in root.xpath("//*[local-name()='uniqueId']"):
        node.set("val", "{new-id}")
    for series in root.xpath("//*[local-name()='ser']"):
        series.getparent().remove(series)
    if normalize_axis_title:
        for text in root.xpath(
            "//*[local-name()='valAx']/*[local-name()='title']//*[local-name()='t']"
        ):
            text.text = "axis title"
    if generated:
        for node in root.xpath(
            "//*[local-name()='plotArea']/*[local-name()='layout']/"
            "*[local-name()='manualLayout'] | "
            "//*[local-name()='plotArea']/*/*[local-name()='title']/"
            "*[local-name()='layout']/*[local-name()='manualLayout'] | "
            "//*[local-name()='catAx']/*[local-name()='tickLblSkip']"
        ):
            node.getparent().remove(node)
        for node in root.xpath("//*[local-name()='valAx']/*[local-name()='numFmt']"):
            node.set("formatCode", "value precision")
    return etree.tostring(root, method="c14n")


def test_piee_standalone_and_editable_caches(references):
    piee, _audit = references
    report = inspect(piee)
    assert len(report.charts) == 31
    assert check_standalone(piee) == []
    assert len({chart.embedded for chart in report.charts}) == 31
    with ZipFile(piee) as archive:
        for chart in report.charts:
            for series in read_series(piee, chart.part):
                assert series.refs is not None
                assert series.refs.values is not None
                actual = _workbook_values(archive, chart.embedded, series.refs.values)
                # XLSX stores 15 significant digits while the XML cache can keep a longer float.
                assert all(
                    math.isclose(value, cached, rel_tol=1e-12, abs_tol=1e-12)
                    if value is not None and cached is not None
                    else value == cached
                    for value, cached in zip(actual, series.values, strict=True)
                )
                if series.refs.name:
                    names = _workbook_values(archive, chart.embedded, series.refs.name)
                    if chart.part == "word/charts/chart31.xml":
                        # Approved golden has a series-name cache whose workbook cell is blank;
                        # a defect of the one-off script, not reproduced by product code.
                        assert names == [None]
                    else:
                        assert names == [series.name]
                if series.refs.categories:
                    categories = _workbook_values(archive, chart.embedded, series.refs.categories)
                    assert [str(value) for value in categories] == series.categories
    copy = OUTPUT / "piee-standalone.docx"
    shutil.copy2(piee, copy)
    word_automation(Settings()).open_check(copy)


def test_audit_embed_clone_and_build(references):
    _piee, audit = references
    original = {chart.part: read_series(audit, chart.part) for chart in inspect(audit).charts}
    current = audit
    for index in range(1, 31):
        next_file = OUTPUT / "audit-embedding-step.docx"
        if current == next_file:
            next_file = OUTPUT / "audit-embedding-step2.docx"
        embed_data(current, f"word/charts/chart{index}.xml", next_file)
        current = next_file
    embedded = OUTPUT / "audit-embedded.docx"
    shutil.copy2(current, embedded)
    assert check_standalone(embedded) == []
    for part, expected in original.items():
        assert read_series(embedded, part) == expected
    word_automation(Settings()).open_check(embedded)
    source_series = original["word/charts/chart11.xml"][0]
    pv = [
        Series("Energie electrică PV", source_series.categories, [float(i) / 8 for i in range(12)])
    ]
    cloned_file = OUTPUT / "audit-clone.docx"
    clone_part = clone_chart(embedded, "word/charts/chart11.xml", pv, None, cloned_file)
    built_file = OUTPUT / "audit-built.docx"
    built_part = build_column_chart(
        cloned_file, "word/charts/chart11.xml", pv, "MWh", built_file, after_part=clone_part
    )
    assert read_series(built_file, clone_part)[0].values == pv[0].values
    assert read_series(built_file, built_part)[0].values == pv[0].values
    assert check_standalone(built_file) == []
    assert _visual_xml(embedded, "word/charts/chart11.xml") == _visual_xml(built_file, clone_part)
    # #56 deliberately changes these three properties of generated monthly charts.
    # Everything else, and the unmodified clone above, still matches her visual XML.
    assert _visual_xml(
        embedded, "word/charts/chart11.xml", normalize_axis_title=True, generated=True
    ) == _visual_xml(built_file, built_part, normalize_axis_title=True, generated=True)
    built_root = etree.fromstring(read_parts(built_file)[built_part])
    assert built_root.xpath("//*[local-name()='catAx']/*[local-name()='tickLblSkip']/@val") == ["1"]
    assert built_root.xpath("//*[local-name()='valAx']/*[local-name()='numFmt']/@formatCode") == [
        "#,##0.00"
    ]
    assert not built_root.xpath("//*[local-name()='manualLayout']")
    report = inspect(built_file)
    with ZipFile(built_file) as archive:
        for part in (clone_part, built_part):
            chart = next(item for item in report.charts if item.part == part)
            product_series = read_series(built_file, part)[0]
            assert product_series.refs is not None and product_series.refs.name is not None
            assert product_series.refs == source_series.refs
            assert _workbook_values(archive, chart.embedded, product_series.refs.name) == [
                product_series.name
            ]
    word_automation(Settings()).open_check(cloned_file)
    word_automation(Settings()).open_check(built_file)


def test_toc_page_numbers_and_pdf(references):
    piee, _audit = references
    original_root, original_toc = _toc(piee)
    assert len(original_toc) == 28
    original_pages = [node.text for _paragraph, node in original_toc]
    formatting = [_formatting(paragraph) for paragraph, _node in original_toc]
    scrambled = OUTPUT / "piee-toc-scrambled.docx"
    parts = read_parts(piee)
    for _paragraph, node in original_toc:
        node.text = "99"
    parts["word/document.xml"] = etree.tostring(original_root, encoding="UTF-8")
    write_parts(parts, scrambled)
    word_automation(Settings()).update_toc_pages(scrambled)
    _updated_root, updated = _toc(scrambled)
    assert [node.text for _paragraph, node in updated] == original_pages
    assert [_formatting(paragraph) for paragraph, _node in updated] == formatting
    pdf = OUTPUT / "piee-toc-updated.pdf"
    word_automation(Settings()).render_pdf(scrambled, pdf)
    assert len(pypdfium2.PdfDocument(pdf)) == 33
