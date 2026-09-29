"""Level 2 reference: delivered energy-manager reports from the real annexes."""

from __future__ import annotations

import math
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from tests.golden.cases import case_path

from ema.energy_data.anexa import parse_anexa
from ema.reporting import generate, write_report

pytestmark = pytest.mark.golden
YEARS = (2023, 2024, 2025)


def _same(actual: object, expected: object) -> bool:
    if isinstance(actual, int | float) and isinstance(expected, int | float):
        return math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-9)
    return actual == expected


def _compare_sheets(actual: Worksheet, expected: Worksheet) -> None:
    assert actual.max_row == expected.max_row
    for row in range(1, expected.max_row + 1):
        for col in range(1, 8):
            assert _same(actual.cell(row, col).value, expected.cell(row, col).value), (
                actual.title,
                actual.cell(row, col).coordinate,
            )


def _blocks(sheet: Worksheet) -> dict[str, tuple[tuple[object, ...], ...]]:
    result: dict[str, tuple[tuple[object, ...], ...]] = {}
    name: str | None = None
    rows: list[tuple[object, ...]] = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        if row[0] == "TOTAL":
            break
        if isinstance(row[0], int):
            if name is not None:
                result[name] = tuple(rows)
            name = str(row[1])
            rows = []
        if name is not None:
            rows.append(tuple(row[:7]))
    if name is not None:
        result[name] = tuple(rows)
    return result


def _exception_category(situation: str) -> str | None:  # noqa: PLR0911
    if "Totalurile lunar și anual diferă" in situation:
        return "monthly_disagreement"
    if "CUI" in situation and "mai multe fișiere" in situation:
        return "duplicate_cui"
    if "M3 este indisponibil" in situation or "totalul în tep/an este indisponibil" in situation:
        return "monthly_fallback"
    if "denumire" in situation.lower() and "adres" in situation.lower():
        return "name_is_address"
    if "Consumul anual este" in situation and "1000 tep" in situation:
        return "below_1000"
    if "Costul investiției lipsește" in situation:
        return "missing_cost"
    if "Nu există măsuri" in situation:
        return "missing_measures"
    return None


def test_36_annexes_equal_delivered_year_sheets_and_control(
    reference_library: Path, tmp_path: Path
) -> None:
    annexes = list((reference_library / "piee" / "anexa-2-3-2025").glob("*.xls*"))
    assert len(annexes) == 36
    result = generate(annexes, YEARS)
    output = write_report(result, tmp_path / "report.xlsx")
    actual = load_workbook(output, data_only=True)
    delivered = load_workbook(
        next((reference_library / case_path("report-2023-2025", "final")).glob("*.xlsx")),
        data_only=True,
    )
    for year in YEARS:
        _compare_sheets(actual[str(year)], delivered[str(year)])
    assert actual.sheetnames == ["2023", "2024", "2025", "Control", "Exceptions"]
    control_actual = list(actual["Control"].values)[4:]
    control_delivered = list(delivered["Control"].values)[4:]
    assert len(control_actual) == len(control_delivered) == 36
    for actual_row, expected_row in zip(control_actual, control_delivered, strict=True):
        for col in (0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16):
            if col in (7, 8, 10, 11, 16) and actual_row[7] == "Date lunare!N3":
                continue
            assert _same(actual_row[col], expected_row[col]), (actual_row[1], col + 1)
        if actual_row[7] == "Date lunare!N3":
            assert actual_row[8] == actual_row[11]
            assert _same(actual_row[10], actual_row[8] - actual_row[9])
    expected_exceptions = {
        (row[1], row[2], category)
        for row in list(delivered["Exceptions"].values)[3:]
        if (category := _exception_category(str(row[3]))) is not None
    }
    actual_exceptions = {
        (row[1], row[2], category)
        for row in list(actual["Exceptions"].values)[3:]
        if (category := _exception_category(str(row[3]))) is not None
    }
    assert all(
        _exception_category(str(row[3])) is not None
        for row in list(actual["Exceptions"].values)[3:]
    )
    assert actual_exceptions == expected_exceptions - {
        item
        for item in expected_exceptions
        if item[2] == "monthly_fallback"
        and any(row[1] == item[0] and row[7] == "Date lunare!N3" for row in control_actual)
    }


def test_two_company_year_blocks_equal_delivered(reference_library: Path, tmp_path: Path) -> None:
    annexes = [case_path("report-2025", "annex_a"), case_path("report-2025", "annex_b")]
    assert len(annexes) == 2
    result = generate(annexes, YEARS)
    output = write_report(result, tmp_path / "report.xlsx")
    actual = load_workbook(output, data_only=True)
    delivered = load_workbook(case_path("report-2025", "workbook"), data_only=True)
    report_client_a = case_path("report-2025", "annex_a")
    report_client_a_name = next(
        company.name for company in result.companies if company.source == report_client_a
    )
    for year in YEARS:
        generated_blocks = _blocks(actual[str(year)])
        delivered_blocks = _blocks(delivered[str(year)])
        assert generated_blocks.keys() == delivered_blocks.keys()
        for name, rows in generated_blocks.items():
            expected_rows = delivered_blocks[name]
            older_omission = year == 2024 and name == report_client_a_name
            if older_omission:
                assert len(rows) == len(expected_rows) + 5
                source = parse_anexa(report_client_a)
                extra_source_rows = [
                    measure
                    for measure in source.existing_measures
                    if measure.description.ref.row in range(26, 31)
                ]
                assert [measure.description.ref.a1 for measure in extra_source_rows] == [
                    f"Solutii EE existente!B{row}" for row in range(26, 31)
                ]
                assert [row[4] for row in rows[10:15]] == [
                    " ".join(str(measure.description.value).split())
                    for measure in extra_source_rows
                ]
                source_book = load_workbook(report_client_a, data_only=True, read_only=True)
                source_sheet = source_book["Solutii EE existente"]
                for report_row, source_row in zip(rows[10:15], range(26, 31), strict=True):
                    assert _same(report_row[5], source_sheet.cell(source_row, 7).value)
                    assert _same(report_row[6], source_sheet.cell(source_row, 5).value)
                source_book.close()
            else:
                assert len(rows) == len(expected_rows)
            compared_rows = rows[:10] if older_omission else rows
            compared_expected = expected_rows[:10] if older_omission else expected_rows
            for index, (row, expected) in enumerate(
                zip(compared_rows, compared_expected, strict=True)
            ):
                assert all(_same(v, w) for v, w in zip(row[1:], expected[1:], strict=True)), (
                    year,
                    name,
                    index,
                )
