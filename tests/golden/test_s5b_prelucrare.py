"""S5b: reference reader and generator-regression writer checks."""

from __future__ import annotations

import re
import runpy
from pathlib import Path

import openpyxl
import pytest
from tests.golden.cases import case_path

from ema.energy_data.calc import co2, energy_intensity, specific_consumption, tep_total
from ema.energy_data.carriers import Carrier
from ema.energy_data.necesar import parse_necesar_info, to_dataset
from ema.energy_data.prelucrare import import_prelucrare
from ema.energy_data.prelucrare_factors import factors_for_output
from ema.energy_data.prelucrare_merge import merge_prelucrare
from ema.energy_data.prelucrare_writer import write_prelucrare

from .s3_book import Workbook
from .s3_workbooks import _factor_after_equals, load_case

pytestmark = pytest.mark.golden
CASES = (
    (
        "piee-case-a",
        case_path("piee-case-a", "prelucrare"),
        (2022, 2023, 2024),
    ),
    (
        "audit-case-c",
        case_path("audit-case-c", "prelucrare"),
        (2022, 2023, 2024),
    ),
    (
        "piee-case-b",
        case_path("piee-case-b", "prelucrare"),
        (2023, 2024, 2025),
    ),
)


@pytest.mark.parametrize("case,relative,years", CASES)
def test_reader_matches_auditor(
    reference_library: Path, case: str, relative: Path, years: tuple[int, ...]
) -> None:
    imported = import_prelucrare(reference_library / relative)
    _, expected, factors, _, _ = load_case(reference_library / relative, case, years)
    assert imported.dataset.years == years
    assert imported.dataset.production == expected.production
    assert imported.dataset.turnover_lei == expected.turnover_lei
    if case == "audit-case-c":
        assert set(imported.dataset.carriers[Carrier.water_potable]) == set(years)
        assert all(
            len(imported.dataset.carriers[Carrier.water_potable][year].months) == 12
            for year in years
        )
    assert imported.filed
    assert imported.located
    assert any(key.startswith("factor_sheet.principali") for key in imported.located)
    assert any(key.startswith("factor_sheet.factori") for key in imported.located)
    for carrier, by_year in expected.carriers.items():
        for year, series in by_year.items():
            actual = imported.dataset.carriers[carrier][year]
            for month, reading in series.months.items():
                assert actual.months[month].value == pytest.approx(reading.value)
                assert actual.months[month].unit == reading.unit
    expected_factors = {(factor.carrier, factor.unit, factor.per_unit) for factor in factors.tep}
    if case == "audit-case-c":
        # B1 preserves labeled, all-missing carriers; their filed factors remain traceable.
        book = Workbook.open(reference_library / relative)
        for carrier, sheet, label, unit in (
            (Carrier.ctl, "Consum Carburanti", "1 t (CTL", "t"),
            (Carrier.purchased_heat, "Consum Energie termica terti", "1 Gcal", "Gcal"),
        ):
            series = imported.dataset.carriers[carrier]
            assert set(series) == set(years)
            assert all(
                item.annual is None
                and len(item.months) == 12
                and all(reading.value is None for reading in item.months.values())
                for item in series.values()
            )
            factor = _factor_after_equals(book.sheet(sheet), label)
            assert factor is not None
            expected_factors.add((carrier, unit, factor))
    assert {
        (factor.carrier, factor.unit, factor.per_unit) for factor in imported.factors.tep
    } == expected_factors
    assert all(
        relative.name in factor.source and "!" in factor.source
        for factor in (*imported.factors.tep, *imported.factors.co2)
    )
    expected_issues = (
        [("cell_error", f"TEP!P{row}") for row in (10, 21, 32)] if case == "piee-case-a" else []
    )
    assert [
        (issue.code, issue.ref.a1 if issue.ref else None) for issue in imported.issues
    ] == expected_issues


def test_piee_case_a_generator_and_writer(  # noqa: C901
    reference_library: Path, tmp_path: Path
) -> None:
    case = reference_library / case_path("piee-case-a")
    prelucrare = import_prelucrare(case_path("piee-case-a", "prelucrare"))
    necesar = to_dataset(parse_necesar_info(case_path("piee-case-a", "necesar")))
    dataset, conflicts = merge_prelucrare(necesar, prelucrare)
    assert not conflicts
    expected = runpy.run_path(str(case_path("piee-case-a", "generator")))
    for name, carrier in (
        ("el", Carrier.electricity_grid),
        ("gas", Carrier.natural_gas),
        ("mot", Carrier.diesel),
        ("benz", Carrier.petrol),
        ("pv", Carrier.electricity_pv),
    ):
        for year in (2023, 2024, 2025):
            series = dataset.carriers[carrier][year]
            if name == "pv" and year == 2025:
                assert not series.months and series.annual is not None
                assert series.annual.value == pytest.approx(expected["pv_total"][year])
            else:
                assert [series.months[month].value for month in range(1, 13)] == pytest.approx(
                    expected[name][year]
                )
    for year in (2023, 2024, 2025):
        series = dataset.production["main"][year]
        assert [series.months[month].value for month in range(1, 13)] == pytest.approx(
            expected["prod"][year]
        )
        assert dataset.turnover_lei[year].value == pytest.approx(expected["turnover_lei"][year])
    output = tmp_path / "prelucrare.xlsx"
    write_prelucrare(
        dataset,
        (2023, 2024, 2025),
        output,
        factors_for_output(prelucrare, (2023, 2024, 2025)),
    )
    rebuilt = import_prelucrare(output)
    assert rebuilt.dataset.years == (2023, 2024, 2025)
    selected_factors = factors_for_output(prelucrare, (2023, 2024, 2025))
    for year in (2023, 2024, 2025):
        assert tep_total(rebuilt.dataset, selected_factors, year).value == pytest.approx(
            expected["tep"]["total"][year]
        )
        assert specific_consumption(
            rebuilt.dataset, selected_factors, year, None, "main"
        ).value == pytest.approx(expected["spec"]["total"][year])
        assert energy_intensity(rebuilt.dataset, selected_factors, year).value == pytest.approx(
            expected["intensity"][year]
        )
        co2_parts = [
            co2(rebuilt.dataset, selected_factors, year, carrier).value
            for carrier in (
                Carrier.electricity_grid,
                Carrier.natural_gas,
                Carrier.diesel,
                Carrier.petrol,
            )
        ]
        assert all(part is not None for part in co2_parts)
        assert sum(part for part in co2_parts if part is not None) == pytest.approx(
            expected["co2"][year]
        )
    for carrier, years in dataset.carriers.items():
        for year in (2023, 2024, 2025):
            if year not in years:
                continue
            first = years[year]
            second = rebuilt.dataset.carriers[carrier][year]
            # B1 retains the writer's labeled row even for an annual-only source.
            assert set(second.months) == (set(first.months) or set(range(1, 13)))
            assert first.months or all(reading.value is None for reading in second.months.values())
            for month, reading in first.months.items():
                assert second.months[month].value == pytest.approx(reading.value)
                assert second.months[month].unit == reading.unit
            assert (second.annual.value if second.annual else None) == pytest.approx(
                first.annual.value if first.annual else None
            )
    workbook = openpyxl.load_workbook(output, data_only=False)
    reconstructed = openpyxl.load_workbook(
        case / "working/calc-workbook-2023-2025.xlsx", read_only=True
    )
    assert set(reconstructed.sheetnames).issubset(set(workbook.sheetnames))
    formulas = [
        cell.value for sheet in workbook for row in sheet for cell in row if cell.data_type == "f"
    ]
    assert formulas and all("#REF!" not in formula for formula in formulas)
    for formula in formulas:
        for sheet_name, coordinate in re.findall(r"'([^']+)'!([A-Z]+[0-9]+)", formula):
            assert workbook[sheet_name][coordinate].value is not None
    assert any(
        "'Consum Electric'!" in formula and "'Principali factori de conversie'!" in formula
        for formula in formulas
    )
    assert any("'Consum electrica fotovoltaic'!" in formula for formula in formulas)
    reconstructed.close()
    workbook.close()
