"""piee_case_a's received files through `ema mcp`: an agent generates and reviews a PIEE draft.

Evidence level 1 (regression): the MCP draft equals the draft the shared use case produces from
the same inputs in the same workspace.
"""

from __future__ import annotations

import os
import sys
from pathlib import Path
from typing import Any

import anyio
import pytest
from docx import Document
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from tests.golden.cases import case_path
from tests.mcp_client import structured

from conftest import artifacts_path
from ema.clients.registry import create_client
from ema.core.workspace import Workspace
from ema.piee.workflow import GenerateRequest
from ema.workflows_registry import generate_draft

pytestmark = [pytest.mark.golden, pytest.mark.word]
EMA = Path(sys.executable).with_name("ema.exe" if os.name == "nt" else "ema")


def _texts(path: Path) -> tuple[list[str], list[str]]:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return paragraphs, cells


def test_agent_generates_and_reviews_piee_case_a_over_stdio(
    reference_library: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    base = next((reference_library / "piee/finished-programs").glob("*MODEL_2026.docx"))
    monkeypatch.setenv("EMA_PIEE_BASE_DOCUMENT", str(base))
    monkeypatch.setenv("EMA_PIEE_BASE_DIRECTORY", str(artifacts_path("s8", "base")))
    monkeypatch.setenv("EMA_WORKSPACE", str(tmp_path / "workspace"))
    received = reference_library / case_path("piee-case-a", "received")
    anexa = next(received.glob("Anexa*.xlsx"))
    necesar = next(received.glob("Necesar*.xls"))
    prelucrare = next(received.glob("*Prelucrare*.xls*"))
    ws = Workspace(tmp_path / "workspace")
    create_client(ws, "Synthetic", "12345678")
    parameters = StdioServerParameters(
        command=str(EMA), args=["mcp", "--import-root", str(received)], env=dict(os.environ)
    )

    async def agent() -> dict[str, Any]:
        seen: dict[str, Any] = {}
        async with (
            stdio_client(parameters) as (read, write),
            ClientSession(read, write) as client,
        ):
            await client.initialize()
            arguments = {
                "client": "12345678",
                "year": 2025,
                "anexa": str(anexa),
                "necesar": str(necesar),
                "prelucrare": str(prelucrare),
            }
            seen["draft"] = structured(await client.call_tool("piee_generate", arguments))
            job = {"job": seen["draft"]["job"]}
            conflicts = await client.call_tool("job_fields", {**job, "status": "conflict"})
            seen["conflicts"] = structured(conflicts)["fields"]
            seen["checks"] = structured(await client.call_tool("job_checks", job))
            field = seen["conflicts"][0]
            choice = {
                "field_id": field["id"],
                "action": "choose",
                "on_revision": field["revision"],
                "alternative": field["alternatives"][0]["id"],
            }
            seen["decision"] = structured(await client.call_tool("job_decide", {**job, **choice}))
        return seen

    seen = anyio.run(agent)

    mcp_draft = Path(seen["draft"]["draft"])
    assert mcp_draft.suffix == ".docx" and mcp_draft.is_file()
    assert Path(seen["draft"]["workbook"]).suffix == ".xlsx"
    assert seen["conflicts"]
    assert any(item["code"] == "conflict" for item in seen["checks"]["blocking"])
    assert seen["decision"]["actor"] == "agent"
    ws = Workspace(tmp_path / "workspace")
    direct = generate_draft(ws, GenerateRequest("12345678", 2025, anexa, necesar, prelucrare, None))
    mcp_paragraphs, mcp_cells = _texts(mcp_draft)
    direct_paragraphs, direct_cells = _texts(direct.draft)
    assert mcp_paragraphs == direct_paragraphs
    assert mcp_cells == direct_cells
    print(
        f"S19 golden: conflicts={len(seen['conflicts'])}; paragraphs={len(mcp_paragraphs)}; "
        f"cells={len(mcp_cells)}; identical=True; evidence=level 1"
    )
