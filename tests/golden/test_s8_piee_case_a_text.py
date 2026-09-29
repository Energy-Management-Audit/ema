"""Pin each semantic difference from the approved piee_case_a text and tables."""

from __future__ import annotations

import shutil
from collections import Counter
from datetime import date
from difflib import SequenceMatcher
from pathlib import Path

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from lxml import etree
from tests.golden.cases import case_path

from conftest import artifacts_path
from ema.core.office.numbers_ro import format_number
from ema.piee.compose import compose_draft
from ema.piee.dataset import load
from ema.piee.tables import _monthly

pytestmark = pytest.mark.golden

# These are document slot IDs, with physical IDs only for unbookmarked base prose.
# No reference wording or client values belong in this inventory.
PARAGRAPHS = {
    "body_9": "source_verbatim",
    "body_11": "source_policy",
    "body_14": "source_verbatim",
    "body_20": "source_verbatim",
    "body_30": "source_verbatim",
    "body_77": "unsupported_prose",
    "body_78": "unsupported_prose",
    "body_79": "unsupported_prose",
    "body_81": "unsupported_prose",
    "body_124": "unsupported_prose",
    "body_165": "unsupported_prose",
    "body_201": "unsupported_prose",
    "body_294": "source_verbatim",
    "body_child_267": "authored_trend_conflicts_with_chart",
    "body_child_290": "authored_trend_conflicts_with_chart",
    "body_343": "source_policy",
    "body_child_309": "source_policy",
    "body_363": "unsupported_prose",
    "body_390": "unsupported_prose",
    "body_414": "unsupported_prose",
    "approved_only_361": "unsupported_prose",
    "body_child_138": "figure_renumbered",
    "identity_423": "figure_renumbered",
    "body_146": "figure_renumbered",
    "table_104_r2_c1": "half_up_rounding",
    **{f"table_374_r{row}_c0": "source_verbatim" for row in (4, 6, 7, 8, 9)},
    **{
        f"table_395_r{row}_c{column}": "source_complete_authored_omission"
        for row in (24, 25, 26)
        for column in range(6)
    },
    "table_420_r0_c0": "source_verbatim",
}


def _text(element: etree._Element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def _body(document: DocxDocument) -> list[etree._Element]:
    return [element for element in document.element.body if element.tag in {qn("w:p"), qn("w:tbl")}]


def _slot(paragraph: etree._Element, index: int) -> str:
    for bookmark in paragraph.iter(qn("w:bookmarkStart")):
        name = bookmark.get(qn("w:name"), "")
        if name.startswith("_ema_"):
            return name.removeprefix("_ema_")
    return f"body_child_{index}"


def _styled_chars(paragraph: etree._Element) -> list[tuple[str, bytes]]:
    result = []
    for run in paragraph.iter(qn("w:r")):
        properties = run.find(qn("w:rPr"))
        style = etree.tostring(properties, method="c14n") if properties is not None else b""
        result.extend((letter, style) for letter in _text(run))
    return result


VISIBLE_RUN_PROPERTIES = {
    "rFonts",
    "sz",
    "b",
    "i",
    "u",
    "color",
    "highlight",
    "vertAlign",
    "caps",
    "smallCaps",
}


def _visible_chars(paragraph: etree._Element) -> list[tuple[str, frozenset[tuple[str, str]]]]:
    result = []
    for run in paragraph.iter(qn("w:r")):
        properties = run.find(qn("w:rPr"))
        style = (
            frozenset(
                (child.tag, str(sorted(child.attrib.items())))
                for child in properties
                if etree.QName(child).localname in VISIBLE_RUN_PROPERTIES
                and not (child.tag == qn("w:color") and child.get(qn("w:val")) == "auto")
            )
            if properties is not None
            else frozenset()
        )
        result.extend((character, style) for character in _text(run))
    return result


def _formatting_same(actual: etree._Element, approved: etree._Element) -> bool:
    if actual.tag != approved.tag:
        return False
    if actual.tag == qn("w:tc"):
        left, right = actual.findall(qn("w:p")), approved.findall(qn("w:p"))
        return len(left) == len(right) and all(
            _formatting_same(a, b) for a, b in zip(left, right, strict=True)
        )
    return _styled_chars(actual) == _styled_chars(approved) or (
        _styled_chars(actual) == _styled_chars(approved)[: len(_styled_chars(actual))]
        and not _text(approved)[len(_text(actual)) :].strip()
    )


def _table_differences(
    actual: etree._Element,
    approved: etree._Element,
    table: int,
    rounding_values: dict[str, float],
) -> dict[str, str]:
    found: dict[str, str] = {}
    left_rows, right_rows = actual.findall(qn("w:tr")), approved.findall(qn("w:tr"))
    if table == 395:
        assert len(left_rows) == 27 and len(right_rows) == 24
    else:
        assert len(left_rows) == len(right_rows)
    for row, (left, right) in enumerate(zip(left_rows, right_rows, strict=False)):
        left_cells, right_cells = left.findall(qn("w:tc")), right.findall(qn("w:tc"))
        assert len(left_cells) == len(right_cells)
        for column, (left_cell, right_cell) in enumerate(zip(left_cells, right_cells, strict=True)):
            if _text(left_cell) == _text(right_cell):
                continue
            key = f"table_{table}_r{row}_c{column}"
            assert key in PARAGRAPHS, key
            if PARAGRAPHS[key] == "half_up_rounding":
                raw = rounding_values[key]
                assert _text(left_cell) == format_number(raw, 2)
                old = f"{raw:,.2f}".replace(",", "\x00").replace(".", ",").replace("\x00", ".")
                assert _text(right_cell) == old
            if PARAGRAPHS[key] == "formatting_only":
                assert _formatting_same(left_cell, right_cell), key
            found[key] = PARAGRAPHS[key]
    for row in range(len(right_rows), len(left_rows)):
        for column, _ in enumerate(left_rows[row].findall(qn("w:tc"))):
            key = f"table_{table}_r{row}_c{column}"
            assert key in PARAGRAPHS, key
            found[key] = PARAGRAPHS[key]
    return found


def _paragraph_difference(actual: etree._Element, approved: etree._Element, key: str) -> str:
    assert key in PARAGRAPHS, key
    category = PARAGRAPHS[key]
    if category == "figure_renumbered":
        assert _text(actual).count("3 c)") == 1
        assert _text(actual).replace("3 c)", "3 d)") == _text(approved)
    if category in {"unsupported_prose", "source_policy"} and key in {
        "body_77",
        "body_78",
        "body_79",
        "body_81",
        "body_124",
        "body_165",
        "body_201",
        "body_363",
        "body_390",
        "body_414",
    }:
        assert _text(actual) == "n.d.", key
    if category == "formatting_only":
        assert _formatting_same(actual, approved), key
    return category


def test_piee_case_a_text_and_tables_have_only_pinned_differences(  # noqa: C901
    reference_library: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    case = reference_library / case_path("piee-case-a")
    data = load(
        2025,
        next(case.rglob("Anexa*.xlsx")),
        next(case.rglob("Necesar*.xls")),
        next(case.rglob("*Prelucrare*.xls*")),
    )
    monkeypatch.setattr("ema.piee.compose._strip_bookmarks", shutil.copyfile)
    output = tmp_path / "bookmarked.docx"
    compose_draft(data, artifacts_path("s8", "base"), output, date(2026, 9, 19))
    approved = next((case / "generated").glob("Program de îmbunătățire*.docx"))
    left, right = _body(Document(output)), _body(Document(approved))
    a = [(node.tag, _text(node)) for node in left]
    b = [(node.tag, _text(node)) for node in right]
    found: dict[str, str] = {}
    formatting_only = 0
    table_numbers = {102: 104, 324: 367, 331: 374, 352: 395, 376: 420}
    rounding = _monthly(data, 6)[0]
    assert rounding is not None
    rounding_values = {"table_104_r2_c1": rounding}
    for operation, start, end, ref_start, ref_end in SequenceMatcher(
        None, a, b, autojunk=False
    ).get_opcodes():
        if operation == "equal":
            for index, reference in zip(range(start, end), range(ref_start, ref_end), strict=True):
                actual_node, approved_node = left[index], right[reference]
                if actual_node.tag != qn("w:p"):
                    continue
                assert _text(actual_node) == _text(approved_node)
                assert _visible_chars(actual_node) == _visible_chars(approved_node), index
                if _styled_chars(actual_node) != _styled_chars(approved_node):
                    formatting_only += 1
            continue
        if operation == "insert":
            assert (start, end, ref_start, ref_end) == (362, 362, 361, 362)
            key = "approved_only_361"
            assert _text(right[ref_start]).strip()
            found[key] = PARAGRAPHS[key]
            continue
        if (operation, start, end, ref_start, ref_end) == ("replace", 307, 309, 307, 308):
            for index in range(start, end):
                key = _slot(left[index], index + 1)
                assert key in {"body_343", "body_child_309"}, key
                assert "Date indisponibile" in _text(left[index]), key
                found[key] = PARAGRAPHS[key]
            continue
        assert operation == "replace" and end - start == ref_end - ref_start
        for index, reference in zip(range(start, end), range(ref_start, ref_end), strict=True):
            actual_node, approved_node = left[index], right[reference]
            if actual_node.tag == qn("w:tbl"):
                found.update(
                    _table_differences(
                        actual_node, approved_node, table_numbers[index], rounding_values
                    )
                )
                continue
            key = _slot(actual_node, index + 1)
            found[key] = _paragraph_difference(actual_node, approved_node, key)
    assert found.keys() == PARAGRAPHS.keys(), sorted(found.keys() ^ PARAGRAPHS.keys())
    assert Counter(found.values()) == Counter(PARAGRAPHS.values())
    # Three renumbered paragraphs now belong to the explicit difference inventory.
    assert formatting_only == 134
    assert "(calculat)" not in "\n".join(_text(node) for node in left)
