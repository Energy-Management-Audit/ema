"""Level-2 audit base checks against local auditor documents and case inputs."""

from __future__ import annotations

import hashlib
import json
import re
from copy import deepcopy
from pathlib import Path
from tempfile import NamedTemporaryFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from ema.audit.base import build_base, build_configured_base
from ema.audit.base_anchor import MARKER, _paragraph_text
from ema.audit.base_identity import derive_identity as _identity
from ema.audit.base_numeric import approved_fixed_text, has_number
from ema.audit.base_package import package_issues
from ema.audit.base_units import UnitPlan, select_units
from ema.audit.catalogue import CATALOGUE
from ema.audit.headings import map_headings
from ema.audit.inventory import inventory
from ema.audit.render_plan import process_count
from ema.core.config import Settings
from ema.core.office.anchors import find
from ema.core.office.word_api import word_automation, word_available
from ema.energy_data.necesar import parse_necesar_info

pytestmark = pytest.mark.golden


def _references(root: Path) -> tuple[Path, Path]:
    audits = root / "audit/finished-audits"
    return next(audits.glob("*AUDIT-01*.docx")), next(audits.glob("*AUDIT-04*.docx"))


def _CLIENT-A1_plan(root: Path) -> UnitPlan:
    received = root / "audit/cases/audit-case-a/received"
    # Six distinct 5.x Flux schemes; the second 5.1 file is a revision.
    assert process_count([path.name for path in received.iterdir()], None) == (6, "schemes")
    panels = list(received.glob("13.[56].Armonici*.pdf"))
    assert len(panels) == 2
    info = parse_necesar_info(next(received.glob("*Necesar info*.xls")))
    families = {"electricity", "gas", "fuel"}
    assert len(info.carriers) == 6
    assert len([name for name in info.tables if name.startswith("echipamente ")]) == 3
    if info.water:
        families.add("water")
    return UnitPlan(received.parent.name, 6, frozenset(families), 2, False, 3, 0)


def _CLIENT-A2_plan(root: Path) -> UnitPlan:
    case = root / "audit/cases/audit-case-b"
    received = case / "received"
    # No 5.x schemes; the received Fisa has two body paragraphs beginning "Flux".
    fisa = next(received.glob("Fisa*.docx"))
    assert process_count([path.name for path in received.iterdir()], fisa) == (2, "fisa")
    assert list(received.glob("*ATR*.pdf"))  # electricity
    assert list(received.glob("*Gaze*.zip"))  # gas
    panels = list((case / "visit/electrical").glob("tablou-electric-*"))
    assert len(panels) == 4
    assert list((case / "visit/thermography").glob("*.jpeg"))
    return UnitPlan(case.name, 2, frozenset({"electricity", "gas"}), 4, True, 0, 0)


def _section_paragraphs(path: Path, section: str) -> list[object]:
    document = Document(path)
    mapping = map_headings(path, "AUDIT-01")
    heading = next(item.heading for item in mapping.mapped if item.section_id == section)
    later = [
        item.heading.index
        for item in mapping.mapped
        if item.heading.index > heading.index and item.heading.level <= heading.level
    ]
    return list(document.paragraphs[heading.index : min(later, default=len(document.paragraphs))])


def _paragraph_signature(paragraph: object) -> str:
    runs = [
        (run.text, etree.tostring(run._r.rPr, method="c14n") if run._r.rPr is not None else b"")
        for run in paragraph.runs  # type: ignore[attr-defined]
    ]
    payload = repr((paragraph.text, runs)).encode()  # type: ignore[attr-defined]
    return hashlib.sha256(payload).hexdigest()


def _style_signature(style: object) -> bytes:
    clone = deepcopy(style.element)  # type: ignore[attr-defined]
    clone.attrib.pop(qn("w:styleId"), None)
    for node in clone.iter():
        if node.tag in {qn("w:basedOn"), qn("w:next"), qn("w:link"), qn("w:numId")}:
            node.attrib.pop(qn("w:val"), None)
    return etree.tostring(clone, method="c14n")


def _prototype_format(base: Path, prototype: Path, output: Path) -> None:
    source = Document(base)
    reference = Document(prototype)
    built = Document(output)
    source_map = map_headings(base, "AUDIT-01")
    reference_map = map_headings(prototype, "AUDIT-04")
    built_map = map_headings(output, "AUDIT-01")
    source_process = next(item for item in source_map.mapped if item.section_id == "ch3.process")
    expected_style = source.paragraphs[source_process.heading.index].style.style_id
    for item in built_map.mapped:
        if item.section_id == "ch3.process":
            assert built.paragraphs[item.heading.index].style.style_id == expected_style
    for section in ("ch5", "ch5.electric", "ch5.electric_fisa", "ch5.termic"):
        original = next(item for item in reference_map.mapped if item.section_id == section)
        rendered = next((item for item in built_map.mapped if item.section_id == section), None)
        if rendered is None:
            continue
        reference_style = reference.paragraphs[original.heading.index].style
        built_style = built.paragraphs[rendered.heading.index].style
        assert _style_signature(reference_style) == _style_signature(built_style)
    numbering = built.part.numbering_part.element
    defined = {node.get(qn("w:numId")) for node in numbering.iter(qn("w:num"))}
    used = {
        node.get(qn("w:val"))
        for root in (built.element, built.styles.element)
        for node in root.iter(qn("w:numId"))
    }
    assert used - {"0"} <= defined


def _fixed_chapters(base: Path, output: Path, identity: tuple[str, ...]) -> None:
    for section in ("ch1", "ch7"):
        original = _section_paragraphs(base, section)
        built = _section_paragraphs(output, section)
        assert len(original) == len(built)
        for index, (source, target) in enumerate(zip(original, built, strict=True)):
            source_text = source.text  # type: ignore[attr-defined]
            if any(term.casefold() in source_text.casefold() for term in identity):
                assert MARKER in target.text or target.text != source_text  # type: ignore[attr-defined]
                continue
            if section == "ch7" and index == 0:
                assert source.style.style_id == target.style.style_id  # type: ignore[attr-defined]
                continue  # its printed chapter number changes when ch5 is inserted
            xml_text = _paragraph_text(source._p)  # type: ignore[attr-defined]
            if has_number(xml_text) and not approved_fixed_text(xml_text):
                assert MARKER in target.text  # type: ignore[attr-defined]
                continue
            assert _paragraph_signature(source) == _paragraph_signature(target)


def _toc_in_step(path: Path) -> None:
    root = Document(path).element
    names = {
        node.get(qn("w:name"))
        for node in root.iter(qn("w:bookmarkStart"))
        if (node.get(qn("w:name")) or "").startswith("_Toc")
    }
    refs = {
        match
        for node in root.iter(qn("w:instrText"))
        for match in re.findall(r"PAGEREF\s+(_Toc\d+)", node.text or "")
    }
    links = {
        node.get(qn("w:anchor"))
        for node in root.iter(qn("w:hyperlink"))
        if (node.get(qn("w:anchor")) or "").startswith("_Toc")
    }
    assert names == refs == links
    assert len(names) == len(set(names))


@pytest.mark.parametrize("case", ("CLIENT-A1", "CLIENT-A2"))
def test_audit_base_for_case(reference_library: Path, tmp_path: Path, case: str) -> None:
    base, prototype = _references(reference_library)
    identity = _identity(base)
    plan = _CLIENT-A1_plan(reference_library) if case == "CLIENT-A1" else _CLIENT-A2_plan(reference_library)
    with NamedTemporaryFile(suffix="-AUDIT-01.docx") as temporary:
        selected = Document(base)
        select_units(selected, base, prototype, plan)
        selected.save(temporary.name)
        units = inventory(Path(temporary.name))
        assert len(units.processes) == plan.processes
        assert len(units.measured_panels) == plan.measured_panels
        assert len(units.equipment_tables) == plan.equipment_tables
        assert len(units.measures) == plan.measures
    output = build_configured_base(
        plan,
        settings=Settings(audit_base_document=base, audit_measurement_prototype=prototype),
        output=tmp_path / f"{case}.docx",
        base_identity=identity,
    )
    mapping = map_headings(output, "AUDIT-01")
    assert not mapping.unmapped
    assert {item.section_id for item in mapping.mapped} <= {item.id for item in CATALOGUE}
    assert sum(item.section_id == "ch3.process" for item in mapping.mapped) == plan.processes
    assert sum(item.section_id == "ch5" for item in mapping.mapped) == 1
    document = Document(output)
    assert all(
        MARKER in document.paragraphs[item.heading.index].text
        for item in mapping.mapped
        if item.section_id == "ch3.process"
    )
    toc_roots = [
        paragraph.text for paragraph in document.paragraphs if paragraph.style.style_id == "TOC1"
    ]
    assert all(
        any(title.startswith(f"{chapter}. ") for title in toc_roots) for chapter in (5, 6, 7)
    )
    assert not package_issues(output, identity)
    _prototype_format(base, prototype, output)
    _fixed_chapters(base, output, identity)
    _toc_in_step(output)
    settings = Settings()
    if not word_available(settings):
        pytest.skip(f"Word unavailable at {settings.word_path}")
    word = word_automation(settings)
    word.open_check(output)
    anchors = json.loads(output.with_suffix(".anchors.json").read_text(encoding="utf-8"))
    roots = [Document(output).element]
    variable = [anchor for anchor in anchors["anchors"] if anchor["classification"] == "variable"]
    assert variable
    for anchor in variable:
        if anchor["part"] == "word/document.xml":
            paragraph = find(roots, anchor["slot"])
            assert MARKER in "".join(node.text or "" for node in paragraph.iter(qn("w:t")))
    print(f"{case}: level 2 structure/format/package")
    print(f"processes={plan.processes}; panels={plan.measured_panels}")


def test_absent_measurements_shift_printed_chapters(
    reference_library: Path, tmp_path: Path
) -> None:
    base, prototype = _references(reference_library)
    plan = UnitPlan("synthetic", 2, frozenset({"electricity", "gas"}), 0, False, 0, 0)
    output = build_base(
        plan,
        base_document=base,
        measurement_prototype=prototype,
        output=tmp_path / "without-measurements.docx",
        base_identity=_identity(base),
    )
    mapping = map_headings(output, "AUDIT-01")
    assert not any(item.section_id == "ch5" for item in mapping.mapped)
    document = Document(output)
    toc_roots = [
        paragraph.text for paragraph in document.paragraphs if paragraph.style.style_id == "TOC1"
    ]
    assert any(title.startswith("5. ") for title in toc_roots)
    assert any(title.startswith("6. ") for title in toc_roots)
    assert not any(title.startswith("7. ") for title in toc_roots)
    _toc_in_step(output)
