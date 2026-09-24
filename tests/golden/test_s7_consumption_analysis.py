"""Reference-level proof for generated chapter-four blocks and pinned exceptions."""

from __future__ import annotations

import os
from pathlib import Path

import pytest
from docx import Document
from lxml import etree
from tests.golden.s7_coverage import retained_categories
from tests.golden.s7_layout import build_blocks
from tests.golden.s7_reader import _monthly, _text, read_chapter

from ema.audit.headings import map_headings
from ema.consumption_analysis.analysis import (
    ChartPlan,
    Metric,
    SectionPlan,
    TablePlan,
    analyze,
)
from ema.core.office.blocks import ElementLocator, Prototypes, Retained
from ema.core.office.charts import embed_all_data
from ema.core.office.numbers_ro import format_number
from ema.core.office.package import C, check_standalone, read_parts
from ema.core.office.region import replace_region
from ema.energy_data.carriers import Carrier
from ema.energy_data.factors import FactorTable

pytestmark = pytest.mark.golden
OUTPUT = Path.home() / "Ema-dev" / "s7"


def _source(name: str) -> Path:
    root = Path(os.environ["EMA_REFERENCE"]) / "audit" / "finished-audits"
    matches = list(root.glob(f"*{name}*.docx"))
    assert len(matches) == 1
    return matches[0]


def _chart_types(parts: dict[str, bytes], part: str) -> tuple[str, ...]:
    root = etree.fromstring(parts[part])
    plot = root.find(f"{{{C}}}chart/{{{C}}}plotArea")
    assert plot is not None
    return tuple(node.tag for node in plot if node.tag.endswith("Chart"))


@pytest.mark.parametrize("name", ["AUDIT-01", "AUDIT-02"])
def test_raw_monthly_tables_are_dataset_backed(name: str) -> None:
    case = read_chapter(_source(name), name.lower())
    for section_id, carrier in (
        ("ch4.electricitate", Carrier.electricity_grid),
        ("ch4.gaz", Carrier.natural_gas),
    ):
        section = case.sections[section_id]
        years, monthly = _monthly(section)
        for year in years:
            for month in range(1, 13):
                reading = case.dataset.carriers[carrier][year].months[month]
                assert reading.value is not None
                assert format_number(reading.value, 2, grouping=False) == format_number(
                    monthly[year][month], 2, grouping=False
                )


@pytest.mark.parametrize("name", ["AUDIT-01", "AUDIT-02"])
def test_gas_table_and_chart_regenerate_from_chapter_four(name: str) -> None:
    source = _source(name)
    case = read_chapter(source, name.lower())
    section = case.sections["ch4.gaz"]
    original_table = section.tables[0][1]
    chart = section.charts[-1]
    years = case.dataset.years
    plan = SectionPlan(
        "ch4.gaz",
        "carrier",
        (
            TablePlan(
                "monthly",
                tuple(
                    Metric("carrier", (Carrier.natural_gas,), month=month) for month in range(1, 7)
                ),
                years,
            ),
            ChartPlan(
                "annual",
                chart.part,
                chart.series.name,
                Metric("carrier", (Carrier.natural_gas,)),
                years,
                tuple(chart.series.categories),
            ),
        ),
    )
    blocks = analyze(case.dataset, FactorTable("reference", min(years), (), ()), (plan,))[0].blocks
    OUTPUT.mkdir(parents=True, exist_ok=True)
    embedded = OUTPUT / f"{name.lower()}-embedded.docx"
    table_out = OUTPUT / f"{name.lower()}-table.docx"
    final = OUTPUT / f"{name.lower()}-gas-regenerated.docx"
    embed_all_data(source, embedded)
    table_index = section.tables[0][0]
    replace_region(
        embedded,
        table_out,
        ElementLocator(table_index),
        ElementLocator(table_index + 2),
        [blocks[0]],
        Prototypes({"monthly": original_table}, 4),
    )
    original_document = Document(str(source))
    chart_paragraph = etree.fromstring(
        etree.tostring(list(original_document.element.body)[chart.body_index])
    )
    replace_region(
        table_out,
        final,
        ElementLocator(chart.body_index),
        ElementLocator(chart.body_index + 2),
        [blocks[1]],
        Prototypes({"annual": chart_paragraph}, 4),
    )
    regenerated = read_chapter(final, name.lower())
    rebuilt_section = regenerated.sections["ch4.gaz"]
    rebuilt_table = rebuilt_section.tables[0][1]
    assert [
        _text(cell)
        for cell in original_table.iter(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc"
        )
    ] == [
        _text(cell)
        for cell in rebuilt_table.iter(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc"
        )
    ]
    rebuilt_chart = rebuilt_section.charts[-1]
    assert rebuilt_chart.series.categories == chart.series.categories
    assert rebuilt_chart.series.values == chart.series.values
    chart_types = []
    for docx, part in ((source, chart.part), (final, rebuilt_chart.part)):
        root = etree.fromstring(read_parts(docx)[part])
        chart_types.append([node.tag for node in root.iter() if node.tag == f"{{{C}}}barChart"])
    assert len(chart_types[0]) == len(chart_types[1]) == 1
    assert check_standalone(final) == []


def test_equivalent_chart_cache_exceptions_are_pinned() -> None:
    expected = {
        ("AUDIT-01", "ch4.echiv_electric", 2023),
        ("AUDIT-01", "ch4.echiv_electric", 2024),
        ("AUDIT-01", "ch4.echiv_electric", 2025),
        ("AUDIT-01", "ch4.echiv_gaz", 2024),
        ("AUDIT-01", "ch4.echiv_gaz", 2025),
        ("AUDIT-01", "ch4.echiv_carburant", 2024),
        ("AUDIT-02", "ch4.echiv_carburant", 2023),
        ("AUDIT-02", "ch4.echiv_carburant", 2024),
    }
    found = set()
    for name in ("AUDIT-01", "AUDIT-02"):
        case = read_chapter(_source(name), name.lower())
        for section_id in (
            "ch4.echiv_electric",
            "ch4.echiv_gaz",
            "ch4.echiv_carburant",
            "ch4.echiv_total",
        ):
            section = case.sections[section_id]
            if len(section.tables) != 2 or not section.charts:
                continue
            years, monthly = _monthly(section)
            values = section.charts[-1].series.values
            for year, chart_value in zip(years, values, strict=True):
                table_sum = sum(monthly[year].values())
                if chart_value is None or abs(chart_value - table_sum) > 0.01:
                    found.add((name, section_id, year))
    assert found == expected


@pytest.mark.parametrize("name,expected_generated", [("AUDIT-01", 63), ("AUDIT-02", 39)])
def test_whole_chapter_roundtrip_with_reported_block_coverage(  # noqa: PLR0915
    name: str, expected_generated: int
) -> None:
    source = _source(name)
    generated, case = build_blocks(name)
    assert len(generated) == expected_generated
    document = Document(str(source))
    mapping = map_headings(source, name.lower()).mapped
    chapter = next(item for item in mapping if item.section_id == "ch4")
    following = next(
        item
        for item in mapping
        if item.heading.index > chapter.heading.index and item.heading.level == 0
    )
    children = list(document.element.body)
    start = children.index(document.paragraphs[chapter.heading.index]._p)
    end = children.index(document.paragraphs[following.heading.index]._p)
    embedded = OUTPUT / f"{name.lower()}-embedded.docx"
    out = OUTPUT / f"{name.lower()}-chapter-four.docx"
    OUTPUT.mkdir(parents=True, exist_ok=True)
    embed_all_data(source, embedded)
    root = etree.fromstring(read_parts(embedded)["word/document.xml"])
    body = root.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body")
    assert body is not None
    original = list(body)
    prototypes = {f"item-{index}": original[index] for index in range(start + 1, end)}
    categories = retained_categories(source, name.lower(), generated)
    assert set(categories) == set(range(start + 1, end)) - set(generated)
    assert all(categories.values())
    blocks = [generated.get(index, Retained(f"item-{index}")) for index in range(start + 1, end)]
    replace_region(
        embedded,
        out,
        ElementLocator(start + 1),
        ElementLocator(end + 1),
        blocks,
        Prototypes(prototypes, 4),
    )
    result_root = etree.fromstring(read_parts(out)["word/document.xml"])
    result_body = result_root.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body"
    )
    assert result_body is not None
    result_nodes = list(result_body)
    for index in range(start + 1, end):
        if original[index].tag == "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p":
            assert _text(original[index]) == _text(result_nodes[index])
    regenerated = read_chapter(out, name.lower())
    source_parts = read_parts(source)
    result_parts = read_parts(out)
    assert list(regenerated.sections) == list(case.sections)
    for section_id in case.sections:
        before, after = case.sections[section_id], regenerated.sections[section_id]
        for table_index, ((_, source_table), (_, result_table)) in enumerate(
            zip(before.tables, after.tables, strict=True)
        ):
            original_cells = [
                [
                    _text(cell)
                    for cell in row.findall(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc"
                    )
                ]
                for row in source_table.findall(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr"
                )
            ]
            result_cells = [
                [
                    _text(cell)
                    for cell in row.findall(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc"
                    )
                ]
                for row in result_table.findall(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr"
                )
            ]
            differences = [
                (row_index, column_index)
                for row_index, (original_row, result_row) in enumerate(
                    zip(original_cells, result_cells, strict=True)
                )
                for column_index, (left, right) in enumerate(
                    zip(original_row, result_row, strict=True)
                )
                if left != right
            ]
            allowed = (
                [(1, 2)]
                if name == "AUDIT-02" and section_id == "ch4.echiv_electric" and table_index == 0
                else []
            )
            assert differences == allowed
        assert len(before.charts) == len(after.charts)
        for source_chart, result_chart in zip(before.charts, after.charts, strict=True):
            assert _chart_types(source_parts, source_chart.part) == _chart_types(
                result_parts, result_chart.part
            )
            assert source_chart.series.name == result_chart.series.name
            assert source_chart.series.categories == result_chart.series.categories
            if name == "AUDIT-01" and section_id in {
                "ch4.specific_electric",
                "ch4.specific_gaz",
            }:
                assert all(
                    left is not None and right is not None and abs(left * 1000 - right) < 1e-10
                    for left, right in zip(
                        source_chart.series.values, result_chart.series.values, strict=True
                    )
                )
            elif section_id in {"ch4.echiv_electric", "ch4.echiv_gaz"}:
                assert all(
                    left is not None and right is not None and abs(left - right) < 1e-10
                    for left, right in zip(
                        source_chart.series.values, result_chart.series.values, strict=True
                    )
                )
            else:
                assert source_chart.series.values == result_chart.series.values
    assert check_standalone(out) == []
