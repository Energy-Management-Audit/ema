"""Chart and numbered-table comparisons for the case C document golden."""

from __future__ import annotations

import math
import re
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from tests.golden.cases import case_path
from tests.golden.piee_case_c_measures import compare_measure_tables

from ema.consumption_analysis.analysis import Metric, value
from ema.core.office.chart_series import read_series
from ema.core.office.package import REL_CHART, C, R, read_parts, relationships, target_part, xml
from ema.core.office.run_range import visible_text
from ema.energy_data.calc import tep
from ema.energy_data.carriers import Carrier
from ema.energy_data.source import normal
from ema.piee.chart_plan import MONTHS
from ema.piee.dataset import PieeData

S16 = frozenset(
    {(10, index) for index in (0, 1, 2, 5, 6, 7, 8, 9)}
    | {(18, index) for index in range(12)}
    | {(22, index) for index in (0, 7, 8, 10, 11)}
    | {(24, 1), (33, 1), (34, 1)}
)
CAPTION_KINDS = {
    "production": r"produc",
    "electricity": r"electric",
    "gas": r"gaz",
    "fuel": r"carbur|motorin|benzin|gpl",
    "cogen": r"cogener",
    "coke": r"cocs",
    "water": r"apa",
    "specific": r"specific",
    "intensity": r"intensit",
    "mix": r"ponder",
    "total": r"total",
    "impact": r"impact|mediu|co2",
    "potable": r"potabil",
    "industrial": r"industrial",
    "storm": r"meteoric",
}


def _caption_kinds(path: Path) -> list[frozenset[str]]:
    root = xml(read_parts(path), "word/document.xml")
    kinds = []
    for paragraph in root.iter(qn("w:p")):
        if not list(paragraph.iter(f"{{{C}}}chart")):
            continue
        following = paragraph.getnext()
        assert following is not None
        caption = normal(visible_text(following))
        assert re.match(r"^(?:fig(?: nr)?|figura numarul) \d+", caption)
        found = {kind for kind, pattern in CAPTION_KINDS.items() if re.search(pattern, caption)}
        ordinal = len(kinds) + 1
        if ordinal <= 4 and "cantitatea" in caption:
            found.add("production")
        if ordinal == 34 and "specific" in caption and "energie" in caption:
            found.add("total")
        kinds.append(frozenset(found))
    return kinds


def _order(path: Path) -> list[str]:
    parts = read_parts(path)
    linked = {
        relation.get("Id"): target_part("word/document.xml", relation.get("Target", ""))
        for relation in relationships(parts, "word/document.xml")
        if relation.get("Type") == REL_CHART
    }
    return [
        linked[node.get(f"{{{R}}}id")]
        for node in xml(parts, "word/document.xml").iter(f"{{{C}}}chart")
    ]


def _source(data: PieeData, ordinal: int, index: int) -> float:
    assert data.prelucrare is not None
    if ordinal in {10, 18, 22}:
        carrier = {
            10: Carrier.natural_gas,
            18: Carrier.diesel,
            22: Carrier.coke,
        }[ordinal]
        reading = data.prelucrare.dataset.carriers[carrier][2024].months[index + 1]
        return reading.value
    if ordinal == 24:
        reading = data.prelucrare.dataset.carriers[Carrier.coke][2023 + index].annual
        assert reading is not None
        return reading.value
    key = "coke" if ordinal == 33 else "total"
    found = data.prelucrare.filed[f"specific.{key}.{2023 + index}"]
    assert isinstance(found.value, int | float)
    return float(found.value)


def _compare_figures(actual: Path, final: Path, data: PieeData) -> None:
    produced, delivered = _order(actual), _order(final)
    assert len(produced) == len(delivered) == 39
    assert len(set(produced)) == 39
    chart_parts = read_parts(actual)
    assert (
        sum(bool(list(xml(chart_parts, part).iter(f"{{{C}}}barChart"))) for part in produced) == 36
    )
    assert (
        sum(bool(list(xml(chart_parts, part).iter(f"{{{C}}}pie3DChart"))) for part in produced) == 3
    )
    differences: set[tuple[int, int]] = set()
    for ordinal, (actual_part, delivered_part) in enumerate(
        zip(produced, delivered, strict=True), 1
    ):
        actual_series = read_series(actual, actual_part)
        delivered_series = read_series(final, delivered_part)
        if ordinal in {17, 18, 19, 20}:
            # S16: the 2024 GPL cache is stale; S17: the other GPL series are zero.
            assert len(actual_series) == 1 and len(delivered_series) == 2
            if ordinal == 18:
                assert any(value not in {None, 0} for value in delivered_series[1].values)
                _zero_gpl_source()
            else:
                assert all(value in {None, 0} for value in delivered_series[1].values)
            delivered_series = delivered_series[:1]
        assert len(actual_series) == len(delivered_series), ordinal
        for actual_item, delivered_item in zip(actual_series, delivered_series, strict=True):
            actual_values, delivered_values = actual_item.values, delivered_item.values
            if ordinal == 30:
                # S17: the authored 2025 pie retains a zero-only GPL category.
                assert len(actual_values) == 4 and len(delivered_values) == 5
                zero = [index for index, value in enumerate(delivered_values) if value == 0]
                assert len(zero) == 1
                assert normal(delivered_item.categories[zero[0]]) == "gpl"
                delivered_values = [
                    value for index, value in enumerate(delivered_values) if index != zero[0]
                ]
                _zero_gpl_source()
            assert len(actual_values) == len(delivered_values), ordinal
            for index, (actual_value, delivered_value) in enumerate(
                zip(actual_values, delivered_values, strict=True)
            ):
                if actual_value is None or delivered_value is None:
                    assert actual_value is delivered_value, (ordinal, index)
                # 1e-8 absorbs OOXML floating-point cache noise, not a display-precision slip.
                elif math.isclose(actual_value, delivered_value, rel_tol=1e-8, abs_tol=1e-8):
                    continue
                else:
                    pair = (ordinal, index)
                    assert pair in S16, pair
                    source = _source(data, ordinal, index)
                    assert math.isclose(actual_value, source, rel_tol=1e-9, abs_tol=1e-9)
                    assert not math.isclose(delivered_value, source, rel_tol=1e-8, abs_tol=1e-8)
                    differences.add(pair)
    assert differences == S16


def _zero_gpl_source() -> None:
    book = load_workbook(case_path("piee-case-c", "prelucrare"), read_only=True, data_only=True)
    for sheet_name in ("Consum Carburanti", "TEP"):
        sheet = book[sheet_name]
        rows = list(sheet.iter_rows())
        found: set[int] = set()
        labels = {"gpl"} if sheet_name == "Consum Carburanti" else {"gpl", "gpl tep"}
        for index, row in enumerate(rows):
            year_cell = next(
                (
                    (column, cell.value)
                    for column, cell in enumerate(row)
                    if cell.value in (2023, 2024, 2025)
                ),
                None,
            )
            if year_cell is None:
                continue
            column, year = year_cell
            if year in found:
                continue
            candidates = [
                later
                for later in rows[index + 1 : index + 11]
                if isinstance(later[column].value, str) and normal(later[column].value) in labels
            ]
            assert len(candidates) == 1, (sheet_name, year)
            assert all(
                cell.value == 0 for cell in candidates[0] if isinstance(cell.value, int | float)
            ), (sheet_name, year)
            found.add(year)
        assert found == {2023, 2024, 2025}, sheet_name
    book.close()


def _number(text: str) -> float:
    assert re.fullmatch(r"[\d. ,]+", text.strip()), text
    return float(text.strip().replace(" ", "").replace(".", "").replace(",", "."))


def _gas_total_is_sourced(data: PieeData) -> None:
    book = load_workbook(case_path("piee-case-c", "prelucrare"), read_only=True, data_only=True)
    sheet = next(book[name] for name in book.sheetnames if normal(name) == "consum gaz")
    for row in sheet.iter_rows():
        if 2024 not in [cell.value for cell in row]:
            continue
        total_col = next(
            (index for index, cell in enumerate(row) if normal(str(cell.value or "")) == "total"),
            None,
        )
        if total_col is None:
            continue
        year_cell = next(cell for cell in row if cell.value == 2024)
        monthly = next(sheet.iter_rows(min_row=year_cell.row + 1, max_row=year_cell.row + 1))
        months_by_name = {normal(name) for name in MONTHS}
        month_cols = [
            index
            for index, cell in enumerate(row)
            if normal(str(cell.value or "")) in months_by_name
        ]
        assert len(month_cols) == 12 and max(month_cols) < total_col
        months = [monthly[index].value for index in month_cols]
        assert len(months) == 12 and all(isinstance(item, int | float) for item in months)
        total = monthly[total_col].value
        assert isinstance(total, int | float)
        assert math.isclose(sum(months), total, rel_tol=1e-9, abs_tol=1e-9)
        sourced = value(
            data.dataset, data.factors, Metric("carrier", (Carrier.natural_gas,)), 2024
        )[0]
        assert sourced is not None and math.isclose(sourced, total, rel_tol=1e-9, abs_tol=1e-9)
        return
    pytest.fail("labelled 2024 gas TOTAL not found")


def _compare_annual_tables(final: Path, data: PieeData) -> None:
    tables = [
        table
        for table in Document(final).tables
        if len(table.rows) >= 3
        and len(table.columns) >= 5
        and normal(table.cell(0, 0).text) == "anul"
        and normal(table.cell(0, 1).text) == "u m"
    ]
    carriers = (
        Carrier.electricity_grid,
        Carrier.natural_gas,
        Carrier.electricity_cogen,
        Carrier.diesel,
        Carrier.coke,
    )
    assert len(tables) == len(carriers)
    slips = {
        (Carrier.natural_gas, 2024, 1): "S18",
        (Carrier.diesel, 2024, 1): "S6",
        (Carrier.diesel, 2024, 2): "S6",
        (Carrier.coke, 2024, 1): "S7",
    }
    for table, carrier in zip(tables, carriers, strict=True):
        for column, year in enumerate(data.dataset.years, 2):
            raw = value(data.dataset, data.factors, Metric("carrier", (carrier,)), year)[0]
            equivalent = tep(data.dataset, data.factors, carrier, year).value
            assert raw is not None and equivalent is not None
            for row, sourced in ((1, raw), (2, equivalent)):
                authored = _number(table.cell(row, column).text)
                same = round(authored, 2) == round(sourced, 2)
                slip = slips.get((carrier, year, row))
                assert same == (slip is None), (carrier, year, row, slip)
    _gas_total_is_sourced(data)


def _compare_centralizer(actual: Path, final: Path, data: PieeData) -> None:
    authored = next(
        table
        for table in Document(final).tables
        if len(table.rows) == 7
        and {normal(row.cells[0].text) for row in table.rows} >= {"cocs", "gaze naturale", "total"}
    )
    produced = next(
        table
        for table in Document(actual).tables
        if len(table.rows) == 4
        and {normal(cell.text) for cell in table.rows[0].cells} >= {"cocs", "gaz natural", "total"}
    )
    produced_cols = {normal(cell.text): index for index, cell in enumerate(produced.rows[0].cells)}
    authored_rows = {normal(row.cells[0].text): row for row in authored.rows}
    carriers = (
        ("energie electrica din sen", "energie electrica", Carrier.electricity_grid),
        ("gaz natural", "gaze naturale", Carrier.natural_gas),
        ("carburant", "motorina", Carrier.diesel),
        ("cocs", "cocs", Carrier.coke),
    )
    for year_index, year in enumerate(data.dataset.years, 1):
        assert int(produced.cell(year_index, 0).text) == year
        for produced_label, authored_label, carrier in carriers:
            source = tep(data.dataset, data.factors, carrier, year).value
            assert source is not None
            column = next(
                index for label, index in produced_cols.items() if label.startswith(produced_label)
            )
            output = _number(produced.cell(year_index, column).text)
            assert round(output, 2) == round(source, 2)
            reference = _number(authored_rows[authored_label].cells[year_index].text)
            slip = carrier == Carrier.electricity_grid and year == 2024  # S2
            assert (round(reference, 2) == round(source, 2)) != slip
        total = value(data.dataset, data.factors, Metric("tep_total"), year)[0]
        assert total is not None
        output = _number(produced.cell(year_index, produced_cols["total"]).text)
        reference = _number(authored_rows["total"].cells[year_index].text)
        assert round(output, 2) == round(total, 2)
        assert (round(reference, 2) == round(total, 2)) != (year == 2024)  # S3


def _display(value: float, decimals: int) -> Decimal:
    quantum = Decimal(1).scaleb(-decimals)
    return Decimal(f"{value:.15g}").quantize(quantum, rounding=ROUND_HALF_UP)


def _compare_numbered_tables(actual: Path, final: Path, data: PieeData) -> None:
    produced, authored = Document(actual).tables, Document(final).tables
    assert len(produced) == 18 and len(authored) == 15
    compare_measure_tables(produced, authored, data)
    for half in range(2):
        for row in range(1, 4):
            for column in range(7):
                left = produced[half].cell(row, column).text
                right = authored[half + 1].cell(row, column).text
                if column == 0:
                    assert left.strip() == right.strip() == str(data.dataset.years[row - 1])
                else:
                    assert _number(left) == _number(right)  # T1
    series = (
        Carrier.electricity_grid,
        Carrier.natural_gas,
        Carrier.electricity_cogen,
        Carrier.diesel,
        Carrier.coke,
    )
    production = next(iter(data.dataset.production.values()))
    for group, by_year in enumerate((production, *(data.dataset.carriers[c] for c in series))):
        for half in range(2):
            table = produced[group * 2 + half]
            for row, year in enumerate(data.dataset.years, 1):
                assert int(table.cell(row, 0).text) == year
                for column in range(1, 7):
                    month = half * 6 + column
                    source = by_year[year].months[month].value
                    assert source is not None
                    output = Decimal(str(_number(table.cell(row, column).text)))
                    assert output == _display(source, 2), (group, year, month)
    for row, year in enumerate(data.dataset.years, 1):
        assert int(produced[15].cell(row, 0).text) == year
        assert int(authored[11].cell(row, 0).text) == year
        assert data.prelucrare is not None
        filed = data.prelucrare.filed[f"co2.total.{year}"]
        assert isinstance(filed.value, int | float)
        output = _number(produced[15].cell(row, 1).text)
        reference = _number(authored[11].cell(row, 1).text)
        assert Decimal(str(output)) == _display(float(filed.value), 0)
        assert (round(reference) == round(output)) == (year != 2023)  # S11
