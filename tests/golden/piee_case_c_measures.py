"""Case C measure tables against the delivered programme and filed Anexa."""

from __future__ import annotations

from collections import Counter
from decimal import ROUND_HALF_UP, Decimal

from docx.table import Table

from ema.energy_data.anexa_cells import Measure
from ema.energy_data.source import normal
from ema.piee.dataset import PieeData
from ema.piee.measures import calculated_payback

MEASURE_PINS: dict[str, str] = {
    "nr_crt_column": "her prior-program layout numbers the rows; the base layout does not",
    "total_row": "her prior-program layout adds TOTAL rows; the base layout does not",
    "split_by_term": (
        "her prior-program layout splits planned measures into short and medium/long term; "
        "the base layout keeps one table"
    ),
    "site_subheader": "her prior-program layout groups medium/long-term rows under a site row",
    "row_order": "Ema keeps the Anexa's row order",
    "auditor_typo": "her final has a malformed number where the Anexa has the value Ema prints",
    "final_not_in_anexa": (
        "her final's medium/long-term rows are not in the supplied Anexa; Ema prints the Anexa's"
    ),
    "audit_report_source": (
        "her audit table carries the audit report's own units and values; "
        "Ema prints the Anexa's Audit energetic values"
    ),
}


def _number(text: str) -> Decimal:
    return Decimal(text.strip().replace(" ", "").replace(".", "").replace(",", "."))


def _source_number(value: float | int) -> Decimal:
    return Decimal(f"{value:.15g}").quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _measure_number(measure: Measure, key: str) -> Decimal:
    value = measure.values[key].value
    assert isinstance(value, int | float)
    return _source_number(value)


def _solution(
    table: Table, row: int, columns: tuple[int, ...]
) -> tuple[str, Decimal, Decimal, Decimal, Decimal, Decimal]:
    cells = table.rows[row].cells
    description, year, payback, cost, mwh, equivalent = (cells[col].text for col in columns)
    return (
        normal(description),
        _number(year),
        _number(payback),
        _number(cost),
        _number(mwh),
        _number(equivalent),
    )


def _anexa_solution(measure: Measure) -> tuple[str, Decimal, Decimal, Decimal, Decimal, Decimal]:
    year = measure.commissioning_year
    assert year is not None and isinstance(year.value, int | float)
    payback = measure.values.get("payback_years")
    calculated = calculated_payback(measure) if payback is None else None
    raw_payback = payback.value if payback is not None else calculated
    assert isinstance(raw_payback, int | float)
    return (
        normal(str(measure.description.value)),
        _source_number(year.value),
        _source_number(raw_payback),
        _measure_number(measure, "investment_thousand_lei"),
        _measure_number(measure, "saving_mwh"),
        _measure_number(measure, "saving_tep"),
    )


def compare_measure_tables(produced: list[Table], authored: list[Table], data: PieeData) -> None:
    used: set[str] = set()

    def pin(name: str, condition: bool) -> None:
        assert condition, f"{name}: {MEASURE_PINS[name]}"
        used.add(name)

    existing, final_existing = produced[14], authored[10]
    pin("nr_crt_column", len(existing.columns) == 6 and len(final_existing.columns) == 7)
    pin(
        "nr_crt_column",
        [row.cells[0].text for row in final_existing.rows[2:14]]
        == [str(index) for index in range(1, 13)],
    )
    pin("total_row", normal(final_existing.rows[-1].cells[1].text) == "total")
    assert len(existing.rows) == 14 and len(final_existing.rows) == 15
    actual_existing = [_solution(existing, row, (0, 1, 2, 3, 4, 5)) for row in range(2, 14)]
    final_rows = [_solution(final_existing, row, (1, 2, 3, 4, 5, 6)) for row in range(2, 14)]
    assert Counter(actual_existing) == Counter(final_rows)
    pin(
        "row_order",
        actual_existing == [_anexa_solution(row) for row in data.anexa.existing_measures],
    )

    planned, short, long = produced[16], authored[12], authored[13]
    pin(
        "split_by_term",
        len(planned.columns) == 6 and len(short.columns) == 6 and len(long.columns) == 12,
    )
    pin("site_subheader", len({normal(cell.text) for cell in long.rows[2].cells}) == 1)
    pin(
        "total_row",
        normal(short.rows[-1].cells[0].text) == "total"
        and normal(long.rows[-1].cells[0].text) == "total",
    )
    assert len(planned.rows) == 11 and len(short.rows) == 10 and len(long.rows) == 6
    actual_planned = [_solution(planned, row, (0, 1, 2, 3, 4, 5)) for row in range(2, 11)]
    assert len(actual_planned) == len(data.anexa.planned_measures) == 9
    assert all(description != "xxx" for description, *_ in actual_planned)
    pin(
        "row_order", actual_planned == [_anexa_solution(row) for row in data.anexa.planned_measures]
    )
    for index in range(7):
        output = actual_planned[index]
        if index == 4:
            assert output[:4] == _solution(short, index + 2, (0, 1, 2, 3, 4, 5))[:4]
            assert output[5] == _number(short.cell(index + 2, 5).text)
            pin(
                "auditor_typo",
                short.cell(index + 2, 4).text == "454.,5"
                and output[4] == _measure_number(data.anexa.planned_measures[index], "saving_mwh"),
            )
        else:
            assert output == _solution(short, index + 2, (0, 1, 2, 3, 4, 5))
    final_long = [_solution(long, row, (0, 2, 4, 6, 8, 10)) for row in (3, 4)]
    pin(
        "final_not_in_anexa",
        all(
            output != delivered
            for output, delivered in zip(actual_planned[7:], final_long, strict=True)
        ),
    )
    pin(
        "final_not_in_anexa",
        actual_planned[7:] == [_anexa_solution(row) for row in data.anexa.planned_measures[7:]],
    )

    audit, final_audit = produced[13], authored[9]
    assert len(audit.rows) == 5 and len(final_audit.rows) == 6
    assert [normal(audit.cell(row, 0).text) for row in range(2, 5)] == [
        normal(final_audit.cell(row, 0).text) for row in range(2, 5)
    ]
    pin("total_row", normal(final_audit.rows[-1].cells[0].text) == "total")
    pin(
        "audit_report_source",
        normal(final_audit.cell(1, 1).text) == "eur"
        and _number(final_audit.cell(4, 4).text)
        != _measure_number(data.anexa.audit_measures[2], "saving_tep"),
    )
    for row, source in enumerate(data.anexa.audit_measures, 2):
        assert normal(audit.cell(row, 0).text) == normal(str(source.description.value))
        assert _number(audit.cell(row, 1).text) == _measure_number(source, "saving_tep")
        # python-docx grid columns repeat merged cells: 4 is investment, 6 is payback.
        assert _number(audit.cell(row, 4).text) == _measure_number(
            source, "investment_thousand_lei"
        )
        payback = source.values.get("payback_years")
        raw = payback.value if payback is not None else calculated_payback(source)
        assert isinstance(raw, int | float)
        assert _number(audit.cell(row, 6).text) == _source_number(raw)
    assert used == set(MEASURE_PINS)
