"""audit_case_a input-level acceptance: deterministic Read and chapter-four package."""

from __future__ import annotations

import re
import shutil
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from tests.conftest import artifacts_path
from tests.golden.cases import case_path
from tests.workspace_jobs import create_job

from ema.audit.base import build_base
from ema.audit.base_package import package_issues
from ema.audit.catalogue import CATALOGUE
from ema.audit.chapter_four import factor_notes, render_chapter_four
from ema.audit.chapter_four_blocks import EMISSIONS_MISSING, emission_carriers
from ema.audit.headings import map_headings
from ema.audit.read import read_dossier
from ema.audit.sections import get_status
from ema.audit.totals_review import record_totals_review, totals_issues
from ema.consumption_analysis.analysis import Metric, value
from ema.core.office.missing_text import TABLE_MISSING_TEXT
from ema.core.review.models import Cell, Evidence, Field
from ema.core.workspace import Workspace
from ema.energy_data.carriers import WATER_CARRIERS, Carrier
from ema.energy_data.factors import AUDIT_FACTORS_2026, FACTORS_2026

from .test_s10b_audit_base import _audit_case_a_plan, _identity, _references

pytestmark = pytest.mark.golden


def _number(text: str) -> Decimal:
    match = re.search(r"[+-]?\d[\d. ]*,\d+", text)
    assert match is not None
    return Decimal(match.group().replace(" ", "").replace(".", "").replace(",", "."))


def _calculated(read: object) -> dict[str, set[Decimal]]:
    dataset = read.dataset  # type: ignore[attr-defined]
    result: dict[str, set[Decimal]] = {}
    emitting = emission_carriers(dataset)
    metrics = [Metric("tep_total"), Metric("intensity"), Metric("co2", emitting)]
    for carrier in dataset.carriers:
        metrics.extend([Metric("carrier", (carrier,)), Metric("tep", (carrier,))])
        metrics.append(Metric("co2", (carrier,)))
        for product in dataset.production:
            metrics.append(Metric("specific", (carrier,), product=product))
            if carrier in WATER_CARRIERS:
                metrics.append(Metric("water_specific", (carrier,), product=product))
    for product in dataset.production:
        metrics.append(Metric("production", product=product))
        metrics.append(Metric("specific", product=product))
    for metric in metrics:
        for year in dataset.years:
            for month in (
                (None, *range(1, 13))
                if metric.kind in {"carrier", "production", "tep", "tep_total"}
                else (None,)
            ):
                candidate = Metric(metric.kind, metric.carriers, metric.product, month)
                number, fact = value(dataset, AUDIT_FACTORS_2026, candidate, year, filed=False)
                if number is not None and fact:
                    result.setdefault(fact, set()).add(Decimal(str(number)))
    return result


def _printed_totals(
    document: Document, mapped: list[Any]
) -> dict[str, dict[int, list[Decimal | None]]]:
    """The "Total anual" lines of every ch. 4 section, as printed; None where it says missing."""
    ends = [item.heading.index for item in mapped] + [len(document.paragraphs)]
    result: dict[str, dict[int, list[Decimal | None]]] = {}
    for item, end in zip(mapped, ends[1:], strict=True):
        if not item.section_id.startswith("ch4."):
            continue
        for paragraph in document.paragraphs[item.heading.index + 1 : end]:
            line = re.match(r"^Total anual (\d{4}): (.*)\.$", paragraph.text)
            if line:
                missing = "date indisponibile" in line.group(2)
                result.setdefault(item.section_id, {}).setdefault(int(line.group(1)), []).append(
                    None if missing else _number(line.group(2))
                )
    return result


def _emission_rows(document: Document) -> dict[str, list[str]]:
    heading = next(p for p in document.paragraphs if p.text == "ANALIZA IMPACTULUI DE MEDIU")
    node = heading._p.getnext()
    while node is not None and node.tag != qn("w:tbl"):
        node = node.getnext()
    assert node is not None
    table = Table(node, document)
    return {row.cells[0].text: [cell.text for cell in row.cells[1:]] for row in table.rows[1:]}


COMPONENTS = ("ch4.echiv_electric", "ch4.echiv_pv", "ch4.echiv_gaz", "ch4.echiv_carburant")


def _assert_total_follows_components(
    document: Document, mapped: list[Any], years: tuple[int, ...], *, complete: bool
) -> None:
    printed = _printed_totals(document, mapped)
    for year in years:
        parts = [part for section in COMPONENTS for part in printed.get(section, {}).get(year, [])]
        (total,) = printed["ch4.echiv_total"][year]
        assert len(parts) >= 3
        if all(part is not None for part in parts):
            assert total is not None
            assert abs(total - sum(parts, Decimal(0))) <= Decimal("0.01") * len(parts)
        else:
            assert total is None  # one carrier without a reading leaves the total missing
        assert (total is not None) == complete


def _without_gpl(dataset: Any) -> Any:
    """The dossier as if the client had no GPL: every other carrier is read and has a factor."""
    carriers = {k: v for k, v in dataset.carriers.items() if k != Carrier.lpg}
    filed = {k: v for k, v in dataset.filed_indicators.items() if k != "tep.lpg"}
    return replace(dataset, carriers=carriers, filed_indicators=filed)


def test_audit_co2_factors_are_the_ones_her_base_states(reference_library: Path) -> None:
    base_source, _ = _references(reference_library)
    text = " ".join(
        "".join(node.text or "" for node in note.iter(qn("w:t")))
        for note in factor_notes(base_source)
    )

    def stated(pattern: str) -> float:
        match = re.search(pattern, text)
        assert match is not None, pattern
        return float(match.group(1).replace(",", "."))

    by_carrier = {factor.carrier: factor.per_unit for factor in AUDIT_FACTORS_2026.co2}
    assert by_carrier[Carrier.electricity_grid] == stated(r"utilizat este ([\d,]+) tone CO2/MWh")
    assert by_carrier[Carrier.natural_gas] == stated(r"conversie de ([\d,]+) kg CO2/kWh")
    for carrier, name in ((Carrier.diesel, "motorină"), (Carrier.petrol, "benzină")):
        per_litre = stated(rf"{name} s-a utilizat factorul de emisie de ([\d,]+) kg CO2e/l")
        density = stated(rf"cantității de {name} s-a utilizat o densitate de ([\d,]+) kg/litru")
        assert by_carrier[carrier] == pytest.approx(per_litre / density)


def test_audit_case_a_read_to_chapter_four(  # noqa: PLR0915
    reference_library: Path, tmp_path: Path
) -> None:
    received = reference_library / case_path("audit-case-a", "received")
    necesar = next(received.glob("*Necesar info*.xls"))
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "audit-case-a", 2026)
    read = read_dossier(ws, job, necesar)
    assert WATER_CARRIERS.intersection(read.dataset.carriers)
    assert Carrier.electricity_pv in read.dataset.carriers
    assert "PV sections in audits taken from the PIEE pattern; confirm" in read.issues
    assert get_status(ws, job, "ch4.electricitate_pv").applicability is True
    assert get_status(ws, job, "ch4.electricitate_pv").status.value == "ready"
    assert any(field.key == "audit.tep_class" and field.derivation for field in read.fields)
    with ws.connect() as db:
        for field in read.fields:
            if not field.key.startswith("audit.employees."):
                continue
            assert field.evidence
            row = db.execute(
                "SELECT data FROM evidence WHERE id=?", (field.evidence[0],)
            ).fetchone()
            evidence = Evidence.model_validate_json(row["data"])
            assert isinstance(evidence.locator, Cell)
            assert evidence.file_sha

    base_source, prototype = _references(reference_library)
    identity = _identity(base_source)
    base = tmp_path / "base.docx"
    output = tmp_path / "filled.docx"
    build_base(
        _audit_case_a_plan(reference_library),
        base_document=base_source,
        measurement_prototype=prototype,
        output=base,
        base_identity=identity,
    )
    digest = artifacts_path("s19", "digest-changes.txt")
    assert digest.relative_to(artifacts_path()).as_posix() == "s19/digest-changes.txt"
    digest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(base.parent / "digest-changes.txt", digest)
    report, skipped = render_chapter_four(
        base,
        output,
        read.dataset,
        identity,
        chart_source=base_source,
        client="Atelier Exemplu SRL",
    )
    assert "ch4.electricitate_pv:electricity_pv" in skipped
    assert not package_issues(output, identity)
    assert report.values
    assert all(use.fact for use in report.values)

    # D3 presents tep/kg as tep/t and intensity as tep/mil lei; facts retain source units.
    assert any(use.scale != 1 for use in report.values)
    expected = _calculated(read)
    for use in report.values:
        if "date indisponibile" in use.text or use.text in {TABLE_MISSING_TEXT, EMISSIONS_MISSING}:
            continue
        assert use.fact is not None
        if not any(
            abs(_number(use.text) - item * Decimal(str(use.scale))) <= Decimal("0.01")
            for item in expected.get(use.fact, ())
        ):
            pytest.fail(
                f"rendered number differs from its input or calculation: {use.fact} {use.text!r}"
            )

    mapped = map_headings(output, "audit-01").mapped
    wanted = [item.id for item in CATALOGUE if item.id.startswith("ch4.")]
    present = [item.section_id for item in mapped if item.section_id.startswith("ch4.")]
    assert present == wanted
    assert {"ch4.electricitate_pv", "ch4.echiv_pv", "ch4.specific_pv"} <= set(present)
    document = Document(str(output))
    toc = {
        link.get(qn("w:anchor"))
        for link in document.element.body.iter(qn("w:hyperlink"))
        if (link.get(qn("w:anchor")) or "").startswith("_Toc")
    }
    for item in mapped:
        if item.section_id not in {"ch4.electricitate_pv", "ch4.echiv_pv", "ch4.specific_pv"}:
            continue
        heading = document.paragraphs[item.heading.index]
        assert any(
            bookmark.get(qn("w:name")) in toc for bookmark in heading._p.iter(qn("w:bookmarkStart"))
        )
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "fotovoltaică" in text and "apă" in text

    # Decision 1: the printed total is the sum of the printed components. Her GPL has no reading,
    # so the total stays missing and the review blocks the final; without GPL it adds up.
    years = read.dataset.years
    _assert_total_follows_components(document, mapped, years, complete=False)
    with ws.connect() as db:
        facts = {
            str(row["key"]): Field.model_validate_json(row["data"])
            for row in db.execute("SELECT key,data FROM fields WHERE job_id=?", (job,))
        }
    assert {issue.code for issue in totals_issues(facts)} == {"carrier_incomplete"}
    assert {k for k in facts if k.startswith("carrier.lpg.") and len(k.split(".")) == 3} == {
        f"carrier.lpg.{year}" for year in years
    }
    completed = _without_gpl(read.dataset)
    record_totals_review(ws, job, completed, FACTORS_2026, "0" * 64)
    with ws.connect() as db:
        conflicts = [
            field.key
            for field in (
                Field.model_validate_json(row["data"])
                for row in db.execute("SELECT data FROM fields WHERE job_id=?", (job,))
            )
            if field.key.startswith("tep_total.") and field.confidence == "conflict"
        ]
    assert conflicts, "a filed total that differs from the components is a review conflict"

    # Decision 3: one row per purchased carrier, a red n.d. where a reading is missing, no PV row.
    rows = _emission_rows(document)
    assert "Total" in rows and "GPL" in rows
    assert rows["GPL"] == ["n.d."] * len(years) == rows["Total"]
    assert not any("fotovoltaic" in label.lower() for label in rows)
    assert all(
        "n.d." not in cells[0] for label, cells in rows.items() if label not in {"GPL", "Total"}
    )
    assert next(iter(rows)) == "Energie electrică din SEN"

    output_complete = tmp_path / "complete.docx"
    render_chapter_four(
        base,
        output_complete,
        completed,
        identity,
        chart_source=base_source,
        client="Atelier Exemplu SRL",
    )
    document_complete = Document(str(output_complete))
    mapped_complete = map_headings(output_complete, "audit-01").mapped
    _assert_total_follows_components(document_complete, mapped_complete, years, complete=True)
    complete_rows = _emission_rows(document_complete)
    assert all("n.d." not in cells[0] for cells in complete_rows.values())
    for index in range(len(years)):
        parts = [
            _number(cells[index]) for label, cells in complete_rows.items() if label != "Total"
        ]
        assert abs(_number(complete_rows["Total"][index]) - sum(parts, Decimal(0))) <= Decimal(
            "0.02"
        )
