"""CLIENT-A1 input-level acceptance: deterministic Read and chapter-four package."""

from __future__ import annotations

import re
import shutil
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from tests.conftest import artifacts_path
from tests.workspace_jobs import create_job

from ema.audit.base import build_base
from ema.audit.base_package import package_issues
from ema.audit.catalogue import CATALOGUE
from ema.audit.chapter_four import render_chapter_four
from ema.audit.headings import map_headings
from ema.audit.read import read_dossier
from ema.audit.sections import get_status
from ema.consumption_analysis.analysis import Metric, value
from ema.core.office.missing_text import TABLE_MISSING_TEXT
from ema.core.review.models import Cell, Evidence
from ema.core.workspace import Workspace
from ema.energy_data.carriers import WATER_CARRIERS, Carrier
from ema.energy_data.factors import FACTORS_2026

from .test_s10b_audit_base import _identity, _CLIENT-A1_plan, _references

pytestmark = pytest.mark.golden


def _number(text: str) -> Decimal:
    match = re.search(r"[+-]?\d[\d. ]*,\d+", text)
    assert match is not None
    return Decimal(match.group().replace(" ", "").replace(".", "").replace(",", "."))


def _calculated(read: object) -> dict[str, set[Decimal]]:
    dataset = read.dataset  # type: ignore[attr-defined]
    result: dict[str, set[Decimal]] = {}
    metrics = [Metric("tep_total"), Metric("intensity"), Metric("co2")]
    for carrier in dataset.carriers:
        metrics.extend([Metric("carrier", (carrier,)), Metric("tep", (carrier,))])
        for product in dataset.production:
            metrics.append(Metric("specific", (carrier,), product=product))
            if carrier in WATER_CARRIERS:
                metrics.append(Metric("water_specific", (carrier,), product=product))
    for product in dataset.production:
        metrics.append(Metric("production", product=product))
        metrics.append(Metric("specific", product=product))
    for metric in metrics:
        for year in dataset.years:
            for month in (
                (None, *range(1, 13))
                if metric.kind in {"carrier", "production", "tep", "tep_total"}
                else (None,)
            ):
                candidate = Metric(metric.kind, metric.carriers, metric.product, month)
                number, fact = value(dataset, FACTORS_2026, candidate, year)
                if number is not None and fact:
                    result.setdefault(fact, set()).add(Decimal(str(number)))
    return result


def test_CLIENT-A1_read_to_chapter_four(  # noqa: PLR0915
    reference_library: Path, tmp_path: Path
) -> None:
    received = reference_library / "audit/cases/audit-case-a/received"
    necesar = next(received.glob("*Necesar info*.xls"))
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "CLIENT-A1", 2026)
    read = read_dossier(ws, job, necesar)
    assert WATER_CARRIERS.intersection(read.dataset.carriers)
    assert Carrier.electricity_pv in read.dataset.carriers
    assert "PV sections in audits taken from the PIEE pattern; confirm" in read.issues
    assert get_status(ws, job, "ch4.electricitate_pv").applicability is True
    assert get_status(ws, job, "ch4.electricitate_pv").status.value == "ready"
    assert any(field.key == "audit.tep_class" and field.derivation for field in read.fields)
    with ws.connect() as db:
        for field in read.fields:
            if not field.key.startswith("audit.employees."):
                continue
            assert field.evidence
            row = db.execute(
                "SELECT data FROM evidence WHERE id=?", (field.evidence[0],)
            ).fetchone()
            evidence = Evidence.model_validate_json(row["data"])
            assert isinstance(evidence.locator, Cell)
            assert evidence.file_sha

    base_source, prototype = _references(reference_library)
    identity = _identity(base_source)
    base = tmp_path / "base.docx"
    output = tmp_path / "filled.docx"
    build_base(
        _CLIENT-A1_plan(reference_library),
        base_document=base_source,
        measurement_prototype=prototype,
        output=base,
        base_identity=identity,
    )
    digest = artifacts_path("s19", "digest-changes.txt")
    assert digest.relative_to(artifacts_path()).as_posix() == "s19/digest-changes.txt"
    digest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(base.parent / "digest-changes.txt", digest)
    report, skipped = render_chapter_four(
        base,
        output,
        read.dataset,
        identity,
        chart_source=base_source,
        client="Atelier Exemplu SRL",
    )
    assert "ch4.electricitate_pv:electricity_pv" in skipped
    assert not package_issues(output, identity)
    assert report.values
    assert all(use.fact for use in report.values)

    # D3 presents tep/kg as tep/t and intensity as tep/mil lei; facts retain source units.
    assert any(use.scale != 1 for use in report.values)
    expected = _calculated(read)
    for use in report.values:
        if "date indisponibile" in use.text or use.text == TABLE_MISSING_TEXT:
            continue
        assert use.fact is not None
        if not any(
            abs(_number(use.text) - item * Decimal(str(use.scale))) <= Decimal("0.01")
            for item in expected.get(use.fact, ())
        ):
            pytest.fail("rendered number differs from its located input or scaled S7 calculation")

    mapped = map_headings(output, "AUDIT-01").mapped
    wanted = [item.id for item in CATALOGUE if item.id.startswith("ch4.")]
    present = [item.section_id for item in mapped if item.section_id.startswith("ch4.")]
    assert present == wanted
    assert {"ch4.electricitate_pv", "ch4.echiv_pv", "ch4.specific_pv"} <= set(present)
    document = Document(str(output))
    toc = {
        link.get(qn("w:anchor"))
        for link in document.element.body.iter(qn("w:hyperlink"))
        if (link.get(qn("w:anchor")) or "").startswith("_Toc")
    }
    for item in mapped:
        if item.section_id not in {"ch4.electricitate_pv", "ch4.echiv_pv", "ch4.specific_pv"}:
            continue
        heading = document.paragraphs[item.heading.index]
        assert any(
            bookmark.get(qn("w:name")) in toc for bookmark in heading._p.iter(qn("w:bookmarkStart"))
        )
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "fotovoltaică" in text and "apă" in text
