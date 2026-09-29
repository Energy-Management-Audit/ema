"""S9a regression and delivered-workbook checks; no reference data enters git."""

from __future__ import annotations

import hashlib
import json
import os
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from tests.workspace_jobs import create_job

from conftest import artifacts_path
from ema.core.errors import EmaError
from ema.core.jobs import run_stage, status, subscribe
from ema.core.workspace import Workspace
from ema.invoices import (
    confirm_client,
    export,
    extract_batch,
)

_BASELINE = artifacts_path("s9a-baseline")
_TEXT_CASES = ("invoice-case-d", "CLIENT-I5", "invoice-case-a")
_OUTCOME_CASES = ("CLIENT-I2", "invoice-case-f")
_SNIPPET_EXCEPTION = (
    "invoice-case-a",
    "Factura energie consum februarie 2025.pdf",
    "Detalii preturi",
    21,
)
_SNIPPET_MISSING_INDEX = 12
_SNIPPET_LEGACY_WORD_COUNT = 13
_SNIPPET_MISSING_SHA256 = "72505985a1fb52d82e588e5472ddbf6a0e9bcf924cfcc95d6c6b4a30e791ceee"
_SOURCE_FIELDS = (
    "category",
    "description",
    "source_quantity",
    "source_unit",
    "source_unit_price",
    "normalized_quantity",
    "normalized_unit",
    "normalized_unit_price",
    "net_value",
)


def _reference() -> Path:
    location = os.environ.get("EMA_REFERENCE")
    assert location, "EMA_REFERENCE is required for this golden test"
    root = Path(location).expanduser()
    assert root.is_dir(), "EMA_REFERENCE does not exist"
    return root / "invoices/cases"


def _baseline(case: str) -> tuple[list[dict[str, object]], Path]:
    directory = _BASELINE / case
    outcome_file = directory / "outcomes.json"
    assert outcome_file.is_file(), f"missing immutable S9a baseline: {case}/outcomes.json"
    outcomes = json.loads(outcome_file.read_text(encoding="utf-8"))
    workbook = directory / "legacy.xlsx"
    if case in _TEXT_CASES:
        assert workbook.is_file(), f"missing immutable S9a baseline: {case}/legacy.xlsx"
    return outcomes, workbook


def _source_paths(case: str, expected: list[dict[str, object]]) -> list[Path]:
    directory = _reference() / case / "received"
    found = {path.name: path for path in directory.rglob("*") if path.suffix.lower() == ".pdf"}
    names = {str(row["source_path"]) for row in expected}
    missing = names - found.keys()
    assert not missing, f"missing reference PDFs in {case}: {sorted(missing)}"
    return [found[name] for name in sorted(names)]


def _run_case(
    tmp_path: Path,
    case: str,
    expected: list[dict[str, object]],
    workspace: Workspace | None = None,
) -> tuple[Workspace, str, list[dict[str, object]]]:
    ws = workspace or Workspace(tmp_path / case)
    job = create_job(ws, "invoices", case, None)
    for index, source in enumerate(_source_paths(case, expected), 1):
        sha = ws.add_file(case, source)
        ws.set_slot(job, f"invoices/{index:04d}", sha, origin=source.name)
    run_id = run_stage(ws, job, "invoices", extract_batch)
    for _event in subscribe(ws, job):
        pass
    run = next(item for item in status(ws, job).runs if item["id"] == run_id)
    assert run["state"] == "ready", f"{case}: invoice stage failed"
    assert run["publication"] == "current", f"{case}: invoice stage became stale"
    with ws.connect() as db:
        artifact = ws.artifact_dir(db, job, "invoices", run_id) / "outcomes.json"
    return ws, job, json.loads(artifact.read_text(encoding="utf-8"))


def _normalized(value: object) -> str:
    return " ".join(str(value).split())


def _one_extra_legacy_word(current: object, legacy: object) -> bool:
    new_words = _normalized(current).split()
    old_words = _normalized(legacy).split()
    return (
        len(old_words) == _SNIPPET_LEGACY_WORD_COUNT
        and old_words[:_SNIPPET_MISSING_INDEX] + old_words[_SNIPPET_MISSING_INDEX + 1 :]
        == new_words
        and hashlib.sha256(old_words[_SNIPPET_MISSING_INDEX].encode()).hexdigest()
        == _SNIPPET_MISSING_SHA256
    )


def _compare_outcomes(
    case: str, actual: list[dict[str, object]], expected: list[dict[str, object]]
) -> None:
    by_name = {str(row["source_path"]): row for row in expected}
    mismatches: list[str] = []
    for row in actual:
        name = str(row["source_path"])
        original = by_name[name]
        if row["status"] != original["status"]:
            mismatches.append(f"{name}:status")
        drafts = row["drafts"]
        reference_drafts = original["drafts"]
        if len(drafts) != len(reference_drafts):
            mismatches.append(f"{name}:draft_count")
        for index, (draft, prior) in enumerate(zip(drafts, reference_drafts, strict=False)):
            mismatches.extend(_draft_mismatches(name, index, draft, prior))
    print(f"{case}: {len(actual)} PDFs, outcomes {dict(Counter(row['status'] for row in actual))}")
    assert not mismatches, f"{case} mismatches: {mismatches}"


def _draft_mismatches(name: str, index: int, draft: dict, prior: dict) -> list[str]:
    mismatches: list[str] = []
    if draft["document_id"] != prior["document_id"]:
        mismatches.append(f"{name}:draft{index}.document_id")
    for field_name, field in draft["fields"].items():
        reference_field = prior["fields"].get(field_name)
        if reference_field is None or (field["value"], field["status"]) != (
            reference_field["value"],
            reference_field["status"],
        ):
            mismatches.append(f"{name}:draft{index}.{field_name}")
    if len(draft["price_details"]) != len(prior["price_details"]):
        mismatches.append(f"{name}:draft{index}.price_count")
    for detail_index, (detail, old) in enumerate(
        zip(draft["price_details"], prior["price_details"], strict=False)
    ):
        for key in _SOURCE_FIELDS:
            if detail[key] != old[key]:
                mismatches.append(f"{name}:draft{index}.price{detail_index}.{key}")
        if detail["evidence"]["page_number"] != old["evidence"]["page_number"]:
            mismatches.append(f"{name}:draft{index}.price{detail_index}.page")
    return mismatches


def _compare_workbooks(case: str, actual_path: Path, baseline_path: Path) -> None:
    actual = load_workbook(actual_path, data_only=False)
    baseline = load_workbook(baseline_path, data_only=False)
    assert actual.sheetnames == baseline.sheetnames
    mismatches: list[str] = []
    normalized_snippets = 0
    pinned_snippets = 0
    for sheet, reference in zip(actual, baseline, strict=True):
        assert (sheet.max_row, sheet.max_column) == (reference.max_row, reference.max_column)
        source_ordinals: Counter[str] = Counter()
        for row in sheet:
            source = str(row[0].value)
            if row[0].row > 1:
                source_ordinals[source] += 1
            for cell in row:
                old = reference[cell.coordinate]
                normalized, pinned, errors = _compare_cell(
                    case, source, source_ordinals[source], sheet.title, cell, old
                )
                normalized_snippets += normalized
                pinned_snippets += pinned
                mismatches.extend(errors)
    print(f"{case}: {normalized_snippets} source snippets equal after whitespace normalization")
    if case == _SNIPPET_EXCEPTION[0]:
        assert pinned_snippets == 1, "pinned ENGIE snippet exception was not exercised exactly once"
    assert not mismatches, f"{case} workbook mismatches: {mismatches}"


def _compare_cell(
    case: str, source: str, ordinal: int, sheet: str, cell, old
) -> tuple[int, int, list[str]]:
    errors: list[str] = []
    result = ""
    if sheet == "Detalii preturi" and cell.column == 17:
        result = _snippet_result((case, source, sheet, ordinal), cell.value, old.value)
        if result == "mismatch":
            errors.append(f"{sheet}:{cell.coordinate}:snippet")
    elif cell.value != old.value:
        errors.append(f"{sheet}:{cell.coordinate}:value")
    if cell.number_format != old.number_format:
        errors.append(f"{sheet}:{cell.coordinate}:number_format")
    return int(result == "normalized"), int(result == "pinned"), errors


def _snippet_result(key: tuple[str, str, str, int], current: object, legacy: object) -> str:
    if current == legacy:
        return "same"
    if _normalized(current) == _normalized(legacy):
        return "normalized"
    if key == _SNIPPET_EXCEPTION and _one_extra_legacy_word(current, legacy):
        return "pinned"
    return "mismatch"


def _formula_shape(value: object) -> object:
    if not isinstance(value, str) or not value.startswith("="):
        return value
    return re.sub(r"(\$?[A-Z]{1,3})\$?\d+", r"\1#", value)


def _row_key(row: tuple[object, ...]) -> tuple[str, ...]:
    return tuple(str(row[index]) for index in (0, 3, 4, 5))


def _compare_delivered(case: str, generated_path: Path) -> None:
    delivered_path = next((_reference() / case / "final").glob("*.xlsx"))
    generated = load_workbook(generated_path, data_only=False)["Centralizator"]
    delivered = load_workbook(delivered_path, data_only=False)["Centralizator"]
    reference_rows = {
        _row_key(tuple(row.value for row in cells)): (index, tuple(row.value for row in cells))
        for index, cells in enumerate(delivered.iter_rows(min_row=2), 2)
        if isinstance(cells[1].value, date | datetime) and cells[1].value.year == 2025
    }
    mismatches: list[str] = []
    matched: set[tuple[str, ...]] = set()
    for cells in generated.iter_rows(min_row=2):
        if not isinstance(cells[1].value, date | datetime) or cells[1].value.year != 2025:
            continue
        row = tuple(cell.value for cell in cells)
        key = _row_key(row)
        if key not in reference_rows:
            mismatches.append(f"{row[0]}:missing delivered row")
            continue
        matched.add(key)
        reference_index, old = reference_rows[key]
        for column, (value, prior) in enumerate(zip(row, old, strict=True), 1):
            if _formula_shape(value) != _formula_shape(prior):
                mismatches.append(f"{row[0]}:column {column}")
            if (
                cells[column - 1].number_format
                != delivered.cell(reference_index, column).number_format
            ):
                mismatches.append(f"{row[0]}:column {column} number_format")
    delivered_only = len(reference_rows) - len(matched)
    print(
        f"{case}: {len(matched)} delivered 2025 rows matched; {delivered_only} delivered-only rows"
    )
    assert not mismatches, f"{case} delivered mismatches: {mismatches}"


@pytest.mark.golden
@pytest.mark.parametrize("case", _TEXT_CASES)
def test_text_pdf_regression_and_delivered_rows(tmp_path: Path, case: str) -> None:
    expected, baseline_workbook = _baseline(case)
    ws, job, actual = _run_case(tmp_path, case, expected)
    assert len(actual) == len(expected)
    _compare_outcomes(case, actual, expected)
    with ws.connect() as db:
        destination = ws.job_path(db, job) / "outputs" / "Facturi.xlsx"
    with pytest.raises(EmaError) as blocked:
        export(ws, job, destination)
    assert blocked.value.code == "invoices_unconfirmed_client"
    confirm_client(ws, job)
    generated = export(ws, job, destination)
    _compare_workbooks(case, generated, baseline_workbook)
    if case in ("invoice-case-d", "CLIENT-I5"):
        _compare_delivered(case, generated)


@pytest.mark.golden
@pytest.mark.parametrize("case", _OUTCOME_CASES)
def test_supported_families_and_incompatible_documents(tmp_path: Path, case: str) -> None:
    expected, _ = _baseline(case)
    _, _, actual = _run_case(tmp_path, case, expected)
    by_name = {str(row["source_path"]): row for row in expected}
    mismatches: list[str] = []
    for row in actual:
        original = by_name[str(row["source_path"])]
        parser = original["metadata"]["parser_name"]
        if parser == "AliveInvoiceParser":
            if row["status"] not in {"exportable", "requires_review"}:
                mismatches.append(f"{row['source_path']}:status")
            continue
        if parser in {
            "EngieInvoiceParser",
            "MetNaturalGasInvoiceParser",
            "SeeExclusiveRefactoringParser",
            "EngieEInvoiceCompanionParser",
            "OmvPetromInvoiceParser",
        }:
            wanted = original["status"]
        else:
            wanted = "requires_review"
        if row["status"] != wanted:
            mismatches.append(f"{row['source_path']}:status")
    print(f"{case}: {len(actual)} PDFs, outcomes {dict(Counter(row['status'] for row in actual))}")
    assert not mismatches, f"{case} mismatches: {mismatches}"
