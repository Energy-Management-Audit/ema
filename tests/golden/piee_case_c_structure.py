"""Case C section order across the approved base and the delivered programme."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from ema.energy_data.source import normal

COMPOSED_HEADINGS = (
    "analiza consumului de energie electrica",
    "analiza consumului de gaz natural",
    "analiza consumului de energie electrica produsa prin cogenerare",
    "analiza consumului de carburanti",
    "analiza consumului de cocs",
    "analiza consumului de apa potabila",
    "analiza consumului de apa industriala",
    "analiza consumului de apa meteorica",
    "analiza consumului echivalent de energie",
    "analiza consumului specific echivalent de gaz natural",
    "consumul specific de cocs",
    "analiza consumului specific echivalent de energie totala",
    "analiza consumului specific echivalent de apa",
    "consumul specific de apa industriala",
    "consumul specific de apa meteorica",
    "intensitatea energetica",
    "analiza auditurilor energetice",
    "analiza activitatii de investitii",
    "impactul de mediu al utilizarii energiei",
)
DELIVERED_MARKERS = (
    "analiza consumului de energie electrica",
    "analiza consumului de gaze naturale",
    "analiza consumului de energie electrica produsa prin cogenerare",
    "analiza consumului de carburanti",
    "analiza consumului de cocs",
    "analiza consumului de apa potabila",
    "analiza consumului de apa industriala",
    "analiza consumului de apa meteorica",
    "in tabelul numarul 7 se prezinta un centralizator al consumurilor echivalente de energie",
    "consumul specific de gaze naturale",
    "consumul specific de cocs",
    "consumul specific total de energie",
    "consumul specific de apa potabila",
    "consumul specific de apa industriala",
    "consumul specific de apa meteorica",
    "intensitatea energetica",
    "analiza auditurilor energetice",
    "analiza activitatii de investitii",
    "impactul de mediu al utilizarii energiei",
)


def assert_section_sequence(composed: Path, final: Path) -> None:
    headings = [
        normal(paragraph.text)
        for paragraph in Document(composed).paragraphs
        if paragraph.style.name == "Heading 3"
    ]
    assert tuple(headings) == COMPOSED_HEADINGS
    # Her equivalent-energy group has a table introduction, not a separate heading.
    authored = [normal(paragraph.text) for paragraph in Document(final).paragraphs]
    positions = [
        [index for index, text in enumerate(authored) if text == marker]
        for marker in DELIVERED_MARKERS
    ]
    assert all(len(matches) == 1 for matches in positions)
    assert [matches[0] for matches in positions] == sorted(matches[0] for matches in positions)
    # S13 retains the approved-base general heading; her rewritten heading differs.
    assert any("adresa sediului social" in normal(p.text) for p in Document(composed).paragraphs)
    assert any(
        normal(p.text).startswith("adresa punctele de lucru") for p in Document(final).paragraphs
    )
