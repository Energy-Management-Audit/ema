"""Level-1 chapter-six package check on the auditor's base and synthetic measures."""

from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from tests.conftest import artifacts_path

from ema.audit.base import build_base
from ema.audit.base_units import UnitPlan, heading_spans_document
from ema.audit.chapter_six import ChapterSixPlan, render_chapter_six
from ema.audit.inventory import inventory
from ema.audit.measures import run_measures
from ema.audit.measures_form import write_measures_form
from ema.core.jobs import create_job
from ema.core.workspace import Workspace

pytestmark = pytest.mark.golden


def _synthetic_form(path: Path) -> Path:
    write_measures_form(path)
    book = load_workbook(path)
    sheet = book["Măsuri propuse"]
    sheet.append(
        [1, "Lighting retrofit", "Less electricity", "Energie electrică", 100, "MWh", 40, 10, None]
    )
    sheet.append(
        [2, "Gas recovery", "Less gas", "Gaze naturale", 10, "MWh", 20, None, "Estimate pending"]
    )
    sheet.append([3, "Fleet change", "Less diesel", "Motorină", 1, "t", 50, 25, None])
    book.save(path)
    return path


@pytest.mark.parametrize("with_ch5", [False, True])
def test_three_measures_in_base(reference_library: Path, tmp_path: Path, with_ch5: bool) -> None:
    root = reference_library / "audit" / "finished-audits"
    AUDIT-01 = next(root.glob("*AUDIT-01*.docx"))
    measurements = next(root.glob("*AUDIT-04*.docx"))
    identity = (AUDIT-01.stem.split("AUDIT ENERGETIC ", 1)[1].rsplit(" - ", 1)[0],)
    output_dir = artifacts_path("s15b", "with-ch5" if with_ch5 else "without-ch5")
    output_dir.mkdir(parents=True, exist_ok=True)
    base = build_base(
        UnitPlan(
            "Synthetic factory",
            0,
            frozenset({"electricity", "gas", "fuel"}),
            1 if with_ch5 else 0,
            with_ch5,
            0,
            1,
        ),
        base_document=AUDIT-01,
        measurement_prototype=measurements,
        output=output_dir / "base.docx",
        base_identity=identity,
    )
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    form = _synthetic_form(tmp_path / "measures.xlsx")
    ws.set_slot(job, "measures", ws.add_file("synthetic", form))
    result = run_measures(ws, job)
    plan = ChapterSixPlan.model_validate_json(result.plan_path.read_text("utf-8"))
    output = output_dir / "chapter-six.docx"
    report = render_chapter_six(base, output, plan, identity)
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    reference_body = list(Document(AUDIT-01).element.body)
    reference_titles = [
        "".join(node.text or "" for node in reference_body[unit.element_range[0]].iter())
        for unit in inventory(AUDIT-01).measures
    ]
    chapter = 6 if with_ch5 else 5
    assert [number.number for number in report.numbers] == [
        f"{chapter}.1",
        f"{chapter}.2",
        f"{chapter}.3",
        f"{chapter}.4",
    ]
    assert "Lighting retrofit" in text and "Gas recovery" in text and "Fleet change" in text
    assert "[de completat]" in text
    assert len(document.tables) >= 4
    spans = heading_spans_document(document)
    first_measure = next(start for item, start, _ in spans if item.section_id == "ch6.measure")
    base_spans = heading_spans_document(Document(base))
    base_measure = next(start for item, start, _ in base_spans if item.section_id == "ch6.measure")
    rendered_style = list(document.element.body)[first_measure].pPr.xml
    prototype_style = list(Document(base).element.body)[base_measure].pPr.xml
    assert rendered_style == prototype_style
    assert "Costurile sunt estimate, iar la implementarea măsurii se recomandă" in text
    assert all(title not in text for title in reference_titles)
    assert result.payback_missing == 1
    assert all(identity[0] not in paragraph.text for paragraph in document.paragraphs)
    assert any("PAGEREF" in paragraph._p.xml for paragraph in document.paragraphs)
