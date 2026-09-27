"""Level 2 chain: indexed report equals S6 output, which matches the delivered workbooks."""

from __future__ import annotations

import math
import threading
import unicodedata
from collections import Counter
from contextlib import ExitStack
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ema.core.workspace import Workspace
from ema.energy_data.annex_index import import_annexes
from ema.reporting import generate, write_report
from ema.reporting.runs import get_run, preview, start_run

pytestmark = pytest.mark.golden


def _done(ws: Workspace, run_id: str) -> dict[str, object]:
    for worker in threading.enumerate():
        if worker.name.startswith("ema-reporting-"):
            worker.join(timeout=120)
            assert not worker.is_alive(), "reporting stage timed out"
    run = get_run(ws, run_id)
    assert run["state"] == "ready"
    return run


def _output(ws: Workspace, output_id: str) -> Path:
    with ws.connect() as db:
        row = db.execute("SELECT relative_path FROM outputs WHERE id=?", (output_id,)).fetchone()
    assert row is not None
    return ws.path(str(row["relative_path"]))


def _same(actual: object, expected: object) -> bool:
    if isinstance(actual, int | float) and isinstance(expected, int | float):
        return math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-9)
    if isinstance(actual, str) and isinstance(expected, str):
        return unicodedata.normalize("NFC", actual) == unicodedata.normalize("NFC", expected)
    return actual == expected


def _compare_workbooks(actual: Path, expected: Path) -> None:
    got = load_workbook(actual, data_only=True, read_only=True)
    want = load_workbook(expected, data_only=True, read_only=True)
    assert got.sheetnames == want.sheetnames
    for name in want.sheetnames:
        left = list(got[name].values)
        right = list(want[name].values)
        assert len(left) == len(right), name
        if name == "Exceptions":

            def normal(row: tuple[object, ...]) -> tuple[object, ...]:
                return tuple(
                    unicodedata.normalize("NFC", cell) if isinstance(cell, str) else cell
                    for cell in row
                )

            assert Counter(normal(row) for row in left[3:]) == Counter(
                normal(row) for row in right[3:]
            )
            continue
        for row_index, (left_row, right_row) in enumerate(zip(left, right, strict=True), 1):
            for col_index, (a, b) in enumerate(zip(left_row, right_row, strict=True), 1):
                assert _same(a, b), (name, row_index, col_index)


def test_imported_reporting_equals_s6_sources(reference_library: Path, tmp_path: Path) -> None:
    paths = list((reference_library / "piee/anexa-2-3-2025").glob("*.xls*"))
    assert len(paths) == 36
    ws = Workspace(tmp_path / "workspace")
    with ExitStack() as stack:
        imported = import_annexes(
            ws, [(path.name, stack.enter_context(path.open("rb"))) for path in paths]
        )
    assert len(imported.imported) == 36
    assert Counter(item.code for item in imported.ignored) == Counter()
    ids = sorted({item.client_id for item in imported.imported})
    for years in ([2025], [2023, 2024, 2025]):
        started = start_run(ws, years, ids)
        run = _done(ws, str(started["id"]))
        shown = preview(ws, str(run["id"]))
        original = generate(paths, tuple(years))
        expected = write_report(original, tmp_path / f"expected-{years[0]}.xlsx")
        _compare_workbooks(_output(ws, str(run["output_id"])), expected)
        if len(years) == 3:
            actual = load_workbook(_output(ws, str(run["output_id"])), data_only=True)
            delivered = load_workbook(
                next(
                    (reference_library / "energy-manager-reporting/cases/2023-2025/final").glob(
                        "*.xlsx"
                    )
                ),
                data_only=True,
            )
            for year in years:
                got_rows = list(actual[str(year)].values)
                expected_rows = list(delivered[str(year)].values)
                assert len(got_rows) == len(expected_rows)
                for got, reference in zip(got_rows, expected_rows, strict=True):
                    assert all(_same(a, b) for a, b in zip(got, reference, strict=True))
            got_control = list(actual["Control"].values)[4:]
            expected_control = list(delivered["Control"].values)[4:]
            assert len(got_control) == len(expected_control) == 36
            for got, reference in zip(got_control, expected_control, strict=True):
                for col in range(17):
                    if got[7] == "Date lunare!N3" and col in (7, 8, 10, 11, 16):
                        continue
                    assert _same(got[col], reference[col]), col
        assert shown["read"] == len(original.companies) == 36
        assert len(run["exceptions"]) == len(original.exceptions)
        assert Counter((item["code"], item["detail"]) for item in run["exceptions"]) == Counter(
            (item.severity, item.situation) for item in original.exceptions
        )
