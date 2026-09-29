"""Dataset-backed block plans for the two authored S7 audit goldens."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from tests.golden.cases import case_path
from tests.golden.s7_reader import Chapter, _text, read_chapter

from ema.audit.headings import map_headings
from ema.consumption_analysis.analysis import (
    ChartPlan,
    MatrixTablePlan,
    Metric,
    NumericSentencePlan,
    SectionPlan,
    TablePlan,
    TrendPlan,
    analyze,
)
from ema.consumption_analysis.phrases import phrase_bank
from ema.core.office.blocks import Block
from ema.energy_data.carriers import Carrier
from ema.energy_data.factors import Factor, FactorTable


def _source(name: str) -> Path:
    return case_path(name)


def build_blocks(name: str) -> tuple[dict[int, Block], Chapter]:  # noqa: C901, PLR0912, PLR0915
    case = read_chapter(_source(name), name.lower())
    factors = FactorTable(
        "reference",
        min(case.dataset.years),
        (
            Factor(Carrier.electricity_grid, "MWh", 0.086, "authored factor"),
            Factor(Carrier.natural_gas, "MWh", 0.086, "authored factor"),
        ),
        (),
    )
    generated: dict[int, Block] = {}
    chart_plans: dict[int, ChartPlan] = {}
    sections = (
        ("ch4.productie", Metric("production", product="turnover")),
        ("ch4.electricitate", Metric("carrier", (Carrier.electricity_grid,))),
        ("ch4.gaz", Metric("carrier", (Carrier.natural_gas,))),
    )
    for section_id, metric in sections:
        section = case.sections[section_id]
        for half, (body_index, table) in enumerate(section.tables):
            columns = tuple(
                Metric(metric.kind, metric.carriers, metric.product, half * 6 + month)
                for month in range(1, 7)
            )
            first_value = _text(
                table.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr")[
                    1
                ].findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc")[1]
            )
            item = TablePlan(
                f"item-{body_index}", columns, case.dataset.years, grouping="." in first_value
            )
            generated[body_index] = analyze(
                case.dataset, factors, (SectionPlan(section_id, "carrier", (item,)),)
            )[0].blocks[0]
        for chart_index, chart in enumerate(section.charts):
            monthly = len(section.charts) == 4 and chart_index < 3
            item = ChartPlan(
                f"item-{chart.body_index}",
                chart.part,
                chart.series.name,
                metric,
                (case.dataset.years[chart_index],) if monthly else case.dataset.years,
                tuple(chart.series.categories),
            )
            generated[chart.body_index] = analyze(
                case.dataset, factors, (SectionPlan(section_id, "carrier", (item,)),)
            )[0].blocks[0]
            chart_plans[chart.body_index] = item
    for section_id, carrier in (
        ("ch4.echiv_electric", Carrier.electricity_grid),
        ("ch4.echiv_gaz", Carrier.natural_gas),
    ):
        section = case.sections[section_id]
        for half, (body_index, table) in enumerate(section.tables):
            columns = tuple(
                Metric("tep", (carrier,), month=half * 6 + month) for month in range(1, 7)
            )
            first_value = _text(
                table.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr")[
                    1
                ].findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc")[1]
            )
            item = TablePlan(
                f"item-{body_index}", columns, case.dataset.years, grouping="." in first_value
            )
            generated[body_index] = analyze(
                case.dataset, factors, (SectionPlan(section_id, "equivalent", (item,)),)
            )[0].blocks[0]
        for chart in section.charts:
            item = ChartPlan(
                f"item-{chart.body_index}",
                chart.part,
                chart.series.name,
                Metric("tep_monthly_sum", (carrier,)),
                case.dataset.years,
                tuple(chart.series.categories),
            )
            generated[chart.body_index] = analyze(
                case.dataset, factors, (SectionPlan(section_id, "equivalent", (item,)),)
            )[0].blocks[0]
            chart_plans[chart.body_index] = item
    intensity = case.sections["ch4.intensitate"]
    intensity_index, intensity_table = intensity.tables[0]
    intensity_rows = intensity_table.findall(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr"
    )
    decimals = len(_text(intensity_rows[2][-1]).rsplit(",", 1)[-1])
    intensity_plan = TablePlan(
        f"item-{intensity_index}",
        (
            Metric("tep_total", decimals=2),
            Metric("filed", product="turnover_thousand", decimals=2, grouping=True),
            Metric("intensity", decimals=decimals),
        ),
        case.dataset.years,
        header_rows=2,
    )
    generated[intensity_index] = analyze(
        case.dataset, factors, (SectionPlan("ch4.intensitate", "intensity", (intensity_plan,)),)
    )[0].blocks[0]
    for chart in intensity.charts:
        item = ChartPlan(
            f"item-{chart.body_index}",
            chart.part,
            chart.series.name,
            Metric("filed", product="intensity_chart"),
            case.dataset.years,
            tuple(chart.series.categories),
        )
        generated[chart.body_index] = analyze(
            case.dataset, factors, (SectionPlan("ch4.intensitate", "intensity", (item,)),)
        )[0].blocks[0]
        chart_plans[chart.body_index] = item
    environment = case.sections["ch4.mediu"]
    co2_index, co2_table = environment.tables[0]
    co2_rows = co2_table.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr")
    if len(co2_rows) == 4:
        co2_plan = TablePlan(f"item-{co2_index}", (Metric("co2"),), case.dataset.years, decimals=0)
    else:
        row_metrics = (
            Metric("co2", (Carrier.electricity_grid,)),
            Metric("co2", (Carrier.natural_gas,)),
            Metric("filed", product="co2.fuel"),
            Metric("co2"),
        )
        co2_plan = MatrixTablePlan(
            f"item-{co2_index}",
            tuple(
                (
                    _text(
                        row.findall(
                            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc"
                        )[0]
                    ),
                    (metric,) * len(case.dataset.years),
                )
                for row, metric in zip(co2_rows[2:], row_metrics, strict=True)
            ),
            case.dataset.years,
            decimals=0,
            header_rows=2,
        )
    generated[co2_index] = analyze(
        case.dataset, factors, (SectionPlan("ch4.mediu", "environment", (co2_plan,)),)
    )[0].blocks[0]
    for chart in environment.charts:
        item = ChartPlan(
            f"item-{chart.body_index}",
            chart.part,
            chart.series.name,
            Metric("filed", product="co2_chart"),
            case.dataset.years,
            tuple(chart.series.categories),
        )
        generated[chart.body_index] = analyze(
            case.dataset, factors, (SectionPlan("ch4.mediu", "environment", (item,)),)
        )[0].blocks[0]
        chart_plans[chart.body_index] = item
    for section_id, carrier in (
        ("ch4.specific_electric", Carrier.electricity_grid),
        ("ch4.specific_gaz", Carrier.natural_gas),
    ):
        for chart in case.sections[section_id].charts:
            item = ChartPlan(
                f"item-{chart.body_index}",
                chart.part,
                chart.series.name,
                Metric("specific", (carrier,), "turnover_specific"),
                case.dataset.years,
                tuple(chart.series.categories),
                decimals=4,
            )
            generated[chart.body_index] = analyze(
                case.dataset, factors, (SectionPlan(section_id, "specific", (item,)),)
            )[0].blocks[0]
            chart_plans[chart.body_index] = item
    document = Document(str(case.docx))
    children = list(document.element.body)
    source_id = "audit-01" if name == "audit-01" else "audit-02"
    bank = {item.paragraph: item for item in phrase_bank() if item.source_document == source_id}
    mapped = map_headings(case.docx, name.lower()).mapped
    chapter = next(item for item in mapped if item.section_id == "ch4")
    next_chapter = next(
        item
        for item in mapped
        if item.heading.index > chapter.heading.index and item.heading.level == 0
    )
    heading_positions = [
        (item.heading.index, item.section_id)
        for item in mapped
        if item.section_id.startswith("ch4")
    ]
    value_metrics = {
        "ch4.productie": Metric("production", product="turnover"),
        "ch4.electricitate": Metric("carrier", (Carrier.electricity_grid,)),
        "ch4.gaz": Metric("carrier", (Carrier.natural_gas,)),
        "ch4.echiv_electric": Metric("tep_monthly_sum", (Carrier.electricity_grid,)),
        "ch4.echiv_gaz": Metric("tep_monthly_sum", (Carrier.natural_gas,)),
        "ch4.specific_electric": Metric(
            "specific", (Carrier.electricity_grid,), "turnover_specific"
        ),
        "ch4.specific_gaz": Metric("specific", (Carrier.natural_gas,), "turnover_specific"),
        "ch4.intensitate": Metric("intensity"),
        "ch4.mediu": Metric("co2"),
    }
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        if not chapter.heading.index < paragraph_index < next_chapter.heading.index:
            continue
        section_id = next(
            section_id
            for position, section_id in reversed(heading_positions)
            if position <= paragraph_index
        )
        pattern = bank.get(paragraph_index)
        if pattern is not None and pattern.direction == "value" and section_id in value_metrics:
            text = " ".join(paragraph.text.split())
            year_match = re.match(r"(?i)^pentru anul (20\d{2})", text)
            number_match = re.search(r"(?:valoare de|s-au înregistrat) (\d[\d.,]*)", text)
            specific_unit_ok = not section_id.startswith("ch4.specific_") or "tep/mil lei" in text
            if year_match is not None and number_match is not None and specific_unit_ok:
                index = children.index(paragraph._p)
                number_text = number_match.group(1)
                item = NumericSentencePlan(
                    f"item-{index}",
                    value_metrics[section_id],
                    int(year_match.group(1)),
                    source_id,
                    paragraph_index,
                    len(number_text.rsplit(",", 1)[1]) if "," in number_text else 0,
                    grouping="." in number_text,
                )
                blocks = analyze(
                    case.dataset,
                    factors,
                    (SectionPlan(section_id, "specific", (item,)),),
                )[0].blocks
                if blocks:
                    generated[index] = blocks[0]
        if (
            pattern is None
            or "{number}" in pattern.pattern
            or "{figure_number}" not in pattern.pattern
        ):
            continue
        text = " ".join(paragraph.text.split())
        figure = re.search(r"Conform figurii numărul\s+(\d+(?:\.\d+)?)", text, re.I)
        year_match = re.search(r"\b20\d{2}\b", text)
        if figure is None:
            continue
        year = int(year_match.group()) if year_match else None
        if pattern.pattern.format(figure_number=figure.group(1), year=year).rstrip() != text:
            continue
        index = children.index(paragraph._p)
        candidates = [position for position in chart_plans if 0 < index - position <= 7]
        if not candidates:
            continue
        nearest = max(candidates)
        item = TrendPlan(
            f"item-{index}",
            "",
            figure.group(1),
            chart_plans[nearest],
            year,
            source_document=source_id,
            source_paragraph=paragraph_index,
        )
        blocks = analyze(case.dataset, factors, (SectionPlan("ch4.trend", "carrier", (item,)),))[
            0
        ].blocks
        if blocks:
            generated[index] = blocks[0]
    return generated, case
