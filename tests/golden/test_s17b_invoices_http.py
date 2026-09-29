"""HTTP invoice journey against the S9b local reference batch (level 1)."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from tests.workspace_jobs import create_job

from ema.api import create_app
from ema.core.jobs import status, subscribe
from ema.core.workspace import Workspace
from ema.invoices import confirm_client, export, readiness
from ema.invoices.identity_review import raw_outcomes

from .test_s9a_invoices import _baseline, _run_case, _source_paths

BASE = "http://127.0.0.1:8766"


def _visible_cells(workbook: object) -> dict[str, list[tuple[object, ...]]]:
    result: dict[str, list[tuple[object, ...]]] = {}
    for sheet in workbook.worksheets:  # type: ignore[attr-defined]
        visible = [
            index
            for index in range(1, sheet.max_column + 1)
            if not sheet.column_dimensions[sheet.cell(1, index).column_letter].hidden
        ]
        result[sheet.title] = [
            tuple(sheet.cell(row, column).value for column in visible)
            for row in range(1, sheet.max_row + 1)
        ]
    return result


@pytest.mark.golden
def test_http_upload_identity_and_workbook_equal_cli(tmp_path: Path) -> None:
    expected, _ = _baseline("CLIENT-I2")
    sources = _source_paths("CLIENT-I2", expected)
    ws = Workspace(tmp_path / "http")
    job = create_job(ws, "invoices", "CLIENT-I2", None)
    client = TestClient(create_app(ws, 8766, launch_code="synthetic-code"), base_url=BASE)
    csrf = client.post("/session", json={"code": "synthetic-code"}).json()["csrf"]
    headers = {"x-ema-csrf": csrf}
    upload = client.post(
        f"/jobs/{job}/invoices/files",
        files=[
            ("files", (source.name, source.read_bytes(), "application/pdf")) for source in sources
        ],
        headers=headers,
    )
    assert upload.status_code == 200, upload.status_code
    assert len(upload.json()["added"]) == len(sources)
    assert upload.json()["rejected"] == []
    read = client.post(
        f"/jobs/{job}/stages/invoices",
        json={"on_revision": 1},
        headers=headers,
    )
    assert read.status_code == 202, read.status_code
    for _event in subscribe(ws, job):
        pass
    assert status(ws, job).runs[-1]["state"] == "ready"
    assert [row["source_path"] for row in raw_outcomes(ws, job)] == [
        source.name for source in sources
    ]
    identity = client.get(f"/jobs/{job}/invoices/identity")
    assert identity.status_code == 200
    confirmed = client.post(
        f"/jobs/{job}/invoices/identity",
        json={
            "client_id": "CLIENT-I2",
            "on_revision": identity.json()["revision"],
            "confirm": True,
        },
        headers=headers,
    )
    assert confirmed.status_code == 200, confirmed.status_code
    assert readiness(ws, job).exportable == 24
    view = client.get(f"/jobs/{job}/invoices")
    assert view.status_code == 200
    assert len({row["slot"] for row in view.json()["rows"] if row["status"] == "exportable"}) == 24
    assert all(row["file_name"] and row["slot"] for row in view.json()["rows"])
    current = client.get(f"/jobs/{job}").json()
    start = client.post(
        f"/jobs/{job}/stages/invoices_workbook",
        json={"on_revision": current["revision"]},
        headers=headers,
    )
    assert start.status_code == 202, start.status_code
    for _event in subscribe(ws, job):
        pass
    assert status(ws, job).runs[-1]["state"] == "ready"
    output = next(
        item
        for item in client.get(f"/jobs/{job}/outputs").json()
        if item["stage"] == "invoices_workbook" and item["kind"] == "final"
    )
    response = client.get(f"/jobs/{job}/outputs/{output['id']}")
    assert response.status_code == 200
    http_book = load_workbook(BytesIO(response.content), data_only=False)

    cli_ws, cli_job, _raw = _run_case(tmp_path / "cli", "CLIENT-I2", expected)
    confirm_client(cli_ws, cli_job)
    with cli_ws.connect() as db:
        cli_path = cli_ws.job_path(db, cli_job) / "outputs" / "Facturi.xlsx"
    cli_book = load_workbook(export(cli_ws, cli_job, cli_path), data_only=False)
    assert _visible_cells(http_book) == _visible_cells(cli_book)
