"""Independent CLIENT-A1 chart oracle from the chart contract and reviewed source values."""

from __future__ import annotations

import math
from dataclasses import dataclass, replace
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook

from ema.audit.chapter_four import MONTHS
from ema.consumption_analysis.analysis import Metric, value
from ema.core.office.chart_series import Series, read_series
from ema.core.office.package import C, R, inspect, read_parts, relationships, target_part, xml
from ema.core.office.workbook import formula_cells
from ema.energy_data.carriers import Carrier
from ema.energy_data.factors import FACTORS_2026
from ema.energy_data.model import EnergyDataset

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
NAMES = {
    Carrier.electricity_grid: "Consumul de energie electrică din rețea",
    Carrier.natural_gas: "Consumul de gaze naturale",
    Carrier.diesel: "Consumul de motorină",
    Carrier.petrol: "Consumul de benzină",
    Carrier.lpg: "Consumul de GPL",
    Carrier.water_potable: "Consumul de apă potabilă",
    Carrier.water_industrial: "Consumul de apă industrială",
}
SINGLE = (
    (
        "ch4.echiv_electric",
        "Evoluția anuală al consumului total echivalent de energie electrică din SEN înregistrat la nivelul {client}",  # noqa: E501
        "tep",
    ),
    (
        "ch4.echiv_gaz",
        "Evoluția anuală al consumului echivalent de gaz natural înregistrat la nivelul {client}",
        "tep",
    ),
    (
        "ch4.echiv_carburant",
        "Evoluția anuală al consumului echivalent de carburant înregistrat la nivelul {client}",
        "tep",
    ),
    (
        "ch4.echiv_total",
        "Evoluția anuală al consumului total echivalent de energie înregistrat la nivelul {client}",
        "tep",
    ),
    (
        "ch4.specific_electric",
        "Evoluția anuală al consumului specific echivalent de energie electrică înregistrat la nivelul {client}",  # noqa: E501
        "tep/kg",
    ),
    (
        "ch4.specific_gaz",
        "Evoluția anuală al consumului specific echivalent de gaz natural înregistrat la nivelul {client}",  # noqa: E501
        "tep/kg",
    ),
    (
        "ch4.specific_carburant",
        "Evoluția anuală al consumului specific echivalent de carburant înregistrat la nivelul {client}",  # noqa: E501
        "tep/kg",
    ),
    (
        "ch4.specific_total",
        "Evoluția anuală a consumului specific echivalent total de energie înregistrat la nivelul {client}",  # noqa: E501
        "tep/kg",
    ),
    ("ch4.intensitate", "Tendința intensității energetice în cadrul {client}", "tep/1000 lei"),
    (
        "ch4.mediu",
        "Evoluția anuală a gazelor cu efect de seră înregistrate la nivelul {client}",
        "t CO₂",
    ),
)


@dataclass(frozen=True)
class Expected:
    caption: str
    series: list[Series]
    axis: str


def _unit(dataset: EnergyDataset, carrier: Carrier) -> str:
    year = next(iter(dataset.carriers[carrier].values()))
    reading = year.annual or next(iter(year.months.values()))
    return reading.unit


def expected_charts(  # noqa: C901, PLR0912, PLR0915
    dataset: EnergyDataset, client: str
) -> list[Expected]:
    """The Contract rows in catalogue order; no call to the production chart planner."""
    found: list[Expected] = []
    fuels = tuple(
        carrier
        for carrier in dataset.carriers
        if carrier
        in {
            Carrier.diesel,
            Carrier.petrol,
            Carrier.lpg,
            Carrier.fuel_oil,
            Carrier.clu,
        }
    )
    product = next(iter(dataset.production))

    def group(
        metrics: list[tuple[str, Metric]],
        years: tuple[int, ...],
        unit: str,
        monthly_subject: str | None,
        annual_text: str,
        *,
        production: bool = False,
        specific: bool = False,
    ) -> None:
        charts: list[tuple[int | None, list[Series]]] = []
        if monthly_subject is not None:
            for year in years:
                series = [
                    Series(
                        name,
                        list(MONTHS),
                        [
                            value(dataset, FACTORS_2026, replace(metric, month=month), year)[0]
                            for month in range(1, 13)
                        ],
                    )
                    for name, metric in metrics
                ]
                if any(point is not None for item in series for point in item.values):
                    charts.append((year, series))
        annual = [
            Series(
                name,
                [str(year) for year in years],
                [value(dataset, FACTORS_2026, metric, year)[0] for year in years],
            )
            for name, metric in metrics
        ]
        if any(point is not None for item in annual for point in item.values):
            charts.append((None, annual))
        if not charts:
            return
        number = len({item.caption.split(" ", 3)[2] for item in found}) + 1
        for index, (year, series) in enumerate(charts):
            letter = f"{chr(ord('a') + index)}) " if any(y is not None for y, _ in charts) else ""
            prefix = f"Fig. nr. 4.{number} {letter}"
            if year is not None:
                assert monthly_subject is not None
                if production:
                    subject = "cifrei lunare de afaceri" if "lei" in unit else "producției lunare"
                    caption = (
                        f"{prefix}Evoluția lunară a {subject} înregistrate de către {client} "
                        f"la nivelul anului {year}"
                    )
                else:
                    caption = (
                        f"{prefix}Evoluția lunară a {monthly_subject} înregistrat de către "
                        f"{client} la nivelul anului {year}"
                    )
            else:
                caption = prefix + annual_text.format(client=client)
            axis = unit if specific else unit + ("/lună" if year is not None else "/an")
            found.append(Expected(caption, series, axis))

    group(
        [(product.replace("_", " "), Metric("production", product=product))],
        tuple(year for year in dataset.years if year in dataset.production[product]),
        dataset.production_unit[product],
        "production",
        "Evoluția anuală a producției înregistrate la nivelul {client}",
        production=True,
    )
    for carrier, subject in (
        (Carrier.electricity_grid, "consumului de energie electrică din SEN"),
        (Carrier.natural_gas, "consumului de gaz natural"),
    ):
        group(
            [(NAMES[carrier], Metric("carrier", (carrier,)))],
            tuple(year for year in dataset.years if year in dataset.carriers[carrier]),
            _unit(dataset, carrier),
            subject,
            "Evoluția anuală al " + subject + " înregistrat la nivelul {client}",
        )
    group(
        [(NAMES[carrier], Metric("carrier", (carrier,))) for carrier in fuels],
        tuple(year for year in dataset.years if any(year in dataset.carriers[c] for c in fuels)),
        _unit(dataset, fuels[0]),
        "consumului de carburant",
        "Evoluția anuală al consumului de carburant înregistrat la nivelul {client}",
    )
    for carrier, subject in (
        (Carrier.water_potable, "consumului de apă"),
        (Carrier.water_industrial, "consumului de apă industrială"),
    ):
        if carrier in dataset.carriers:
            group(
                [(NAMES[carrier], Metric("carrier", (carrier,)))],
                tuple(year for year in dataset.years if year in dataset.carriers[carrier]),
                _unit(dataset, carrier),
                subject,
                "Evoluția anuală al " + subject + " înregistrat la nivelul {client}",
            )
    for section, text, unit in SINGLE:
        if section == "ch4.intensitate":
            for carrier, water_text in (
                (
                    Carrier.water_potable,
                    "Evoluția anuală a consumului specific de apă înregistrat la nivelul {client}",
                ),
                (
                    Carrier.water_industrial,
                    "Evoluția anuală a consumului specific de apă industrială înregistrat la nivelul {client}",  # noqa: E501
                ),
            ):
                if carrier in dataset.carriers:
                    water_years = tuple(
                        year
                        for year in dataset.years
                        if year in dataset.carriers[carrier] and year in dataset.production[product]
                    )
                    group(
                        [(NAMES[carrier], Metric("water_specific", (carrier,), product=product))],
                        water_years,
                        "m³/" + dataset.production_unit[product],
                        None,
                        water_text,
                        specific=True,
                    )
        chart_unit = "tep/" + dataset.production_unit[product] if unit == "tep/kg" else unit
        if section.endswith("electric"):
            carriers = (Carrier.electricity_grid,)
        elif section.endswith("gaz"):
            carriers = (Carrier.natural_gas,)
        elif section.endswith("carburant"):
            carriers = fuels
        else:
            carriers = ()
        if section.startswith("ch4.echiv_") and section != "ch4.echiv_total":
            metrics = [(NAMES[carrier], Metric("tep", (carrier,))) for carrier in carriers]
        elif section == "ch4.echiv_total":
            metrics = [("Consum total echivalent", Metric("tep_total"))]
        elif section.startswith("ch4.specific_") and section != "ch4.specific_total":
            metrics = [
                (NAMES[carrier], Metric("specific", (carrier,), product=product))
                for carrier in carriers
            ]
        elif section == "ch4.specific_total":
            metrics = [("Consum total", Metric("specific", product=product))]
        elif section == "ch4.intensitate":
            metrics = [("Intensitate energetică", Metric("intensity"))]
        else:
            metrics = [("Emisii", Metric("co2"))]
        years = tuple(
            year
            for year in dataset.years
            if (not carriers or any(year in dataset.carriers[c] for c in carriers))
            and (not section.startswith("ch4.specific_") or year in dataset.production[product])
        )
        group(
            metrics,
            years,
            chart_unit,
            None,
            text,
            specific=section.startswith("ch4.specific_") or section == "ch4.intensitate",
        )
    return found


def _text(element: etree._Element) -> str:
    return "".join(node.text or "" for node in element.iter(W + "t"))


def _cell_values(book: object, formula: str) -> list[object]:
    sheet, cells = formula_cells(formula)
    return [book[sheet].cell(row, col).value for row, col in cells]  # type: ignore[index,attr-defined]


def assert_final_charts(docx: Path, expected: list[Expected]) -> None:
    package = inspect(docx)
    assert len(package.charts) == len(expected)
    parts = read_parts(docx)
    rels = {
        rel.get("Id"): target_part("word/document.xml", rel.get("Target", ""))
        for rel in relationships(parts, "word/document.xml")
        if rel.get("Type", "").endswith("/chart")
    }
    body = xml(parts, "word/document.xml").find(W + "body")
    assert body is not None
    actual: list[tuple[str, str]] = []
    for paragraph in body:
        for chart in paragraph.iter(f"{{{C}}}chart"):
            following = paragraph.getnext()
            assert following is not None and following.tag == W + "p"
            actual.append((rels[chart.get(f"{{{R}}}id")], _text(following)))
    assert len(actual) == len(expected)
    with ZipFile(docx) as archive:
        for index, ((part, caption), wanted) in enumerate(zip(actual, expected, strict=True)):
            assert caption == wanted.caption, f"chart {index + 1}: caption"
            root = xml(parts, part)
            axis = root.xpath("string(.//c:valAx/c:title)", namespaces={"c": C})
            axis_text = "".join(
                root.xpath(
                    ".//c:valAx/c:title//a:t/text()",
                    namespaces={
                        "c": C,
                        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                    },
                )
            )
            assert axis and axis_text == wanted.axis, f"chart {index + 1}: axis"
            chart_ref = next(item for item in package.charts if item.part == part)
            assert chart_ref.embedded and chart_ref.external is None
            assert len(package.chart_workbook_relationships[part]) == 1
            book = load_workbook(BytesIO(archive.read(chart_ref.embedded)), data_only=True)
            series = read_series(docx, part)
            assert len(series) == len(wanted.series), f"chart {index + 1}: series count"
            for actual_series, target in zip(series, wanted.series, strict=True):
                assert (actual_series.name, actual_series.categories) == (
                    target.name,
                    target.categories,
                )
                assert actual_series.refs is not None
                assert _cell_values(book, actual_series.refs.name or "") == [target.name]
                assert _cell_values(book, actual_series.refs.categories or "") == target.categories
                workbook_values = _cell_values(book, actual_series.refs.values or "")
                for cached, embedded, wanted_value in zip(
                    actual_series.values, workbook_values, target.values, strict=True
                ):
                    if wanted_value is None:
                        assert cached is None and embedded is None, f"chart {index + 1}: gap"
                    else:
                        assert cached is not None and embedded is not None
                        assert math.isclose(cached, wanted_value, rel_tol=1e-9)
                        assert math.isclose(float(embedded), wanted_value, rel_tol=1e-9)
