"""Both reference inputs exercise the PIEE review gate through generated jobs."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from tests.golden.cases import case_path
from tests.golden.piee_case_b_review import review_case_b_reconciliation
from typer.testing import CliRunner

from conftest import artifacts_path
from ema.cli import _app
from ema.clients.registry import create_client
from ema.core.errors import EmaError
from ema.core.jobs import status, subscribe
from ema.core.review import decide, fields
from ema.core.workspace import Workspace
from ema.piee.review_workflow import PieeWorkflow, _latest_draft
from ema.piee.workflow import GenerateRequest, start_generate, start_generate_for_job

pytestmark = [pytest.mark.golden, pytest.mark.word]


@pytest.mark.parametrize("case_name", ["piee-case-a", "piee-case-b"])
def test_review_decision_and_final_gate(  # noqa: C901, PLR0915
    reference_library: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, case_name: str
) -> None:
    base = case_path("piee-01")
    monkeypatch.setenv("EMA_PIEE_BASE_DOCUMENT", str(base))
    monkeypatch.setenv("EMA_PIEE_BASE_DIRECTORY", str(artifacts_path("s8", "base")))
    case = case_path(case_name)
    anexa = next(case.rglob("Anexa*.xlsx"))
    necesar = next(case.rglob("Necesar*.xls"), None)
    prelucrare = next(case.rglob("*Prelucrare*.xls*"))
    previous_piee = (
        reference_library / case_path("piee-case-b", "final")
        if case_name == "piee-case-b"
        else None
    )
    ws = Workspace(tmp_path / "workspace")
    registered = create_client(ws, "Synthetic", "12345678")
    job, run = start_generate(
        ws, GenerateRequest(str(registered["id"]), 2025, anexa, necesar, prelucrare, previous_piee)
    )
    for _ in subscribe(ws, job):
        pass
    assert next(item for item in status(ws, job).runs if item["id"] == run)["state"] == "ready"
    if previous_piee is not None:
        with ws.connect() as db:
            assert (
                db.execute(
                    "SELECT 1 FROM slots WHERE job_id=? AND name='previous_piee'", (job,)
                ).fetchone()
                is not None
            )

    workflow = PieeWorkflow()
    if case_name == "piee-case-b":
        review_case_b_reconciliation(ws, job)
    conflict = next(field for field in fields(ws, job) if field.confidence == "conflict")
    assert any(issue.field_id == conflict.id for issue in workflow.readiness(ws, job).blocking)
    with pytest.raises(EmaError) as refused:
        workflow.render(ws, job, "final")
    assert refused.value.code == "not_ready"
    choice = next(item for item in conflict.alternatives if item.value == conflict.value)
    decide(ws, job, conflict.id, "choose", conflict.revision, "user", alternative=choice.id)
    assert not any(issue.field_id == conflict.id for issue in workflow.readiness(ws, job).blocking)
    for item in fields(ws, job, status="conflict"):
        selected = next(
            candidate for candidate in item.alternatives if candidate.value == item.value
        )
        decide(ws, job, item.id, "choose", item.revision, "user", alternative=selected.id)
    calculated = [
        item
        for item in fields(ws, job)
        if item.key.endswith(".payback_years")
        and item.state == "calculated"
        and item.review == "pending"
    ]
    assert calculated
    assert any(
        issue.code == "calculated_unconfirmed" for issue in workflow.readiness(ws, job).blocking
    )
    for item in calculated:
        decide(ws, job, item.id, "accept", item.revision, "user")
    assert any(issue.code == "stale" for issue in workflow.readiness(ws, job).blocking)
    regenerated = start_generate_for_job(ws, job)
    for _ in subscribe(ws, job):
        pass
    assert (
        next(item for item in status(ws, job).runs if item["id"] == regenerated)["state"] == "ready"
    )
    if case_name == "piee-case-b":
        _, _, draft = _latest_draft(ws, job)
        pv_year = Document(draft).tables[4].rows[2]
        assert pv_year.cells[0].text == "2024"
        missing_pv = pv_year.cells[1]
        assert missing_pv.text == "n.d."
        assert any(
            run.text == "n.d." and run.font.color and str(run.font.color.rgb) == "FF0000"
            for paragraph in missing_pv.paragraphs
            for run in paragraph.runs
        )
    assert workflow.readiness(ws, job).final_ok

    class StubWord:
        def __init__(self, **_kwargs: object) -> None:
            pass

        def update_toc_pages(self, _docx: Path) -> None:
            pass

        def render_pdf(self, _docx: Path, pdf: Path) -> None:
            pdf.write_bytes(b"%PDF-1.4\n")

        def open_check(self, _docx: Path) -> None:
            pass

    monkeypatch.setattr("ema.piee.review_workflow.word_automation", lambda _settings: StubWord())
    final_id = workflow.render(ws, job, "final")
    with ws.connect() as db:
        outputs = db.execute(
            "SELECT kind,relative_path FROM outputs WHERE job_id=? ORDER BY seq DESC LIMIT 3",
            (job,),
        ).fetchall()
    assert len(outputs) == 3
    assert [row["kind"] for row in outputs] == ["final", "draft", "draft"]
    assert final_id
    reviewed = next(item for item in fields(ws, job) if item.id == conflict.id)
    alternative = next(item for item in reviewed.alternatives if item.value != reviewed.value)
    decide(ws, job, reviewed.id, "choose", reviewed.revision, "user", alternative=alternative.id)
    assert any(issue.code == "stale" for issue in workflow.readiness(ws, job).blocking)


def test_cli_generate_review_and_refuse_unresolved_final(
    reference_library: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    base = case_path("piee-01")
    monkeypatch.setenv("EMA_WORKSPACE", str(tmp_path / "workspace"))
    monkeypatch.setenv("EMA_PIEE_BASE_DOCUMENT", str(base))
    monkeypatch.setenv("EMA_PIEE_BASE_DIRECTORY", str(artifacts_path("s8", "base")))
    case = reference_library / case_path("piee-case-b")
    anexa = next(case.rglob("Anexa*.xlsx"))
    prelucrare = next(case.rglob("*Prelucrare*.xls*"))
    create_client(Workspace(tmp_path / "workspace"), "Synthetic", "12345678")
    runner = CliRunner()
    generated = runner.invoke(
        _app,
        [
            "piee",
            "generate",
            "--client",
            "12345678",
            "--year",
            "2025",
            "--anexa",
            str(anexa),
            "--prelucrare",
            str(prelucrare),
        ],
    )
    assert generated.exit_code == 0
    job = json.loads(generated.output)["job"]
    checked = runner.invoke(_app, ["job", "checks", job])
    assert checked.exit_code == 0
    assert any(item["code"] == "conflict" for item in json.loads(checked.output)["blocking"])
    refused = runner.invoke(
        _app, ["job", "export", job, "--dest", str(tmp_path / "final.docx"), "--final"]
    )
    assert isinstance(refused.exception, EmaError)
    assert refused.exception.code == "not_ready"
    ws = Workspace(tmp_path / "workspace")
    conflicting = fields(ws, job, status="conflict")
    first = conflicting[0]
    selected = next(candidate for candidate in first.alternatives if candidate.value == first.value)
    decided = runner.invoke(
        _app,
        [
            "job",
            "decide",
            job,
            first.id,
            "choose",
            "--on-revision",
            str(first.revision),
            "--alternative",
            selected.id,
        ],
    )
    assert decided.exit_code == 0
    assert len(fields(ws, job, status="conflict")) == len(conflicting) - 1
