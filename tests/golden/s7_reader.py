"""Test-only S7 input reader for authored chapter-four tables and chart caches."""

# pyright: reportPrivateUsage=false

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from lxml import etree

from ema.audit.headings import map_headings
from ema.core.office.chart_series import Series, read_series
from ema.core.office.package import C, R, inspect
from ema.energy_data.carriers import Carrier
from ema.energy_data.model import CarrierSeries, EnergyDataset, FiledValue, Reading

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


@dataclass(frozen=True)
class Chart:
    body_index: int
    part: str
    series: Series


@dataclass
class Section:
    tables: list[tuple[int, etree._Element]] = field(default_factory=list)
    charts: list[Chart] = field(default_factory=list)


@dataclass(frozen=True)
class Chapter:
    docx: Path
    sections: dict[str, Section]
    dataset: EnergyDataset


def _text(node: etree._Element) -> str:
    return "".join(item.text or "" for item in node.iter(f"{{{W}}}t")).strip()


def _number(text: str) -> float:
    return float(text.replace(".", "").replace(",", "."))


def _filed(text: str, unit: str, source: str) -> FiledValue:
    return FiledValue(
        _number(text), unit, len(text.rsplit(",", 1)[1]) if "," in text else 0, source
    )


def _cells(row: etree._Element) -> list[str]:
    return [_text(cell) for cell in row.findall(f"{{{W}}}tc")]


def _column(headers: list[str], label: str) -> int:
    matches = [index for index, text in enumerate(headers) if label in text.casefold()]
    if len(matches) != 1:
        raise ValueError(f"expected one {label} column")
    return matches[0]


def _chapter_indicators(  # noqa: C901, PLR0912
    sections: dict[str, Section],
) -> dict[str, dict[int, FiledValue]]:
    filed: dict[str, dict[int, FiledValue]] = {}

    def record(key: str, year: int, text: str, unit: str, location: str) -> None:
        filed.setdefault(key, {})[year] = _filed(text, unit, location)

    intensity_index, intensity_table = sections["ch4.intensitate"].tables[0]
    intensity_rows = intensity_table.findall(f"{{{W}}}tr")
    headers = _cells(intensity_rows[0])
    columns = (
        (_column(headers, "ceq"), "tep_total", "tep"),
        (
            _column(headers, "producției")
            if "producției" in " ".join(headers).casefold()
            else _column(headers, "afaceri"),
            "turnover_thousand",
            "mii lei",
        ),
        (_column(headers, "intensitatea"), "intensity", "tep/1000 lei"),
    )
    year_column = _column(headers, "an")
    for row_index, row in enumerate(intensity_rows[1:], 1):
        cells = _cells(row)
        if not cells[year_column].isdigit():
            continue
        year = int(cells[year_column])
        for column, key, unit in columns:
            record(key, year, cells[column], unit, f"ch4:{intensity_index}:{row_index}:{column}")

    co2_index, co2_table = sections["ch4.mediu"].tables[0]
    rows = co2_table.findall(f"{{{W}}}tr")
    if "anul" in _cells(rows[0])[0].casefold():
        for row_index, row in enumerate(rows[1:], 1):
            cells = _cells(row)
            if not cells[0].isdigit():
                continue
            record("co2.total", int(cells[0]), cells[1], "t CO₂", f"ch4:{co2_index}:{row_index}:1")
    else:
        year_cells = next((_cells(row) for row in rows if _cells(row)[0] == ""), None)
        if year_cells is None:
            raise ValueError("environment table has no year labels")
        years = [int(text) for text in year_cells[1:]]
        keys = {
            "energia electrică": "co2.electricity_grid",
            "gazul natural": "co2.natural_gas",
            "carburant": "co2.fuel",
            "total": "co2.total",
        }
        for row_index, row in enumerate(rows):
            cells = _cells(row)
            key = keys.get(cells[0].casefold())
            if key is None:
                continue
            for column, year in enumerate(years, 1):
                record(key, year, cells[column], "t CO₂", f"ch4:{co2_index}:{row_index}:{column}")
        if any(key not in filed for key in keys.values()):
            raise ValueError("environment table is missing a carrier row")
    for section_id, key, unit in (
        ("ch4.intensitate", "intensity_chart", "tep/1000 lei"),
        ("ch4.mediu", "co2_chart", "t CO₂"),
    ):
        for chart in sections[section_id].charts:
            for year_text, value in zip(chart.series.categories, chart.series.values, strict=True):
                if value is not None:
                    filed.setdefault(key, {})[int(year_text)] = FiledValue(
                        value, unit, 12, f"ch4:chart:{chart.body_index}"
                    )
    return filed


def _monthly(section: Section) -> tuple[tuple[int, ...], dict[int, dict[int, float]]]:
    if len(section.tables) != 2:
        raise ValueError("expected two authored monthly half-year tables")
    months: dict[int, dict[int, float]] = {}
    years: tuple[int, ...] = ()
    for half, (_, table) in enumerate(section.tables):
        rows = table.findall(f"{{{W}}}tr")
        if len(rows) != 4:
            raise ValueError("expected a heading and three year rows")
        found = tuple(int(_text(row.findall(f"{{{W}}}tc")[0])) for row in rows[1:])
        if half == 0:
            years = found
            months = {year: {} for year in years}
        elif found != years:
            raise ValueError("half-year tables disagree on years")
        for year, row in zip(years, rows[1:], strict=True):
            cells = row.findall(f"{{{W}}}tc")
            if len(cells) != 7:
                raise ValueError("monthly table must have six months")
            for offset, cell in enumerate(cells[1:], 1):
                months[year][half * 6 + offset] = _number(_text(cell))
    return years, months


def _series(section: Section, years: tuple[int, ...], unit: str) -> dict[int, CarrierSeries]:
    _, tables = _monthly(section)
    monthly_charts = section.charts[:3] if len(section.charts) == 4 else []
    annual_chart = section.charts[-1] if section.charts else None
    result: dict[int, CarrierSeries] = {}
    for index, year in enumerate(years):
        values = (
            monthly_charts[index].series.values if monthly_charts else list(tables[year].values())
        )
        if len(values) != 12:
            raise ValueError("monthly chart cache must contain twelve points")
        months = {month: Reading(value, unit) for month, value in enumerate(values, 1)}
        annual = (
            annual_chart.series.values[index]
            if annual_chart is not None and len(annual_chart.series.values) == len(years)
            else sum(value for value in values if value is not None)
        )
        result[year] = CarrierSeries(months, Reading(annual, unit))
    return result


def read_chapter(path: Path, audit_id: str) -> Chapter:
    document = Document(path)
    children = list(document.element.body)
    mapping = map_headings(path, audit_id).mapped
    heading = next(item for item in mapping if item.section_id == "ch4")
    following = next(
        item
        for item in mapping
        if item.heading.index > heading.heading.index and item.heading.level == 0
    )
    start = children.index(document.paragraphs[heading.heading.index]._p)
    end = children.index(document.paragraphs[following.heading.index]._p)
    ids = {
        id(document.paragraphs[item.heading.index]._p): item.section_id
        for item in mapping
        if item.section_id.startswith("ch4")
    }
    chart_parts = {chart.rel_id: chart.part for chart in inspect(path).charts}
    sections: dict[str, Section] = {}
    current = "ch4"
    for index in range(start, end):
        node = children[index]
        current = ids.get(id(node), current)
        section = sections.setdefault(current, Section())
        if node.tag == f"{{{W}}}tbl":
            section.tables.append((index, node))
        for chart in node.iter(f"{{{C}}}chart"):
            part = chart_parts[chart.get(f"{{{R}}}id")]
            series = read_series(path, part)
            if len(series) != 1:
                raise ValueError("S7 reference chart must have one series")
            section.charts.append(Chart(index, part, series[0]))
    electric_years, _ = _monthly(sections["ch4.electricitate"])
    gas_years, _ = _monthly(sections["ch4.gaz"])
    if electric_years != gas_years:
        raise ValueError("electricity and gas periods differ")
    carriers = {
        Carrier.electricity_grid: _series(sections["ch4.electricitate"], electric_years, "MWh"),
        Carrier.natural_gas: _series(sections["ch4.gaz"], gas_years, "MWh"),
    }
    production: dict[str, dict[int, CarrierSeries]] = {}
    if len(sections["ch4.productie"].tables) == 2:
        prod_years, _ = _monthly(sections["ch4.productie"])
        production["turnover"] = _series(sections["ch4.productie"], prod_years, "lei")
    elif sections["ch4.productie"].charts:
        values = sections["ch4.productie"].charts[0].series.values
        production["turnover"] = {
            year: CarrierSeries(annual=Reading(value, "lei"))
            for year, value in zip(electric_years, values, strict=True)
        }
    scale = 1_000_000
    production["turnover_specific"] = {
        year: CarrierSeries(annual=Reading(series.annual.value / scale, "mil lei"))
        for year, series in production["turnover"].items()
        if series.annual is not None and series.annual.value is not None
    }
    turnover_lei = {
        year: Reading(series.annual.value, "lei")
        for year, series in production["turnover"].items()
        if series.annual is not None and series.annual.value is not None
    }
    dataset = EnergyDataset(
        electric_years,
        carriers,
        production,
        {"turnover": "lei", "turnover_specific": "mil lei"},
        turnover_lei,
        filed_indicators=_chapter_indicators(sections),
        energy_inventory_complete=False,
    )
    return Chapter(path, sections, dataset)
