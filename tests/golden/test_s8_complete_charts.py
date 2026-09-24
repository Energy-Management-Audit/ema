"""All CLIENT-P1 chart caches and editable workbooks follow the approved figures."""

from __future__ import annotations

import math
import re
from collections import defaultdict
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ema.core.office.chart_series import Series, read_series
from ema.core.office.package import REL_PACKAGE, read_parts, relationships, target_part
from ema.core.office.workbook import formula_cells
from ema.piee.compose import compose_draft
from ema.piee.dataset import load

pytestmark = pytest.mark.golden
CHART = re.compile(r"word/charts/chart\d+\.xml")


def _chart_key(series: list[Series]) -> tuple[tuple[str, tuple[str, ...]], ...]:
    return tuple((item.name, tuple(item.categories)) for item in series)


def _workbook(parts: dict[str, bytes], chart: str) -> dict[tuple[str, str], Any]:
    package = next(
        relation for relation in relationships(parts, chart) if relation.get("Type") == REL_PACKAGE
    )
    target = target_part(chart, package.get("Target", ""))
    book = load_workbook(BytesIO(parts[target]), read_only=True, data_only=False)
    return {
        (sheet.title, cell.coordinate): cell.value
        for sheet in book
        for row in sheet
        for cell in row
        if cell.value is not None
    }


def _same_value(left: Any, right: Any) -> bool:
    if isinstance(left, int | float) and isinstance(right, int | float):
        return math.isclose(left, right, rel_tol=1e-9, abs_tol=0)
    return bool(left == right)


def test_all_CLIENT-P1_charts_and_workbooks_match_approved_semantics(
    reference_library: Path, tmp_path: Path
) -> None:
    case = reference_library / "piee/cases/piee-case-a"
    data = load(
        2025,
        next(case.rglob("Anexa*.xlsx")),
        next(case.rglob("Necesar*.xls")),
        next(case.rglob("*Prelucrare*.xls*")),
    )
    output = tmp_path / "draft.docx"
    compose_draft(data, Path.home() / "Ema-dev/s8/base", output, date(2026, 9, 24))
    approved = next((case / "generated").glob("Program de îmbunătățire*.docx"))
    generated_parts, approved_parts = read_parts(output), read_parts(approved)
    expected: dict[tuple[tuple[str, tuple[str, ...]], ...], list[tuple[str, list[Series]]]] = (
        defaultdict(list)
    )
    for chart in approved_parts:
        if CHART.fullmatch(chart):
            series = read_series(approved, chart)
            expected[_chart_key(series)].append((chart, series))
    matched = 0
    for chart in generated_parts:
        if not CHART.fullmatch(chart):
            continue
        series = read_series(output, chart)
        key = _chart_key(series)
        assert expected[key], f"chart series has no approved match: {chart}"
        cells = _workbook(generated_parts, chart)
        if chart == "word/charts/chart31.xml":
            name_formula = series[0].refs.name if series[0].refs else None
            assert name_formula is not None
            sheet, coordinates = formula_cells(name_formula)
            assert len(coordinates) == 1
            row, column = coordinates[0]
            name_cell = (sheet, f"{get_column_letter(column)}{row}")
            assert name_cell in cells
            del cells[name_cell]
        match = next(
            (
                (index, approved_chart, approved_series, workbook)
                for index, (approved_chart, approved_series) in enumerate(expected[key])
                if (workbook := _workbook(approved_parts, approved_chart)).keys() == cells.keys()
            ),
            None,
        )
        assert match is not None, f"workbook range has no approved match: {chart}"
        index, _approved_chart, approved_series, reference_cells = match
        expected[key].pop(index)
        for actual, authored in zip(series, approved_series, strict=True):
            assert actual.name == authored.name and actual.categories == authored.categories
            assert len(actual.values) == len(authored.values)
            assert all(
                _same_value(left, right)
                for left, right in zip(actual.values, authored.values, strict=True)
            ), f"numeric cache differs: {chart}"
        assert cells.keys() == reference_cells.keys(), f"workbook range differs: {chart}"
        assert all(_same_value(value, reference_cells[key]) for key, value in cells.items())
        matched += 1
    assert matched == 31
    assert not any(expected.values())
