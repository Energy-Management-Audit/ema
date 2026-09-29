"""Render synthetic chapter-five readings into the local auditor base."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from tests.conftest import artifacts_path
from tests.golden.test_s10b_audit_base import _identity, _references

from ema.audit.base import build_base
from ema.audit.base_package import package_issues
from ema.audit.base_units import UnitPlan, heading_spans_document
from ema.audit.chapter_five import (
    ChapterFivePlan,
    PlannedPanel,
    PlannedPhoto,
    PlannedReading,
    PlannedThermal,
)
from ema.audit.chapter_five_render import render_chapter_five
from ema.core.office.package import read_parts

pytestmark = pytest.mark.golden


def test_chapter_five_in_real_base_with_synthetic_readings(
    reference_library: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    base_document, _ = _references(reference_library)
    AUDIT-03 = next((reference_library / "audit/finished-audits").glob("*AUDIT-03*.docx"))
    model = (
        reference_library / "audit/section-models/Fișa de măsuratori electroenergetice - model.docx"
    )
    assert model.is_file()
    identity = (*_identity(base_document), *_identity(AUDIT-03))
    base = build_base(
        UnitPlan("Atelier Exemplu", 1, frozenset({"electricity"}), 1, True, 0, 0),
        base_document=base_document,
        measurement_prototype=AUDIT-03,
        output=tmp_path / "base.docx",
        base_identity=_identity(base_document),
    )
    image = tmp_path / "display.png"
    thermal = tmp_path / "thermal.png"
    Image.new("RGB", (100, 50), "white").save(image)
    Image.new("RGB", (50, 100), "black").save(thermal)
    slot = "visit/meter/Panel 1/display.png"
    thermal_slot = "visit/thermal/thermal.png"
    plan = ChapterFivePlan(
        panels=[
            PlannedPanel(
                id="panel-1",
                label="Panel 1",
                device="Synthetic meter",
                photos=[
                    PlannedPhoto(
                        slot=slot,
                        sha="a" * 64,
                        name="display.png",
                        caption="Valorile Tensiunii de fază",
                        display="voltage_ln",
                        readings=[
                            PlannedReading(
                                key="meter.panel-1.aaaaaaaa.voltage_ln.l1",
                                label="U1",
                                value="230.00",
                                unit="V",
                            )
                        ],
                        norm="voltage",
                        narrative_key=None,
                    )
                ],
            )
        ],
        thermal=[
            PlannedThermal(slot=thermal_slot, sha="b" * 64, name="thermal.png", component="tablou")
        ],
        visit_date="2026-09-27",
        client="Atelier Exemplu",
        missing_narratives=["narrative.ch5.termic_rezultate"],
    )
    monkeypatch.setenv("EMA_AUDIT_MEASUREMENT_SHEET_MODEL", str(model))
    monkeypatch.setenv("EMA_AUDIT_MEASUREMENT_PROTOTYPE", str(AUDIT-03))
    output = artifacts_path("s15a", "chapter-five-synthetic.docx")
    output.parent.mkdir(parents=True, exist_ok=True)
    report = render_chapter_five(base, output, plan, {slot: image, thermal_slot: thermal}, identity)
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Fig. 5.1" in text and "Fig. 5.2" in text
    assert "230,00 V" in text
    assert "Synthetic meter" in text
    assert "[de completat]" in text
    assert not package_issues(output, identity)
    assert not report.issues
    source = Document(AUDIT-03)
    source_positions = {item.section_id: start for item, start, _ in heading_spans_document(source)}
    result_positions = {
        item.section_id: start for item, start, _ in heading_spans_document(document)
    }
    source_body = list(source.element.body)
    result_body = list(document.element.body)

    def visible(node: object) -> str:
        return "".join(
            part.text or ""
            for part in node.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")
        ).translate(str.maketrans("şţŞŢ", "șțȘȚ"))  # type: ignore[attr-defined]

    method = [
        visible(node)
        for node in source_body[
            source_positions["ch5.electric"] + 1 : source_positions["ch5.electric_fisa"]
        ]
    ]
    instrument = next(
        index
        for index, item in enumerate(method)
        if item.strip().startswith("Pentru realizarea măsurătorilor electrice")
    )
    rendered_method = [
        visible(node)
        for node in result_body[
            result_positions["ch5.electric"] + 1 : result_positions["ch5.electric_fisa"]
        ]
    ]
    assert rendered_method == method[:instrument]
    harmonics = [
        visible(node)
        for node in source_body[
            source_positions["ch5.electric_concluzii"] + 2 : source_positions["ch5.termic"]
        ]
    ]
    harmonics = [
        item
        for item in harmonics
        if not any(term.casefold() in item.casefold() for term in identity)
    ]
    rendered_harmonics = [
        visible(node)
        for node in result_body[
            result_positions["ch5.electric_concluzii"] + 1 : result_positions["ch5.termic"]
        ]
    ]
    assert rendered_harmonics[1:] == harmonics
    thermal_method = [
        visible(node)
        for node in source_body[
            source_positions["ch5.termic"] + 1 : source_positions["ch5.termic_fisa"]
        ]
    ]
    rendered_thermal = [
        visible(node)
        for node in result_body[
            result_positions["ch5.termic"] + 1 : result_positions["ch5.termic_fisa"]
        ]
    ]
    assert rendered_thermal == thermal_method
    package_text = "\n".join(paragraph.text for paragraph in document.paragraphs).casefold()
    assert not any(term.casefold() in package_text for term in identity)
    assert [entry.number for entry in report.numbers[:2]] == ["5.1", "5.2"]
    media_hashes = {
        hashlib.sha256(data).digest()
        for name, data in read_parts(output).items()
        if name.startswith("word/media/")
    }
    assert hashlib.sha256(image.read_bytes()).digest() in media_hashes
    assert hashlib.sha256(thermal.read_bytes()).digest() in media_hashes
