"""piee-case-c: level 2 source readings from the auditor's Prelucrare workbook."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook
from tests.golden.cases import case_path

from ema.energy_data.calc import specific_consumption, tep, tep_total, water_specific
from ema.energy_data.carriers import Carrier
from ema.energy_data.prelucrare import import_prelucrare

pytestmark = pytest.mark.golden


def test_piee_case_c_source_series_and_cells() -> None:
    source = case_path("piee-case-c", "prelucrare")
    imported = import_prelucrare(source)
    book = load_workbook(source, read_only=True, data_only=True)
    assert imported.dataset.years == (2023, 2024, 2025)
    assert imported.dataset.production_unit == {"main": "tone"}
    rows = {
        Carrier.electricity_cogen: ("energi electrica din cogenerar ", (8, 12, 16)),
        Carrier.coke: ("Consum Cocs", (8, 12, 16)),
        Carrier.diesel: ("Consum Carburanti", (9, 19, 29)),
    }
    for carrier, (sheet_name, years_rows) in rows.items():
        sheet = book[sheet_name]
        for year, row in zip(imported.dataset.years, years_rows, strict=True):
            series = imported.dataset.carriers[carrier][year]
            unit = "t" if carrier == Carrier.diesel else "MWh"
            assert series.annual is not None
            assert series.annual.value == pytest.approx(sheet[f"Q{row}"].value)
            assert series.annual.unit == unit
            assert (
                imported.located[f"carrier.{carrier.value}.{year}"].ref.a1 == f"{sheet_name}!Q{row}"
            )
            for month in range(1, 13):
                column = month + 3
                assert series.months[month].value == pytest.approx(sheet.cell(row, column).value)
                assert imported.located[f"carrier.{carrier.value}.{year}.{month:02d}"].ref.a1 == (
                    f"{sheet_name}!{sheet.cell(row, column).coordinate}"
                )
    for year, monthly_row in zip(imported.dataset.years, (7, 10, 13), strict=True):
        series = imported.dataset.production["main"][year]
        assert series.annual is not None
        assert series.annual.value == pytest.approx(book["Productii"][f"P{monthly_row}"].value)
        assert imported.located[f"production.main.{year}"].ref.a1 == f"Productii!P{monthly_row}"
        for month in range(1, 13):
            cell = book["Productii"].cell(monthly_row, month + 3)
            assert series.months[month].value == pytest.approx(cell.value)
            assert imported.located[f"production.main.{year}.{month:02d}"].ref.a1 == (
                f"Productii!{cell.coordinate}"
            )
        assert (
            specific_consumption(imported.dataset, imported.factors, year, None, "main").value
            is not None
        )

    water = {
        Carrier.water_potable: (4, 5, 6),
        Carrier.water_industrial: (28, 29, 30),
        Carrier.water_storm: (54, 55, 56),
    }
    for carrier, years_rows in water.items():
        for year, row in zip(imported.dataset.years, years_rows, strict=True):
            series = imported.dataset.carriers[carrier][year]
            assert series.annual is not None
            assert series.annual.value == pytest.approx(book["Consum apa "][f"O{row}"].value)
            assert series.annual.unit == "m3"
            assert (
                imported.located[f"carrier.{carrier.value}.{year}"].ref.a1 == f"Consum apa !O{row}"
            )
            for month in range(1, 13):
                cell = book["Consum apa "].cell(row, month + 2)
                assert series.months[month].value == pytest.approx(cell.value)
                assert imported.located[f"carrier.{carrier.value}.{year}.{month:02d}"].ref.a1 == (
                    f"Consum apa !{cell.coordinate}"
                )
            assert water_specific(imported.dataset, year, carrier, "main").value is not None


def test_piee_case_c_factors_and_annual_cross_check() -> None:
    source = case_path("piee-case-c", "prelucrare")
    imported = import_prelucrare(source)
    book = load_workbook(source, read_only=True, data_only=True)
    for carrier, unit, sheet, cell in (
        (Carrier.electricity_cogen, "MWh", "energi electrica din cogenerar ", "F4"),
        (Carrier.coke, "t", "Consum Cocs", "F4"),
        (Carrier.coke, "MWh", "Consum Cocs", "F5"),
    ):
        actual = imported.factors.tep_factor(carrier, unit, 2025)
        assert actual is not None and actual.per_unit == book[sheet][cell].value
        assert actual.source.endswith(f"!{cell}")
        assert imported.located[f"factor.tep.{carrier.value}.{unit}"].ref.a1.endswith(cell)
    assert "factor.tep.coke" not in imported.located
    assert imported.located["factor.tep.electricity_cogen"].ref.a1.endswith("!F4")
    for year, tep_rows, cogen_row, coke_co2_row in (
        (
            2023,
            {
                Carrier.natural_gas: 5,
                Carrier.electricity_grid: 6,
                Carrier.diesel: 7,
                Carrier.coke: 9,
            },
            9,
            7,
        ),
        (
            2024,
            {
                Carrier.natural_gas: 16,
                Carrier.electricity_grid: 17,
                Carrier.diesel: 18,
                Carrier.coke: 20,
            },
            13,
            16,
        ),
        (
            2025,
            {
                Carrier.natural_gas: 27,
                Carrier.electricity_grid: 28,
                Carrier.diesel: 29,
                Carrier.coke: 31,
            },
            17,
            26,
        ),
    ):
        for carrier, row in tep_rows.items():
            computed = tep(imported.dataset, imported.factors, carrier, year)
            filed = book["TEP"][f"Q{row}"].value
            assert computed.value is not None and round(computed.value, 2) == round(filed, 2)
            assert imported.filed[f"tep.{carrier.value}.{year}"].ref.a1 == f"TEP!Q{row}"
            for month in range(1, 13):
                monthly = tep(imported.dataset, imported.factors, carrier, year, month)
                cell = book["TEP"].cell(row, month + 3)
                assert monthly.value is not None and round(monthly.value, 2) == round(cell.value, 2)
                assert imported.filed[f"tep.{carrier.value}.{year}.{month:02d}"].ref.a1 == (
                    f"TEP!{cell.coordinate}"
                )
        cogen_tep = tep(imported.dataset, imported.factors, Carrier.electricity_cogen, year)
        assert cogen_tep.value is not None
        assert round(cogen_tep.value, 2) == round(
            book["energi electrica din cogenerar "][f"Q{cogen_row}"].value, 2
        )
        total = tep_total(imported.dataset, imported.factors, year)
        filed_total = imported.filed[f"tep.total.{year}"]
        assert total.value is not None and round(total.value, 2) == round(filed_total.value, 2)
        assert filed_total.ref.a1.startswith("TEP!Q")
        coke_co2 = imported.factors.co2_factor(Carrier.coke, "t", year)
        assert coke_co2 is not None
        assert coke_co2.per_unit == book["impact de mediu"][f"H{coke_co2_row}"].value
        assert imported.located[f"factor.co2.coke.{year}"].ref.a1 == (
            f"impact de mediu!H{coke_co2_row}"
        )
    assert Carrier.lpg not in imported.dataset.carriers
    assert any(
        issue.code == "carrier_all_zero" and issue.detail == "lpg" for issue in imported.issues
    )
    for year, row in zip(imported.dataset.years, (22, 23, 24), strict=True):
        gas = specific_consumption(
            imported.dataset, imported.factors, year, Carrier.natural_gas, "main"
        )
        total = specific_consumption(imported.dataset, imported.factors, year, None, "main")
        production = imported.dataset.production["main"][year].annual
        filed_total = imported.filed[f"tep.total.{year}"]
        filed_partial = imported.filed[f"specific.gas_grid_partial.{year}"]
        assert gas.value is not None
        assert round(gas.value * 1000, 4) == round(book["Productii"][f"E{row}"].value * 1000, 4)
        assert total.value is not None and production is not None
        assert total.value == pytest.approx(filed_total.value / production.value)
        assert filed_partial.ref.a1 == f"Productii!G{row}"
        assert filed_partial.value == pytest.approx(
            book["Productii"][f"E{row}"].value + book["Productii"][f"F{row}"].value
        )
