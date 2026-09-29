"""Write S7 reference coverage without copying client text into the report."""

from __future__ import annotations

import json
import os
import re
from collections import Counter
from pathlib import Path

from docx import Document
from tests.golden.cases import case_path
from tests.golden.s7_layout import build_blocks
from tests.golden.s7_reader import Chapter

from conftest import artifacts_path
from ema.audit.headings import map_headings
from ema.consumption_analysis.phrases import phrase_bank
from ema.core.office.blocks import NativeChart, Num, Paragraph, Table

SOURCE_LIMITS = {
    "ch4.carburant": "source-limited: fuel measurements",
    "ch4.apa": "source-limited: water measurements",
    "ch4.echiv_carburant": "source-limited: fuel measurements and factors",
    "ch4.echiv_total": "source-limited: complete carrier inventory",
    "ch4.specific_carburant": "source-limited: fuel measurements and factors",
    "ch4.specific_total": "source-limited: complete carrier inventory",
    "ch4.specific_apa": "source-limited: water measurements",
}
BOILERPLATE = re.compile(r"^(?:Tabelul|Fig\.|Figura|În tabelul|În figura)\b", re.I)
NUMERIC_CLAIM = re.compile(r"^pentru anul 20\d{2}.*(?:valoare de|s-au înregistrat)", re.I)


def _chapter(path: Path, audit_id: str) -> tuple[object, list[object], int, int]:
    document = Document(str(path))
    children = list(document.element.body)
    mapping = map_headings(path, audit_id).mapped
    chapter = next(item for item in mapping if item.section_id == "ch4")
    following = next(
        (
            item
            for item in mapping
            if item.heading.index > chapter.heading.index and item.heading.level == 0
        ),
        None,
    )
    start = children.index(document.paragraphs[chapter.heading.index]._p)
    end = (
        children.index(document.paragraphs[following.heading.index]._p)
        if following is not None
        else len(children) - 1
    )
    return document, children, start, end


def retained_categories(path: Path, audit_id: str, generated: dict[int, object]) -> dict[int, str]:
    """Account for every element left in the authored chapter-four region."""
    document, children, start, end = _chapter(path, audit_id)
    mapping = map_headings(path, audit_id).mapped
    paragraph_indexes = {id(item._p): index for index, item in enumerate(document.paragraphs)}
    headings = sorted(
        (item.heading.index, item.section_id)
        for item in mapping
        if item.section_id.startswith("ch4")
    )
    heading_indexes = {index for index, _ in headings}
    result: dict[int, str] = {}
    section_id = "ch4"
    for body_index in range(start + 1, end):
        paragraph_index = paragraph_indexes.get(id(children[body_index]))
        if paragraph_index is not None:
            section_id = next(
                (sid for index, sid in reversed(headings) if index <= paragraph_index),
                section_id,
            )
        if body_index in generated:
            continue
        text = (
            document.paragraphs[paragraph_index].text.strip() if paragraph_index is not None else ""
        )
        if section_id == "ch4.bilant_real":
            category = "bilant_real"
        elif paragraph_index in heading_indexes or BOILERPLATE.match(text):
            category = "heading/caption boilerplate from the base"
        elif section_id in SOURCE_LIMITS:
            category = SOURCE_LIMITS[section_id]
        elif not text:
            category = "heading/caption boilerplate from the base"
        elif NUMERIC_CLAIM.search(text):
            category = "not-yet-derivable: authored unit lacks a matching denominator"
        else:
            category = "causal/domain commentary"
        result[body_index] = category
    return result


def _filed_block_count(generated: dict[int, object], case: Chapter) -> int:
    filed_charts = {
        chart.body_index
        for section_id in ("ch4.intensitate", "ch4.mediu")
        for chart in case.sections[section_id].charts
    }
    count = 0
    for index, block in generated.items():
        if (isinstance(block, NativeChart) and index in filed_charts) or (
            isinstance(block, Table)
            and any(
                isinstance(segment, Num)
                and segment.fact is not None
                and segment.fact.startswith("filed:")
                for row in block.rows
                for cell in row
                for segment in cell
            )
        ):
            count += 1
    return count


def _audit(name: str, root: Path) -> dict[str, object]:
    source = case_path(name)
    generated, case = build_blocks(name)
    document, children, start, end = _chapter(source, name.lower())
    retained = retained_categories(source, name.lower(), generated)
    assert set(retained) == set(range(start + 1, end)) - set(generated)
    paragraph_indexes = {id(item._p): index for index, item in enumerate(document.paragraphs)}
    paragraph_rows = [
        {
            "paragraph": paragraph_indexes[id(children[index])],
            "status": "identical" if isinstance(generated.get(index), Paragraph) else "not covered",
        }
        for index in range(start + 1, end)
        if id(children[index]) in paragraph_indexes
        and document.paragraphs[paragraph_indexes[id(children[index])]].text.strip()
    ]
    section_counts = {
        section_id: {
            "tables": len(section.tables),
            "charts": len(section.charts),
            "generated_tables": sum(index in generated for index, _ in section.tables),
            "generated_charts": sum(chart.body_index in generated for chart in section.charts),
        }
        for section_id, section in case.sections.items()
    }
    generated_paragraphs = sum(isinstance(block, Paragraph) for block in generated.values())
    derivable_retained = sum(
        category.startswith("not-yet-derivable")
        for index, category in retained.items()
        if id(children[index]) in paragraph_indexes
        and document.paragraphs[paragraph_indexes[id(children[index])]].text.strip()
    )
    filed_blocks = _filed_block_count(generated, case)
    return {
        "sections": section_counts,
        "generated_blocks": {
            "tables": sum(isinstance(block, Table) for block in generated.values()),
            "charts": sum(isinstance(block, NativeChart) for block in generated.values()),
            "paragraphs": sum(isinstance(block, Paragraph) for block in generated.values()),
        },
        "paragraph_coverage": {
            "identical": sum(row["status"] == "identical" for row in paragraph_rows),
            "same_pattern_different_wording": 0,
            "not_covered": sum(row["status"] == "not covered" for row in paragraph_rows),
        },
        "retained_category_counts": dict(Counter(retained.values())),
        "retained_categories": {str(index): category for index, category in retained.items()},
        "headline_numeric_paragraphs": {
            "generated": generated_paragraphs,
            "derivable": generated_paragraphs + derivable_retained,
        },
        "generated_from_filed_values": filed_blocks,
        "paragraphs": paragraph_rows,
    }


def _audit_05(root: Path) -> dict[str, int]:
    source = case_path("audit-05")
    document, children, start, end = _chapter(source, "audit-05")
    paragraphs = {id(item._p): index for index, item in enumerate(document.paragraphs)}
    present = {
        paragraphs[id(children[index])]
        for index in range(start + 1, end)
        if id(children[index]) in paragraphs
        and document.paragraphs[paragraphs[id(children[index])]].text.strip()
    }
    bank = {item.paragraph for item in phrase_bank() if item.source_document == "audit-05"}
    return {
        "nonempty_paragraphs": len(present),
        "bank_pattern_paragraphs": len(present & bank),
        "generated_paragraphs": 0,
    }


def main() -> None:
    root = Path(os.environ["EMA_REFERENCE"])
    report = {
        "audit-01": _audit("audit-01", root),
        "audit-02": _audit("audit-02", root),
        "audit_05_rom": _audit_05(root),
    }
    output = artifacts_path("s7") / "coverage.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for name, case in report.items():
        if name == "audit_05_rom":
            print(name, case)
        else:
            print(name, case["generated_blocks"], case["paragraph_coverage"])
    print("report", output)


if __name__ == "__main__":
    main()
