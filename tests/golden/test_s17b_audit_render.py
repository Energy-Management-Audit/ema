"""S17b acceptance: the CLIENT-A1 dossier rendered into the auditor's base through Word (level 2/3).

The draft and its marker inventory; the final, approved and exported, is test_s17b_audit_final.
"""

from __future__ import annotations

import json
import os
import re
import shutil
import zipfile
from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from tests.audit_replay import CH2_DRAFT, draft_recording, support_recording
from tests.conftest import artifacts_path
from tests.golden.s17b_audit_ui_fixture import prepare_all_but
from tests.golden.test_s10b_audit_base import _references
from tests.golden.test_s15b_chapter_six_base import _synthetic_form

from ema.audit.base_identity import derive_identity
from ema.audit.base_package import package_issues
from ema.audit.base_units import heading_spans_document
from ema.audit.catalogue import CATALOGUE
from ema.audit.draft_schema import DraftText, SectionDraft
from ema.audit.draft_stage import draft_section
from ema.audit.headings import map_headings
from ema.audit.intake import audit_intake
from ema.audit.measures import run_measures
from ema.audit.read import read_job
from ema.audit.render import RenderSummary, start_audit_render
from ema.audit.render_steps import ai_wording_hits
from ema.audit.sections import Status, refresh_staleness, set_status, statuses
from ema.audit.sections_bulk import patch_sections
from ema.audit.workflow import AuditWorkflow
from ema.core.jobs import create_job, run_stage, status, subscribe
from ema.core.review import decide, fields, propose
from ema.core.review.models import Evidence, Manual
from ema.core.workspace import Workspace

pytestmark = pytest.mark.golden

CH3_DRAFT = SectionDraft(
    section="ch3.flux",
    status="drafted",
    paragraphs=[
        DraftText(
            text="Fluxul {{f:audit.process_sections}} este descris.",
            fact_ids=["audit.process_sections"],
        )
    ],
)
ORDER = {section.id: index for index, section in enumerate(CATALOGUE)}
CHAPTER = {section.id: section.chapter for section in CATALOGUE}


def _wait(ws: Workspace, job: str, run: str) -> dict[str, object]:
    for _ in subscribe(ws, job):
        pass
    return next(item for item in status(ws, job).runs if item["id"] == run)


def _supply(ws: Workspace, job: str, key: str, value: object) -> None:
    evidence = Evidence(
        id="synthetic:" + key,
        provenance="manual",
        locator=Manual(who="synthetic"),
        method="manual",
        retrieved_at=datetime.now(UTC),
        highlight="exact",
    )
    propose(ws, job, key, value, [evidence], state="supplied")


def seed_job(
    root: Path, workspace: Path, tmp_path: Path, form: Path | None = None
) -> tuple[Workspace, str]:
    """The CLIENT-A1 dossier read, two replayed drafts and a synthetic measures form (no render)."""
    received = root / "audit/cases/audit-case-a/received"
    ws = Workspace(workspace)
    job = create_job(ws, "audit", "CLIENT-A1-golden", 2026)
    for source in sorted(path for path in received.iterdir() if path.is_file()):
        ws.set_slot(job, f"dossier/{source.name}", ws.add_file("CLIENT-A1-golden", source))
    for stage, fn in (("intake", audit_intake), ("read", read_job)):
        record = _wait(ws, job, run_stage(ws, job, stage, fn))
        assert record["state"] == "ready", (stage, record["error"])
    # Synthetic identity and process facts: no client value enters the drafts.
    _supply(ws, job, "audit.company_name", "Atelier Exemplu SRL")
    _supply(ws, job, "audit.process_sections", "asamblare")
    for index, draft in enumerate((CH2_DRAFT, CH3_DRAFT)):
        draft_section(
            ws,
            job,
            draft.section,
            draft_recording=draft_recording(ws, job, draft, tmp_path / f"draft-{index}.json"),
            support_recording=support_recording(
                ws, job, draft, tmp_path / f"support-{index}.json", None
            ),
        )
    form = form or _synthetic_form(tmp_path / "m.xlsx")
    ws.set_slot(job, "measures", ws.add_file("CLIENT-A1-golden", form))
    run_measures(ws, job)
    return ws, job


def base_settings(root: Path, tmp_path: Path) -> dict[str, str]:
    """The D1 settings as environment values; the identity file is written by derive_identity."""
    base, prototype = _references(root)
    path = tmp_path / "base-identity.json"
    path.write_text(json.dumps(list(derive_identity(base)), ensure_ascii=False), encoding="utf-8")
    model = root / "audit/section-models/Fișa de măsuratori electroenergetice - model.docx"
    return {
        "EMA_AUDIT_BASE_DOCUMENT": str(base),
        "EMA_AUDIT_MEASUREMENT_PROTOTYPE": str(prototype),
        "EMA_AUDIT_MEASUREMENT_SHEET_MODEL": str(model),
        "EMA_AUDIT_BASE_IDENTITY": str(path),
    }


def _settings(root: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[str, ...]:
    values = base_settings(root, tmp_path)
    for key, value in values.items():
        monkeypatch.setenv(key, value)
    return tuple(json.loads(Path(values["EMA_AUDIT_BASE_IDENTITY"]).read_text("utf-8")))


def review_inputs(ws: Workspace, job: str) -> dict[str, str]:
    """Step 2 of the final (D8): her texts written, one reading corrected and one rejected.

    Returns the two reviewed keys; their ch. 4 cells are what the final golden checks.
    """
    for field in fields(ws, job):
        if field.key.startswith("narrative.") and field.value is None:
            decide(ws, job, field.id, "correct", field.revision, "user", value="Text verificat.")
    run_measures(ws, job)  # ch. 6 is composed again with the texts just written
    reviewed: dict[str, str] = {}
    for carrier, action in (("electricity_grid", "correct"), ("natural_gas", "reject")):
        field = min(
            (
                item
                for item in fields(ws, job)
                if re.fullmatch(rf"carrier\.{carrier}\.\d{{4}}\.\d{{2}}", item.key)
                and item.value is not None
            ),
            key=lambda item: item.key,
        )
        value = {"value": 1234.5} if action == "correct" else {}
        decide(ws, job, field.id, action, field.revision, "user", **value)  # type: ignore[arg-type]
        reviewed[action] = field.key
    return reviewed


def confirm_chapters(ws: Workspace, job: str) -> None:
    """Step 4 of the final (D8): per chapter, one confirmation of its drafts; the rest n/a."""
    refresh_staleness(ws, job)
    chapters: dict[int, list[dict[str, object]]] = {}
    for state in statuses(ws, job):
        if state.status == Status.DRAFTED and not state.stale:
            chapters.setdefault(CHAPTER[state.section_id], []).append(
                {
                    "section_id": state.section_id,
                    "status": "done",
                    "on_revision": state.revision,
                    "confirm": True,
                }
            )
    for items in chapters.values():
        patch_sections(ws, job, items)
    for state in statuses(ws, job):
        if state.status not in (Status.DONE, Status.NA):
            set_status(ws, job, state.section_id, Status.NA, "user", "golden")


def prepare_final(ws: Workspace, job: str) -> None:
    """Final readiness the way the use cases allow: texts, reviewed data, chapters confirmed."""
    review_inputs(ws, job)
    confirm_chapters(ws, job)


def _outputs(ws: Workspace, job: str, run: str) -> dict[str, Path]:
    with ws.connect() as db:
        rows = db.execute(
            "SELECT relative_path FROM outputs WHERE run_id=? ORDER BY seq", (run,)
        ).fetchall()
    paths = [ws.path(str(row["relative_path"])) for row in rows]
    return {path.name.split("-", 1)[1]: path for path in paths}


def _package_checks(docx: Path, identity: tuple[str, ...], summary: RenderSummary) -> None:
    assert not package_issues(docx, identity)
    with zipfile.ZipFile(docx) as archive:
        names = archive.namelist()
        texts = {
            name: archive.read(name).decode("utf-8", "ignore")
            for name in names
            if name.endswith((".xml", ".rels"))
        }
    for name, text in texts.items():
        assert not any(term.casefold() in text.casefold() for term in identity), name
        if name.endswith(".rels"):
            assert 'TargetMode="External"' not in text, name
    charts = [name for name in names if re.fullmatch(r"word/charts/chart\d+\.xml", name)]
    assert len(charts) == summary.charts
    for chart in charts:
        rels = texts.get(chart.replace("charts/", "charts/_rels/") + ".rels", "")
        assert len(re.findall(r"embeddings/[^\"]+\.xlsx", rels)) == 1, chart


def _heading_checks(docx: Path, base: Path, summary: RenderSummary) -> None:
    """Every heading of the plan's base is still there, in the base's own order.

    Ch. 4 is rebuilt from the catalogue, so its headings follow the catalogue; ch. 6 rewrites
    its "ch6.measure" headings with the measures' own titles, so those are left out by name.
    """

    def headings(path: Path) -> list[str]:
        return [
            item.section_id
            for item in map_headings(path, "AUDIT-01").mapped
            if item.section_id != "ch6.measure"
        ]

    rendered, planned = headings(docx), headings(base)
    assert set(rendered) <= set(ORDER)
    assert [item for item in rendered if not item.startswith("ch4.")] == [
        item for item in planned if not item.startswith("ch4.")
    ]
    four = [ORDER[item] for item in rendered if item.startswith("ch4.")]
    assert four and four == sorted(four)
    assert [chapter.section_id for chapter in summary.chapters] == [
        item for item in rendered if "." not in item
    ]


def _toc_pages(docx: Path) -> list[str]:
    body = Document(str(docx)).element.body
    results = []
    for paragraph in body.iter(qn("w:p")):
        style = paragraph.find(f"./{qn('w:pPr')}/{qn('w:pStyle')}")
        if style is not None and (style.get(qn("w:val")) or "").startswith("TOC"):
            texts = [node.text or "" for node in paragraph.iter(qn("w:t"))]
            results.append(texts[-1] if texts else "")
    return results


def _marker_inventory(docx: Path) -> list[dict[str, object]]:
    """Structure only: where each marker sits and which anchor holds it, never its text."""
    document = Document(str(docx))
    spans = heading_spans_document(document)
    rows: dict[tuple[str, str, str], int] = {}
    for index, element in enumerate(document.element.body):
        owner = max(
            (span for span in spans if span[1] <= index < span[2]),
            key=lambda span: span[1],
            default=None,
        )
        section = owner[0].section_id if owner else "front"
        for paragraph in element.iter(qn("w:p")):
            count = "".join(node.text or "" for node in paragraph.iter(qn("w:t"))).count(
                "[de completat]"
            )
            if not count:
                continue
            style = paragraph.find(f"./{qn('w:pPr')}/{qn('w:pStyle')}")
            kind = (
                "toc"
                if style is not None and (style.get(qn("w:val")) or "").startswith("TOC")
                else "heading"
                if owner is not None and owner[1] == index
                else "table-cell"
                if element.tag == qn("w:tbl")
                else "paragraph"
            )
            slots = ",".join(
                (node.get(qn("w:name")) or "")[5:]
                for node in paragraph.iter(qn("w:bookmarkStart"))
                if (node.get(qn("w:name")) or "").startswith("_ema_")
            )
            key = (section, kind, slots)
            rows[key] = rows.get(key, 0) + count
    return [
        {"section": section, "kind": kind, "anchor": slots, "count": count}
        for (section, kind, slots), count in rows.items()
    ]


def _delivered() -> list[Path]:
    root = os.environ.get("EMA_REFERENCE")
    folder = Path(root) / "audit/finished-audits" if root else None
    return sorted(folder.glob("*.docx")) if folder and folder.is_dir() else []


@pytest.mark.parametrize("audit", _delivered(), ids=lambda path: path.name)
def test_no_ai_wording_in_delivered_audits(audit: Path) -> None:
    """The final gate's AI rule on every audit the auditor delivered: no false alarm (fix round 1)."""
    assert ai_wording_hits(audit, []) == []


def test_CLIENT-A1_audit_render_through_word(
    reference_library: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    identity = _settings(reference_library, tmp_path, monkeypatch)
    ws, job = seed_job(reference_library, tmp_path / "workspace", tmp_path)
    record = _wait(ws, job, start_audit_render(ws, job, "draft"))
    assert record["state"] == "ready", record["error"]
    run = str(record["id"])
    with ws.connect() as db:
        folder = ws.artifact_dir(db, job, "audit_render", run)
    summary = RenderSummary.model_validate_json((folder / "render.json").read_text("utf-8"))
    outputs = _outputs(ws, job, run)
    assert list(outputs) == ["Audit-ciorna.pdf", "Audit-ciorna.docx"]
    docx, pdf = outputs["Audit-ciorna.docx"], outputs["Audit-ciorna.pdf"]
    # D9.1: Word opened it without repair (open_check), numbered the TOC and made the PDF.
    assert summary.pdf and summary.toc_pages_set and pdf.stat().st_size > 0
    assert pdf.read_bytes().startswith(b"%PDF")
    pages = _toc_pages(docx)
    assert pages and all(text.strip().isdigit() and int(text) > 0 for text in pages)
    assert all(chapter.page for chapter in summary.chapters)
    _package_checks(docx, identity, summary)

    _heading_checks(docx, folder / "base.docx", summary)
    assert ai_wording_hits(docx, []) == []  # her own base text never trips the AI rule
    assert not summary.failures
    assert summary.unit_plan.processes_source == "schemes"
    inventory = _marker_inventory(docx)
    out = artifacts_path("s17b-audit-report", "golden")
    out.mkdir(parents=True, exist_ok=True)
    (out / "marker-inventory.json").write_text(json.dumps(inventory, indent=2), "utf-8")
    shutil.copyfile(docx, out / "Audit-ciorna.docx")
    if parity := os.environ.get("EMA_PARITY_OUT"):
        Path(parity).mkdir(parents=True, exist_ok=True)
        for name in ("Audit-ciorna.docx", "Audit-ciorna.pdf"):
            shutil.copyfile(outputs[name], Path(parity) / name)

    print(
        f"S17b audit render: chapters={len(summary.chapters)} tables={summary.tables} "
        f"charts={summary.charts} markers={len(summary.markers)} "
        f"fields={summary.fields_confirmed}/{summary.fields_total} pages={len(pages)}"
    )


if __name__ == "__main__":
    # The UI golden seeds through this: `python -m tests.golden.test_s17b_audit_render
    # seed|final <workspace> <scratch>`; it prints only the job id and the settings' paths.
    import sys

    command, workspace, scratch = sys.argv[1], Path(sys.argv[2]), Path(sys.argv[3])
    library = Path(os.environ["EMA_REFERENCE"])
    if command == "seed":
        from tests.golden.test_s17b_audit_final import _complete_form

        settings = base_settings(library, scratch)
        os.environ.update(settings)
        _, seeded = seed_job(library, workspace, scratch, _complete_form(scratch / "all.xlsx"))
        print(json.dumps({"job": seeded, "env": settings}))
    elif command == "review-ui":
        from tests.golden.test_s17b_audit_final import _cover_photo

        from ema.audit.render_bindings import COVER_SLOT

        ws = Workspace(workspace)
        _supply(ws, sys.argv[4], "audit.address", "Str. Exemplu nr. 1, Localitatea Exemplu")
        ws.set_slot(
            sys.argv[4], COVER_SLOT, ws.add_file("CLIENT-A1-golden", _cover_photo(scratch / "cover.png"))
        )
        review_inputs(ws, sys.argv[4])
        print(json.dumps({"reviewed": True}))
    elif command == "final-ui":
        ws = Workspace(workspace)
        left = prepare_all_but(ws, sys.argv[4], "ch1")
        print(
            json.dumps(
                {
                    "left": [
                        {"section_id": item.section_id, "revision": item.revision} for item in left
                    ],
                    "final_ok": False,
                }
            )
        )
    else:
        prepare_final(Workspace(workspace), sys.argv[4])
        print(
            json.dumps(
                {"final_ok": AuditWorkflow().readiness(Workspace(workspace), sys.argv[4]).final_ok}
            )
        )
