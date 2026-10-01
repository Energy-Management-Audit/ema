"""Source-label evidence for case C's missing identity and production fields."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from tests.golden.cases import case_path

from ema.energy_data.source import normal
from ema.piee.dataset import PieeData


def _labelled(kind: str, labels: tuple[str, ...]) -> list[tuple[object, str]]:
    wanted = {normal(label) for label in labels}
    book = load_workbook(case_path("piee-case-c", kind), read_only=True, data_only=True)
    found: list[tuple[object, str]] = []
    for sheet in book:
        for row in sheet.iter_rows():
            for index, cell in enumerate(row):
                if not isinstance(cell.value, str) or normal(cell.value) not in wanted:
                    continue
                for following in row[index + 1 : index + 4]:
                    if following.value is not None and str(following.value).strip():
                        found.append((following.value, f"{sheet.title}!{following.coordinate}"))
                        break
    book.close()
    return found


def assert_missing_identity_is_unsourced(data: PieeData) -> None:
    assert "registrul_comertului" not in data.anexa.identity
    assert not _labelled("anexa", ("Registrul Comerțului", "Nr. Registrul Comerțului"))
    assert not _labelled("prelucrare", ("Registrul Comerțului", "Nr. Registrul Comerțului"))
    assert "ownership_state" not in data.anexa.identity
    state = _labelled("anexa", ("Stat",))
    assert len(state) == 1 and normal(str(state[0][0])) == "xxx"
    assert not _labelled("prelucrare", ("Stat",))
    assert "ownership_private" in data.anexa.identity

    assert "site_2_address" not in data.anexa.identity
    assert "site_1_production_share" not in data.anexa.identity
    assert "site_2_production_share" not in data.anexa.identity
    site_two = normal(str(data.anexa.identity["site_2_name"].value))
    addresses = _labelled("anexa", ("Adresa poștală", "Adresa poştală"))
    assert len(addresses) == 1 and site_two not in normal(str(addresses[0][0]))
    assert not _labelled("prelucrare", ("Adresa poștală", "Adresa poştală"))
    for kind in ("anexa", "prelucrare"):
        assert not _labelled(
            kind,
            (
                "Adresa sucursalei 2",
                "Ponderea producției sucursalei 1",
                "Ponderea producției sucursalei 2",
            ),
        )


def assert_production_name_is_sourced(data: PieeData, output: Path) -> None:
    assert data.prelucrare is not None
    found = _labelled("prelucrare", ("Productie",))
    assert len(found) == 1
    source, reference = found[0]
    assert isinstance(source, str) and normal(source).startswith("productia de ")
    chosen = data.prelucrare.located["production_name.main"]
    assert chosen.ref.a1 == reference
    assert data.dataset.production_name["main"] == chosen.value
    by_slot = {
        mark.get(qn("w:name"), "").removeprefix("_ema_"): paragraph.text
        for paragraph in Document(output).paragraphs
        for mark in paragraph._p.iter(qn("w:bookmarkStart"))
    }
    for slot in ("body_30", "body_294"):
        assert str(chosen.value) in by_slot[slot]
        assert "n.d." not in by_slot[slot]
