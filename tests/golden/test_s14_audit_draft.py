"""S14 synthetic request-bound replay and structure comparison with auditor audits."""

import json
import os
from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from tests.audit_replay import draft_recording, render_draft_section, support_recording
from tests.golden.s14_structure import auditor_structure
from tests.golden.test_s10b_audit_base import _identity, _references
from tests.workspace_jobs import create_job

from ema.audit.base import build_base
from ema.audit.base_units import UnitPlan
from ema.audit.catalogue import CATALOGUE
from ema.audit.draft_agent import (
    draft_section_replay,
    recorded_facts,
)
from ema.audit.draft_checks import check_draft
from ema.audit.draft_render import render_section
from ema.audit.draft_schema import DraftText, SectionDraft
from ema.audit.sections import get_status, refresh_staleness
from ema.core.llm import Limits
from ema.core.office.anchors import find
from ema.core.review.fields import propose
from ema.core.review.models import Evidence, Field, Manual
from ema.core.workspace import Workspace

pytestmark = [pytest.mark.golden, pytest.mark.word]
TITLES = {section.id: section.title for section in CATALOGUE}


def _base(tmp_path: Path) -> Path:
    """Each drafted section's own region: its heading, then what it replaces."""
    base = tmp_path / "base.docx"
    document = Document()
    for text, style in (
        (TITLES["ch2"], "Heading 1"),
        (TITLES["ch2.date_generale"], "Heading 2"),
        ("[de completat]", None),
        ("[de completat]", None),
        (TITLES["ch3"], "Heading 1"),
        (TITLES["ch3.flux"], "Heading 2"),
        ("[de completat]", None),
        (TITLES["ch3.utilitati"], "Heading 2"),
    ):
        document.add_paragraph(text, style=style)
    document.save(str(base))
    return base


def test_synthetic_chapters_replay_render_and_compare(tmp_path: Path) -> None:
    reference = os.environ.get("EMA_REFERENCE")
    assert reference and Path(reference).is_dir(), "EMA_REFERENCE is required for this golden"
    structure = auditor_structure(Path(reference))
    assert len(structure) == 5
    assert all(item["sections"] for item in structure.values())
    ws = Workspace(tmp_path / "workspace")
    job = create_job(ws, "audit", "synthetic", 2026)
    for key, value in (
        ("audit.company_name", "Atelier Exemplu"),
        ("audit.employees", 85),
        ("audit.process_sections", "asamblare"),
    ):
        evidence = Evidence(
            id="synthetic:" + key,
            provenance="manual",
            locator=Manual(who="synthetic"),
            method="manual",
            retrieved_at=datetime.now(UTC),
            highlight="exact",
        )
        propose(ws, job, key, value, [evidence], state="supplied")
    drafts = [
        SectionDraft(
            section="ch2.date_generale",
            status="drafted",
            paragraphs=[
                DraftText(
                    text="Societatea {{f:audit.company_name}} are {{f:audit.employees}} angajați.",
                    fact_ids=["audit.company_name", "audit.employees"],
                ),
                DraftText(
                    text=(
                        "Societatea {{f:audit.company_name}} utilizează "
                        "echipamente moderne, eficiente."
                    ),
                    fact_ids=["audit.company_name"],
                ),
            ],
        ),
        SectionDraft(
            section="ch3.flux",
            status="drafted",
            paragraphs=[
                DraftText(
                    text="Fluxul {{f:audit.process_sections}} este descris.",
                    fact_ids=["audit.process_sections"],
                )
            ],
        ),
    ]
    rejected = check_draft(
        SectionDraft(
            section="ch2.date_generale",
            status="drafted",
            paragraphs=[
                DraftText(
                    text="Societatea {{f:audit.company_name}} are 99 angajați.",
                    fact_ids=["audit.company_name"],
                )
            ],
        ),
        recorded_facts(ws, job, "ch2.date_generale"),
        job,
    )
    assert [issue.code for issue in rejected.fatal] == ["literal_number"]
    rendered = _base(tmp_path)
    coverage = []
    for index, draft in enumerate(drafts):
        state, replayed, checked, flags = draft_section_replay(
            ws,
            job,
            draft.section,
            draft_recording(ws, job, draft, tmp_path / f"draft-{index}.json"),
            support_recording(
                ws, job, draft, tmp_path / f"support-{index}.json", 1 if index == 0 else None
            ),
            Limits(8),
        )
        assert state.status == "done" and not checked.fatal
        assert len(flags) == (1 if index == 0 else 0)
        coverage.append(checked.coverage)
        output = tmp_path / f"rendered-{index}.docx"
        render_draft_section(
            ws,
            job,
            rendered,
            output,
            draft=replayed,
            facts=recorded_facts(ws, job, draft.section),
            flags=flags,
        )
        review = json.loads(output.with_suffix(".draft-review.json").read_text(encoding="utf-8"))
        assert review["coverage"] == 1.0
        assert len(review["review"]) == (1 if index == 0 else 0)
        rendered = output
    texts = [p.text for p in Document(str(rendered)).paragraphs]
    general = texts.index(TITLES["ch2.date_generale"]) + 1
    assert "Atelier Exemplu" in texts[general] and "85" in texts[general]
    assert texts[general + 1] == "[de completat]"  # unsupported prose withheld
    assert "asamblare" in texts[texts.index(TITLES["ch3.flux"]) + 1]
    assert coverage == [1.0, 1.0]
    synthetic_structure = {
        draft.section: {"paragraphs": len(draft.paragraphs), "tables": len(draft.tables)}
        for draft in drafts
    }
    propose(
        ws,
        job,
        "audit.process_sections",
        "asamblare și ambalare",
        [
            Evidence(
                id="synthetic:process-update",
                provenance="manual",
                locator=Manual(who="synthetic"),
                method="manual",
                retrieved_at=datetime.now(UTC),
                highlight="exact",
            )
        ],
        state="supplied",
    )
    refresh_staleness(ws, job)
    assert get_status(ws, job, "ch3.flux").stale
    assert not get_status(ws, job, "ch2.date_generale").stale
    print(
        f"S14 synthetic: coverage={coverage}; support_flags=1; "
        f"comparator_audits={len(structure)}; evidence=level 1; "
        f"structure={json.dumps({'synthetic': synthetic_structure, 'auditor': structure})}"
    )


def test_synthetic_paragraph_in_auditor_base(tmp_path: Path) -> None:
    reference = os.environ.get("EMA_REFERENCE")
    assert reference and Path(reference).is_dir(), "EMA_REFERENCE is required for this golden"
    source, prototype = _references(Path(reference))
    base, output = tmp_path / "base.docx", tmp_path / "rendered.docx"
    build_base(
        UnitPlan("Atelier Exemplu", 1, frozenset({"electricity", "gas"}), 0, False, 0, 0),
        base_document=source,
        measurement_prototype=prototype,
        output=base,
        base_identity=_identity(source),
    )
    anchors = base.with_suffix(".anchors.json")
    mapping = json.loads(anchors.read_text(encoding="utf-8"))
    other_slot = next(
        item["slot"]
        for item in mapping["anchors"]
        if item["section"] == "ch3.flux" and item["classification"] == "variable"
    )
    before = find([Document(str(base)).element], other_slot)
    before_text = "".join(node.text or "" for node in before.iter(qn("w:t")))
    fact = {
        "audit.company_name": Field(
            id="synthetic",
            job_id="synthetic",
            key="audit.company_name",
            label="Company",
            value_type="text",
            value="Atelier Exemplu",
            state="supplied",
            presence="found",
            evidence=["synthetic-evidence"],
        )
    }
    draft = SectionDraft(
        section="ch2.date_generale",
        status="drafted",
        paragraphs=[
            DraftText(text="Societatea {{f:audit.company_name}}.", fact_ids=["audit.company_name"])
        ],
    )
    render_section(base, output, draft, fact, (), job="synthetic")
    after_document = Document(str(output))
    after = find([after_document.element], other_slot)
    assert "".join(node.text or "" for node in after.iter(qn("w:t"))) == before_text
    assert any(
        "Societatea Atelier Exemplu." in paragraph.text for paragraph in after_document.paragraphs
    )
