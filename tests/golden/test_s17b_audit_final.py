"""S17b: the case A final passes audit_markers, approval and export (D8, Word).

Every marker is cleared by a source: her reviewed base text, a field, the render date, a
human-confirmed section or a text she wrote. Client values are compared, never printed.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import zipfile
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from tests.audit_replay import CH2_DRAFT, draft_recording, support_recording
from tests.conftest import artifacts_path
from tests.golden.cases import case_path
from tests.golden.s17b_chart_oracle import assert_final_charts, expected_charts
from tests.golden.s17c_chart_layout_oracle import assert_production_values_fit_pdf
from tests.golden.s17c_structure_oracle import assert_final_structure
from tests.golden.test_s17b_audit_render import (
    _outputs,
    _package_checks,
    _settings,
    _supply,
    _toc_pages,
    _wait,
    confirm_chapters,
    review_inputs,
    seed_job,
)

from ema.api import create_app
from ema.audit.base_numeric import _ALLOWED  # pyright: ignore[reportPrivateUsage]
from ema.audit.base_units import heading_spans_document
from ema.audit.catalogue import CATALOGUE
from ema.audit.chapter_four import MONTHS
from ema.audit.chapter_six import MEASURE_HEADER, SYNTHESIS_HEADER
from ema.audit.draft_stage import draft_section
from ema.audit.measures_form import write_measures_form
from ema.audit.render import RenderSummary, start_audit_render
from ema.audit.render_bindings import COVER_SLOT
from ema.audit.render_dataset import reviewed_dataset
from ema.audit.render_sections import RENDER_DRAFTED
from ema.audit.render_steps import body_counts, section_ids
from ema.audit.sections import Status, statuses
from ema.audit.workflow import AuditWorkflow
from ema.core.review import fields
from ema.core.workspace import Workspace
from ema.energy_data.necesar import parse_necesar_info, to_dataset

pytestmark = [pytest.mark.golden, pytest.mark.word]
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
FIXED = {section.id for section in CATALOGUE if section.kind == "fixed"}
ADDRESS = "Str. Exemplu nr. 1, Localitatea Exemplu"


def _complete_form(path: Path) -> Path:
    """Three synthetic measures with every column filled: no payback is left missing."""
    write_measures_form(path)
    book = load_workbook(path)
    sheet = book["Măsuri propuse"]
    sheet.append(
        [
            1,
            "Lighting retrofit",
            "Less electricity",
            "Energie electrică",
            100,
            "MWh",
            40,
            10,
            "Offer",
        ]
    )
    sheet.append([2, "Gas recovery", "Less gas", "Gaze naturale", 10, "MWh", 20, 5, "Estimate"])
    sheet.append([3, "Fleet change", "Less diesel", "Motorină", 1, "t", 50, 25, "Quote"])
    book.save(path)
    return path


def _assert_toc_is_upright(path: Path) -> None:
    document = Document(str(path))
    styles = {
        style.get(qn("w:styleId")): style
        for style in document.styles.element.findall(qn("w:style"))
    }
    entries = [
        p
        for p in document.element.body.iter(qn("w:p"))
        if p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}") is not None
        and p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}").get(qn("w:val"), "").startswith("TOC")
    ]
    assert entries
    for paragraph in entries:
        style_id = paragraph.find(f"{qn('w:pPr')}/{qn('w:pStyle')}").get(qn("w:val"))
        style = styles[style_id]
        italic_nodes = style.xpath("./w:rPr/w:i | ./w:rPr/w:iCs") + paragraph.xpath(
            ".//w:i | .//w:iCs"
        )
        assert all(node.get(qn("w:val")) not in {None, "1", "true", "on"} for node in italic_nodes)


def _text(element: Any) -> str:
    return "".join(node.text or "" for node in element.iter(W + "t"))


def _region(document: Any, section_id: str) -> list[Any]:
    spans = heading_spans_document(document)
    start = next(start for item, start, _ in spans if item.section_id == section_id)
    end = min((begin for _, begin, _ in spans if begin > start), default=len(document.element.body))
    return list(document.element.body)[start + 1 : end]


def _cell(document: Any, section_id: str, key: str) -> str:
    """The ch. 4 monthly table cell of `carrier.<c>.<year>.<month>`."""
    year, month = (int(part) for part in key.split(".")[2:4])
    tables = [item for item in _region(document, section_id) if item.tag == W + "tbl"]
    table = tables[0] if month <= 6 else tables[1]
    row = next(
        tr for tr in table.findall(W + "tr") if _text(tr.findall(W + "tc")[0]).strip() == str(year)
    )
    return _text(row.findall(W + "tc")[(month - 1) % 6 + 1]).strip()


def _header(table: Any, rows: int) -> list[list[str]]:
    return [
        [_text(tc).strip() for tc in tr.findall(W + "tc")] for tr in table.findall(W + "tr")[:rows]
    ]


def _bound(docx: Path, anchors: Path) -> dict[str, list[str]]:
    """Each binding's filled paragraphs, found by the anchor bookmark the base recorded."""
    wanted = {
        f"_ema_{item['slot']}": item["binding"]
        for item in json.loads(anchors.read_text(encoding="utf-8"))["anchors"]
        if item.get("binding")
    }
    document = Document(str(docx))
    roots = [document.element] + [section.header._element for section in document.sections]
    found: dict[str, list[str]] = {}
    for root in roots:
        for node in root.iter(qn("w:bookmarkStart")):
            if (binding := wanted.get(node.get(qn("w:name"), ""))) is not None:
                found.setdefault(binding, []).append(_text(node.getparent()).strip())
    return found


def _cover_photo(path: Path) -> Path:
    Image.new("RGB", (640, 480), "seagreen").save(path, "PNG")
    return path


def _images_are_reviewed_or_uploaded(docx: Path, photo: Path) -> None:
    """Round 1: every picture in the final is her reviewed bytes or the job's own photo."""
    allowed = set().union(*_ALLOWED.values()) | {hashlib.sha256(photo.read_bytes()).hexdigest()}
    with zipfile.ZipFile(docx) as archive:
        media = {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("word/media/")
        }
    assert media and all(digest in allowed for digest in media.values()), sorted(
        name for name, digest in media.items() if digest not in allowed
    )
    assert hashlib.sha256(photo.read_bytes()).hexdigest() in media.values()


def _draft_marks(ws: Workspace, job: str) -> tuple[list[str], set[str]]:
    """Step 3: the draft renders without failure and marks what it wrote itself."""
    # Review changes recorded inputs; replay CH2 against the current facts before rendering.
    directory = ws.root / "final-recordings"
    directory.mkdir(exist_ok=True)
    draft_section(
        ws,
        job,
        CH2_DRAFT.section,
        draft_recording=draft_recording(ws, job, CH2_DRAFT, directory / "ch2.json"),
        support_recording=support_recording(
            ws, job, CH2_DRAFT, directory / "ch2-support.json", None
        ),
    )
    before = {state.section_id: state.status for state in statuses(ws, job)}
    draft = _wait(ws, job, start_audit_render(ws, job, "draft"))
    assert draft["state"] == "ready", draft["error"]
    with ws.connect() as db:
        folder = ws.artifact_dir(db, job, "audit_render", str(draft["id"]))
    summary = RenderSummary.model_validate_json(
        (folder / "render.json").read_text(encoding="utf-8")
    )
    assert summary.failures == []
    present = set(section_ids(folder / "Audit-ciorna.docx"))
    after = {state.section_id: state.status for state in statuses(ws, job)}
    marked = [key for key in RENDER_DRAFTED if key in present and before[key] == Status.READY]
    assert marked and all(after[key] == Status.DRAFTED for key in marked)
    return marked, present


def _document_checks(
    ws: Workspace, job: str, docx: Path, anchors: Path, reviewed: dict[str, str]
) -> None:
    """Step 5: every bound cell holds its source, her headers and titles, the reviewed cells."""
    _assert_toc_is_upright(docx)
    document = Document(str(docx))
    today = date.today()
    address = next(field for field in fields(ws, job) if field.key == "audit.address")
    with ws.connect() as db:
        evidence = [
            json.loads(row["data"])["provenance"]
            for row in db.execute("SELECT data FROM evidence WHERE job_id=?", (job,))
            if json.loads(row["data"])["id"] in address.evidence
        ]
    assert (address.value, evidence) == (ADDRESS, ["manual"])
    bound = _bound(docx, anchors)
    assert bound["client_name"][0] == "Atelier Exemplu SRL"
    assert bound["address"] == [ADDRESS]
    assert bound["report_month"] == [f"{MONTHS[today.month - 1]} {today.year}"]
    assert bound["report_year"] and all(str(today.year) in text for text in bound["report_year"])
    with zipfile.ZipFile(docx) as archive:
        header = archive.read("word/header1.xml").decode("utf-8")
    assert "Atelier Exemplu SRL" in header and "[de completat]" not in header
    spans = heading_spans_document(document)
    flux = next(start for item, start, _ in spans if item.section_id == "ch3.flux")
    title = re.sub(r"^\s*[\d.]+\s*", "", _text(document.element.body[flux]))
    assert title == next(section.title for section in CATALOGUE if section.id == "ch3.flux").upper()
    measure = next(item for item in _region(document, "ch6.measure") if item.tag == W + "tbl")
    synthesis = next(item for item in _region(document, "ch6.sinteza") if item.tag == W + "tbl")
    assert _header(measure, 2) == [list(row) for row in MEASURE_HEADER]
    assert _header(synthesis, 2) == [list(row) for row in SYNTHESIS_HEADER]
    corrected = _cell(document, "ch4.electricitate", reviewed["correct"]) == "1.234,50"
    rejected = _cell(document, "ch4.gaz", reviewed["reject"]) == "—"
    assert corrected and rejected, (corrected, rejected)


def _export(ws: Workspace, job: str) -> Path:
    """Step 6: approved and exported through the API in a human session."""
    client = TestClient(
        create_app(ws, 8766, launch_code="golden"), base_url="http://127.0.0.1:8766"
    )
    headers = {"X-Ema-CSRF": client.post("/session", json={"code": "golden"}).json()["csrf"]}
    output_id = AuditWorkflow().render(ws, job, "final")
    checks = client.get(f"/jobs/{job}/export/checks").json()
    assert checks["readiness"]["final_ok"], checks["readiness"]["blocking"]
    body = {
        "dest_dir": None,
        "output_id": output_id,
        "readiness_hash": checks["readiness_hash"],
        "confirm": True,
    }
    exported = client.post(f"/jobs/{job}/export", json=body, headers=headers)
    assert exported.status_code == 200, exported.json()
    assert len(client.get(f"/jobs/{job}/approvals").json()) == 1
    copy = Path(
        next(file["path"] for file in exported.json()["files"] if file["name"].endswith(".docx"))
    )
    assert copy.read_bytes() == client.get(f"/jobs/{job}/outputs/{output_id}").content
    return copy


def test_audit_case_a_final_approved_and_exported(
    reference_library: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    identity = _settings(reference_library, tmp_path, monkeypatch)
    form = _complete_form(tmp_path / "all.xlsx")
    ws, job = seed_job(reference_library, tmp_path / "workspace", tmp_path, form)
    # The case came without an Anexa: the auditor enters the address in Revizuire.
    _supply(ws, job, "audit.address", ADDRESS)
    photo = _cover_photo(tmp_path / "cover.png")  # the job's own cover photo, synthetic
    ws.set_slot(job, COVER_SLOT, ws.add_file("audit_case_a-golden", photo))
    reviewed = review_inputs(ws, job)  # 2. her texts, one reading corrected, one rejected
    marked, present = _draft_marks(ws, job)  # 3.
    confirm_chapters(ws, job)  # 4. one confirmation per chapter, the rest n/a by the user
    states = {state.section_id: state.status for state in statuses(ws, job)}
    na = {key for key, value in states.items() if value == Status.NA}
    kept = (FIXED | {"ch3", "ch6"} | {key for key in present if key.startswith("ch4")}) & present
    assert not na & kept, sorted(na & kept)

    readiness = AuditWorkflow().readiness(ws, job)  # 5. the final
    assert states[CH2_DRAFT.section] == Status.DONE
    assert readiness.final_ok, [issue.message for issue in readiness.blocking]
    final = _wait(ws, job, AuditWorkflow().start_final(ws, job))
    assert final["state"] == "ready", final["error"]
    run = str(final["id"])
    with ws.connect() as db:
        folder = ws.artifact_dir(db, job, "audit_final", run)
    summary = RenderSummary.model_validate_json(
        (folder / "render.json").read_text(encoding="utf-8")
    )
    outputs = _outputs(ws, job, run)
    docx, pdf = outputs["Audit-final.docx"], outputs["Audit-final.pdf"]
    assert body_counts(docx).markers == []
    _package_checks(docx, identity, summary)
    toc_pages = _toc_pages(docx)
    assert summary.toc_pages_set and toc_pages and all(text.strip().isdigit() for text in toc_pages)
    pages, toc_entries = assert_final_structure(docx, pdf)
    assert_production_values_fit_pdf(docx, pdf)
    assert pages > 0 and toc_entries == len(toc_pages)
    _document_checks(ws, job, docx, folder / "base.anchors.json", reviewed)
    _images_are_reviewed_or_uploaded(docx, photo)
    copy = _export(ws, job)  # 6.

    out = artifacts_path("s17b-audit-final", "golden")  # 7. the documents and counts only
    out.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(docx, out / "Audit-final.docx")
    shutil.copyfile(pdf, out / "Audit-final.pdf")
    review = artifacts_path("s17c", "structure")
    review.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(pdf, review / "Audit-final.pdf")
    shutil.copyfile(docx, review / "Audit-final.docx")
    shutil.copyfile(folder / "digest-changes.txt", review.parent / "digest-changes.txt")
    p1_digest = artifacts_path("s19", "digest-changes.txt")
    p1_digest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(folder / "digest-changes.txt", p1_digest)
    done = sum(value == Status.DONE for value in states.values())
    evidence = {
        "chapters": len(summary.chapters),
        "tables": summary.tables,
        "charts": summary.charts,
        "markers": len(summary.markers),
        "done": done,
        "na": len(na),
        "pages": pages,
        "toc_entries": toc_entries,
        "exported": True,
        "render_drafted": len(marked),
        "reviewed_keys": reviewed,
    }
    (out / "final-checks.json").write_text(json.dumps(evidence, indent=2), "utf-8")
    print(
        f"S17b audit final: chapters={len(summary.chapters)} tables={summary.tables} "
        f"charts={summary.charts} markers=0 done={done} na={len(na)} "
        f"pages={pages} toc_entries={toc_entries} "
        f"exported={copy.name}"
    )


def test_audit_case_a_final_charts(
    reference_library: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _settings(reference_library, tmp_path, monkeypatch)
    form = _complete_form(tmp_path / "all.xlsx")
    ws, job = seed_job(reference_library, tmp_path / "workspace", tmp_path, form)
    _supply(ws, job, "audit.address", ADDRESS)
    photo = _cover_photo(tmp_path / "cover.png")
    ws.set_slot(job, COVER_SLOT, ws.add_file("audit_case_a-golden", photo))
    reviewed = review_inputs(ws, job)
    _draft_marks(ws, job)
    confirm_chapters(ws, job)
    final = _wait(ws, job, AuditWorkflow().start_final(ws, job))
    assert final["state"] == "ready", final["error"]
    with ws.connect() as db:
        folder = ws.artifact_dir(db, job, "audit_final", str(final["id"]))
    summary = RenderSummary.model_validate_json(
        (folder / "render.json").read_text(encoding="utf-8")
    )
    docx = _outputs(ws, job, str(final["id"]))["Audit-final.docx"]
    source = next(
        (reference_library / case_path("audit-case-a", "received")).glob("*Necesar info*.xls")
    )
    dataset = reviewed_dataset(to_dataset(parse_necesar_info(source)), fields(ws, job))
    expected = expected_charts(dataset, "Atelier Exemplu SRL")
    assert_final_charts(docx, expected)
    assert summary.charts == len(expected)
    assert summary.charts_skipped == [
        "ch4.electricitate_pv:electricity_pv",
        "ch4.echiv_pv:electricity_pv",
        "ch4.specific_pv:electricity_pv",
        "ch4.carburant:2023:no_data",
        "ch4.carburant:2024:no_data",
        "ch4.carburant:2025:no_data",
        "ch4.apa:water_potable:2023:no_data",
        "ch4.apa:water_potable:2024:no_data",
        "ch4.apa:water_potable:2025:no_data",
        "ch4.apa:water_potable:annual:no_data",
        "ch4.specific_apa:water_potable:annual:no_data",
        "ch4.mediu:annual:no_data",
    ]
    for action, carrier, point in (("correct", "electricitate", 1234.5), ("reject", "gaz", None)):
        year, month = (int(part) for part in reviewed[action].split(".")[-2:])
        chart = next(
            item
            for item in expected
            if f"anului {year}" in item.caption
            and ("din SEN" if carrier == "electricitate" else "gaz natural") in item.caption
        )
        assert chart.series[0].values[month - 1] == point
