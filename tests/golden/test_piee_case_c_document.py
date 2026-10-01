"""PIEE case C composition against the delivered native figures and sourced slips."""

from __future__ import annotations

import json
import re
import shutil
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from tests.golden.cases import case_path
from tests.golden.piee_case_c_comparison import (
    _caption_kinds,
    _compare_annual_tables,
    _compare_centralizer,
    _compare_figures,
    _compare_numbered_tables,
    _order,
)
from tests.golden.piee_case_c_sources import (
    assert_missing_identity_is_unsourced,
    assert_production_name_is_sourced,
)
from tests.golden.piee_case_c_structure import assert_section_sequence

from conftest import artifacts_path
from ema.consumption_analysis.analysis import Metric, value
from ema.core.office.chart_series import read_series
from ema.core.office.package import (
    C,
    check_standalone,
    read_parts,
    xml,
)
from ema.energy_data.carriers import Carrier
from ema.energy_data.source import normal
from ema.piee.compose import compose_draft, load_approved_base
from ema.piee.dataset import PieeData, load
from ema.piee.number import prototype_number

pytestmark = pytest.mark.golden

NARRATIVE_SLOTS = {
    "body_77": "client electricity-use narrative, no source in the inputs",
    "body_78": "client equipment-use narrative, no source in the inputs",
    "body_79": "client lighting narrative, no source in the inputs",
    "body_81": "client ventilation narrative, no source in the inputs",
    "body_165": "client gas-use narrative, no source in the inputs",
    "body_201": "client fuel-use narrative, no source in the inputs",
    "body_363": "client audit narrative, no source in the inputs",
    "body_390": "client management narrative, no source in the inputs",
    "body_414": "client audit-history narrative, no full source in the inputs",  # #47
}
MISSING_FIELD_KEYS = {
    "body_11": "identity.registrul_comertului",
    "body_14": "identity.ownership_state",  # #47: private share exists; state share is missing.
    "site_2_address": "identity.site_2_address",
    "site_1_share": "identity.site_1_production_share",
    "site_2_share": "identity.site_2_production_share",
}
CLONE_ROLES = frozenset(
    {
        "heading",
        "intro",
        "monthly",
        "table_intro",
        "table_caption",
        "table",
        "annual_intro",
        "annual_chart",
        "annual_caption",
        "annual_observation",
        "annual_list_item",
        "monthly_intro",
        "monthly_observation",
        "balance_intro",
        "annual_list_intro",
        "specific_heading",
        "specific_intro",
        "specific_observation",
        "coke_centralizer",
        "monthly_chart",
        "monthly_caption",
        "specific_chart",
        "specific_caption",
        "mix_observation",
    }
)
CLONE_PREFIXES = (
    "electricity_grid",
    "electricity_pv",
    "natural_gas",
    "electricity_cogen",
    "diesel",
    "petrol",
    "lpg",
    "ctl",
    "coke",
    "sunflower_husks",
    "wood",
    "biomass",
    "purchased_heat",
    "water_potable",
    "water_industrial",
    "water_storm",
    "total_energy",
)


def _data(previous: Path | None) -> PieeData:
    return load(
        2025,
        case_path("piee-case-c", "anexa"),
        None,
        case_path("piee-case-c", "prelucrare"),
        previous,
    )


def _headings(path: Path) -> list[str]:
    return [
        normal(paragraph.text)
        for paragraph in Document(path).paragraphs
        if paragraph.style.name == "Heading 3"
    ]


def _assert_heading_order(path: Path, phrases: tuple[str, ...]) -> None:
    headings = _headings(path)
    positions = [
        next(index for index, heading in enumerate(headings) if phrase in heading)
        for phrase in phrases
    ]
    assert positions == sorted(set(positions))


def _bookmark_names(path: Path) -> set[str]:
    return {
        name
        for mark in Document(path).element.body.iter(qn("w:bookmarkStart"))
        if (name := mark.get(qn("w:name"), "")).startswith("_ema_")
    }


def _assert_clone_bookmarks(base: Path, composed: Path) -> None:
    approved = _bookmark_names(base / "piee-master.docx")
    names = _bookmark_names(composed)
    added = names - approved
    assert added
    assert not any(name.startswith("_ema_electricity_pv_") for name in added)
    for name in added:
        matches = [prefix for prefix in CLONE_PREFIXES if name.startswith(f"_ema_{prefix}_")]
        assert len(matches) == 1, name
        role = name.removeprefix(f"_ema_{matches[0]}_")
        role = re.sub(r"_\d+$", "", role)
        assert role in CLONE_ROLES, name


def test_delivered_layout_composes_case_c_figures_and_missing_markers(  # noqa: PLR0915
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    final = case_path("piee-case-c", "final")
    data = _data(final)
    base = artifacts_path("s8", "base")
    load_approved_base(base)  # The new prototype manifest did not change the approved map.
    prototypes = json.loads((base / "carrier-prototypes.json").read_text(encoding="utf-8"))
    assert set(prototypes) == {"boundaries", "gas", "water", "specific_gas", "gas_annual"}
    assert "annual_list_item" in prototypes["gas"] and "annual_list_item" in prototypes["water"]
    assert {item.key for item in data.disagreements} == {
        "co2.coke.2023",
        "co2.coke.2024",
        "co2.coke.2025",
    }  # S11
    assert data.prelucrare is not None
    assert all(
        data.prelucrare.located[f"turnover.{year}"].label is not None for year in data.dataset.years
    )
    output = tmp_path / "draft.docx"
    status = compose_draft(data, base, output, date(2026, 10, 1))
    assert status.final_ready
    assert not check_standalone(output)
    _compare_figures(output, final, data)
    assert _caption_kinds(output) == _caption_kinds(final)
    _compare_annual_tables(final, data)
    _compare_centralizer(output, final, data)
    _compare_numbered_tables(output, final, data)
    assert_section_sequence(output, final)
    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    production = next(iter(data.dataset.production))
    source_production = value(
        data.dataset, data.factors, Metric("production", product=production), 2023
    )[0]
    assert source_production is not None
    produced_amount = prototype_number(source_production, 2)
    assert f"pentru anul 2023 s-au înregistrat {produced_amount} tone/an" in text  # S1
    assert "Evaluția cantității totale" not in text  # S15
    assert (
        "Conform informațiilor prezentate, energia electrică consumată din cele două surse, "
        "achiziționată din SEN și produsă local prin cogenerare, este în cantitate totală de:"
    ) in text
    assert (
        "În această pondere nu este prezentată distinct cantitatea de energie electrică "
        "produsă prin cogenerare întrucât, pentru obținerea ei se utilizează gazul natural "
        "care este prezentat."
    ) in text
    assert "Fig. nr. 17 Evoluția cantității totale de energie echivalentă" in text
    for year in data.dataset.years:
        grid = data.dataset.carriers[Carrier.electricity_grid][year]
        cogen = data.dataset.carriers[Carrier.electricity_cogen][year]
        purchased = (
            grid.annual.value
            if grid.annual is not None
            else sum(reading.value for reading in grid.months.values() if reading.value is not None)
        )
        produced = (
            cogen.annual.value
            if cogen.annual is not None
            else sum(
                reading.value for reading in cogen.months.values() if reading.value is not None
            )
        )
        assert purchased is not None and produced is not None
        total = purchased + produced
        share = round(100 * produced / total)
        ending = "." if year == data.year else ";"
        assert (
            f"{prototype_number(total, 2)} MWh/an din care {share}% o reprezintă ponderea "
            f"energiei electrice produsă intern prin cogenerare în {year}{ending}"
        ) in text
    for carrier, tip in (
        (Carrier.water_potable, "potabilă"),
        (Carrier.water_industrial, "industrială"),
        (Carrier.water_storm, "meteorică"),
    ):
        assert f"Consumul de apă {tip} a fost:" in text
        for year in data.dataset.years:
            reading = data.dataset.carriers[carrier][year].annual
            assert reading is not None and reading.value is not None
            ending = "." if year == data.year else ","
            assert f"în anul {year} de {prototype_number(reading.value, 0)} m3/an{ending}" in text
    assert "tep/mii tone" in text and "m3/tone" in text
    company = str(data.anexa.identity["name"].value)
    assert f"Centralizator al consumurilor de energie la {company}" in text  # S10
    assert f"Impactul de mediu al modului de utilizare a energiei în cadrul {company}" in text
    assert "Adresa sediului social." in text  # S13 keeps the approved base heading.
    for field_key in MISSING_FIELD_KEYS.values():
        assert field_key.startswith("identity.")
        assert field_key.removeprefix("identity.") not in data.anexa.identity
        assert field_key not in data.prelucrare.located
    assert_missing_identity_is_unsourced(data)
    assert not any(
        "specific echivalent de energie electrica" in heading
        or "specific echivalent de carbur" in heading
        for heading in _headings(output)
    )  # S9 follows the delivered layout profile.
    parts = read_parts(output)
    assert (
        sum(bool(list(xml(parts, part).iter(f"{{{C}}}pie3DChart"))) for part in _order(output)) == 3
    )

    # The bookmarked copy identifies each red gap without retaining client material.
    monkeypatch.setattr("ema.piee.compose._strip_bookmarks", shutil.copyfile)
    bookmarked = tmp_path / "bookmarked.docx"
    compose_draft(data, artifacts_path("s8", "base"), bookmarked, date(2026, 10, 1))
    _assert_clone_bookmarks(base, bookmarked)
    assert_production_name_is_sourced(data, bookmarked)
    mapping = load_approved_base(base)
    narrative = {entry.slot: entry for entry in mapping.elements if entry.slot in NARRATIVE_SLOTS}
    assert set(narrative) == set(NARRATIVE_SLOTS)
    assert all(
        item.kind == "paragraph" and item.classification == "variable"
        for item in narrative.values()
    )
    slots: set[str] = set()
    absent_identity = 0
    for paragraph in Document(bookmarked).paragraphs:
        if "n.d." not in paragraph.text:
            continue
        assert all(
            str(run.font.color.rgb) == "FF0000"
            for run in paragraph.runs
            if "n.d." in run.text and run.font.color is not None
        )
        names = {
            name.removeprefix("_ema_")
            for mark in paragraph._p.iter(qn("w:bookmarkStart"))
            if (name := mark.get(qn("w:name"), "")).startswith("_ema_")
            and not name.startswith("_ema_prototype_")
        }
        if names:
            slots.update(names)
        else:
            absent_identity += paragraph.text.count("n.d.")
    assert slots == set(NARRATIVE_SLOTS) | {"body_11", "body_14"}
    assert set(MISSING_FIELD_KEYS) == {
        "body_11",
        "body_14",
        "site_2_address",
        "site_1_share",
        "site_2_share",
    }
    assert absent_identity == 3
    assert (
        sum("n.d." in p.text for p in Document(bookmarked).paragraphs if "Sucursala" in p.text) == 1
    )
    assert (
        next(
            p.text.count("n.d.")
            for p in Document(bookmarked).paragraphs
            if p.text.startswith("Din totalul producției pe anul 2025,")
        )
        == 2
    )


def test_base_layout_composes_all_sourced_case_c_sections(tmp_path: Path) -> None:
    data = _data(None)
    output = tmp_path / "draft.docx"
    status = compose_draft(data, artifacts_path("s8", "base"), output, date(2026, 10, 1))
    assert status.final_ready
    assert not check_standalone(output)
    # 24 production/carrier + 12 water + 3 mix + 8 specific + intensity + CO2 = 49.
    assert len(_order(output)) == 49
    _assert_heading_order(
        output,
        (
            "energie electrica",
            "gaz natural",
            "cogenerare",
            "carburanti",
            "cocs",
            "apa potabila",
            "apa industriala",
            "apa meteorica",
            "echivalent de energie",
            "specific echivalent de energie electrica",
            "specific echivalent de gaz",
            "specific echivalent de carbur",
            "specific de cocs",
            "specific echivalent de energie totala",
            "specific echivalent de apa",
            "specific de apa industriala",
            "specific de apa meteorica",
            "intensitatea energetica",
        ),
    )
    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "Analiza consumului de cocs" in text
    assert "Analiza consumului de apă industrială" in text
    assert "Analiza consumului de apă meteorică" in text
    assert "energie electrică produsă prin cogenerare" in text
    assert not any("fotovolta" in heading for heading in _headings(output))
    assert "fotovolta" not in normal(text)  # PV section, caption, table, and prose are gone.
    parts = read_parts(output)
    body = xml(parts, "word/document.xml").find(qn("w:body"))
    assert body is not None
    toc = list(body)[1]
    assert "fotovolta" not in normal("".join(node.text or "" for node in toc.iter(qn("w:t"))))
    for part in _order(output):
        for series in read_series(output, part):
            assert "fotovolta" not in normal(series.name)
            assert all("fotovolta" not in normal(category) for category in series.categories)
