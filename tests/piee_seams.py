"""Synthetic stand-ins for the PIEE document seams; start_generate, core.jobs and outputs stay real.

The document internals are covered by the S8/S16/S19 goldens on the reference library.
"""

from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook

from ema.core.review.fields import propose
from ema.core.review.models import Evidence, Manual
from ema.core.workspace import Workspace


def _data(*_args: object) -> SimpleNamespace:
    return SimpleNamespace(
        factors=SimpleNamespace(version="synthetic-factors"),
        dataset=SimpleNamespace(years=(2025,)),
        production_conversion=None,
        pie_representation="synthetic",
        pie_representation_source="base",
        separate_pv_figures=False,
    )


def _compose(_data: object, _base: Path, output: Path, _today: object) -> SimpleNamespace:
    document = Document()
    document.add_paragraph("Program de îmbunătățire a eficienței energetice (sintetic)")
    document.save(str(output))
    return SimpleNamespace(untouched=[], package_issues=[], leftover_parts=[], final_ready=True)


def _workbook(_dataset: object, _years: object, path: Path, _factors: object) -> None:
    Workbook().save(str(path))


def _import(ws: Workspace, job: str, *_args: object, **_kwargs: object) -> None:
    evidence = Evidence(
        id="synthetic:identity.name:" + job,
        provenance="manual",
        locator=Manual(who="synthetic"),
        method="manual",
        retrieved_at=datetime.now(UTC),
        highlight="exact",
    )
    propose(ws, job, "identity.name", "Atelier Exemplu SRL", [evidence], state="supplied")


def stub_piee_seams(monkeypatch: pytest.MonkeyPatch, base: Path) -> None:
    """Replace the six document seams of ema.piee.workflow; base is any existing folder."""
    monkeypatch.setattr("ema.piee.workflow.load", _data)
    monkeypatch.setattr("ema.piee.workflow.base_directory", lambda _ws: base)
    monkeypatch.setattr(
        "ema.piee.workflow.load_approved_base", lambda _base: SimpleNamespace(base_sha="0" * 64)
    )
    monkeypatch.setattr("ema.piee.workflow.import_piee_into_job", _import)
    monkeypatch.setattr("ema.piee.workflow.compose_draft", _compose)
    monkeypatch.setattr("ema.piee.workflow.write_prelucrare", _workbook)


def synthetic_inputs(folder: Path) -> tuple[Path, Path, Path]:
    """Blank anexa, necesar and prelucrare workbooks; the stubbed load never parses them."""
    folder.mkdir(parents=True, exist_ok=True)
    paths = (folder / "Anexa.xlsx", folder / "Necesar.xlsx", folder / "Prelucrare.xlsx")
    for path in paths:
        Workbook().save(str(path))
    return paths
